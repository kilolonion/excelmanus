"""analyze：join / 日期分组回归 + pivot 二维 + conditions=[] / operator=not。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from excelmanus.security import FileAccessGuard
from excelmanus.tools import workbook_tools, reference_tools
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.workbook_tools import analyze_spreadsheet, observe_spreadsheet


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    workbook_tools.init_guard(workspace)
    reference_tools.init_guard(workspace)


def _save(path: Path, sheets: dict[str, list[list[object]]]) -> Path:
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
    wb.save(path)
    wb.close()
    return path


def test_aggregate_join_and_year_month(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _save(
        tmp_path / "orders.xlsx",
        {
            "订单": [
                ["产品ID", "数量", "下单日期"],
                ["P1", 2, "2024-01-15"],
                ["P1", 3, "2024-01-20"],
                ["P2", 1, "2024-02-01"],
            ],
            "产品": [["产品ID", "名称"], ["P1", "苹果"], ["P2", "香蕉"]],
        },
    )
    joined = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        sheet="订单",
        join={"sheet": "产品", "on": "产品ID", "columns": ["名称"]},
        group_by="名称",
        aggregations={"数量": "sum"},
    )
    assert joined.success, joined.model_text
    groups = {row["名称"]: row["数量_sum"] for row in joined.value["groups"]}
    assert groups["苹果"] == 5
    assert groups["香蕉"] == 1

    by_month = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        sheet="订单",
        group_by={"column": "下单日期", "transform": "year_month"},
        aggregations={"数量": "sum"},
    )
    assert by_month.success, by_month.model_text
    months = {row["下单日期__year_month"]: row["数量_sum"] for row in by_month.value["groups"]}
    assert months["2024-01"] == 5
    assert months["2024-02"] == 1


def test_aggregate_sort_by_source_column_maps_to_derived_key(tmp_path: Path) -> None:
    """sort_by 传源列名时映射到 派生键输出列（源列__transform）。"""
    _bind(tmp_path)
    path = _save(
        tmp_path / "orders.xlsx",
        {
            "订单": [
                ["产品ID", "数量", "下单日期"],
                ["P1", 2, "2024-02-15"],
                ["P1", 3, "2024-01-20"],
                ["P2", 1, "2024-03-01"],
            ],
        },
    )
    result = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        sheet="订单",
        group_by={"column": "下单日期", "transform": "year_month"},
        aggregations={"数量": "sum"},
        sort_by="下单日期",
    )
    assert result.success, result.model_text
    keys = [row["下单日期__year_month"] for row in result.value["groups"]]
    assert keys == sorted(keys)


def test_analyze_pivot_matrix(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _save(
        tmp_path / "pivot.xlsx",
        {
            "Sheet1": [
                ["部门", "月", "金额"],
                ["销售", 1, 10],
                ["销售", 2, 20],
                ["研发", 1, 5],
            ]
        },
    )
    result = analyze_spreadsheet(
        mode="pivot",
        file_path=str(path),
        index="部门",
        columns="月",
        values="金额",
        aggfunc="sum",
    )
    assert result.success, result.model_text
    assert result.value["mode"] == "pivot"
    assert result.value["returned"] == result.value["total"]
    matrix = result.value["matrix"]
    assert matrix[0][0] == "部门"
    header = matrix[0]
    body = {row[0]: row for row in matrix[1:]}
    jan = header.index("1")
    feb = header.index("2")
    assert body["销售"][jan] == 10
    assert body["销售"][feb] == 20
    assert body["研发"][jan] == 5


def test_filter_empty_conditions_and_not_operator(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _save(
        tmp_path / "filter.xlsx",
        {"Sheet1": [["部门", "金额"], ["销售", 10], ["研发", 20]]},
    )
    all_rows = analyze_spreadsheet(
        mode="filter",
        file_path=str(path),
        conditions=[],
    )
    assert all_rows.success, all_rows.model_text
    assert all_rows.value["filtered_rows"] == 2
    assert all_rows.value["returned_rows"] == 2

    not_sales = analyze_spreadsheet(
        mode="filter",
        file_path=str(path),
        column="部门",
        operator="not",
        value="销售",
    )
    assert not_sales.success, not_sales.model_text
    depts = [row["部门"] for row in not_sales.value["data"]]
    assert depts == ["研发"]


def test_inspect_range_max_rows_warns(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _save(
        tmp_path / "range.xlsx",
        {"Sheet1": [["部门", "金额"], ["销售", 10], ["研发", 20]]},
    )
    result = observe_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet="Sheet1",
        range="A1:B3",
    )
    assert result.success, result.model_text
    assert result.value["regions"][0]["rect"] == {"r0":1,"c0":1,"r1":3,"c1":2}
import json


def test_structured_arguments_compose(tmp_path: Path) -> None:
    """V2 structured arguments compose across join, filter and pivot."""
    _bind(tmp_path)
    path = _save(
        tmp_path / "json_args.xlsx",
        {
            "订单": [
                ["产品ID", "数量", "下单日期"],
                ["P1", 2, "2024-01-15"],
                ["P2", 1, "2024-02-01"],
            ],
            "产品": [["产品ID", "名称"], ["P1", "苹果"], ["P2", "香蕉"]],
        },
    )
    joined = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        sheet="订单",
        join={"sheet": "产品", "on": "产品ID", "columns": ["名称"]},
        group_by="名称",
        aggregations={"数量": "sum"},
    )
    assert joined.success, joined.model_text
    groups = {row["名称"]: row["数量_sum"] for row in joined.value["groups"]}
    assert groups["苹果"] == 2

    filtered = analyze_spreadsheet(
        mode="filter",
        file_path=str(path),
        sheet="订单",
        conditions=[{"column": "产品ID", "operator": "==", "value": "P1"}],
    )
    assert filtered.success, filtered.model_text
    assert filtered.value["filtered_rows"] == 1

    pivoted = analyze_spreadsheet(
        mode="pivot",
        file_path=str(path),
        sheet="订单",
        index="产品ID",
        columns=json.dumps({"column": "下单日期", "transform": "year_month"}),
        values=json.dumps(["数量"]),
        aggfunc="sum",
    )
    assert pivoted.success, pivoted.model_text
    header = pivoted.value["matrix"][0]
    assert "2024-01" in header and "2024-02" in header


def test_parse_bound_selection_accepts_json_string(tmp_path: Path) -> None:
    from excelmanus.workbook.snapshot import parse_bound_selection

    raw = json.dumps({
        "file": "a.xlsx",
        "sheet": "Sheet1",
        "rows": [2, 3],
        "content_version": "sha256:abc",
        "snapshot_id": "ws|id|v",
    })
    sel = parse_bound_selection(raw)
    assert sel is not None
    assert list(sel.rows) == [2, 3]
    assert parse_bound_selection("not json") is None
    assert parse_bound_selection("[1,2]") is None


def test_filter_logic_not_single_condition(tmp_path: Path) -> None:
    """logic='not' 对单条件整体取反；多条件时明确拒绝（德摩根歧义）。"""
    _bind(tmp_path)
    path = _save(
        tmp_path / "staff.xlsx",
        {
            "员工": [
                ["姓名", "部门"],
                ["张三", "销售"],
                ["李四", "技术"],
                ["王五", "销售"],
            ],
        },
    )
    res = analyze_spreadsheet(
        mode="filter",
        file_path=str(path),
        sheet="员工",
        column="部门",
        operator="eq",
        value="销售",
        logic="not",
    )
    assert res.success, res.model_text
    assert res.value["filtered_rows"] == 1
    assert res.value["records"][0]["姓名"] == "李四"

    # 多条件 + not：德摩根歧义，必须明确拒绝并给出指引
    bad = analyze_spreadsheet(
        mode="filter",
        file_path=str(path),
        sheet="员工",
        conditions=[
            {"column": "部门", "operator": "eq", "value": "销售"},
            {"column": "姓名", "operator": "contains", "value": "三"},
        ],
        logic="not",
    )
    assert not bad.success
    assert "not" in bad.model_text and "单个条件" in bad.model_text


def test_analyze_pivot_missing_index_required_for_mode(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _save(
        tmp_path / "pivot_miss.xlsx",
        {"Sheet1": [["部门", "月", "金额"], ["销售", 1, 10]]},
    )
    result = analyze_spreadsheet(
        mode="pivot",
        file_path=str(path),
        columns="月",
        values="金额",
    )
    assert not result.success
    payload = result.value if isinstance(result.value, dict) else {}
    required = payload.get("required_for_mode") or []
    assert "index" in required
    assert "columns" in required
    assert "values" in required
    assert "index" in (result.model_text or "")


def test_analyze_overview_points_to_inspect(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _save(
        tmp_path / "ov.xlsx",
        {"Sheet1": [["部门"], ["销售"]]},
    )
    result = analyze_spreadsheet(mode="overview", file_path=str(path))
    assert not result.success
    text = result.model_text or ""
    assert "overview" in text
    assert "observe_spreadsheet" in text


# ── 回归：aggregations/group_by 等结构化参数传 JSON 字符串 ──────────────────
# 真实 agent 会话中 analyze_spreadsheet(mode="aggregate") 连续失败 4 次，
# agent 传的 4 种 aggregations 形状全部是 JSON 字符串，被 schema 以
# "is not of type 'object','array'" 整段 jsonschema 原文拒绝。
# 统一合同：结构化对象/数组与等价 JSON 字符串等价（运行时自动解析）；
# 无法解析时返回含正确示例的中文错误。


_INCIDENT_HEADER = ["区域", "年份", "销售额(万元)"]
_INCIDENT_ROWS = [
    ["华东", 2023, 10],
    ["华东", 2023, 20],
    ["华北", 2023, 5],
    ["华北", 2024, 8],
    ["华东", 2024, 30],
]


def _incident_csv(tmp_path: Path) -> Path:
    """列名含中文与括号的 CSV（复现事故数据形状）。"""
    path = tmp_path / "销售数据.csv"
    lines = [",".join(_INCIDENT_HEADER)]
    lines += [",".join(str(value) for value in row) for row in _INCIDENT_ROWS]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def test_incident_case1_object_with_func_array_json_string(tmp_path: Path) -> None:
    """确切调用 1：aggregations='{"销售额(万元)": ["sum","mean","max","min"]}'。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    result = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        aggregations='{"销售额(万元)": ["sum", "mean", "max", "min"]}',
        group_by='["区域", "年份"]',
    )
    assert result.success, result.model_text
    assert result.value["group_count"] == 4
    east_2023 = next(
        row for row in result.value["groups"]
        if row["区域"] == "华东" and str(row["年份"]) == "2023"
    )
    assert east_2023["销售额(万元)_sum"] == 30
    assert east_2023["销售额(万元)_mean"] == 15
    assert east_2023["销售额(万元)_max"] == 20
    assert east_2023["销售额(万元)_min"] == 10


def test_incident_case2_object_with_single_func_json_string(tmp_path: Path) -> None:
    """确切调用 2：aggregations='{"销售额(万元)": "sum"}'。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    result = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        aggregations='{"销售额(万元)": "sum"}',
        group_by='["区域"]',
    )
    assert result.success, result.model_text
    sums = {row["区域"]: row["销售额(万元)_sum"] for row in result.value["groups"]}
    assert sums == {"华东": 60, "华北": 13}


def test_incident_case3_column_func_array_json_string(tmp_path: Path) -> None:
    """确切调用 3：aggregations='[{"column":"销售额(万元)","func":"sum"}, …]'。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    result = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        aggregations=(
            '[{"column": "销售额(万元)", "func": "sum"},'
            ' {"column": "销售额(万元)", "func": "mean"},'
            ' {"column": "销售额(万元)", "func": "max"},'
            ' {"column": "销售额(万元)", "func": "min"}]'
        ),
        group_by='["区域"]',
    )
    assert result.success, result.model_text
    east = next(row for row in result.value["groups"] if row["区域"] == "华东")
    assert east["销售额(万元)_sum"] == 60
    assert east["销售额(万元)_max"] == 30
    assert east["销售额(万元)_min"] == 10


def test_incident_case4_single_column_func_json_string(tmp_path: Path) -> None:
    """确切调用 4：aggregations='[{"column":"销售额(万元)","func":"sum"}]'。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    result = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        aggregations='[{"column": "销售额(万元)", "func": "sum"}]',
        group_by='["区域"]',
    )
    assert result.success, result.model_text
    sums = {row["区域"]: row["销售额(万元)_sum"] for row in result.value["groups"]}
    assert sums == {"华东": 60, "华北": 13}


def test_json_string_forms_equivalent_to_structured_forms(tmp_path: Path) -> None:
    """{列: 函数数组} 对象、[{column, func}] 数组与两种 JSON 字符串产出一致。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    forms = [
        {"销售额(万元)": ["sum", "mean"]},
        [{"column": "销售额(万元)", "func": "sum"}, {"column": "销售额(万元)", "func": "mean"}],
        '{"销售额(万元)": ["sum", "mean"]}',
        '[{"column": "销售额(万元)", "func": "sum"}, {"column": "销售额(万元)", "func": "mean"}]',
    ]
    baselines = []
    for aggregations in forms:
        result = analyze_spreadsheet(
            mode="aggregate", file_path=str(path),
            aggregations=aggregations, group_by='["区域"]',
        )
        assert result.success, f"aggregations={aggregations!r}: {result.model_text}"
        baselines.append(result.value["groups"])
    assert all(groups == baselines[0] for groups in baselines)


def test_group_by_multi_column_json_string_and_max_rows(tmp_path: Path) -> None:
    """group_by 多列 JSON 字符串生效，且 max_rows 截断组数并标注 truncated。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    result = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        aggregations='{"销售额(万元)": "sum"}',
        group_by='["区域", "年份"]',
        max_rows=2,
    )
    assert result.success, result.model_text
    assert result.value["group_count"] == 4
    assert result.value["returned_groups"] == 2
    assert len(result.value["groups"]) == 2
    assert result.value.get("truncated") is True


def test_conditions_and_join_json_strings_work(tmp_path: Path) -> None:
    """conditions / join 与同级字段同策略：JSON 字符串自动解析后生效。"""
    _bind(tmp_path)
    csv_path = _incident_csv(tmp_path)
    filtered = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(csv_path),
        aggregations='{"销售额(万元)": "sum"}',
        group_by='["区域"]',
        conditions='[{"column": "年份", "operator": "eq", "value": 2023}]',
    )
    assert filtered.success, filtered.model_text
    sums = {row["区域"]: row["销售额(万元)_sum"] for row in filtered.value["groups"]}
    assert sums == {"华东": 30, "华北": 5}

    _bind(tmp_path)
    xlsx_path = _save(
        tmp_path / "orders_json.xlsx",
        {
            "订单": [
                ["产品ID", "数量", "下单日期"],
                ["P1", 2, "2024-01-15"],
                ["P1", 3, "2024-01-20"],
                ["P2", 1, "2024-02-01"],
            ],
            "产品": [["产品ID", "名称"], ["P1", "苹果"], ["P2", "香蕉"]],
        },
    )
    joined = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(xlsx_path),
        sheet="订单",
        join='{"sheet": "产品", "on": "产品ID", "columns": ["名称"]}',
        group_by='["名称"]',
        aggregations='{"数量": "sum"}',
    )
    assert joined.success, joined.model_text
    groups = {row["名称"]: row["数量_sum"] for row in joined.value["groups"]}
    assert groups == {"苹果": 5, "香蕉": 1}


def test_pivot_index_columns_values_json_strings(tmp_path: Path) -> None:
    """pivot 的 index/columns/values 与 aggregate 同策略。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    result = analyze_spreadsheet(
        mode="pivot",
        file_path=str(path),
        index='["区域"]',
        columns='["年份"]',
        values='["销售额(万元)"]',
    )
    assert result.success, result.model_text
    assert result.value["mode"] == "pivot"


def test_unparseable_json_string_gives_actionable_chinese_error(tmp_path: Path) -> None:
    """坏 JSON 不再倾倒 jsonschema：中文说明期望形状并给正确示例。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    result = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        aggregations='{"销售额(万元)": ',
        group_by='["区域"]',
    )
    assert not result.success
    text = result.model_text or ""
    assert "aggregations" in text
    assert "示例：" in text
    assert "is not of type" not in text
    assert "Failed validating" not in text


def test_bad_shape_aggregations_error_mentions_expected_shape(tmp_path: Path) -> None:
    """形状不对（如裸字符串）同样给出期望形状 + 正确示例，而非 jsonschema 原文。"""
    _bind(tmp_path)
    path = _incident_csv(tmp_path)
    result = analyze_spreadsheet(
        mode="aggregate",
        file_path=str(path),
        aggregations="sum",
        group_by='["区域"]',
    )
    assert not result.success
    text = result.model_text or ""
    assert "aggregations" in text
    assert "示例：" in text
    assert "is not of type" not in text
    assert "Failed validating" not in text
