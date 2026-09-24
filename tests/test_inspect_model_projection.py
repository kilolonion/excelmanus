"""E01：inspect/analyze 正文必须带上行列、隐藏表、小窗口 values 与未缓存公式。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.workbook_tools import analyze_spreadsheet, init_guard, observe_spreadsheet
from excelmanus.workbook.data import init_guard as init_data_guard
from excelmanus.tools.workbook_tools import init_guard as init_sheets_guard
from excelmanus.workbook_commit import seed_seen_versions


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    init_guard(workspace)
    init_data_guard(workspace)
    init_sheets_guard(workspace)
    seed_seen_versions({})


def _messy_book(path: Path) -> Path:
    wb = Workbook()
    detail = wb.active
    assert detail is not None
    detail.title = "明细"
    headers = ["区域", "数量", "单价", "金额", "备注", "日期", "销售", "渠道"]
    detail.append(headers)
    detail["A2"] = "华东"
    detail["B2"] = 2
    detail["C2"] = 3
    detail["D2"] = "=B2*C2"
    for row in range(3, 15):
        detail[f"A{row}"] = "西  部" if row % 2 else "east"
        detail[f"B{row}"] = row
        detail[f"C{row}"] = 10
        detail[f"D{row}"] = row * 10
    notes = wb.create_sheet("随便写写")
    notes["A1"] = "口头备忘"
    hidden = wb.create_sheet("隐藏底稿")
    hidden["A1"] = "华东"
    hidden["B1"] = 860000
    hidden.sheet_state = "hidden"
    wb.save(path)
    wb.close()
    return path


def test_overview_model_projection_keeps_shapes_hidden_styles_and_formulas(tmp_path: Path) -> None:
    import json
    from tests.workbook_support import model_projection, projected_payload
    _bind(tmp_path)
    path = _messy_book(tmp_path / "q3.xlsx")
    result = observe_spreadsheet(mode="overview", file_path=str(path), facets=["data","presentation"])
    assert result.success
    payload = projected_payload(result,tmp_path)
    assert len(payload["sheets"]) == 3
    assert payload["sheets"][0]["used"] == {"rows":14,"cols":8}
    assert payload["sheets"][2]["state"] == "hidden"
    assert payload["regions"][0]["cells"]["2,4"]["f"] == "=B2*C2"
    assert "s" in payload["regions"][0]["cells"]["1,1"]


def test_overview_default_includes_layout_dimensions(tmp_path: Path) -> None:
    _bind(tmp_path)
    path=tmp_path/"layout.xlsx"
    wb=Workbook(); ws=wb.active; ws["A1"]="标题"; ws.column_dimensions["A"].width=31; ws.row_dimensions[1].height=18; wb.save(path); wb.close()
    result=observe_spreadsheet(file_path=str(path))
    geometry=result.value["regions"][0]["geometry"]
    assert geometry["columns"][0]["native"]==31
    assert geometry["rows"][0]["native"]==18
    assert geometry["defaults"]["column_width"]["native"] is None
    assert geometry["defaults"]["column_width"]["source"]=="estimated"


def test_overview_lists_multiple_formula_samples(tmp_path: Path) -> None:
    from tests.workbook_support import model_projection, projected_payload
    _bind(tmp_path)
    path=tmp_path/"kpi.xlsx"
    wb=Workbook(); ws=wb.active
    for row in range(2,6): ws.cell(row,2,f"=SUM(A{row}:A{row+1})")
    wb.save(path); wb.close()
    result=observe_spreadsheet(file_path=str(path),facets=["data"])
    text=model_projection(result,tmp_path)
    for row in range(2,6): assert f"=SUM(A{row}:A{row+1})" in text


def test_small_range_projects_full_window_and_uncached_formula(tmp_path: Path) -> None:
    import json
    from tests.workbook_support import model_projection, projected_payload
    _bind(tmp_path)
    path=_messy_book(tmp_path/"q3.xlsx")
    result=observe_spreadsheet(file_path=str(path),sheet="明细",mode="range",range="A1:H14")
    payload=projected_payload(result,tmp_path)
    cells=payload["regions"][0]["cells"]
    assert cells["2,4"]["f"]=="=B2*C2"
    assert cells["2,4"]["v"] is None and cells["2,4"]["cached"]=="no"
    assert cells["2,1"]["v"]=="华东"
    assert cells["3,1"]["v"]=="西  部"
    assert payload["coverage"]["loaded"][0]["r1"]==14


def test_quality_model_text_scopes_sheet_and_names_columns(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _messy_book(tmp_path / "q3.xlsx")
    result = analyze_spreadsheet(
        mode="quality",
        file_path=str(path),
        sheet="明细",
    )
    assert result.success
    text = result.model_text
    assert "本次只扫「明细」" in text
    assert "missing_data(明细." in text or "empty_column(明细." in text
    assert '"quality_signals"' not in text
