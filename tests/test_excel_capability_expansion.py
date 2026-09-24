from io import BytesIO
from zipfile import ZipFile

import pytest
import xlsxwriter
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

from excelmanus.tools.context import use_workspace
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes, apply_spreadsheet_changes, observe_spreadsheet, analyze_spreadsheet
from excelmanus.workbook_commit import content_version_of_file
from excelmanus.tools.spreadsheet_data_tools import validate_spreadsheet, query_spreadsheet
from excelmanus.tools.spreadsheet_engine_tools import calculate_spreadsheet, convert_spreadsheet, render_spreadsheet


@pytest.fixture(autouse=True)
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("EXCELMANUS_FORMULA_RECALC", "off")
    with use_workspace(tmp_path):
        yield


def book(path, rows, name="Data"):
    wb=Workbook(); ws=wb.active; ws.title=name
    for row in rows: ws.append(row)
    wb.save(path); wb.close()


def edit(path, operations):
    result=apply_spreadsheet_changes(file_path=path.name,expected_version=content_version_of_file(path),operations=operations)
    assert result.success, result.model_text
    return result


def objects(path, operations):
    result=apply_spreadsheet_changes(file_path=path.name,expected_version=content_version_of_file(path),operations=operations)
    assert result.success, result.model_text
    return result


def test_insert_rewrites_formula_table_name_rules_and_chart(tmp_path):
    path=tmp_path/"book.xlsx"; wb=Workbook(); ws=wb.active; ws.title="Data"
    for row in [["id","amount"],[1,10],[2,20]]: ws.append(row)
    ws.add_table(Table(displayName="Sales",ref="A1:B3"))
    wb.defined_names.add(DefinedName("Amounts",attr_text="Data!$B$2:$B$3"))
    other=wb.create_sheet("Summary"); other["A1"]="=SUM(Data!$B$2:$B$3)"
    chart=BarChart(); chart.add_data(Reference(ws,min_col=2,min_row=1,max_row=3),titles_from_data=True); other.add_chart(chart,"D1")
    rule=DataValidation(type="whole",formula1="1",formula2="100"); rule.add("B2:B3"); ws.add_data_validation(rule)
    ws.conditional_formatting.add("B2:B3",CellIsRule(operator="greaterThan",formula=["15"]))
    ws.print_area="A1:B3"; wb.save(path); wb.close()
    edit(path,[{"kind":"insert","sheet":"Data","axis":"row","at":2,"count":1}])
    wb=load_workbook(path)
    assert wb["Data"].tables["Sales"].ref=="A1:B4"
    assert wb.defined_names["Amounts"].attr_text=="Data!$B$3:$B$4"
    assert wb["Summary"]["A1"].value=="=SUM(Data!$B$3:$B$4)"
    assert wb["Summary"]._charts[0].series[0].val.numRef.f=="'Data'!$B$3:$B$4"
    assert str(wb["Data"].data_validations.dataValidation[0].sqref)=="B3:B4"
    assert str(next(iter(wb["Data"].conditional_formatting)).sqref)=="B3:B4"
    assert "$A$1:$B$4" in wb["Data"].print_area
    wb.close()


def test_rename_ignores_literal_and_delete_shrinks_ref(tmp_path):
    path=tmp_path/"book.xlsx"; book(path,[[1],[2],[3]])
    wb=load_workbook(path); other=wb.create_sheet("Summary"); other["A1"]='=SUM(Data!$A$1:$A$3)'; other["B1"]='="Data!A1"'; wb.save(path); wb.close()
    edit(path,[{"kind":"sheet","sheet":"Data","action":"rename","new_name":"New Data"},{"kind":"delete_rows","sheet":"New Data","at":2,"count":1}])
    wb=load_workbook(path)
    assert wb["Summary"]["A1"].value=="=SUM('New Data'!$A$1:$A$2)"
    assert wb["Summary"]["B1"].value=='="Data!A1"'
    wb.close()


def test_mixed_charts_and_property_patch(tmp_path):
    path=tmp_path/"book.xlsx"; book(path,[["amount"],[10],[20]])
    objects(path,[{"kind":"chart","sheet":"Data","chart_type":"bar","data_range":"A1:A3","target_cell":"D1","title":"Original"},{"kind":"chart","sheet":"Data","chart_type":"line","data_range":"A1:A3","target_cell":"D16"}])
    objects(path,[{"kind":"chart","sheet":"Data","chart_type":"line","data_range":"A1:A3","target_cell":"D1"},{"kind":"delete_chart","sheet":"Data","index":0}])
    wb=load_workbook(path); assert len(wb.active._charts)==2; wb.active._charts[0].legend.position="b"; wb.save(path); wb.close()
    objects(path,[{"kind":"update_chart","sheet":"Data","index":0,"title":"Changed"}])
    wb=load_workbook(path); assert len(wb.active._charts)==2; assert wb.active._charts[0].legend.position=="b"; wb.close()


def test_formula_pivot_uses_cache_and_rejects_stale_cache(tmp_path):
    path=tmp_path/"book.xlsx"
    with xlsxwriter.Workbook(path) as wb:
        ws=wb.add_worksheet("Data"); ws.write_row(0,0,["group","period","amount"])
        ws.write_row(1,0,["A","Q1"]); ws.write_formula(1,2,"=10*2",None,20)
        ws.write_row(2,0,["A","Q1"]); ws.write_formula(2,2,"=5*3",None,15)
    pivot={"kind":"pivot","sheet":"Data","target_sheet":"Summary","index":["group"],"columns":["period"],"values":["amount"]}
    before=path.read_bytes()
    result=apply_spreadsheet_changes(file_path=path.name,expected_version=content_version_of_file(path),operations=[{"kind":"write","sheet":"Data","start_cell":"A2","values":[["B"]]},pivot])
    assert not result.success and result.value["error_code"]=="FORMULA_CACHE_STALE"
    assert path.read_bytes()==before
    edit(path,[pivot]); wb=load_workbook(path); assert wb["Summary"]["B2"].value==35; wb.close()


def test_units_and_percent_conversion(tmp_path):
    path=tmp_path/"book.xlsx"; book(path,[["group","amount"],["A","1万"],["A","2000元"]])
    result=analyze_spreadsheet(file_path=path.name,sheet="Data",mode="aggregate",group_by="group",aggregations={"amount":"sum"})
    assert result.success and result.value["groups"][0]["amount_sum"]==12000
    from excelmanus.workbook.numeric import parse_number
    assert parse_number("10%") == .1
    assert parse_number("1.2亿元")==120000000


def test_range_operations_preserve_formula_style_and_literal_values(tmp_path):
    path=tmp_path/"book.xlsx"; book(path,[["key","amount"],[2,20],[1,10]])
    edit(path,[{"kind":"sort","sheet":"Data","range":"A1:B3","by":["key"]},{"kind":"fill","sheet":"Data","range":"C2:C3","mode":"series","start":5,"step":2},{"kind":"replace","sheet":"Data","range":"A1","find":"key","replacement":"id"},{"kind":"paste_special","sheet":"Data","source_range":"A1:C3","target_start":"E1","mode":"values"},{"kind":"clear","sheet":"Data","range":"C2:C3"}])
    wb=load_workbook(path); ws=wb.active
    assert [ws["A2"].value,ws["A3"].value]==[1,2]
    assert [ws["G2"].value,ws["G3"].value]==[5,7]
    assert ws["C2"].value is None; assert ws["E1"].value=="id"; wb.close()


def test_table_name_comment_hyperlink_append_and_inventory(tmp_path):
    path=tmp_path/"book.xlsx"; book(path,[["id","amount"],[1,10]])
    objects(path,[{"kind":"table","sheet":"Data","name":"Sales","ref":"A1:B2"},{"kind":"defined_name","sheet":"Data","name":"Amounts","refers_to":"Data!$B$2:$B$2"},{"kind":"comment","sheet":"Data","cell":"B1","text":"Amount in CNY"},{"kind":"hyperlink","sheet":"Data","cell":"D1","target":"https://example.com","display":"Source"}])
    edit(path,[{"kind":"append","sheet":"Data","table":"Sales","values":[[2,20]]}])
    wb=load_workbook(path); assert wb.active.tables["Sales"].ref=="A1:B3"; assert wb.active["B1"].comment.text=="Amount in CNY"; wb.close()
    result=observe_spreadsheet(file_path=path.name,mode="objects")
    assert result.success and {x["kind"] for x in result.value["regions"][0]["objects"]}>={"table","defined_name","comment","hyperlink"}


def test_native_pivot_serializes_cache_and_can_refresh(tmp_path):
    path=tmp_path/"book.xlsx"; book(path,[["group","period","amount"],["A","Q1",20],["A","Q1",15]])
    edit(path,[{"kind":"sheet","action":"create","new_name":"Summary"}])
    op={"kind":"pivot_table","sheet":"Summary","name":"SalesPivot","source_sheet":"Data","source_range":"A1:C3","rows":["group"],"columns":["period"],"values":[{"field":"amount","aggregation":"sum"}]}
    objects(path,[op])
    with ZipFile(path) as z:
        assert "xl/pivotTables/pivotTable1.xml" in z.namelist()
        assert "xl/pivotCache/pivotCacheRecords1.xml" in z.namelist()
    wb=load_workbook(path); assert wb["Summary"]["B2"].value==35; assert wb["Summary"]._pivots[0].cache.recordCount==2; wb.close()
    edit(path,[{"kind":"write","sheet":"Data","start_cell":"C2","values":[[30]]}])
    objects(path,[{"kind":"pivot_table","sheet":"Summary","name":"SalesPivot","action":"refresh"}])
    wb=load_workbook(path); assert wb["Summary"]["B2"].value==45; wb.close()


def test_validation_and_disk_query_export(tmp_path):
    path=tmp_path/"book.xlsx"; book(path,[["id","quantity","price","amount"],[1,2,10,20],[1,3,5,14]])
    check=validate_spreadsheet(path.name,[{"kind":"unique","columns":["id"]},{"kind":"row_expression","expression":"amount == quantity * price"},{"kind":"total","column":"amount","expected":35}])
    assert check.success and not check.value["valid"] and check.value["failure_count"]==3
    result=query_spreadsheet([{"file_path":path.name,"name":"orders"}], 'SELECT id, sum(quantity * price) AS amount FROM orders GROUP BY id',output_path="result.xlsx")
    assert result.success, result.model_text
    assert result.value["values"]==[[1,35]]
    wb=load_workbook(tmp_path/"result.xlsx"); assert wb.active["B2"].value==35; wb.close()
    denied=query_spreadsheet([{"file_path":path.name}], "SELECT load_extension('bad')")
    assert not denied.success


def test_validation_reports_uncached_formula_as_partial(tmp_path):
    path = tmp_path / "formula.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["quantity", "price", "amount"])
    ws.append([2, 10, "=A2*B2"])
    wb.save(path)
    wb.close()

    default_partial = validate_spreadsheet(
        path.name,
        [{"kind": "total", "column": "amount", "expected": 20}],
    )
    assert default_partial.success
    assert default_partial.value["validation_status"] == "partial"
    assert default_partial.value["valid"] is None

    blocked = validate_spreadsheet(
        path.name,
        [{"kind": "total", "column": "amount", "expected": 20}],
        allow_uncached=False,
    )
    assert not blocked.success
    assert blocked.error.code == "FORMULA_CACHE_MISSING"
    assert blocked.value["cells"]

    partial = validate_spreadsheet(
        path.name,
        [{"kind": "total", "column": "amount", "expected": 20}],
        allow_uncached=True,
    )
    assert partial.success
    assert partial.value["status"] == "partial"
    assert partial.value["valid"] is None
    assert partial.value["validation_status"] == "partial"
    assert partial.value["uncalculated_cells"]


def test_engine_failures_do_not_modify_source_and_data_conversion(tmp_path,monkeypatch):
    path=tmp_path/"book.xlsx"; book(path,[["id"],[1]])
    before=path.read_bytes()
    assert not calculate_spreadsheet(path.name).success
    monkeypatch.setattr("excelmanus.runtime_capabilities.office_executable",lambda:None)
    assert not render_spreadsheet(path.name,"preview.pdf").success
    assert not (tmp_path/"preview.pdf").exists()
    assert path.read_bytes()==before
    converted=convert_spreadsheet(path.name,"copy.xlsx",mode="data_only")
    assert converted.success, converted.model_text
    assert converted.value["loss_report"]["warnings"]
    assert path.read_bytes()==before


def test_reference_parser_ignores_literals_and_binds_names_and_tables():
    from excelmanus.reference_graph.formula_parser import FormulaRefExtractor
    wb=Workbook(); ws=wb.active; ws.title="Data"; ws.append(["id","amount"]); ws.append([1,10]); table=Table(displayName="Sales",ref="A1:B2"); table._initialise_columns(); table.tableColumns[0].name="id"; table.tableColumns[1].name="amount"; ws.add_table(table)
    wb.defined_names.add(DefinedName("Amounts",attr_text="Data!$B$2:$B$2"))
    extractor=FormulaRefExtractor()
    assert "B2" not in [r.cell_or_range for r in extractor.extract('=IF(A1="B2",C1,D1)')]
    assert extractor.extract("=SUM(Amounts)",wb,"Data")
    assert extractor.extract("=SUM(Sales[amount])",wb,"Data")
