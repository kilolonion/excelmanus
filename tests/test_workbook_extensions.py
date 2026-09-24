"""Preserve LibreOffice workbook metadata through the normal atomic editor."""

from io import BytesIO
from types import ModuleType
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.tools.context import use_workspace
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes, observe_spreadsheet
from excelmanus.workbook_commit import content_version_of_file


_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_LO = "http://schemas.libreoffice.org/"
_URI = "{7626C862-2A13-11E5-B345-FEFF819CDC9F}"
_EXTENSION = f'<extLst><ext uri="{_URI}"><loext:extCalcPr stringRefSyntax="ExcelA1"/></ext></extLst>'
_HEIGHTS = {str(row): height for row, height in enumerate(
    [58, 31, 11, 26, 26, 36.5, 26, 30, 26, 26, 26, 41, 11, 26, 30, 11, 26], 1,
)}


def _rewrite_part(data, part, transform):
    out = BytesIO()
    with ZipFile(BytesIO(data)) as src, ZipFile(out, "w") as dst:
        dst.comment = src.comment
        for item in src.infolist():
            raw = src.read(item.filename)
            dst.writestr(item, transform(raw) if item.filename == part else raw)
    return out.getvalue()


def _extension(data):
    with ZipFile(BytesIO(data)) as package:
        root = ET.fromstring(package.read("xl/workbook.xml"))
    extensions = root.findall(f"{{{_MAIN}}}extLst")
    assert len(extensions) == 1
    return ET.tostring(extensions[0])


@pytest.fixture
def book(tmp_path):
    path = tmp_path / "receipt.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "收款收据"
    ws["A1"] = "收款收据"
    ws.merge_cells("A1:C1")
    ws["D2"] = "=SUM(D3:D4)"
    ws["D3"], ws["D4"] = 100, 200
    wb.create_sheet("说明")["A1"] = "Keep this sheet"
    wb.save(path)
    wb.close()
    # LibreOffice declares loext on the root, not on extLst itself.
    path.write_bytes(_rewrite_part(path.read_bytes(), "xl/workbook.xml", lambda raw: (
        raw.replace(b"<workbook ", f'<workbook xmlns:loext="{_LO}" '.encode(), 1)
        .replace(b"</workbook>", _EXTENSION.encode() + b"</workbook>")
    )))
    path.write_bytes(_rewrite_part(path.read_bytes(), "_rels/.rels", lambda raw: raw.replace(
        b"</Relationships>",
        b'<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>',
    )))
    path.write_bytes(_rewrite_part(path.read_bytes(), "[Content_Types].xml", lambda raw: raw.replace(
        b"</Types>",
        b'<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
    )))
    with ZipFile(path, "a") as package:
        package.writestr("docProps/custom.xml", '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"/>')
    with use_workspace(tmp_path):
        yield path


def _change(book, **kwargs):
    return apply_spreadsheet_changes(
        file_path=book.name, expected_version=content_version_of_file(book),
        operations=[{"kind": "size", "sheet": "收款收据", "row_heights": _HEIGHTS}],
        **kwargs,
    )


def test_row_heights_preserve_libreoffice_calculation_extension(book):
    before = book.read_bytes()
    observed = observe_spreadsheet(file_path=book.name)
    assert observed.success
    assert observed.value["capabilities"]["edit_support"] == "supported_subset"
    result = _change(book)
    assert result.success, result.model_text
    assert result.value["committed"]
    assert result.value["previous_version"] != result.value["content_version"]
    assert _extension(book.read_bytes()) == _extension(before)
    with ZipFile(BytesIO(before)) as src, ZipFile(book) as dst:
        from excelmanus.workbook.adapter import relationship_edges

        assert src.read("docProps/custom.xml") == dst.read("docProps/custom.xml")
        assert relationship_edges(src) <= relationship_edges(dst)
    wb = load_workbook(book)
    try:
        assert {str(row): wb["收款收据"].row_dimensions[row].height for row in range(1, 18)} == _HEIGHTS
        assert str(wb["收款收据"].merged_cells) == "A1:C1"
        assert wb["收款收据"]["D2"].value == "=SUM(D3:D4)"
        assert wb["说明"]["A1"].value == "Keep this sheet"
    finally:
        wb.close()
    assert all(entry["verified"] for entry in result.value["observation"]["geometry_changes"])
    # The output stays editable and does not acquire duplicate extension lists.
    again = _change(book)
    assert again.success, again.model_text
    assert _extension(book.read_bytes()) == _extension(before)


def test_extension_preserving_dry_run_keeps_original_bytes(book):
    before = book.read_bytes()
    result = _change(book, dry_run=True)
    assert result.success, result.model_text
    assert not result.value["committed"]
    assert book.read_bytes() == before
    assert result.value["files"][0]["observation"]["geometry_changes"][0]["verified"]


@pytest.mark.parametrize("extension", [
    _EXTENSION.replace(_URI, "{UNKNOWN}"),
    _EXTENSION.replace("extCalcPr", "unknownCalcFeature"),
    _EXTENSION.replace('stringRefSyntax="ExcelA1"', 'stringRefSyntax="ExcelA1" unknown="data"'),
    _EXTENSION.replace("ExcelA1", "UnknownSyntax"),
])
def test_unrecognized_extensions_remain_blocked_without_data_loss(book, extension):
    book.write_bytes(_rewrite_part(book.read_bytes(), "xl/workbook.xml", lambda raw: (
        raw.replace(_EXTENSION.encode(), extension.encode())
    )))
    before = book.read_bytes()
    result = _change(book)
    assert not result.success
    assert result.error.code == "UNSUPPORTED_PRESERVATION"
    assert result.value["failure_class"] == "unsupported"
    assert not result.value["committed"]
    assert book.read_bytes() == before


def test_worksheet_extensions_are_not_allowlisted_as_workbook_metadata(book):
    book.write_bytes(_rewrite_part(book.read_bytes(), "xl/worksheets/sheet1.xml", lambda raw: (
        raw.replace(b"<worksheet ", f'<worksheet xmlns:loext="{_LO}" '.encode(), 1)
        .replace(b"</worksheet>", _EXTENSION.encode() + b"</worksheet>")
    )))
    before = book.read_bytes()
    result = _change(book)
    assert not result.success and result.error.code == "UNSUPPORTED_PRESERVATION"
    assert book.read_bytes() == before


def test_extension_payload_mismatch_aborts_before_publish(book, monkeypatch):
    from excelmanus.workbook import adapter

    preserve = adapter.preserve_workbook_extensions

    def corrupted(before, after):
        restored = preserve(before, after)
        return _rewrite_part(restored, "xl/workbook.xml", lambda raw: raw.replace(b"ExcelA1", b"CalcA1"))

    monkeypatch.setattr(adapter, "preserve_workbook_extensions", corrupted)
    before = book.read_bytes()
    result = _change(book)
    assert not result.success and result.error.code == "UNSUPPORTED_PRESERVATION"
    assert book.read_bytes() == before


def test_signed_package_stays_blocked(book):
    with ZipFile(book, "a") as package:
        package.writestr("_xmlsignatures/sig1.xml", "<Signature/>")
    before = book.read_bytes()
    result = _change(book)
    assert not result.success and result.error.code == "SIGNED_WORKBOOK"
    assert result.value["failure_class"] == "blocked"
    assert book.read_bytes() == before


def test_sdk_missing_version_helper_explains_observed_version_field():
    from excelmanus.code_mode import render_sdk_source
    from excelmanus.tools.workbook_tools import get_tools

    sdk = ModuleType("em")
    exec(compile(render_sdk_source(get_tools()), "<em>", "exec"), sdk.__dict__)
    with pytest.raises(AttributeError, match="observe_spreadsheet") as error:
        sdk.content_version("receipt.xlsx")
    assert "expected_version" in str(error.value)
    assert "content_version" not in sdk.__all__


def test_probe_failure_points_to_target_dry_run(book):
    result = apply_spreadsheet_changes(
        file_path="outputs/_norm_test.xlsx", create=True,
        operations=[{"kind": "write", "sheet": "收款收据", "start_cell": "A1", "values": [[1]]}],
    )
    assert not result.success and result.error.code == "PROBE_FILE_FORBIDDEN"
    assert "dry_run" in result.value["remediation"]
    assert not (book.parent / "outputs/_norm_test.xlsx").exists()
