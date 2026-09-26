"""2026-09-25 continuation: validation intent and formula acceptance contracts.

The existing cell rule is reused; no alternate rule or cache semantics are added.
Formula recalculation is stubbed at its engine boundary, never via an installed LO.
"""

import json

import pytest
from openpyxl import Workbook

from excelmanus.engine_core.error_payload import (
    ERROR_CODES,
    FORMULA_ERRORS,
    failure_class_for_error_code,
    make_error_payload,
)
from excelmanus.tools.context import use_workspace
from excelmanus.tools.spreadsheet_data_tools import get_tools, validate_spreadsheet
from excelmanus.tools.spreadsheet_engine_tools import calculate_spreadsheet
from excelmanus.workbook_commit import content_version_of_file

SLOPE = 6.707067669172933
INTERCEPT = 1.433759398496254


@pytest.fixture(autouse=True)
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("EXCELMANUS_FORMULA_RECALC", "off")
    with use_workspace(tmp_path):
        yield


@pytest.fixture
def regression_book(tmp_path):
    path = tmp_path / "regression.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "回归分析"
    for row in [
        ["指标", "值"], ["样本量", 5], ["R²", 0.999],
        ["标准误差", 0.31], ["t 统计量", 21.64],
        ["系数", "估计值"], ["斜率", SLOPE], ["截距", INTERCEPT],
        ["说明", "ABC"], ["启用", True], ["禁用", False],
        ["留空", None], ["空白", "  "],
    ]:
        ws.append(row)
    data = wb.create_sheet("数据")
    data.append(["x", "y"])
    data.append([1, SLOPE + INTERCEPT])
    wb.save(path)
    wb.close()
    return path


def check_cell(path, expected, **kwargs):
    rule = {"kind": "cell", "sheet": "回归分析", "cell": "B7", "expected": expected}
    rule.update(kwargs)
    return validate_spreadsheet(path.name, [rule])


def test_seq179_a1_recovery_preserves_single_cell_intent(regression_book):
    result = validate_spreadsheet(regression_book.name, [{
        "kind": "row_expression", "sheet": "回归分析",
        "expression": f"abs(B7-{SLOPE})<0.0001 and abs(B8-{INTERCEPT})<0.0001",
    }])
    assert not result.success
    payload = result.value
    assert payload["error_code"] == "INVALID_ARGS"
    assert payload["failure_class"] == "invalid_args"
    assert payload["a1_style_references"] == ["B7", "B8"]
    for text in (payload["message"], payload["remediation"]):
        assert "cell" in text and "expected" in text and "tolerance" in text
        assert text.index("cell") < text.index("B 或 B{row}")
        assert "仅" in text and "逐行" in text
        assert "不要" in text and "total" in text
    corrected = validate_spreadsheet(regression_book.name, [
        {"kind": "cell", "sheet": "回归分析", "cell": cell, "expected": value, "tolerance": 0.0001}
        for cell, value in [("B7", SLOPE), ("B8", INTERCEPT)]
    ])
    assert corrected.success and corrected.value["valid"] is True
    assert [rule["checked"] for rule in corrected.value["rules"]] == [1, 1]


@pytest.mark.parametrize("expected", [SLOPE, 999])
def test_seq230_total_recovery_does_not_reinterpret_single_cell(regression_book, expected):
    result = validate_spreadsheet(regression_book.name, [{
        "kind": "total", "sheet": "回归分析", "column": "B",
        "expected": expected, "tolerance": 0.0001,
    }])
    assert result.success and result.value["valid"] is False
    failure = next(item for item in result.value["failures"] if item["reason"] == "total_mismatch")
    text = failure["message"]
    assert "整列" in text and "cell" in text and "tolerance" in text
    assert "仅" in text and "逐行" in text
    assert "或用 row_expression" not in text
    if expected == SLOPE:
        assert failure["expected_matches_cell"] == ["B7"]
    else:
        assert "expected_matches_cell" not in failure
    assert check_cell(regression_book, SLOPE, tolerance=0.0001).value["valid"] is True


@pytest.mark.parametrize("kind", ["cell", "total"])
@pytest.mark.parametrize("expected", [float("nan"), float("inf"), -float("inf"), 10 ** 400])
def test_numeric_expected_must_be_finite(regression_book, kind, expected):
    rule = {"kind": kind, "sheet": "回归分析", "expected": expected}
    rule.update({"cell": "B7"} if kind == "cell" else {"column": "B"})
    result = validate_spreadsheet(regression_book.name, [rule])
    assert not result.success, result.model_text
    assert result.value["error_code"] == "INVALID_ARGS"
    assert result.value["failure_class"] == "invalid_args"
    assert result.value["field"] == "expected"
    assert result.value["rule_index"] == 0
    json.dumps(result.value, allow_nan=False)


@pytest.mark.parametrize("kind", ["cell", "total"])
@pytest.mark.parametrize("tolerance", [True, False, -1, float("nan"), float("inf"), -float("inf"), "Infinity", "0.1", None, 10 ** 400])
def test_tolerance_is_finite_nonnegative_number_not_boolean(regression_book, kind, tolerance):
    rule = {"kind": kind, "sheet": "回归分析", "expected": 999, "tolerance": tolerance}
    rule.update({"cell": "B7"} if kind == "cell" else {"column": "B"})
    result = validate_spreadsheet(regression_book.name, [rule])
    assert not result.success, result.model_text
    assert result.value["error_code"] == "INVALID_ARGS"
    assert result.value["failure_class"] == "invalid_args"
    assert result.value["field"] == "tolerance"
    json.dumps(result.value, allow_nan=False)


@pytest.mark.parametrize("cell,expected,extra", [
    ("B7", SLOPE, {"tolerance": 0}),
    ("B7", 6.707, {"tolerance": 0.0001}),
    ("B9", "ABC", {}), ("B9", "abc", {"case_sensitive": False}),
    ("B10", True, {}), ("B11", False, {}),
    ("B12", None, {}), ("B13", None, {}),
])
def test_valid_scalar_assertions_keep_existing_contract(regression_book, cell, expected, extra):
    result = check_cell(regression_book, expected, cell=cell, **extra)
    assert result.success, result.model_text
    assert result.value["valid"] is True
    assert result.value["rules"][0]["checked"] == 1
    json.dumps(result.value, allow_nan=False)


def test_boolean_remains_distinct_from_numeric_assertion(regression_book):
    for cell, expected in [("B10", 1), ("B7", True)]:
        result = check_cell(regression_book, expected, cell=cell)
        assert result.success and result.value["valid"] is False
    result = validate_spreadsheet(regression_book.name, [{
        "kind": "total", "sheet": "回归分析", "column": "B", "expected": True,
    }])
    assert not result.success and result.value["field"] == "expected"


@pytest.mark.parametrize("allow_uncached", [True, False])
def test_cell_cache_partial_and_blocked_contract_unchanged(tmp_path, allow_uncached):
    path = tmp_path / "uncached.xlsx"
    wb = Workbook()
    wb.active.title = "回归分析"
    wb.active["B7"] = "=1+2"
    wb.save(path)
    wb.close()
    result = validate_spreadsheet(path.name, [{
        "kind": "cell", "cell": "B7", "expected": 3,
    }], allow_uncached=allow_uncached)
    if allow_uncached:
        assert result.success and result.value["valid"] is None
        assert result.value["validation_status"] == "partial"
        assert result.value["uncalculated_cells"] == ["回归分析!B7"]
    else:
        assert not result.success and result.error.code == "FORMULA_CACHE_MISSING"
        assert result.value["validation_status"] == "blocked"
        assert result.value["committed"] is False


@pytest.mark.parametrize("output_path", [None, "recalculated.xlsx"])
def test_seq112_formula_acceptance_failure_never_publishes(tmp_path, monkeypatch, output_path):
    path = tmp_path / "header-offset.xlsx"
    wb = Workbook()
    wb.active.title = "回归分析"
    wb.active["B28"] = "销售额"
    wb.active["B29"] = "=B28*2"
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    version = content_version_of_file(path)
    errors = ["回归分析!B29:#VALUE!"]
    monkeypatch.setattr("excelmanus.workbook_commit.recalculate_workbook_bytes", lambda *a, **kw: (
        b"recalculated bytes must not be published",
        {"status": "recalculated", "errors": errors, "error_count": 1},
    ))
    def unexpected_publish(*args, **kwargs):
        pytest.fail("Formula acceptance failure must not reach publication")
    monkeypatch.setattr("excelmanus.tools.spreadsheet_engine_tools._publish", unexpected_publish)
    result = calculate_spreadsheet(path.name, expected_version=version, output_path=output_path)
    assert not result.success
    payload = result.value
    assert payload["error_code"] == FORMULA_ERRORS
    assert payload["failure_class"] == "invalid_args"
    assert payload["validation_status"] == "failed"
    assert payload["committed"] is False
    assert payload["source_version"] == version and payload["file_path"] == path.name
    assert payload["formula_recalculation"]["errors"] == errors
    assert path.read_bytes() == before and content_version_of_file(path) == version
    if output_path:
        assert not (tmp_path / output_path).exists()
    hint = payload["remediation"]
    for required in ("B29", "#VALUE!", "表头", "行号偏移", "observe_spreadsheet", "source_version", "expected_version", "新 content_version", "committed=false"):
        assert required in hint
    assert "不要" in hint and "allow_formula_errors=true" in hint


def test_formula_errors_taxonomy_and_bounded_coordinate_preview():
    assert FORMULA_ERRORS in ERROR_CODES
    assert failure_class_for_error_code(FORMULA_ERRORS) == "invalid_args"
    errors = [f"回归分析!B{row}:#VALUE!" for row in range(29, 33)]
    payload = make_error_payload("公式验收失败", error_code=FORMULA_ERRORS, formula_recalculation={"errors": errors})
    assert payload["failure_class"] == "invalid_args"
    assert "B29:#VALUE!" in payload["remediation"]
    assert "共 4 项" in payload["remediation"]
    assert "完整列表见 formula_recalculation.errors" in payload["remediation"]
    assert payload["formula_recalculation"]["errors"] == errors


def test_schema_exposes_existing_cell_once_and_finite_numeric_contract():
    tool = next(tool for tool in get_tools() if tool.name == "validate_spreadsheet")
    props = tool.input_schema["properties"]["rules"]["items"]["properties"]
    assert props["kind"]["enum"].count("cell") == 1
    assert "有限" in props["expected"]["description"]
    assert "有限" in props["tolerance"]["description"]
