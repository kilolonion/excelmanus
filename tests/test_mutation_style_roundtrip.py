"""New cells on engine-converted workbooks must survive the serialization check.

Excel and LibreOffice spell the OOXML defaults out in the workbook's first xf
(``<alignment horizontal="general" vertical="bottom"/>``), while openpyxl reports
a freshly created cell with bare defaults. The post-serialization verification in
``excelmanus.workbook.mutation`` therefore compares resolved styles through
``_style_signature`` instead of ``==`` on the style proxies. These tests pin both
halves of that contract: equivalent spellings pass, real drift still aborts.
"""

import re
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

from excelmanus.tools.context import use_workspace
from excelmanus.tools.spreadsheet_engine_tools import calculate_spreadsheet, convert_spreadsheet
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes
from excelmanus.workbook.mutation import _style_signature
from excelmanus.workbook_commit import content_version_of_file


def _office_engine():
    from excelmanus.runtime_capabilities import office_executable

    return office_executable()


def _rewrite_part(data: bytes, part: str, transform) -> bytes:
    out = BytesIO()
    with ZipFile(BytesIO(data)) as src, ZipFile(out, "w") as dst:
        dst.comment = src.comment
        for item in src.infolist():
            raw = src.read(item.filename)
            dst.writestr(item, transform(raw) if item.filename == part else raw)
    return out.getvalue()


def _explicit_default_alignment(data: bytes) -> bytes:
    """Mimic the engine xf: spell horizontal=general / vertical=bottom out."""
    def patch(raw: bytes) -> bytes:
        return re.sub(
            rb"(<cellXfs[^>]*>\s*<xf\b[^>]*?)/>",
            rb'\1><alignment horizontal="general" vertical="bottom"/></xf>',
            raw,
            count=1,
        )

    return _rewrite_part(data, "xl/styles.xml", patch)


@pytest.fixture
def workspace(tmp_path):
    with use_workspace(tmp_path):
        yield tmp_path


def _sample_book(path, *, engine_defaults: bool = False):
    wb = Workbook()
    ws = wb.active
    ws.title = "input"
    ws.append(["item", "qty", "price", "total"])
    ws.append(["widget", 2, 3.5, "=B2*C2"])
    ws.append(["gadget", 4, 1.25, "=B3*C3"])
    wb.save(path)
    wb.close()
    if engine_defaults:
        path.write_bytes(_explicit_default_alignment(path.read_bytes()))
    return path


def _edit(relative_path, path, operations):
    return apply_spreadsheet_changes(
        file_path=relative_path,
        expected_version=content_version_of_file(path),
        operations=operations,
    )


def test_style_signature_folds_engine_default_alignment(tmp_path):
    """The old false positive: fresh cell None/None vs reloaded general/bottom."""
    path = _sample_book(tmp_path / "converted.xlsx", engine_defaults=True)
    wb = load_workbook(path)
    fresh = wb["input"].cell(1, 5)
    fresh.value = "margin"
    assert fresh.alignment.horizontal is None and fresh.alignment.vertical is None
    out = BytesIO()
    wb.save(out)
    wb.close()
    reloaded = load_workbook(BytesIO(out.getvalue()))["input"]["E1"]
    # The reloaded cell carries what the file spells out; both are OOXML defaults.
    assert reloaded.alignment.horizontal == "general"
    assert reloaded.alignment.vertical == "bottom"
    assert _style_signature(fresh) == _style_signature(reloaded)


@pytest.mark.parametrize(
    "mutate, facet",
    [
        (lambda cell: setattr(cell, "number_format", "0.00"), "number_format"),
        (lambda cell: setattr(cell, "number_format", "yyyy-mm-dd"), "number_format"),
        (lambda cell: setattr(cell, "alignment", Alignment(horizontal="center")), "alignment"),
        (lambda cell: setattr(cell, "alignment", Alignment(vertical="top")), "alignment"),
        (lambda cell: setattr(cell, "alignment", Alignment(wrapText=True)), "alignment"),
        (lambda cell: setattr(cell, "font", Font(bold=True)), "font"),
        (lambda cell: setattr(cell, "font", Font(name="Courier New")), "font"),
        (lambda cell: setattr(cell, "font", Font(color="FFFF0000")), "font"),
        (lambda cell: setattr(cell, "fill", PatternFill(fill_type="solid", fgColor="FFFF00")), "fill"),
        (lambda cell: setattr(cell, "border", Border(left=Side(style="thin"))), "border"),
        (lambda cell: setattr(cell, "border", Border(bottom=Side(style="double"))), "border"),
    ],
)
def test_style_signature_still_detects_real_drift(tmp_path, mutate, facet):
    path = _sample_book(tmp_path / "drift.xlsx")
    wb = load_workbook(path)
    cell = wb["input"]["A1"]
    before = _style_signature(cell)
    mutate(cell)
    after = _style_signature(cell)
    assert before[facet] != after[facet], f"{facet} drift must stay visible"
    wb.close()


def test_style_signature_tolerates_equivalent_spellings(tmp_path):
    path = _sample_book(tmp_path / "equivalent.xlsx")
    wb = load_workbook(path)
    cell = wb["input"]["A1"]
    before = _style_signature(cell)
    # Documented OOXML alignment defaults written out explicitly.
    cell.alignment = Alignment(
        horizontal="general", vertical="bottom", textRotation=0, wrapText=False, indent=0
    )
    # Default number format and class defaults written out explicitly.
    cell.number_format = "General"
    cell.fill = PatternFill(patternType=None)
    cell.border = Border(
        left=Side(style=None), right=Side(style=None), top=Side(style=None), bottom=Side(style=None)
    )
    cell.font = Font(name="Calibri", sz=11.0, family=2.0, scheme="minor", color=Color(theme=1))
    assert _style_signature(cell) == before
    wb.close()


def test_new_cells_commit_on_plain_workbook(workspace, tmp_path):
    book = _sample_book(tmp_path / "plain.xlsx")
    result = _edit(book.name, book, [
        {"kind": "write", "sheet": "input", "start_cell": "E1", "values": [["margin"], ["=D2*0.1"], ["=D3*0.1"]]},
    ])
    assert result.success, result.model_text
    assert result.value["committed"]
    wb = load_workbook(book)
    try:
        assert [wb["input"][f"E{row}"].value for row in (1, 2, 3)] == ["margin", "=D2*0.1", "=D3*0.1"]
    finally:
        wb.close()


@pytest.mark.skipif(not _office_engine(), reason="no local LibreOffice engine")
def test_new_column_commits_and_recalculates_on_converted_csv(workspace, tmp_path):
    """csv -> convert_spreadsheet -> next round new column -> recalculate."""
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    (uploads / "sales.csv").write_text(
        "item,qty,price,total\nwidget,2,3.5,=B2*C2\ngadget,4,1.25,=B3*C3\n",
        encoding="utf-8",
    )
    converted = convert_spreadsheet("uploads/sales.csv", "outputs/sales.xlsx")
    assert converted.success, converted.model_text
    target = tmp_path / "outputs" / "sales.xlsx"

    result = _edit("outputs/sales.xlsx", target, [
        {
            "kind": "write",
            "sheet": "input",
            "start_cell": "E1",
            "values": [["margin"], ["=D2*0.1"], ["=D3*0.1"]],
        },
    ])
    assert result.success, result.model_text
    assert result.value["committed"]

    wb = load_workbook(target)
    try:
        ws = wb["input"]
        assert [ws[f"E{row}"].value for row in (1, 2, 3)] == ["margin", "=D2*0.1", "=D3*0.1"]
        # Nothing else moved.
        assert [[ws.cell(row, column).value for column in range(1, 5)] for row in range(1, 4)] == [
            ["item", "qty", "price", "total"],
            ["widget", 2, 3.5, "=B2*C2"],
            ["gadget", 4, 1.25, "=B3*C3"],
        ]
    finally:
        wb.close()

    recalculated = calculate_spreadsheet(
        "outputs/sales.xlsx", expected_version=content_version_of_file(target)
    )
    assert recalculated.success, recalculated.model_text
    assert recalculated.value["formula_recalculation"]["status"] == "recalculated"
    assert recalculated.value["formula_recalculation"]["error_count"] == 0
    wb = load_workbook(target, data_only=True)
    try:
        ws = wb["input"]
        # D2 = 2*3.5 = 7, D3 = 4*1.25 = 5, E2/E3 = 10% of those.
        assert [ws["D2"].value, ws["D3"].value, ws["E2"].value, ws["E3"].value] == [7, 5, 0.7, 0.5]
    finally:
        wb.close()


@pytest.mark.skipif(not _office_engine(), reason="no local LibreOffice engine")
def test_new_cells_commit_on_converted_xlsx(workspace, tmp_path):
    """The pre-existing repro: xlsx -> engine xlsx already failed for new cells."""
    source = _sample_book(tmp_path / "source.xlsx")
    converted = convert_spreadsheet(source.name, "outputs/roundtrip.xlsx", mode="preserve")
    assert converted.success, converted.model_text
    target = tmp_path / "outputs" / "roundtrip.xlsx"
    result = _edit("outputs/roundtrip.xlsx", target, [
        {"kind": "write", "sheet": "input", "start_cell": "A4", "values": [["sprocket", 8, 2.0, "=B4*C4"]]},
    ])
    assert result.success, result.model_text
    wb = load_workbook(target)
    try:
        assert [wb["input"][f"A4"].value, wb["input"][f"D4"].value] == ["sprocket", "=B4*C4"]
    finally:
        wb.close()


def _drift_during_save(monkeypatch, mutate):
    original_save = OpenpyxlWorkbook.save

    def drifting_save(workbook, filename):
        sheet = workbook["input"] if "input" in workbook.sheetnames else None
        if sheet is not None and sheet["D2"].value is not None:
            mutate(sheet["D2"])
        return original_save(workbook, filename)

    monkeypatch.setattr(OpenpyxlWorkbook, "save", drifting_save)


@pytest.mark.parametrize(
    "mutate, expected_drift",
    [
        (lambda cell: setattr(cell, "value", "=B2*C2+1"), "value"),
        (lambda cell: setattr(cell, "number_format", "0.00"), "style.number_format"),
        (lambda cell: setattr(cell, "alignment", Alignment(horizontal="center")), "style.alignment"),
    ],
)
def test_real_drift_still_aborts_before_publish(workspace, tmp_path, monkeypatch, mutate, expected_drift):
    book = _sample_book(tmp_path / "commit.xlsx")
    before = book.read_bytes()
    _drift_during_save(monkeypatch, mutate)
    result = _edit(book.name, book, [
        {"kind": "write", "sheet": "input", "start_cell": "D2", "values": [["=B2*C2"]]},
    ])
    assert not result.success
    assert result.error.code == "SERIALIZATION_MISMATCH"
    assert expected_drift in result.error.message
    assert result.value["committed"] is False
    # Nothing was published; the verified sample is what protects the caller.
    assert book.read_bytes() == before
