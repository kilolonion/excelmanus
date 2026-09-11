"""P3 invariant: writes stay on the user path; history is RevisionStore."""

from __future__ import annotations

import sys
from pathlib import Path

from excelmanus.security.guard import FileAccessGuard
from excelmanus.tools import code_tools
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.workbook_commit import commit_bytes
from excelmanus.workspace.revisions import RevisionStore


def _payload(result) -> dict:
    return dict(result.value) if isinstance(result.value, dict) else {}


def test_run_code_write_stays_on_user_path(tmp_path: Path) -> None:
    set_guard(FileAccessGuard(str(tmp_path)))
    target = tmp_path / "book.txt"
    target.write_text("original", encoding="utf-8")
    result = _payload(
        code_tools.run_code(
            code=f"open(r'{target}', 'w').write('updated')",
            python_command=sys.executable,
            sandbox_tier="GREEN",
        )
    )
    assert result["status"] == "success"
    assert "cow_mapping" not in result
    assert "cow_hint" not in result
    assert target.read_text(encoding="utf-8") == "updated"
    assert not (tmp_path / "outputs" / "backups").exists()


def test_commit_bytes_writes_revisions_not_backups(tmp_path: Path) -> None:
    guard = FileAccessGuard(str(tmp_path))
    first = commit_bytes(guard=guard, file_path="book.xlsx", data=b"v1", expected_version=None)
    commit_bytes(guard=guard, file_path="book.xlsx", data=b"v2", expected_version=first.content_version)
    recs = RevisionStore(tmp_path).list("book.xlsx")
    assert [r.reason for r in recs] == ["afterEdit", "beforeEdit", "afterEdit"] or [
        r.reason for r in recs
    ].count("afterEdit") == 2
    assert (tmp_path / "book.xlsx").read_bytes() == b"v2"
    assert not (tmp_path / "outputs" / "backups").exists()
    assert (tmp_path / ".excelmanus" / "revisions").is_dir()


def test_run_code_openpyxl_save_publishes_to_user_path(tmp_path: Path) -> None:
    set_guard(FileAccessGuard(str(tmp_path)))
    target = tmp_path / "book.xlsx"
    from openpyxl import Workbook, load_workbook
    from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions

    wb = Workbook()
    wb.active["A1"] = "original"
    wb.save(str(target))
    wb.close()
    seed_seen_versions({"book.xlsx": content_version_of_file(target)})
    result = _payload(
        code_tools.run_code(
            code=(
                "from openpyxl import load_workbook\n"
                f"wb = load_workbook(r'{target}')\n"
                "wb.active['A1'] = 'updated'\n"
                f"wb.save(r'{target}')\n"
            ),
            python_command=sys.executable,
            sandbox_tier="GREEN",
        )
    )
    assert result["status"] == "success", result
    assert "cow_mapping" not in result
    published = result.get("published") or []
    assert any(
        item.get("path") == "book.xlsx" and item.get("status") == "committed"
        for item in published
    )
    wb2 = load_workbook(str(target))
    assert wb2.active["A1"].value == "updated"
    wb2.close()
    recs = RevisionStore(tmp_path).list("book.xlsx")
    assert any(r.reason == "afterEdit" for r in recs)
    assert not (tmp_path / "outputs" / "backups").exists()
    pending = tmp_path / ".excelmanus" / "pending"
    if pending.is_dir():
        assert not any(pending.iterdir())
