"""Unified tool routing through the real loop, SDK subprocess, and workbook commit.

Only model responses are scripted; no external model or provider is contacted.
"""

from __future__ import annotations

import json
import sys
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.agent.session import AgentEngine
from excelmanus.config import ExcelManusConfig
from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.tools.registry import ToolDef, ToolRegistry


def _response(name: str | None = None, arguments: dict | None = None, call_id: str = ""):
    calls = None if name is None else [SimpleNamespace(
        id=call_id,
        function=SimpleNamespace(name=name, arguments=json.dumps(arguments or {})),
    )]
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(
        content="完成" if name is None else "", tool_calls=calls,
    ))])


@pytest.mark.asyncio
async def test_direct_inspect_programmatic_write_direct_inspect_in_one_turn(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active.append(["item", "amount"])
    workbook.active.append(["rent", 12])
    workbook.save(tmp_path / "book.xlsx")
    workbook.close()
    registry = ToolRegistry()
    registry.register_builtin_tools(str(tmp_path))
    engine = AgentEngine(ExcelManusConfig(
        api_key="test", base_url="https://test.invalid/v1", model="test-model",
        workspace_root=str(tmp_path), main_model_vision="false", max_iterations=20,
    ), registry)
    engine._session_id = "unified-flow"
    engine._full_access_enabled = True
    code = (
        "import em\n"
        "before = em.observe_spreadsheet(file_path='book.xlsx', mode='range', range='A1:B2')\n"
        "result = em.apply_spreadsheet_changes(file_path='book.xlsx', "
        "expected_version=before['content_version'], operations=["
        "{'kind': 'write', 'sheet': 'Sheet1', 'start_cell': 'B2', 'values': [[42]]}])\n"
        "print(result['content_version'])\n"
    )
    inspect_args = {"file_path": "book.xlsx", "mode": "range", "range": "A1:B2"}
    responses = iter([
        _response("observe_spreadsheet", inspect_args, "direct-before"),
        _response("run_code", {
            "code": code, "python_command": sys.executable,
            "timeout_seconds": 30, "require_excel_deps": False,
        }, "program"),
        _response("observe_spreadsheet", inspect_args, "direct-after"),
        _response(),
    ])
    wire_tool_sets = []

    async def create(**kwargs):
        wire_tool_sets.append({t["function"]["name"] for t in kwargs.get("tools", [])})
        return next(responses)

    engine._client.chat.completions.create = AsyncMock(side_effect=create)
    events = []
    result = await engine.followup(
        "先检查 book.xlsx，在程序中把 Sheet1!B2 更新为 42，再直接检查结果。",
        on_event=events.append,
    )
    assert result.reply == "完成"
    assert not result.truncated
    assert result.tool_calls and all(t.success for t in result.tool_calls), [
        (t.tool_name, t.result) for t in result.tool_calls
    ]
    assert all({"observe_spreadsheet", "apply_spreadsheet_changes", "run_code"} <= names
               for names in wire_tool_sets)
    actual = load_workbook(tmp_path / "book.xlsx")
    try:
        assert actual["Sheet1"]["B2"].value == 42
    finally:
        actual.close()
    assert any(getattr(event, "parent_call_id", None) == "program" for event in events)


@pytest.mark.asyncio
async def test_mcp_detail_loads_next_request_and_persists_across_turns(tmp_path):
    call = Mock(return_value=ToolResult(success=True, model_text="count: 7", value={"count": 7}))
    registry = ToolRegistry()
    registry.register_tools([ToolDef(
        name="mcp_demo_count", description="Count demo entries.",
        input_schema={"type": "object", "properties": {}},
        func=call, write_effect="none",
    )])
    engine = AgentEngine(ExcelManusConfig(
        api_key="test", base_url="https://test.invalid/v1", model="test-model",
        workspace_root=str(tmp_path), main_model_vision="false",
    ), registry)
    engine._session_id = "discovery-flow"
    responses = iter([
        _response("introspect_capability", {
            "query_type": "tool_detail", "query": "mcp_demo_count.missing",
        }, "missing-field"),
        _response("introspect_capability", {
            "query_type": "tool_detail", "query": "mcp_demo_count",
        }, "load-detail"),
        _response("mcp_demo_count", {}, "call-discovered-tool"),
        _response(),
        _response(),
    ])
    visible = []

    async def create(**kwargs):
        visible.append("mcp_demo_count" in {t["function"]["name"] for t in kwargs.get("tools", [])})
        return next(responses)

    engine._client.chat.completions.create = AsyncMock(side_effect=create)
    first = await engine.followup("查询并使用 demo 计数工具")
    assert first.reply == "完成"
    assert all(tool.success for tool in first.tool_calls)
    call.assert_called_once_with()
    assert visible == [False, False, True, True]
    second = await engine.followup("开始独立的新任务")
    assert second.reply == "完成"
    # 工具披露属于会话级状态：新一轮仍保留已披露工具，保持 tools 前缀稳定。
    assert visible == [False, False, True, True, True]


@pytest.mark.asyncio
async def test_builtin_version_tool_stays_loaded_after_real_compaction(tmp_path, monkeypatch):
    from excelmanus.compaction import COMPACTION_SYSTEM_PROMPT
    workbook = Workbook()
    workbook.active["A1"] = "unchanged"
    workbook.save(tmp_path / "book.xlsx")
    workbook.close()
    registry = ToolRegistry()
    registry.register_builtin_tools(str(tmp_path))
    engine = AgentEngine(ExcelManusConfig(
        api_key="test", base_url="https://test.invalid/v1", model="test-model",
        workspace_root=str(tmp_path), main_model_vision="false",
        compaction_keep_recent_turns=2,
    ), registry)
    engine._session_id = "builtin-discovery-compaction"
    for index in range(10):
        engine._memory.add_user_message(f"已完成的历史问题 {index}")
        # Real compaction keeps a structured handoff as well as the summary;
        # enough source history must exist for the replacement to be smaller.
        engine._memory.add_assistant_message(
            f"历史问题 {index} 已完成。" + "已核对工作表中的列名和数据，原文件保持不变。" * 50
        )
    responses = iter([
        _response("introspect_capability", {
            "query_type": "tool_detail", "query": "manage_spreadsheet_versions",
        }, "load-versions"),
        _response("manage_spreadsheet_versions", {
            "file_path": "book.xlsx", "action": "list",
        }, "list-versions"),
        _response(),
        _response(),
    ])
    # Trigger the normal Driver pre-step boundary exactly after discovery.
    monkeypatch.setattr(engine._compaction_manager, "should_compact", lambda *_:
        "manage_spreadsheet_versions" in getattr(engine, "_loaded_tool_names", set())
        and engine._compaction_manager.stats.compaction_count == 0
    )
    visible = []
    summaries = []

    async def create(**kwargs):
        last = kwargs["messages"][-1].get("content", "")
        if isinstance(last, str) and last.startswith(COMPACTION_SYSTEM_PROMPT):
            summaries.append(last)
            return _response()
        names = {tool["function"]["name"] for tool in kwargs.get("tools", [])}
        assert "run_code" not in names  # discovery must not expand read permissions
        visible.append("manage_spreadsheet_versions" in names)
        if len(visible) == 2:
            assert engine._compaction_manager.stats.compaction_count == 1
            assert "manage_spreadsheet_versions" in engine._loaded_tool_names
        return next(responses)

    engine._client.chat.completions.create = AsyncMock(side_effect=create)
    first = await engine.followup("查看 book.xlsx 的版本记录", chat_mode="read")
    assert first.reply == "完成"
    assert all(tool.success for tool in first.tool_calls)
    assert visible == [False, True, True]
    assert len(summaries) == 1
    await engine.followup("新的只读任务", chat_mode="read")
    # 披露会话内保留；仍与当前授权目录求交，run_code 不会泄露进只读模式。
    assert visible == [False, True, True, True]
    actual = load_workbook(tmp_path / "book.xlsx")
    try:
        assert actual.active["A1"].value == "unchanged"
    finally:
        actual.close()
