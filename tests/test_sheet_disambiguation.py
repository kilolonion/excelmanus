"""P0-B 列感知 sheet 消歧：网关决策表 + 4 入口回填 + 非沉默三件套。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from excelmanus.security import FileAccessGuard
from excelmanus.tools import intent_tools, reference_tools
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.intent_tools import analyze_spreadsheet
from excelmanus.workbook.snapshot import (
    SheetRequired,
    resolve_sheet_by_columns,
    resolve_sheet_by_visibility,
)


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    intent_tools.init_guard(workspace)
    reference_tools.init_guard(workspace)


def _save(path: Path, sheets: dict[str, list[list[object]]], hidden: set[str] | None = None) -> Path:
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        if hidden and name in hidden:
            ws.sheet_state = "hidden"
    wb.save(path)
    wb.close()
    return path


def _probe(mapping: dict[str, tuple[list[str] | None, str | None, bool]]):
    def _run(sheet: str):
        return mapping[sheet]
    return _run


# ── 网关纯单测 ─────────────────────────────────────────────


class TestGateway:
    def test_unique_exact_hit_binds(self):
        sheet, info = resolve_sheet_by_columns(
            ["订单", "产品"],
            ["数量"],
            _probe({"订单": (["订单号", "数量"], None, False), "产品": (["产品ID", "名称"], None, False)}),
        )
        assert sheet == "订单"
        assert info["strategy"] == "column_unique_match"
        assert info["normalized_match"] is False

    def test_no_hints_fails(self):
        try:
            resolve_sheet_by_columns(["订单", "产品"], [], _probe({}))
        except SheetRequired as exc:
            assert exc.fields["available_sheets"] == ["订单", "产品"]
        else:  # pragma: no cover
            raise AssertionError("应抛 SHEET_REQUIRED")

    def test_multi_hit_fails_with_matrix(self):
        try:
            resolve_sheet_by_columns(
                ["订单", "退货"],
                ["状态"],
                _probe({"订单": (["订单号", "状态"], None, False), "退货": (["订单号", "状态"], None, False)}),
            )
        except SheetRequired as exc:
            assert "hint_match_matrix" in exc.fields
        else:  # pragma: no cover
            raise AssertionError("应抛 SHEET_REQUIRED")

    def test_alias_not_matched(self):
        """金额 vs 金额(元)：别名不算命中 → 零命中失败。"""
        try:
            resolve_sheet_by_columns(
                ["订单", "产品"],
                ["金额"],
                _probe({"订单": (["订单号", "金额(元)"], None, False), "产品": (["产品ID", "名称"], None, False)}),
            )
        except SheetRequired:
            pass
        else:  # pragma: no cover
            raise AssertionError("别名不应命中")

    def test_hidden_collision_fails(self):
        """可见+隐藏同时命中 → 歧义失败（永不自动选中隐藏表）。"""
        try:
            resolve_sheet_by_columns(
                ["订单", "底稿"],
                ["数量"],
                _probe({"订单": (["订单号", "数量"], None, False), "底稿": (["订单号", "数量"], None, True)}),
            )
        except SheetRequired as exc:
            assert "隐藏" in exc.fields["disambiguation_skipped_reason"]
        else:  # pragma: no cover
            raise AssertionError("隐藏表碰撞应失败")

    def test_hidden_only_hit_fails(self):
        try:
            resolve_sheet_by_columns(
                ["订单", "底稿"],
                ["内部码"],
                _probe({"订单": (["订单号", "数量"], None, False), "底稿": (["内部码"], None, True)}),
            )
        except SheetRequired:
            pass
        else:  # pragma: no cover
            raise AssertionError("仅隐藏表命中应失败")

    def test_form_empty_excluded(self):
        sheet, info = resolve_sheet_by_columns(
            ["订单", "表单", "空表"],
            ["数量"],
            _probe({
                "订单": (["订单号", "数量"], None, False),
                "表单": (None, "form", False),
                "空表": (None, "empty", False),
            }),
        )
        assert sheet == "订单"
        assert info["excluded"] == {"form": ["表单"], "empty": ["空表"]}

    def test_normalized_case_insensitive_binds_with_flag(self):
        sheet, info = resolve_sheet_by_columns(
            ["订单", "产品"],
            ["amount"],
            _probe({"订单": (["Amount", "qty"], None, False), "产品": (["产品ID", "名称"], None, False)}),
        )
        assert sheet == "订单"
        assert info["normalized_match"] is True

    def test_normalized_collision_fails(self):
        """归一后多表碰撞（金额 vs '金额 '）→ 失败。"""
        try:
            resolve_sheet_by_columns(
                ["A", "B"],
                ["金额"],
                _probe({"A": (["金额"], None, False), "B": (["金额 "], None, False)}),
            )
        except SheetRequired:
            pass
        else:  # pragma: no cover
            raise AssertionError("归一碰撞应失败")

    def test_too_many_sheets_skips(self):
        names = [f"S{i}" for i in range(11)]
        try:
            resolve_sheet_by_columns(names, ["数量"], _probe({n: (["数量"], None, False) for n in names}))
        except SheetRequired as exc:
            assert "disambiguation_skipped_reason" in exc.fields
        else:  # pragma: no cover
            raise AssertionError("超限应降级失败")

    def test_sole_visible_sheet_binds_without_column_hints(self):
        sheet, info = resolve_sheet_by_visibility(
            ["订单", "底稿"], {"订单": False, "底稿": True},
        )
        assert sheet == "订单"
        assert info["strategy"] == "sole_visible_sheet"
        assert info["hidden_sheets"] == ["底稿"]

    def test_multiple_visible_sheets_do_not_bind_without_hints(self):
        try:
            resolve_sheet_by_visibility(
                ["订单", "产品", "底稿"],
                {"订单": False, "产品": False, "底稿": True},
            )
        except SheetRequired as exc:
            assert exc.fields["visible_sheets"] == ["订单", "产品"]
        else:  # pragma: no cover
            raise AssertionError("多张可见表不应自动绑定")


# ── 工作簿级（工具全链路） ─────────────────────────────────


def _two_sheet_workbook(path: Path) -> Path:
    return _save(
        path,
        {
            "订单": [["订单号", "数量", "状态"], ["S1", 2, "已完成"], ["S2", 1, "已取消"]],
            "产品": [["产品ID", "名称"], ["P1", "苹果"]],
        },
    )


class TestWorkbookLevel:
    def test_aggregate_unique_hit_binds_with_declaration(self, tmp_path: Path):
        """正例 + 非沉默三件套：resolved_sheet + warnings + model_text ⚠️。"""
        _bind(tmp_path)
        path = _two_sheet_workbook(tmp_path / "w.xlsx")
        res = analyze_spreadsheet(
            mode="aggregate", file_path=str(path),
            group_by="状态", aggregations={"数量": "sum"},
        )
        assert res.success, res.model_text
        assert res.value["resolved_sheet"] == "订单"
        assert res.value["warnings"], "warnings 不能为空（显式声明）"
        assert "订单" in res.value["warnings"][0]
        assert "⚠️" in res.model_text and "订单" in res.model_text
        info = res.value["sheet_disambiguation"]
        assert info["strategy"] == "column_unique_match"
        assert set(info["hints"]) >= {"状态", "数量"}

    def test_ambiguous_still_fails(self, tmp_path: Path):
        """两表都有 状态 列 → 仍 SHEET_REQUIRED。"""
        _bind(tmp_path)
        path = _save(
            tmp_path / "w.xlsx",
            {
                "订单": [["订单号", "状态"], ["S1", "已完成"]],
                "退货": [["订单号", "状态"], ["R1", "已退"]],
            },
        )
        res = analyze_spreadsheet(
            mode="distinct", file_path=str(path), column="状态",
        )
        assert not res.success
        assert "SHEET_REQUIRED" in res.model_text

    def test_explicit_sheet_unchanged_no_warning(self, tmp_path: Path):
        """回归：显式指定 sheet 不产生消歧声明。"""
        _bind(tmp_path)
        path = _two_sheet_workbook(tmp_path / "w.xlsx")
        res = analyze_spreadsheet(
            mode="aggregate", file_path=str(path), sheet_name="订单",
            group_by="状态", aggregations={"数量": "sum"},
        )
        assert res.success, res.model_text
        assert res.value["resolved_sheet"] == "订单"
        assert "sheet_disambiguation" not in res.value
        assert "sheet_disambiguation_warning" not in res.value

    def test_single_sheet_omitted_no_warning(self, tmp_path: Path):
        """回归：单表省略走 S0，无消歧 warning。"""
        _bind(tmp_path)
        path = _save(tmp_path / "w.xlsx", {"订单": [["订单号", "数量"], ["S1", 2]]})
        res = analyze_spreadsheet(
            mode="aggregate", file_path=str(path),
            group_by="订单号", aggregations={"数量": "sum"},
        )
        assert res.success, res.model_text
        assert "sheet_disambiguation" not in res.value

    def test_distinct_unique_hit_binds(self, tmp_path: Path):
        _bind(tmp_path)
        path = _two_sheet_workbook(tmp_path / "w.xlsx")
        res = analyze_spreadsheet(mode="distinct", file_path=str(path), column="状态")
        assert res.success, res.model_text
        assert res.value["resolved_sheet"] == "订单"
        assert "⚠️" in res.model_text

    def test_filter_unique_hit_binds(self, tmp_path: Path):
        _bind(tmp_path)
        path = _two_sheet_workbook(tmp_path / "w.xlsx")
        res = analyze_spreadsheet(
            mode="filter", file_path=str(path),
            column="状态", operator="eq", value="已完成",
        )
        assert res.success, res.model_text
        assert res.value["resolved_sheet"] == "订单"

    def test_hidden_sheet_collision_fails_workbook(self, tmp_path: Path):
        _bind(tmp_path)
        path = _save(
            tmp_path / "w.xlsx",
            {
                "订单": [["订单号", "数量"], ["S1", 2]],
                "底稿": [["订单号", "数量"], ["S9", 9]],
            },
            hidden={"底稿"},
        )
        res = analyze_spreadsheet(
            mode="aggregate", file_path=str(path),
            group_by="订单号", aggregations={"数量": "sum"},
        )
        assert not res.success
        assert "SHEET_REQUIRED" in res.model_text

    def test_empty_sheet_excluded(self, tmp_path: Path):
        _bind(tmp_path)
        wb = Workbook()
        ws = wb.active
        ws.title = "订单"
        ws.append(["订单号", "数量"])
        ws.append(["S1", 2])
        wb.create_sheet("空表")
        path = tmp_path / "w.xlsx"
        wb.save(path)
        wb.close()
        res = analyze_spreadsheet(
            mode="aggregate", file_path=str(path),
            group_by="订单号", aggregations={"数量": "sum"},
        )
        assert res.success, res.model_text
        assert res.value["resolved_sheet"] == "订单"

    def test_aggregate_without_hints_binds_sole_visible_sheet(self, tmp_path: Path):
        _bind(tmp_path)
        path = _save(
            tmp_path / "w.xlsx",
            {
                "订单": [["订单号", "数量"], ["S1", 2]],
                "底稿": [["内部码"], ["X1"]],
            },
            hidden={"底稿"},
        )
        res = analyze_spreadsheet(mode="aggregate", file_path=str(path))
        assert res.success, res.model_text
        assert res.value["resolved_sheet"] == "订单"
        assert res.value["sheet_disambiguation"]["strategy"] == "sole_visible_sheet"
        assert "唯一可见工作表" in res.value["warnings"][0]
        assert "⚠️" in res.model_text

    def test_aggregate_without_hints_keeps_visible_ambiguity(self, tmp_path: Path):
        _bind(tmp_path)
        path = _two_sheet_workbook(tmp_path / "w.xlsx")
        res = analyze_spreadsheet(mode="aggregate", file_path=str(path))
        assert not res.success
        assert "SHEET_REQUIRED" in res.model_text

    def test_pivot_unique_column_match_binds(self, tmp_path: Path):
        _bind(tmp_path)
        path = _two_sheet_workbook(tmp_path / "w.xlsx")
        res = analyze_spreadsheet(
            mode="pivot", file_path=str(path),
            index="状态", columns="订单号", values="数量", aggfunc="sum",
        )
        assert res.success, res.model_text
        assert res.value["resolved_sheet"] == "订单"
        assert res.value["sheet_disambiguation"]["strategy"] == "column_unique_match"
