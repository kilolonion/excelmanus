"""Workbook questions: scope, optimistic versions, event delivery and HTTP resume."""
from __future__ import annotations

import asyncio
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook

from excelmanus.engine_core.interaction_handler import InteractionHandler
from excelmanus.engine_core.tool_handlers import ShowWorkbookHandler
from excelmanus.interaction import InteractionRegistry
from excelmanus.question_flow import QuestionFlowManager
from excelmanus.workbook.interaction import bind_target, prepare_questions, validate_selection_answer
from excelmanus.workbook.snapshot import SnapshotStale
from excelmanus.workspace.refs import WorkspaceRef


@pytest.fixture
def engine(tmp_path):
    wb = Workbook()
    wb.active.title = "明细"
    wb.active.append(["产品", "金额"])
    wb.active.append(["A", 12])
    wb.save(tmp_path / "sales.xlsx")
    registry = InteractionRegistry()
    engine = SimpleNamespace(
        _workspace_ref=WorkspaceRef.from_root(tmp_path, workspace_id="w1"),
        _question_flow=QuestionFlowManager(),
        _interaction_registry=registry, interaction_registry=registry,
        _driver=SimpleNamespace(running=True, _persist_runtime_state=lambda: None),
        _emit=lambda callback, event: callback(event) if callback else None,
    )
    engine._interaction_handler = InteractionHandler(engine)
    return engine


def target(engine):
    return bind_target(engine, {"file_path": "sales.xlsx", "sheet": "明细", "ranges": ["$A$2:$B$2"]})


def pending(engine):
    question = engine._question_flow.enqueue({"text": "请选择范围", "selection": target(engine)}, "call1")
    engine._interaction_handler._recovery = {
        "kind": "question", "phase": "waiting", "consumed": False,
        "tool_call_id": "call1", "questions": engine._question_flow.snapshot(), "answers": {},
    }
    return question


def test_bound_target_snapshot_and_event_survive_restore(engine):
    q = pending(engine)
    assert q.selection["ranges"] == ["A2:B2"]
    saved = engine._question_flow.snapshot()
    engine._question_flow.clear()
    engine._question_flow.restore(saved)
    assert engine._question_flow.current().selection == q.selection
    events = []
    engine._interaction_handler.emit_user_question_event(question=q, on_event=events.append, iteration=0)
    from excelmanus.api_sse import sse_event_to_sse
    event = sse_event_to_sse(events[0])
    data = json.loads(event.split("data: ", 1)[1])
    assert data["selection"] == q.selection
    assert data["tool_call_id"] == "call1"


@pytest.mark.parametrize("patch", [
    {"workspace_id": "other"}, {"file_path": "other.xlsx"}, {"sheet": "Other"},
    {"ranges": ["XFE1"]}, {"ranges": []}, {"ranges": ["Other!A1"]},
    {"ranges": ["A0"]}, {"content_version": ""},
])
def test_wrong_scope_or_invalid_selection_is_rejected(engine, patch):
    t = target(engine)
    with pytest.raises(ValueError):
        validate_selection_answer(engine, t, {**t, **patch})


def test_stale_selection_rejected_but_reselected_current_version_accepted(engine):
    old = target(engine)
    path = engine._workspace_ref.root / "sales.xlsx"
    wb = load_workbook(path)
    wb["明细"]["B2"] = 99
    wb.save(path)
    with pytest.raises(SnapshotStale):
        validate_selection_answer(engine, old, old)
    current = target(engine)
    assert validate_selection_answer(engine, old, current)["content_version"] == current["content_version"]


def test_bad_second_question_does_not_leave_partial_queue(engine):
    with pytest.raises(Exception):
        prepared = prepare_questions(engine, [{"text": "请选择", "selection": target(engine)}, {"text": "错误", "selection": {"file_path": "missing.xlsx", "sheet": "明细"}}])
        engine._question_flow.enqueue_batch(prepared, "batch")
    assert engine._question_flow.queue_size() == 0
    with pytest.raises(ValueError):
        engine._question_flow.enqueue_batch([{"text": "好", "options": [{"label": "A"}]}, {"text": "错"}], "batch")
    assert engine._question_flow.queue_size() == 0


def test_strict_provider_null_optional_fields_can_request_selection(engine):
    questions = prepare_questions(engine, [{"text": "选择区域", "options": None, "multiSelect": None,
        "selection": {"file_path": "sales.xlsx", "sheet": "明细", "ranges": None, "content_version": None}}])
    q = engine._question_flow.enqueue_batch(questions, "strict")[0]
    assert q.selection["ranges"] == [] and q.multi_select is False


@pytest.mark.asyncio
async def test_answer_endpoint_resolves_only_current_question_once(engine, monkeypatch):
    from excelmanus import api_routes_chat
    q = pending(engine)
    future = engine.interaction_registry.create(q.question_id)
    manager = SimpleNamespace(get_or_restore_engine=AsyncMock(return_value=engine))
    monkeypatch.setattr(api_routes_chat, "get_session_manager", lambda: manager)
    monkeypatch.setattr(api_routes_chat, "_has_session_access", AsyncMock(return_value=True))
    app = FastAPI()
    app.include_router(api_routes_chat.router)
    async with AsyncClient(transport=ASGITransport(app), base_url="http://test") as client:
        url = "/api/v1/chat/s1/answer"
        wrong = await client.post(url, json={"question_id": q.question_id, "selection": {**q.selection, "sheet": "Wrong"}})
        assert wrong.status_code == 422 and not future.done()
        stale = await client.post(url, json={"question_id": q.question_id, "selection": {**q.selection, "content_version": "sha256:outdated"}})
        assert stale.status_code == 409 and not future.done()
        body = {"question_id": q.question_id, "selection": {**q.selection, "ranges": ["B2"]}}
        accepted = await client.post(url, json=body)
        assert accepted.status_code == 200
        answer = await asyncio.wait_for(future, 1)
        assert answer["status"] == "confirmed" and answer["selection"]["ranges"] == ["B2"]
        assert (await client.post(url, json=body)).status_code == 200
        assert (await client.post(url, json={**body, "selection": q.selection})).status_code == 422
        assert engine._interaction_handler.snapshot()["answers"][q.question_id] == answer


@pytest.mark.asyncio
async def test_presentation_is_read_only_and_changed_requires_write_version(engine):
    handler = ShowWorkbookHandler(engine, None)
    path = engine._workspace_ref.root / "sales.xlsx"
    before = path.read_bytes()
    t = target(engine)
    for stage in ["planned", "changed"]:
        result = await handler.handle("show_workbook", "show1", {"target": t, "stage": stage, "summary": "处理金额"})
        payload = json.loads(result.result_str)
        assert payload["target"]["content_version"] == t["content_version"]
        assert payload["stage"] == stage
    assert path.read_bytes() == before
    with pytest.raises(ValueError, match="content_version"):
        await handler.handle("show_workbook", "show2", {"target": {"file_path": "sales.xlsx", "sheet": "明细", "ranges": ["B2"]}, "stage": "changed"})
