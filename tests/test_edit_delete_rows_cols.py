"""edit_spreadsheet：delete_rows / delete_columns（1-based at，与 insert 同风格）。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.intent_tools import edit_spreadsheet, get_tools, init_guard
from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    init_guard(workspace)
    seed_seen_versions({})


def _err(result: ToolResult) -> str:
    if result.error is not None:
        return f"{result.error.code} {result.error.message}"
    if isinstance(result.value, dict):
        return str(result.value)
    return result.model_text


def _grid_book(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "A1"
    ws["B1"] = "B1"
    ws["C1"] = "C1"
    ws["A2"] = "A2"
    ws["B2"] = "B2"
    ws["C2"] = "C2"
    ws["A3"] = "A3"
    ws["B3"] = "B3"
    ws["C3"] = "C3"
    wb.save(path)
    wb.close()
    return path


def test_delete_rows_shifts_below_up(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _grid_book(tmp_path / "rows.xlsx")
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "delete_rows",
            "sheet": "Sheet1",
            "at": 2,
            "count": 1,
        }],
    )
    assert result.success, _err(result)
    payload = result.value
    assert isinstance(payload, dict)
    applied = payload.get("applied") or []
    assert any("delete_rows@2" in str(item) for item in applied)
    wb = load_workbook(path)
    try:
        assert wb.active["A1"].value == "A1"
        assert wb.active["A2"].value == "A3"
        assert wb.active["B2"].value == "B3"
        assert wb.active["A3"].value is None
    finally:
        wb.close()


def test_delete_columns_shifts_right_left(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _grid_book(tmp_path / "cols.xlsx")
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "delete_columns",
            "sheet": "Sheet1",
            "at": "B",
            "count": 1,
        }],
    )
    assert result.success, _err(result)
    wb = load_workbook(path)
    try:
        assert wb.active["A1"].value == "A1"
        assert wb.active["B1"].value == "C1"
        assert wb.active["C1"].value is None
    finally:
        wb.close()


def test_delete_rows_count_deletes_multiple(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _grid_book(tmp_path / "multi.xlsx")
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "delete_rows",
            "sheet": "Sheet1",
            "at": 1,
            "count": 2,
        }],
    )
    assert result.success, _err(result)
    wb = load_workbook(path)
    try:
        assert wb.active["A1"].value == "A3"
        assert wb.active["A2"].value is None
    finally:
        wb.close()


def test_delete_rows_rewrites_formulas(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "formula.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 10
    ws["B1"] = "=A1"
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "delete_rows",
            "sheet": "Sheet1",
            "at": 1,
            "count": 1,
        }],
    )
    assert result.success, result.model_text
    reopened = load_workbook(path, data_only=False)
    assert reopened["Sheet1"]["B1"].value is None
    reopened.close()


def _sheet_delete(path: Path, sheet: str) -> ToolResult:
    return edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "sheet", "action": "delete", "sheet": sheet}],
    )


def test_sheet_delete_allows_unreferenced_with_same_sheet_formulas(tmp_path: Path) -> None:
    """R06 场景：新表写了同表 SUM 公式后删除无关源表应放行。"""
    _bind(tmp_path)
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws["A1"] = 1
    wb.create_sheet("汇总")["A1"] = "=SUM(B1:C1)"
    wb.save(path)
    wb.close()
    result = _sheet_delete(path, "订单")
    assert result.success, _err(result)
    loaded = load_workbook(path)
    try:
        assert loaded.sheetnames == ["汇总"]
        assert loaded["汇总"]["A1"].value == "=SUM(B1:C1)"
    finally:
        loaded.close()


def test_sheet_delete_refuses_when_formula_references_target(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws["A1"] = 1
    wb.create_sheet("汇总")["A1"] = "=订单!A1*2"
    wb.save(path)
    wb.close()
    result = _sheet_delete(path, "订单")
    assert not result.success
    assert "订单" in _err(result)


def test_sheet_delete_refuses_quoted_sheet_reference(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "My Data"
    ws["A1"] = 1
    wb.create_sheet("汇总")["A1"] = "='My Data'!A1"
    wb.save(path)
    wb.close()
    result = _sheet_delete(path, "My Data")
    assert not result.success


def test_sheet_delete_refuses_table_on_target(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws["A1"] = "h"
    ws["A2"] = 1
    from openpyxl.worksheet.table import Table

    ws.add_table(Table(displayName="T1", ref="A1:A2"))
    wb.create_sheet("其他")
    wb.save(path)
    wb.close()
    result = _sheet_delete(path, "订单")
    assert not result.success
    assert "表对象" in _err(result)


def test_sheet_delete_refuses_chart_series_reference(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    for i in range(1, 5):
        ws.cell(row=i, column=1, value=i)
    other = wb.create_sheet("图表")
    from openpyxl.chart import BarChart, Reference

    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=4))
    other.add_chart(chart, "B2")
    wb.save(path)
    wb.close()
    result = _sheet_delete(path, "订单")
    assert not result.success
    assert "图表" in _err(result)


def test_delete_kind_in_schema_with_additional_properties_false() -> None:
    tools = {tool.name: tool for tool in get_tools()}
    schema = tools["edit_spreadsheet"].input_schema
    items = schema["properties"]["operations"]["items"]
    assert items.get("additionalProperties") is False
    assert {
        "write", "insert", "sheet", "copy", "delete_rows", "delete_columns",
        "pivot", "transform",
    } <= set(items["properties"]["kind"]["enum"])
