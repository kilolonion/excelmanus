"""提示词重构配套契约：目录信封、诚实回读、合并填充、源码隔离、探测文件。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from excelmanus.prompt.canonical import TOOL_DESCRIPTIONS
from excelmanus.prompt.skill_catalog import (
    catalog_entries_digest,
    collect_skill_entries,
    render_available_skills,
)
from excelmanus.replica_spec import StyleClass, compile_replica_to_bytes, workbook_spec_to_replica
from excelmanus.replica_spec import (
    BorderSpec,
    FillSpec,
    MergedRange,
    SheetSpec,
    WorkbookSpec,
)
from excelmanus.security.guard import FileAccessGuard, SecurityViolationError
from excelmanus.security.source_isolation import (
    PROBE_FILE_FORBIDDEN,
    PRODUCT_SOURCE_FORBIDDEN,
    is_probe_path,
)
from excelmanus.tools.file_tools import copy_file, init_guard as init_file_guard
from excelmanus.tools.intent_tools import (
    _MODEL_CAPABILITIES,
    analyze_spreadsheet,
    edit_spreadsheet,
    format_spreadsheet,
    init_guard as init_intent_guard,
    inspect_spreadsheet,
)
from excelmanus.tools.shell_tools import init_guard as init_shell_guard, run_shell
from excelmanus.workbook.data import init_guard as init_data_guard
from excelmanus.workbook.sheets import init_guard as init_sheets_guard, list_sheets


def _bind(tmp_path: Path) -> Path:
    root = tmp_path
    init_intent_guard(str(root))
    init_file_guard(str(root))
    init_shell_guard(str(root))
    init_sheets_guard(str(root))
    init_data_guard(str(root))
    return root


def _payload(result) -> dict:
    return result.value if result.success else (result.error.fields if result.error else {})


def test_skill_catalog_uses_system_reminder_and_entry_digest() -> None:
    text = render_available_skills({"format_basic": type("S", (), {"description": "样式", "disable_model_invocation": False})()})
    assert text.startswith("<system-reminder>")
    assert "<available_skills>" in text
    assert "summaries only" in text
    assert "/name" in text
    assert "before taking task actions" not in text
    assert "call the `skill` tool" not in text
    entries = collect_skill_entries({"a": type("S", (), {"description": "x", "disable_model_invocation": False})()})
    first = catalog_entries_digest(entries)
    second = catalog_entries_digest(entries)
    assert first == second
    assert first != catalog_entries_digest([("a", "y")])


def test_meta_tools_use_names_not_chinese_catalog() -> None:
    from types import SimpleNamespace

    from excelmanus.engine_core.meta_tools import MetaToolBuilder

    engine = SimpleNamespace(
        _skill_router=SimpleNamespace(
            list_skill_names=lambda blocked_skillpacks=None: ["data_basic"],
            build_skill_catalog=lambda blocked_skillpacks=None: ("可用技能：\n- data_basic：分析", ["data_basic"]),
        ),
        _skill_resolver=SimpleNamespace(blocked_skillpacks=lambda: set()),
        _subagent_registry=SimpleNamespace(build_catalog=lambda: ("", ["explorer"])),
    )
    tools = MetaToolBuilder(engine).build_meta_tools()
    names = [item["function"]["name"] for item in tools]
    assert "skill" in names
    assert "activate_skill" not in names
    blob = str(tools)
    assert "可用技能：" not in blob
    manage = next(item for item in tools if item["function"]["name"] == "manage_skills")
    assert "适用场景" not in manage["function"]["description"]


def test_capabilities_notes_are_facts_not_playbooks() -> None:
    notes = "\n".join(_MODEL_CAPABILITIES["notes"])
    assert "截断" in notes
    assert "content_version" in notes
    assert "结束工具" in notes
    assert "rebase" not in notes.lower()
    assert "对照源图" not in notes
    assert "run_code" not in notes
    assert "并行" not in notes


def test_system_skill_texts_are_not_playbooks() -> None:
    root = Path(__file__).resolve().parent.parent / "excelmanus" / "skillpacks" / "system"
    forbidden = (
        "标准流程",
        "必须先",
        "默认优先",
        "不够大就退回",
        "普通任务不要走这里",
        "禁止分多次取交集",
        "不是从图片还原",
        "before taking task actions",
    )
    for skill_md in sorted(root.glob("*/SKILL.md")):
        text = skill_md.read_text(encoding="utf-8")
        for phrase in forbidden:
            assert phrase not in text, f"{skill_md.parent.name}: {phrase}"


def test_offer_download_description_is_canonical() -> None:
    from excelmanus.tools.file_tools import get_tools

    tools = {tool.name: tool.description for tool in get_tools()}
    assert tools["offer_download"] == TOOL_DESCRIPTIONS["offer_download"]
    assert "探测" in tools["offer_download"]


def test_overview_include_styles_merges_formulas(tmp_path: Path) -> None:
    _bind(tmp_path)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "收据"
    ws.merge_cells("A1:B1")
    ws["A1"] = "合计"
    ws["A2"] = "=B2"
    ws["A2"].fill = PatternFill("solid", fgColor="FFFF00")
    path = tmp_path / "form.xlsx"
    wb.save(path)

    result = inspect_spreadsheet(
        mode="overview",
        file_path=str(path),
        include=["styles", "merges", "formulas", "not_a_dim"],
    )
    assert result.success
    payload = _payload(result)
    sheet = payload["sheets"][0]
    assert "styles" in sheet
    assert sheet["merges"]["count"] >= 1
    assert sheet["formulas"]["count"] >= 1
    assert "include_warning" in payload
    assert "include_warning" in result.model_text or "未知的 include" in result.model_text


def test_range_include_declares_ignored(tmp_path: Path) -> None:
    _bind(tmp_path)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = 1
    path = tmp_path / "grid.xlsx"
    wb.save(path)
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        range="A1:A1",
        include=["styles"],
    )
    assert result.success
    assert "include_warning" in _payload(result)
    assert "忽略" in result.model_text


def test_analyze_form_does_not_emit_missing_data(tmp_path: Path) -> None:
    _bind(tmp_path)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.merge_cells("A1:D1")
    ws.merge_cells("A2:B2")
    ws.merge_cells("C2:D2")
    ws["A1"] = "收款收据"
    ws["A2"] = "备注"
    ws["C2"] = "日期"
    ws["A3"] = "编号"
    ws["B3"] = "001"
    path = tmp_path / "receipt.xlsx"
    wb.save(path)
    result = analyze_spreadsheet(mode="quality", file_path=str(path))
    assert result.success
    signals = _payload(result).get("quality_signals") or []
    types = {item.get("type") for item in signals}
    assert "missing_data" not in types
    assert "type_mixed" not in types
    assert "empty_column" not in types
    assert "版式表单" in result.model_text


def test_format_skips_merged_non_anchors(tmp_path: Path) -> None:
    _bind(tmp_path)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.merge_cells("A1:B1")
    ws["A1"] = "标题"
    path = tmp_path / "merged.xlsx"
    wb.save(path)
    listed = inspect_spreadsheet(mode="overview", file_path=str(path))
    version = _payload(listed)["content_version"]
    result = format_spreadsheet(
        file_path=str(path),
        expected_version=version,
        operations=[
            {
                "kind": "format",
                "range": "A1:B1",
                "fill": {"color": "FFCC00"},
            }
        ],
    )
    assert result.success
    payload = _payload(result)
    assert payload["applied"] == ["format:A1:B1"]
    assert "B1" in payload.get("skipped_merged_non_anchors", [])


def test_compiler_applies_border_and_rejects_bad_merge() -> None:
    spec = WorkbookSpec(
        sheets=[
            SheetSpec(
                name="S",
                dimensions={"rows": 2, "cols": 2},
                value_blocks=[{"start": "A1", "values": [["x", "y"]]}],
                styles={"box": StyleClass(border=BorderSpec(style="thin", color="000000"), fill=FillSpec(color="#EEEEEE"))},
                style_regions=[{"range": "A1:B1", "style_id": "box"}],
                merged_ranges=[MergedRange(range="A1:B1")],
            )
        ],
        uncertainties=[],
    )
    data, summary = compile_replica_to_bytes(workbook_spec_to_replica(spec))
    assert data[:2] == b"PK"
    assert summary["merges_applied"] == 1

    bad = WorkbookSpec(
        sheets=[
            SheetSpec(
                name="S",
                dimensions={"rows": 2, "cols": 2},
                value_blocks=[{"start": "A1", "values": [["x", "y"], ["z", "w"]]}],
                merged_ranges=[MergedRange(range="A1:A0")],
            )
        ],
        uncertainties=[],
    )
    try:
        compile_replica_to_bytes(workbook_spec_to_replica(bad))
        raised = False
    except ValueError as exc:
        raised = True
        assert "合并失败" in str(exc)
    assert raised


def test_guard_denies_product_source(tmp_path: Path) -> None:
    (tmp_path / "excelmanus").mkdir()
    (tmp_path / "excelmanus" / "replica_spec.py").write_text("x", encoding="utf-8")
    guard = FileAccessGuard(str(tmp_path))
    try:
        guard.resolve_and_validate("excelmanus/replica_spec.py")
        denied = False
    except SecurityViolationError as exc:
        denied = True
        assert PRODUCT_SOURCE_FORBIDDEN in str(exc)
    assert denied


def test_shell_denies_product_source(tmp_path: Path) -> None:
    (tmp_path / "excelmanus").mkdir()
    (tmp_path / "excelmanus" / "replica_spec.py").write_text("secret", encoding="utf-8")
    init_shell_guard(str(tmp_path))
    result = run_shell("cat excelmanus/replica_spec.py", workdir=str(tmp_path))
    payload = result.value or {}
    assert payload.get("status") == "blocked"
    assert PRODUCT_SOURCE_FORBIDDEN in str(payload.get("reason") or "")


def test_probe_paths_are_rejected(tmp_path: Path) -> None:
    _bind(tmp_path)
    src = tmp_path / "ok.xlsx"
    Workbook().save(src)
    assert is_probe_path("outputs/_probe_merge.xlsx")
    copied = copy_file("ok.xlsx", "outputs/_probe_merge.xlsx")
    assert not copied.success or (copied.value or {}).get("code") == PROBE_FILE_FORBIDDEN
    if copied.error:
        assert copied.error.code == PROBE_FILE_FORBIDDEN or PROBE_FILE_FORBIDDEN in str(copied.value)
    created = edit_spreadsheet(
        file_path="outputs/_probe.xlsx",
        workbook_spec={
            "sheets": [{"name": "S", "dimensions": {"rows": 1, "cols": 1}}],
            "uncertainties": [],
        },
    )
    assert not created.success
    assert created.error is not None
    assert created.error.code == PROBE_FILE_FORBIDDEN

    missing = edit_spreadsheet(
        file_path="",
        workbook_spec={
            "sheets": [{"name": "S", "dimensions": {"rows": 1, "cols": 1}}],
            "uncertainties": [],
        },
    )
    assert not missing.success
    assert missing.error is not None
    assert missing.error.code == "PATH_REQUIRED"


def test_list_sheets_include_warning_in_model_text(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "a.xlsx"
    Workbook().save(path)
    result = list_sheets(str(path), include=["nope"])
    assert result.success
    assert "include_warning" in (result.value or {})
    assert "未知的 include" in result.model_text
