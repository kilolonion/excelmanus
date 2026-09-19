"""RefCache 测试。"""
from __future__ import annotations

import time
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excelmanus.reference_graph.cache import RefCache
from excelmanus.reference_graph.models import WorkbookRefIndex


def _make_index(path: str = "test.xlsx") -> WorkbookRefIndex:
    return WorkbookRefIndex(
        file_path=path, sheets={}, cross_sheet_edges=[],
        external_refs=[], named_ranges={}, built_at=time.time(),
    )


class TestRefCache:
    def test_put_and_get(self) -> None:
        cache = RefCache()
        idx = _make_index("a.xlsx")
        cache.put_tier1("a.xlsx", idx)
        assert cache.get_tier1("a.xlsx") is idx

    def test_miss_returns_none(self) -> None:
        cache = RefCache()
        assert cache.get_tier1("missing.xlsx") is None

    def test_invalidate(self) -> None:
        cache = RefCache()
        cache.put_tier1("a.xlsx", _make_index())
        cache.invalidate("a.xlsx")
        assert cache.get_tier1("a.xlsx") is None

    def test_invalidate_all(self) -> None:
        cache = RefCache()
        cache.put_tier1("a.xlsx", _make_index("a.xlsx"))
        cache.put_tier1("b.xlsx", _make_index("b.xlsx"))
        cache.invalidate_all()
        assert cache.get_tier1("a.xlsx") is None
        assert cache.get_tier1("b.xlsx") is None

    def test_tier2_cache(self) -> None:
        cache = RefCache()
        from excelmanus.reference_graph.models import CellNode

        node = CellNode(sheet="S1", address="A1")
        cache.put_tier2("f.xlsx", "S1", "A1", node)
        assert cache.get_tier2("f.xlsx", "S1", "A1") is node

    def test_tier2_invalidate_by_file(self) -> None:
        cache = RefCache()
        from excelmanus.reference_graph.models import CellNode

        node = CellNode(sheet="S1", address="A1")
        cache.put_tier2("f.xlsx", "S1", "A1", node)
        cache.invalidate("f.xlsx")
        assert cache.get_tier2("f.xlsx", "S1", "A1") is None

    def test_missing_path_none_fingerprints_still_hit(self) -> None:
        """假路径两侧 fingerprint 均为 None 时仍命中，兼容现有测试。"""
        cache = RefCache()
        idx = _make_index("a.xlsx")
        cache.put_tier1("a.xlsx", idx)
        assert cache.get_tier1("a.xlsx") is idx

    def test_put_tier1_evicts_other_versions(self) -> None:
        cache = RefCache()
        old = _make_index("book.xlsx")
        new = _make_index("book.xlsx")
        cache.put_tier1("ws|book.xlsx|sha256:old", old)
        cache.put_tier1("ws|book.xlsx|sha256:new", new)
        assert cache.get_tier1("ws|book.xlsx|sha256:old") is None
        assert cache.get_tier1("ws|book.xlsx|sha256:new") is new
        assert "ws|book.xlsx|sha256:old" not in cache.all_tier1()


class TestSessionCacheIsolation:
    """contextvar 多会话隔离测试。"""

    def test_different_contexts_get_different_caches(self) -> None:
        import contextvars
        from excelmanus.reference_graph.cache import (
            get_session_cache,
            set_session_cache,
            reset_session_cache,
        )

        cache_a = RefCache()
        cache_b = RefCache()
        cache_a.put_tier1("a.xlsx", _make_index("a.xlsx"))

        token_a = set_session_cache(cache_a)
        assert get_session_cache() is cache_a
        assert get_session_cache().get_tier1("a.xlsx") is not None

        reset_session_cache(token_a)

        token_b = set_session_cache(cache_b)
        assert get_session_cache() is cache_b
        assert get_session_cache().get_tier1("a.xlsx") is None

        reset_session_cache(token_b)

    def test_same_context_reuses_cache(self) -> None:
        from excelmanus.reference_graph.cache import get_session_cache

        c1 = get_session_cache()
        c2 = get_session_cache()
        assert c1 is c2

    def test_set_and_reset(self) -> None:
        from excelmanus.reference_graph.cache import (
            get_session_cache,
            set_session_cache,
            reset_session_cache,
        )

        original = get_session_cache()
        custom = RefCache()
        token = set_session_cache(custom)
        assert get_session_cache() is custom
        reset_session_cache(token)
        assert get_session_cache() is original


class TestEnsureIndexContentIdentity:
    """外部改写 xlsx 后，_ensure_index 必须重扫而不是复用旧对象。"""

    def test_ensure_index_rescans_after_external_formula_change(self, tmp_path: Path) -> None:
        from excelmanus.reference_graph.cache import (
            RefCache,
            reset_session_cache,
            set_session_cache,
        )
        from excelmanus.tools.context import bind_workspace
        from excelmanus.tools.reference_tools import _ensure_index, init_guard

        bind_workspace(str(tmp_path))
        init_guard(str(tmp_path))
        path = tmp_path / "book.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "S1"
        ws["A1"] = "=1+1"
        wb.save(path)
        wb.close()

        cache = RefCache()
        token = set_session_cache(cache)
        try:
            first = _ensure_index("book.xlsx")
            assert sum(s.formula_count for s in first.sheets.values()) == 1

            edited = load_workbook(path)
            edited.active["B1"] = "=2+2"
            edited.save(path)
            edited.close()

            second = _ensure_index("book.xlsx")
            assert second is not first
            assert second.sheets["S1"].formula_count == 2
            assert sum(s.formula_count for s in second.sheets.values()) == 2
        finally:
            reset_session_cache(token)
