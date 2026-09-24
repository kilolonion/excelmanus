"""Style import/writeback and whole-file revision fidelity."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Color, Font, PatternFill, Side
from openpyxl.writer.theme import theme_xml

from excelmanus.tools._style_extract import extract_cell_style, resolve_color
from excelmanus.workbook.mutation import execute_operation
from excelmanus.workspace.file_service import WorkspaceFileService


@pytest.mark.parametrize("name,expected", [
    ("thin", 1), ("hair", 2), ("dotted", 3), ("dashed", 4),
    ("dashDot", 5), ("dashDotDot", 6), ("double", 7), ("medium", 8),
    ("mediumDashed", 9), ("mediumDashDot", 10), ("mediumDashDotDot", 11),
    ("slantDashDot", 12), ("thick", 13),
])
def test_border_matches_univer_and_survives_writeback(name, expected):
    ws = Workbook().active
    ws["A1"].border = Border(bottom=Side(style=name, color="000000"))
    style = extract_cell_style(ws["A1"])
    assert style["bd"]["b"] == {"s": expected, "cl": {"rgb": "#000000"}}
    execute_operation(ws.parent, {"kind":"format","sheet":ws.title,"range":"B1","border":{"bottom":{"style":name,"color":"000000"}}})
    assert ws["B1"].border.bottom.style == name


def test_styles_use_workbook_theme_and_full_indexed_palette():
    wb = Workbook()
    wb.loaded_theme = theme_xml.replace('val="4F81BD"', 'val="123456"').encode()
    ws = wb.active
    ws["A1"].font = Font(name="Calibri", size=11, color=Color(theme=0))
    ws["A1"].fill = PatternFill("solid", fgColor=Color(theme=4))
    ws["A1"].border = Border(bottom=Side(style="double", color=Color(indexed=40)))
    data = BytesIO()
    wb.save(data)
    data.seek(0)
    reopened = load_workbook(data)
    try:
        style = extract_cell_style(reopened.active["A1"])
        assert style["ff"] == "Calibri"
        assert style["fs"] == 11
        assert style["cl"]["rgb"] == "#FFFFFF"
        assert style["bg"]["rgb"] == "#123456"
        assert style["bd"]["b"]["cl"]["rgb"] == "#00CCFF"
        reopened.active["B1"].font = Font(color=Color(theme=2))
        reopened.active["C1"].font = Font(color=Color(theme=3))
        assert extract_cell_style(reopened.active["B1"])["cl"]["rgb"] == "#EEECE1"
        assert extract_cell_style(reopened.active["C1"])["cl"]["rgb"] == "#1F497D"
        assert resolve_color(Color(theme=2)) == "#EEECE1"
    finally:
        reopened.close()


def test_revision_restore_preserves_styles_merges_dimensions_and_formula(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Receipt"
    ws.merge_cells("A1:F1")
    ws["A1"].font = Font(name="Calibri", size=24, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="155A8A")
    ws["A1"].border = Border(bottom=Side(style="double", color="155A8A"))
    ws.column_dimensions["B"].width = 22
    ws.row_dimensions[1].height = 40
    ws["D6"], ws["E6"], ws["F6"] = 6, 128, "=D6*E6"
    ws["F6"].number_format = "#,##0.00"
    buffer = BytesIO()
    wb.save(buffer)
    original = buffer.getvalue()
    svc = WorkspaceFileService(tmp_path)
    created = svc.create("receipt.xlsx", original)
    svc.raise_if_failed(created)
    revision = svc.list_history("receipt.xlsx")[-1]
    ws.unmerge_cells("A1:F1")
    ws["A1"] = "changed"
    ws["A1"].fill = PatternFill()
    buffer = BytesIO()
    wb.save(buffer)
    updated = svc.update("receipt.xlsx", buffer.getvalue(), expected_version=created.primary_version())
    svc.raise_if_failed(updated)
    restored = svc.restore("receipt.xlsx", revision.id, expected_version=updated.primary_version())
    svc.raise_if_failed(restored)
    # Restore must preserve the original archive, not re-create its cell values.
    assert (tmp_path / "receipt.xlsx").read_bytes() == original
    wb = load_workbook(tmp_path / "receipt.xlsx")
    try:
        ws = wb.active
        assert str(next(iter(ws.merged_cells.ranges))) == "A1:F1"
        assert ws["A1"].fill.fgColor.rgb == "00155A8A"
        assert ws["A1"].font.size == 24
        assert ws["A1"].border.bottom.style == "double"
        assert ws.column_dimensions["B"].width == 22
        assert ws.row_dimensions[1].height == 40
        assert ws["F6"].value == "=D6*E6"
        assert ws["F6"].number_format == "#,##0.00"
    finally:
        wb.close()
