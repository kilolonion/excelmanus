"""D1：spill 外置与写后语义校验。"""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from excelmanus.engine_core.spill import (
    DEFAULT_SPILL_BYTE_THRESHOLD,
    DEFAULT_SPILL_CHAR_THRESHOLD,
    DEFAULT_SPILL_TOKEN_THRESHOLD,
    SpillStore,
    apply_spill,
    extract_spill_locator,
    is_spill_locator,
    project_for_wire,
    retrieve_spill,
    retrieve_spill_result,
    should_spill,
)
from excelmanus.engine_core.tool_result import ToolResult

_HOST_DRIVE = re.compile(r"^[A-Za-z]:\\")
_HOST_USERS = "/Users/"
_HOST_HOME = "/home/"


def _assert_no_host_path(handle: str) -> None:
    assert not _HOST_DRIVE.search(handle)
    assert _HOST_USERS not in handle
    assert _HOST_HOME not in handle
    assert "\\" not in handle


def _book(path: Path, cells: dict[str, object], *, sheet: str = "Sheet1") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for addr, value in cells.items():
        ws[addr] = value
    wb.save(path)
    wb.close()
    return path


class TestSpillThresholds:
    def test_small_result_not_spilled_at_boundary(self, tmp_path: Path) -> None:
        store = SpillStore(tmp_path)
        text = "x" * 10
        assert should_spill(text, char_threshold=10, byte_threshold=10_000, token_threshold=10_000) is False
        proj = project_for_wire(
            text,
            store=store,
            char_threshold=10,
            byte_threshold=10_000,
            token_threshold=10_000,
        )
        assert proj.spilled is False
        assert proj.locator is None
        assert proj.model_text == text
        assert list(tmp_path.joinpath(".excelmanus", "spill").glob("*")) == []

    def test_just_over_char_threshold_spills_with_opaque_handle(self, tmp_path: Path) -> None:
        store = SpillStore(tmp_path)
        text = "y" * 11
        assert should_spill(text, char_threshold=10, byte_threshold=10_000, token_threshold=10_000) is True
        proj = project_for_wire(
            text,
            store=store,
            char_threshold=10,
            byte_threshold=10_000,
            token_threshold=10_000,
        )
        assert proj.spilled is True
        assert proj.locator is not None
        assert is_spill_locator(proj.locator)
        _assert_no_host_path(str(proj.locator))
        _assert_no_host_path(proj.model_text)
        payload = json.loads(proj.model_text)
        assert payload["spilled"] is True
        assert payload["spill"] == proj.locator
        assert "preview" in payload
        assert str(tmp_path) not in proj.model_text
        # 磁盘文件名只有 hex
        files = list(store.directory.iterdir())
        assert len(files) == 1
        assert re.fullmatch(r"[0-9a-f]{64}", files[0].name)
        _assert_no_host_path(files[0].name)


class TestSpillRetrieve:
    def test_retrieve_matches_original(self, tmp_path: Path) -> None:
        store = SpillStore(tmp_path)
        original = "hello-spill\n第二行 " + ("z" * 50)
        locator = store.put(original)
        _assert_no_host_path(str(locator))
        assert retrieve_spill(locator, workspace_root=tmp_path) == original
        assert store.get(locator) == original

    def test_retrieve_tool_result_preserves_text(self, tmp_path: Path) -> None:
        store = SpillStore(tmp_path)
        original = json.dumps({"status": "success", "rows": list(range(8))}, ensure_ascii=False)
        locator = store.put(original)
        result = retrieve_spill_result(locator, workspace_root=tmp_path)
        assert result.success is True
        assert result.model_text == original
        assert result.coverage and result.coverage.get("spill_retrieve") is True

    def test_apply_spill_keeps_value_contract(self, tmp_path: Path) -> None:
        store = SpillStore(tmp_path)
        original = "w" * 200
        result = ToolResult(
            success=True,
            model_text=original,
            value={"status": "success", "file_path": "book.xlsx", "applied": ["A1"]},
        )
        spilled = apply_spill(
            result,
            store=store,
            char_threshold=50,
            byte_threshold=10_000,
            token_threshold=10_000,
        )
        assert spilled.model_text != original
        assert "spill:" in spilled.model_text
        assert spilled.value["status"] == "success"
        assert spilled.value["file_path"] == "book.xlsx"
        assert spilled.value["applied"] == ["A1"]
        assert spilled.value["meta"]["spill"]["locator"]
        locator = spilled.value["meta"]["spill"]["locator"]
        _assert_no_host_path(str(locator))
        assert store.get(locator) == original

    def test_extract_locator_from_file_path(self) -> None:
        locator = "spill:" + "ab" * 32
        assert is_spill_locator(locator)
        assert extract_spill_locator({"file_path": locator, "sheet": "Sheet1"}) == locator
        assert extract_spill_locator({"file_path": "book.xlsx"}) is None






def test_default_thresholds_documented() -> None:
    assert DEFAULT_SPILL_CHAR_THRESHOLD == 8000
    assert DEFAULT_SPILL_BYTE_THRESHOLD == 24_000
    assert DEFAULT_SPILL_TOKEN_THRESHOLD == 2000
    assert DEFAULT_SPILL_CHAR_THRESHOLD < 12000
