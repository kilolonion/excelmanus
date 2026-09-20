"""§6.4 pandas 工作区直读：真实 run_code 包装，不在 pytest 进程里直接 pandas 充当证据。"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from excelmanus.tools import code_tools


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    code_tools.init_guard(str(tmp_path))
    (tmp_path / "scripts" / "temp").mkdir(parents=True, exist_ok=True)
    return tmp_path


def _payload(result):
    if hasattr(result, "value") and isinstance(result.value, dict):
        return result.value
    return {}


def _book(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "项目"
    ws["B1"] = "金额"
    ws["A2"] = "租金"
    ws["B2"] = 1200
    wb.save(path)
    wb.close()
    return path


def _require_symlink_privilege(tmp_dir: Path) -> None:
    """Windows 上创建符号链接需要特权；无权限时跳过用例。"""
    target = tmp_dir / ".symlink_probe_target"
    link = tmp_dir / ".symlink_probe_link"
    target.touch()
    try:
        link.symlink_to(target)
    except OSError:
        pytest.skip("当前环境无创建符号链接权限")
    finally:
        link.unlink(missing_ok=True)
        target.unlink(missing_ok=True)


def test_run_code_pandas_reads_workspace_xlsx(workspace: Path) -> None:
    target = _book(workspace / "sales.xlsx")
    original = target.read_bytes()
    packed = code_tools.run_code(
        code=(
            "import pandas as pd\n"
            "df = pd.read_excel('sales.xlsx')\n"
            "print(int(df['金额'].sum()))\n"
        ),
        python_command=sys.executable,
        require_excel_deps=True,
        sandbox_tier="GREEN",
    )
    result = _payload(packed)
    assert packed.success, packed.model_text
    assert result.get("status") == "success"
    assert "1200" in str(result.get("stdout_tail") or packed.model_text)
    assert target.read_bytes() == original


def test_run_code_pandas_to_excel_does_not_write_workspace(workspace: Path) -> None:
    target = workspace / "out.xlsx"
    packed = code_tools.run_code(
        code=(
            "import pandas as pd\n"
            "pd.DataFrame({'a': [1]}).to_excel('out.xlsx', index=False)\n"
        ),
        python_command=sys.executable,
        require_excel_deps=True,
        sandbox_tier="GREEN",
    )
    assert packed.success is False
    text = str(packed.model_text or "") + str((_payload(packed).get("stderr_tail") or ""))
    assert "to_excel" in text or "edit_spreadsheet" in text or "format_spreadsheet" in text
    assert not target.exists()


def test_run_code_pandas_product_source_xlsx_denied(workspace: Path) -> None:
    product = workspace / "excelmanus"
    product.mkdir()
    trapped = _book(product / "secret.xlsx")
    packed = code_tools.run_code(
        code=(
            "import pandas as pd\n"
            "pd.read_excel('excelmanus/secret.xlsx')\n"
        ),
        python_command=sys.executable,
        require_excel_deps=True,
        sandbox_tier="GREEN",
    )
    assert packed.success is False
    assert trapped.exists()


def test_run_code_pandas_read_outside_workspace_denied(workspace: Path) -> None:
    """与 Native FileAccessGuard 同口径：区外数据文件在 GREEN 沙盒不可读。"""
    import tempfile

    outside = Path(tempfile.mkdtemp()) / "secret.xlsx"
    _book(outside)
    packed = code_tools.run_code(
        code=(
            "import pandas as pd\n"
            f"pd.read_excel(r'{outside}')\n"
        ),
        python_command=sys.executable,
        require_excel_deps=True,
        sandbox_tier="GREEN",
    )
    assert packed.success is False
    text = str(packed.model_text or "") + str(_payload(packed).get("stderr_tail") or "")
    assert "PATH_OUTSIDE_WORKSPACE" in text or "PermissionError" in text


def test_run_code_pandas_excelmanus_internal_denied(workspace: Path) -> None:
    """.excelmanus 内部文件（如状态库）在 GREEN 沙盒不可读。"""
    internal_dir = workspace / ".excelmanus"
    internal_dir.mkdir(exist_ok=True)
    internal = internal_dir / "excelmanus.db"
    internal.write_bytes(b"SQLITE")
    packed = code_tools.run_code(
        code=(
            f"open(r'{internal}', 'rb').read()\n"
            "print('should_not_reach')\n"
        ),
        python_command=sys.executable,
        require_excel_deps=True,
        sandbox_tier="GREEN",
    )
    assert packed.success is False
    text = str(packed.model_text or "") + str(_payload(packed).get("stderr_tail") or "")
    assert "SENSITIVE_FILE" in text or "RESERVED_NAMESPACE" in text or "PermissionError" in text
    assert internal.exists()


def test_run_code_pandas_symlink_escape_denied(workspace: Path) -> None:
    """工作区内指向外部的符号链接按真实解析路径拒绝。"""
    _require_symlink_privilege(workspace)
    import tempfile

    outside = Path(tempfile.mkdtemp()) / "secret.xlsx"
    _book(outside)
    link = workspace / "escape.xlsx"
    link.symlink_to(outside)
    packed = code_tools.run_code(
        code=(
            "import pandas as pd\n"
            "pd.read_excel('escape.xlsx')\n"
        ),
        python_command=sys.executable,
        require_excel_deps=True,
        sandbox_tier="GREEN",
    )
    assert packed.success is False
    text = str(packed.model_text or "") + str(_payload(packed).get("stderr_tail") or "")
    assert "PATH_OUTSIDE_WORKSPACE" in text or "PermissionError" in text


def test_run_code_openpyxl_load_workspace_still_allowed(workspace: Path) -> None:
    """正向：工作区内工作簿经 openpyxl 直读不受影响。"""
    target = _book(workspace / "data.xlsx")
    packed = code_tools.run_code(
        code=(
            "from openpyxl import load_workbook\n"
            "wb = load_workbook('data.xlsx')\n"
            "print(wb.active['B2'].value)\n"
        ),
        python_command=sys.executable,
        require_excel_deps=True,
        sandbox_tier="GREEN",
    )
    result = _payload(packed)
    assert packed.success, packed.model_text
    assert "1200" in str(result.get("stdout_tail") or packed.model_text)
