"""History boundaries: corruption, cleanup, identities, concurrency and HTTP reads."""

from __future__ import annotations

import asyncio
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from pathlib import Path
import threading
from types import SimpleNamespace
from unittest.mock import MagicMock
from urllib.parse import quote

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import pytest

from excelmanus import api_routes_workspace as routes
from excelmanus.workbook_commit import CommitError, content_version_of
from excelmanus.workspace.file_service import WorkspaceFileService
from excelmanus.workspace.revisions import RevisionIntegrityError, RevisionStore


@pytest.fixture
def service(tmp_path, monkeypatch):
    monkeypatch.setattr(routes, "get_config", lambda: SimpleNamespace(workspace_root=str(tmp_path), data_root=""))
    monkeypatch.setattr(routes, "get_session_manager", lambda: MagicMock())
    return WorkspaceFileService(tmp_path)


@pytest.mark.parametrize("corruption", ["array", "identity", "digest", "sequence", "label", "committed"])
def test_corrupt_record_does_not_hide_good_history_or_modify_live_file(service, corruption):
    service.create("book.xlsx", b"good")
    bad = service.checkpoint("book.xlsx", expected_version=content_version_of(b"good"), label="bad")
    path = service.store.root / service.store.path_key(bad.path) / "records" / f"{bad.id}.json"
    data = bad.to_json_dict()
    if corruption == "array":
        data = []
    elif corruption == "identity":
        data["id"] = "another-record"
    elif corruption == "digest":
        data["sha256"] = "../../outside"
    elif corruption == "sequence":
        data["sequence"] = -1
    elif corruption == "label":
        data["label"] = {"bad": "type"}
    else:
        data["committed"] = "false"
    path.write_text(json.dumps(data), encoding="utf-8")
    history = service.list_history("book.xlsx")
    assert history and all(r.id != bad.id for r in history)
    with pytest.raises((RevisionIntegrityError, CommitError)):
        service.restore("book.xlsx", bad.id, expected_version=content_version_of(b"good"))
    assert (service.root / "book.xlsx").read_bytes() == b"good"


def test_uncommitted_revision_cannot_be_restored(service):
    service.create("book.xlsx", b"live")
    hidden = service.store.add_record(path="book.xlsx", data=b"unpublished", reason="beforeRestore", transaction_id="pending", committed=False)
    with pytest.raises(CommitError) as error:
        service.restore("book.xlsx", hidden.id, expected_version=content_version_of(b"live"))
    assert error.value.code == "NOT_FOUND"
    assert (service.root / "book.xlsx").read_bytes() == b"live"


def test_prune_preserves_failed_deletion_and_pending_recovery_blobs(tmp_path, monkeypatch):
    store = RevisionStore(tmp_path)
    held = store.add_record(path="book.xlsx", data=b"held", reason="afterEdit", transaction_id="held")
    dropped = store.add_record(path="book.xlsx", data=b"drop", reason="afterEdit", transaction_id="drop")
    pending = store.add_record(path="book.xlsx", data=b"pending", reason="beforeRestore", transaction_id="pending", committed=False)
    checkpoint = store.checkpoint("book.xlsx", b"checkpoint")
    original = Path.unlink

    def locked_unlink(path, *args, **kwargs):
        if path.name == f"{held.id}.json":
            raise PermissionError("record is in use")
        return original(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", locked_unlink)
    assert store.prune("book.xlsx", keep=0) == 1
    for rec in (held, pending, checkpoint):
        assert store.get(rec.path, rec.id) == rec
        assert store.read_blob(rec.path, rec.sha256) is not None
    assert store.get(dropped.path, dropped.id) is None
    assert store.read_blob(dropped.path, dropped.sha256) is None


def test_recreated_file_gets_own_checkpoint_even_with_identical_bytes(service):
    version = content_version_of(b"same")
    service.create("slot.xlsx", b"same")
    old = service.checkpoint("slot.xlsx", expected_version=version, label="saved")
    service.delete("slot.xlsx", expected_version=version)
    service.create("slot.xlsx", b"same")
    new = service.checkpoint("slot.xlsx", expected_version=version, label="saved")
    assert new.id != old.id and new.lineage_id != old.lineage_id
    assert old.id not in {r.id for r in service.list_history("slot.xlsx")}
    assert service.checkpoint("slot.xlsx", expected_version=version, label="saved").id == new.id


def test_corrupt_new_lineage_does_not_reveal_previous_file_history(service):
    service.create("slot.xlsx", b"old")
    service.delete("slot.xlsx", expected_version=content_version_of(b"old"))
    service.create("slot.xlsx", b"new")
    new = service.list_history("slot.xlsx")
    for rec in new:
        (service.store.root / service.store.path_key(rec.path) / "records" / f"{rec.id}.json").write_text("[]")
    assert service.list_history("slot.xlsx") == []


def test_delete_checkpoint_keeps_shared_blob_and_automatic_history(service):
    service.create("book.xlsx", b"same")
    auto = service.list_history("book.xlsx")[0]
    checkpoint = service.checkpoint("book.xlsx", expected_version=content_version_of(b"same"))
    with pytest.raises(CommitError):
        service.delete_checkpoint("book.xlsx", auto.id)
    service.delete_checkpoint("book.xlsx", checkpoint.id)
    assert service.read_history("book.xlsx", auto.id)[1] == b"same"
    assert checkpoint.id not in {r.id for r in service.list_history("book.xlsx")}


def test_simultaneous_restores_have_exactly_one_winner(service):
    service.create("book.xlsx", b"first")
    first = service.list_history("book.xlsx")[-1]
    service.update("book.xlsx", b"second", expected_version=content_version_of(b"first"))
    second = service.list_history("book.xlsx")[-1]
    service.update("book.xlsx", b"live", expected_version=content_version_of(b"second"))
    barrier = threading.Barrier(2)

    def restore(revision):
        barrier.wait(timeout=5)
        try:
            return service.restore("book.xlsx", revision.id, expected_version=content_version_of(b"live")).state
        except CommitError as error:
            return error.code

    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(restore, [first, second]))
    assert sorted(results) == ["VERSION_CONFLICT", "committed"]
    assert (service.root / "book.xlsx").read_bytes() in (b"first", b"second")
    before = [r for r in service.list_history("book.xlsx") if r.reason == "beforeRestore"]
    assert len(before) == 1
    assert service.read_history("book.xlsx", before[0].id)[1] == b"live"


@pytest.mark.asyncio
async def test_preview_is_styled_immutable_and_can_select_other_sheet(service):
    book = Workbook()
    book.active.title = "收据"
    book.active["A1"] = "收款收据"
    book.active.merge_cells("A1:C1")
    book.active["A1"].fill = PatternFill("solid", fgColor="195D85")
    book.active.column_dimensions["A"].width = 25
    book.active.row_dimensions[1].height = 40
    book.create_sheet("明细")["B2"] = "=2*3"
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    raw = buffer.getvalue()
    service.create("收据.xlsx", raw)
    revision = service.list_history("收据.xlsx")[-1]
    service.update("收据.xlsx", b"different live bytes", expected_version=content_version_of(raw))
    before = list(service.list_history("收据.xlsx"))
    payload = json.loads((await routes.preview_revision(path="收据.xlsx", revision_id=revision.id)).body)
    win = payload["regions"][0]
    assert win["cells"]["1,1"]["s"]["bg"]["rgb"] == "#195D85"
    assert win["merges"] == [{"min_row": 1, "max_row": 1, "min_col": 1, "max_col": 3, "anchor":"A1"}]
    assert win["geometry"]["columns"][0]["native"] == 25 and win["geometry"]["rows"][0]["native"] == 40
    other = json.loads((await routes.preview_revision(path="收据.xlsx", revision_id=revision.id, sheet="明细")).body)
    assert other["regions"][0]["cells"]["2,2"]["f"] == "=2*3"
    assert service.list_history("收据.xlsx") == before
    assert (service.root / "收据.xlsx").read_bytes() == b"different live bytes"


@pytest.mark.asyncio
async def test_preview_rejects_unbounded_range_before_reading(service, monkeypatch):
    read = MagicMock(side_effect=AssertionError("must reject before reading blob"))
    monkeypatch.setattr(WorkspaceFileService, "read_history", read)
    with pytest.raises(HTTPException) as error:
        await routes.preview_revision(path="book.xlsx", revision_id="irrelevant", rect="A1:XFD1048576")
    assert error.value.status_code == 400
    read.assert_not_called()


@pytest.mark.asyncio
async def test_malformed_workbook_has_clear_preview_error(service):
    service.create("broken.xlsx", b"not a zip file")
    revision = service.list_history("broken.xlsx")[-1]
    with pytest.raises(HTTPException) as error:
        await routes.preview_revision(path="broken.xlsx", revision_id=revision.id)
    assert error.value.status_code == 422


@pytest.mark.asyncio
async def test_unicode_download_filename_and_macro_mime(service):
    path = "收款 收据.xlsm"
    service.create(path, b"original macro workbook bytes")
    revision = service.list_history(path)[-1]
    response = await routes.read_revision_content(path=path, revision_id=revision.id)
    assert response.headers["content-disposition"] == f"inline; filename*=UTF-8''{quote(path, safe='')}"
    assert response.media_type == "application/vnd.ms-excel.sheet.macroEnabled.12"
    assert b"".join([chunk async for chunk in response.body_iterator]) == b"original macro workbook bytes"


@pytest.mark.asyncio
async def test_word_revision_preview_uses_closed_temp_file(service):
    from docx import Document

    document = Document()
    document.add_paragraph("历史文档")
    buffer = BytesIO()
    document.save(buffer)
    service.create("文档.docx", buffer.getvalue())
    revision = service.list_history("文档.docx")[-1]
    response = await routes.preview_revision(path="文档.docx", revision_id=revision.id)
    assert json.loads(response.body)["paragraphs"][0]["text"] == "历史文档"


def test_explicit_session_never_falls_back_to_another_workspace(service, monkeypatch, tmp_path):
    manager = MagicMock()
    monkeypatch.setattr(routes, "get_session_manager", lambda: manager)
    manager.session_exists.return_value = False
    with pytest.raises(HTTPException) as error:
        routes._workspace_root("missing")
    assert error.value.status_code == 404
    manager.session_exists.return_value = True
    manager.workspace_path_for_session.return_value = str(tmp_path)
    manager.resolve_workspace_binding.return_value = (str(tmp_path / "other"), "other")
    with pytest.raises(HTTPException) as error:
        routes._workspace_root("session", "other")
    assert error.value.status_code == 409


@pytest.mark.asyncio
@pytest.mark.parametrize("method,helper,args", [
    ("list_revisions", "_list_revisions", {"path": "book.xlsx"}),
    ("preview_revision", "_preview_revision", {"path": "book.xlsx", "revision_id": "old"}),
])
async def test_history_io_does_not_block_event_loop(monkeypatch, method, helper, args):
    release = threading.Event()
    entered = threading.Event()

    def slow_io(*_args):
        entered.set()
        assert release.wait(3), "history I/O blocked the event loop"
        return "done"

    monkeypatch.setattr(routes, helper, slow_io)
    task = asyncio.create_task(getattr(routes, method)(**args))
    try:
        for _ in range(100):
            if entered.is_set():
                break
            await asyncio.sleep(0.005)
        assert entered.is_set() and not task.done()
    finally:
        release.set()
    assert await task == "done"


def _audited_edit(service, paths=("a.xlsx",)):
    from excelmanus.approval import ApprovalManager

    for path in paths:
        service.create(path, b"before")
    manager = ApprovalManager(str(service.root))
    # This synthetic multi-file operation supplies its explicit audit targets.
    manager._resolve_target_paths = lambda *_: [service.root / path for path in paths]

    def execute(*_args):
        for path in paths:
            service.update(path, b"after", expected_version=content_version_of(b"before"))
        return "ok"

    _, record = manager.execute_and_audit(approval_id=manager.new_approval_id(), tool_name="run_code",
        arguments={"code": "# audited edit"}, tool_scope=["run_code"], execute=execute, undoable=True,
        created_at_utc=manager.utc_now())
    assert {change.path for change in record.changes} == set(paths)
    return manager, record


def test_operation_undo_matches_exact_snapshot_pair_and_is_repeat_safe(service):
    manager, record = _audited_edit(service)
    assert manager.undo(record.approval_id).startswith("已回滚 `")
    assert (service.root / "a.xlsx").read_bytes() == b"before"
    assert record.undoable is False
    assert "不支持自动回滚" in manager.undo(record.approval_id)


def test_operation_undo_does_not_guess_latest_transaction_when_record_is_missing(service):
    manager, record = _audited_edit(service)
    # Simulate an old approval whose corresponding afterEdit was pruned.
    for rec in service.store.list("a.xlsx"):
        if rec.reason == "afterEdit" and rec.sha256 == content_version_of(b"after").removeprefix("sha256:"):
            service.store.delete_record(rec)
    service.update("a.xlsx", b"later", expected_version=content_version_of(b"after"))
    assert manager.undo(record.approval_id).startswith("未回滚")
    assert (service.root / "a.xlsx").read_bytes() == b"later"
    assert record.undoable is True


def test_multifile_undo_preflights_all_versions_before_writing_any_file(service):
    manager, record = _audited_edit(service, ("a.xlsx", "b.xlsx"))
    service.update("b.xlsx", b"human edit", expected_version=content_version_of(b"after"))
    assert manager.undo(record.approval_id).startswith("未回滚")
    assert (service.root / "a.xlsx").read_bytes() == b"after"
    assert (service.root / "b.xlsx").read_bytes() == b"human edit"
    assert record.undoable is True
