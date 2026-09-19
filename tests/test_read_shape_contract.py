"""C2：读类成功载荷统一形状、并集/命名/表引用、schema 由 refs 生成。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.intent_tools import get_tools, init_guard, inspect_spreadsheet
from excelmanus.workbook.data import init_guard as init_data_guard, read_excel
from excelmanus.workbook.refs import describe_for_schema, reference_examples
from excelmanus.workbook_commit import seed_seen_versions


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    init_guard(workspace)
    init_data_guard(workspace)
    seed_seen_versions({})


def _err(result: ToolResult) -> str:
    if result.error is not None:
        return f"{result.error.code} {result.error.message}"
    if isinstance(result.value, dict):
        return str(result.value)
    return result.model_text


def _payload(result: ToolResult) -> dict:
    assert isinstance(result.value, dict), _err(result)
    return result.value


def _book(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "部门"
    ws["B1"] = "金额"
    ws["A2"] = "销售"
    ws["B2"] = 10
    ws["A3"] = "研发"
    ws["B3"] = 20
    ws["C3"] = "尾"
    ws["D4"] = 8
    wb.defined_names.add(DefinedName(name="MyName", attr_text="Sheet1!$A$1:$B$2"))
    ws.add_table(Table(displayName="Table1", ref="A1:B3"))
    wb.save(path)
    wb.close()
    return path


def test_range_read_has_unified_status_and_fields(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "range.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="A1:B2",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    assert payload["status"] == "success"
    assert payload.get("range") == "A1:B2"
    assert payload.get("values") == [["部门", "金额"], ["销售", 10]]
    assert payload.get("data") == payload["values"]
    meta = payload.get("meta") or {}
    assert meta.get("sheet") == "Sheet1"
    assert meta.get("kind") == "range"
    assert meta.get("truncated") is False
    assert meta.get("formulas_uncached") in (False, "unknown")
    assert "header_row" in meta
    assert "sampled" in meta


def test_overview_read_has_unified_status_and_fields(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "overview.xlsx")
    result = read_excel(file_path=str(path), sheet_name="Sheet1")
    assert result.success, _err(result)
    payload = _payload(result)
    assert payload["status"] == "success"
    assert payload.get("preview")
    assert payload.get("values")
    meta = payload.get("meta") or {}
    assert meta.get("kind") == "overview"
    assert meta.get("sheet") == "Sheet1"
    assert meta.get("sampled") is False
    assert meta.get("formulas_uncached") == "unknown"


def test_multi_area_read_covers_all_regions(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "union.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="A1:A2,C3:D4",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    assert payload["status"] == "success"
    areas = payload.get("areas") or []
    assert len(areas) == 2
    left = areas[0].get("values") or areas[0].get("data")
    right = areas[1].get("values") or areas[1].get("data")
    assert left == [["部门"], ["销售"]]
    assert right[0][0] == "尾"
    assert right[-1][-1] == 8
    values = payload.get("values")
    assert values == [left, right]
    assert "A1:A2" in str(payload.get("range"))
    assert "C3:D4" in str(payload.get("range"))


def test_named_range_binds_workbook_metadata(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "named.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="MyName",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    assert payload["status"] == "success"
    assert payload.get("values") == [["部门", "金额"], ["销售", 10]]
    assert payload.get("range") == "A1:B2"


def test_table_ref_binds_workbook_metadata(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "table.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="Table1[#All]",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    assert payload["status"] == "success"
    assert payload.get("values") == [
        ["部门", "金额"],
        ["销售", 10],
        ["研发", 20],
    ]


def test_named_range_missing_uses_named_range_not_found(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "named_miss.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="MissingName",
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code == "NAMED_RANGE_NOT_FOUND"


def test_table_missing_uses_table_not_found(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "table_miss.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="NoSuchTable[列]",
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code == "TABLE_NOT_FOUND"


def test_schema_range_description_is_generated_from_refs() -> None:
    """schema 内 range 说明是短版（common_errors）；完整语法由 introspect 按需查询。"""
    tools = {tool.name: tool for tool in get_tools()}
    desc = str(tools["inspect_spreadsheet"].input_schema["properties"]["range"]["description"])
    hint = describe_for_schema()
    assert hint["common_errors"] in desc
    # 完整语法不经 schema 内嵌，由 introspect_capability(tool_detail=...range) 按需给出
    import excelmanus.tools.introspection_tools as _intro

    class _StubCatalog:
        def introspection_source(self):
            return tools

    token = _intro._call_catalog.set(_StubCatalog())
    try:
        detail = _intro._handle_tool_detail("inspect_spreadsheet.range")
    finally:
        _intro._call_catalog.reset(token)
    assert hint["syntax"] in detail
    assert hint["examples"] in detail


def test_invalid_ref_error_maps_to_range_invalid(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "badref.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="=A1",
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code == "RANGE_INVALID"
    assert "等号" in _err(result) or "A1" in _err(result)


def test_versions_list_is_declared_non_mutating() -> None:
    from excelmanus.tools.policy import has_readonly_action, write_effect_for_call

    tools = {tool.name: tool for tool in get_tools()}
    assert tools["manage_spreadsheet_versions"].write_effect == "workspace_write"
    assert has_readonly_action("manage_spreadsheet_versions")
    assert write_effect_for_call(
        "manage_spreadsheet_versions",
        {"action": "list"},
        declared="workspace_write",
    ) == "none"
