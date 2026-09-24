"""读取性能辅助模块的合同测试：CSV 记录索引、读缓存合并、工作区快速扫描。"""

from __future__ import annotations

import threading
from pathlib import Path

import pytest
from openpyxl import Workbook

from excelmanus.workbook.csv_index import csv_index
from excelmanus.workbook.read_cache import ReadCache
from excelmanus.workbook.refs import parse_rect
from excelmanus.workbook.snapshot import open_snapshot_at
from excelmanus.workbook.observation import observe_snapshot
from excelmanus.workspace.listing import scan_workspace
from excelmanus.workspace.refs import WorkspaceRef


def _snapshot(tmp_path: Path, name: str, data: bytes):
    path = tmp_path / name
    path.write_bytes(data)
    return open_snapshot_at(path, relative=name, workspace=WorkspaceRef.from_root(tmp_path))


def _xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1
    ws["B1"] = "x"
    wb.save(path)
    wb.close()
    return path


class TestCsvIndex:
    def test_quoted_newline_counts_as_one_record(self, tmp_path: Path) -> None:
        snap = _snapshot(
            tmp_path, "rows.csv",
            'h1,h2\n"line1\nline2",x\nlast,y\n'.encode("utf-8"),
        )
        index = csv_index(snap)
        assert index.rows == 3
        rows = list(index.window(2, 2))
        assert rows == [["line1\nline2", "x"]]

    def test_index_is_reused_for_the_same_content_version(self, tmp_path: Path) -> None:
        snap = _snapshot(tmp_path, "rows.csv", b"a,b\n1,2\n")
        first = csv_index(snap)
        again = csv_index(_snapshot(tmp_path, "rows.csv", b"a,b\n1,2\n"))
        assert again is first

    def test_content_change_produces_a_new_index(self, tmp_path: Path) -> None:
        snap = _snapshot(tmp_path, "rows.csv", b"a,b\n1,2\n")
        first = csv_index(snap)
        changed = csv_index(_snapshot(tmp_path, "rows.csv", b"a,b\n1,2\n3,4\n"))
        assert changed is not first
        assert changed.rows == 3

    def test_view_window_reads_only_requested_rows(self, tmp_path: Path) -> None:
        text = "h1,h2\n" + "".join(f"{i},v{i}\n" for i in range(1, 500))
        snap = _snapshot(tmp_path, "rows.csv", text.encode("utf-8"))
        rect = parse_rect("A400:B402")
        view = observe_snapshot(snap, range=rect.to_a1(include_sheet=False))
        cells = view["regions"][0]["cells"]
        assert cells["400,1"]["v"] == "399"
        assert cells["402,2"]["v"] == "v401"
        assert view["sheets"][0]["used"]["rows"] == 500


class TestReadCache:
    def test_concurrent_misses_share_one_build(self) -> None:
        cache: ReadCache[str] = ReadCache(max_bytes=1024)
        builds = 0
        gate = threading.Event()

        def build() -> tuple[str, int]:
            nonlocal builds
            builds += 1
            assert gate.wait(2)
            return "value", 8

        results: list[str] = []

        def worker() -> None:
            results.append(cache.get_or_create("k", build))

        threads = [threading.Thread(target=worker) for _ in range(4)]
        for t in threads:
            t.start()
        gate.set()
        for t in threads:
            t.join(2)
        assert builds == 1
        assert results == ["value"] * 4

    def test_eviction_respects_byte_budget(self) -> None:
        cache: ReadCache[str] = ReadCache(max_bytes=16, max_entries=8)
        cache.get_or_create("a", lambda: ("x" * 8, 8))
        cache.get_or_create("b", lambda: ("y" * 8, 8))
        cache.get_or_create("c", lambda: ("z" * 8, 8))
        assert "a" not in cache._entries
        assert "c" in cache._entries


class TestWorkbookPairCache:
    def test_observation_reuses_parsed_workbooks_per_version(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        _xlsx(tmp_path / "book.xlsx")
        snap = open_snapshot_at(
            tmp_path / "book.xlsx", relative="book.xlsx",
            workspace=WorkspaceRef.from_root(tmp_path),
        )
        loads = 0
        import openpyxl

        real_load = openpyxl.load_workbook

        def counted(*args, **kwargs):
            nonlocal loads
            loads += 1
            return real_load(*args, **kwargs)

        monkeypatch.setattr(openpyxl, "load_workbook", counted)
        observe_snapshot(snap, range="A1:B5", facets=["data"])
        assert loads == 2
        observe_snapshot(snap, range="C1:D5", facets=["data"])
        assert loads == 2


class TestScanWorkspace:
    def test_filters_internals_and_returns_display_names(self, tmp_path: Path) -> None:
        (tmp_path / "uploads").mkdir()
        (tmp_path / "uploads" / "abcd1234_sales.xlsx").write_bytes(b"x")
        (tmp_path / "uploads" / "plain.csv").write_text("a,b", encoding="utf-8")
        (tmp_path / "outputs" / "backups").mkdir(parents=True)
        (tmp_path / "outputs" / "backups" / "x_20260911T091344_f525.xlsx").write_bytes(b"b")
        (tmp_path / ".excelmanus").mkdir()
        (tmp_path / ".excelmanus" / "state.db").write_bytes(b"s")
        (tmp_path / "_rc_12345678.py").write_text("print(1)", encoding="utf-8")
        (tmp_path / ".hidden").write_text("h", encoding="utf-8")
        (tmp_path / "notes").mkdir()
        (tmp_path / "notes" / "memo.txt").write_text("n", encoding="utf-8")

        files, truncated = scan_workspace(tmp_path)
        by_path = {f["path"]: f for f in files}
        assert not truncated
        assert "uploads" in by_path and by_path["uploads"]["is_dir"]
        assert by_path["uploads/abcd1234_sales.xlsx"]["filename"] == "sales.xlsx"
        assert "notes/memo.txt" in by_path
        for path in by_path:
            assert "outputs/backups" not in path
            assert ".excelmanus" not in path
            assert not Path(path).name.startswith("_rc_")
            assert ".hidden" not in path

    def test_limit_marks_result_truncated(self, tmp_path: Path) -> None:
        for i in range(5):
            (tmp_path / f"f{i}.txt").write_text("x", encoding="utf-8")
        files, truncated = scan_workspace(tmp_path, limit=3)
        assert truncated
        assert len(files) == 3

    def test_excel_only_keeps_traversing_dirs(self, tmp_path: Path) -> None:
        (tmp_path / "sub").mkdir()
        (tmp_path / "sub" / "data.csv").write_text("a,b", encoding="utf-8")
        (tmp_path / "readme.txt").write_text("x", encoding="utf-8")
        files, _ = scan_workspace(tmp_path, excel_only=True)
        assert [f["path"] for f in files] == ["sub/data.csv"]

    def test_symlink_dirs_are_not_followed(self, tmp_path: Path) -> None:
        outside = tmp_path.parent / f"{tmp_path.name}_outside"
        outside.mkdir(exist_ok=True)
        (outside / "secret.txt").write_text("s", encoding="utf-8")
        try:
            (tmp_path / "link").symlink_to(outside, target_is_directory=True)
        except (OSError, NotImplementedError):
            pytest.skip("当前平台不允许创建符号链接")
        files, _ = scan_workspace(tmp_path)
        assert all(not f["path"].startswith("link") for f in files)
