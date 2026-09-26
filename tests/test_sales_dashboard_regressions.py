"""Regressions extracted from the 2026-09-24 monthly sales dashboard trace."""
from io import BytesIO
import os
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
import pytest
import xlsxwriter

from excelmanus.tools.context import use_workspace
from excelmanus.tools.spreadsheet_data_tools import validate_spreadsheet
from excelmanus.tools.spreadsheet_engine_tools import calculate_spreadsheet, render_spreadsheet
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes
from excelmanus.workbook.ooxml import MAIN, NS, WORKBOOK, merge_formula_caches, render_selection, rewrite, serialize, xml
from excelmanus.workbook_commit import content_version_of_file


def dashboard(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "经营看板"
    ws["A1"] = "经营看板"
    ws.print_area = "A1:H30"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    data = wb.create_sheet("数据")
    data.append(["月份", "2024", "2025"])
    data.append(["1月", 10, "=B2*1.2"])
    data.append(["2月", 20, "=B3*1.2"])
    data.print_area = "A1:C3"
    wb.save(path)
    wb.close()


def cached_book():
    stream = BytesIO()
    with xlsxwriter.Workbook(stream) as wb:
        wb.add_worksheet("经营看板").write("A1", "engine changed this title")
        ws = wb.add_worksheet("数据")
        ws.write_row("A1", ["月份", "2024", "2025"])
        ws.write_row("A2", ["1月", 10])
        ws.write_row("A3", ["2月", 20])
        ws.write_formula("C2", "=B2*1.2", None, 12)
        ws.write_formula("C3", "=B3*1.2", None, 24)
    return stream.getvalue()


def test_recalculation_copies_only_caches_and_preserves_package(tmp_path):
    path = tmp_path / "book.xlsx"
    dashboard(path)
    original = path.read_bytes()
    output, count = merge_formula_caches(original, cached_book())
    assert count == 2
    with ZipFile(BytesIO(original)) as src, ZipFile(BytesIO(output)) as dst:
        assert src.namelist() == dst.namelist()
        assert [n for n in src.namelist() if src.read(n) != dst.read(n)] == ["xl/worksheets/sheet2.xml"]
    wb = load_workbook(BytesIO(output), data_only=True)
    assert wb["经营看板"]["A1"].value == "经营看板"
    assert [wb["数据"][f"C{r}"].value for r in (2, 3)] == [12, 24]
    wb.close()
    wb = load_workbook(BytesIO(output), data_only=False)
    assert wb["数据"]["C2"].value == "=B2*1.2"
    wb.close()


@pytest.mark.parametrize("value,kind,expected", [("华东", "s", "华东"), ("", "str", ""), ("1", "b", True), ("#DIV/0!", "e", "#DIV/0!")])
def test_formula_result_types_are_preserved(tmp_path, value, kind, expected):
    path = tmp_path / "book.xlsx"
    dashboard(path)
    raw = cached_book()
    with ZipFile(BytesIO(raw)) as package:
        root = xml(package.read("xl/worksheets/sheet2.xml"))
        cell = root.find(".//s:c[@r='C2']", NS)
        cell.set("t", kind)
        cell.find("s:v", NS).text = "0" if kind == "s" else value
        updates = {"xl/worksheets/sheet2.xml": serialize(root)}
        if kind == "s":
            updates["xl/sharedStrings.xml"] = f'<sst xmlns="{MAIN}"><si><t>{value}</t></si></sst>'.encode()
        converted = rewrite(package, updates)
    output, _ = merge_formula_caches(path.read_bytes(), converted)
    wb = load_workbook(BytesIO(output), data_only=True)
    # openpyxl represents an explicitly cached empty string as None.
    assert wb["数据"]["C2"].value == (None if expected == "" else expected)
    wb.close()


def test_missing_cache_is_rejected_instead_of_claiming_success(tmp_path):
    path = tmp_path / "book.xlsx"
    dashboard(path)
    with pytest.raises(ValueError, match="Missing formula cache"):
        merge_formula_caches(path.read_bytes(), path.read_bytes())


def test_existing_array_caches_are_updated_without_rewriting_the_formula():
    def array_book(values):
        stream = BytesIO()
        with xlsxwriter.Workbook(stream) as wb:
            ws = wb.add_worksheet("Array")
            ws.write_array_formula("A1:A2", "{=ROW(A1:A2)}", None, values[0])
            ws.write_number("A2", values[1])
        return stream.getvalue()
    output, count = merge_formula_caches(array_book([0, 0]), array_book([1, 2]))
    wb = load_workbook(BytesIO(output), data_only=True)
    assert [wb.active[c].value for c in ("A1", "A2")] == [1, 2]
    assert count == 1
    wb.close()


def test_changed_array_spill_range_fails_before_publication(tmp_path):
    path = tmp_path / "book.xlsx"
    dashboard(path)
    with ZipFile(BytesIO(cached_book())) as package:
        root = xml(package.read("xl/worksheets/sheet2.xml"))
        formula = root.find(".//s:c[@r='C2']/s:f", NS)
        formula.set("t", "array")
        formula.set("ref", "C2:C4")
        converted = rewrite(package, {"xl/worksheets/sheet2.xml": serialize(root)})
    with pytest.raises(ValueError, match="spill range"):
        merge_formula_caches(path.read_bytes(), converted)


def test_render_selection_keeps_all_non_workbook_parts_and_only_one_print_area(tmp_path):
    path = tmp_path / "book.xlsx"
    dashboard(path)
    raw = path.read_bytes()
    rendered = render_selection(raw, "经营看板", "A3:H30")
    with ZipFile(BytesIO(raw)) as src, ZipFile(BytesIO(rendered)) as dst:
        assert [n for n in src.namelist() if src.read(n) != dst.read(n)] == [WORKBOOK]
        names = xml(dst.read(WORKBOOK)).findall("s:definedNames/s:definedName", NS)
        areas = [n for n in names if n.get("name") == "_xlnm.Print_Area"]
        assert len(areas) == 1
        assert areas[0].text == "'经营看板'!$A$3:$H$30"
    assert path.read_bytes() == raw


def test_workbook_formula_validation_and_composite_unique_key(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "看板"
    wb.active["A1"] = "#DIV/0!"
    data = wb.create_sheet("明细")
    for row in [["年份", "月份", "区域"], [2024, 1, "华东"], [2025, 1, "华东"]]:
        data.append(row)
    wb.save(path)
    wb.close()
    with use_workspace(tmp_path):
        result = validate_spreadsheet(path.name, [{"kind": "formula_errors"}, {"kind": "unique", "sheet": "明细", "columns": ["年份", "月份", "区域"]}])
        assert result.success, result.model_text
        assert result.value["valid"] is False
        assert result.value["failure_count"] == 1
        assert result.value["failures"][0]["sheet"] == "看板"
        assert result.value["rules"][-1]["failed"] == 0
        ambiguous = validate_spreadsheet(path.name, [{"kind": "unique", "column": "区域"}])
        assert ambiguous.error.code == "SHEET_REQUIRED"
        assert ambiguous.value["available_sheets"] == ["看板", "明细"]


def test_delete_sdk_accepts_path_without_bypassing_confirmation(tmp_path):
    from excelmanus.code_mode import render_sdk_source
    from excelmanus.tools.file_tools import get_tools
    from excelmanus.tools.registry import ToolRegistry

    registry = ToolRegistry()
    tool = next(t for t in get_tools() if t.name == "delete_file")
    registry.register_tool(tool)
    sdk = {}
    exec(render_sdk_source([tool]), sdk)
    sdk["_call_host"] = lambda name, args: registry.call_tool(name, {k: v for k, v in args.items() if v is not sdk["_EM_UNSET"]}).value
    path = tmp_path / "report.txt"
    path.write_text("keep me")
    with use_workspace(tmp_path):
        result = sdk["delete_file"](path="report.txt")
    assert result["status"] == "confirmation_required"
    assert path.exists()
    with pytest.raises(TypeError, match="同时给出"):
        sdk["delete_file"](file_path="a.txt", path="b.txt")


@pytest.mark.skipif(os.name == "nt", reason="POSIX resource limits")
@pytest.mark.parametrize("allow_subprocess", [False, True])
def test_resource_limits_match_subprocess_authorization(monkeypatch, allow_subprocess):
    import resource
    from excelmanus.tools.code_tools import _build_unix_limits_preexec
    calls = []
    monkeypatch.setattr(resource, "getrlimit", lambda code: (resource.RLIM_INFINITY, resource.RLIM_INFINITY))
    monkeypatch.setattr(resource, "setrlimit", lambda code, values: calls.append(code))
    hook, enabled, _ = _build_unix_limits_preexec(30, allow_subprocess=allow_subprocess)
    assert enabled
    hook()
    assert (resource.RLIMIT_NPROC in calls) is not allow_subprocess
    assert resource.RLIMIT_CPU in calls


def test_page_limit_returns_actual_count_and_recovery(tmp_path, monkeypatch):
    from excelmanus.tools import spreadsheet_engine_tools as engine
    path = tmp_path / "book.xlsx"
    dashboard(path)
    def convert(source, destination, extension):
        pdf = destination / "input.pdf"
        pdf.write_bytes(b"pdf")
        return pdf, "soffice"
    monkeypatch.setattr(engine, "_office_convert", convert)
    monkeypatch.setattr(engine.shutil, "which", lambda name: name)
    monkeypatch.setattr(engine, "_run", lambda *args: "Pages: 7")
    with use_workspace(tmp_path):
        result = render_spreadsheet(path.name, "charts.png", format="png", max_pages=3)
    assert not result.success
    assert result.error.code == "LIMIT_EXCEEDED"
    assert result.value["suggested_max_pages"] == 7
    assert "preview_spreadsheet" in result.value["remediation"]
    assert not (tmp_path / "charts.png").exists()


def test_create_calculate_edit_render_with_real_office(tmp_path, monkeypatch):
    from excelmanus.runtime_capabilities import office_executable
    import shutil
    if not office_executable() or not shutil.which("pdfinfo"):
        pytest.skip("LibreOffice and pdfinfo required")
    monkeypatch.setenv("EXCELMANUS_FORMULA_RECALC", "auto")
    path = tmp_path / "book.xlsx"
    dashboard(path)
    with use_workspace(tmp_path):
        result = apply_spreadsheet_changes(path.name, expected_version=content_version_of_file(path), operations=[{
            "kind": "chart", "sheet": "数据", "target_sheet": "经营看板", "target_cell": "A3",
            "chart_type": "line", "data_range": "A1:C3", "categories_range": "A2:A3",
        }])
        assert result.success, result.model_text
        before = path.read_bytes()
        calc = calculate_spreadsheet(path.name, expected_version=content_version_of_file(path))
        assert calc.success, calc.model_text
        with ZipFile(BytesIO(before)) as a, ZipFile(path) as b:
            assert a.read("xl/charts/chart1.xml") == b.read("xl/charts/chart1.xml")
        wb = load_workbook(path, data_only=True)
        assert wb["数据"]["C3"].value == 24
        assert len(wb["经营看板"]._charts[0].series) == 2
        wb.close()
        edit = apply_spreadsheet_changes(path.name, expected_version=calc.value["content_version"], operations=[{"kind": "print_layout", "sheet": "经营看板", "print_layout": {"print_area": "A1:H31"}}])
        assert edit.success, edit.model_text
        pdf = render_spreadsheet(path.name, "dashboard.pdf", sheet="经营看板", range="A1:H31", max_pages=1)
        assert pdf.success, pdf.model_text
        assert pdf.value["page_count"] == 1
