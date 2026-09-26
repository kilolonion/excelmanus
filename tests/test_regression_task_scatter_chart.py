"""回归测试：原生散点图系列名、点样式默认值、行列方向与类别数量校验。

背景（会话 3，日志 (3)）：调用方按合同传入含表头的 data_range（B1:C21）+
categories_range（A2:A21），但图例最终显示 `Column B / Column C`，且散点序列没有
显式点样式，Excel 与 LibreOffice 各自套用默认值。这里锁定修复后的行为。
"""

from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties

from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.workbook.charts import init_guard as init_chart_guard
from excelmanus.tools.workbook_tools import init_guard, apply_spreadsheet_changes
from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions


@pytest.fixture(autouse=True)
def _bind_workspace(tmp_path: Path) -> None:
    workspace = str(tmp_path)
    set_guard(FileAccessGuard(workspace))
    init_guard(workspace)
    init_chart_guard(workspace)
    seed_seen_versions({})


def _payload(result: ToolResult) -> dict:
    assert isinstance(result, ToolResult)
    assert isinstance(result.value, dict)
    return result.value


def _make_regression_data(tmp_path: Path, name: str = "regression.xlsx") -> Path:
    """20 行广告投入/销售额 + C 列拟合值，复刻本次会话的数据形状。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws["A1"], ws["B1"], ws["C1"] = "广告投入(万元)", "销售额(万元)", "预测销售额(万元)"
    for i in range(20):
        x = 2.5 + 0.5 * i
        y = 18.3 + 3.5 * i
        ws.cell(row=i + 2, column=1, value=x)
        ws.cell(row=i + 2, column=2, value=y)
        ws.cell(row=i + 2, column=3, value=round(1.4338 + 6.7071 * x, 4))
    fp = tmp_path / name
    wb.save(fp)
    wb.close()
    return fp


def _apply(path: Path, **payload: object) -> ToolResult:
    return apply_spreadsheet_changes(
        file_path=str(path),
        operations=[{"sheet": "数据", **payload}],
        expected_version=content_version_of_file(path),
    )


def _tx_ref(series) -> str:
    label = series.tx
    assert isinstance(label, SeriesLabel), f"序列缺少显式标题: {label!r}"
    return str(getattr(getattr(label, "strRef", None), "f", "") or "")


class TestScatterSeriesNames:
    def test_series_titles_come_from_header_row(self, tmp_path: Path) -> None:
        fp = _make_regression_data(tmp_path)
        result = _apply(
            fp, kind="chart", chart_type="scatter",
            data_range="B1:C21", categories_range="A2:A21", target_cell="F2",
        )
        assert result.success is True, result.model_text
        wb = load_workbook(fp)
        charts = wb["数据"]._charts
        assert len(charts) == 1
        series = charts[0].series
        assert len(series) == 2, "B、C 两列应各成为一个序列"
        refs = [_tx_ref(item) for item in series]
        assert refs == ["'数据'!$B$1", "'数据'!$C$1"], refs
        wb.close()

    def test_series_titles_survive_serialization(self, tmp_path: Path) -> None:
        fp = _make_regression_data(tmp_path)
        _apply(fp, kind="chart", chart_type="scatter",
               data_range="B1:C21", categories_range="A2:A21", target_cell="F2")
        with ZipFile(fp) as z:
            xml = z.read("xl/charts/chart1.xml").decode("utf-8")
        assert "Column B" not in xml and "Column C" not in xml
        assert "'数据'!$B$1" in xml and "'数据'!$C$1" in xml
        assert len(re.findall(r"<tx>", xml)) == 2, xml[:2000]
        assert '<a:noFill/>' in xml and '<symbol val="circle"/>' in xml


class TestScatterPointStyle:
    def test_default_is_marker_only_not_connected(self, tmp_path: Path) -> None:
        fp = _make_regression_data(tmp_path)
        _apply(fp, kind="chart", chart_type="scatter",
               data_range="B1:C21", categories_range="A2:A21", target_cell="F2")
        wb = load_workbook(fp)
        for series in wb["数据"]._charts[0].series:
            assert series.marker is not None and series.marker.symbol == "circle"
            line = getattr(series.graphicalProperties, "ln", None)
            assert line is not None and line.noFill is True, "散点默认不应连线"
        wb.close()

    def test_explicit_series_style_is_preserved_on_update(self, tmp_path: Path) -> None:
        fp = _make_regression_data(tmp_path)
        _apply(fp, kind="chart", chart_type="scatter",
               data_range="B1:C21", categories_range="A2:A21", target_cell="F2")
        # 用户显式改过点样式（方块）
        wb = load_workbook(fp)
        for series in wb["数据"]._charts[0].series:
            series.marker.symbol = "square"
        wb.save(fp)
        wb.close()
        _apply(fp, kind="update_chart", chart_type="scatter", index=0,
               data_range="B1:C21", categories_range="A2:A21")
        wb = load_workbook(fp)
        for series in wb["数据"]._charts[0].series:
            assert series.marker.symbol == "square", "显式样式必须保留"
        wb.close()

    def test_legacy_default_style_is_repaired_on_update(self, tmp_path: Path) -> None:
        """旧版本写出的 symbol=none + 实线是历史默认值，更新时不应被复制回来。"""
        fp = _make_regression_data(tmp_path)
        _apply(fp, kind="chart", chart_type="scatter",
               data_range="B1:C21", categories_range="A2:A21", target_cell="F2")
        wb = load_workbook(fp)
        from openpyxl.chart.marker import Marker
        from openpyxl.drawing.line import LineProperties

        for series in wb["数据"]._charts[0].series:
            series.marker = Marker(symbol="none")
            series.graphicalProperties = GraphicalProperties(ln=LineProperties())
        wb.save(fp)
        wb.close()
        _apply(fp, kind="update_chart", chart_type="scatter", index=0,
               data_range="B1:C21", categories_range="A2:A21")
        wb = load_workbook(fp)
        for series in wb["数据"]._charts[0].series:
            assert series.marker.symbol == "circle", "历史默认 marker=none 不应被继承"
        wb.close()


class TestScatterOrientationAndShape:
    def test_from_rows_orients_series_by_row(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "数据"
        ws.append(["月份", "实际", "预测"])
        for i, month in enumerate(["1月", "2月", "3月"], start=2):
            ws.append([month, 100 + i, 90 + i])
        fp = tmp_path / "rows.xlsx"
        wb.save(fp)
        wb.close()
        result = _apply(fp, kind="chart", chart_type="scatter",
                        data_range="A1:C4", from_rows=True, target_cell="E1")
        assert result.success is True, result.model_text
        wb = load_workbook(fp)
        series = wb["数据"]._charts[0].series
        assert len(series) == 3, "每一行（2..4）应成为一个序列"
        assert [_tx_ref(item) for item in series] == ["'数据'!$A$2", "'数据'!$A$3", "'数据'!$A$4"]
        wb.close()

    def test_mismatched_categories_count_is_rejected(self, tmp_path: Path) -> None:
        fp = _make_regression_data(tmp_path)
        result = _apply(fp, kind="chart", chart_type="scatter",
                        data_range="B1:C21", categories_range="A2:A10", target_cell="F2")
        assert result.success is False
        assert "数量必须一致" in (result.model_text or "")
        # 拒绝的图表不能落盘
        wb = load_workbook(fp)
        assert wb["数据"]._charts == []
        wb.close()

    def test_wrong_orientation_categories_is_rejected(self, tmp_path: Path) -> None:
        fp = _make_regression_data(tmp_path)
        result = _apply(fp, kind="chart", chart_type="scatter",
                        data_range="B1:C21", categories_range="A2:C2", target_cell="F2")
        assert result.success is False
        assert "必须是单列" in (result.model_text or "")
