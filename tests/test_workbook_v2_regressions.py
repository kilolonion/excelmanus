"""Boundary regressions for the complete workbook/2 path (no model mocks)."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import ColorScaleRule
from excelmanus.tools.context import use_workspace
from excelmanus.tools.workbook_tools import (
    observe_spreadsheet,
    apply_spreadsheet_changes,
    preview_spreadsheet,
    split_spreadsheet,
)
from excelmanus.workbook_commit import content_version_of_file


@pytest.fixture
def book(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["item", "amount"])
    ws.append(["apple", 2])
    ws.append(["apple", 3])
    ws.append(["pear", 4])
    wb.save(path)
    wb.close()
    with use_workspace(tmp_path):
        yield path


def change(book, operations, **kwargs):
    return apply_spreadsheet_changes(
        file_path=book.name,
        expected_version=content_version_of_file(book),
        operations=operations,
        **kwargs,
    )


def test_search_cursor_covers_all_matches_from_same_version(book):
    first = observe_spreadsheet(
        file_path=book.name, mode="search", query="apple", limit=1
    )
    assert first.success
    second = observe_spreadsheet(
        file_path=book.name,
        mode="search",
        query="apple",
        limit=1,
        offset=first.value["coverage"]["next_offset"],
        expected_version=first.value["content_version"],
    )
    assert [
        first.value["matches"][0]["cell_ref"],
        second.value["matches"][0]["cell_ref"],
    ] == ["A2", "A3"]
    assert second.value["coverage"]["next_offset"] is None
    assert first.value["request"]["mode"] == "search"


def test_csv_respects_unrequested_facets_and_sparse_shape(book):
    (book.parent / "table.csv").write_text("a,b\n1,2\n")
    result = observe_spreadsheet(
        file_path="table.csv", facets=["geometry"], mode="range", range="A1:B2"
    )
    assert result.value["regions"][0]["cells"] == {}
    coverage = result.value["regions"][0]["coverage"]
    assert coverage["data"]["status"] == "not_requested"
    assert coverage["geometry"]["status"] == "unsupported"


def test_disjoint_coverage_does_not_claim_gap_is_loaded(book):
    result = observe_spreadsheet(file_path=book.name, range="A1:B1,A3:B4", mode="range")
    assert result.success, result.model_text
    assert {"sheet": "S", "r0": 2, "r1": 2, "c0": 1, "c1": 2} in result.value[
        "coverage"
    ]["unloaded"]


def test_dry_run_same_compiler_leaves_bytes_and_history_unchanged(book):
    before = book.read_bytes()
    ops = [
        {"kind": "geometry.scale", "sheet": "S", "range": "A1:B4", "x": 1.2, "y": 0.8}
    ]
    result = change(book, ops, dry_run=True)
    assert result.success and not result.value["committed"]
    assert result.value["files"][0]["observation"]["geometry_changes"][0]["verified"]
    assert book.read_bytes() == before
    assert not list((book.parent / ".excelmanus/revisions").glob("**/manifest.json"))
    assert change(book, ops).value["committed"]


def test_multiple_resizes_verify_final_composed_dimensions(book):
    result = change(
        book,
        [
            {"kind": "size", "sheet": "S", "column_widths": {"A": 20}},
            {"kind": "size", "sheet": "S", "column_widths": {"A": 30}},
        ],
    )
    assert result.success, result.model_text
    entries = result.value["observation"]["geometry_changes"]
    assert all(e["verified"] for e in entries)
    assert entries[0]["dimensions_requested"]["column_widths"]["A"] == 20
    assert entries[0]["dimensions_persisted"]["columns"][0]["native"] == 30


def test_preserve_outside_rejects_shared_axis_and_keeps_file(book):
    before = book.read_bytes()
    op = {
        "kind": "geometry.scale",
        "sheet": "S",
        "range": "A1:B2",
        "x": 1.2,
        "y": 0.8,
        "preserve_outside": True,
    }
    result = change(book, [op])
    assert not result.success and result.error.code == "LAYOUT_SCOPE_CONFLICT"
    assert book.read_bytes() == before
    result = change(book, [{**op, "preserve_outside": False}])
    assert result.success
    assert result.value["observation"]["operations"][0]["impact"]["outside_count"] == 4


def test_pixel_resize_and_grouped_native_resize_share_geometry(book):
    wb = load_workbook(book)
    ws = wb["S"]
    ws.column_dimensions.group("A", "C", hidden=False)
    ws.column_dimensions["A"].width = 24
    wb.save(book)
    wb.close()
    result = change(
        book,
        [
            {
                "kind": "geometry.resize",
                "sheet": "S",
                "axis": "column",
                "sizes": {"2": 145},
                "unit": "css_px",
            }
        ],
    )
    assert result.success, result.model_text
    observed = observe_spreadsheet(
        file_path=book.name, sheet="S", range="A1:C4", mode="range", facets=["geometry"]
    )
    cols = observed.value["regions"][0]["geometry"]["columns"]
    assert [c["native"] for c in cols] == [24, 20, 24]
    assert cols[1]["pixels"] == 145


def test_cells_patch_style_cannot_redirect_write(book):
    before = book.read_bytes()
    result = change(
        book,
        [
            {
                "kind": "cells.patch",
                "sheet": "S",
                "cells": [{"cell": "A1", "style": {"kind": "clear", "range": "A1:B4"}}],
            }
        ],
    )
    assert not result.success and book.read_bytes() == before


def test_probe_creation_rejected(book):
    result = apply_spreadsheet_changes(
        file_path="outputs/_probe.xlsx",
        create=True,
        operations=[
            {"kind": "write", "sheet": "S", "start_cell": "A1", "values": [[1]]}
        ],
    )
    assert not result.success and not (book.parent / "outputs/_probe.xlsx").exists()


def test_formula_provenance_and_rule_xml_survive_observation(book):
    wb = load_workbook(book)
    ws = wb["S"]
    ws["C2"] = "=B2*10"
    ws.conditional_formatting.add(
        "B2:B4",
        ColorScaleRule(
            start_type="min", start_color="FF0000", end_type="max", end_color="00FF00"
        ),
    )
    wb.save(book)
    wb.close()
    result = observe_spreadsheet(
        file_path=book.name,
        range="A1:C4",
        mode="range",
        facets=["data", "presentation"],
    )
    region = result.value["regions"][0]
    assert region["cells"]["2,3"]["value_source"] == "formula_text"
    assert region["cells"]["2,3"]["calculation_provenance"]["source"] == "missing_cache"
    assert "FF0000" in region["conditional_formatting"]["rules"][0]["raw_xml"]
    assert region["conditional_formatting"]["rules"][0]["evaluated"] is False


def test_manifest_not_mutable_and_dependencies_do_not_extend_bounds(book):
    result = observe_spreadsheet(file_path=book.name)
    result.value["sheets"][0]["used"]["rows"] = 99999
    dep = observe_spreadsheet(
        file_path=book.name, mode="dependencies", range="A801:B805"
    )
    assert dep.success
    assert (
        observe_spreadsheet(file_path=book.name).value["sheets"][0]["used"]["rows"] == 4
    )


def test_unknown_ooxml_parts_are_observable_but_not_overwritten(book):
    original = book.read_bytes()
    out = BytesIO()
    with ZipFile(BytesIO(original)) as src, ZipFile(out, "w") as dest:
        for name in src.namelist():
            dest.writestr(name, src.read(name))
        dest.writestr("customXml/item1.xml", "<important>keep me</important>")
    book.write_bytes(out.getvalue())
    before = book.read_bytes()
    observed = observe_spreadsheet(file_path=book.name)
    assert (
        observed.success
        and "customXml/item1.xml" in observed.value["capabilities"]["unknown_parts"]
    )
    result = change(
        book, [{"kind": "write", "sheet": "S", "start_cell": "A1", "values": [["new"]]}]
    )
    assert not result.success and result.error.code == "UNSUPPORTED_PRESERVATION"
    assert book.read_bytes() == before


def test_spec_builds_chart_atomically(book):
    spec = {
        "sheets": [
            {
                "name": "S",
                "dimensions": {"rows": 3, "cols": 2},
                "value_blocks": [
                    {"start": "A1", "values": [["x", "y"], ["a", 1], ["b", 2]]}
                ],
                "objects": {
                    "charts": [
                        {
                            "chart_type": "bar",
                            "data_range": "B1:B3",
                            "target_cell": "D1",
                        }
                    ]
                },
            }
        ],
        "uncertainties": [],
    }
    result = apply_spreadsheet_changes(file_path="chart.xlsx", workbook_spec=spec)
    assert result.success, result.model_text
    wb = load_workbook(book.parent / "chart.xlsx")
    assert len(wb["S"]._charts) == 1
    wb.close()


def test_split_schema_reaches_domain(book):
    result = split_spreadsheet(
        file_path=book.name, sheet="S", by_column="item", header_row=1
    )
    assert result.success, result.model_text
    assert len(result.value["files"]) == 2


def test_preview_coordinates_display_and_cache(book, monkeypatch):
    import excelmanus.workbook.preview as preview
    from excelmanus.attachments.store import reset_attachment_store

    reset_attachment_store()
    first = preview_spreadsheet(file_path=book.name, sheet="S", range="A1:B4")
    assert first.success, first.model_text
    value = first.value
    assert value["cell_to_pixel_map"]["columns"][0]["x"] == 0
    assert value["source_pixel_size"]["width"] > 0
    assert value["measured"]["displayed_cells"]["2,1"]["text"] == "apple"
    assert value["font_fingerprint"] and value["renderer_version"]
    monkeypatch.setattr(
        preview,
        "_run",
        lambda *a, **kw: pytest.fail("cache hit must not start a renderer"),
    )
    first.value["geometry"]["width_px"] = -1
    second = preview_spreadsheet(file_path=book.name, sheet="S", range="A1:B4")
    assert second.success and second.value["geometry"]["width_px"] > 0
    assert second.value["render_id"] == value["render_id"]


def test_generated_types_are_current():
    from scripts.generate_workbook_contracts import render

    root = Path(__file__).resolve().parents[1]
    assert (
        root / "web/src/lib/workbook-contracts.generated.ts"
    ).read_text() == render()


def test_small_scale_never_reverses_native_direction(book):
    before = observe_spreadsheet(
        file_path=book.name, sheet="S", range="A1:B4", mode="range"
    ).value["regions"][0]["geometry"]
    result = change(
        book,
        [
            {
                "kind": "geometry.scale",
                "sheet": "S",
                "range": "A1:B4",
                "x": 1.00001,
                "y": 0.99999,
            }
        ],
    )
    assert result.success, result.model_text
    after = result.value["observation"]["geometry_changes"][0]["persisted"]
    assert after["columns"][0]["native"] > before["columns"][0]["native"]
    assert after["rows"][0]["native"] < before["rows"][0]["native"]


def test_analysis_refuses_missing_formula_cache(book):
    from excelmanus.tools.workbook_tools import analyze_spreadsheet

    wb = load_workbook(book)
    wb["S"]["B2"] = "=1+2"
    wb.save(book)
    wb.close()
    result = analyze_spreadsheet(
        file_path=book.name, sheet="S", mode="aggregate", aggregations={"amount": "sum"}
    )
    assert not result.success and result.error.code == "FORMULA_CACHE_MISSING"
    profile = analyze_spreadsheet(file_path=book.name, sheet="S", mode="quality")
    assert profile.success
    column = next(
        c for c in profile.value["sheets"][0]["columns"] if c["name"] == "amount"
    )
    assert column["statistics_status"] == "unavailable"
    assert "null_count" not in column


def test_blank_presentation_uses_inherited_column_style_without_mutating_snapshot(book):
    from excelmanus.workbook.snapshot import open_snapshot, _cached_workbook_pair

    wb = load_workbook(book)
    wb["S"].column_dimensions["B"].fill = PatternFill("solid", fgColor="AABBCC")
    wb.save(book)
    wb.close()
    snapshot = open_snapshot(book.name)
    parsed, _, _ = _cached_workbook_pair(snapshot, True)
    count = len(parsed["S"]._cells)
    result = observe_spreadsheet(
        file_path=book.name,
        sheet="S",
        range="A801:B805",
        mode="range",
        facets=["presentation"],
    )
    assert result.value["regions"][0]["cells"]["801,2"]["s"]["bg"]["rgb"] == "#AABBCC"
    assert len(parsed["S"]._cells) == count


def test_preview_marks_uncalculated_formula_without_filling_offscreen_inputs(book):
    wb = load_workbook(book)
    ws = wb["S"]
    ws["B1"] = "=A100+A101"
    ws["A100"] = 100
    ws["A101"] = 200
    wb.save(book)
    wb.close()
    result = preview_spreadsheet(file_path=book.name, sheet="S", range="A1:B1")
    assert result.success, result.model_text
    display = result.value["measured"]["displayed_cells"]["1,2"]
    assert display["status"] == "uncalculated"
    assert display["text"] == "⟨未计算⟩"


def test_delete_last_comment_and_hyperlink_keeps_receipt_truthful(book):
    from openpyxl.comments import Comment

    wb = load_workbook(book)
    wb["S"]["A1"].comment = Comment("note", "author")
    wb["S"]["B1"].hyperlink = "https://example.com"
    wb.save(book)
    wb.close()
    result = change(
        book,
        [
            {"kind": "comment", "sheet": "S", "cell": "A1", "action": "delete"},
            {"kind": "hyperlink", "sheet": "S", "cell": "B1", "action": "delete"},
        ],
    )
    assert result.success, result.model_text
    wb = load_workbook(book)
    try:
        assert wb["S"]["A1"].comment is None and wb["S"]["B1"].hyperlink is None
    finally:
        wb.close()


def test_resize_loaded_image_and_chart_changes_serialized_anchors(book):
    from PIL import Image

    Image.new("RGB", (24, 12), "red").save(book.parent / "image.png")
    created = change(
        book,
        [
            {
                "kind": "image",
                "sheet": "S",
                "image_path": "image.png",
                "target_cell": "D1",
            },
            {
                "kind": "chart",
                "sheet": "S",
                "chart_type": "bar",
                "data_range": "B1:B4",
                "target_cell": "D8",
            },
        ],
    )
    assert created.success, created.model_text
    result = change(
        book,
        [
            {
                "kind": "image",
                "sheet": "S",
                "action": "resize",
                "index": 0,
                "width": 120,
                "height": 60,
            },
            {
                "kind": "update_chart",
                "sheet": "S",
                "index": 0,
                "width": 20,
                "height": 12,
            },
        ],
    )
    assert result.success, result.model_text
    wb = load_workbook(book)
    try:
        assert wb["S"]._images[0].anchor.ext.cx == 120 * 9525
        assert wb["S"]._images[0].anchor.ext.cy == 60 * 9525
        assert wb["S"]._charts[0].anchor.ext.cx == 20 * 360000
        assert wb["S"]._charts[0].anchor.ext.cy == 12 * 360000
    finally:
        wb.close()


def test_array_formula_observation_is_json_and_keeps_cache_state(book):
    from openpyxl.worksheet.formula import ArrayFormula

    wb = load_workbook(book)
    wb["S"]["C1"] = ArrayFormula(ref="C1:C4", text="=B1:B4*2")
    wb.save(book)
    wb.close()
    observed = observe_spreadsheet(
        file_path=book.name,
        sheet="S",
        range="A1:C4",
        mode="range",
        facets=["data", "dependencies"],
    )
    assert observed.success, observed.model_text
    cell = observed.value["regions"][0]["cells"]["1,3"]
    assert cell["f"] == "=B1:B4*2" and cell["cached"] == "no" and cell["v"] is None
    assert cell["formula_metadata"] == {"t": "array", "ref": "C1:C4"}
    json.dumps(observed.value)
