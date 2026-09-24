from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from excelmanus.workbook.merge import review_merge
from excelmanus.workbook_commit import content_version_of, content_version_of_file
from excelmanus.workspace.file_service import WorkspaceFileService
from tests.test_api import _make_transport, _setup_api_globals, _test_config


def workbook(**values):
    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    for cell, value in {"A1": 1, "A2": 2, "B1": "=A1+1", "B2": "keep"}.items():
        ws[cell] = value
    for cell, value in values.items():
        ws[cell] = value
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def ops(**values):
    return [{"kind": "cells.patch", "sheet": "销售", "cells": [{"cell": cell, "value": value} for cell, value in values.items()]}]


def test_merges_non_overlapping_values_and_formulas():
    base, remote = workbook(), workbook(A2=20)
    review, merged = review_merge(base, remote, ops(A1=10, B1="=SUM(A1:A2)"), apply=True)
    assert review["conflict_count"] == 0 and review["safe_count"] == 2
    wb = load_workbook(BytesIO(merged))
    assert [wb.active[cell].value for cell in ("A1", "A2", "B1", "B2")] == [10, 20, "=SUM(A1:A2)", "keep"]
    wb.close()


@pytest.mark.parametrize("choice,expected", [("local", 10), ("remote", 30)])
def test_conflicts_require_choices_and_preserve_unrelated_edits(choice, expected):
    base, remote = workbook(), workbook(A1=30, B2="agent edit")
    review, unchanged = review_merge(base, remote, ops(A1=10, A2=20))
    assert unchanged is None and review["conflict_count"] == 1
    conflict = next(cell for cell in review["cells"] if cell["conflict"])
    assert (conflict["base"], conflict["local"], conflict["remote"]) == (1, 10, 30)
    with pytest.raises(ValueError, match="每个冲突"):
        review_merge(base, remote, ops(A1=10, A2=20), apply=True)
    _, merged = review_merge(base, remote, ops(A1=10, A2=20), choices={conflict["id"]: choice}, apply=True)
    wb = load_workbook(BytesIO(merged))
    assert wb.active["A1"].value == expected
    assert wb.active["A2"].value == 20 and wb.active["B2"].value == "agent edit"
    wb.close()


def test_style_and_value_changes_can_merge_on_the_same_cell():
    base = workbook()
    remote = load_workbook(BytesIO(base))
    remote.active["A1"] = 30
    remote.active["A1"].fill = PatternFill("solid", fgColor="00FF00")
    data = BytesIO()
    remote.save(data)
    remote.close()
    style_ops = [{"kind": "cells.patch", "sheet": "销售", "cells": [{"cell": "A1", "style": {"font": {"bold": True}}}]}]
    review, merged = review_merge(base, data.getvalue(), style_ops, apply=True)
    assert review["conflict_count"] == 0
    wb = load_workbook(BytesIO(merged))
    assert wb.active["A1"].font.bold
    assert wb.active["A1"].value == 30 and wb.active["A1"].fill.fgColor.rgb.endswith("00FF00")
    wb.close()


def test_structural_changes_and_commands_require_replanning():
    base = workbook()
    review, merged = review_merge(base, workbook(A3=3), ops(A1=10), apply=True)
    assert review["status"] == "replan" and merged is None
    review, merged = review_merge(base, base, [{"kind": "insert", "sheet": "销售", "axis": "row", "at": 1, "count": 1}], apply=True)
    assert review["status"] == "replan" and merged is None


def test_boolean_changes_are_not_treated_as_equal_numbers():
    review, merged = review_merge(workbook(), workbook(A2=20), ops(A1=True), apply=True)
    assert review["safe_count"] == 1
    wb = load_workbook(BytesIO(merged))
    assert wb.active["A1"].value is True and wb.active["A1"].data_type == "b"
    wb.close()


@pytest.mark.asyncio
async def test_api_review_then_compare_and_swap_merge(tmp_path):
    path = tmp_path / "book.xlsx"
    base = workbook()
    path.write_bytes(base)
    original = content_version_of(base)
    svc = WorkspaceFileService(tmp_path)
    receipt = svc.update("book.xlsx", workbook(A1=30, B2="agent edit"), expected_version=original)
    assert receipt.state == "committed"
    with _setup_api_globals(config=_test_config(workspace_root=str(tmp_path))):
        async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
            sid = (await client.post("/api/v1/sessions", json={})).json()["id"]
            request = {"path": "book.xlsx", "session_id": sid, "batches": [{"expected_version": original, "operations": ops(A1=10, A2=20)}]}
            review_response = await client.post("/api/v1/workbooks/merge-review", json=request)
            assert review_response.status_code == 200, review_response.text
            review = review_response.json()
            assert review["conflict_count"] == 1
            assert content_version_of_file(path) == receipt.primary_version()  # preview never writes
            missing = await client.post("/api/v1/workbooks/merge-review", json={**request, "apply": True, "expected_version": review["content_version"]})
            assert missing.status_code == 400
            conflict_id = next(c["id"] for c in review["cells"] if c["conflict"])
            apply = {**request, "apply": True, "expected_version": review["content_version"], "choices": {conflict_id: "local"}}
            # A further Agent write invalidates the entire review, even with selected choices.
            later = svc.update("book.xlsx", workbook(A1=40, B2="new agent edit"), expected_version=review["content_version"])
            stale = await client.post("/api/v1/workbooks/merge-review", json=apply)
            assert stale.status_code == 409
            assert content_version_of_file(path) == later.primary_version()
            refreshed = (await client.post("/api/v1/workbooks/merge-review", json=request)).json()
            done = await client.post("/api/v1/workbooks/merge-review", json={**apply, "expected_version": refreshed["content_version"]})
            assert done.status_code == 200 and done.json()["status"] == "merged", done.text
    wb = load_workbook(path)
    assert [wb.active[c].value for c in ("A1", "A2", "B2")] == [10, 20, "new agent edit"]
    wb.close()


@pytest.mark.asyncio
async def test_api_requires_scope_and_retains_unavailable_baseline(tmp_path):
    (tmp_path / "book.xlsx").write_bytes(workbook())
    with _setup_api_globals(config=_test_config(workspace_root=str(tmp_path))):
        async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
            request = {"path": "book.xlsx", "batches": [{"expected_version": "sha256:" + "a" * 64, "operations": ops(A1=10)}]}
            assert (await client.post("/api/v1/workbooks/merge-review", json=request)).status_code == 400
            sid = (await client.post("/api/v1/sessions", json={})).json()["id"]
            response = await client.post("/api/v1/workbooks/merge-review", json={**request, "session_id": sid})
            assert response.json()["status"] == "replan"
            escaped = await client.post("/api/v1/workbooks/merge-review", json={**request, "session_id": sid, "path": "../book.xlsx"})
            assert escaped.status_code == 404
