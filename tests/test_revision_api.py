"""Revision HTTP API: list after write, restore with CAS."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus import api_app_state
from excelmanus.workbook_commit import content_version_of_file


def _harness(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg = SimpleNamespace(workspace_root=str(tmp_path), data_root="")
    api_app_state.set_config(cfg)
    api_app_state.set_session_manager(MagicMock())
    import excelmanus.api_routes_files as files_mod

    monkeypatch.setattr(
        files_mod,
        "_resolve_workspace_root",
        lambda _req, session_id=None: str(tmp_path),
    )


def _book(path: Path, value: str = "v0") -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = value
    wb.save(path)
    wb.close()
    return path


def _json(resp) -> dict:
    return json.loads(resp.body)


@pytest.mark.asyncio
async def test_list_revisions_empty_before_any_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _harness(tmp_path, monkeypatch)
    _book(tmp_path / "book.xlsx")
    from excelmanus.api_routes_workspace import list_revisions

    resp = await list_revisions(path="book.xlsx")
    assert resp.status_code == 200
    body = _json(resp)
    assert body["revisions"] == []
    assert str(body["content_version"]).startswith("sha256:")


@pytest.mark.asyncio
async def test_write_records_history_then_restore(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _harness(tmp_path, monkeypatch)
    _book(tmp_path / "book.xlsx")
    ver0 = content_version_of_file(tmp_path / "book.xlsx")
    assert ver0 is not None

    from excelmanus import api as api_module
    from excelmanus.api_routes_workspace import (
        RevisionRestoreRequest,
        list_revisions,
        restore_revision,
    )

    write_resp = await api_module.write_excel_cells(
        api_module.ExcelWriteRequest(
            path="book.xlsx",
            changes=[{"cell": "A1", "value": "v1"}],
            expected_version=ver0,
        ),
        MagicMock(),
    )
    assert write_resp.status_code == 200
    ver1 = _json(write_resp)["content_version"]

    listed = _json(await list_revisions(path="book.xlsx"))
    reasons = [item["reason"] for item in listed["revisions"]]
    assert "beforeEdit" in reasons
    assert "afterEdit" in reasons
    assert all(item.get("created_at") for item in listed["revisions"])
    before = next(item for item in listed["revisions"] if item["reason"] == "beforeEdit")

    restore_resp = await restore_revision(
        RevisionRestoreRequest(
            path="book.xlsx",
            revision_id=before["revision_id"],
            expected_version=ver1,
        )
    )
    assert restore_resp.status_code == 200
    wb = load_workbook(tmp_path / "book.xlsx")
    try:
        assert wb.active["A1"].value == "v0"
    finally:
        wb.close()

    listed2 = _json(await list_revisions(path="book.xlsx"))
    reasons2 = [item["reason"] for item in listed2["revisions"]]
    assert "beforeRestore" in reasons2
    assert reasons2[-1] == "afterEdit"


@pytest.mark.asyncio
async def test_restore_stale_version_returns_409(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _harness(tmp_path, monkeypatch)
    _book(tmp_path / "book.xlsx")
    ver0 = content_version_of_file(tmp_path / "book.xlsx")
    assert ver0 is not None

    from excelmanus import api as api_module
    from excelmanus.api_routes_workspace import (
        RevisionRestoreRequest,
        list_revisions,
        restore_revision,
    )

    write_resp = await api_module.write_excel_cells(
        api_module.ExcelWriteRequest(
            path="book.xlsx",
            changes=[{"cell": "A1", "value": "v1"}],
            expected_version=ver0,
        ),
        MagicMock(),
    )
    assert write_resp.status_code == 200
    listed = _json(await list_revisions(path="book.xlsx"))
    before = next(item for item in listed["revisions"] if item["reason"] == "beforeEdit")

    restore_resp = await restore_revision(
        RevisionRestoreRequest(
            path="book.xlsx",
            revision_id=before["revision_id"],
            expected_version="sha256:" + "0" * 64,
        )
    )
    assert restore_resp.status_code == 409
    wb = load_workbook(tmp_path / "book.xlsx")
    try:
        assert wb.active["A1"].value == "v1"
    finally:
        wb.close()


@pytest.mark.asyncio
async def test_restore_deleted_file_does_not_require_live_version(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _harness(tmp_path, monkeypatch)
    from excelmanus.workspace.file_service import WorkspaceFileService

    svc = WorkspaceFileService(tmp_path)
    svc.create("book.xlsx", b"v0")
    listed_before = svc.list_history("book.xlsx")
    revision_id = listed_before[-1].id if listed_before else None
    assert revision_id
    svc.delete("book.xlsx", expected_version=content_version_of_file(tmp_path / "book.xlsx"))
    assert not (tmp_path / "book.xlsx").exists()

    from excelmanus.api_routes_workspace import RevisionRestoreRequest, restore_revision

    response = await restore_revision(
        RevisionRestoreRequest(path="book.xlsx", revision_id=revision_id)
    )
    assert response.status_code == 200
    assert (tmp_path / "book.xlsx").read_bytes()


@pytest.mark.asyncio
async def test_list_revisions_rejects_reserved_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _harness(tmp_path, monkeypatch)
    from excelmanus.api_routes_workspace import list_revisions
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as ei:
        await list_revisions(path=".excelmanus/revisions/x.xlsx")
    assert ei.value.status_code == 400


@pytest.mark.asyncio
async def test_preview_revision_reads_immutable_workbook_bytes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _harness(tmp_path, monkeypatch)
    from io import BytesIO
    from excelmanus.workspace.file_service import WorkspaceFileService

    wb = Workbook()
    wb.active["A1"] = "historical"
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    svc = WorkspaceFileService(tmp_path)
    svc.create("book.xlsx", buffer.getvalue())
    revision_id = svc.list_history("book.xlsx")[-1].id

    from excelmanus.api_routes_workspace import preview_revision

    response = await preview_revision(path="book.xlsx", revision_id=revision_id)
    assert response.status_code == 200
    body = _json(response)
    assert body["revision_id"] == revision_id
    assert body["windows"][0]["cells"]["1,1"]["v"] == "historical"
