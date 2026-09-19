"""workbook_commit 单元测试：锁、版本冲突、原子替换、越权路径。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.security.guard import FileAccessGuard
from excelmanus.workbook_commit import (
    CommitError,
    commit_bytes,
    commit_workbook,
    content_version_of,
    content_version_of_file,
)


def _guard(tmp_path: Path) -> FileAccessGuard:
    return FileAccessGuard(str(tmp_path))


def test_commit_bytes_create_and_conflict(tmp_path: Path) -> None:
    guard = _guard(tmp_path)
    data = b"hello-workbook"
    result = commit_bytes(guard=guard, file_path="a.bin", data=data, expected_version=None)
    assert result.status == "committed"
    assert result.content_version == content_version_of(data)
    assert (tmp_path / "a.bin").read_bytes() == data

    with pytest.raises(CommitError) as ei:
        commit_bytes(guard=guard, file_path="a.bin", data=b"other", expected_version=None)
    assert ei.value.code == "VERSION_CONFLICT"


def test_commit_bytes_stale_version(tmp_path: Path) -> None:
    guard = _guard(tmp_path)
    first = commit_bytes(guard=guard, file_path="a.bin", data=b"v1", expected_version=None)
    with pytest.raises(CommitError) as ei:
        commit_bytes(
            guard=guard,
            file_path="a.bin",
            data=b"v2",
            expected_version="sha256:" + "0" * 64,
        )
    assert ei.value.code == "VERSION_CONFLICT"
    second = commit_bytes(
        guard=guard, file_path="a.bin", data=b"v2", expected_version=first.content_version
    )
    assert second.previous_version == first.content_version
    assert (tmp_path / "a.bin").read_bytes() == b"v2"


def test_commit_revision_failure_marks_history_pending(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from excelmanus.workspace import revisions as rev_mod

    def boom(self, *args, **kwargs):
        raise RuntimeError("revision store down")

    monkeypatch.setattr(rev_mod.RevisionStore, "add_record", boom)
    guard = _guard(tmp_path)
    result = commit_bytes(guard=guard, file_path="book.xlsx", data=b"ok", expected_version=None)
    assert result.status == "committed"
    assert (tmp_path / "book.xlsx").read_bytes() == b"ok"
    assert "pending_recover" in result.warnings
    receipt = result.extra.get("receipt") or {}
    assert receipt.get("history_state") == "pending_recover"


def test_commit_bytes_rejects_product_secrets(tmp_path: Path) -> None:
    guard = _guard(tmp_path)
    with pytest.raises(CommitError) as ei:
        commit_bytes(
            guard=guard,
            file_path=".secret_key",
            data=b"nope",
            expected_version=None,
        )
    assert ei.value.code == "PATH_INVALID"


def test_commit_bytes_records_revision_pair_not_backups(tmp_path: Path) -> None:
    from excelmanus.workspace.revisions import RevisionStore

    guard = _guard(tmp_path)
    first = commit_bytes(guard=guard, file_path="book.xlsx", data=b"v1", expected_version=None)
    second = commit_bytes(
        guard=guard, file_path="book.xlsx", data=b"v2", expected_version=first.content_version
    )
    store = RevisionStore(tmp_path)
    recs = store.list("book.xlsx")
    reasons = [r.reason for r in recs]
    assert reasons.count("afterEdit") == 2
    assert "beforeEdit" in reasons
    assert (tmp_path / "book.xlsx").read_bytes() == b"v2"
    assert not (tmp_path / "outputs" / "backups").exists()
    assert second.previous_version == first.content_version
    blob_dir = store.root / store.path_key("book.xlsx") / "blobs"
    blobs = [p for p in blob_dir.iterdir() if p.is_file()]
    assert blobs


def test_commit_workbook_records_after_edit(tmp_path: Path) -> None:
    from excelmanus.workspace.revisions import RevisionStore

    guard = _guard(tmp_path)

    def mutate(wb) -> None:
        wb.active["A1"] = "hello"

    commit_workbook(guard=guard, file_path="new.xlsx", mutate_fn=mutate, create=True)
    store = RevisionStore(tmp_path)
    recs = store.list("new.xlsx")
    assert [r.reason for r in recs] == ["afterEdit"]
    assert not (tmp_path / "outputs" / "backups").exists()
    assert not list((tmp_path / "outputs").glob("**/*")) or all(
        "backups" not in str(p) for p in (tmp_path / "outputs").rglob("*")
    )


def test_commit_workbook_mutate(tmp_path: Path) -> None:
    guard = _guard(tmp_path)
    seed = Workbook()
    seed.active["A1"] = "old"
    seed.save(tmp_path / "book.xlsx")
    ver = content_version_of_file(tmp_path / "book.xlsx")
    assert ver is not None

    def mutate(wb) -> None:
        wb.active["A1"] = "new"

    result = commit_workbook(
        guard=guard,
        file_path="book.xlsx",
        mutate_fn=mutate,
        expected_version=ver,
    )
    assert result.status == "committed"
    wb = load_workbook(tmp_path / "book.xlsx")
    assert wb.active["A1"].value == "new"
    assert content_version_of_file(tmp_path / "book.xlsx") == result.content_version
    from excelmanus.workspace.revisions import RevisionStore

    recs = RevisionStore(tmp_path).list("book.xlsx")
    assert [r.reason for r in recs] == ["beforeEdit", "afterEdit"]
    assert recs[0].transaction_id == recs[1].transaction_id


def test_commit_rejects_path_traversal(tmp_path: Path) -> None:
    guard = _guard(tmp_path)
    with pytest.raises(CommitError) as ei:
        commit_bytes(guard=guard, file_path="../secret.bin", data=b"x", expected_version=None)
    assert ei.value.code == "PATH_INVALID"


def test_commit_rejects_reserved_namespace(tmp_path: Path) -> None:
    guard = _guard(tmp_path)
    with pytest.raises(CommitError) as ei:
        commit_bytes(
            guard=guard,
            file_path=".excelmanus/revisions/x.xlsx",
            data=b"x",
            expected_version=None,
        )
    assert ei.value.code == "PATH_INVALID"
    with pytest.raises(CommitError) as ei:
        commit_bytes(
            guard=guard,
            file_path="outputs/backups/copy.xlsx",
            data=b"x",
            expected_version=None,
        )
    assert ei.value.code == "PATH_INVALID"


def test_commit_workbook_create(tmp_path: Path) -> None:
    guard = _guard(tmp_path)

    def mutate(wb) -> None:
        wb.active["B2"] = 100

    result = commit_workbook(
        guard=guard,
        file_path="new.xlsx",
        mutate_fn=mutate,
        create=True,
    )
    assert result.previous_version is None
    wb = load_workbook(tmp_path / "new.xlsx")
    assert wb.active["B2"].value == 100


def test_commit_workbook_requires_expected_version(tmp_path: Path) -> None:
    guard = _guard(tmp_path)
    seed = Workbook()
    seed.active["A1"] = "old"
    seed.save(tmp_path / "book.xlsx")

    def mutate(wb) -> None:
        wb.active["A1"] = "new"

    with pytest.raises(CommitError) as ei:
        commit_workbook(
            guard=guard,
            file_path="book.xlsx",
            mutate_fn=mutate,
        )
    assert ei.value.code == "VERSION_CONFLICT"


def test_create_excel_chart_goes_through_commit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """create_excel_chart 必须经 commit_workbook，且返回 content_version。"""
    import excelmanus.workbook_commit as wc
    from excelmanus.workbook.charts import create_excel_chart, init_guard

    init_guard(str(tmp_path))
    seed = Workbook()
    ws = seed.active
    ws.title = "数据"
    ws["A1"] = "月份"
    ws["B1"] = "营收"
    ws["A2"] = "1月"
    ws["B2"] = 100
    ws["A3"] = "2月"
    ws["B3"] = 150
    seed.save(tmp_path / "chart.xlsx")
    seed.close()

    seen: list[dict] = []
    real = wc.commit_workbook

    def _wrapped(**kwargs):
        seen.append(kwargs)
        return real(**kwargs)

    monkeypatch.setattr(wc, "commit_workbook", _wrapped)

    from excelmanus.workbook_commit import content_version_of_file

    charted = create_excel_chart(
        file_path="chart.xlsx",
        chart_type="bar",
        data_range="B1:B3",
        categories_range="A2:A3",
        target_cell="D1",
        expected_version=content_version_of_file(tmp_path / "chart.xlsx"),
    )
    payload = charted.value
    assert charted.success
    assert payload["status"] == "success"
    assert seen, "create_excel_chart 应调用 commit_workbook"
    assert seen[0]["file_path"] == "chart.xlsx"
    assert str(payload.get("content_version", "")).startswith("sha256:")


def test_tool_write_conflicts_after_read_then_external_edit(tmp_path: Path) -> None:
    """先读记住版本，外部改盘后再写，工具层必须 VERSION_CONFLICT。"""
    from excelmanus.security import FileAccessGuard
    from excelmanus.tools._guard_ctx import set_guard
    from excelmanus.tools.intent_tools import edit_spreadsheet, init_guard
    from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions

    workspace = str(tmp_path)
    set_guard(FileAccessGuard(workspace))
    init_guard(workspace)
    seed = Workbook()
    seed.active["A1"] = "seen"
    seed.save(tmp_path / "book.xlsx")
    seed.close()
    seen_ver = content_version_of_file(tmp_path / "book.xlsx")
    seed_seen_versions({"book.xlsx": seen_ver})
    try:
        outsider = load_workbook(tmp_path / "book.xlsx")
        outsider.active["A1"] = "external"
        outsider.save(tmp_path / "book.xlsx")
        outsider.close()

        result = edit_spreadsheet(
            file_path="book.xlsx",
            operations=[{"kind": "write", "sheet": "Sheet", "start_cell": "A1", "values": [["from-tool"]]}],
        )
        assert result.success is False
        assert result.error is not None
        assert result.error.code == "VERSION_CONFLICT"
        wb = load_workbook(tmp_path / "book.xlsx")
        assert wb.active["A1"].value == "external"
    finally:
        seed_seen_versions({})


@pytest.mark.asyncio
async def test_write_excel_cells_stale_version_returns_409(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """API write 在提供错误 expected_version 时返回 HTTP 409。"""
    from types import SimpleNamespace
    from unittest.mock import MagicMock

    from excelmanus import api as api_module
    from excelmanus import api_app_state
    import excelmanus.api_routes_files as files_mod

    seed = Workbook()
    seed.active["A1"] = "old"
    seed.save(tmp_path / "book.xlsx")
    seed.close()

    cfg = SimpleNamespace(workspace_root=str(tmp_path))
    api_app_state.set_config(cfg)
    api_app_state.set_session_manager(None)
    monkeypatch.setattr(files_mod, "_resolve_workspace_root", lambda _req, session_id=None: str(tmp_path))

    req = api_module.ExcelWriteRequest(
        path="book.xlsx",
        changes=[{"cell": "A1", "value": "new"}],
        expected_version="sha256:" + "0" * 64,
    )
    raw = MagicMock()
    raw.app.state.auth_enabled = False
    resp = await api_module.write_excel_cells(req, raw)
    assert resp.status_code == 409
    body = resp.body.decode() if isinstance(resp.body, (bytes, bytearray)) else str(resp.body)
    assert "VERSION_CONFLICT" in body or "版本" in body or "conflict" in body.lower()


def _write_harness(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from types import SimpleNamespace
    from unittest.mock import MagicMock

    from excelmanus import api as api_module
    from excelmanus import api_app_state
    import excelmanus.api_routes_files as files_mod

    cfg = SimpleNamespace(workspace_root=str(tmp_path))
    api_app_state.set_config(cfg)
    api_app_state.set_session_manager(None)
    monkeypatch.setattr(files_mod, "_resolve_workspace_root", lambda _req, session_id=None: str(tmp_path))
    raw = MagicMock()
    raw.app.state.auth_enabled = False
    return api_module, raw


@pytest.mark.asyncio
async def test_write_excel_cells_missing_version_returns_409(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    api_module, raw = _write_harness(tmp_path, monkeypatch)
    seed = Workbook()
    seed.active["A1"] = "old"
    seed.save(tmp_path / "book.xlsx")
    seed.close()

    req = api_module.ExcelWriteRequest(
        path="book.xlsx",
        changes=[{"cell": "A1", "value": "new"}],
    )
    resp = await api_module.write_excel_cells(req, raw)
    assert resp.status_code == 409
    wb = load_workbook(tmp_path / "book.xlsx")
    assert wb.active["A1"].value == "old"


@pytest.mark.asyncio
async def test_snapshot_then_external_edit_then_stale_write_409(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """读快照后磁盘被改，再用快照版本提交应 409，不得覆盖外部修改。"""
    from unittest.mock import MagicMock

    api_module, raw = _write_harness(tmp_path, monkeypatch)
    seed = Workbook()
    seed.active["A1"] = "seen"
    seed.save(tmp_path / "book.xlsx")
    seed.close()

    snap_req = MagicMock()
    snap_req.query_params = {
        "path": "book.xlsx",
        "all_sheets": "1",
        "max_rows": "50",
        "with_styles": "0",
    }
    snap_req.app.state.auth_enabled = False
    snap = await api_module.get_excel_snapshot(snap_req)
    assert snap.status_code == 200
    import json as _json
    body = _json.loads(snap.body)
    seen_ver = body.get("content_version")
    assert isinstance(seen_ver, str) and seen_ver.startswith("sha256:")

    outsider = load_workbook(tmp_path / "book.xlsx")
    outsider.active["A1"] = "external"
    outsider.save(tmp_path / "book.xlsx")
    outsider.close()

    req = api_module.ExcelWriteRequest(
        path="book.xlsx",
        changes=[{"cell": "A1", "value": "from-ui"}],
        expected_version=seen_ver,
    )
    resp = await api_module.write_excel_cells(req, raw)
    assert resp.status_code == 409
    wb = load_workbook(tmp_path / "book.xlsx")
    assert wb.active["A1"].value == "external"


@pytest.mark.asyncio
async def test_write_excel_cells_clears_and_writes_other_sheet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    api_module, raw = _write_harness(tmp_path, monkeypatch)
    seed = Workbook()
    seed.active.title = "Sheet1"
    seed.active["A1"] = "keep-me"
    ws2 = seed.create_sheet("Sheet2")
    ws2["A1"] = "old-two"
    seed.save(tmp_path / "book.xlsx")
    seed.close()
    ver = content_version_of_file(tmp_path / "book.xlsx")

    req = api_module.ExcelWriteRequest(
        path="book.xlsx",
        changes=[
            {"cell": "A1", "value": None, "sheet": "Sheet1"},
            {"cell": "A1", "value": "new-two", "sheet": "Sheet2"},
        ],
        expected_version=ver,
    )
    resp = await api_module.write_excel_cells(req, raw)
    assert resp.status_code == 200
    wb = load_workbook(tmp_path / "book.xlsx")
    assert wb["Sheet1"]["A1"].value is None
    assert wb["Sheet2"]["A1"].value == "new-two"


def test_commit_unlink_requires_expected_version(tmp_path: Path) -> None:
    from excelmanus.workbook_commit import commit_unlink

    guard = _guard(tmp_path)
    target = tmp_path / "a.bin"
    target.write_bytes(b"keep")
    with pytest.raises(CommitError) as ei:
        commit_unlink(guard=guard, file_path="a.bin")
    assert ei.value.code == "VERSION_CONFLICT"
    assert target.read_bytes() == b"keep"

    result = commit_unlink(
        guard=guard,
        file_path="a.bin",
        expected_version=content_version_of(b"keep"),
    )
    assert result.status == "committed"
    assert not target.exists()


def test_commit_move_requires_expected_version(tmp_path: Path) -> None:
    from excelmanus.workbook_commit import commit_move

    guard = _guard(tmp_path)
    (tmp_path / "a.bin").write_bytes(b"keep")
    with pytest.raises(CommitError) as ei:
        commit_move(guard=guard, source="a.bin", destination="b.bin")
    assert ei.value.code == "VERSION_CONFLICT"
    assert (tmp_path / "a.bin").read_bytes() == b"keep"

    result = commit_move(
        guard=guard,
        source="a.bin",
        destination="b.bin",
        expected_version=content_version_of(b"keep"),
    )
    assert result.status == "committed"
    assert not (tmp_path / "a.bin").exists()
    assert (tmp_path / "b.bin").read_bytes() == b"keep"
