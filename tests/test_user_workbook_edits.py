from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.agent.inbox import Inbox
from excelmanus.api_app_state import AppRuntime, bind_runtime, reset_runtime
from excelmanus.config import ExcelManusConfig
from excelmanus.engine import AgentEngine
from excelmanus.session import SessionManager
from excelmanus.tools.registry import ToolDef, ToolRegistry
from excelmanus.workbook.user_edits import PROMPT_KIND, summarize_operations
from excelmanus.workbook_commit import content_version_of_file
from excelmanus.workspace.file_service import WorkspaceFileService


def engine_at(root: Path) -> AgentEngine:
    return AgentEngine(ExcelManusConfig(api_key="test", model="test-model", base_url="https://test.example/v1",
        workspace_root=str(root), max_iterations=8), ToolRegistry())


def manager_for(**engines):
    manager = object.__new__(SessionManager)
    manager._database = None
    manager._migrated_workspace_paths = set()
    manager._sessions = {sid: SimpleNamespace(engine=engine) for sid, engine in engines.items()}
    return manager


def edit_event(index=1):
    return {"event_id": f"edit-{index}", "path": "book.xlsx", "after_version": f"v{index}",
        "context": {"source": "user", "summary": summarize_operations([
            {"op": "set_values", "sheet": "Sales", "cells": [{"cell": "A1", "value": f"=B1+{index}"}]}])}}


def response(text=None, tools=None):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text or "", tool_calls=tools))])


def call(name="ping", ident="c1"):
    return SimpleNamespace(id=ident, type="function", function=SimpleNamespace(name=name, arguments="{}"))


@pytest.fixture
def write_env(tmp_path, monkeypatch):
    import excelmanus.api_routes_files as files
    seed = Workbook()
    seed.active.title = "Sales"
    seed.active["A1"] = "original"
    seed.save(tmp_path / "book.xlsx")
    seed.close()
    engine = engine_at(tmp_path)
    manager = manager_for(s1=engine)
    token = bind_runtime(AppRuntime(config=SimpleNamespace(workspace_root=str(tmp_path)), session_manager=manager))
    monkeypatch.setattr(files, "_file_workspace_root", lambda *args: (str(tmp_path), None))
    monkeypatch.setattr(files, "_resolve_excel_path", lambda *args, **kwargs: str(tmp_path / "book.xlsx"))
    yield files, engine, manager
    reset_runtime(token)


@pytest.mark.asyncio
async def test_saved_user_changes_reach_idle_agent_and_other_sessions_only_in_workspace(write_env, tmp_path):
    files, engine, manager = write_env
    same = engine_at(tmp_path)
    other = engine_at(tmp_path / "other")
    manager._sessions.update(s2=SimpleNamespace(engine=same), foreign=SimpleNamespace(engine=other))
    version = content_version_of_file(tmp_path / "book.xlsx")
    engine._state.file_content_versions.update({"./book.xlsx": version, "unrelated.xlsx": "keep"})
    req = files.ExcelWriteRequest(path="book.xlsx", session_id="s1", expected_version=version,
        changes=[{"sheet": "Sales", "cell": "A1", "value": "=B1+2"}], operation_id="user-edit")
    result = await files.write_excel_cells(req, MagicMock())
    assert result.status_code == 200
    assert engine._driver.status == "idle"
    assert len(engine._driver.inbox.next_step) == 1
    brief = engine._driver.inbox.next_step[0]
    assert brief.extra["prompt_kind"] == PROMPT_KIND
    assert "A1" in brief.content and "=B1+2" in brief.content and "Sales" in brief.content
    assert engine._state.file_content_versions == {"unrelated.xlsx": "keep"}
    assert len(same._driver.inbox.next_step) == 1
    assert not other._driver.inbox.next_step
    manager.drain_workspace_events()
    replay = await files.write_excel_cells(req, MagicMock())
    assert replay.status_code == 200
    assert brief.extra["workbook_event_count"] == 1
    # A session loaded later has an independent durable cursor.
    restored = engine_at(tmp_path)
    manager._sessions["restored"] = SimpleNamespace(engine=restored)
    manager.drain_workspace_events()
    assert len(restored._driver.inbox.next_step) == 1
    engine._client.chat.completions.create = AsyncMock(return_value=response("done"))
    await engine.followup("continue")
    sent = engine._client.chat.completions.create.call_args.kwargs["messages"]
    assert any("用户表格改动简报" in str(m.get("content")) for m in sent)


@pytest.mark.asyncio
async def test_conflicting_write_preserves_agent_change_and_emits_no_user_brief(write_env, tmp_path):
    files, engine, _ = write_env
    path = tmp_path / "book.xlsx"
    old = content_version_of_file(path)
    wb = load_workbook(path)
    wb.active["A1"] = "agent change"
    wb.save(path)
    wb.close()
    result = await files.write_excel_cells(files.ExcelWriteRequest(path="book.xlsx", expected_version=old,
        changes=[{"cell": "A1", "value": "stale user change"}]), MagicMock())
    assert result.status_code == 409
    assert not engine._driver.inbox.next_step
    wb = load_workbook(path)
    assert wb.active["A1"].value == "agent change"
    wb.close()


def test_edit_briefs_coalesce_bound_size_and_survive_inbox_restore(tmp_path):
    engine = engine_at(tmp_path)
    for index in range(50):
        engine._driver.inject_workbook_change(edit_event(index))
    item = engine._driver.inbox.next_step[0]
    assert len(engine._driver.inbox.next_step) == 1
    assert item.extra["workbook_event_count"] == 50
    assert len(item.content) < 13000
    assert "=B1+49" in item.content
    engine._driver.inject_workbook_change(edit_event(49))
    assert item.extra["workbook_event_count"] == 50
    restored = Inbox()
    restored.load_reconstructed(engine._driver.inbox.reconstruct())
    assert restored.next_step[0].snapshot() == item.snapshot()


@pytest.mark.asyncio
@pytest.mark.parametrize("during", ["model", "tool", "final_reply", "preparing"])
async def test_user_changes_arriving_in_flight_are_consumed_before_next_decision(tmp_path, during):
    engine = engine_at(tmp_path)
    executed = []
    def ping():
        executed.append("ping")
        if during == "tool":
            engine._driver.inject_workbook_change(edit_event())
        return "pong"
    engine._registry.register_tool(ToolDef(name="ping", description="ping", input_schema={"type": "object", "properties": {}}, func=ping))
    if during == "preparing":
        original = engine._apply_claimed_followup
        async def preparing(item):
            result = await original(item)
            engine._driver.inject_workbook_change(edit_event())
            return result
        engine._apply_claimed_followup = preparing
    calls = []
    async def model(**kwargs):
        calls.append(kwargs["messages"])
        if len(calls) == 1 and during in {"model", "final_reply"}:
            engine._driver.inject_workbook_change(edit_event())
            return response("old answer") if during == "final_reply" else response(tools=[call()])
        if len(calls) == 1 and during == "tool":
            return response(tools=[call(), call(ident="c2")])
        return response("updated answer")
    engine._client.chat.completions.create = AsyncMock(side_effect=model)
    result = await engine.followup("update workbook")
    assert result.reply == "updated answer"
    assert any("用户表格改动简报" in str(m.get("content")) for m in calls[-1])
    assert executed == (["ping"] if during == "tool" else [])
    assert not engine._driver.inbox.next_step
    if during == "model":
        assert any(m.get("tool_call_id") == "c1" and "USER_EDIT_PENDING" in str(m.get("content")) for m in calls[-1])


def test_context_survives_transaction_recovery(tmp_path, monkeypatch):
    svc = WorkspaceFileService(tmp_path)
    receipt = svc.create("a.txt", b"old")
    original = svc._project_history
    monkeypatch.setattr(svc, "_project_history", lambda *args: "pending_recover")
    edited = svc.update_with_builder("a.txt", lambda _: b"new", expected_version=receipt.primary_version(),
        event_context={"source": "user", "summary": "A1 changed"}, intent={"edit": 1})
    assert edited.history_state == "pending_recover"
    monkeypatch.setattr(svc, "_project_history", original)
    svc.recover()
    received = []
    svc.deliver_outbox(received.append, consumer_id="test")
    assert received[-1]["context"] == {"source": "user", "summary": "A1 changed"}


def test_structural_and_large_paste_summaries_are_bounded():
    summary = summarize_operations([{"op": "insert_axis", "sheet": "Sales", "axis": "row", "index": 2, "count": 4},
        {"op": "set_values", "sheet": "Sales", "cells": [{"cell": f"A{i}", "value": "x" * 10000} for i in range(100)]}])
    assert "insert_axis" in summary and '"cell_count": 100' in summary
    assert len(summary) <= 2420

@pytest.mark.asyncio
async def test_real_save_during_model_request_defers_stale_tool(write_env, tmp_path):
    import asyncio
    files, engine, _ = write_env
    entered, release = asyncio.Event(), asyncio.Event()
    executed = []
    engine._registry.register_tool(ToolDef(name="ping", description="ping", input_schema={"type":"object", "properties":{}},
        func=lambda: executed.append("stale tool") or "pong"))
    model_calls = []
    async def model(**kwargs):
        model_calls.append(kwargs["messages"])
        if len(model_calls) == 1:
            entered.set()
            await release.wait()
            return response(tools=[call()])
        return response("read latest user changes")
    engine._client.chat.completions.create = AsyncMock(side_effect=model)
    task = asyncio.create_task(engine.followup("update workbook"))
    try:
        await asyncio.wait_for(entered.wait(), 5)
        saved = await files.write_excel_cells(files.ExcelWriteRequest(path="book.xlsx", session_id="s1",
            expected_version=content_version_of_file(tmp_path / "book.xlsx"),
            changes=[{"sheet":"Sales", "cell":"B2", "value":777}]), MagicMock())
        assert saved.status_code == 200
    finally:
        release.set()
        result = await task
    assert result.reply == "read latest user changes"
    assert not executed
    assert any("B2" in str(m.get("content")) and "777" in str(m.get("content")) for m in model_calls[-1])


@pytest.mark.asyncio
async def test_restoring_workbook_history_notifies_agent(write_env, tmp_path, monkeypatch):
    from excelmanus import api_routes_workspace as revisions
    files, engine, _ = write_env
    path = tmp_path / "book.xlsx"
    saved = await files.write_excel_cells(files.ExcelWriteRequest(path="book.xlsx",
        expected_version=content_version_of_file(path), changes=[{"cell":"A1", "value":"new"}]), MagicMock())
    assert saved.status_code == 200
    before = next(r for r in WorkspaceFileService(tmp_path).list_history("book.xlsx") if r.reason == "beforeEdit")
    monkeypatch.setattr(revisions, "_workspace_root", lambda *args: tmp_path)
    restored = await revisions.restore_revision(revisions.RevisionRestoreRequest(path="book.xlsx",
        revision_id=before.id, expected_version=content_version_of_file(path)))
    assert restored.status_code == 200
    item = engine._driver.inbox.next_step[0]
    assert item.extra["workbook_event_count"] == 2
    assert "恢复了整个文件" in item.content
    wb = load_workbook(path)
    assert wb.active["A1"].value == "original"
    wb.close()
