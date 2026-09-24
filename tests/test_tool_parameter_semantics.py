"""Adjacent tool contracts: no silently ignored coordinates, options or aliases."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.tools.context import use_workspace
from excelmanus.tools.registry import ToolRegistry
from excelmanus.workbook_commit import content_version_of_file


@pytest.fixture
def book(tmp_path: Path) -> Path:
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    for name in ("Data", "Other"):
        ws = wb.active if name == "Data" else wb.create_sheet()
        ws.title = name
        ws.append(["id", "value", "note"])
        ws.append([1, 10, "a long description"])
        ws.append([2, 20, "tiny"])
        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 40
        ws.row_dimensions[1].height = 40
        ws.row_dimensions[2].height = 40
        ws.row_dimensions[3].height = 40
        ws.freeze_panes = "C3"
    wb.save(path)
    wb.close()
    return path


def _call(book: Path, operations: list[dict], *, tool: str = "apply_spreadsheet_changes", enforce: bool = True):
    registry = ToolRegistry()
    registry.register_builtin_tools(str(book.parent))
    registry.configure_schema_validation(mode="enforce" if enforce else "off", canary_percent=100, strict_path=False)
    with use_workspace(book.parent):
        return registry.call_tool(tool, {"file_path": book.name, "operations": operations,
                                       "expected_version": content_version_of_file(book)})


@pytest.mark.parametrize("axis", ["column", "row"])
def test_auto_fit_only_changes_requested_range(book: Path, axis: str) -> None:
    result = _call(book, [{"kind": "size", "sheet": "Data", "range": "B2:B2", "auto_fit": True, "axis": axis}])
    assert result.success, result.model_text
    wb = load_workbook(book)
    try:
        ws = wb["Data"]
        if axis == "column":
            assert ws.column_dimensions["A"].width == ws.column_dimensions["C"].width == 40
            assert ws.column_dimensions["B"].width != 40
        else:
            assert ws.row_dimensions[1].height == ws.row_dimensions[3].height == 40
            assert ws.row_dimensions[2].height != 40
    finally:
        wb.close()


def test_column_auto_fit_and_explicit_row_sizes_apply_in_one_batch(book: Path) -> None:
    result = _call(book, [
        {"kind": "size", "sheet": "Data", "range": "B2:B2", "auto_fit": True, "axis": "column"},
        {"kind": "size", "sheet": "Data", "row_heights": {"2": 27}},
    ])
    assert result.success, result.model_text
    wb = load_workbook(book)
    try:
        assert wb["Data"].column_dimensions["B"].width != 40
        assert wb["Data"].row_dimensions[2].height == 27
        assert wb["Data"].column_dimensions["A"].width == 40
    finally:
        wb.close()


@pytest.mark.parametrize("target", ["C2", None])
def test_freeze_cell_contract(book: Path, target: str | None) -> None:
    result = _call(book, [{"kind": "freeze", "sheet": "Data", "freeze_panes": target or ""}])
    assert result.success, result.model_text
    wb = load_workbook(book)
    try:
        assert wb["Data"].freeze_panes == target
    finally:
        wb.close()


def test_missing_freeze_target_guides_to_canonical_field(book: Path) -> None:
    result = _call(book, [{"kind": "freeze", "sheet": "Data"}])
    assert not result.success
    assert "freeze_panes" in result.error.message
    assert "rows/cols" not in result.error.message
    assert _call(book, [{"kind": "freeze", "sheet": "Data", "freeze_panes": "A2"}]).success


def test_empty_dimensions_recovery_example_is_accepted(book: Path) -> None:
    result = _call(book, [{"kind": "size", "sheet": "Data", "column_widths": {}}])
    assert not result.success
    assert "column_widths/row_heights" in result.error.message
    assert "columns/rows" not in result.error.message
    example = result.value["example"]
    assert _call(book, [example]).success


@pytest.mark.parametrize("op", [
    {"kind": "freeze"},
    {"kind": "freeze", "rows": 1.5},
    {"kind": "freeze", "rows": -1},
    {"kind": "freeze", "freeze_panes": "A2", "rows": 2},
    {"kind": "size", "columns": [12, "bad", 14]},
    {"kind": "size", 'column_widths': {"A": -1}},
    {"kind": "size", 'row_heights': {"0": 12}},
    {"kind": "size", 'column_widths': {"A": 12}, "font": {"bold": True}},
    {"kind": "format", "range": "A1", "rule": {"type": "formula", "formula": "=TRUE"}},
    {"kind": "format", "range": "A1", "font": {"bold": True}, "typo": 1},
    {"kind": "format", "range": "A1", "cell_range": "B1", "font": {"bold": True}},
    {"kind": "format", "range": "A1", "sheet_name": "Other", "font": {"bold": True}},
    {"kind": "format", "range": "A1", "font": {"unknown": True}},
])
def test_invalid_operation_is_actionable_and_atomic(book: Path, op: dict) -> None:
    before = book.read_bytes()
    result = _call(book, [
        {"kind": "size", "sheet": "Data", 'column_widths': {"C": 15}},
        {"sheet": "Data", **op},
    ], enforce=False)
    assert not result.success, result.model_text
    assert result.error.code == "INVALID_ARGS", result.model_text
    assert result.value["operation_index"] == 1
    assert result.value["committed"] is False
    assert book.read_bytes() == before


def test_write_alias_conflict_rejects_before_commit(book: Path) -> None:
    before = book.read_bytes()
    result = _call(book, [{"kind": "write", "sheet": "Data", "start_cell": "A1", "startCell": "B1", "values": [[99]]}], tool="apply_spreadsheet_changes")
    assert not result.success
    assert result.value["operation_index"] == 0
    assert book.read_bytes() == before


@pytest.mark.parametrize("prior_change", [False, True])
def test_chart_target_does_not_disambiguate_source(book: Path, prior_change: bool) -> None:
    before = book.read_bytes()
    operations = [{"kind": "size", "sheet": "Data", "column_widths": {"A": 12}}] if prior_change else []
    operations.append({"kind": "chart", "chart_type": "bar", "data_range": "B1:B3", "target_cell": "Other!E1"})
    result = _call(book, operations)
    assert not result.success, result.model_text
    assert result.error.code == "SHEET_REQUIRED"
    assert result.value["operation_index"] == int(prior_change)
    assert result.value["available_sheets"] == ["Data", "Other"]
    assert result.value["committed"] is False
    assert result.value["applied"] == []
    assert book.read_bytes() == before


def test_chart_canonical_fields_work_with_enforced_schema(book: Path) -> None:
    result = _call(book, [{"kind": "chart", "sheet": "Data", "chart_type": "bar", "data_range": "B1:B3", "categories_range": "A2:A3", "target_sheet": "Other", "target_cell": "E1", "from_rows": False}])
    assert result.success, result.model_text
    wb = load_workbook(book)
    try:
        assert len(wb["Other"]._charts) == 1
    finally:
        wb.close()


def test_retired_chart_aliases_are_rejected_atomically(book: Path) -> None:
    before = book.read_bytes()
    result = _call(book, [{
        "kind": "create_chart", "sheet_name": "Data", "chartType": "bar",
        "dataRange": "B1:B3", "targetSheet": "Other", "targetCell": "E1",
    }])
    assert not result.success
    assert result.value["operation_index"] == 0
    assert result.value["committed"] is False
    assert book.read_bytes() == before


def test_conditional_text_escapes_excel_string_literals(book: Path) -> None:
    text = 'say "hello"'
    result = _call(book, [{"kind": "conditional_format", "sheet": "Data", "range": "C2:C3", "rule": {"type": "text", "text": text, "fill": {"color": "FFFF00"}}}])
    assert result.success, result.model_text
    wb = load_workbook(book)
    try:
        ws = wb["Data"]
        rule = ws.conditional_formatting[next(iter(ws.conditional_formatting))][0]
        assert rule.formula == ['NOT(ISERROR(SEARCH("say ""hello""",C2)))']
    finally:
        wb.close()


def test_spill_json_array_stays_array(book: Path) -> None:
    from excelmanus.engine_core.spill import SpillStore
    from excelmanus.tools.file_tools import read_text_file
    matrix = [[1, None, '"quoted"'], [2, 3, 4]]
    locator = SpillStore(book.parent).put(json.dumps(matrix))
    with use_workspace(book.parent):
        result = read_text_file(locator)
    assert result.success
    assert result.value == matrix


def test_inspect_range_without_range_or_window_is_rejected(book: Path) -> None:
    from excelmanus.tools.workbook_tools import observe_spreadsheet

    with use_workspace(book.parent):
        result = observe_spreadsheet(file_path=book.name, mode="range")
    assert not result.success
    assert result.error.code == "INVALID_ARGS"
    assert "range" in result.value["missing_fields"]


def test_relationships_accepts_single_file_path(book: Path) -> None:
    from excelmanus.tools.workbook_tools import analyze_spreadsheet

    with use_workspace(book.parent):
        result = analyze_spreadsheet(mode="relationships", file_path=book.name)
    assert result.success, result.model_text
    assert result.value["files_analyzed"] == 1


def test_impact_scope_filters_sheets(book: Path) -> None:
    from excelmanus.tools.workbook_tools import trace_spreadsheet_formulas
    wb = load_workbook(book)
    wb["Data"]["B2"] = "=A2"
    wb["Other"]["B2"] = "=Data!A2"
    wb.save(book)
    wb.close()
    with use_workspace(book.parent):
        result = trace_spreadsheet_formulas(
            file_path=book.name, mode="impact", target="Data!A2", scope="sheet",
        )
    assert result.success, result.model_text
    assert result.value["scope"] == "sheet"
    assert all(item["sheet"] == "Data" for item in result.value["direct_impact"])


@pytest.mark.parametrize("rule", [
    {"type": "cell_value", "operator": "greaterThan", "value": 1, "formula1": 2},
    {"type": "list", "values": ["A", "B"], "formula1": "Other!A1:A2"},
])
def test_format_semantic_alias_conflict_is_invalid(book: Path, rule: dict) -> None:
    kind = "data_validation" if rule["type"] == "list" else "conditional_format"
    before = book.read_bytes()
    result = _call(book, [{"kind": kind, "sheet": "Data", "range": "A2:A3", "rule": rule}])
    assert not result.success
    assert result.error.code == "INVALID_ARGS"
    assert book.read_bytes() == before
