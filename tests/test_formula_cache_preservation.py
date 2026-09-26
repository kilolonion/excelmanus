"""值不变写入保留公式缓存；改动计算输入才要求显式重算。

复盘来源：收款收据还原任务里一个纯数字格式补丁被迫再走一次 calculate_spreadsheet。
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
import xlsxwriter

from excelmanus.security import FileAccessGuard
from excelmanus.tools import workbook_tools
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.workbook.ooxml import merge_formula_caches
from excelmanus.workbook_commit import content_version_of_file


def _make_calculated_book(path: Path) -> None:
    """创建带公式的簿，再把“引擎算好的缓存”抄进包（模拟 calculate_spreadsheet 之后）。"""
    set_guard(FileAccessGuard(str(path.parent)))
    workbook_tools.init_guard(str(path.parent))
    result = workbook_tools.apply_spreadsheet_changes(
        file_path=str(path),
        create=True,
        workbook_spec={
            "sheets": [{
                "name": "Data",
                "dimensions": {"rows": 3, "cols": 2},
                "value_blocks": [{"start": "A1", "values": [["x", "y"], [2, None], [3, None]]}],
                "formula_blocks": [{"start": "B2", "formulas": [["=A2*10"], ["=A3*10"]]}],
            }],
            "uncertainties": [],
        },
    )
    assert result.success, result.model_text
    stream = BytesIO()
    with xlsxwriter.Workbook(stream) as book:
        sheet = book.add_worksheet("Data")
        sheet.write_row("A1", ["x", "y"])
        sheet.write("A2", 2)
        sheet.write("A3", 3)
        sheet.write_formula("B2", "=A2*10", None, 20)
        sheet.write_formula("B3", "=A3*10", None, 30)
    calculated, count = merge_formula_caches(path.read_bytes(), stream.getvalue())
    assert count == 2
    path.write_bytes(calculated)


def test_value_preserving_write_keeps_formula_caches(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    _make_calculated_book(path)
    result = workbook_tools.apply_spreadsheet_changes(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "format", "sheet": "Data", "range": "A1:B1",
            "font": {"bold": True},
        }],
    )
    assert result.success, result.model_text
    observation = result.value["observation"]
    assert observation["calculation_inputs"] == "unchanged"
    assert observation["formula_cache"] == "preserved_unchanged_calculation_inputs"
    workbook = load_workbook(path, data_only=True)
    try:
        assert [workbook["Data"][f"B{row}"].value for row in (2, 3)] == [20, 30]
    finally:
        workbook.close()


def test_value_change_drops_caches_and_requests_recalculation(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    _make_calculated_book(path)
    result = workbook_tools.apply_spreadsheet_changes(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "write", "sheet": "Data", "start_cell": "A2",
            "values": [[5]],
        }],
    )
    assert result.success, result.model_text
    observation = result.value["observation"]
    assert observation["calculation_inputs"] == "changed"
    assert observation["formula_cache"].startswith("invalidated")
    workbook = load_workbook(path, data_only=True)
    try:
        assert workbook["Data"]["B2"].value is None  # 缺缓存不能当作 20/0
    finally:
        workbook.close()


def test_uncached_formula_book_still_reports_invalidated(tmp_path: Path) -> None:
    """源文件本来就没有缓存时不能声称保留（沿用失效声明，要求显式重算）。"""
    path = tmp_path / "book.xlsx"
    set_guard(FileAccessGuard(str(tmp_path)))
    workbook_tools.init_guard(str(tmp_path))
    result = workbook_tools.apply_spreadsheet_changes(
        file_path=str(path),
        create=True,
        workbook_spec={
            "sheets": [{
                "name": "Data",
                "dimensions": {"rows": 2, "cols": 1},
                "formula_blocks": [{"start": "A1", "formulas": [["=1+1"], ["=2+2"]]}],
            }],
            "uncertainties": [],
        },
    )
    assert result.success, result.model_text
    result = workbook_tools.apply_spreadsheet_changes(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "format", "sheet": "Data", "range": "A1:A2",
                     "font": {"bold": True}}],
    )
    assert result.success, result.model_text
    observation = result.value["observation"]
    assert observation["calculation_inputs"] == "unchanged"
    assert observation["formula_cache"].startswith("invalidated")
