"""工具契约第 4 批：SDK/schema 别名、矩形 round-trip、批次失败不改版本。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excelmanus.code_mode import render_sdk_source
from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.workbook_tools import (
    analyze_spreadsheet,
    compare_spreadsheets,
    apply_spreadsheet_changes,
    apply_spreadsheet_changes,
    get_tools,
    init_guard,
    observe_spreadsheet,
)
from excelmanus.tools.registry import ToolRegistry, normalize_tool_aliases
from tests.workbook_support import region_matrix
from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    init_guard(workspace)
    seed_seen_versions({})


def _err(result: ToolResult) -> str:
    if result.error is not None:
        return f"{result.error.code} {result.error.message}"
    if isinstance(result.value, dict):
        return str(result.value)
    return result.model_text


def _book(path: Path, *, title: str = "Sheet1") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws["A1"] = "id"
    ws["B1"] = "name"
    ws["C1"] = "qty"
    ws["A2"] = 1
    ws["B2"] = "a"
    ws["C2"] = 10
    ws["A3"] = 2
    ws["B3"] = "b"
    ws["C3"] = None
    wb.save(path)
    wb.close()
    return path


def test_sdk_exposes_canonical_names_not_aliases() -> None:
    source = render_sdk_source(get_tools())
    assert "def observe_spreadsheet(" in source
    assert "file_path" in source
    assert "def observe_spreadsheet(path" not in source
    assert "content_version=" not in source or "expected_version" in source
    assert "首次用某工具前可 introspect_capability" in source
    assert "operations.kind:" in source or "mode:" in source


def test_path_alias_is_rejected_by_v2_contract(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "book.xlsx")
    registry = ToolRegistry()
    registry.register_tools(get_tools())
    result = registry.call_tool(
        "observe_spreadsheet",
        {"mode": "overview", "path": str(path)},
    )
    assert isinstance(result, ToolResult)
    assert not result.success
    assert result.error and result.error.code == "TOOL_ARGUMENT_VALIDATION_ERROR"


def test_conflicting_aliases_are_rejected() -> None:
    result = normalize_tool_aliases({"path": "a.xlsx", "file_path": "b.xlsx"})
    assert isinstance(result, ToolResult)
    assert not result.success


def test_unknown_nested_kind_is_invalid(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "book.xlsx")
    result = apply_spreadsheet_changes(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "glow", "sheet": "Sheet1", "range": "A1"}],
    )
    assert not result.success
    assert "kind" in _err(result)


def test_read_filter_write_same_sheet(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "book.xlsx")
    inspected = observe_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet="Sheet1",
        range="A1:C3",
    )
    assert inspected.success, _err(inspected)
    assert inspected.value["regions"][0]["sheet"] == "Sheet1"
    filtered = analyze_spreadsheet(
        mode="filter",
        file_path=str(path),
        sheet="Sheet1",
        column="id",
        operator="eq",
        value=1,
    )
    assert filtered.success, _err(filtered)
    assert filtered.value.get("resolved_sheet") == "Sheet1"
    written = apply_spreadsheet_changes(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "start_cell": "A4",
            "values": [[3, "c", 30]],
        }],
    )
    assert written.success, _err(written)


def test_range_read_write_keeps_trailing_nulls(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "rect.xlsx")
    read = observe_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet="Sheet1",
        range="A2:C3",
    )
    assert read.success, _err(read)
    rows = region_matrix(read.value["regions"][0])
    assert rows[1] == [2, "b", None]
    written = apply_spreadsheet_changes(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "start_cell": "E2",
            "values": rows,
        }],
    )
    assert written.success, _err(written)
    wb = load_workbook(path)
    assert [wb.active.cell(2, c).value for c in range(5, 8)] == [1, "a", 10]
    assert wb.active["G3"].value is None
    wb.close()


def test_inspect_style_subset_then_format(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "styled.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "t"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FF0000")
    wb.save(path)
    wb.close()
    overview = observe_spreadsheet(
        mode="overview",
        file_path=str(path),
        facets=['presentation'],
    )
    assert overview.success, _err(overview)
    style = overview.value["regions"][0]["cells"]["1,1"]["s"]
    font = {"name":style["ff"],"size":style["fs"],"bold":bool(style["bl"]),"color":style["cl"]["rgb"]}
    formatted = apply_spreadsheet_changes(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "format",
            "sheet": "Sheet1",
            "range": "A1",
            "font": {key: font[key] for key in font if key in {"name", "size", "bold", "color"}},
        }],
    )
    assert formatted.success, _err(formatted)


def test_batch_second_item_failure_leaves_version(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(tmp_path / "batch.xlsx")
    before = content_version_of_file(path)
    result = apply_spreadsheet_changes(
        file_path=str(path),
        expected_version=before,
        operations=[
            {"kind": "write", "sheet": "Sheet1", "start_cell": "A1", "values": [["x"]]},
            {"kind": "insert", "sheet": "Missing", "axis": "row", "at": 1, "count": 1},
        ],
    )
    assert not result.success
    assert content_version_of_file(path) == before
    wb = load_workbook(path)
    assert wb.active["A1"].value == "id"
    wb.close()


def test_compare_swapped_rows_and_style_only(tmp_path: Path) -> None:
    _bind(tmp_path)
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    wb = Workbook()
    wb.active.append(["id", "v"])
    wb.active.append([1, 2])
    wb.active.append([3, 4])
    wb.save(a)
    wb.close()
    wb = Workbook()
    wb.active.append(["id", "v"])
    wb.active.append([3, 4])
    wb.active.append([1, 2])
    wb.save(b)
    wb.close()
    swapped = compare_spreadsheets(file_a=str(a), file_b=str(b), alignment="position")
    assert swapped.success, _err(swapped)
    assert swapped.value["summary"]["cells_different"] > 0

    styled = tmp_path / "styled.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A1"].font = Font(bold=True)
    wb.save(styled)
    wb.close()
    same = compare_spreadsheets(file_a=str(a), file_b=str(a), ignore_style=False)
    assert same.success, _err(same)
    assert "appearance" in same.value


def test_partial_formula_scan_has_coverage(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "f.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "=1+1"
    wb.save(path)
    wb.close()
    result = observe_spreadsheet(
        mode="overview",
        file_path=str(path),
        facets=['data'],
    )
    assert result.success, _err(result)
    region = result.value["regions"][0]
    assert region["coverage"]["data"]["status"] == "complete"
    assert region["cells"]["1,1"]["f"] == "=1+1"
    assert region["cells"]["1,1"]["cached"] == "no"
