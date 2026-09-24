"""第 3 批验收：WorkbookSnapshot、selection 写入、旧路径删除。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from excelmanus.engine_core.spill import retrieve_spill_result
from excelmanus.engine_core.tool_result import ToolResult, finalize_content
from excelmanus.mentions.parser import Mention
from excelmanus.mentions.resolver import MentionResolver
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes, init_guard
from excelmanus.workbook.data import (
    filter_data,
    init_guard as init_data_guard,
    inspect_excel_files,
    search_excel_values,
)
from excelmanus.workbook.snapshot import open_snapshot
from excelmanus.workbook_commit import content_version_of_file, remember_content_version, seed_seen_versions


def _bind(root: Path) -> FileAccessGuard:
    guard = FileAccessGuard(str(root))
    set_guard(guard)
    init_guard(str(root))
    init_data_guard(str(root))
    seed_seen_versions({})
    return guard


def _people_book(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Name"
    ws["B1"] = "Amount"
    ws["A2"] = "Alice"
    ws["B2"] = 10
    ws["A3"] = "Bob"
    ws["B3"] = 20
    ws["A4"] = "Carol"
    ws["B4"] = 5
    ws["A5"] = "Bob"
    ws["B5"] = 30
    wb.save(path)
    wb.close()
    return path


def _title_grid(path: Path, *, sep: str | None = None) -> Path:
    rows = ["Report,,", "Name,Amount,Code", "Alice,10,001", "Bob,20,002"]
    if path.suffix.lower() == ".xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Report", None, None])
        ws.append(["Name", "Amount", "Code"])
        ws.append(["Alice", 10, "001"])
        ws.append(["Bob", 20, "002"])
        wb.save(path)
        wb.close()
        return path
    text = "\n".join(rows) + "\n"
    if sep == "\t" or path.suffix.lower() == ".tsv":
        text = text.replace(",", "\t")
    path.write_text(text, encoding="utf-8")
    return path


def test_filter_selection_write_and_discontiguous_delete(tmp_path: Path) -> None:
    _bind(tmp_path)
    _people_book(tmp_path / "people.xlsx")
    filtered = filter_data(
        file_path="people.xlsx",
        sheet_name="Sheet1",
        column="Name",
        operator="eq",
        value="Bob",
        sort_by="Amount",
        ascending=False,
    )
    assert filtered.success, filtered.model_text
    payload = filtered.value
    assert payload["source_rows"] == [5, 3]
    assert payload["result_kind"] == "records"
    selection = payload["selection"]
    assert selection["content_version"]
    records_write = apply_spreadsheet_changes(
        file_path="people.xlsx",
        expected_version=selection["content_version"],
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "start_cell": "A3",
            "values": payload["data"],
        }],
    )
    assert not records_write.success
    assert records_write.error is not None
    assert records_write.error.code == "INVALID_ARGS"

    written = apply_spreadsheet_changes(
        file_path="people.xlsx",
        expected_version=selection["content_version"],
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "selection": selection,
            "values": [["Bob", 99], ["Bob", 20]],
        }],
    )
    assert written.success, written.model_text
    wb_after = load_workbook(tmp_path / "people.xlsx", data_only=False)
    assert wb_after.active["B5"].value in (99, "99")
    assert wb_after.active["B3"].value in (20, "20")
    wb_after.close()

    deleted = apply_spreadsheet_changes(
        file_path="people.xlsx",
        expected_version=written.value["content_version"],
        operations=[{
            "kind": "delete_rows",
            "selection": {
                **selection,
                "content_version": written.value["content_version"],
                "rows": [5, 3],
            },
        }],
    )
    assert deleted.success, deleted.model_text
    wb_del = load_workbook(tmp_path / "people.xlsx", data_only=False)
    names = [
        wb_del.active.cell(row=r, column=1).value
        for r in range(1, (wb_del.active.max_row or 0) + 1)
        if wb_del.active.cell(row=r, column=1).value not in (None, "")
    ]
    wb_del.close()
    assert names == ["Name", "Alice", "Carol"]


def test_selection_stale_rejects_peek_seen_rebase(tmp_path: Path) -> None:
    _bind(tmp_path)
    _people_book(tmp_path / "people.xlsx")
    filtered = filter_data(
        file_path="people.xlsx",
        sheet_name="Sheet1",
        column="Name",
        operator="eq",
        value="Bob",
    )
    selection = filtered.value["selection"]
    old_version = selection["content_version"]

    wb = load_workbook(tmp_path / "people.xlsx")
    wb.active.insert_rows(2)
    wb.save(tmp_path / "people.xlsx")
    wb.close()
    remember_content_version("people.xlsx", content_version_of_file(tmp_path / "people.xlsx"))

    stale = apply_spreadsheet_changes(
        file_path="people.xlsx",
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "selection": selection,
            "values": [["X", 999]],
        }],
    )
    assert not stale.success
    assert stale.error is not None
    assert stale.error.code in {"SELECTION_STALE", "VERSION_CONFLICT"}

    peek_write = apply_spreadsheet_changes(
        file_path="people.xlsx",
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "source_rows": [3],
            "values": [["Hijack", 1]],
        }],
    )
    assert not peek_write.success
    assert peek_write.error is not None
    assert peek_write.error.code == "INVALID_ARGS"
    assert old_version


def test_named_start_cell_and_whole_axis_write(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "named.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "x"
    ws["B1"] = "y"
    wb.defined_names.add(DefinedName(name="MyName", attr_text="Sheet1!$B$1"))
    wb.save(path)
    wb.close()
    version = content_version_of_file(path)
    remember_content_version("named.xlsx", version)
    ok = apply_spreadsheet_changes(
        file_path="named.xlsx",
        expected_version=version,
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "start_cell": "MyName",
            "values": [[42]],
        }],
    )
    assert ok.success, ok.model_text
    wb2 = load_workbook(path, data_only=False)
    assert wb2.active["B1"].value == 42
    wb2.close()
    version2 = content_version_of_file(path)
    bad = apply_spreadsheet_changes(
        file_path="named.xlsx",
        expected_version=version2,
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "start_cell": "A:A",
            "values": [[1]],
        }],
    )
    assert not bad.success
    assert bad.error is not None
    assert bad.error.code == "REF_UNSUPPORTED"


def test_mention_named_range_does_not_expand_colon(tmp_path: Path) -> None:
    guard = _bind(tmp_path)
    path = tmp_path / "named.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "hello"
    wb.defined_names.add(DefinedName(name="MyName", attr_text="Sheet1!$A$1"))
    wb.save(path)
    wb.close()
    resolver = MentionResolver(str(tmp_path), guard)
    resolved = resolver._resolve_file(
        Mention(kind="file", value="named.xlsx", raw="@file:named.xlsx", start=0, end=16, range_spec="MyName")
    )
    assert not resolved.error, resolved.error
    assert "MyName:MyName" not in (resolved.context_block or "")
    assert "hello" in (resolved.context_block or "")


def test_inspect_files_use_public_header_row(tmp_path: Path) -> None:
    _bind(tmp_path)
    _title_grid(tmp_path / "title.xlsx")
    result = inspect_excel_files(directory=".")
    assert result.success
    item = next(f for f in result.value["files"] if f["file"] == "title.xlsx")
    assert "header_row_hint" not in item["sheets"][0]
    assert item["sheets"][0]["header_row"] == 2
    assert item.get("content_version")


def test_search_uses_snapshot_version(tmp_path: Path) -> None:
    _bind(tmp_path)
    _people_book(tmp_path / "people.xlsx")
    found = search_excel_values(file_path="people.xlsx", query="Alice")
    assert found.success, found.model_text
    version = found.value.get("content_version")
    assert version
    snap = open_snapshot("people.xlsx")
    assert version == snap.content_version


def test_finalize_does_not_keep_complete_after_truncate() -> None:
    result = ToolResult(
        success=True,
        model_text="x" * 200,
        value={"ok": True},
        coverage={"kind": "complete"},
    )
    out = finalize_content(result, max_chars=40)
    assert out.truncated
    assert out.coverage["kind"] == "truncated"


def test_spill_retrieve_not_retruncated() -> None:
    spilled = retrieve_spill_result  # type: ignore[truthy-function]
    result = ToolResult(
        success=True,
        model_text="y" * 200,
        value={"ok": True},
        coverage={"spill_retrieve": True, "kind": "complete"},
    )
    out = finalize_content(result, max_chars=40)
    assert out.coverage["kind"] == "complete"
    assert len(out.model_text) == 200
    assert spilled


def test_published_mutation_observation_is_sampled(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "people.xlsx"
    _people_book(path)
    from excelmanus.tools.workbook_tools import apply_spreadsheet_changes
    from excelmanus.workbook_commit import content_version_of_file
    result = apply_spreadsheet_changes(file_path=path.name, expected_version=content_version_of_file(path),
        operations=[{"kind":"write","sheet":"Sheet1","start_cell":"B2","values":[[10]]}])
    assert result.success, result.model_text
    assert result.value["observation"]["coverage"]["kind"] == "sampled"
    assert result.value["content_version"] == content_version_of_file(path)


def test_parallel_materialize_backing_same_digest(tmp_path: Path) -> None:
    """同一 digest 并发落盘时，Windows 不能因 os.replace 抢锁失败。"""
    from concurrent.futures import ThreadPoolExecutor, as_completed

    from excelmanus.workbook.snapshot import _materialize_backing
    from excelmanus.workbook_commit import content_version_of
    from excelmanus.workspace.refs import WorkspaceRef

    path = _people_book(tmp_path / "people.xlsx")
    data = path.read_bytes()
    version = content_version_of(data)
    workspace = WorkspaceRef.from_root(tmp_path)

    def _one() -> int:
        backing = _materialize_backing(workspace, data, ".xlsx", version)
        assert backing.blob_path is not None
        assert backing.blob_path.is_file()
        return backing.blob_path.stat().st_size

    with ThreadPoolExecutor(max_workers=8) as pool:
        futs = [pool.submit(_one) for _ in range(16)]
        sizes = []
        for fut in as_completed(futs):
            exc = fut.exception()
            assert exc is None, exc
            sizes.append(fut.result())
    assert set(sizes) == {len(data)}
