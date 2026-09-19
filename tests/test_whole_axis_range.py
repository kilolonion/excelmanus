"""整行/整列地址：inspect range、copy source_range，以及 filter 不把 A:A 当区域。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.intent_tools import (
    analyze_spreadsheet,
    edit_spreadsheet,
    init_guard,
    inspect_spreadsheet,
)
from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    init_guard(workspace)
    seed_seen_versions({})


def _err(result: ToolResult) -> str:
    if result.error is not None:
        return f"{result.error.code} {result.error.message}"
    if isinstance(result.value, dict):
        return str(result.value)
    return result.model_text


def _payload(result: ToolResult) -> dict:
    assert isinstance(result.value, dict), _err(result)
    return result.value


def _column_book(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row, value in enumerate(["部门", "销售", "研发", "财务"], start=1):
        ws.cell(row=row, column=1, value=value)
        ws.cell(row=row, column=2, value=row * 10)
    wb.save(path)
    wb.close()
    return path


def test_inspect_whole_column_clips_to_used_range(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _column_book(tmp_path / "cols.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="A:A",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    assert payload.get("resolved_range") == "A1:A4"
    assert payload.get("data") == [["部门"], ["销售"], ["研发"], ["财务"]]
    assert payload.get("start_row") == 1
    assert payload.get("end_row") == 4


def test_inspect_whole_row_clips_to_used_range(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _column_book(tmp_path / "rows.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="1:1",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    assert payload.get("resolved_range") == "A1:B1"
    assert payload.get("data") == [["部门", 10]]


def test_inspect_whole_column_truncates_over_cap(tmp_path: Path, monkeypatch) -> None:
    from excelmanus.workbook import address as address_mod

    monkeypatch.setattr(address_mod, "MAX_WHOLE_RANGE_CELLS", 2)
    _bind(tmp_path)
    path = _column_book(tmp_path / "cap.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="A:A",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    assert payload.get("resolved_range") == "A1:A2"
    assert payload.get("data") == [["部门"], ["销售"]]
    assert payload.get("is_truncated") is True


def test_inspect_multi_area_range_reads_all_regions(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _column_book(tmp_path / "multi.xlsx")
    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        sheet_name="Sheet1",
        range="A1:A2,B1:B2",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    areas = payload.get("areas") or []
    assert len(areas) == 2
    assert (areas[0].get("values") or areas[0].get("data")) == [["部门"], ["销售"]]
    assert (areas[1].get("values") or areas[1].get("data")) == [[10], [20]]


def test_copy_whole_column_clips_to_used_range(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _column_book(tmp_path / "copy.xlsx")
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "copy",
            "source_sheet": "Sheet1",
            "source_range": "A:A",
            "target_sheet": "Sheet1",
            "target_start": "C1",
        }],
    )
    assert result.success, _err(result)
    applied = _payload(result).get("applied") or []
    assert any("A1:A4" in str(item) for item in applied)
    wb = load_workbook(path)
    try:
        assert [wb.active.cell(row=r, column=3).value for r in range(1, 5)] == [
            "部门", "销售", "研发", "财务",
        ]
        assert wb.active["C5"].value is None
    finally:
        wb.close()


def test_copy_whole_row_clips_to_used_range(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _column_book(tmp_path / "copy_row.xlsx")
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "copy",
            "source_sheet": "Sheet1",
            "source_range": "1:1",
            "target_sheet": "Sheet1",
            "target_start": "A5",
        }],
    )
    assert result.success, _err(result)
    wb = load_workbook(path)
    try:
        assert wb.active["A5"].value == "部门"
        assert wb.active["B5"].value == 10
        assert wb.active["A6"].value is None
    finally:
        wb.close()


def test_filter_rejects_whole_column_address_as_column_name(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _column_book(tmp_path / "filter.xlsx")
    result = analyze_spreadsheet(
        mode="filter",
        file_path=str(path),
        sheet_name="Sheet1",
        column="A:A",
        operator="eq",
        value="销售",
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code in {"INVALID_ARGS", "NOT_FOUND"}
    assert "A:A" in _err(result)
    assert "不存在" in _err(result)


def test_filter_by_header_still_works_on_full_column_sheet(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _column_book(tmp_path / "filter_ok.xlsx")
    result = analyze_spreadsheet(
        mode="filter",
        file_path=str(path),
        sheet_name="Sheet1",
        column="部门",
        operator="eq",
        value="销售",
    )
    assert result.success, _err(result)
    payload = _payload(result)
    rows = payload.get("data") or payload.get("preview") or payload.get("rows")
    assert rows is not None
    text = result.model_text + str(payload)
    assert "销售" in text
