#!/usr/bin/env python3
"""Generate (and overwrite) the P0 fixed-scenario workbooks.

Usage:
    python tests/fixtures/scenarios/generate.py

Idempotent: re-running replaces the three input workbooks in place.
Does not call any LLM. Workbooks are kept under 50KB each.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _write_wb(path: Path, wb: Workbook) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    size = path.stat().st_size
    if size >= 50 * 1024:
        raise SystemExit(f"{path} is {size} bytes; keep each workbook under 50KB")
    print(f"wrote {path.relative_to(ROOT)} ({size} bytes)")


def generate_single_cell() -> None:
    """3-sheet book. Task: set Sheet1!B2 to 100; only that cell should change."""
    wb = Workbook()

    s1 = wb.active
    s1.title = "Sheet1"
    s1["A1"] = "项目"
    s1["B1"] = "数值"
    s1["A2"] = "目标格"
    s1["B2"] = 10
    s1["A3"] = "旁格"
    s1["B3"] = 20
    s1["C1"] = "备注"
    s1["C2"] = "不要改这一列"

    s2 = wb.create_sheet("Sheet2")
    s2["A1"] = "类别"
    s2["B1"] = "数量"
    s2["A2"] = "库存"
    s2["B2"] = 42
    s2["A3"] = "在途"
    s2["B3"] = 7

    s3 = wb.create_sheet("Sheet3")
    s3["A1"] = "说明"
    s3["A2"] = "此表应保持原样"
    s3["B2"] = "sentinel"

    _write_wb(ROOT / "single_cell" / "input.xlsx", wb)


def generate_sparse_template() -> None:
    """Sparse template with merged titles and large empty regions.

    Task: fill exactly four cells — C8, D8, C12, E20.
    Expectation: no fabricated values in blank cells; blanks must not error.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "报销模板"

    title_fill = PatternFill("solid", fgColor="1F4E79")
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    subtitle_fill = PatternFill("solid", fgColor="D6DCE4")
    header_fill = PatternFill("solid", fgColor="2E75B6")
    header_font = Font(bold=True, color="FFFFFF")
    gray_fill = PatternFill("solid", fgColor="F2F2F2")
    target_fill = PatternFill("solid", fgColor="FFF2CC")

    ws.merge_cells("A1:H1")
    ws["A1"] = "2026年度费用报销模板"
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = "灰色结构区请勿改动；仅填写黄色标注的 4 个单元格"
    ws["A2"].fill = subtitle_fill
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A4:C4")
    ws["A4"] = "部门信息（结构区）"
    ws["A4"].font = Font(bold=True)
    ws["A4"].fill = gray_fill

    headers = ["序号", "科目", "金额", "币种", "经办人", "日期", "备注", "附件号"]
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(5, col, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Stretch used range with empty cells so readers must tolerate sparsity.
    for row in range(6, 61):
        for col in range(1, 9):
            cell = ws.cell(row, col, None)
            cell.border = THIN
            if col == 1:
                cell.value = row - 5  # 序号，结构列，不应被任务改写

    # Four target cells left blank on purpose.
    for coord in ("C8", "D8", "C12", "E20"):
        ws[coord].fill = target_fill
        ws[coord].value = None

    # A second sheet: merged section title + empty body (must stay empty).
    notes = wb.create_sheet("填写说明")
    notes.merge_cells("A1:D1")
    notes["A1"] = "本页无需填写"
    notes["A1"].font = Font(bold=True, size=14)
    notes.merge_cells("A3:D5")
    notes["A3"] = "空白与合并单元格不是错误。不要向未指定的格子填造数据。"
    notes["A3"].alignment = Alignment(wrap_text=True, vertical="top")

    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 12

    _write_wb(ROOT / "sparse_template" / "input.xlsx", wb)


def generate_key_match() -> None:
    """Two-sheet key match. Prices sheet is shuffled, has 1 duplicate and extras.

    Task: fill Master.Price from Prices by ID.
    Expectation: report duplicate key A002 and unmatched master keys A003, A005.
    """
    wb = Workbook()

    master = wb.active
    master.title = "主表"
    master["A1"] = "ID"
    master["B1"] = "名称"
    master["C1"] = "价格"
    for cell in ("A1", "B1", "C1"):
        master[cell].font = Font(bold=True)

    master_rows = [
        ("A001", "苹果", None),
        ("A002", "香蕉", None),
        ("A003", "橙子", None),  # unmatched
        ("A004", "葡萄", None),
        ("A005", "西瓜", None),  # unmatched
        ("A006", "梨", None),
    ]
    for i, (key, name, price) in enumerate(master_rows, start=2):
        master.cell(i, 1, key)
        master.cell(i, 2, name)
        master.cell(i, 3, price)

    prices = wb.create_sheet("价格表")
    prices["A1"] = "ID"
    prices["B1"] = "价格"
    prices["A1"].font = Font(bold=True)
    prices["B1"].font = Font(bold=True)

    # Shuffled vs master order; A002 duplicated; B999 not in master.
    price_rows = [
        ("A004", 12.5),
        ("A001", 8.0),
        ("A002", 3.5),
        ("A002", 4.0),  # duplicate key
        ("A006", 6.2),
        ("B999", 99.0),  # extra key, not in master
    ]
    for i, (key, price) in enumerate(price_rows, start=2):
        prices.cell(i, 1, key)
        prices.cell(i, 2, price)

    _write_wb(ROOT / "key_match" / "input.xlsx", wb)


def main() -> None:
    generate_single_cell()
    generate_sparse_template()
    generate_key_match()


if __name__ == "__main__":
    main()
