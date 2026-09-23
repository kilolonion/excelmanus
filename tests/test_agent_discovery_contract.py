"""Discovery must reveal executable examples and preserve the real contract."""

from __future__ import annotations

import copy
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from excelmanus.config import ExcelManusConfig
from excelmanus.engine import AgentEngine
from excelmanus.engine_core.meta_tools import MetaToolBuilder
from excelmanus.tools.context import (
    CallerCapability,
    ToolCallContext,
    bind_call,
    binding_from_engine,
    reset_call,
)
from excelmanus.tools.registry import ToolDef, ToolRegistry
from excelmanus.tools.schema_walk import walk_schema_path


@pytest.fixture
def engine(tmp_path: Path) -> AgentEngine:
    registry = ToolRegistry()
    registry.register_builtin_tools(str(tmp_path))
    registry.configure_schema_validation(mode="enforce", canary_percent=100, strict_path=False)
    return AgentEngine(ExcelManusConfig(
        api_key="test", base_url="https://test.invalid/v1", model="test",
        workspace_root=str(tmp_path), jev_enabled="off",
    ), registry)


def _schemas(engine: AgentEngine, **kwargs) -> dict[str, dict]:
    return {row["function"]["name"]: row["function"]
            for row in MetaToolBuilder(engine).build_v5_tools(**kwargs)}


def _detail(engine: AgentEngine, query: str) -> str:
    _schemas(engine)
    token = bind_call(ToolCallContext(
        binding=binding_from_engine(engine), tool_name="introspect_capability",
        loaded_tool_names=engine._loaded_tool_names,
    ))
    try:
        result = engine.registry.call_tool("introspect_capability", {
            "query_type": "tool_detail", "query": query,
        })
        assert result.success, result.model_text
        return result.model_text
    finally:
        reset_call(token)


def _detail_node(text: str) -> dict:
    # Only the JSON node is a contract. Prose and whitespace may evolve.
    start = text.index("\n{") + 1
    node, _ = json.JSONDecoder().raw_decode(text[start:])
    return node


@pytest.mark.asyncio
async def test_first_schema_example_executes_without_discovery(
    engine: AgentEngine, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    initial = _schemas(engine)
    spec_schema = initial["edit_spreadsheet"]["parameters"]["properties"]["workbook_spec"]
    examples = spec_schema.get("examples")
    assert isinstance(examples, list) and examples
    example = copy.deepcopy(examples[0])
    assert example["uncertainties"] == []
    sheet = example["sheets"][0]
    assert {"name", "dimensions", "value_blocks", "formula_blocks", "styles",
            "style_regions", "merged_ranges", "print_layout"} <= sheet.keys()

    def no_extra_discovery(**kwargs):
        raise AssertionError("The first disclosed example must execute without another schema query")

    monkeypatch.setattr(engine.registry.get_tool("introspect_capability"), "func", no_extra_discovery)
    arguments = {"file_path": "receipt.xlsx", "workbook_spec": example}
    tool_call = SimpleNamespace(id="create", function=SimpleNamespace(
        name="edit_spreadsheet", arguments=json.dumps(arguments, ensure_ascii=False),
    ))
    result = await engine._tool_runtime.execute(tool_call, None, None, 1)
    assert result.success, result.result
    assert result.structured.value["content_version"]
    wb = load_workbook(tmp_path / "receipt.xlsx", data_only=False)
    try:
        ws = wb[sheet["name"]]
        assert ws["B3"].value == 2 and ws["C3"].value == 10
        assert ws["D3"].value == "=B3*C3"
        assert ws["D4"].value == "=SUM(D3:D3)"
        assert "A1:D1" in {str(r) for r in ws.merged_cells.ranges}
        assert ws["A2"].font.bold
        assert ws.column_dimensions["A"].width == 24
        assert ws.row_dimensions[1].height == 26
        assert ws.page_setup.fitToWidth == ws.page_setup.fitToHeight == 1
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True
        assert "$A$1:$D$4" in str(ws.print_area)
    finally:
        wb.close()
    assert examples[0] == example  # Native execution cannot mutate the disclosed schema.


@pytest.mark.parametrize("mode,access", [("write", "may_write"), ("read", "may_write"),
                                          ("plan", "may_write"), ("write", "read_only")])
def test_default_file_readers_are_visible_without_widening_modes(
    engine: AgentEngine, mode: str, access: str,
) -> None:
    engine._current_chat_mode = mode
    visible = _schemas(engine, tool_access=access)
    assert {"read_image", "read_text_file"} <= visible.keys()
    assert "run_shell" not in visible
    if mode != "write" or access == "read_only":
        assert not {"edit_spreadsheet", "format_spreadsheet", "write_text_file",
                    "delete_file", "manage_spreadsheet_objects"}.intersection(visible)
        denied = _detail(engine, "edit_spreadsheet.workbook_spec") if mode != "write" else None
        if denied is not None:
            assert "不可用" in denied
    else:
        assert {"edit_spreadsheet", "format_spreadsheet"} <= visible.keys()


@pytest.mark.parametrize("allowed_readers", [frozenset(), frozenset({"read_text_file"}),
                                             frozenset({"read_image", "read_text_file"})])
def test_core_disclosure_cannot_grant_child_read_permissions(
    engine: AgentEngine, allowed_readers: frozenset[str],
) -> None:
    engine._is_host_session = False
    engine._fixed_capability = CallerCapability(
        allowed_tools=allowed_readers | {"introspect_capability"},
    )
    visible = _schemas(engine)
    assert set(visible) == allowed_readers | {"introspect_capability"}
    for name in {"read_image", "read_text_file"} - allowed_readers:
        assert "不可用" in _detail(engine, name)
        assert name not in engine._loaded_tool_names
    assert set(_schemas(engine)) == set(visible)


def test_explicit_denial_wins_over_default_file_reader_disclosure(engine: AgentEngine) -> None:
    engine._fixed_capability = CallerCapability(disallowed_tools=frozenset({"read_image"}))
    visible = _schemas(engine)
    assert "read_text_file" in visible and "read_image" not in visible
    assert "不可用" in _detail(engine, "read_image")
    assert "read_image" not in _schemas(engine)


@pytest.mark.parametrize("path,required_keys", [
    ("$defs.StyleClass", {"font", "fill", "border", "alignment", "number_format"}),
    ("$defs.StyleClass.font", {"name", "size", "bold", "italic", "color"}),
    ("$defs.Uncertainty", {"location", "reason", "candidate_values", "confidence"}),
])
def test_definition_navigation_uses_the_live_registered_schema(
    engine: AgentEngine, path: str, required_keys: set[str],
) -> None:
    schema = engine.registry.get_tool("edit_spreadsheet").input_schema
    node, available, error = walk_schema_path(schema, path)
    assert node is not None, error
    assert required_keys <= set(available)
    text = _detail(engine, "edit_spreadsheet." + path)
    assert "字段不存在" not in text
    assert len(text) < 1800
    for key in required_keys:
        assert key in text


def test_unknown_definition_reports_valid_names(engine: AgentEngine) -> None:
    text = _detail(engine, "edit_spreadsheet.$defs.MissingStyle")
    assert "字段不存在" in text
    assert "StyleClass" in text and "FontSpec" in text
    assert len(text) < 1500


def test_creation_field_query_keeps_required_inputs_and_executable_example(engine: AgentEngine) -> None:
    text = _detail(engine, "edit_spreadsheet.workbook_spec")
    node = _detail_node(text)
    assert {"sheets", "uncertainties"} <= set(node["required"])
    assert {"object", "string"} <= set(node["type"])
    assert node["examples"][0] == _schemas(engine)["edit_spreadsheet"]["parameters"]["properties"]["workbook_spec"]["examples"][0]
    assert "location" in text and "reason" in text
    assert len(text) < 5500
    assert "Python SDK" not in text and "常见错误" not in text


@pytest.mark.parametrize("path,type_name,field", [
    ("$defs.StyleClass.font", "StyleClass", "font"),
    ("workbook_spec.sheets.print_layout", "SheetSpec", "print_layout"),
    ("workbook_spec.sheets.styles.header.font.bold", "FontSpec", "bold"),
])
def test_compact_leaf_details_preserve_nullable_raw_schema(
    engine: AgentEngine, path: str, type_name: str, field: str,
) -> None:
    text = _detail(engine, "edit_spreadsheet." + path)
    node = _detail_node(text)
    original = engine.registry.get_tool("edit_spreadsheet").input_schema["$defs"][type_name]["properties"][field]
    assert node["anyOf"] == original["anyOf"]
    assert node["default"] is None
    assert len(text) < 2000
    assert "Python SDK" not in text and "常见错误" not in text


@pytest.mark.parametrize("path,type_name", [("$defs.Uncertainty", "Uncertainty"),
                                           ("$defs.FormulaBlock", "FormulaBlock")])
def test_compact_object_details_preserve_required_fields(
    engine: AgentEngine, path: str, type_name: str,
) -> None:
    text = _detail(engine, "edit_spreadsheet." + path)
    node = _detail_node(text)
    original = engine.registry.get_tool("edit_spreadsheet").input_schema["$defs"][type_name]
    assert node["required"] == original["required"]
    assert set(node["required"]) <= node["properties"].keys()
    assert len(text) < 3000


def test_named_nullable_definition_keeps_null_branch(engine: AgentEngine) -> None:
    nullable = {"anyOf": [
        {"type": "object", "properties": {"label": {"type": "string"}}, "required": ["label"]},
        {"type": "null"},
    ]}
    engine.registry.register_tool(ToolDef(
        name="nullable_contract", description="Nullable definition lookup", write_effect="none",
        input_schema={"type": "object", "$defs": {"NullableLabel": nullable},
                      "properties": {"label": {"$ref": "#/$defs/NullableLabel"}}},
        func=lambda **kwargs: kwargs,
    ))
    node = _detail_node(_detail(engine, "nullable_contract.$defs.NullableLabel"))
    assert node["anyOf"] == nullable["anyOf"]


def test_compact_output_leaf_preserves_nullable(engine: AgentEngine) -> None:
    raw = {"anyOf": [{"type": "string"}, {"type": "null"}]}
    engine.registry.register_tool(ToolDef(
        name="output_contract", description="Nullable output lookup", write_effect="none",
        input_schema={"type": "object", "properties": {}},
        output_schema={"type": "object", "properties": {"value": raw}, "required": ["value"]},
        func=lambda: {"value": None},
    ))
    text = _detail(engine, "output_contract.output.value")
    assert _detail_node(text)["anyOf"] == raw["anyOf"]
    assert len(text) < 600
