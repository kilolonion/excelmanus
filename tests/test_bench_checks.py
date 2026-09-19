"""bench_checks 新加固选项的单元测试：charts.min_series / file_exists.min_rows。"""

from __future__ import annotations

from pathlib import Path

from excelmanus.bench_checks import load_answers, run_output_checks


def _result(reply: str = "") -> dict:
    return {"result": {"reply": reply}}


def _make_workbook(path: Path, n_rows: int) -> Path:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["省份", "金额"])
    for i in range(n_rows):
        ws.append([f"省{i:02d}", i * 10])
    wb.save(path)
    wb.close()
    return path


class TestFileExistsMinRows:
    def test_empty_files_do_not_count(self, tmp_path):
        _make_workbook(tmp_path / "outputs" / "full.xlsx", 30)
        _make_workbook(tmp_path / "outputs" / "empty.xlsx", 1)
        out = run_output_checks(
            tmp_path, _result(),
            [{"type": "file_exists", "glob": "outputs/**/*.xlsx", "min": 2, "min_rows": 20}],
            {},
        )
        assert len(out) == 1
        assert out[0].passed is False  # 只有 1 个达标

    def test_qualified_count_passes(self, tmp_path):
        _make_workbook(tmp_path / "outputs" / "a.xlsx", 30)
        _make_workbook(tmp_path / "outputs" / "b.xlsx", 25)
        out = run_output_checks(
            tmp_path, _result(),
            [{"type": "file_exists", "glob": "outputs/**/*.xlsx", "min": 2, "max": 19, "min_rows": 20}],
            {},
        )
        assert out[0].passed is True

    def test_no_min_rows_keeps_old_semantics(self, tmp_path):
        _make_workbook(tmp_path / "outputs" / "empty.xlsx", 1)
        out = run_output_checks(
            tmp_path, _result(),
            [{"type": "file_exists", "glob": "outputs/**/*.xlsx", "min": 1}],
            {},
        )
        assert out[0].passed is True


def _make_chart_workbook(path: Path, *, with_data: bool) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    for row in (["月", "金额"], ["1月", 10], ["2月", 20], ["3月", 30]):
        ws.append(row)
    chart = LineChart()
    if with_data:
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(chart, "D1")
    wb.save(path)
    wb.close()
    return path


class TestChartsMinSeries:
    def test_chart_with_data_passes(self, tmp_path):
        _make_chart_workbook(tmp_path / "outputs" / "r.xlsx", with_data=True)
        out = run_output_checks(
            tmp_path, _result(),
            [{"type": "charts", "min": 1, "types": ["line"], "min_series": 1}],
            {},
        )
        assert out[0].passed is True

    def test_empty_chart_frame_fails_min_series(self, tmp_path):
        _make_chart_workbook(tmp_path / "outputs" / "r.xlsx", with_data=False)
        strict = run_output_checks(
            tmp_path, _result(),
            [{"type": "charts", "min": 1, "types": ["line"], "min_series": 1}],
            {},
        )
        assert strict[0].passed is False
        # 不带 min_series 时保持旧语义：只验数量与类型
        loose = run_output_checks(
            tmp_path, _result(),
            [{"type": "charts", "min": 1, "types": ["line"]}],
            {},
        )
        assert loose[0].passed is True

    def test_load_answers_smoke(self):
        assert load_answers(None) == {}


class TestOutputCheckSeverityAndReplyRegex:
    def test_default_severity_is_error(self, tmp_path):
        out = run_output_checks(
            tmp_path, _result(reply="测试回复"),
            [{"type": "reply_regex", "pattern": "找不到"}],
            {},
        )
        assert len(out) == 1
        assert out[0].passed is False
        assert out[0].severity == "error"

    def test_custom_severity_warning(self, tmp_path):
        out = run_output_checks(
            tmp_path, _result(reply="测试回复"),
            [{"type": "reply_regex", "pattern": "找不到", "severity": "warning"}],
            {},
        )
        assert len(out) == 1
        assert out[0].passed is False
        assert out[0].severity == "warning"

    def test_r31_reply_regex_matches_original_file_expression(self, tmp_path):
        pattern = "恢复|还原|撤|回退|版本|原[件样状貌始来]|重置|初始|已是原"
        for phrase in [
            "内容已是完整原件",
            "已重置为原貌",
            "这是初始状态的数据",
            "已撤销修改并还原",
            "原件已确认",
        ]:
            out = run_output_checks(
                tmp_path, _result(reply=phrase),
                [{"type": "reply_regex", "pattern": pattern}],
                {},
            )
            assert out[0].passed is True, f"应当匹配: {phrase}"

