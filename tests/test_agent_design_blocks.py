"""excelmanus-agent-design.md 第 12 章各块验收。"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes, init_guard as init_intent_guard
from excelmanus.workbook_commit import content_version_of_file


def _book(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"
    ws["A1"] = "部门"
    ws["B1"] = "金额"
    wb.save(path)
    return path


def test_edit_model_text_is_bounded_and_keeps_version(tmp_path: Path) -> None:
    """第二块：写入摘要有界，value / ui_meta 仍带路径和版本。"""
    import json

    set_guard(FileAccessGuard(str(tmp_path)))
    init_intent_guard(str(tmp_path))
    path = _book(tmp_path / "sales.xlsx")
    version = content_version_of_file(path)
    result = apply_spreadsheet_changes(
        file_path=str(path),
        operations=[{"kind": "write", "sheet": "Sheet1", "start_cell": "B2", "values": [[99]]}],
        expected_version=version,
    )
    assert isinstance(result, ToolResult)
    assert result.success
    dumped = json.dumps(result.value, ensure_ascii=False)
    assert result.model_text != dumped
    assert result.ui_meta.content_version
    assert result.ui_meta.content_version == result.value["content_version"]
    assert result.ui_meta.files
    assert "__tool_result_image__" not in result.model_text


def test_existing_file_without_version_conflicts(tmp_path: Path) -> None:
    """第三块：改已有文件不带版本必冲突。"""
    set_guard(FileAccessGuard(str(tmp_path)))
    init_intent_guard(str(tmp_path))
    path = _book(tmp_path / "sales.xlsx")
    result = apply_spreadsheet_changes(
        file_path=str(path),
        operations=[{"kind": "write", "sheet": "Sheet1", "start_cell": "B2", "values": [[1]]}],
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code == "VERSION_CONFLICT"
    assert result.error.fields.get("path")
    assert result.error.fields.get("content_version")


def test_workbook_spec_creates_via_commit(tmp_path: Path) -> None:
    """第三块：workbook_spec 新建，磁盘版本来自 commit_*。"""
    set_guard(FileAccessGuard(str(tmp_path)))
    init_intent_guard(str(tmp_path))
    dest = tmp_path / "outputs" / "from_spec.xlsx"
    dest.parent.mkdir(parents=True, exist_ok=True)
    spec = {
        "name": "from_spec",
        "sheets": [{
            "name": "Sheet1",
            "dimensions": {"rows": 2, "cols": 2},
            "value_blocks": [{"start": "A1", "values": [["Name", "Age"], ["Ada", 1]]}],
            "formula_blocks": [],
            "styles": {},
            "style_regions": [],
            "merged_ranges": [],
        }],
        "uncertainties": [],
    }
    result = apply_spreadsheet_changes(
        file_path=str(dest),
        workbook_spec=spec,
        create=True,
    )
    assert result.success, result.model_text
    assert dest.is_file()
    assert result.ui_meta.content_version == content_version_of_file(dest)
    assert result.value["document"].get("uncertainties") == []
    wb = load_workbook(dest)
    assert wb["Sheet1"]["A1"].value == "Name"


def test_spec_rejects_existing_path(tmp_path: Path) -> None:
    set_guard(FileAccessGuard(str(tmp_path)))
    init_intent_guard(str(tmp_path))
    path = _book(tmp_path / "exists.xlsx")
    result = apply_spreadsheet_changes(
        file_path=str(path),
        workbook_spec={
            "name": "x",
            "sheets": [{
                "name": "S",
                "dimensions": {"rows": 1, "cols": 1},
                "value_blocks": [{"start": "A1", "values": [["1"]]}],
            }],
            "uncertainties": [],
        },
        create=True,
    )
    assert not result.success
    assert result.error and result.error.code == "VERSION_CONFLICT"


@pytest.mark.asyncio
async def test_chat_does_not_open_xlsx_before_first_model(tmp_path: Path) -> None:
    """第一块：第一次模型请求之前不读该 xlsx。"""
    from excelmanus.config import ExcelManusConfig
    from excelmanus.engine import AgentEngine, ChatResult
    from excelmanus.tools import ToolRegistry

    sales = _book(tmp_path / "sales.xlsx")
    opened: list[str] = []
    real_load = load_workbook

    def _track(src, *args, **kwargs):
        opened.append(str(src))
        return real_load(src, *args, **kwargs)

    config = ExcelManusConfig(
        api_key="test",
        base_url="https://example.test/v1",
        model="test-model",
        workspace_root=str(tmp_path),
        main_model_vision="false",
    )
    engine = AgentEngine(config, ToolRegistry())
    engine.start_registry_scan = lambda **_k: False  # type: ignore[method-assign]

    async def _loop(*_a, **_k):
        assert sales.resolve() not in {Path(p).resolve() for p in opened}
        return ChatResult(reply="先 observe_spreadsheet")

    with patch("openpyxl.load_workbook", side_effect=_track), patch(
        "excelmanus.agent.loop.run_tool_loop", new=_loop
    ):
        result = await engine.followup("看一下 sales.xlsx")
    assert "先 observe_spreadsheet" in result.reply
    assert not any(Path(p).name == "sales.xlsx" for p in opened)


@pytest.mark.asyncio
async def test_pending_approval_does_not_block_unrelated_chat(tmp_path: Path) -> None:
    """第一块：有 pending 审批不整段挡住无关消息。"""
    from excelmanus.config import ExcelManusConfig
    from excelmanus.engine import AgentEngine, ChatResult
    from excelmanus.tools import ToolRegistry

    config = ExcelManusConfig(
        api_key="test",
        base_url="https://example.test/v1",
        model="test-model",
        workspace_root=str(tmp_path),
    )
    engine = AgentEngine(config, ToolRegistry())
    engine._approval.has_pending = lambda: True  # type: ignore[method-assign]
    engine._approval.pending_block_message = lambda: "不应出现"  # type: ignore[method-assign]
    with patch(
        "excelmanus.agent.loop.run_tool_loop",
        new_callable=AsyncMock,
        return_value=ChatResult(reply="继续"),
    ):
        result = await engine.followup("先看另一张表")
    assert result.reply == "继续"
    assert "不应出现" not in result.reply


def test_interrupt_queue_drains_on_next_step(tmp_path: Path) -> None:
    """飞行中 followup 进 next-turn；drain 兼容旧 API。"""
    from excelmanus.config import ExcelManusConfig
    from excelmanus.engine import AgentEngine
    from excelmanus.tools import ToolRegistry

    config = ExcelManusConfig(
        api_key="test",
        base_url="https://example.test/v1",
        model="test-model",
        workspace_root=str(tmp_path),
    )
    engine = AgentEngine(config, ToolRegistry())
    engine.push_interrupt_message("改用 B 表")
    assert engine.drain_interrupt_messages() == ["改用 B 表"]
    assert engine.drain_interrupt_messages() == []


@pytest.mark.asyncio
async def test_enqueue_interrupt_when_in_flight(tmp_path: Path) -> None:
    from excelmanus.config import ExcelManusConfig
    from excelmanus.session import SessionManager

    config = ExcelManusConfig(
        api_key="test",
        base_url="https://example.test/v1",
        model="test-model",
        workspace_root=str(tmp_path),
    )
    manager = SessionManager(max_sessions=4, ttl_seconds=60, config=config, registry=MagicMock())
    session_id, engine = await manager.acquire_for_chat(None)
    queued = await manager.enqueue_user_interrupt(session_id, "改用 B 表")
    assert queued is True
    assert engine.drain_interrupt_messages() == ["改用 B 表"]
    await manager.release_for_chat(session_id)


def test_code_mode_disclaimer_and_parent_call_field() -> None:
    """第七块：本机围栏声明；事件带 parent_call_id。"""
    from excelmanus.code_mode import LOCAL_SANDBOX_DISCLAIMER
    from excelmanus.events import EventType, ToolCallEvent

    assert "本机" in LOCAL_SANDBOX_DISCLAIMER
    assert "禁网络" in LOCAL_SANDBOX_DISCLAIMER
    assert "禁起进程" in LOCAL_SANDBOX_DISCLAIMER
    assert "禁出工作区" in LOCAL_SANDBOX_DISCLAIMER
    event = ToolCallEvent(
        event_type=EventType.TOOL_CALL_START,
        tool_call_id="child",
        tool_name="apply_spreadsheet_changes",
        parent_call_id="parent-run-code",
    )
    assert event.parent_call_id == "parent-run-code"


def test_cancelled_parent_blocks_new_sub_write(tmp_path: Path) -> None:
    """第七块：父 run_code 取消后不再出现新的子写入。"""
    from excelmanus.code_mode import CodeModeSession

    dispatcher = MagicMock()
    dispatcher.is_cancelled.return_value = True
    session = CodeModeSession(
        dispatcher=dispatcher,
        root_call_id="parent-run-code",
        bridge_dir=tmp_path / "bridge",
    )
    result = session._call_dispatcher(
        tool_name="apply_spreadsheet_changes",
        arguments={"file_path": "a.xlsx", "operations": []},
        root_call_id="parent-run-code",
    )
    assert result.success is False
    assert result.error is not None
    assert result.error.code == "CANCELLED"
    dispatcher.execute_subcall.assert_not_called()
    dispatcher.call_registry_tool.assert_not_called()


def test_workspace_resolve_is_single_root(tmp_path) -> None:
    """第五块：创建会话不再写出 users/<id>/... 根。"""
    from excelmanus.workspace import IsolatedWorkspace

    ws_a = IsolatedWorkspace.resolve(str(tmp_path), data_root="")
    ws_b = IsolatedWorkspace.resolve(str(tmp_path), data_root="")
    assert ws_a.root_dir == ws_b.root_dir
    assert ws_a.root_dir == tmp_path.resolve()
    assert not (tmp_path / "users").exists()
