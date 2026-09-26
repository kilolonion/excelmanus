"""WorkbookSpec 校验与私有编译辅助。"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from excelmanus.workbook.spec import SpecValidationError, validate_workbook_spec


def _minimal_workbook_spec(**overrides: object) -> dict:
    spec: dict = {
        "name": "replica",
        "sheets": [{
            "name": "Sheet1",
            "dimensions": {"rows": 2, "cols": 2},
            "value_blocks": [{
                "start": "A1",
                "values": [["Name", "Age"], ["Alice", 30]],
            }],
            "formula_blocks": [],
            "styles": {
                "header": {"font": {"bold": True}},
            },
            "style_regions": [{"range": "A1:B1", "style_id": "header"}],
            "merged_ranges": [],
            "column_widths": [12, 8],
            "row_heights": {"1": 20},
        }],
        "uncertainties": [],
    }
    spec.update(overrides)
    return spec


class TestWorkbookSpecValidation:
    def test_missing_uncertainties_reports_path(self) -> None:
        data = _minimal_workbook_spec()
        del data["uncertainties"]
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        paths = [e["path"] for e in excinfo.value.errors]
        assert "uncertainties" in paths
        payload = excinfo.value.to_payload()
        assert payload["status"] == "error"
        assert payload["error_code"] == "SPEC_VALIDATION_FAILED"
        assert any(e["path"] == "uncertainties" for e in payload["errors"])

    def test_empty_uncertainties_ok(self) -> None:
        spec = validate_workbook_spec(_minimal_workbook_spec())
        assert spec.uncertainties == []
        assert spec.sheets[0].value_blocks[0].start == "A1"

    def test_out_of_bounds_value_block(self) -> None:
        data = _minimal_workbook_spec()
        data["sheets"][0]["value_blocks"] = [{
            "start": "A1",
            "values": [["a", "b", "c"], ["d", "e", "f"]],
        }]
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        err = excinfo.value.errors[0]
        assert err["path"] == "sheets.0.value_blocks.0"
        assert "超出" in err["message"]

    def test_out_of_bounds_formula_block(self) -> None:
        data = _minimal_workbook_spec()
        data["sheets"][0]["formula_blocks"] = [{
            "start": "A3",
            "formulas": [["=A1+A2"]],
        }]
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        assert any("formula_blocks" in e["path"] for e in excinfo.value.errors)

    def test_bad_style_region_ref(self) -> None:
        data = _minimal_workbook_spec()
        data["sheets"][0]["style_regions"] = [
            {"range": "A1:B1", "style_id": "missing_style"},
        ]
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        err = next(e for e in excinfo.value.errors if e["path"].endswith("style_id"))
        assert err["path"] == "sheets.0.style_regions.0.style_id"
        assert "missing_style" in err["message"]

    def test_does_not_silently_patch(self) -> None:
        data = _minimal_workbook_spec()
        del data["uncertainties"]
        with pytest.raises(SpecValidationError):
            validate_workbook_spec(data)


class TestCompileFromWorkbookSpec:
    def test_compile_writes_xlsx_from_blocks(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        from excelmanus.workbook.spec import validate_workbook_spec
        from tests.workbook_support import create_document_bytes

        spec = validate_workbook_spec(_minimal_workbook_spec())
        data, summary = create_document_bytes(spec)
        xlsx_path = tmp_path / "out.xlsx"
        xlsx_path.write_bytes(data)
        assert xlsx_path.is_file()
        assert summary["applied"].count("write") >= 4
        wb = load_workbook(str(xlsx_path))
        ws = wb["Sheet1"]
        assert ws["A1"].value == "Name"
        assert ws["B2"].value == 30
        assert ws["A1"].font.bold is True


class TestValidateWorkbookSpec:
    def test_accepts_minimal_spec(self) -> None:
        workbook = validate_workbook_spec(_minimal_workbook_spec())
        assert workbook.uncertainties == []
        assert len(workbook.sheets) == 1

    def test_validation_errors_include_field_paths(self) -> None:
        data = _minimal_workbook_spec()
        del data["uncertainties"]
        with pytest.raises(SpecValidationError) as ei:
            validate_workbook_spec(data)
        payload = ei.value.to_payload()
        assert payload["error_code"] == "SPEC_VALIDATION_FAILED"
        assert any(e["path"] == "uncertainties" for e in payload["errors"])

    def test_image_tools_only_register_read_image(self) -> None:
        from excelmanus.tools.image_tools import get_tools

        names = {t.name for t in get_tools()}
        assert names == {"read_image"}
        assert "extract_table_spec" not in names


def test_workbook_spec_commit_includes_verification(tmp_path: Path) -> None:
    from excelmanus.security import FileAccessGuard
    from excelmanus.tools import workbook_tools, reference_tools
    from excelmanus.tools._guard_ctx import set_guard
    from excelmanus.tools.workbook_tools import apply_spreadsheet_changes

    workspace = str(tmp_path)
    set_guard(FileAccessGuard(workspace))
    workbook_tools.init_guard(workspace)
    reference_tools.init_guard(workspace)
    result = apply_spreadsheet_changes(
        file_path=str(tmp_path / "created.xlsx"),
        workbook_spec=_minimal_workbook_spec(),
    )
    assert result.success, result.model_text
    observation = result.value["observation"]
    assert all(item["verified"] for item in observation["cell_checks"])
    assert all(item["verified"] for item in observation["geometry_changes"])


def test_json_string_is_rejected_without_compatibility_coercion() -> None:
    """WorkbookSpec V2 accepts an object only; callers must decode JSON first."""
    truncated = '{"name":"x","sheets":[{"name":"S1","dimensions":{"rows":2,"cols":2'
    with pytest.raises(SpecValidationError) as excinfo:
        validate_workbook_spec(truncated)
    message = str(excinfo.value)
    assert "必须是对象" in message


class TestSpecShorthandAliases:
    """任务复盘回归：收款收据还原任务中模型用同义键/简写被 SPEC_VALIDATION_FAILED 拒绝。"""

    def test_fill_accepts_changeset_style_aliases(self) -> None:
        data = _minimal_workbook_spec()
        data["sheets"][0]["styles"]["header"] = {
            "fill": {"color": "1F4E79", "fill_type": "solid"},
            "font": {"bold": True, "strikethrough": True},
        }
        spec = validate_workbook_spec(data)
        style = spec.sheets[0].styles["header"]
        assert style.fill is not None
        assert (style.fill.type, style.fill.color) == ("solid", "1F4E79")
        assert style.font is not None and style.font.strike is True

    def test_fill_accepts_pattern_type_and_fg_color(self) -> None:
        data = _minimal_workbook_spec()
        data["sheets"][0]["styles"]["header"] = {"fill": {"fgColor": "DCE6F1", "patternType": "solid"}}
        spec = validate_workbook_spec(data)
        fill = spec.sheets[0].styles["header"].fill
        assert fill is not None
        assert (fill.type, fill.color) == ("solid", "DCE6F1")

    def test_uncertainty_accepts_field_note_and_candidates(self) -> None:
        data = _minimal_workbook_spec()
        data["uncertainties"] = [{
            "field": "A1:F2 标题及表头填充色",
            "note": "图片深蓝为近似取色 #1F4E79，非精确像素取色",
            "candidates": ["#1F4E79", "#2F5F8F"],
        }]
        spec = validate_workbook_spec(data)
        item = spec.uncertainties[0]
        assert item.location == "A1:F2 标题及表头填充色"
        assert item.reason == "图片深蓝为近似取色 #1F4E79，非精确像素取色"
        assert item.candidate_values == ["#1F4E79", "#2F5F8F"]

    def test_merge_string_and_size_shorthand_match_documented_hints(self) -> None:
        data = _minimal_workbook_spec()
        sheet = data["sheets"][0]
        sheet["merged_ranges"] = ["A1:B1"]
        sheet["column_widths"] = {"A": 18, "B": 9}
        sheet["row_heights"] = [22, 15]
        spec = validate_workbook_spec(data)
        parsed = spec.sheets[0]
        assert parsed.merged_ranges[0].range == "A1:B1"
        assert parsed.column_widths == [18, 9]
        assert parsed.row_heights == {"1": 22, "2": 15}

    def test_default_font_accepts_font_name_string(self) -> None:
        data = _minimal_workbook_spec()
        data["default_font"] = "微软雅黑"
        spec = validate_workbook_spec(data)
        assert spec.default_font is not None
        assert spec.default_font.name == "微软雅黑"

    def test_extra_fill_field_error_names_legal_shape(self) -> None:
        data = _minimal_workbook_spec()
        data["sheets"][0]["styles"]["header"] = {"fill": {"color": "FF0000", "patternFillType": "solid"}}
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        err = next(e for e in excinfo.value.errors if e["path"].endswith("patternFillType"))
        assert "fill 只接受" in err["message"]
        assert "patternFillType" in err["message"]

    def test_uncertainty_missing_required_fields_is_friendly(self) -> None:
        data = _minimal_workbook_spec()
        data["uncertainties"] = [{"candidate_values": ["1200"]}]
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        paths = {e["path"]: e["message"] for e in excinfo.value.errors}
        assert "uncertainties.0.location" in paths
        assert "location" in paths["uncertainties.0.location"]
        assert "reason" in paths["uncertainties.0.location"]

    def test_remediation_preview_reports_total_error_count(self) -> None:
        data = _minimal_workbook_spec()
        data["sheets"][0]["styles"].update({
            name: {"fill": {"color": "1F4E79", "patternFillType": "solid"}}
            for name in ("title", "subtitle", "total_left", "total_amount")
        })
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        payload = excinfo.value.to_payload()
        assert len(payload["errors"]) > 3
        assert "完整列表见 errors" in payload["remediation"]
