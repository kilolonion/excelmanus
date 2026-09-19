"""模型常见 Excel 写法应被表格工具接受，而不是 SAVE_FAILED / SPEC_VALIDATION_FAILED。"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excelmanus.replica_spec import compile_workbook_spec_to_bytes, validate_workbook_spec
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.intent_tools import (
    edit_spreadsheet,
    format_spreadsheet,
    inspect_spreadsheet,
    init_guard as init_intent_guard,
)
from excelmanus.workbook.styles import _build_fill
from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions


def _bind(tmp_path: Path) -> None:
    workspace = str(tmp_path)
    set_guard(FileAccessGuard(workspace))
    init_intent_guard(workspace)
    seed_seen_versions({})


def test_format_accepts_sheet_qualified_range(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "销售明细"
    ws = wb.create_sheet("区域汇总")
    ws["A5"] = "合计"
    ws["B5"] = 1
    ws["C5"] = 0.1
    wb.save(path)
    wb.close()

    result = format_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[
            {
                "kind": "format",
                "range": "区域汇总!A5:C5",
                "font": {"bold": True},
                "fill": {"color": "DDEBF7", "pattern": "solid"},
            }
        ],
    )
    assert result.success, result.model_text
    loaded = load_workbook(path)
    try:
        cell = loaded["区域汇总"]["A5"]
        assert cell.font.bold is True
        assert "DDEBF7" in (cell.fill.start_color.rgb or "")
    finally:
        loaded.close()


def test_format_accepts_quoted_sheet_range(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "Sheet 2"
    wb.active["A1"] = "x"
    wb.save(path)
    wb.close()

    result = format_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "format", "range": "'Sheet 2'!A1", "font": {"bold": True}}],
    )
    assert result.success, result.model_text


def test_workbook_spec_accepts_font_string_and_width_map(tmp_path: Path) -> None:
    spec = validate_workbook_spec({
        "name": "演示_月度销售分析",
        "locale": "zh-CN",
        "default_font": "微软雅黑",
        "theme_hint": "business_blue",
        "uncertainties": [],
        "sheets": [{
            "name": "区域汇总",
            "dimensions": {"rows": 5, "cols": 3},
            "value_blocks": [{"start": "A1", "values": [["区域", "销售额合计", "毛利率"]]}],
            "column_widths": {"A": 12, "B": 16, "C": 11},
            "row_heights": [20, 18],
            "merged_ranges": ["A1:A1"],
            "cells": [{"address": "B2", "formula": "=1+1"}],
            "styles": {
                "header": {
                    "font": {"bold": True, "color": "FFFFFF"},
                    "fill": {"color": "4472C4", "pattern": "solid"},
                }
            },
            "style_regions": [{"range": "区域汇总!A1:C1", "style_id": "header"}],
        }],
    })
    assert spec.default_font is not None
    assert spec.default_font.name == "微软雅黑"
    assert spec.sheets[0].column_widths == [12.0, 16.0, 11.0]
    assert spec.sheets[0].row_heights["1"] == 20.0
    assert spec.sheets[0].merged_ranges[0].range == "A1"
    assert spec.sheets[0].cells[0].value == "=1+1"
    assert spec.sheets[0].cells[0].value_type == "formula"
    assert spec.sheets[0].style_regions[0].range == "A1:C1"

    data, _summary = compile_workbook_spec_to_bytes(spec)
    out = tmp_path / "out.xlsx"
    out.write_bytes(data)
    loaded = load_workbook(out)
    try:
        ws = loaded["区域汇总"]
        assert ws.column_dimensions["A"].width == 12.0
        assert ws["B2"].value == "=1+1"
        assert ws["A2"].font.name == "微软雅黑"
        assert ws["A1"].font.bold is True
        assert ws["A1"].font.name == "微软雅黑"
    finally:
        loaded.close()


def test_edit_spreadsheet_compiles_user_shaped_spec(tmp_path: Path) -> None:
    _bind(tmp_path)
    dest = tmp_path / "outputs" / "演示_月度销售分析.xlsx"
    dest.parent.mkdir(parents=True, exist_ok=True)
    result = edit_spreadsheet(
        file_path=str(dest),
        create_workbook=True,
        workbook_spec={
            "name": "演示_月度销售分析",
            "locale": "zh-CN",
            "default_font": "微软雅黑",
            "uncertainties": [],
            "sheets": [{
                "name": "区域汇总",
                "dimensions": {"rows": 5, "cols": 3},
                "value_blocks": [
                    {"start": "A1", "values": [["区域", "销售额合计", "毛利率"]]},
                    {"start": "A2", "values": [["华东"], ["华北"], ["华南"], ["合计"]]},
                ],
                "column_widths": {"A": 12, "B": 16, "C": 11},
            }],
        },
    )
    assert result.success, result.model_text
    assert dest.is_file()
    formatted = format_spreadsheet(
        file_path=str(dest),
        expected_version=result.value["content_version"],
        operations=[{
            "kind": "format",
            "range": "区域汇总!A5:C5",
            "font": {"bold": True},
            "fill": {"color": "DDEBF7", "pattern": "solid"},
        }],
    )
    assert formatted.success, formatted.model_text


def test_inspect_range_accepts_sheet_bang(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "销售明细"
    wb.active["A1"] = "月份"
    wb.save(path)
    wb.close()

    result = inspect_spreadsheet(
        mode="range",
        file_path=str(path),
        range="销售明细!A1:A1",
    )
    assert result.success, result.model_text
    payload = result.value
    assert isinstance(payload, dict)
    assert payload.get("data") == [["月份"]] or payload.get("range") == "A1:A1"


def test_write_accepts_sheet_qualified_start(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "销售明细"
    wb.active["A1"] = "x"
    wb.save(path)
    wb.close()

    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "write", "start_cell": "销售明细!B2", "values": [[99]]}],
    )
    assert result.success, result.model_text
    loaded = load_workbook(path)
    try:
        assert loaded["销售明细"]["B2"].value == 99
    finally:
        loaded.close()


def test_build_fill_accepts_pattern_and_type() -> None:
    from_pattern = _build_fill({"color": "DDEBF7", "pattern": "solid"})
    from_type = _build_fill({"color": "DDEBF7", "type": "solid"})
    assert from_pattern.patternType == "solid"
    assert from_type.patternType == "solid"
    assert from_pattern.start_color.rgb.endswith("DDEBF7")


def test_bad_range_returns_range_invalid(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(path)
    wb.close()

    result = format_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "format", "sheet": "Sheet", "range": "=A1", "font": {"bold": True}}],
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code == "RANGE_INVALID"
    assert "SAVE_FAILED" not in result.model_text
    assert "语法可解析不等于" not in result.model_text
    assert "write.start_cell" not in result.model_text


def test_conflicting_sheet_and_range_is_invalid(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "销售明细"
    wb.active["A5"] = "x"
    wb.create_sheet("区域汇总")["A5"] = "y"
    wb.save(path)
    wb.close()

    result = format_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "format",
            "sheet": "销售明细",
            "range": "区域汇总!A5:C5",
            "font": {"bold": True},
        }],
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code == "INVALID_ARGS"
    assert "不一致" in result.model_text


def test_empty_size_is_invalid_args(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(path)
    wb.close()

    result = format_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "size", "sheet": "Sheet"}],
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code == "INVALID_ARGS"


def test_inspect_range_without_mode(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "销售明细"
    wb.active["A1"] = "月份"
    wb.save(path)
    wb.close()

    result = inspect_spreadsheet(file_path=str(path), range="销售明细!A1:A1")
    assert result.success, result.model_text
    payload = result.value
    assert isinstance(payload, dict)
    assert payload.get("data") == [["月份"]] or payload.get("range") == "A1:A1"


def test_format_union_applies_same_style(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A4"] = "a"
    ws["A5"] = "skip"
    ws["A11"] = "b"
    ws["A12"] = "c"
    ws["D12"] = "d"
    wb.save(path)
    wb.close()

    cells = format_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "format", "sheet": "Sheet", "range": "A4,A11", "font": {"bold": True}}],
    )
    assert cells.success, cells.model_text
    blocks = format_spreadsheet(
        file_path=str(path),
        expected_version=cells.value["content_version"],
        operations=[{
            "kind": "format",
            "sheet": "Sheet",
            "range": "A5:D5,A12:D12",
            "font": {"italic": True},
        }],
    )
    assert blocks.success, blocks.model_text
    wb = load_workbook(path)
    try:
        assert wb.active["A4"].font.bold is True
        assert wb.active["A11"].font.bold is True
        assert wb.active["A5"].font.bold is not True
        assert wb.active["A5"].font.italic is True
        assert wb.active["D12"].font.italic is True
        assert wb.active["A11"].font.italic is not True
    finally:
        wb.close()


def test_format_merge_rejects_union(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(path)
    wb.close()

    result = format_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "merge", "sheet": "Sheet", "range": "A1:B1,C1:D1"}],
    )
    assert not result.success
    assert result.error is not None
    assert result.error.code == "REF_UNSUPPORTED"
    assert "只要一个矩形" in result.model_text
    assert "语法可解析不等于" not in result.model_text


def test_edit_sheet_create_accepts_sheet_name_and_json_operations(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(path)
    wb.close()

    created = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=json.dumps([{"kind": "sheet", "action": "create", "sheet_name": "摘要"}]),
    )
    assert created.success, created.model_text
    wb = load_workbook(path)
    try:
        assert "摘要" in wb.sheetnames
    finally:
        wb.close()

    malformed = edit_spreadsheet(
        file_path=str(path),
        expected_version=created.value["content_version"],
        operations='[{"kind": "write", "kind2": null}',
    )
    assert not malformed.success
    assert malformed.error is not None
    assert malformed.error.code == "INVALID_ARGS"
    assert "JSON 字符串" in malformed.model_text
    # 末尾截断的 JSON 给出截断识别与分块指引
    assert "截断" in malformed.model_text
