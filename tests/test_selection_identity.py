from openpyxl import Workbook, load_workbook
import pytest

from excelmanus.tools.context import bind_workspace, reset_call
from excelmanus.tools.intent_tools import edit_spreadsheet
from excelmanus.workbook.data import filter_data


@pytest.mark.parametrize("other_workspace", [False, True])
def test_same_bytes_do_not_make_selection_portable(tmp_path, other_workspace):
    root_a = tmp_path / "a"
    root_b = tmp_path / "b" if other_workspace else root_a
    root_a.mkdir()
    root_b.mkdir(exist_ok=True)
    wb = Workbook()
    wb.active.title = "Data"
    wb.active.append(["Name", "Amount"])
    wb.active.append(["Bob", 20])
    wb.save(root_a / "source.xlsx")
    target = "source.xlsx" if other_workspace else "copy.xlsx"
    (root_b / target).write_bytes((root_a / "source.xlsx").read_bytes())
    token = bind_workspace(root_a)
    try:
        result = filter_data("source.xlsx", sheet_name="Data", header_row=1, column="Name", operator="eq", value="Bob")
        selection = result.value["selection"]
    finally:
        reset_call(token)
    token = bind_workspace(root_b)
    try:
        result = edit_spreadsheet(target, operations=[{"kind": "write", "selection": selection, "values": [["Bob", 999]]}])
        assert not result.success
        assert result.error.code == "SELECTION_STALE"
        assert load_workbook(root_b / target)["Data"]["B2"].value == 20
    finally:
        reset_call(token)
