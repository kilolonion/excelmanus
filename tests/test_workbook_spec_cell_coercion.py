"""WorkbookSpec 单元格归一化回归：公式/值矩阵的 null、数字、布尔、日期与颜色写法。

这些行为在收款收据还原任务里被实测依赖：公式矩阵混排 null/数字、值矩阵混排
null/ISO 日期、颜色接受 "FF0000"/"#FF0000"/"FFFF0000" 多种写法。这里用
apply_spreadsheet_changes 全链路编译落盘后以 openpyxl 回读锁定。
"""

from __future__ import annotations

from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from excelmanus.tools.context import use_workspace
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes
from excelmanus.workbook.spec import SpecValidationError, validate_workbook_spec


def _spec(sheet_overrides: dict, **top_overrides: object) -> dict:
    sheet = {
        "name": "S",
        "dimensions": {"rows": 5, "cols": 5},
        "value_blocks": [],
        "formula_blocks": [],
    }
    sheet.update(sheet_overrides)
    spec: dict = {
        "version": "2",
        "purpose": "data",
        "sheets": [sheet],
        "uncertainties": [],
    }
    spec.update(top_overrides)
    return spec


def _day(value):
    return value.date() if isinstance(value, datetime) else value


class TestFormulaBlockCellCoercion:
    def test_null_blank_number_literal_and_formula_cells_land(self, tmp_path) -> None:
        data = _spec({
            "formula_blocks": [{"start": "A4", "formulas": [
                ["=SUM(C1:C1)", None, 5, 5.0, 5.5],
                [True, "标签", "", None, 7],
            ]}],
        })
        spec = validate_workbook_spec(data)
        parsed = spec.sheets[0].formula_blocks[0].formulas
        assert parsed[0][0] == "=SUM(C1:C1)"
        assert parsed[0][1] is None  # null 保留为 None（跳过）
        assert parsed[0][2] == "5"  # 数字转字符串字面量
        assert parsed[0][3] == "5"  # 5.0 不得留下 ".0"
        assert parsed[0][4] == "5.5"
        assert parsed[1][0] == "TRUE"  # 布尔转字面量

        with use_workspace(tmp_path):
            result = apply_spreadsheet_changes(file_path="book.xlsx", workbook_spec=data)
            assert result.success, result.model_text
            wb = load_workbook(str(tmp_path / "book.xlsx"))
            ws = wb["S"]
            # "=..." 字符串 → 公式
            assert ws["A4"].value == "=SUM(C1:C1)"
            assert ws["A4"].data_type == "f"
            # null → 该格无公式、留空（不写空字符串）
            assert ws["B4"].value is None
            assert ws["B4"].data_type != "f"
            # 数字/布尔 → 字符串字面量，不是公式
            assert ws["C4"].value == "5" and ws["C4"].data_type == "s"
            assert ws["D4"].value == "5" and ws["D4"].data_type == "s"
            assert ws["E4"].value == "5.5" and ws["E4"].data_type == "s"
            assert ws["A5"].value == "TRUE" and ws["A5"].data_type == "s"
            assert ws["E5"].value == "7" and ws["E5"].data_type == "s"
            # 非 "=" 字符串 → 文本字面量；空串按 null 处理留空
            assert ws["B5"].value == "标签" and ws["B5"].data_type == "s"
            assert ws["C5"].value is None
            wb.close()

    def test_null_cell_does_not_overwrite_existing_value(self, tmp_path) -> None:
        data = _spec({
            "dimensions": {"rows": 2, "cols": 2},
            "value_blocks": [{"start": "A1", "values": [["标题", 123]]}],
            "formula_blocks": [{"start": "B1", "formulas": [[None], ["=B1*2"]]}],
        })
        assert validate_workbook_spec(data)
        with use_workspace(tmp_path):
            result = apply_spreadsheet_changes(file_path="book.xlsx", workbook_spec=data)
            assert result.success, result.model_text
            wb = load_workbook(str(tmp_path / "book.xlsx"))
            ws = wb["S"]
            assert ws["B1"].value == 123 and ws["B1"].data_type == "n"
            assert ws["B2"].value == "=B1*2" and ws["B2"].data_type == "f"
            wb.close()

    def test_formula_block_rejects_container_cells_with_shape_hint(self) -> None:
        data = _spec({
            "formula_blocks": [{"start": "A1", "formulas": [[["nested"]]]}],
            "dimensions": {"rows": 1, "cols": 1},
        })
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        err = excinfo.value.errors[0]
        assert err["path"] == "sheets.0.formula_blocks.0.formulas.0.0"
        assert "公式矩阵单元格只接受" in err["message"]
        assert "null" in err["message"]


class TestValueBlockCellCoercion:
    def test_null_blank_and_date_cells_land(self, tmp_path) -> None:
        data = _spec({
            "dimensions": {"rows": 3, "cols": 5},
            "value_blocks": [{"start": "A1", "values": [
                ["月份", None, 1200, True, ""],
                ["2024-01-15", "2024-02-01T08:30:00", 5.0, "文本", None],
            ]}],
        })
        assert validate_workbook_spec(data)
        with use_workspace(tmp_path):
            result = apply_spreadsheet_changes(file_path="book.xlsx", workbook_spec=data)
            assert result.success, result.model_text
            wb = load_workbook(str(tmp_path / "book.xlsx"))
            ws = wb["S"]
            # null / 空串 → 留空，不残留空字符串
            assert ws["B1"].value is None
            assert ws["E1"].value is None
            # 数字 / 布尔保持原类型
            assert ws["C1"].value == 1200 and ws["C1"].data_type == "n"
            assert ws["D1"].value is True and ws["D1"].data_type == "b"
            # ISO 日期 / 时间戳 → 日期单元格
            assert _day(ws["A2"].value) == date(2024, 1, 15)
            assert ws["A2"].data_type == "d"
            assert ws["B2"].value == datetime(2024, 2, 1, 8, 30)
            assert ws["B2"].data_type == "d"
            # 5.0 是数字单元格，不带 ".0" 文本
            assert ws["C2"].value == 5 and ws["C2"].data_type == "n"
            assert ws["D2"].value == "文本" and ws["D2"].data_type == "s"
            wb.close()

    @pytest.mark.parametrize("bad", [
        [[["nested"]]],          # 列表单元格
        [["ok", {"a": 1}]],      # 对象单元格
        [[["x", "y"]]],          # 二维列表单元格
    ])
    def test_container_cells_report_shape_error(self, bad) -> None:
        data = _spec({
            "dimensions": {"rows": 1, "cols": 2},
            "value_blocks": [{"start": "A1", "values": bad}],
        })
        with pytest.raises(SpecValidationError) as excinfo:
            validate_workbook_spec(data)
        err = next(
            e for e in excinfo.value.errors
            if e["path"].startswith("sheets.0.value_blocks.0.values")
        )
        assert "值矩阵单元格只接受" in err["message"]
        assert "null" in err["message"]


class TestColorSpellingsAndFillAliases:
    def test_color_spellings_normalize_to_same_fill_on_disk(self, tmp_path) -> None:
        data = _spec({
            "dimensions": {"rows": 1, "cols": 4},
            "styles": {
                "plain6": {"fill": {"color": "FF0000", "fill_type": "solid"}},
                "hash6": {"fill": {"color": "#FF0000", "fill_type": "solid"}},
                "argb8": {"fill": {"color": "FFFF0000", "fill_type": "solid"}},
                "aliased": {
                    "fill": {"fgColor": "00FF00", "patternType": "solid"},
                    "font": {"color": "#0000FF", "bold": True},
                },
            },
            "style_regions": [
                {"range": "A1", "style_id": "plain6"},
                {"range": "B1", "style_id": "hash6"},
                {"range": "C1", "style_id": "argb8"},
                {"range": "D1", "style_id": "aliased"},
            ],
            "value_blocks": [{"start": "A1", "values": [["a", "b", "c", "d"]]}],
        })
        spec = validate_workbook_spec(data)
        # 归一发生在执行层（styles._resolve_color）；模型层原样接受三种写法
        assert spec.sheets[0].styles["hash6"].fill is not None
        assert spec.sheets[0].styles["hash6"].fill.color == "#FF0000"

        with use_workspace(tmp_path):
            result = apply_spreadsheet_changes(file_path="book.xlsx", workbook_spec=data)
            assert result.success, result.model_text
            wb = load_workbook(str(tmp_path / "book.xlsx"))
            ws = wb["S"]

            def rgb(cell) -> str:
                return str(cell.fill.fgColor.rgb).upper()

            # "FF0000" 与 "#FF0000" 归一后完全一致
            assert rgb(ws["A1"]) == rgb(ws["B1"])
            assert rgb(ws["A1"]).endswith("FF0000")
            # 8 位 ARGB 原样保留
            assert rgb(ws["C1"]) == "FFFF0000"
            # fill_type/fgColor/patternType 别名同样落盘
            assert rgb(ws["D1"]).endswith("00FF00")
            assert str(ws["D1"].font.color.rgb).upper().endswith("0000FF")
            assert ws["D1"].font.bold is True
            wb.close()


class TestUncertaintyRealPayloadShape:
    def test_field_note_candidates_payload_round_trips(self, tmp_path) -> None:
        data = _spec(
            {
                "dimensions": {"rows": 1, "cols": 1},
                "value_blocks": [{"start": "A1", "values": [["标题"]]}],
            },
            uncertainties=[{
                "field": "A1:F1 标题及表头填充色",
                "note": "图片深蓝为近似取色 #1F4E79，非精确像素取色",
                "candidates": ["#1F4E79", "#2F5F8F"],
            }],
        )
        spec = validate_workbook_spec(data)
        item = spec.uncertainties[0]
        assert item.location == "A1:F1 标题及表头填充色"
        assert item.reason == "图片深蓝为近似取色 #1F4E79，非精确像素取色"
        assert item.candidate_values == ["#1F4E79", "#2F5F8F"]

        with use_workspace(tmp_path):
            result = apply_spreadsheet_changes(file_path="book.xlsx", workbook_spec=data)
            assert result.success, result.model_text
            assert result.value["observation"]["document"]["uncertainties"][0]["location"].startswith("A1:F1")
