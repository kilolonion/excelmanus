"""Regression coverage for the native read/write and object-contract repairs."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from openpyxl import Workbook, load_workbook

from excelmanus.tools import workbook_tools
from excelmanus.tools.context import use_workspace
from excelmanus.tools.workbook_tools import (
    analyze_spreadsheet,
    apply_spreadsheet_changes,
    observe_spreadsheet,
    split_spreadsheet,
)
from excelmanus.tools.registry import ToolRegistry
from tests.workbook_support import region_matrix
from excelmanus.workbook_commit import content_version_of_file


def _book(path: Path, rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def test_filter_projection_selection_keeps_source_column(tmp_path: Path) -> None:
    _book(tmp_path / "book.xlsx", [["id", "name", "amount"], [1, "a", 10], [2, "b", 20]])
    with use_workspace(tmp_path):
        result = analyze_spreadsheet(
            file_path="book.xlsx",
            mode="filter",
            sheet="Sheet1",
            conditions=[],
            columns=["amount"],
        )
        assert result.success
        selection = result.value["selection"]
        assert selection["cols"] == [3]
        assert "可写回 selection" in result.model_text
        written = apply_spreadsheet_changes(
            file_path="book.xlsx",
            expected_version=result.value["content_version"],
            operations=[
                {"kind": "write", "selection": selection, "values": [[99], [88]]}
            ],
        )
        assert written.success, written.model_text
    wb = load_workbook(tmp_path / "book.xlsx")
    assert [wb.active["C2"].value, wb.active["C3"].value] == [99, 88]
    assert [wb.active["A2"].value, wb.active["A3"].value] == [1, 2]
    wb.close()


def test_range_selection_keeps_non_a_column(tmp_path: Path) -> None:
    _book(tmp_path / "book.xlsx", [["a", "b", "c"], [1, 2, 3]])
    with use_workspace(tmp_path):
        result = observe_spreadsheet(file_path="book.xlsx", mode="range", range="C2:C2")
        assert result.success
        assert result.value["regions"][0]["selection"]["cols"] == [3]
        written = apply_spreadsheet_changes(
            file_path="book.xlsx",
            expected_version=result.value["content_version"],
            operations=[{"kind": "write", "selection": result.value["regions"][0]["selection"], "values": [[9]]}],
        )
        assert written.success, written.model_text
    wb = load_workbook(tmp_path / "book.xlsx")
    assert wb.active["A2"].value == 1
    assert wb.active["C2"].value == 9
    wb.close()


def test_analyze_schema_fields_bind_and_execute(tmp_path: Path) -> None:
    _book(tmp_path / "book.xlsx", [["group", "value"], ["A", 1]])
    registry = ToolRegistry()
    registry.register_tools(workbook_tools.get_tools())
    with use_workspace(tmp_path):
        result = registry.call_tool(
            "analyze_spreadsheet",
            {"mode": "files", "directory": ".", "max_files": 5, "query": "book"},
        )
    assert result.success, result.model_text


def test_profile_quality_accept_sample_rows_as_scan_budget(tmp_path: Path) -> None:
    """profile/quality 接受 sample_rows/limit 作采样上限别名；其它 mode 仍拒绝。"""
    (tmp_path / "sales.csv").write_text(
        "month,amount\n" + "\n".join(f"2024-{i:02d},{i}" for i in range(1, 21)),
        encoding="utf-8",
    )
    with use_workspace(tmp_path):
        for mode in ("quality", "profile"):
            sampled = analyze_spreadsheet(file_path="sales.csv", mode=mode, sample_rows=5)
            assert sampled.success, sampled.model_text
            sheet = sampled.value["sheets"][0]
            assert sheet["sampled"] is True
            assert sheet["sample_size"] == 5
            limited = analyze_spreadsheet(file_path="sales.csv", mode=mode, max_rows=5)
            assert limited.success, limited.model_text
        rejected = analyze_spreadsheet(file_path="sales.csv", mode="files", sample_rows=5)
        assert not rejected.success


def test_canonical_schemas_have_declared_required_fields():
    from jsonschema import Draft202012Validator
    for tool in workbook_tools.get_tools():
        Draft202012Validator.check_schema(tool.input_schema)
        assert set(tool.input_schema.get("required",[])) <= set(tool.input_schema["properties"])
    from excelmanus.workbook.protocol import ObservationRequest
    observe=next(t for t in workbook_tools.get_tools() if t.name=="observe_spreadsheet")
    assert set(observe.input_schema["properties"])-{"file_path","expected_version"} <= set(ObservationRequest.model_fields)


def test_split_tsv_and_distinct_keys(tmp_path: Path) -> None:
    (tmp_path / "book.tsv").write_text("group\tv\nA/B\t1\nA:B\t2\n", encoding="utf-8")
    with use_workspace(tmp_path):
        result = split_spreadsheet(file_path="book.tsv", by_column="group")
    assert result.success, result.model_text
    assert result.value["groups"] == 2
    outputs = sorted((tmp_path / "outputs").glob("*.xlsx"))
    assert len(outputs) == 2
    values = sorted(load_workbook(path).active["A2"].value for path in outputs)
    assert values == ["A/B", "A:B"]


def test_profile_explicit_sheet_beyond_cap_is_not_empty(tmp_path: Path) -> None:
    wb = Workbook()
    wb.active.title = "S1"
    for index in range(2, 12):
        ws = wb.create_sheet(f"S{index}")
        ws.append(["id", "value"])
        ws.append([index, index])
    path = tmp_path / "many.xlsx"
    wb.save(path)
    wb.close()
    with use_workspace(tmp_path):
        result = analyze_spreadsheet(file_path="many.xlsx", mode="profile", sheet="S11")
    assert result.success, result.model_text
    assert result.value["sheets"]
    assert result.value["sheets"][0]["name"] == "S11"


def test_transform_rejects_formula_rewrite(tmp_path: Path) -> None:
    _book(tmp_path / "formula.xlsx", [["id", "value"], [1, 2], [1, "=B2*2"]])
    with use_workspace(tmp_path):
        result = apply_spreadsheet_changes(
            file_path="formula.xlsx",
            expected_version=content_version_of_file(tmp_path / "formula.xlsx"),
            operations=[
                {
                    "kind": "transform",
                    "sheet": "Sheet1",
                    "action": "dedupe",
                    "key_columns": ["id"],
                }
            ],
        )
    assert not result.success
    assert result.error is not None
    assert "公式" in result.error.message


def test_pivot_same_source_requires_explicit_overwrite(tmp_path: Path) -> None:
    _book(tmp_path / "pivot.xlsx", [["group", "kind", "value"], ["A", "x", 1]])
    with use_workspace(tmp_path):
        result = apply_spreadsheet_changes(
            file_path="pivot.xlsx",
            expected_version=content_version_of_file(tmp_path / "pivot.xlsx"),
            operations=[
                {
                    "kind": "pivot",
                    "sheet": "Sheet1",
                    "target_sheet": "Sheet1",
                    "index": "group",
                    "columns": "kind",
                    "values": "value",
                }
            ],
        )
    assert not result.success
    assert result.error is not None
    assert "overwrite=true" in result.error.message


def test_compare_coordinates_counts_versions_and_uncached_formula(tmp_path: Path) -> None:
    from excelmanus.tools.workbook_tools import compare_spreadsheets
    _book(tmp_path / 'a.xlsx', [['id', 'v'], [1, '=1+1'], [2, 10], [3, 20]])
    _book(tmp_path / 'b.xlsx', [['v', 'id'], [1, '=1+2'], [2, 11], [3, 21]])
    with use_workspace(tmp_path):
        result = compare_spreadsheets(file_a='a.xlsx', file_b='b.xlsx', max_diffs=1)
    assert result.success, result.model_text
    assert result.value['summary']['cells_different'] == 5
    assert result.value['coverage']['counts_complete'] is True
    assert len(result.value['sample_diffs']) == 1
    assert result.value['sample_diffs'][0]['cell'] == 'A1'
    assert result.value['content_version_a'] == content_version_of_file(tmp_path / 'a.xlsx')
    assert result.value['formula_status']['a'] == 'uncached'
    assert result.value['compared_sheets'] == {'a': ['Sheet1'], 'b': ['Sheet1']}


def test_union_range_selections_are_independent(tmp_path: Path) -> None:
    _book(tmp_path / 'book.xlsx', [['A', 'B', 'C'], [1, 2, 3], [4, 5, 6]])
    with use_workspace(tmp_path):
        read = observe_spreadsheet(file_path='book.xlsx', mode='range', range='B2:B3,C2:C3')
        assert read.success, read.model_text
        assert 'selection' not in read.value
        sel = read.value['regions'][1]['selection']
        assert sel['cols'] == [3]
        written = apply_spreadsheet_changes(file_path='book.xlsx', expected_version=read.value['content_version'], operations=[{'kind': 'write', 'selection': sel, 'values': [[30], [60]]}])
        assert written.success, written.model_text
    wb = load_workbook(tmp_path / 'book.xlsx')
    assert wb.active['B2'].value == 2
    assert wb.active['C3'].value == 60
    wb.close()


def test_split_offset_header_styles_literal_formula_and_key_identity(tmp_path: Path) -> None:
    from openpyxl.styles import Font
    p = tmp_path / 'book.xlsx'
    _book(p, [['Title'], ['group', 'v'], ['A/B', '=literal'], ['A:B', 2]])
    wb = load_workbook(p)
    wb.active['B3'].data_type = 's'
    wb.active['B3'].font = Font(bold=True)
    wb.save(p)
    wb.close()
    with use_workspace(tmp_path):
        result = split_spreadsheet(file_path='book.xlsx', by_column='group', header_row=2)
    assert result.success, result.model_text
    assert {f['key'] for f in result.value['files']} == {'A/B', 'A:B'}
    path = next(f['file_path'] for f in result.value['files'] if f['key'] == 'A/B')
    wb = load_workbook(tmp_path / path)
    assert wb.active['A1'].value == 'group'
    assert wb.active['B2'].value == '=literal'
    assert wb.active['B2'].data_type == 's'
    assert wb.active['B2'].font.bold
    wb.close()


def test_split_transaction_partial_is_truthful_and_recoverable(tmp_path: Path, monkeypatch) -> None:
    from excelmanus.workspace.file_service import WorkspaceFileService
    _book(tmp_path / 'book.xlsx', [['group', 'v'], ['A', 1], ['B', 2]])
    publish = WorkspaceFileService._publish_one
    def injected(self, target):
        if target.rel == 'outputs/B.xlsx':
            raise OSError('injected disk failure')
        return publish(self, target)
    with use_workspace(tmp_path), monkeypatch.context() as m:
        m.setattr(WorkspaceFileService, '_publish_one', injected)
        result = split_spreadsheet(file_path='book.xlsx', by_column='group')
    assert not result.success
    assert result.value['partial'] is True
    assert result.value['committed_files'] == ['outputs/A.xlsx']
    assert result.value['operation_id']
    assert (tmp_path / 'outputs/A.xlsx').is_file()
    WorkspaceFileService(tmp_path).recover()
    assert (tmp_path / 'outputs/B.xlsx').is_file()


def test_schema_alias_reaches_sdk_and_registry(tmp_path: Path) -> None:
    from excelmanus.code_mode import render_sdk_source
    namespace = {}
    exec(render_sdk_source(workbook_tools.get_tools()), namespace)
    calls = []
    namespace['_call_host'] = lambda name, args: calls.append((name, args)) or args
    namespace['apply_spreadsheet_changes'](file_path='book.xlsx', operations=[{'kind': 'format', 'range': 'A1', 'font': {'bold': True}}])
    assert calls[0][1]['file_path'] == 'book.xlsx'
    registry = ToolRegistry()
    registry.register_tools(workbook_tools.get_tools())
    registry.configure_schema_validation(mode='enforce', canary_percent=100, strict_path=False)
    _book(tmp_path / 'book.xlsx', [['group'], ['A']])
    with use_workspace(tmp_path):
        result = registry.call_tool('split_spreadsheet', {'file_path': 'book.xlsx', 'by_column': 'group'})
    assert result.success, result.model_text


@pytest.mark.asyncio
async def test_native_dispatcher_read_filter_write_result_is_model_usable(tmp_path: Path) -> None:
    from excelmanus.engine import AgentEngine
    from excelmanus.config import ExcelManusConfig
    from excelmanus.engine_core.spill import SpillStore
    _book(tmp_path / 'book.xlsx', [['id', 'amount'], [1, 2], [2, 4], [3, 6], [4, 8]])
    registry = ToolRegistry()
    registry.register_builtin_tools(str(tmp_path))
    engine = AgentEngine(ExcelManusConfig(api_key='test-key', base_url='https://test.example/v1', model='test', workspace_root=str(tmp_path)), registry)
    async def call(name, args, call_id):
        tc = SimpleNamespace(id=call_id, function=SimpleNamespace(name=name, arguments=json.dumps(args)))
        out = await engine._tool_runtime.execute(tc, None, None, 1)
        assert out.success, out.result
        payload = json.loads(out.result)
        if payload.get('result_spill'):
            payload = json.loads(SpillStore(tmp_path).get(payload['result_spill']))
        return payload
    filtered = await call('analyze_spreadsheet', {'mode': 'filter', 'file_path': 'book.xlsx', 'conditions': [], 'columns': ['amount']}, 'filter')
    assert len(filtered['values']) == 4  # actual native text, beyond the old three-row preview
    assert filtered['model_field_aliases']['data'] == 'values'
    assert filtered['selection']['cols'] == [2]
    written = await call('apply_spreadsheet_changes', {'file_path': 'book.xlsx', 'expected_version':filtered['content_version'], 'operations': [{'kind': 'write', 'selection': filtered['selection'], 'values': [[10], [20], [30], [40]]}]}, 'write')
    assert written['content_version']
    read = await call('observe_spreadsheet', {'file_path': 'book.xlsx', 'mode': 'range', 'range': 'B2:B5'}, 'read')
    assert region_matrix(read['regions'][0]) == [[10], [20], [30], [40]]


def test_large_native_payload_spills_actual_data(tmp_path: Path) -> None:
    from excelmanus.engine_core.spill import SpillStore, expose_spreadsheet_value
    from excelmanus.engine_core.tool_result import ToolResult
    value = {'status': 'success', 'data': [{'key': str(i)*20} for i in range(1000)], 'warnings': ['保留此警告']}
    store = SpillStore(tmp_path)
    output = expose_spreadsheet_value(ToolResult(success=True, value=value, model_text='only preview'), store=store)
    model = json.loads(output.model_text)
    assert model['warnings'] == value['warnings']
    assert json.loads(store.get(model['result_spill'])) == value


def test_overview_target_header_object_metadata_and_version(tmp_path: Path) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation
    _book(tmp_path / 'book.xlsx', [['Title'], ['id', None, 'amount'], [1, None, 20]])
    wb = load_workbook(tmp_path / 'book.xlsx')
    rule = DataValidation(type='list', formula1='"a,b"')
    rule.add('A3:A10')
    wb.active.add_data_validation(rule)
    wb.create_sheet('Other')
    wb.save(tmp_path / 'book.xlsx')
    wb.close()
    with use_workspace(tmp_path):
        read = observe_spreadsheet(file_path='book.xlsx', mode='range', range='A2:C3', sheet='Sheet1',
                                   facets=['data', 'presentation'])
        assert read.success, read.model_text
        assert len(read.value['regions']) == 1
        sheet = read.value['regions'][0]
        assert region_matrix(sheet)[0] == ['id', None, 'amount']
        assert sheet['data_validation']
        assert 'print_settings' in sheet
        stale = observe_spreadsheet(file_path='book.xlsx', mode='overview', expected_version='outdated')
        assert not stale.success
        assert stale.error.code == 'STALE_SNAPSHOT'


def test_clean_value_column_keeps_formulas_elsewhere_and_dedupe_moves_styles(tmp_path: Path) -> None:
    from openpyxl.styles import Font
    p = tmp_path / 'book.xlsx'
    _book(p, [['Title'], ['id', 'phone', 'derived'], ['A', '138-0013-8000', '=1+1'], ['B', '13900001111', 3]])
    with use_workspace(tmp_path):
        result = apply_spreadsheet_changes(file_path='book.xlsx', expected_version=content_version_of_file(p),
            operations=[{'kind': 'transform', 'sheet': 'Sheet1', 'header_row': 2, 'action': 'normalize_phone', 'column': 'phone'}])
    assert result.success, result.model_text
    wb = load_workbook(p)
    assert wb.active['B3'].value == '13800138000'
    assert wb.active['C3'].value == '=1+1'
    assert wb.active['A1'].value == 'Title'
    wb.close()
    p = tmp_path / 'dedupe.xlsx'
    _book(p, [['id', 'v'], ['A', 'first'], ['A', 'discard'], ['B', 'kept']])
    wb = load_workbook(p)
    wb.active['B4'].font = Font(bold=True, color='FF0000')
    wb.save(p)
    wb.close()
    with use_workspace(tmp_path):
        result = apply_spreadsheet_changes(file_path='dedupe.xlsx', expected_version=content_version_of_file(p),
                                 operations=[{'kind': 'transform', 'action': 'dedupe', 'key_columns': ['id']}])
    assert result.success, result.model_text
    wb = load_workbook(p)
    assert wb.active['B3'].value == 'kept'
    assert wb.active['B3'].font.bold
    wb.close()


def test_format_appearance_uses_address_sheet_and_preserves_font_details(tmp_path: Path) -> None:
    from excelmanus.tools.workbook_tools import apply_spreadsheet_changes
    from openpyxl.styles import Font
    p = tmp_path / 'book.xlsx'
    _book(p, [['a']])
    wb = load_workbook(p)
    ws = wb.create_sheet('Other')
    ws['A1'] = 'text'
    ws['A1'].font = Font(name='Arial', family=2, scheme='minor')
    wb.save(p)
    wb.close()
    with use_workspace(tmp_path):
        result = apply_spreadsheet_changes(file_path='book.xlsx', expected_version=content_version_of_file(p),
            operations=[{'kind': 'format', 'range': 'Other!A1', 'font': {'bold': True}, 'alignment': {'indent': 1, 'text_rotation': 30}}])
    assert result.success, result.model_text
    assert any(s['name']=='Other' for s in result.value['observation']['sheets'])
    wb = load_workbook(p)
    assert wb['Other']['A1'].font.family == 2
    assert wb['Other']['A1'].font.scheme == 'minor'
    assert wb['Other']['A1'].alignment.indent == 1
    assert wb['Other']['A1'].alignment.textRotation == 30
    wb.close()


def test_trace_absolute_cell_and_checkpoint_preserve_version(tmp_path: Path) -> None:
    from excelmanus.tools.workbook_tools import trace_spreadsheet_formulas, manage_spreadsheet_versions
    _book(tmp_path / 'book.xlsx', [['id', 'v'], [1, '=A2+1']])
    with use_workspace(tmp_path):
        result = trace_spreadsheet_formulas(file_path='book.xlsx', mode='trace', target='Sheet1!$B$2')
        assert result.success, result.model_text
        assert result.value['formula'] == '=A2+1'
        assert result.value['coverage']['kind'] == 'partial'
        checkpoint = manage_spreadsheet_versions(file_path='book.xlsx', action='checkpoint', expected_version=result.value['content_version'])
        assert checkpoint.success
        stale = manage_spreadsheet_versions(file_path='book.xlsx', action='checkpoint', expected_version='outdated')
        assert not stale.success


def test_output_branch_rejects_filter_without_selection() -> None:
    from excelmanus.tools.output_contracts import validate_output, contract_summary
    assert validate_output('analyze_spreadsheet', {'status': 'success', 'meta': {'kind': 'filter'}})
    assert 'meta.kind=filter' in contract_summary('analyze_spreadsheet')


def test_irrelevant_edit_field_and_conflicting_request_are_rejected(tmp_path: Path) -> None:
    p = tmp_path / 'book.xlsx'
    _book(p, [['id'], [1]])
    registry = ToolRegistry()
    registry.register_tools(workbook_tools.get_tools())
    with use_workspace(tmp_path):
        result = apply_spreadsheet_changes(file_path='book.xlsx', expected_version=content_version_of_file(p),
                                 operations=[{'kind': 'write', 'start_cell': 'A2', 'values': [[2]], 'font': {'bold': True}}])
        assert not result.success
        assert result.value['invalid_fields'] == ['font']
        result = registry.call_tool('observe_spreadsheet', {'file_path': 'book.xlsx', 'request': {'path': 'different.xlsx'}})
        assert not result.success
        assert result.error.code == 'TOOL_ARGUMENT_VALIDATION_ERROR'


def test_copy_literal_whitespace_and_split_case_collisions(tmp_path: Path) -> None:
    p = tmp_path / 'book.xlsx'
    _book(p, [['group', 'v'], ['A/B', '  =literal'], ['a/b', 'text']])
    with use_workspace(tmp_path):
        written = apply_spreadsheet_changes(file_path='book.xlsx', expected_version=content_version_of_file(p),
            operations=[{'kind': 'copy', 'source_sheet': 'Sheet1', 'source_range': 'B2', 'target_sheet': 'Sheet1', 'target_start': 'C2'}])
        assert written.success, written.model_text
        split = split_spreadsheet(file_path='book.xlsx', by_column='group', header_row=1)
        assert split.success, split.model_text
    wb = load_workbook(p)
    assert wb.active['C2'].value == '  =literal'
    assert wb.active['C2'].data_type == 's'
    wb.close()
    files = [f['file_path'] for f in split.value['files']]
    assert len({p.casefold() for p in files}) == 2
