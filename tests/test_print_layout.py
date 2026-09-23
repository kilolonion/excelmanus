"""Print settings survive compilation, atomic edits, and persisted inspection."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from excelmanus.replica_spec import compile_workbook_spec_to_bytes, validate_workbook_spec
from excelmanus.security import FileAccessGuard
from excelmanus.tools import intent_tools, reference_tools
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.workbook.layout import PrintLayout, apply_print_layout
from excelmanus.workbook_commit import content_version_of_file


def _bind_workspace(root: Path) -> None:
    set_guard(FileAccessGuard(str(root)))
    intent_tools.init_guard(str(root))
    reference_tools.init_guard(str(root))


def _spec(layout: dict) -> dict:
    return {
        "sheets": [{
            "name": "收款单",
            "dimensions": {"rows": 4, "cols": 3},
            "value_blocks": [{"start": "A1", "values": [["客户", "数量", "金额"], ["甲", 2, 20]]}],
            "formula_blocks": [{"start": "C4", "formulas": [["=SUM(C2:C3)"]]}],
            "print_layout": layout,
        }],
        "uncertainties": [],
    }


def _book(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "收款单"
    ws["A1"] = "原始数据"
    ws["B1"] = "=1+1"
    ws.print_area = "A1:C4"
    ws.page_setup.scale = 75
    # Excel can retain fit counts while fit-to-page is disabled.
    ws.page_setup.fitToWidth = 2
    ws.page_setup.fitToHeight = 3
    ws.column_dimensions.group("A", "C")
    ws.column_dimensions["A"].width = 19
    ws.row_dimensions[1].height = 27
    wb.create_sheet("说明")
    wb.save(path)
    wb.close()
    return path


@pytest.mark.parametrize("paper_size,expected", [("A3", "8"), ("A4", "9"), ("A5", "11"), ("Letter", "1"), ("Legal", "5")])
def test_compile_print_layout_survives_reopen(paper_size: str, expected: str) -> None:
    spec = validate_workbook_spec(_spec({
        "print_area": " $a$1:$c$4 ",
        "orientation": "landscape",
        "paper_size": paper_size,
        "fit_to_width": 1,
        "fit_to_height": 1,
    }))
    data, _stats = compile_workbook_spec_to_bytes(spec)
    wb = load_workbook(io.BytesIO(data))
    try:
        ws = wb["收款单"]
        assert ws["A2"].value == "甲"
        assert ws["C4"].value == "=SUM(C2:C3)"
        assert "$A$1:$C$4" in str(ws.print_area)
        assert ws.page_setup.orientation == "landscape"
        assert str(ws.page_setup.paperSize) == expected
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True
        assert ws.page_setup.fitToWidth == ws.page_setup.fitToHeight == 1
        assert ws.page_setup.scale is None
    finally:
        wb.close()


@pytest.mark.parametrize("layout", [
    {"print_area": "Other!A1:B2"},
    {"print_area": "A1:B2,D1:E2"},
    {"print_area": "A:A"},
    {"print_area": "1:3"},
    {"print_area": "A0:B2"},
    {"print_area": "B2:A1"},
    {"print_area": "XFE1"},
    {"print_area": "A1048577"},
    {"orientation": "sideways"},
    {"paper_size": "A0"},
    {"fit_to_width": -1},
    {"fit_to_height": 32768},
    {"fit_to_width": True},
    {"fit_to_height": 1.5},
    {"scale": 9},
    {"scale": 401},
    {"scale": 100, "fit_to_page": True},
    {"scale": 100, "fit_to_width": 1},
    {"scale": 100, "fit_to_height": 0},
    {"fit_to_wdith": 1},
])
def test_invalid_print_layout_rejected_before_use(layout: dict) -> None:
    with pytest.raises(ValidationError):
        PrintLayout.model_validate(layout)


@pytest.mark.parametrize("layout,expected_width,expected_height", [
    ({"fit_to_page": True}, 1, 0),
    ({"fit_to_width": 0}, 0, 0),
    ({"fit_to_height": 2}, 1, 2),
])
def test_fit_defaults_are_explicit_in_persisted_workbook(layout: dict, expected_width: int, expected_height: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.sheet_properties.pageSetUpPr = None
    ws.page_setup.scale = 80
    apply_print_layout(ws, PrintLayout.model_validate(layout))
    buffer = io.BytesIO()
    wb.save(buffer)
    loaded = load_workbook(io.BytesIO(buffer.getvalue()))
    try:
        persisted = loaded.active
        assert persisted.sheet_properties.pageSetUpPr.fitToPage is True
        assert persisted.page_setup.scale is None
        assert persisted.page_setup.fitToWidth == expected_width
        assert persisted.page_setup.fitToHeight == expected_height
    finally:
        loaded.close()
        wb.close()


def test_format_print_layout_switches_scaling_modes_without_changing_content(tmp_path: Path) -> None:
    _bind_workspace(tmp_path)
    path = _book(tmp_path / "receipt.xlsx")
    fitted = intent_tools.format_spreadsheet(
        file_path=str(path), expected_version=content_version_of_file(path),
        operations=[{"kind": "print_layout", "sheet": "收款单", "print_layout": {
            "orientation": "landscape", "paper_size": "A4", "fit_to_width": 1, "fit_to_height": 0,
        }}],
    )
    assert fitted.success, fitted.model_text
    wb = load_workbook(path)
    try:
        ws = wb["收款单"]
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True
        assert ws.page_setup.scale is None
        assert ws.page_setup.fitToWidth == 1
        assert ws.page_setup.fitToHeight == 0
        assert ws["A1"].value == "原始数据"
        assert ws["B1"].value == "=1+1"
        assert wb["说明"].page_setup.orientation is None
    finally:
        wb.close()
    scaled = intent_tools.format_spreadsheet(
        file_path=str(path), expected_version=fitted.ui_meta.content_version,
        operations=[{"kind": "print_layout", "sheet": "收款单", "print_layout": {"scale": 125, "print_area": ""}}],
    )
    assert scaled.success, scaled.model_text
    wb = load_workbook(path)
    try:
        ws = wb["收款单"]
        assert ws.sheet_properties.pageSetUpPr.fitToPage is False
        assert ws.page_setup.scale == 125
        assert not ws.print_area
        assert ws.page_setup.orientation == "landscape"
        assert ws["B1"].value == "=1+1"
    finally:
        wb.close()


@pytest.mark.parametrize("layout", [{}, {"fit_to_width": -1}, {"scale": 100, "fit_to_page": True}, {"print_area": "Other!A1:B2"}])
def test_invalid_print_edit_does_not_save_earlier_operations(tmp_path: Path, layout: dict) -> None:
    _bind_workspace(tmp_path)
    path = _book(tmp_path / "receipt.xlsx")
    before = path.read_bytes()
    result = intent_tools.format_spreadsheet(
        file_path=str(path), expected_version=content_version_of_file(path),
        operations=[
            {"kind": "format", "sheet": "收款单", "range": "A1", "font": {"bold": True}},
            {"kind": "print_layout", "sheet": "收款单", "print_layout": layout},
        ],
    )
    assert not result.success
    assert result.error.code == "INVALID_ARGS"
    assert path.read_bytes() == before
    assert result.value["operation_index"] == 1


def test_invalid_creation_print_layout_leaves_no_output(tmp_path: Path) -> None:
    _bind_workspace(tmp_path)
    path = tmp_path / "receipt.xlsx"
    result = intent_tools.edit_spreadsheet(
        file_path=str(path), workbook_spec=_spec({"fit_to_width": 1, "scale": 50}),
    )
    assert not result.success
    assert result.error.code == "SPEC_VALIDATION_FAILED"
    assert not path.exists()


def test_inspection_reports_persisted_dimensions_and_print_mode(tmp_path: Path) -> None:
    _bind_workspace(tmp_path)
    path = _book(tmp_path / "receipt.xlsx")
    result = intent_tools.inspect_spreadsheet(
        file_path=str(path), mode="overview", include=["print_settings", "row_heights", "column_widths"],
    )
    assert result.success, result.model_text
    sheet = next(sheet for sheet in result.value["sheets"] if sheet["name"] == "收款单")
    assert sheet["row_heights"] == {"1": 27}
    assert sheet["column_widths"] == {"A": 19, "B": 19, "C": 19}
    assert sheet["print_settings"]["scaling_mode"] == "scale"
    assert sheet["print_settings"]["fit_to_page"] is False
    assert sheet["print_settings"]["scale"] == 75


def test_partial_print_edit_preserves_other_layout_settings(tmp_path: Path) -> None:
    _bind_workspace(tmp_path)
    path = _book(tmp_path / "receipt.xlsx")
    wb = load_workbook(path)
    ws = wb["收款单"]
    ws.print_title_rows = "1:2"
    ws.print_title_cols = "A:B"
    ws.page_margins.left = 0.25
    ws.oddHeader.center.text = "收款明细"
    ws.page_setup.fitToWidth = 2
    ws.page_setup.fitToHeight = 3
    wb.save(path)
    wb.close()
    result = intent_tools.format_spreadsheet(
        file_path=str(path), expected_version=content_version_of_file(path),
        operations=[{"kind": "print_layout", "sheet": "收款单", "print_layout": {"orientation": "landscape"}}],
    )
    assert result.success, result.model_text
    wb = load_workbook(path)
    try:
        ws = wb["收款单"]
        assert ws.page_setup.orientation == "landscape"
        assert ws.page_setup.scale == 75
        assert ws.page_setup.fitToWidth == 2
        assert ws.page_setup.fitToHeight == 3
        assert ws.print_title_rows == "$1:$2"
        assert ws.print_title_cols == "$A:$B"
        assert ws.page_margins.left == 0.25
        assert ws.oddHeader.center.text == "收款明细"
        assert "$A$1:$C$4" in str(ws.print_area)
    finally:
        wb.close()


def test_print_edit_requires_sheet_only_when_ambiguous(tmp_path: Path) -> None:
    _bind_workspace(tmp_path)
    path = _book(tmp_path / "receipt.xlsx")
    before = path.read_bytes()
    operation = {"kind": "print_layout", "print_layout": {"orientation": "landscape"}}
    rejected = intent_tools.format_spreadsheet(
        file_path=str(path), expected_version=content_version_of_file(path), operations=[operation],
    )
    assert not rejected.success
    assert rejected.error.code == "SHEET_REQUIRED"
    assert path.read_bytes() == before
    wb = load_workbook(path)
    del wb["说明"]
    wb.save(path)
    wb.close()
    accepted = intent_tools.format_spreadsheet(
        file_path=str(path), expected_version=content_version_of_file(path), operations=[operation],
    )
    assert accepted.success, accepted.model_text
    wb = load_workbook(path)
    try:
        assert wb.active.page_setup.orientation == "landscape"
    finally:
        wb.close()


@pytest.mark.parametrize("scaling", [
    {"fit_to_width": 1, "fit_to_height": 0},
    {"scale": 125},
])
def test_print_settings_can_be_used_directly_to_format_another_sheet(tmp_path: Path, scaling: dict) -> None:
    _bind_workspace(tmp_path)
    path = _book(tmp_path / "receipt.xlsx")
    result = intent_tools.format_spreadsheet(
        file_path=str(path), expected_version=content_version_of_file(path),
        operations=[{"kind": "print_layout", "sheet": "收款单", "print_layout": {
            "print_area": "A1:C4", "paper_size": "A4", "orientation": "landscape", **scaling,
        }}],
    )
    assert result.success, result.model_text
    settings = result.value["appearance"]["sheets"][0]["print_settings"]
    layout = settings["print_layout"]
    assert PrintLayout.model_validate(layout).print_area == "A1:C4"
    assert layout["paper_size"] == "A4"
    assert layout["orientation"] == "landscape"
    if "scale" in scaling:
        assert layout["scale"] == 125
        assert "fit_to_width" not in layout
        assert "fit_to_height" not in layout
    else:
        assert layout["fit_to_width"] == 1
        assert layout["fit_to_height"] == 0
        assert "scale" not in layout
    readback = intent_tools.inspect_spreadsheet(
        file_path=str(path), mode="overview", include=["print_settings"],
    )
    assert readback.success, readback.model_text
    source = next(sheet for sheet in readback.value["sheets"] if sheet["name"] == "收款单")
    assert source["print_settings"]["print_layout"] == layout
    copied = intent_tools.format_spreadsheet(
        file_path=str(path), expected_version=readback.ui_meta.content_version,
        operations=[{"kind": "print_layout", "sheet": "说明", "print_layout": layout}],
    )
    assert copied.success, copied.model_text
    wb = load_workbook(path)
    try:
        ws = wb["说明"]
        assert "$A$1:$C$4" in str(ws.print_area)
        assert ws.page_setup.orientation == "landscape"
        assert str(ws.page_setup.paperSize) == "9"
        if "scale" in scaling:
            assert ws.sheet_properties.pageSetUpPr.fitToPage is False
            assert ws.page_setup.scale == 125
        else:
            assert ws.sheet_properties.pageSetUpPr.fitToPage is True
            assert ws.page_setup.fitToWidth == 1
            assert ws.page_setup.fitToHeight == 0
    finally:
        wb.close()


def test_print_inspection_preserves_unsupported_multiple_areas_and_paper(tmp_path: Path) -> None:
    _bind_workspace(tmp_path)
    path = _book(tmp_path / "receipt.xlsx")
    wb = load_workbook(path)
    ws = wb["收款单"]
    ws.print_area = ["A1:B4", "D1:F4"]
    ws.page_setup.paperSize = "12"
    wb.save(path)
    wb.close()
    result = intent_tools.inspect_spreadsheet(
        file_path=str(path), mode="overview", include=["print_settings"],
    )
    assert result.success, result.model_text
    source = next(sheet for sheet in result.value["sheets"] if sheet["name"] == "收款单")
    settings = source["print_settings"]
    assert "$A$1:$B$4" in str(settings["print_area"])
    assert "$D$1:$F$4" in str(settings["print_area"])
    assert str(settings["paper_size"]) == "12"
    assert "print_area" not in settings["print_layout"]
    assert "paper_size" not in settings["print_layout"]
    PrintLayout.model_validate(settings["print_layout"])


@pytest.mark.parametrize("sheet_name", ["明细,账单", "客户's 收据", "逗号,引号' 混合"])
def test_readback_print_area_handles_quoted_sheet_names(sheet_name: str) -> None:
    from excelmanus.workbook.data import _collect_print_settings

    wb = Workbook()
    try:
        ws = wb.active
        ws.title = sheet_name
        ws.print_area = "A1:C4"
        settings = _collect_print_settings(ws)
        assert settings["print_layout"]["print_area"] == "A1:C4"
        PrintLayout.model_validate(settings["print_layout"])
    finally:
        wb.close()


def test_readback_without_fit_flag_does_not_activate_retained_fit_counts() -> None:
    from excelmanus.workbook.data import _collect_print_settings

    wb = Workbook()
    try:
        source = wb.active
        source.sheet_properties.pageSetUpPr = None
        source.page_setup.fitToWidth = 1
        source.page_setup.fitToHeight = 1
        source.page_setup.scale = 80
        layout = _collect_print_settings(source)["print_layout"]
        assert layout["fit_to_page"] is False
        assert layout["scale"] == 80
        assert "fit_to_width" not in layout
        assert "fit_to_height" not in layout
        destination = wb.create_sheet("copy")
        apply_print_layout(destination, PrintLayout.model_validate(layout))
        assert destination.sheet_properties.pageSetUpPr.fitToPage is False
        assert destination.page_setup.scale == 80
    finally:
        wb.close()
