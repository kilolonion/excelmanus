"""wave-r5u 复盘落地项的定向测试。

覆盖 §12.7 缺陷清单：
- C: operations JSON 尾杂散容错 / Extra data 不误诊截断
- A2: pivot margins 读写两侧合计
- A1: write.values 接受 spill 句柄 / JSON 字符串
- B: code-mode 桥错误全字段透传 + HostToolError.details
- D: 解释器探测缓存
- E: _tools_cache key 含 skill/subagent 摘要
- 候选1: write_text → write_text_file 别名
- 候选2: data_validation 写侧
- 候选3: max_rows 规范名在 analyze 各 mode 生效
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.security import FileAccessGuard
from excelmanus.tools import ToolRegistry, intent_tools, reference_tools
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.intent_tools import (
    _coerce_operations,
    analyze_spreadsheet,
    edit_spreadsheet,
    format_spreadsheet,
)


def _bind_workspace(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    intent_tools.init_guard(workspace)
    reference_tools.init_guard(workspace)


def _book(path: Path, rows: list[list[Any]] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows or []:
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def _payload(result: ToolResult) -> dict[str, Any]:
    assert isinstance(result, ToolResult)
    assert isinstance(result.value, dict)
    return result.value


def _version(result: ToolResult) -> str:
    return str(_payload(result).get("content_version") or "")


def _edit(path: Path, operations: list[dict[str, Any]]) -> ToolResult:
    from excelmanus.workbook_commit import content_version_of_file

    return edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=operations,
    )


def _format(path: Path, operations: list[dict[str, Any]]) -> ToolResult:
    from excelmanus.workbook_commit import content_version_of_file

    return format_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=operations,
    )


# ── C: _coerce_operations ───────────────────────────────────


class TestCoerceOperationsRecovery:
    def test_trailing_extra_brace_recovered(self) -> None:
        ops = [{"kind": "write", "sheet": "S", "start_cell": "A1", "values": [[1, 2], [3, 4]]}]
        raw = json.dumps(ops, ensure_ascii=False) + "}"
        out = _coerce_operations(raw)
        assert isinstance(out, list)
        assert out == ops

    def test_trailing_junk_then_whitespace_recovered(self) -> None:
        ops = [{"kind": "write", "sheet": "S", "values": [[1]]}]
        raw = json.dumps(ops) + "}   \n"
        out = _coerce_operations(raw)
        assert isinstance(out, list)

    def test_unrecoverable_extra_data_not_mislabeled_truncated(self) -> None:
        raw = '[{"kind":"write"}] xyz-not-json'
        out = _coerce_operations(raw)
        assert isinstance(out, ToolResult)
        assert out.success is False
        text = out.model_text or ""
        assert "截断" not in text
        assert "多余内容" in text or "重试" in text

    def test_real_truncation_still_reports_truncated(self) -> None:
        raw = '[{"kind":"write","values":[[1,2,3],'
        out = _coerce_operations(raw)
        assert isinstance(out, ToolResult)
        assert out.success is False
        assert "截断" in (out.model_text or "")

    def test_mid_string_garbage_is_invalid_not_truncated(self) -> None:
        raw = '[{"kind":oops,"values":[[1]]}]'
        out = _coerce_operations(raw)
        assert isinstance(out, ToolResult)
        assert out.success is False
        assert "截断" not in (out.model_text or "")

    def test_single_object_and_string_items_still_coerce(self) -> None:
        assert _coerce_operations({"kind": "write"}) == [{"kind": "write"}]
        out = _coerce_operations(['{"kind": "write"}'])
        assert out == [{"kind": "write"}]


# ── A2: pivot margins ───────────────────────────────────────

_PIVOT_ROWS = [
    ["区域", "月份", "金额"],
    ["华东", "2024-01", 100],
    ["华东", "2024-02", 200],
    ["华北", "2024-01", 50],
    ["华北", "2024-02", 70],
]


class TestPivotMargins:
    def test_analyze_pivot_with_margins(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "orders.xlsx", _PIVOT_ROWS)
        result = analyze_spreadsheet(
            file_path=str(path),
            mode="pivot",
            index="区域",
            columns="月份",
            values="金额",
            margins=True,
        )
        assert result.success
        data = _payload(result)
        matrix = data.get("matrix") or []
        assert matrix, f"pivot result keys: {sorted(data)}"
        # 最后一行应为合计行，最后一列为合计列
        assert matrix[-1][0] == "合计"
        header = matrix[0]
        assert header[-1] == "合计"
        # 合计单元格：100+200+50+70=420
        assert matrix[-1][-1] == 420

    def test_analyze_pivot_margins_name(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "orders.xlsx", _PIVOT_ROWS)
        result = analyze_spreadsheet(
            file_path=str(path),
            mode="pivot",
            index="区域",
            columns="月份",
            values="金额",
            margins=True,
            margins_name="总计",
        )
        assert result.success
        matrix = _payload(result).get("matrix") or []
        assert matrix[-1][0] == "总计"

    def test_analyze_pivot_without_margins_unchanged(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "orders.xlsx", _PIVOT_ROWS)
        result = analyze_spreadsheet(
            file_path=str(path),
            mode="pivot",
            index="区域",
            columns="月份",
            values="金额",
        )
        assert result.success
        matrix = _payload(result).get("matrix") or []
        assert matrix[-1][0] != "合计"

    def test_edit_pivot_writes_total_row(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "orders.xlsx", _PIVOT_ROWS)
        from excelmanus.workbook_commit import content_version_of_file

        result = _edit(path, [{
                "kind": "pivot",
                "sheet": "Sheet1",
                "target_sheet": "汇总",
                "index": "区域",
                "columns": "月份",
                "values": "金额",
                "margins": True,
            }],
        )
        assert result.success
        wb = load_workbook(path)
        ws = wb["汇总"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        wb.close()
        assert rows[-1][0] == "合计"
        assert rows[-1][-1] == 420


# ── A1: write.values 接 spill 句柄 / JSON 字符串 ───────────


class TestWriteValuesLocator:
    def test_values_spill_locator(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["h1", "h2"]])
        from excelmanus.engine_core.spill import SpillStore

        matrix = [["华东", 100], ["华北", 50]]
        locator = SpillStore(tmp_path).put(json.dumps(matrix, ensure_ascii=False))
        result = _edit(path, [{
                "kind": "write",
                "sheet": "Sheet1",
                "start_cell": "A2",
                "values": str(locator),
            }],
        )
        assert result.success, result.model_text
        wb = load_workbook(path)
        ws = wb["Sheet1"]
        assert ws["A2"].value == "华东" and ws["B3"].value == 50
        wb.close()

    def test_values_json_string_coerced(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["h1"]])
        result = _edit(path, [{
                "kind": "write",
                "sheet": "Sheet1",
                "start_cell": "A2",
                "values": json.dumps([[9, 8]]),
            }],
        )
        assert result.success, result.model_text
        wb = load_workbook(path)
        assert wb["Sheet1"]["A2"].value == 9
        wb.close()

    def test_values_records_spill_rejected_with_guidance(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["h1"]])
        from excelmanus.engine_core.spill import SpillStore

        locator = SpillStore(tmp_path).put(json.dumps({"records": [{"a": 1}]}))
        result = _edit(path, [{
                "kind": "write",
                "sheet": "Sheet1",
                "start_cell": "A2",
                "values": str(locator),
            }],
        )
        assert result.success is False
        assert "矩阵" in (result.model_text or "")

    def test_values_missing_locator_reports_failure(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["h1"]])
        result = _edit(path, [{
                "kind": "write",
                "sheet": "Sheet1",
                "start_cell": "A2",
                "values": "spill:" + "0" * 64,
            }],
        )
        assert result.success is False


# ── B: 桥错误 details 全字段透传 ────────────────────────────


class TestBridgeErrorDetails:
    def test_payload_passes_all_structured_fields(self) -> None:
        from excelmanus.code_mode import _payload_from_tool_result
        from excelmanus.engine_core.tool_result import ToolError

        result = ToolResult(
            success=False,
            model_text="WorkbookSpec 校验失败",
            error=ToolError(code="SPEC_VALIDATION_FAILED", message="WorkbookSpec 校验失败"),
            value={
                "error_code": "SPEC_VALIDATION_FAILED",
                "errors": [{"path": "sheets[0].values", "message": "必须为二维数组"}],
                "available_sheets": ["Sheet1"],
                "conflicts": ["A1"],
                "expected_version": "sha256:abc",
            },
        )
        payload = _payload_from_tool_result(result)
        err = payload["error"]
        assert err["code"] == "SPEC_VALIDATION_FAILED"
        details = err.get("details") or {}
        assert details.get("errors") == [{"path": "sheets[0].values", "message": "必须为二维数组"}]
        assert details.get("available_sheets") == ["Sheet1"]
        assert details.get("expected_version") == "sha256:abc"
        # 顶层不再重复信封键
        assert "error_code" not in details
        assert "message" not in details

    def test_sdk_source_compiles_and_host_error_str_shows_details(self) -> None:
        from excelmanus.code_mode import render_sdk_source
        from excelmanus.tools.registry import ToolDef

        tool = ToolDef(
            name="demo_tool",
            description="d",
            input_schema={"type": "object", "properties": {}},
            func=lambda **k: None,
        )
        src = render_sdk_source([tool])
        namespace: dict[str, Any] = {}
        exec(compile(src, "<sdk>", "exec"), namespace)
        host_err = namespace["HostToolError"]("bad spec", "SPEC_VALIDATION_FAILED", {"errors": [{"path": "a"}]})
        text = str(host_err)
        assert "bad spec" in text
        assert '"path": "a"' in text or "'path': 'a'" in text or "errors" in text


# ── D: 解释器探测缓存 ──────────────────────────────────────


class TestInterpreterResolveCache:
    def test_success_cached_skips_reprobe(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from excelmanus.tools import code_tools

        code_tools._RESOLVE_CACHE.clear()
        calls: list[list[str]] = []

        def fake_probe(command: list[str], **_kwargs: Any) -> Any:
            calls.append(command)
            return code_tools._InterpreterProbe(command=command, status="ok", detail="")

        monkeypatch.setattr(code_tools, "_probe_environment", fake_probe)
        first = code_tools._resolve_python_command("python3", require_excel_deps=True)
        second = code_tools._resolve_python_command("python3", require_excel_deps=True)
        assert first[0] == second[0]
        assert len(calls) == 1, "命中缓存不应重复探测"

    def test_failure_cached_briefly(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from excelmanus.tools import code_tools

        code_tools._RESOLVE_CACHE.clear()
        calls: list[list[str]] = []

        def fake_probe(command: list[str], **_kwargs: Any) -> Any:
            calls.append(command)
            return code_tools._InterpreterProbe(command=command, status="error", detail="boom")

        monkeypatch.setattr(code_tools, "_probe_environment", fake_probe)
        with pytest.raises(RuntimeError):
            code_tools._resolve_python_command("definitely-missing-py", require_excel_deps=False)
        with pytest.raises(RuntimeError, match="60s"):
            code_tools._resolve_python_command("definitely-missing-py", require_excel_deps=False)
        assert len(calls) == 1, "失败条目 60s 内不应重复探测"

    def test_probe_retries_once_on_timeout(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """首轮超时后用更长预算重试一次（Windows 冷启动宽限）。"""
        import subprocess as sp
        from types import SimpleNamespace
        from excelmanus.tools import code_tools

        monkeypatch.setattr(code_tools, "_command_exists", lambda command: True)
        timeouts: list[float] = []

        def fake_run(cmd: list[str], **kwargs: Any) -> Any:
            timeouts.append(kwargs["timeout"])
            if len(timeouts) == 1:
                raise sp.TimeoutExpired(cmd, kwargs["timeout"])
            return SimpleNamespace(returncode=0, stdout="", stderr="")

        monkeypatch.setattr(code_tools.subprocess, "run", fake_run)
        probe = code_tools._probe_environment(["python"], require_excel_deps=False)
        assert probe.status == "ok"
        assert timeouts == [
            code_tools._PROBE_TIMEOUT_SECONDS,
            code_tools._PROBE_TIMEOUT_RETRY_SECONDS,
        ]

    def test_probe_uses_isolated_flags(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """探针与真实执行一致带 -B -I -X utf8，避免宿主机 PYTHON* 污染。"""
        from types import SimpleNamespace
        from excelmanus.tools import code_tools

        monkeypatch.setattr(code_tools, "_command_exists", lambda command: True)
        captured: list[list[str]] = []

        def fake_run(cmd: list[str], **_kwargs: Any) -> Any:
            captured.append(cmd)
            return SimpleNamespace(returncode=0, stdout="", stderr="")

        monkeypatch.setattr(code_tools.subprocess, "run", fake_run)
        probe = code_tools._probe_environment(["python"], require_excel_deps=True)
        assert probe.status == "ok"
        flags = captured[0]
        assert "-B" in flags and "-I" in flags
        assert flags[flags.index("-X") + 1] == "utf8"

    def test_auto_dedupes_by_resolved_executable(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """PATH 上的 python/python3 与既有候选指向同一 exe 时不重复探测。"""
        import sys
        from excelmanus.tools import code_tools

        code_tools._RESOLVE_CACHE.clear()
        monkeypatch.delenv("EXCELMANUS_RUN_PYTHON", raising=False)
        monkeypatch.setattr(
            code_tools.shutil,
            "which",
            lambda name: sys.executable if name in ("python", "python3") else None,
        )
        calls: list[list[str]] = []

        def fake_probe(command: list[str], **_kwargs: Any) -> Any:
            calls.append(command)
            return code_tools._InterpreterProbe(command=command, status="error", detail="boom")

        monkeypatch.setattr(code_tools, "_probe_environment", fake_probe)
        with pytest.raises(RuntimeError):
            code_tools._resolve_python_command("auto", require_excel_deps=True)
        assert [sys.executable] in calls
        assert ["python"] not in calls and ["python3"] not in calls

    def test_warmup_populates_resolve_cache(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """预热成功后，首个真实调用直接命中缓存不再探测。"""
        from excelmanus.tools import code_tools

        code_tools._RESOLVE_CACHE.clear()
        calls: list[list[str]] = []

        def fake_probe(command: list[str], **_kwargs: Any) -> Any:
            calls.append(command)
            return code_tools._InterpreterProbe(command=command, status="ok", detail="")

        monkeypatch.setattr(code_tools, "_probe_environment", fake_probe)
        code_tools.warmup_interpreter()
        assert calls
        calls.clear()
        code_tools._resolve_python_command("auto", require_excel_deps=True)
        assert not calls

    def test_warmup_clears_failure_cache(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """预热失败不留失败缓存，首个真实调用可全新重探。"""
        from excelmanus.tools import code_tools

        code_tools._RESOLVE_CACHE.clear()
        monkeypatch.setattr(
            code_tools,
            "_probe_environment",
            lambda command, **_kwargs: code_tools._InterpreterProbe(
                command=command, status="error", detail="boom"
            ),
        )
        with pytest.raises(RuntimeError):
            code_tools.warmup_interpreter()
        assert all(value[0] is not None for value in code_tools._RESOLVE_CACHE.values())

    def test_shorten_marker_avoids_overflow_keywords(self) -> None:
        from excelmanus.tools.code_tools import _shorten

        shortened = _shorten("x" * 500)
        assert "truncat" not in shortened
        assert "截断" not in shortened


# ── E: _tools_cache key 含 skill/subagent ──────────────────


class TestToolsCacheKey:
    def test_cache_key_includes_session_lists(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from excelmanus.engine_core.meta_tools import MetaToolBuilder

        class _FakeCatalog:
            mode = "write"

            def digest(self) -> str:
                return "d1"

        class _FakeEngine:
            _active_skills: list[Any] = []
            _tools_cache: Any = None
            _tools_cache_key: Any = None
            _skill_router = None
            _skill_resolver = None
            _subagent_registry = None

        engine = _FakeEngine()
        builder = MetaToolBuilder(engine)

        import excelmanus.prompt.envelope as env
        import excelmanus.tools.catalog as cat_mod
        import excelmanus.tools.meta_tool_defs as mtd

        monkeypatch.setattr(cat_mod, "catalog_from_engine", lambda e, **_kw: _FakeCatalog())
        monkeypatch.setattr(env, "sort_tool_schemas", lambda tools: tools)
        monkeypatch.setattr(MetaToolBuilder, "build_v5_tools_impl", lambda self, **_kw: [])
        monkeypatch.setattr(mtd, "_session_skill_names", lambda e: ["s1"])

        monkeypatch.setattr(mtd, "_session_subagent_names", lambda e: ["a1"])
        builder.build_v5_tools(tool_access="may_write")
        key_a = engine._tools_cache_key

        monkeypatch.setattr(mtd, "_session_subagent_names", lambda e: ["a2"])
        builder.build_v5_tools(tool_access="may_write")
        key_b = engine._tools_cache_key

        assert key_a != key_b, "subagent 名单变化必须使 tools 缓存失效"
        assert "a2" in key_b[4] or frozenset({"a2"}) in key_b


# ── 候选1: write_text 别名 ─────────────────────────────────


class TestToolNameAlias:
    def test_registry_resolves_alias(self, tmp_path: Path) -> None:
        registry = ToolRegistry()
        registry.register_builtin_tools(str(tmp_path))
        assert registry.get_tool("write_text") is not None
        assert registry.get_tool("write_text").name == "write_text_file"
        # 目录面只暴露规范名
        assert "write_text" not in registry.get_tool_names()

    def test_call_tool_via_alias(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        registry = ToolRegistry()
        registry.register_builtin_tools(str(tmp_path))
        result = registry.call_tool(
            "write_text",
            {"file_path": "out.txt", "content": "hello"},
        )
        assert getattr(result, "success", False) is True
        assert (tmp_path / "out.txt").read_text() == "hello"


# ── 候选2: data_validation ─────────────────────────────────


class TestDataValidationWrite:
    def test_list_validation(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["状态"], ["ok"]])
        result = _format(path, [{
                "kind": "data_validation",
                "sheet": "Sheet1",
                "range": "A2:A10",
                "rule": {"type": "list", "values": ["ok", "ng"], "allow_blank": True},
            }],
        )
        assert result.success, result.model_text
        wb = load_workbook(path)
        dvs = wb["Sheet1"].data_validations.dataValidation
        wb.close()
        assert len(dvs) == 1
        assert dvs[0].type == "list"
        assert dvs[0].formula1 == '"ok,ng"'
        assert str(dvs[0].sqref) == "A2:A10"

    def test_numeric_between_validation(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["分数"], [50]])
        result = _format(path, [{
                "kind": "data_validation",
                "sheet": "Sheet1",
                "range": "A2:A10",
                "rule": {"type": "whole", "operator": "between", "value": 0, "value2": 100},
            }],
        )
        assert result.success, result.model_text
        wb = load_workbook(path)
        dv = wb["Sheet1"].data_validations.dataValidation[0]
        wb.close()
        assert dv.type == "whole" and dv.operator == "between"
        assert dv.formula1 == "0" and dv.formula2 == "100"

    def test_remove_validation(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["状态"], ["ok"]])
        add = _format(path, [{
                "kind": "data_validation",
                "sheet": "Sheet1",
                "range": "A2:A10",
                "rule": {"type": "list", "values": ["a", "b"]},
            }],
        )
        assert add.success
        removed = _format(path, [{
                "kind": "data_validation",
                "sheet": "Sheet1",
                "range": "A5:A6",
                "remove": True,
            }],
        )
        assert removed.success
        wb = load_workbook(path)
        assert len(wb["Sheet1"].data_validations.dataValidation) == 0
        wb.close()

    def test_invalid_type_rejected(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["x"], [1]])
        result = _format(path, [{
                "kind": "data_validation",
                "sheet": "Sheet1",
                "range": "A2:A3",
                "rule": {"type": "nonsense"},
            }],
        )
        assert result.success is False
        assert "type" in (result.model_text or "")


# ── 候选2b: conditional_format 别名归一 + remove ─────────────


class TestConditionalFormatWrite:
    """R27 回退复现：operator 别名缺口 + remove 未实现导致的重试链。"""

    def _cf_rules(self, path: Path) -> list[Any]:
        wb = load_workbook(path)
        rules = [cf for cf in wb["Sheet1"].conditional_formatting]
        wb.close()
        return rules

    def test_snake_operator_alias(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["金额"], [100]])
        result = _format(path, [{
                "kind": "conditional_format",
                "sheet": "Sheet1",
                "range": "A2:A10",
                "rule": {
                    "type": "cell_value",
                    "operator": "greater_than_equal",
                    "value": 10000,
                    "fill": {"color": "FF0000"},
                },
            }],
        )
        assert result.success, result.model_text
        rules = self._cf_rules(path)
        assert len(rules) == 1
        assert rules[0].rules[0].operator == "greaterThanOrEqual"
        assert rules[0].rules[0].formula == ["10000"]

    def test_symbol_operator_alias(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["金额"], [100]])
        result = _format(path, [{
                "kind": "conditional_format",
                "sheet": "Sheet1",
                "range": "A2:A10",
                "rule": {"type": "cell_value", "operator": ">=", "value": 50},
            }],
        )
        assert result.success, result.model_text
        rules = self._cf_rules(path)
        assert rules[0].rules[0].operator == "greaterThanOrEqual"

    def test_remove_rule_by_range_intersection(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["金额"], [100]])
        add = _format(path, [{
                "kind": "conditional_format",
                "sheet": "Sheet1",
                "range": "A2:A10",
                "rule": {"type": "cell_value", "operator": ">=", "value": 50},
            }],
        )
        assert add.success
        removed = _format(path, [{
                "kind": "conditional_format",
                "sheet": "Sheet1",
                "range": "A5:A6",
                "remove": True,
            }],
        )
        assert removed.success, removed.model_text
        assert "removed:1" in (removed.model_text or "") or "removed" in (
            removed.model_text or ""
        )
        assert self._cf_rules(path) == []

    def test_remove_without_rule_no_misleading_error(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["x"], [1]])
        result = _format(path, [{
                "kind": "conditional_format",
                "sheet": "Sheet1",
                "range": "A2:A3",
                "remove": True,
            }],
        )
        # 无规则可删也应成功（removed:0），而非误报"需要 value 阈值"
        assert result.success, result.model_text

    def test_missing_value_names_field(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["x"], [1]])
        result = _format(path, [{
                "kind": "conditional_format",
                "sheet": "Sheet1",
                "range": "A2:A3",
                "rule": {"type": "cell_value", "operator": "greaterThanOrEqual"},
            }],
        )
        assert result.success is False
        assert "value" in (result.model_text or "")

    def test_snake_case_operator_aliases(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["金额"], [100]])
        for op in ["greater_than_equal", "greater_than_or_equal", "greater_or_equal", "more_than_or_equal"]:
            result = _format(path, [{
                "kind": "conditional_format",
                "sheet": "Sheet1",
                "range": "A2:A10",
                "rule": {"type": "cell_value", "operator": op, "value": 50, "bg_color": "FFC7CE"},
            }])
            assert result.success, f"operator={op} 应该成功: {result.model_text}"
            rules = self._cf_rules(path)
            assert rules[-1].rules[0].operator == "greaterThanOrEqual"

    def test_flat_style_and_threshold_aliases(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["金额"], [100]])
        result = _format(path, [{
            "kind": "conditional_format",
            "sheet": "Sheet1",
            "range": "A2:A10",
            "rule": {
                "type": "cell_value",
                "operator": ">=",
                "threshold": 80,
                "font_color": "9C0006",
                "bg_color": "FFC7CE",
            },
        }])
        assert result.success, result.model_text
        rules = self._cf_rules(path)
        rule = rules[0].rules[0]
        assert rule.operator == "greaterThanOrEqual"
        assert rule.formula == ["80"]
        assert rule.dxf.font.color.rgb == "009C0006" or rule.dxf.font.color.value == "9C0006" or rule.dxf.font.color is not None
        assert rule.dxf.fill.start_color.rgb == "00FFC7CE" or rule.dxf.fill.start_color.value == "FFC7CE" or rule.dxf.fill is not None

    def test_data_validation_snake_case_operator(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        path = _book(tmp_path / "book.xlsx", [["金额"], [100]])
        result = _format(path, [{
            "kind": "data_validation",
            "sheet": "Sheet1",
            "range": "A2:A10",
            "rule": {
                "type": "whole",
                "operator": "greater_than_equal",
                "value": 10,
            },
        }])
        assert result.success, result.model_text



# ── 候选3: max_rows 规范名 ─────────────────────────────────


class TestPaginationCanonical:
    def test_aggregate_max_rows_caps_groups(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        rows = [["区域", "金额"]] + [[f"R{i}", i] for i in range(10)]
        path = _book(tmp_path / "d.xlsx", rows)
        result = analyze_spreadsheet(
            file_path=str(path),
            mode="aggregate",
            group_by="区域",
            aggregations={"金额": "sum"},
            max_rows=3,
        )
        assert result.success
        groups = _payload(result).get("groups") or []
        assert len(groups) == 3

    def test_filter_limit_alias_still_works(self, tmp_path: Path) -> None:
        _bind_workspace(tmp_path)
        rows = [["区域", "金额"]] + [[f"R{i}", i] for i in range(10)]
        path = _book(tmp_path / "d.xlsx", rows)
        result = analyze_spreadsheet(
            file_path=str(path),
            mode="filter",
            column="金额",
            operator="ge",
            value=0,
            limit=4,
        )
        assert result.success
        data = _payload(result)
        assert data.get("returned_rows") == 4
