"""Search -> exact source -> specification/examples through real session tools."""
from __future__ import annotations

import ast
from dataclasses import replace
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from jsonschema import validate

from excelmanus.knowledge.documents import TOPIC_BY_ID
from excelmanus.knowledge.examples import EXAMPLES, steps_for
from excelmanus.knowledge.portal import next_call
from excelmanus.knowledge.reading import Document, find_literal
from excelmanus.tools.context import ToolCallContext, bind_call, binding_from_engine, reset_call
from excelmanus.tools.registry import ToolDef
from tests.test_knowledge_portal import make_engine, invoke, read, all_items


def body(engine, call):
    chunks, results = [], []
    while call:
        result = invoke(engine, call)
        assert result["status"] == "ok", result
        chunks.append(result.get("content", ""))
        results.append(result)
        call = result.get("next_call")
    return "".join(chunks), results


def quoted(text, citation):
    lines = text.splitlines(keepends=True)
    start = sum(map(len, lines[:citation["line_start"] - 1])) + citation["start_column"] - 1
    end = sum(map(len, lines[:citation["line_end"] - 1])) + citation["end_column_exclusive"] - 1
    return text[start:end]


def test_headings_ignore_code_and_have_unique_stable_anchors():
    doc = Document("doc:test", "test", "# Root\n## A {#a}\none\n```md\n## fake\n```\n### Child\ntwo\n## A-2 {#a-2}\n## A {#a}\nthree\n", "docs", "test")
    assert [s.anchor for s in doc.sections] == ["root", "a", "child", "a-2", "a-3"]
    start, end = doc.bounds("a")
    assert "two" in doc.text[start:end] and "three" not in doc.text[start:end]
    assert doc.sections[2].hierarchy == ("Root", "A", "Child")
    with pytest.raises(ValueError):
        doc.bounds(line_start=0)
    with pytest.raises(ValueError):
        doc.bounds(line_end=False)


def test_toc_anchor_and_citation_roundtrip(make_engine):
    engine = make_engine()
    toc = all_items(engine, next_call("knowledge_toc", "doc:execution"))
    section = next(item for item in toc if item["ref"] == "doc:execution#sdk")
    text, pages = body(engine, section["next_call"])
    original = TOPIC_BY_ID["execution"].read()
    assert text == quoted(original, section["citation"])
    assert "import em" in text and "## 并发" not in text
    assert all(page["content"] == quoted(original, page["citation"]) for page in pages)
    missing = read(engine, "doc:execution#missing")
    assert missing["status"] == "not_found"
    assert invoke(engine, missing["toc_call"])["items"]


@pytest.mark.parametrize("query,scope,target", [
    ("observe_spreadsheet", "tools", "tool:observe_spreadsheet"),
    ("VERSION_CONFLICT", "errors", "error:VERSION_CONFLICT"),
    ("上下文压缩", "docs", "doc:context"),
    ("workspace sandbox permission", "docs", "doc:execution"),
])
def test_ranked_search_returns_useful_first_page(make_engine, query, scope, target):
    engine = make_engine()
    result = invoke(engine, next_call("knowledge_search", query, scope=scope, limit=5))
    assert target in {item["ref"].partition("#")[0] for item in result["items"]}
    assert all(item["kind"] == scope for item in result["items"])
    scores = [item["score"] for item in result["items"]]
    assert scores == sorted(scores, reverse=True)
    if "_" in query:
        assert result["items"][0]["ref"] == target
    for hit in result["items"]:
        exact, _ = body(engine, hit["next_call"])
        assert hit["snippet"] in exact
    assert not engine._active_skills


def test_search_does_not_load_tools_or_read_unavailable_skill_bodies(make_engine):
    engine = make_engine()
    skill = engine._skill_router._loader.get_skillpack("data_basic")
    engine._skill_router._loader._skillpacks["data_basic"] = replace(skill,
        disable_model_invocation=True, instructions="NON_INVOKABLE_PRIVATE_TEXT")
    before = set(engine._loaded_tool_names)
    invoke(engine, next_call("knowledge_search", "difference", scope="all"))
    assert engine._loaded_tool_names == before
    result = invoke(engine, next_call("knowledge_search", "NON_INVOKABLE_PRIVATE_TEXT", scope="skills"))
    assert result["items"] == []


def test_search_covers_nested_schema_and_supplemental_resource(make_engine):
    engine = make_engine()
    engine.registry.register_tool(ToolDef(name="mcp_metrics", description="Metrics",
        input_schema={"type": "object", "properties": {"options": {"type": "object", "properties": {
            "resolution": {"type": "string", "description": "SPECTRAL_BIN_WIDTH"}}}}},
        func=lambda **kwargs: {}, write_effect="none"))
    hits = invoke(engine, next_call("knowledge_search", "SPECTRAL_BIN_WIDTH", scope="tools"))["items"]
    assert hits and hits[0]["ref"] == "schema:mcp_metrics"
    assert "SPECTRAL_BIN_WIDTH" in body(engine, hits[0]["next_call"])[0]
    loader = engine._skill_router._loader
    skill = loader.get_skillpack("data_basic")
    loader._skillpacks["data_basic"] = replace(skill, resource_contents={"references/manual.md": "# 资源\n稀疏矩阵归一化流程仅见于补充资源\n"})
    hits = invoke(engine, next_call("knowledge_search", "稀疏矩阵归一化", scope="skills"))["items"]
    assert any(hit["ref"].startswith("resource:data_basic/references/manual.md") for hit in hits)


def test_find_is_literal_and_cites_exact_original_offsets(make_engine):
    engine = make_engine()
    result = invoke(engine, next_call("knowledge_find", "Workbook.save", ref="doc:execution#sdk"))
    assert result["items"]
    original = TOPIC_BY_ID["execution"].read()
    for hit in result["items"]:
        assert quoted(original, hit["citation"]) == "Workbook.save"
        assert quoted(original, hit["snippet_citation"]) == hit["snippet"]
    assert not invoke(engine, next_call("knowledge_find", ".*", ref="doc:execution"))["items"]
    assert invoke(engine, next_call("knowledge_find", "", ref="doc:execution"))["status"] == "invalid_query"
    doc = Document("doc:unicode", "Unicode", "İ中文\nSTRASSE Straße", "docs", "test")
    for lo, hi in find_literal(doc, "straße"):
        assert quoted(doc.text, doc.citation(lo, hi)) == doc.text[lo:hi]


def test_read_exact_lines_and_stale_citations_restart(make_engine):
    engine = make_engine()
    result = invoke(engine, next_call("knowledge_search", "审批", scope="docs", limit=2))
    hit = result["items"][0]
    root = hit["ref"].partition("#")[0]
    topic = TOPIC_BY_ID[root[4:]]
    from unittest.mock import patch
    with patch.object(type(topic), "read", lambda self: "# Changed\nnew content\n"):
        stale = invoke(engine, hit["next_call"])
    assert stale["status"] == "stale" and stale["restart_call"]
    original = TOPIC_BY_ID["execution"].read()
    sliced = invoke(engine, next_call("knowledge_read", "doc:execution", line_start=3, line_end=5))
    assert sliced["content"] == "".join(original.splitlines(keepends=True)[2:5])
    assert invoke(engine, next_call("knowledge_read", "doc:execution", line_start=5, line_end=3))["status"] == "invalid_query"


def test_search_pagination_preserves_filters_and_is_bounded(make_engine):
    engine = make_engine()
    call = next_call("knowledge_search", "schema", scope="tools", limit=2)
    seen = set()
    while call:
        result = invoke(engine, call)
        assert len(result["items"]) <= 2
        assert len(json.dumps(result, ensure_ascii=False)) < 12000
        for item in result["items"]:
            assert item["ref"] not in seen
            seen.add(item["ref"])
        call = result.get("next_call")
        if call:
            assert call["arguments"]["scope"] == "tools" and call["arguments"]["limit"] == 2
    assert len(seen) > 2


def test_related_links_have_forward_and_reverse_navigation(make_engine):
    engine = make_engine()
    relations = all_items(engine, next_call("knowledge_related", "doc:execution"))
    assert any(item["relation"] == "references" and item["ref"] == "tool:run_code" for item in relations)
    assert any(item["relation"] == "referenced_by" and item["ref"] == "doc:workflows" for item in relations)
    for item in relations:
        assert invoke(engine, item["next_call"])["status"] in {"ok", "unavailable"}


def test_specs_preserve_live_schema_and_filter_example_language(make_engine):
    from excelmanus.tools.reference_contract import augment_reference_schema
    engine = make_engine()
    raw, pages = body(engine, next_call("knowledge_spec", "apply_spreadsheet_changes", language="python"))
    spec = json.loads(raw)
    assert spec["input_schema"] == augment_reference_schema(engine.registry.get_tool("apply_spreadsheet_changes").input_schema)
    assert spec["examples"]
    for example in spec["examples"]:
        assert example["validation"] == {"schema": "passed", "execution": "not_run"}
        assert set(example["snippets"]) == {"python"}
        ast.parse(example["snippets"]["python"])
    raw, _ = body(engine, next_call("knowledge_spec", "apply_spreadsheet_changes", language="json", examples_only=True))
    examples = json.loads(raw)
    assert "input_schema" not in examples
    assert all(set(item["snippets"]) == {"json"} for item in examples["examples"])
    assert all(page["citation"]["content_revision"] == pages[0]["citation"]["content_revision"] for page in pages)
    for page in pages[:2]:
        assert invoke(engine, page["citation"]["fetch_call"])["content"] == page["content"]


def test_example_keyword_search_keeps_language_and_content_revision_consistent(make_engine):
    engine = make_engine()
    result = invoke(engine, next_call("knowledge_examples", "修改", language="python"))
    assert result["items"]
    for hit in result["items"]:
        fetched = invoke(engine, hit["next_call"])
        assert fetched["status"] == "ok", fetched
        assert hit["next_call"]["arguments"]["language"] == "python"
    toc = invoke(engine, next_call("knowledge_toc", "doc:execution#sdk"))
    assert len(toc["items"]) == 1 and toc["items"][0]["ref"] == "doc:execution#sdk"


def test_all_examples_validate_and_observe_edit_really_executes(make_engine):
    engine = make_engine()
    for example in EXAMPLES:
        raw, _ = body(engine, next_call("knowledge_read", "example:" + example.id))
        data = json.loads(raw)
        assert data["validation"]["schema"] == "passed"
        for step in data["steps"]:
            validate(step["call"]["arguments"], engine.registry.get_tool(step["call"]["name"]).input_schema)
    create = steps_for("create-workbook")[0]["call"]
    create["arguments"]["file_path"] = "example.xlsx"
    results = {}
    token = bind_call(ToolCallContext(binding_from_engine(engine)))
    try:
        assert engine.registry.call_tool(create["name"], create["arguments"]).success
        for step in steps_for("observe-edit"):
            arguments = dict(step["call"]["arguments"])
            for key, binding in step.get("bindings", {}).items():
                arguments[key] = results[binding["step"]][binding["field"]]
            result = engine.registry.call_tool(step["call"]["name"], arguments)
            assert result.success, result.model_text
            results[step["id"]] = result.value
    finally:
        reset_call(token)
    from openpyxl import load_workbook
    workbook = load_workbook(Path(engine.config.workspace_root) / "example.xlsx")
    try:
        assert workbook["明细"]["B3"].value == 5
    finally:
        workbook.close()


def test_incompatible_and_denied_examples_never_offer_stale_calls(make_engine):
    engine = make_engine()
    engine._current_chat_mode = "read"
    assert read(engine, "example:observe-edit")["status"] == "unavailable"
    assert invoke(engine, next_call("knowledge_spec", "apply_spreadsheet_changes"))["status"] == "unavailable"
    engine._current_chat_mode = "write"
    engine.registry.get_tool("observe_spreadsheet").input_schema = {"type": "object", "required": ["new_required_field"]}
    blocked = read(engine, "example:observe-range")
    assert blocked["status"] == "unavailable"
    assert "snippets" not in blocked["data"]


@pytest.mark.asyncio
@pytest.mark.parametrize("arguments", [
    {"query_type": "knowledge_search", "query": "版本", "scope": "docs", "limit": 2},
    {"query_type": "knowledge_toc", "query": "doc:execution"},
    {"query_type": "knowledge_find", "query": "权限", "ref": "doc:execution"},
    {"query_type": "knowledge_related", "query": "doc:execution"},
    {"query_type": "knowledge_spec", "query": "observe_spreadsheet", "examples_only": True},
    {"query_type": "knowledge_examples", "query": "observe_spreadsheet", "language": "python"},
])
async def test_docs_operations_pass_real_dispatch_output_contract(make_engine, arguments):
    engine = make_engine()
    call = SimpleNamespace(id="docs-operation", function=SimpleNamespace(name="introspect_capability", arguments=json.dumps(arguments)))
    result = await engine._tool_runtime.execute(call, None, None, 1)
    assert result.success, result.result
    payload = json.loads(result.result)
    assert payload["status"] == "ok" and payload["portal_version"] == 2
