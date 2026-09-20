"""Output contracts through the real dispatcher, SDK subprocess and discovery."""
from __future__ import annotations

import json
from itertools import count
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest

from excelmanus.agent.session import AgentEngine
from excelmanus.code_mode import _payload_from_tool_result, _sdk_signature_line, render_sdk_source
from excelmanus.config import ExcelManusConfig
from excelmanus.engine_core.tool_result import ToolResult, coerce_legacy_result, ok_result, result_value
from excelmanus.events import EventType
from excelmanus.tools.output_contracts import output_schema_for, validate_declared_output_schema
from excelmanus.tools.registry import ToolDef, ToolRegistry

_CALL_IDS = count()


def make_engine(tmp_path):
    registry = ToolRegistry()
    registry.register_builtin_tools(str(tmp_path))
    engine = AgentEngine(ExcelManusConfig(
        api_key="test", base_url="https://test.invalid/v1", model="test-model",
        workspace_root=str(tmp_path), main_model_vision="false", memory_enabled=False,
        tool_schema_validation_mode="enforce", code_policy_enabled=False,
    ), registry)
    engine._full_access_enabled = True
    return engine


async def execute(engine, name, arguments=None, events=None, cid=None):
    from excelmanus.tools.catalog import catalog_from_engine

    # Normal Driver request compilation refreshes this binding before dispatch.
    catalog_from_engine(engine)
    tc = SimpleNamespace(id=cid or f"contract-call-{next(_CALL_IDS)}", function=SimpleNamespace(name=name, arguments=json.dumps(arguments or {})))
    return await engine._execute_tool_call(tc, None, events.append if events is not None else None, 0)


def register(engine, name, func, schema, *, effect="none"):
    tool = ToolDef(name=name, description="output fixture", func=func, write_effect=effect,
                   input_schema={"type": "object", "properties": {}}, output_schema=schema)
    engine.registry.register_tool(tool)
    return tool


def test_complete_builtin_and_session_catalog_has_declared_outputs(tmp_path):
    engine = make_engine(tmp_path)
    from excelmanus.tools.search_tools import get_tools

    tools = engine.registry.get_all_tools() + get_tools(SimpleNamespace())
    missing = [tool.name for tool in tools if not output_schema_for(tool.name, tool_def=tool)]
    assert missing == []


@pytest.mark.parametrize("raw", [None, False, 7, 2.5, [1, None, {"key": "value"}]])
def test_legacy_json_values_keep_their_type_in_sdk(raw):
    result = coerce_legacy_result(raw)
    assert result_value(result) == raw
    assert _payload_from_tool_result(result)["value"] == raw


@pytest.mark.parametrize("text", ['{"raw_input":"hello"}', '[1, 2]', '123', 'null'])
def test_explicit_text_is_not_reparsed_as_json(text):
    result = ToolResult.from_text(text)
    assert result_value(result) == text
    assert _payload_from_tool_result(result, tool_name="memory_read_topic")["value"] == text


def test_full_array_union_ref_and_constraints_are_validated():
    schema = {"type": "array", "items": {"$ref": "#/$defs/entry"},
              "$defs": {"entry": {"type": ["integer", "null"], "minimum": 0}}}
    assert validate_declared_output_schema([None] + list(range(150)), schema) == []
    errors = validate_declared_output_schema(list(range(150)) + [-1], schema)
    assert any("150" in error for error in errors)
    assert validate_declared_output_schema(True, {"type": "integer"})
    assert validate_declared_output_schema({"unexpected": 1}, {"type": "object", "additionalProperties": False})
    assert validate_declared_output_schema("other", {"enum": ["expected"]})
    assert validate_declared_output_schema(2, {"const": 1})
    assert validate_declared_output_schema(1, {"oneOf": [{"type": "integer"}, {"type": "number"}]})


def test_invalid_and_remote_schemas_fail_without_network(monkeypatch):
    import urllib.request

    def forbidden(*args, **kwargs):
        raise AssertionError("output validation must not fetch a schema")

    monkeypatch.setattr(urllib.request, "urlopen", forbidden)
    assert validate_declared_output_schema({}, {"$ref": "https://test.invalid/schema"})
    assert validate_declared_output_schema({}, {"type": "made-up"})


def test_generated_return_annotations_resolve_enums_refs_and_spill_variants(tmp_path):
    from typing import get_type_hints, Literal

    engine = make_engine(tmp_path)
    tool = register(engine, "enum_result", lambda: "yes", {
        "$defs": {"answer": {"type": "string", "enum": ["yes", "no"]}}, "$ref": "#/$defs/answer",
    })
    assert _sdk_signature_line(tool).endswith(" -> str")
    namespace = {}
    exec(compile(render_sdk_source([tool, engine.registry.get_tool("read_text_file")]), "em.py", "exec"), namespace)
    assert get_type_hints(namespace["enum_result"], globalns=namespace)["return"] == Literal["yes", "no"]
    assert get_type_hints(namespace["read_text_file"], globalns=namespace)["return"] == (dict | list | str)


@pytest.mark.asyncio
async def test_custom_contract_discovery_native_and_real_sdk_share_types(tmp_path):
    engine = make_engine(tmp_path)
    schema = {"type": "array", "items": {"type": ["integer", "null"]}}
    tool = register(engine, "typed_rows", lambda: [1, None, 3], schema)
    register(engine, "text_number", lambda: ToolResult.from_text("123"), {"type": "string"})
    register(engine, "null_result", lambda: None, {"type": "null"})
    native = await execute(engine, "typed_rows")
    assert native.success and native.structured.value == [1, None, 3]
    detail = await execute(engine, "introspect_capability", {"query_type": "tool_detail", "query": "typed_rows.output"})
    assert detail.success and '"integer"' in detail.result and '"null"' in detail.result
    assert "-> list" in _sdk_signature_line(tool)
    source = render_sdk_source([tool])
    namespace = {}
    exec(compile(source, "em.py", "exec"), namespace)
    assert namespace["typed_rows"].__annotations__["return"] == "list[int | None]"
    sdk = await execute(engine, "run_code", {"code": "import em, json\nprint(json.dumps([em.typed_rows(), em.text_number(), em.null_result()]))"})
    assert sdk.success, sdk.result
    assert json.loads(sdk.structured.value["stdout_tail"]) == [[1, None, 3], "123", None]


@pytest.mark.asyncio
async def test_output_violation_preserves_commit_and_failure_event(tmp_path):
    engine = make_engine(tmp_path)
    count = 0
    def write():
        nonlocal count
        count += 1
        (tmp_path / "committed.txt").write_text("saved")
        return ok_result({"file_path": "committed.txt", "content_version": "sha256:observed", "operation_id": "op-1"})
    register(engine, "bad_receipt", write, {"type": "object", "required": ["receipt"]}, effect="workspace_write")
    events = []
    result = await execute(engine, "bad_receipt", events=events)
    assert count == 1 and (tmp_path / "committed.txt").read_text() == "saved"
    assert not result.success and result.error == "SDK_CONTRACT_VIOLATION"
    assert result.structured.value["execution_completed"] is True
    assert result.structured.value["operation_id"] == "op-1"
    assert result.structured.ui_meta.files == ["committed.txt"]
    ends = [event for event in events if event.event_type == EventType.TOOL_CALL_END]
    assert len(ends) == 1 and not ends[0].success and ends[0].error == result.error
    assert ends[0].ui["content_version"] == "sha256:observed"
    sdk = await execute(engine, "run_code", {"code": "import em, json\ntry:\n    em.bad_receipt()\nexcept em.HostToolError as e:\n    print(json.dumps([e.code, e.details['execution_completed'], e.details['operation_id']]))"})
    assert sdk.success, sdk.result
    assert json.loads(sdk.structured.value["stdout_tail"]) == ["SDK_CONTRACT_VIOLATION", True, "op-1"]
    assert count == 2  # One Native call and one explicit SDK call, no automatic retry.


@pytest.mark.asyncio
async def test_output_override_and_catalog_digest_follow_actual_definition(tmp_path):
    engine = make_engine(tmp_path)
    tool = engine.registry.get_tool("list_directory")
    before = engine.registry.catalog_digest()
    tool.output_schema = {"type": "string"}
    assert engine.registry.catalog_digest() != before
    assert _sdk_signature_line(tool).splitlines()[0].endswith(" -> str")
    detail = await execute(engine, "introspect_capability", {"query_type": "tool_detail", "query": "list_directory.output"})
    assert '"string"' in detail.result
    result = await execute(engine, "list_directory")
    assert result.error == "SDK_CONTRACT_VIOLATION"


@pytest.mark.asyncio
async def test_builtin_text_word_and_run_code_results_through_dispatcher(tmp_path):
    from docx import Document
    from excelmanus.workbook_commit import content_version_of_file

    engine = make_engine(tmp_path)
    doc = Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("original paragraph")
    doc.save(tmp_path / "report.docx")
    for name, args in [
        ("task_create", {"title": "contract test", "subtasks": ["check"]}),
        ("task_update", {"task_index": 0, "status": "in_progress"}),
        ("read_word", {"file_path": "report.docx"}),
        ("inspect_word", {"file_path": "report.docx"}),
        ("search_word", {"file_path": "report.docx", "query": "original"}),
        ("write_word", {"file_path": "report.docx", "expected_version": content_version_of_file(tmp_path / "report.docx"),
                        "operations": [{"action": "append", "text": "added paragraph"}]}),
    ]:
        result = await execute(engine, name, args, cid=name)
        assert result.success, (name, result.result)
    assert Document(tmp_path / "report.docx").paragraphs[-1].text == "added paragraph"


@pytest.mark.asyncio
async def test_completed_multi_question_answers_are_structured_and_timeout_fails(tmp_path):
    engine = make_engine(tmp_path)
    engine._question_resolver = AsyncMock(side_effect=["1", "1"])
    args = {"questions": [{"text": text, "options": [{"label": "yes"}]} for text in ("first?", "second?")]}
    result = await execute(engine, "ask_user", args)
    assert result.success and isinstance(result.structured.value, list)
    assert len(result.structured.value) == 2
    engine.handle_ask_user_blocking = AsyncMock(return_value="等待用户回答超时，已取消问题。")
    result = await execute(engine, "ask_user", args)
    assert not result.success and result.structured.error.code == "CANCELLED"


@pytest.mark.asyncio
async def test_output_field_discovery_preserves_nullable_array(tmp_path):
    engine = make_engine(tmp_path)
    register(engine, "nullable_rows", lambda: None, {
        "type": "object", "properties": {"rows": {"type": ["array", "null"], "items": {"type": "integer"}}},
    })
    result = await execute(engine, "introspect_capability", {"query_type": "tool_detail", "query": "nullable_rows.output.rows"})
    assert result.success
    assert '"array", "null"' in result.result


@pytest.mark.asyncio
async def test_pending_approval_is_not_validated_as_completed_output(tmp_path):
    engine = make_engine(tmp_path)
    engine._full_access_enabled = False
    (tmp_path / "keep.txt").write_text("keep")
    result = await execute(engine, "delete_file", {"file_path": "keep.txt"})
    assert result.pending_approval and result.success
    assert (tmp_path / "keep.txt").exists()


@pytest.mark.asyncio
@pytest.mark.parametrize("use_sync", [False, True])
async def test_mcp_declared_output_keeps_structured_value_and_remote_errors(tmp_path, use_sync):
    from excelmanus.mcp.manager import make_tool_def

    engine = make_engine(tmp_path)
    client = SimpleNamespace(call_tool=AsyncMock())
    schema = {"type": "object", "required": ["count"], "properties": {"count": {"type": "integer"}}, "additionalProperties": False}
    definition = SimpleNamespace(name="count", description="fixture", inputSchema={"type": "object"}, outputSchema=schema)
    tool = make_tool_def("fixture", client, definition, scope="readonly")
    if use_sync:
        tool.async_func = None
    engine.registry.register_tool(tool)
    client.call_tool.return_value = SimpleNamespace(
        content=[SimpleNamespace(type="text", text="one item found")], structuredContent={"count": 1}, isError=False,
    )
    result = await execute(engine, tool.name)
    assert result.success and result.structured.value == {"count": 1}
    assert result.result == "one item found"
    detail = await execute(engine, "introspect_capability", {"query_type": "tool_detail", "query": tool.name + ".output"})
    assert '"count"' in detail.result
    sdk = await execute(engine, "run_code", {"code": "import em\nprint(em.mcp_fixture_count()['count'])"})
    assert sdk.success and sdk.structured.value["stdout_tail"].strip() == "1"
    client.call_tool.return_value.structuredContent = {"count": "invalid"}
    engine._tool_dispatcher.begin_call_budget(None)  # Next turn: do not reuse the prior read result.
    result = await execute(engine, tool.name)
    assert not result.success and result.error == "SDK_CONTRACT_VIOLATION"
    # Diagnose before another read, so the existing repeated-failure guard does
    # not deliberately reject the identical call before it reaches the server.
    await execute(engine, "introspect_capability", {"query_type": "tool_detail", "query": tool.name + ".output"})
    client.call_tool.return_value.isError = True
    client.call_tool.return_value.structuredContent = {"count": 1}
    calls_before = client.call_tool.call_count
    result = await execute(engine, tool.name)
    assert not result.success and result.structured.error.code == "TOOL_ERROR"
    assert client.call_tool.call_count > calls_before  # Invalid output was not cached as a successful read.


@pytest.mark.asyncio
async def test_sync_delegation_preserves_json_looking_text(tmp_path):
    from excelmanus.subagent.models import SubagentResult

    engine = make_engine(tmp_path)
    engine.delegate_to_subagent = AsyncMock(return_value=SubagentResult(
        stop_reason="completed", output='{"answer":42}', subagent_name="subagent", permission_mode="read", conversation_id="child",
    ))
    result = await execute(engine, "delegate", {"task": "return text"})
    assert result.success and result.structured.value == '{"answer":42}'
    assert _payload_from_tool_result(result.structured, "delegate")["value"] == '{"answer":42}'


@pytest.mark.asyncio
@pytest.mark.parametrize("payload", [{"count": 1}, None])
async def test_shadow_input_diagnostics_do_not_invalidate_strict_output(tmp_path, payload):
    engine = make_engine(tmp_path)
    schema = ({"type": "object", "properties": {"count": {"type": "integer"}}, "required": ["count"], "additionalProperties": False}
              if payload is not None else {"type": "null"})
    tool = register(engine, "shadow_output", lambda **kw: ToolResult(success=True, model_text="one", value=payload, value_is_set=True), schema)
    tool.input_schema = {"type": "object", "properties": {"label": {"type": "string"}}, "additionalProperties": False}
    engine.registry.configure_schema_validation(mode="shadow", canary_percent=100, strict_path=False)
    result = await execute(engine, tool.name, {"unknown": 1})
    assert result.success
    assert "schema_violations" in result.result
    assert result.structured.value == payload
    bridged = _payload_from_tool_result(result.structured, tool.name, tool_def=tool)
    assert bridged["ok"] is True and bridged["value"] == payload
