"""第 5 批：工作台视图合同、作用域 fail-closed、operations 一次提交。"""

from __future__ import annotations

from pathlib import Path
import asyncio
import threading

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from excelmanus.workbook_commit import content_version_of_file
from tests.test_api import _make_transport, _setup_api_globals, _test_config


pytestmark = pytest.mark.asyncio


def _xlsx(path: Path, *, rows: int = 4, formula: bool = False) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 2024
    ws["B1"] = "=1+1" if formula else "name"
    for i in range(2, rows + 1):
        ws.cell(i, 1, i)
        ws.cell(i, 2, f"r{i}")
    wb.save(path)
    wb.close()
    return path


async def _session_id(client: AsyncClient) -> str:
    response = await client.post("/api/v1/sessions", json={})
    assert response.status_code == 200
    return response.json()["id"]


class TestWorkbookViewContract:
    async def test_first_window_is_active_sheet_and_sparse(self, tmp_path: Path, monkeypatch) -> None:
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        from openpyxl.styles import PatternFill

        path = _xlsx(tmp_path / "book.xlsx", formula=True)
        wb = load_workbook(path)
        ws = wb.create_sheet("Actual active")
        ws["A1"] = "first paint"
        ws["C2"] = "=1+2"
        ws["F5"].fill = PatternFill("solid", fgColor="FF0000")
        ws.merge_cells("A1:B1")
        ws.column_dimensions["A"].width = 22
        ws.row_dimensions[1].height = 32
        wb.active = 1
        wb.save(path)
        wb.close()

        def no_random_read(*args, **kwargs):
            raise AssertionError("Read-only XML must be scanned once, not once per cell")

        monkeypatch.setattr(ReadOnlyWorksheet, "cell", no_random_read)
        with _setup_api_globals(config=_test_config(workspace_root=str(tmp_path))):
            async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
                sid = await _session_id(client)
                base = {"path": "book.xlsx", "session_id": sid}
                first = (await client.get("/api/v1/files/excel/view", params={**base, "with_styles": "0"})).json()
                styled = (await client.get("/api/v1/files/excel/view", params={**base, "expected_version": first["content_version"]})).json()
                page = (await client.get("/api/v1/files/excel/view", params={**base, "sheet": "Sheet1", "rect": "A201:AX400"})).json()
        assert len(first["sheets"]) == 2
        assert len(first["windows"]) == 1
        assert first["windows"][0]["sheet"] == "Actual active"
        assert set(first["windows"][0]["cells"]) == {"1,1", "2,3"}
        assert first["windows"][0]["cells"]["2,3"]["f"] == "=1+2"
        assert first["windows"][0]["cells"]["2,3"]["cached"] == "no"
        win = styled["windows"][0]
        assert "s" in win["cells"]["5,6"]
        assert win["merges"] == [{"min_row": 1, "min_col": 1, "max_row": 1, "max_col": 2}]
        assert win["col_widths"]["A"] == 22
        assert win["row_heights"]["1"] == 32
        assert page["windows"][0]["cells"] == {}
        assert page["coverage"]["loaded"][0]["r1"] == 400

    @pytest.mark.parametrize("operation", ["view", "write"])
    async def test_workbook_work_does_not_block_event_loop(self, tmp_path: Path, monkeypatch, operation: str) -> None:
        import excelmanus.workbook.snapshot as snapshots
        from excelmanus.workspace.file_service import WorkspaceFileService

        path = _xlsx(tmp_path / "book.xlsx")
        entered, release = threading.Event(), threading.Event()
        owner, attr = (snapshots, "project_view") if operation == "view" else (WorkspaceFileService, "update_with_builder")
        original = getattr(owner, attr)

        def held(*args, **kwargs):
            entered.set()
            assert release.wait(2), "workbook operation blocked the ASGI event loop"
            return original(*args, **kwargs)

        monkeypatch.setattr(owner, attr, held)
        with _setup_api_globals(config=_test_config(workspace_root=str(tmp_path))):
            async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
                sid = await _session_id(client)
                if operation == "view":
                    request = client.get("/api/v1/files/excel/view", params={"path": "book.xlsx", "session_id": sid})
                else:
                    request = client.post("/api/v1/files/excel/write", json={"path": "book.xlsx", "session_id": sid,
                        "expected_version": content_version_of_file(path), "changes": [{"sheet": "Sheet1", "cell": "A1", "value": 9}]})
                pending = asyncio.create_task(request)
                try:
                    assert await asyncio.to_thread(entered.wait, 1)
                    # This coroutine must progress while openpyxl/commit is held.
                    await asyncio.sleep(0)
                finally:
                    release.set()
                assert (await pending).status_code == 200

    async def test_view_budget_and_stale_version(self, tmp_path: Path) -> None:
        _xlsx(tmp_path / "book.xlsx")
        with _setup_api_globals(config=_test_config(workspace_root=str(tmp_path))):
            async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
                sid = await _session_id(client)
                base = {"path": "book.xlsx", "session_id": sid}
                too_large = await client.get("/api/v1/files/excel/view", params={**base, "rect": "A1:AX1000"})
                stale = await client.get("/api/v1/files/excel/view", params={**base, "expected_version": "sha256:old"})
        assert too_large.status_code == 400
        assert too_large.json()["code"] == "VIEW_TOO_LARGE"
        assert stale.status_code == 409
        assert stale.json()["code"] == "STALE_VIEW"

    async def test_view_requires_scope(self, tmp_path: Path) -> None:
        _xlsx(tmp_path / "book.xlsx")
        config = _test_config(workspace_root=str(tmp_path))
        with _setup_api_globals(config=config):
            async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
                response = await client.get("/api/v1/files/excel/view", params={"path": "book.xlsx"})
        assert response.status_code == 400
        assert response.json().get("code") == "FILE_SCOPE_REQUIRED"

    async def test_view_returns_r1_formula_and_used(self, tmp_path: Path) -> None:
        _xlsx(tmp_path / "book.xlsx", rows=400, formula=True)
        config = _test_config(workspace_root=str(tmp_path))
        with _setup_api_globals(config=config):
            async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
                sid = await _session_id(client)
                response = await client.get(
                    "/api/v1/files/excel/view",
                    params={"path": "book.xlsx", "session_id": sid, "rect": "A1:B200"},
                )
        assert response.status_code == 200
        data = response.json()
        assert data["file"]["relative"].endswith("book.xlsx")
        assert data["file"]["workspaceKey"]
        assert data["sheets"][0]["used"]["rows"] == 400
        first = data["windows"][0]["cells"]["1,1"]
        assert first["t"] == "n"
        assert first["v"] == 2024
        formula = data["windows"][0]["cells"]["1,2"]
        assert formula["f"] == "=1+1"
        assert any(item["r0"] == 201 for item in data["coverage"]["unloaded"])

    async def test_csv_view_counts_all_rows(self, tmp_path: Path) -> None:
        lines = ["h1,h2"] + [f"{i},x" for i in range(400)]
        (tmp_path / "rows.csv").write_text("\n".join(lines), encoding="utf-8")
        config = _test_config(workspace_root=str(tmp_path))
        with _setup_api_globals(config=config):
            async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
                sid = await _session_id(client)
                response = await client.get(
                    "/api/v1/files/excel/view",
                    params={"path": "rows.csv", "session_id": sid, "rect": "A1:B50"},
                )
        data = response.json()
        assert response.status_code == 200
        assert data["sheets"][0]["used"]["rows"] == 401
        assert data["windows"][0]["cells"]["1,1"]["v"] == "h1"

    async def test_write_operations_one_commit(self, tmp_path: Path) -> None:
        path = _xlsx(tmp_path / "book.xlsx")
        version = content_version_of_file(path)
        config = _test_config(workspace_root=str(tmp_path))
        with _setup_api_globals(config=config):
            async with AsyncClient(transport=_make_transport(), base_url="http://test") as client:
                sid = await _session_id(client)
                response = await client.post(
                    "/api/v1/files/excel/write",
                    json={
                        "path": "book.xlsx",
                        "session_id": sid,
                        "expected_version": version,
                        "operations": [
                            {
                                "op": "set_values",
                                "sheet": "Sheet1",
                                "cells": [
                                    {"cell": "A1", "value": 9},
                                    {"cell": "C1", "value": "styled", "style": {"bl": 1}},
                                ],
                            },
                            {"op": "set_dims", "sheet": "Sheet1", "columns": {"A": 20}},
                        ],
                    },
                )
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "success"
        assert body["content_version"]
        assert body["operation_id"]
        wb = load_workbook(path)
        assert wb.active["A1"].value == 9
        assert wb.active["C1"].value == "styled"
        assert wb.active["C1"].font.bold
        wb.close()
