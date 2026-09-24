"""edit.kind=pivot / transform：透视写入与四类清洗。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from excelmanus.security import FileAccessGuard
from excelmanus.tools import workbook_tools, reference_tools
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes
from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    workbook_tools.init_guard(workspace)
    reference_tools.init_guard(workspace)
    seed_seen_versions({})


def _book(path: Path, rows: list[list[object]], title: str = "Sheet1") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    wb.save(path)
    wb.close()
    return path


def _edit(path: Path, operations: list[dict]) -> object:
    return apply_spreadsheet_changes(
        file_path=str(path),
        operations=operations,
        expected_version=content_version_of_file(path),
    )


def test_edit_pivot_writes_target_sheet(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(
        tmp_path / "src.xlsx",
        [["部门", "月", "金额"], ["销售", 1, 10], ["销售", 2, 20], ["研发", 1, 5]],
    )
    result = _edit(
        path,
        [{
            "kind": "pivot",
            "sheet": "Sheet1",
            "target_sheet": "透视",
            "index": "部门",
            "columns": "月",
            "values": "金额",
            "aggfunc": "sum",
        }],
    )
    assert result.success, result.model_text
    wb = load_workbook(path)
    assert "透视" in wb.sheetnames
    values = [[cell.value for cell in row] for row in wb["透视"].iter_rows()]
    wb.close()
    header = values[0]
    assert header[0] == "部门"
    body = {row[0]: row for row in values[1:]}
    assert "销售" in body
    assert "研发" in body


def test_transform_dedupe_split_date_phone(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(
        tmp_path / "dirty.xlsx",
        [
            ["编号", "姓名电话", "日期"],
            ["A", "张三,138-0013-8000", "2024/1/2"],
            ["A", "重复,13800138000", "2024-01-02"],
            ["B", "李四,86 13900001111", "20240103"],
        ],
    )
    deduped = _edit(
        path,
        [{"kind": "transform", "sheet": "Sheet1", "action": "dedupe", "key_columns": ["编号"], "keep": "first"}],
    )
    assert deduped.success, deduped.model_text
    split = _edit(
        path,
        [{
            "kind": "transform",
            "sheet": "Sheet1",
            "action": "split",
            "column": "姓名电话",
            "delimiter": ",",
            "into": ["姓名", "电话"],
        }],
    )
    assert split.success, split.model_text
    dated = _edit(
        path,
        [{"kind": "transform", "sheet": "Sheet1", "action": "normalize_date", "column": "日期"}],
    )
    assert dated.success, dated.model_text
    phoned = _edit(
        path,
        [{"kind": "transform", "sheet": "Sheet1", "action": "normalize_phone", "column": "电话"}],
    )
    assert phoned.success, phoned.model_text
    wb = load_workbook(path)
    rows = [[cell.value for cell in row] for row in wb.active.iter_rows()]
    wb.close()
    # split 原位替换源列：姓名电话(列1) → 姓名/电话 占据原位置
    assert rows[0] == ["编号", "姓名", "电话", "日期"]
    assert len(rows) == 3
    assert rows[1][0] == "A"
    assert rows[1][1] == "张三"
    assert str(rows[1][2]).replace(" ", "") == "13800138000"
    assert rows[1][3] == "2024-01-02"
    assert rows[2][0] == "B"
    assert str(rows[2][2]).endswith("13900001111")
    assert rows[2][3] == "2024-01-03"


def test_transform_dedupe_normalized_phone_keeps_earliest_date(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(
        tmp_path / "customers.xlsx",
        [
            ["客户", "手机", "注册日期"],
            ["较晚", "+86 138 0013 8000", "2024年1月2日"],
            ["其他", "13900001111", "2023-05-06T00:00:00"],
            ["最早", "138-0013-8000", "2022.8.14"],
            ["同日后行", "13800138000", "2022/8/14"],
        ],
    )
    wb = load_workbook(path)
    wb.active["A5"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    wb.save(path)
    wb.close()
    result = _edit(
        path,
        [{
            "kind": "transform",
            "sheet": "Sheet1",
            "action": "dedupe",
            "key_columns": ["手机"],
            "key_normalizers": {"手机": "phone"},
            "keep": "earliest",
            "order_by": "注册日期",
        }],
    )
    assert result.success, result.model_text
    assert "rows=4→2 removed=2" in str(result.value["observation"]["operations"])
    assert "output_range=A1:C3" in str(result.value["observation"]["operations"])
    wb = load_workbook(path)
    assert wb.active.max_row == 3
    rows = [[cell.value for cell in row] for row in wb.active.iter_rows()]
    wb.close()
    assert rows == [
        ["客户", "手机", "注册日期"],
        ["其他", "13900001111", "2023-05-06T00:00:00"],
        ["最早", "138-0013-8000", "2022.8.14"],
    ]


def test_transform_dedupe_earliest_requires_parseable_order_column(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = _book(
        tmp_path / "bad-date.xlsx",
        [["手机", "日期"], ["13800138000", "未知"], ["13800138000", "2024-01-01"]],
    )
    result = _edit(
        path,
        [{
            "kind": "transform",
            "sheet": "Sheet1",
            "action": "dedupe",
            "key_columns": ["手机"],
            "key_normalizers": {"手机": "phone"},
            "keep": "earliest",
            "order_by": "日期",
        }],
    )
    assert not result.success
    assert "无法解析" in result.model_text


def test_transform_dedupe_schema_exposes_normalized_ordered_keep() -> None:
    tool = next(item for item in workbook_tools.get_tools() if item.name == "apply_spreadsheet_changes")
    from excelmanus.workbook.contracts import OPERATION_SCHEMAS
    props = OPERATION_SCHEMAS["transform"]["properties"]
    assert props["keep"]["enum"] == ["first", "last", "earliest", "latest"]
    assert props["key_normalizers"]["additionalProperties"]["enum"] == ["phone"]
    assert "order_by" in props


def test_transform_split_infers_action_and_ignores_into_anchor(tmp_path: Path) -> None:
    """R10 复现：省略 action 按字段推断 split；into={start_cell} 是锚点不是列名。"""
    _bind(tmp_path)
    path = _book(
        tmp_path / "split.xlsx",
        [["姓名-工号", "部门"], ["张三-E001", "研发"], ["李四-E002", "市场"]],
    )
    result = _edit(
        path,
        [{
            "kind": "transform",
            "sheet": "Sheet1",
            "column": "姓名-工号",
            "delimiter": "-",
            "new_columns": ["姓名", "工号"],
            "into": {"start_cell": "A1"},
        }],
    )
    assert result.success, result.model_text
    wb = load_workbook(path)
    rows = [[cell.value for cell in row] for row in wb.active.iter_rows()]
    wb.close()
    # split 原位替换：姓名-工号(列0) → 姓名/工号 占首列
    assert rows[0] == ["姓名", "工号", "部门"]
    assert rows[1] == ["张三", "E001", "研发"]


def test_create_workbook_drops_untouched_default_sheet(tmp_path: Path) -> None:
    """R24 复现：create_workbook 建簿后只建命名表，默认 Sheet 不应残留。"""
    _bind(tmp_path)
    target = tmp_path / "new_book.xlsx"
    result = apply_spreadsheet_changes(
        file_path=str(target),
        create=True,
        operations=[
            {"kind": "sheet", "action": "create", "new_name": "订单"},
            {"kind": "write", "sheet": "订单", "start_cell": "A1", "values": [["a", "b"], [1, 2]]},
        ],
    )
    assert result.success, result.model_text
    wb = load_workbook(target)
    assert wb.sheetnames == ["订单"]
    wb.close()


def test_create_workbook_keeps_default_sheet_when_written(tmp_path: Path) -> None:
    _bind(tmp_path)
    target = tmp_path / "direct.xlsx"
    result = apply_spreadsheet_changes(
        file_path=str(target),
        create=True,
        operations=[
            {"kind": "write", "sheet": "Sheet", "start_cell": "A1", "values": [["x"], [9]]},
            {"kind": "sheet", "action": "create", "new_name": "其它"},
        ],
    )
    assert result.success, result.model_text
    wb = load_workbook(target)
    assert "Sheet" in wb.sheetnames
    wb.close()
