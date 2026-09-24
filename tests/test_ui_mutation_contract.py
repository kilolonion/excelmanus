from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from tests.workbook_support import apply_operations_in_memory
from excelmanus.workspace.file_service import WorkspaceFileService


def test_ui_updates_structure_references_before_publication(tmp_path):
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A2"] = 3
    wb.active["B2"] = "=A2*2"
    out = BytesIO(); wb.save(out)
    svc = WorkspaceFileService(tmp_path)
    created = svc.create("book.xlsx", out.getvalue())
    def builder(before):
        book = load_workbook(BytesIO(before))
        apply_operations_in_memory(book, [{"kind":"insert","sheet":"Data","axis":"row","at":2,"count":1}])
        target = BytesIO(); book.save(target); return target.getvalue()
    receipt = svc.update_with_builder("book.xlsx", builder, expected_version=created.primary_version())
    assert receipt.state == "committed"
    loaded = load_workbook(tmp_path / "book.xlsx", data_only=False)
    assert loaded["Data"]["B3"].value == "=A3*2"
    loaded.close()


def test_style_patch_can_disable_without_erasing_unrelated_fields():
    wb = Workbook(); ws = wb.active
    ws["A1"].font = Font(underline="single", strike=True, scheme="minor", bold=True)
    ws["A1"].alignment = Alignment(wrap_text=True, horizontal="left")
    apply_operations_in_memory(wb, [{"kind":"cells.patch", "sheet":ws.title,
        "cells":[{"cell":"A1","style":{"font":{"underline":None,"strike":False},"alignment":{"horizontal":"center"}}}]}])
    assert ws["A1"].font.underline is None and ws["A1"].font.strike is False
    assert ws["A1"].font.bold and ws["A1"].font.scheme == "minor"
    assert ws["A1"].alignment.wrap_text and ws["A1"].alignment.horizontal == "center"


def test_unknown_operation_or_fields_are_rejected():
    with pytest.raises(ValueError):
        apply_operations_in_memory(Workbook(), [{"kind":"unknown"}])
    with pytest.raises(ValueError):
        apply_operations_in_memory(Workbook(), [{"kind":"cells.patch","celss":[]}])
