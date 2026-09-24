"""V2 acceptance through the registered surface, serialized bytes and shared UI observation."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from excelmanus.tools.context import use_workspace
from excelmanus.tools.registry import ToolRegistry
from excelmanus.tools.workbook_tools import observe_spreadsheet, apply_spreadsheet_changes, preview_spreadsheet
from excelmanus.workbook_commit import content_version_of_file


@pytest.fixture
def workspace(tmp_path):
    with use_workspace(tmp_path):
        yield tmp_path


def book(root: Path, *, grouped=False) -> Path:
    path = root / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["名称", "数量", "金额"])
    ws.append(["项目甲", 2, "=B2*10"])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 16
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 28
    if grouped:
        ws.column_dimensions.group("D", "F")
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["D"].hidden = False
        ws.row_dimensions[3].hidden = True
    wb.save(path)
    wb.close()
    return path


def test_one_transaction_write_style_geometry_and_roundtrip(workspace):
    path = book(workspace)
    before = observe_spreadsheet(file_path=path.name, sheet="数据", mode="range", range="A1:C2", facets=["data", "presentation", "geometry"])
    assert before.success
    r = apply_spreadsheet_changes(file_path=path.name, expected_version=before.value["content_version"], operations=[
        {"kind": "cells.patch", "sheet": "数据", "cells": [{"cell": "A2", "value": "新项目", "style": {"font": {"bold": True}}}]},
        {"kind": "geometry.scale", "sheet": "数据", "range": "A1:C2", "x": 1.3, "y": .7},
    ])
    assert r.success, r.model_text
    assert len(r.value["receipt"]["targets"]) == 1
    geometry = r.value["observation"]["geometry_changes"][0]
    assert geometry["verified"]
    assert geometry["persisted"]["width_px"] > geometry["before"]["width_px"]
    assert geometry["persisted"]["height_px"] < geometry["before"]["height_px"]
    assert r.value["observation"]["cell_checks"][0]["verified"]
    wb = load_workbook(path)
    assert wb["数据"]["A2"].value == "新项目"
    assert wb["数据"]["A2"].font.bold
    assert wb["数据"]["C2"].value == "=B2*10"
    wb.close()


def test_late_range_styles_grouped_dimensions_and_hidden_axes(workspace):
    path = book(workspace, grouped=True)
    wb = load_workbook(path)
    ws = wb["数据"]
    ws["D801"] = "深处的标题"
    ws["D801"].fill = PatternFill("solid", fgColor="112233")
    ws.merge_cells("D801:F801")
    wb.save(path)
    wb.close()
    r = observe_spreadsheet(file_path=path.name, sheet="数据", mode="range", range="E801:F802", facets=["data", "presentation", "geometry"])
    assert r.success, r.model_text
    region = r.value["regions"][0]
    assert [d["native"] for d in region["geometry"]["columns"]] == [24, 24]
    assert region["merge_anchors"]["D801"]["value"] == "深处的标题"
    assert region["merge_anchors"]["D801"]["s"]["bg"]["rgb"] == "#112233"
    hidden = observe_spreadsheet(file_path=path.name, sheet="数据", range="A1:C3", mode="range")
    assert hidden.value["regions"][0]["geometry"]["rows"][-1]["pixels"] == 0


def test_observing_blank_windows_does_not_change_bounds(workspace):
    path = book(workspace)
    before = path.read_bytes()
    r = observe_spreadsheet(file_path=path.name, sheet="数据", range="A400:C420", mode="range", facets=["data"])
    assert r.success
    again = observe_spreadsheet(file_path=path.name)
    assert again.value["sheets"][0]["used"] == {"rows": 2, "cols": 3}
    assert path.read_bytes() == before


def test_failed_batch_is_atomic_and_version_bound(workspace):
    path = book(workspace)
    before = path.read_bytes()
    r = apply_spreadsheet_changes(file_path=path.name, expected_version=content_version_of_file(path), operations=[
        {"kind": "write", "sheet": "数据", "start_cell": "A1", "values": [["bad"]]},
        {"kind": "geometry.scale", "sheet": "数据", "range": "A1:C2", "x": 10000, "y": 1},
    ])
    assert not r.success
    assert before == path.read_bytes()
    stale = apply_spreadsheet_changes(file_path=path.name, expected_version="sha256:stale", operations=[{"kind": "size", "sheet": "数据", "column_widths": {"A": 30}}])
    assert not stale.success and stale.error.code == "VERSION_CONFLICT"
    assert before == path.read_bytes()


def test_registered_surface_has_no_retired_tools(workspace):
    registry = ToolRegistry()
    registry.register_builtin_tools(str(workspace))
    assert {"observe_spreadsheet", "preview_spreadsheet", "apply_spreadsheet_changes"} <= set(registry.get_tool_names())
    assert not {"inspect_spreadsheet", "edit_spreadsheet", "format_spreadsheet", "manage_spreadsheet_objects"} & set(registry.get_tool_names())


def test_visual_document_requires_geometry_and_data_document_does_not(workspace):
    spec = {"version": "2", "purpose": "visual_replica", "sheets": [{"name": "S", "dimensions": {"rows": 2, "cols": 2},
            "value_blocks": [{"start": "A1", "values": [["标题", None], ["数量", 3]]}]}], "uncertainties": []}
    r = apply_spreadsheet_changes(file_path="visual.xlsx", workbook_spec=spec)
    assert not r.success and not (workspace / "visual.xlsx").exists()
    spec["purpose"] = "data"
    r = apply_spreadsheet_changes(file_path="data.xlsx", workbook_spec=spec)
    assert r.success, r.model_text
    wb = load_workbook(workspace / "data.xlsx")
    assert wb["S"]["B2"].value == 3
    assert wb.sheetnames == ["S"]
    wb.close()


def test_scope_labels_are_independent(workspace):
    path = book(workspace)
    r = observe_spreadsheet(file_path=path.name, sheet="数据", range="A1:C2", mode="range", facets=["data"])
    scope = r.value["regions"][0]["coverage"]
    assert scope["data"]["status"] == "complete"
    assert scope["geometry"]["status"] == "not_requested"
    assert scope["objects"]["status"] == "not_requested"
    c = r.value["regions"][0]["cells"]["2,3"]
    assert c["f"] == "=B2*10" and c["cached"] == "no"


def test_new_contract_rejects_old_dimension_names(workspace):
    path = book(workspace)
    r = apply_spreadsheet_changes(file_path=path.name, expected_version=content_version_of_file(path), operations=[{"kind": "size", "sheet": "数据", "columns": {"A": 30}}])
    assert not r.success


def test_registry_execution_and_sdk_share_schema(workspace):
    registry = ToolRegistry()
    registry.register_builtin_tools(str(workspace))
    registry.configure_schema_validation(mode="enforce", canary_percent=100, strict_path=False)
    r = registry.call_tool("apply_spreadsheet_changes", {"file_path": "new.xlsx", "create": True,
        "operations": [{"kind": "write", "sheet": "S", "start_cell": "A1", "values": [[1, 2]]}]})
    assert r.success, r.model_text
    o = registry.call_tool("observe_spreadsheet", {"file_path": "new.xlsx", "sheet": "S", "mode": "range", "range": "A1:B1"})
    assert o.success, o.model_text
    assert o.value["regions"][0]["cells"]["1,2"]["v"] == 2


def test_preview_produces_attached_image_from_same_version(workspace):
    path = book(workspace)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    r = preview_spreadsheet(file_path=path.name, sheet="数据", range="A1:C2", expected_version=content_version_of_file(path))
    assert r.success, r.model_text
    assert r.ui_meta.image["attachment"]["attachmentId"] == r.value["attachment_id"]
    assert r.value["measured"]["content_version"] == r.value["content_version"]
    assert r.value["measured"]["clip"]["width"] > 0
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
    assert list(workspace.glob("*.png")) == []
