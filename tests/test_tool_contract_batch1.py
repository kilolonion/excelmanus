"""工具契约第 1 批：写入保真、copy 快照、结构拒绝、chart 原子、compare 对齐。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.security import FileAccessGuard
from excelmanus.tools._guard_ctx import set_guard
from excelmanus.tools.intent_tools import (
    compare_spreadsheets,
    edit_spreadsheet,
    init_guard,
    manage_spreadsheet_objects,
)
from excelmanus.workbook_commit import content_version_of_file, seed_seen_versions


def _bind(root: Path) -> None:
    workspace = str(root)
    set_guard(FileAccessGuard(workspace))
    init_guard(workspace)
    seed_seen_versions({})


def _err(result: ToolResult) -> str:
    if result.error is not None:
        return f"{result.error.code} {result.error.message}"
    if isinstance(result.value, dict):
        return str(result.value)
    return result.model_text


def test_null_clears_existing_cell(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "id"
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "write", "sheet": "Sheet1", "start_cell": "A1", "values": [[None]]}],
    )
    assert result.success, _err(result)
    wb = load_workbook(path)
    assert wb.active["A1"].value is None
    wb.close()


def test_write_preserves_explicit_strings(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "start_cell": "A1",
            "values": [["00123", "1E10", 42]],
        }],
    )
    assert result.success, _err(result)
    wb = load_workbook(path)
    assert wb.active["A1"].value == "00123"
    assert wb.active["B1"].value == "1E10"
    assert wb.active["C1"].value == 42
    wb.close()


def test_merged_values_collide(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "merged.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.merge_cells("A1:B1")
    ws["A1"] = "keep"
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "write",
            "sheet": "Sheet1",
            "start_cell": "A1",
            "values": [["wanted", "overwrites"]],
        }],
    )
    assert not result.success
    assert "锚点" in _err(result)
    wb = load_workbook(path)
    assert wb.active["A1"].value == "keep"
    wb.close()


def test_overlapping_copy_uses_snapshot(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "copy.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for col, value in enumerate([1, 2, 3, 4], start=1):
        ws.cell(row=1, column=col, value=value)
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "copy",
            "source_sheet": "Sheet1",
            "source_range": "A1:C1",
            "target_sheet": "Sheet1",
            "target_start": "B1",
        }],
    )
    assert result.success, _err(result)
    wb = load_workbook(path)
    assert [wb.active.cell(row=1, column=c).value for c in range(1, 5)] == [1, 1, 2, 3]
    wb.close()


def test_copy_whole_column_rejected(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "cols.xlsx"
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = 1
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "copy",
            "source_sheet": "Sheet1",
            "source_range": "A:A",
            "target_sheet": "Sheet1",
            "target_start": "B1",
        }],
    )
    assert not result.success
    assert "INVALID_ARGS" in _err(result) or "整列" in _err(result)


def test_copy_sheet_conflict_rejected(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "conflict.xlsx"
    wb = Workbook()
    wb.active.title = "First"
    wb.create_sheet("Second")
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{
            "kind": "copy",
            "source_sheet": "First",
            "source_range": "Second!A1:A1",
            "target_sheet": "First",
            "target_start": "B1",
        }],
    )
    assert not result.success
    assert "工作表不一致" in _err(result)


def test_insert_refuses_when_formulas_exist(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "formula.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 10
    ws["B1"] = "=A1"
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "insert", "sheet": "Sheet1", "axis": "row", "at": 1, "count": 1}],
    )
    assert not result.success
    wb = load_workbook(path)
    assert wb.active["B1"].value == "=A1"
    wb.close()


def test_rename_refuses_when_formulas_exist(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "rename.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "First"
    ws["A1"] = "=1+1"
    wb.save(path)
    wb.close()
    result = edit_spreadsheet(
        file_path=str(path),
        expected_version=content_version_of_file(path),
        operations=[{"kind": "sheet", "action": "rename", "sheet": "First", "new_name": "Renamed"}],
    )
    assert not result.success
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "First"
    wb.close()


def test_object_batch_failure_does_not_commit_first_chart(tmp_path: Path) -> None:
    _bind(tmp_path)
    path = tmp_path / "charts.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "n"
    ws["B1"] = "v"
    ws["A2"] = "a"
    ws["B2"] = 1
    wb.save(path)
    wb.close()
    before = content_version_of_file(path)
    result = manage_spreadsheet_objects(
        file_path=str(path),
        expected_version=before,
        operations=[
            {
                "kind": "chart",
                "chart_type": "bar",
                "data_range": "B1:B2",
                "categories_range": "A2:A2",
                "sheet": "Sheet1",
                "target_cell": "D1",
            },
            {
                "kind": "chart",
                "chart_type": "not-a-chart",
                "data_range": "B1:B2",
                "sheet": "Sheet1",
                "target_cell": "D15",
            },
        ],
    )
    assert not result.success
    assert content_version_of_file(path) == before
    wb = load_workbook(path)
    assert len(wb.active._charts) == 0
    wb.close()


def test_alignment_position_rejects_key_columns(tmp_path: Path) -> None:
    _bind(tmp_path)
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    for path in (a, b):
        wb = Workbook()
        wb.active.append(["id", "value"])
        wb.active.append([1, 2])
        wb.save(path)
        wb.close()
    result = compare_spreadsheets(
        file_a=str(a),
        file_b=str(b),
        alignment="position",
        key_columns=["id"],
    )
    assert not result.success


def test_ignore_style_false_rejected(tmp_path: Path) -> None:
    _bind(tmp_path)
    a = tmp_path / "a.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(a)
    wb.close()
    result = compare_spreadsheets(file_a=str(a), file_b=str(a), ignore_style=False)
    assert not result.success
    assert "样式" in _err(result)


def test_write_patterns_no_longer_claims_short_write_deletes() -> None:
    text = (
        Path(__file__).resolve().parents[1]
        / "excelmanus/skillpacks/system/run_code_templates/references/write_patterns.md"
    ).read_text(encoding="utf-8")
    assert "覆盖写不等于删除" in text
    assert "null" in text
