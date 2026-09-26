"""空串写入契约回归：写 "" 必须等同 Excel 的“清空单元格”，且回执与落盘一致。

背景：openpyxl 落盘写不出字面空串——``cell.value = ""`` 在内存里是 ``t="s"``，
写出的 XML 只有空文本节点（``<c r="B2" t="inlineStr"></c>``），重新读取只能得到
``None``/``inlineStr``。内存里留着 "" 会让提交前的序列化校验读到并不存在的差异，
报成 ``SERIALIZATION_MISMATCH``（failure_class=internal，用户看到“内部错误”）。
这里锁定“空串 = 清空该格”语义：写入前归一，格子既不残留值，也不残留幽灵空文本节点，
回执里该格 ``value`` 为 null 且与文件回读一致。
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excelmanus.tools.context import bind_workspace, reset_call, use_workspace
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes
from excelmanus.workbook.data import filter_data
from excelmanus.workbook_commit import content_version_of_file


def _seed(path: Path, rows: list[list], *, sheet: str = "S") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def _sheet_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("xl/worksheets/sheet1.xml").decode("utf-8")


def _apply(tmp_path: Path, operations: list[dict], *, file_name: str = "b.xlsx"):
    with use_workspace(tmp_path):
        version = content_version_of_file(tmp_path / file_name)
        return apply_spreadsheet_changes(
            file_path=file_name, expected_version=version, operations=operations
        )


def _cell_checks(result) -> dict[str, object]:
    return {
        entry["cell"]: entry["value"]
        for entry in result.value["observation"]["cell_checks"]
    }


def test_write_empty_string_clears_cell_and_keeps_neighbors(tmp_path: Path) -> None:
    """write 写 ""：该格被清空、回读为 null，相邻单元格不受影响。"""
    _seed(tmp_path / "b.xlsx", [["h", None, None], [1, "old", "keep"]])
    result = _apply(
        tmp_path,
        [{"kind": "write", "sheet": "S", "start_cell": "B2", "values": [[""]]}],
    )
    assert result.success, result.error and result.error.message

    # 回执：该格如实报 null，且与落盘版本一致。
    assert _cell_checks(result) == {"B2": None}
    assert result.value["content_version"] == content_version_of_file(tmp_path / "b.xlsx")

    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        ws = wb["S"]
        assert ws["B2"].value is None
        assert ws["B2"].data_type == "n"  # 不是 inlineStr 幽灵空文本
        assert ws["A1"].value == "h" and ws["A2"].value == 1
        assert ws["C2"].value == "keep"
    finally:
        wb.close()
    # 落盘 XML 里不留空文本节点。
    assert '<c r="B2"' not in _sheet_xml(tmp_path / "b.xlsx")


def test_write_empty_string_into_blank_cell_stays_absent(tmp_path: Path) -> None:
    """write 往从未有值的格子写 ""：不产生空 inlineStr 节点。"""
    _seed(tmp_path / "b.xlsx", [["h"]])
    result = _apply(
        tmp_path,
        [{"kind": "write", "sheet": "S", "start_cell": "B2", "values": [[""]]}],
    )
    assert result.success, result.error and result.error.message
    assert '<c r="B2"' not in _sheet_xml(tmp_path / "b.xlsx")
    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        assert wb["S"]["B2"].value is None
    finally:
        wb.close()


def test_write_matrix_mixes_empty_and_filled_cells(tmp_path: Path) -> None:
    """同一矩形里 "" 只清空自己的格，同排/邻居的真实值照常落盘。"""
    _seed(tmp_path / "b.xlsx", [["h", None, None], [1, "old", "old2"]])
    result = _apply(
        tmp_path,
        [
            {
                "kind": "write",
                "sheet": "S",
                "start_cell": "B2",
                "values": [["", "新"]],
            }
        ],
    )
    assert result.success, result.error and result.error.message
    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        ws = wb["S"]
        assert ws["B2"].value is None
        assert ws["C2"].value == "新"
        assert ws["A2"].value == 1
    finally:
        wb.close()


def test_write_empty_string_into_merged_region_is_not_a_collision(
    tmp_path: Path,
) -> None:
    """合并区内多个 "" 都是“清空”，不该被当成多个非空值抢占锚点。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "title"
    ws.merge_cells("A1:B1")
    wb.save(tmp_path / "b.xlsx")
    wb.close()
    result = _apply(
        tmp_path,
        [{"kind": "write", "sheet": "S", "start_cell": "A1", "values": [["", ""]]}],
    )
    assert result.success, result.error and result.error.message
    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        assert wb["S"]["A1"].value is None
    finally:
        wb.close()


def test_cells_patch_empty_string_clears_cell(tmp_path: Path) -> None:
    """cells.patch 写 ""：同样按清空语义成功提交，回执与文件一致。"""
    _seed(tmp_path / "b.xlsx", [[None, None], ["x", "old"]], sheet="Sheet")
    result = _apply(
        tmp_path,
        [
            {
                "kind": "cells.patch",
                "sheet": "Sheet",
                "cells": [{"cell": "B2", "value": ""}],
            }
        ],
    )
    assert result.success, result.error and result.error.message
    assert _cell_checks(result) == {"B2": None}
    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        assert wb["Sheet"]["B2"].value is None
        assert wb["Sheet"]["A2"].value == "x"
    finally:
        wb.close()
    assert '<c r="B2"' not in _sheet_xml(tmp_path / "b.xlsx")


def test_selection_write_empty_string_clears_selected_cells(tmp_path: Path) -> None:
    """selection 写入（filter 结果）里的 "" 同样按清空处理，行内其他格照常。"""
    _seed(tmp_path / "source.xlsx", [["Name", "Amount"], ["Bob", 20], ["Ann", 30]], sheet="Data")
    token = bind_workspace(tmp_path)
    try:
        filtered = filter_data(
            "source.xlsx", sheet_name="Data", header_row=1, column="Name", operator="eq", value="Bob"
        )
        assert filtered.success, filtered.error and filtered.error.message
        selection = filtered.value["selection"]
        result = apply_spreadsheet_changes(
            file_path="source.xlsx",
            expected_version=selection["content_version"],
            operations=[{"kind": "write", "selection": selection, "values": [["", None]]}],
        )
    finally:
        reset_call(token)
    assert result.success, result.error and result.error.message
    wb = load_workbook(tmp_path / "source.xlsx")
    try:
        ws = wb["Data"]
        assert ws["A2"].value is None
        assert ws["B2"].value is None
        assert ws["A3"].value == "Ann" and ws["B3"].value == 30
    finally:
        wb.close()


def test_whitespace_only_string_is_a_real_value(tmp_path: Path) -> None:
    """纯空白串 " " 是 Excel 里的真实值，能正常往返，不得被当成清空。"""
    _seed(tmp_path / "b.xlsx", [["h"]])
    result = _apply(
        tmp_path,
        [{"kind": "write", "sheet": "S", "start_cell": "B2", "values": [[" "]]}],
    )
    assert result.success, result.error and result.error.message
    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        assert wb["S"]["B2"].value == " "
    finally:
        wb.close()


def test_replace_with_empty_replacement_clears_cell(tmp_path: Path) -> None:
    """replace 的 replacement="" 是同一根因：清空命中格，不报内部错误。"""
    _seed(tmp_path / "b.xlsx", [["hello"], ["hello"]])
    result = _apply(
        tmp_path,
        [
            {
                "kind": "replace",
                "sheet": "S",
                "range": "A1:A2",
                "find": "hello",
                "replacement": "",
            }
        ],
    )
    assert result.success, result.error and result.error.message
    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        assert wb["S"]["A1"].value is None
        assert wb["S"]["A2"].value is None
    finally:
        wb.close()
    xml = _sheet_xml(tmp_path / "b.xlsx")
    assert '<c r="A1"' not in xml and '<c r="A2"' not in xml


def test_fill_empty_value_clears_range(tmp_path: Path) -> None:
    """fill mode=value 且 value="" 也按清空处理，不报内部错误。"""
    _seed(tmp_path / "b.xlsx", [["a", "b", "x"], ["c", "d", "y"]])
    result = _apply(
        tmp_path,
        [
            {
                "kind": "fill",
                "sheet": "S",
                "range": "C1:C2",
                "mode": "value",
                "value": "",
            }
        ],
    )
    assert result.success, result.error and result.error.message
    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        assert wb["S"]["C1"].value is None
        assert wb["S"]["C2"].value is None
    finally:
        wb.close()


def test_transform_split_empty_part_leaves_no_phantom_cell(tmp_path: Path) -> None:
    """split 派生出的空字段按清空落盘，不留空 inlineStr 幽灵格。"""
    _seed(tmp_path / "b.xlsx", [["raw"], ["a,,b"]])
    result = _apply(
        tmp_path,
        [
            {
                "kind": "transform",
                "sheet": "S",
                "action": "split",
                "column": "raw",
                "delimiter": ",",
                "new_columns": ["p1", "p2", "p3"],
            }
        ],
    )
    assert result.success, result.error and result.error.message
    wb = load_workbook(tmp_path / "b.xlsx")
    try:
        ws = wb["S"]
        assert ws["A2"].value == "a"
        assert ws["B2"].value is None
        assert ws["C2"].value == "b"
    finally:
        wb.close()
    assert 't="inlineStr"></c>' not in _sheet_xml(tmp_path / "b.xlsx")
