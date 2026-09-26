"""CSV imports carry a LibreOffice CalcA1 calc marker that must stay editable.

The only route into a CSV-only workspace is ``convert_spreadsheet``: the engine
materializes the first xlsx, so the next round advertises an xlsx profile and
``apply_spreadsheet_changes`` unlocks.  LibreOffice writes
``<loext:extCalcPr stringRefSyntax="CalcA1"/>`` when it imports CSV/ODS, and
``stringRefSyntax="ExcelA1"`` when it imports xlsx.  Both values describe A1
references; they only differ in the string form used by text-to-reference
conversion (``Sheet1!A1`` vs ``Sheet1.A1``).  R1C1 variants stay rejected.
"""

import re
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.tools.context import use_workspace
from excelmanus.tools.spreadsheet_engine_tools import calculate_spreadsheet, convert_spreadsheet
from excelmanus.tools.workbook_tools import apply_spreadsheet_changes
from excelmanus.workbook.adapter import (
    _preservable_workbook_extensions,
    ensure_editable,
    package_inventory,
    preserve_workbook_extensions,
)
from excelmanus.workbook_commit import CommitError, content_version_of_file

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_URI = "{7626C862-2A13-11E5-B345-FEFF819CDC9F}"


def _extension(syntax: str, *, extra_attribute: str = "") -> str:
    # LibreOffice declares loext on <ext>, not on <extLst> or the root.
    return (
        f'<extLst><ext xmlns:loext="http://schemas.libreoffice.org/" uri="{_URI}">'
        f'<loext:extCalcPr stringRefSyntax="{syntax}"{extra_attribute}/></ext></extLst>'
    )


def _rewrite_part(data: bytes, part: str, transform) -> bytes:
    out = BytesIO()
    with ZipFile(BytesIO(data)) as src, ZipFile(out, "w") as dst:
        dst.comment = src.comment
        for item in src.infolist():
            raw = src.read(item.filename)
            dst.writestr(item, transform(raw) if item.filename == part else raw)
    return out.getvalue()


def _marked(path, syntax: str | None, *, ref_mode: str | None = "A1") -> bytes:
    """Write an openpyxl workbook whose workbook.xml looks like an engine import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "input"
    ws.append(["item", "qty", "price", "total"])
    ws.append(["widget", 2, 3.5, "=B2*C2"])
    ws.append(["gadget", 4, 1.25, "=B3*C3"])
    wb.save(path)
    wb.close()
    data = path.read_bytes()

    def patch(raw: bytes) -> bytes:
        if ref_mode is not None:
            # Replace the reference mode on the existing calcPr: a workbook may
            # declare at most one, and openpyxl already writes one.
            def set_ref_mode(match):
                tag = match.group(0)
                if b"refMode=" in tag:
                    return re.sub(rb'refMode="[^"]*"', f'refMode="{ref_mode}"'.encode(), tag)
                return tag[:-2] + f' refMode="{ref_mode}"/>'.encode()

            raw = re.sub(rb"<calcPr\b[^>]*/>", set_ref_mode, raw, count=1)
        if syntax is not None:
            raw = raw.replace(b"</workbook>", _extension(syntax).encode() + b"</workbook>")
        return raw

    data = _rewrite_part(data, "xl/workbook.xml", patch)
    path.write_bytes(data)
    return data


def _workbook_extension(data: bytes) -> ET.Element | None:
    with ZipFile(BytesIO(data)) as package:
        root = ET.fromstring(package.read("xl/workbook.xml"))
    lists = root.findall(f"{{{_MAIN}}}extLst")
    assert len(lists) <= 1
    return lists[0] if lists else None


def _syntax_of(extension: ET.Element | None) -> str | None:
    if extension is None:
        return None
    return extension[0][0].get("stringRefSyntax")


@pytest.fixture
def workspace(tmp_path):
    with use_workspace(tmp_path):
        yield tmp_path


def test_calca1_csv_import_marker_is_admitted_and_preserved(tmp_path):
    data = _marked(tmp_path / "csv_import.xlsx", "CalcA1")
    inventory = package_inventory(data)
    assert inventory["unsupported_features"] == []
    assert inventory["edit_support"] == "supported_subset"
    assert inventory["preserved_features"] == [
        {
            "part": "xl/workbook.xml",
            "feature": "extCalcPr",
            "uri": _URI,
            "string_ref_syntax": "CalcA1",
        }
    ]
    # The residual risk is stated instead of hidden: only string-encoded
    # references follow the Calc grammar recorded by the marker.
    assert any("CalcA1" in line and "INDIRECT" in line for line in inventory["limitations"])
    assert ensure_editable(data)["edit_support"] == "supported_subset"
    assert _syntax_of(_workbook_extension(data)) == "CalcA1"


def test_excel_a1_marker_stays_admitted_without_calc_caveat(tmp_path):
    data = _marked(tmp_path / "xlsx_import.xlsx", "ExcelA1")
    inventory = package_inventory(data)
    assert inventory["unsupported_features"] == []
    assert inventory["preserved_features"][0]["string_ref_syntax"] == "ExcelA1"
    assert not any("CalcA1" in line for line in inventory["limitations"])


@pytest.mark.parametrize(
    "syntax",
    ["CalcR1C1", "ExcelR1C1", "UnknownSyntax", "calca1", ""],
)
def test_non_a1_or_unknown_syntax_stays_blocked(tmp_path, syntax):
    data = _marked(tmp_path / "blocked.xlsx", syntax)
    inventory = package_inventory(data)
    assert inventory["unsupported_features"] == [
        {"part": "xl/workbook.xml", "feature": "extLst"}
    ]
    assert inventory["edit_support"] == "unsupported"
    with pytest.raises(CommitError) as error:
        ensure_editable(data)
    assert error.value.code == "UNSUPPORTED_PRESERVATION"
    assert _preservable_workbook_extensions(
        ET.fromstring(ZipFile(BytesIO(data)).read("xl/workbook.xml"))
    ) is None


def test_extra_attribute_or_foreign_uri_stays_blocked(tmp_path):
    inflated = _marked(tmp_path / "inflated.xlsx", "CalcA1", ref_mode=None)
    inflated = _rewrite_part(inflated, "xl/workbook.xml", lambda raw: raw.replace(
        _extension("CalcA1").encode(),
        _extension("CalcA1", extra_attribute=' unknown="data"').encode(),
    ))
    assert package_inventory(inflated)["edit_support"] == "unsupported"
    foreign = _marked(tmp_path / "foreign.xlsx", "CalcA1", ref_mode=None)
    foreign = _rewrite_part(foreign, "xl/workbook.xml", lambda raw: raw.replace(
        _URI.encode(), b"{UNKNOWN}",
    ))
    assert package_inventory(foreign)["edit_support"] == "unsupported"


def test_r1c1_workbook_is_not_admitted_even_with_known_marker(tmp_path):
    data = _marked(tmp_path / "r1c1.xlsx", "CalcA1", ref_mode="R1C1")
    inventory = package_inventory(data)
    assert inventory["unsupported_features"] == [
        {"part": "xl/workbook.xml", "feature": "extLst"}
    ]
    with pytest.raises(CommitError):
        ensure_editable(data)


def test_preserve_workbook_extensions_restores_calca1_after_openpyxl_save(tmp_path):
    marked = _marked(tmp_path / "marked.xlsx", "CalcA1")
    stripped = _rewrite_part(marked, "xl/workbook.xml", lambda raw: raw.replace(
        _extension("CalcA1").encode(), b"",
    ))
    assert _workbook_extension(stripped) is None
    restored = preserve_workbook_extensions(marked, stripped)
    assert _syntax_of(_workbook_extension(restored)) == "CalcA1"
    # Only the extension is restored; the serialized package is otherwise kept.
    with ZipFile(BytesIO(stripped)) as before, ZipFile(BytesIO(restored)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "xl/workbook.xml":
                assert before.read(name) == after.read(name)


def _edit(relative_path, path, operations):
    return apply_spreadsheet_changes(
        file_path=relative_path,
        expected_version=content_version_of_file(path),
        operations=operations,
    )


def test_edit_keeps_calca1_marker_formulas_and_values(workspace, tmp_path):
    book = tmp_path / "converted.xlsx"
    _marked(book, "CalcA1")
    before = _workbook_extension(book.read_bytes())
    result = _edit(book.name, book, [
        {"kind": "write", "sheet": "input", "start_cell": "A2", "values": [["widget pro"]]},
        {"kind": "write", "sheet": "input", "start_cell": "B2", "values": [[3]]},
        {"kind": "write", "sheet": "input", "start_cell": "D3", "values": [["=B3*C3*2"]]},
    ])
    assert result.success, result.model_text
    assert result.value["committed"]
    after = _workbook_extension(book.read_bytes())
    assert after is not None and _syntax_of(after) == "CalcA1"
    assert ET.tostring(after) == ET.tostring(before)
    wb = load_workbook(book)
    try:
        ws = wb["input"]
        assert [ws.cell(row, column).value for row in (2, 3) for column in range(1, 5)] == [
            "widget pro", 3, 3.5, "=B2*C2",
            "gadget", 4, 1.25, "=B3*C3*2",
        ]
    finally:
        wb.close()
    # A second edit must not duplicate the extension list.
    again = _edit(book.name, book, [{"kind": "write", "sheet": "input", "start_cell": "C2", "values": [[4]]}])
    assert again.success, again.model_text
    assert ET.tostring(_workbook_extension(book.read_bytes())) == ET.tostring(before)


def _office_engine():
    from excelmanus.runtime_capabilities import office_executable

    return office_executable()


@pytest.mark.skipif(not _office_engine(), reason="no local LibreOffice engine")
def test_csv_convert_then_edit_round_trip_with_local_engine(workspace, tmp_path):
    """csv -> convert_spreadsheet -> next round apply_spreadsheet_changes."""
    source = tmp_path / "uploads"
    source.mkdir()
    (source / "sales.csv").write_text(
        "item,qty,price,total\nwidget,2,3.5,=B2*C2\ngadget,4,1.25,=B3*C3\n",
        encoding="utf-8",
    )
    converted = convert_spreadsheet("uploads/sales.csv", "outputs/sales.xlsx")
    assert converted.success, converted.model_text
    target = tmp_path / "outputs" / "sales.xlsx"
    inventory = package_inventory(target.read_bytes())
    # The engine import is the only way into a CSV-only workspace; it must be
    # an editable base, otherwise the workspace can never produce a workbook.
    assert inventory["edit_support"] == "supported_subset", inventory
    assert inventory["preserved_features"][0]["string_ref_syntax"] == "CalcA1"
    marker = _workbook_extension(target.read_bytes())
    assert _syntax_of(marker) == "CalcA1"

    # In-place edits: existing cells, including the imported formula cell.
    result = _edit("outputs/sales.xlsx", target, [
        {"kind": "write", "sheet": "input", "start_cell": "A2", "values": [["widget pro"]]},
        {"kind": "write", "sheet": "input", "start_cell": "B2", "values": [[3]]},
        {"kind": "write", "sheet": "input", "start_cell": "D3", "values": [["=B3*C3*2"]]},
    ])
    assert result.success, result.model_text
    assert result.value["committed"]
    assert ET.tostring(_workbook_extension(target.read_bytes())) == ET.tostring(marker)

    wb = load_workbook(target)
    try:
        ws = wb["input"]
        assert [[ws.cell(row, column).value for column in range(1, 5)] for row in range(1, 4)] == [
            ["item", "qty", "price", "total"],
            ["widget pro", 3, 3.5, "=B2*C2"],
            ["gadget", 4, 1.25, "=B3*C3*2"],
        ]
    finally:
        wb.close()

    recalculated = calculate_spreadsheet(
        "outputs/sales.xlsx", expected_version=content_version_of_file(target)
    )
    assert recalculated.success, recalculated.model_text
    assert recalculated.value["formula_recalculation"]["status"] == "recalculated"
    assert recalculated.value["formula_recalculation"]["error_count"] == 0
    wb = load_workbook(target, data_only=True)
    try:
        ws = wb["input"]
        # =B2*C2 with B2=3 -> 10.5, =B3*C3*2 with 4 and 1.25 -> 10.
        assert [ws["D2"].value, ws["D3"].value] == [10.5, 10]
    finally:
        wb.close()
    assert _syntax_of(_workbook_extension(target.read_bytes())) == "CalcA1"
