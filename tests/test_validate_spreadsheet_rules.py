"""validate_spreadsheet 规则语义回归测试（agent 会话真实失败逐条修复）。

覆盖七类修复：
1. 列引用支持列字母/1-based 列号（column/columns/reference.columns 一致），解析失败报可用表头名；
2. total 规则跳过空值计入 skipped，非数值文本不再中止整个调用；
3. row_expression 支持 {row} 占位符与大写函数名，语法错误返回结构化 INVALID_ARGS（中文说明+示例），不泄漏 SyntaxError；
4. input_schema 描述与实际能力对齐（列引用写法、示例表达式、各 kind 所需字段）；
5. A1 单元格引用（abs(B7-1)<0.001）必须报 INVALID_ARGS 并说明本表达式语言不支持 A1 引用，给出替代写法
   （列字母作用于当前行 / cell 规则 / observe_spreadsheet 回读 / 常量写进数据行）；
6. 失败负载机器可读：total 的非数值文本项 actual/expected 只放数值或 null，说明进 reason/message/raw_value，
   并返回 excluded_non_numeric 计数与被排除单元格列表。
7. 新增 cell 规则（真实会话 #104/#105 的能力缺口）：sheet + cell(A1 单格) + expected(+tolerance/case_sensitive)
   断言具体单元格；失败项 {rule,kind,sheet,cell,expected,actual,delta?} 只放机器可读值；
   非法/越界坐标 INVALID_ARGS、多表未指定 sheet 报 SHEET_REQUIRED、公式缓存缺失按既有口径 partial。
"""

import zipfile
import json
from datetime import date

import pytest
from openpyxl import Workbook

from excelmanus.tools.context import use_workspace
from excelmanus.tools.spreadsheet_data_tools import (
    ExpressionSyntaxError,
    evaluate_expression,
    get_tools,
    validate_spreadsheet,
)


@pytest.fixture(autouse=True)
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("EXCELMANUS_FORMULA_RECALC", "off")
    with use_workspace(tmp_path):
        yield


def book(path, rows, name="Data"):
    wb = Workbook()
    ws = wb.active
    ws.title = name
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()
    return path


# ── Bug 1：列字母 / 列号列引用 ──────────────────────────────────────────────


def test_column_letters_and_numbers_resolve(tmp_path):
    path = book(tmp_path / "book.xlsx", [["id", "qty", "amount"], [1, 2, 10], [2, 3, 15]])
    by_letter = validate_spreadsheet(path.name, [{"kind": "total", "column": "C", "expected": 25}])
    assert by_letter.success, by_letter.model_text
    assert by_letter.value["valid"] is True
    multi = validate_spreadsheet(path.name, [{"kind": "required", "columns": ["A", "C"]}])
    assert multi.success and multi.value["valid"] is True
    by_number = validate_spreadsheet(path.name, [{"kind": "total", "column": 3, "expected": 25}])
    assert by_number.success and by_number.value["valid"] is True
    by_digit_string = validate_spreadsheet(path.name, [{"kind": "total", "column": "3", "expected": 25}])
    assert by_digit_string.success and by_digit_string.value["valid"] is True
    wrong_total = validate_spreadsheet(path.name, [{"kind": "total", "column": "C", "expected": 99}])
    assert wrong_total.success and wrong_total.value["valid"] is False
    assert wrong_total.value["failures"][0]["cell"] == "C2"


def test_date_values_in_failures_are_json_safe(tmp_path):
    path = book(tmp_path / "dates.xlsx", [["when", "amount"], [date(2024, 1, 2), 3]])
    result = validate_spreadsheet(
        path.name,
        [{"kind": "row_expression", "expression": "amount > 10"}],
    )
    assert result.success and result.value["valid"] is False
    # The failed row contains an openpyxl date internally; the public payload
    # must still be directly serializable by API clients.
    json.dumps(result.value, ensure_ascii=False)
    assert result.value["failures"][0]["actual"][0].startswith("2024-01-02")


def test_header_name_wins_over_column_letter(tmp_path):
    # 第 2 列的表头名叫 "A"：按规则表头名优先，"A" 指第 2 列而不是 A 列。
    path = book(tmp_path / "book.xlsx", [["x", "A"], [1, 10], [2, 20]])
    check = validate_spreadsheet(path.name, [{"kind": "total", "column": "A", "expected": 30}])
    assert check.success, check.model_text
    assert check.value["valid"] is True


def test_unresolvable_column_reference_reports_rule_sheet_and_headers(tmp_path):
    path = book(tmp_path / "book.xlsx", [["id", "qty", "amount"], [1, 2, 10]])
    result = validate_spreadsheet(path.name, [{"kind": "total", "column": "zzz", "expected": 0}])
    assert not result.success
    payload = result.value
    assert payload["error_code"] == "INVALID_ARGS"
    assert payload["failure_class"] == "invalid_args"
    message = payload["message"]
    assert "rules[0]" in message
    assert "total" in message
    assert "Data" in message
    assert "'zzz'" in message
    assert "amount" in message and "qty" in message  # 列出可用表头名
    assert "列字母" in message and "1-based" in message  # 说明解析规则
    assert "invalid literal" not in message  # 不泄漏 Python int() 异常
    assert payload["rule_index"] == 0 and payload["available_columns"][:3] == ["id", "qty", "amount"]

    out_of_range = validate_spreadsheet(path.name, [{"kind": "unique", "columns": ["Z"]}])
    assert not out_of_range.success
    assert "rules[0]" in out_of_range.value["message"] and "无法解析列引用" in out_of_range.value["message"]


def test_foreign_key_reference_columns_accept_column_letters(tmp_path):
    main = book(tmp_path / "main.xlsx", [["id", "code"], [1, "E"], [2, "S"]])
    ref = book(tmp_path / "ref.xlsx", [["region", "code"], ["华东", "E"], ["华南", "S"]], name="Ref")
    rule = {"kind": "foreign_key", "columns": ["B"], "reference": {"file_path": ref.name, "columns": ["B"]}}
    ok = validate_spreadsheet(main.name, [rule])
    assert ok.success, ok.model_text
    assert ok.value["valid"] is True

    bad = book(tmp_path / "bad.xlsx", [["id", "code"], [1, "E"], [2, "X"]])
    broken = validate_spreadsheet(bad.name, [{"kind": "foreign_key", "columns": ["code"], "reference": {"file_path": ref.name, "columns": ["B"]}}])
    assert broken.success, broken.model_text
    assert broken.value["valid"] is False
    assert broken.value["failure_count"] == 1
    assert broken.value["failures"][0]["cell"] == "B3"


def test_mixed_column_reference_forms(tmp_path):
    path = book(tmp_path / "book.xlsx", [["id", "qty", "amount"], [1, 2, 10]])
    result = validate_spreadsheet(path.name, [{"kind": "required", "columns": ["id", "C", 3]}])
    assert result.success, result.model_text
    assert result.value["valid"] is True


# ── Bug 2：total 规则对空值/非数值文本的处理 ─────────────────────────────────


def test_total_skips_blanks_counts_them_and_reports_text_cells(tmp_path):
    path = book(tmp_path / "book.xlsx", [
        ["名称", "金额"],
        ["a", 10],
        ["b", 20],
        ["合计", "合计"],
        ["d", None],
        ["e", "  "],
    ])
    result = validate_spreadsheet(path.name, [{"kind": "total", "column": "金额", "expected": 30}])
    # 数据内容问题不中止调用、不作为参数错误返回
    assert result.success, result.model_text
    assert result.value["status"] == "success"
    assert "failure_class" not in result.value
    assert result.value["valid"] is False
    assert result.value["failure_count"] == 1
    failure = result.value["failures"][0]
    assert failure["rule"] == 0
    assert failure["kind"] == "total"
    assert failure["sheet"] == "Data"
    assert failure["cell"] == "B4"
    # 契约修复：actual/expected 只放机器可读值，原文与说明移到 raw_value/message/reason
    assert failure["actual"] is None
    assert failure["expected"] is None
    assert failure["reason"] == "non_numeric_text"
    assert failure["raw_value"] == "合计"
    assert "非数值文本" in failure["message"]
    rule_result = result.value["rules"][0]
    assert rule_result["skipped"] == 2
    assert rule_result["skipped_cells"] == ["B5", "B6"]
    assert rule_result["checked"] == 3


def test_total_mismatch_and_text_both_reported_with_context(tmp_path):
    path = book(tmp_path / "book.xlsx", [["名称", "金额"], ["a", 10], ["b", "小计"], ["c", 5]])
    result = validate_spreadsheet(path.name, [{"kind": "total", "column": "B", "expected": 99}])
    assert result.success
    assert result.value["valid"] is False
    cells = {f["cell"] for f in result.value["failures"]}
    assert cells == {"B3", "B2"}  # 非数值单元格 + 合计不符
    assert all(f["rule"] == 0 and f["sheet"] == "Data" for f in result.value["failures"])
    assert all(f["reason"] in {"non_numeric_text", "total_mismatch"} for f in result.value["failures"])


def test_total_text_cells_are_machine_readable_and_counted(tmp_path):
    """回归（真实会话消息 #105）：失败项不得把中文说明串塞进 actual/expected，并须给出排除计数。"""
    path = book(tmp_path / "book.xlsx", [
        ["名称", "金额"],
        ["广告投入", 6.7071],
        ["说明", "销售额 = 1.4338 + 6.7071 × 广告投入"],
        ["其他", 10],
    ])
    result = validate_spreadsheet(path.name, [{"kind": "total", "column": "B", "expected": 6.7071}])
    assert result.success, result.model_text
    assert result.value["valid"] is False
    by_cell = {f["cell"]: f for f in result.value["failures"]}
    assert set(by_cell) == {"B3", "B2"}

    text_failure = by_cell["B3"]
    # actual/expected 必须是机器可读值：非数值文本 → null（原因码与原文放到独立字段）
    assert text_failure["actual"] is None and text_failure["expected"] is None
    assert text_failure["reason"] == "non_numeric_text"
    assert text_failure["raw_value"] == "销售额 = 1.4338 + 6.7071 × 广告投入"
    assert "非数值文本" in text_failure["message"]
    assert "数值（非数值文本不计入合计）" not in str(text_failure.get("expected"))

    # 合计不符项仍是数值，并带可执行的语义修正建议（total=整列求和，单点断言换工具）
    mismatch = by_cell["B2"]
    assert mismatch["actual"] == 16.7071 and mismatch["expected"] == 6.7071
    assert mismatch["reason"] == "total_mismatch"
    assert mismatch["data_range"] == "B2:B4"
    assert "整列" in mismatch["message"] and "row_expression" in mismatch["message"]

    rule_result = result.value["rules"][0]
    assert rule_result["semantics"] == "column_total"
    assert rule_result["total"] == 16.7071
    assert rule_result["excluded_non_numeric"] == 1
    assert rule_result["excluded_non_numeric_cells"] == ["B3"]
    # 任何失败项的 actual/expected 都不再承载自然语言说明
    for failure in result.value["failures"]:
        for value in (failure["actual"], failure["expected"]):
            assert value is None or isinstance(value, (int, float, list))


def test_total_expected_equal_to_cell_value_is_marked_as_cell_assertion(tmp_path):
    """P2：纯数值列里 expected 恰等于某个单元格值时，回执须给出机器可读的单点断言误用信号。"""
    path = book(tmp_path / "book.xlsx", [["名称", "金额"], ["a", 6.7071], ["b", 10], ["c", 5]])
    result = validate_spreadsheet(path.name, [
        {"kind": "total", "column": "B", "expected": 6.7071, "tolerance": 0.0001},
    ])
    assert result.success, result.model_text
    assert result.value["valid"] is False
    assert result.value["failure_count"] == 1
    failure = result.value["failures"][0]
    assert failure["reason"] == "total_mismatch"
    assert failure["expected_matches_cell"] == ["B2"]  # expected 等于 B2 的值 → 可能是单点断言误用
    assert "单元格" in failure["message"] and "row_expression" in failure["message"]
    assert result.value["rules"][0]["expected_matches_cell"] == ["B2"]

    # 真合计错误（expected 不等于任何单元格值）不得带该标记：仍是产物缺陷
    genuine = validate_spreadsheet(path.name, [{"kind": "total", "column": "B", "expected": 999}])
    assert genuine.value["valid"] is False
    assert "expected_matches_cell" not in genuine.value["failures"][0]
    assert "expected_matches_cell" not in genuine.value["rules"][0]


# ── Bug 3：row_expression 的 {row} 占位符、大写函数、结构化语法错误 ───────────


def test_row_expression_accepts_excel_style_placeholder_and_uppercase_functions(tmp_path):
    path = book(tmp_path / "book.xlsx", [
        ["id", "note", "total", "x", "part1", "part2", "part3"],
        [1, "n", 10, 0, 3, 3, 4],
        [2, "n", 12, 0, 5, 5, 1],
    ])
    expression = "ABS(E{row}+F{row}+G{row}-C{row})<0.05"
    result = validate_spreadsheet(path.name, [{"kind": "row_expression", "expression": expression}])
    assert result.success, result.model_text
    assert result.value["valid"] is False
    assert result.value["failure_count"] == 1
    assert result.value["failures"][0]["cell"] == "A3"

    assert evaluate_expression(expression, {"E": 3, "F": 3, "G": 4, "C": 10}) is True
    assert evaluate_expression("ABS(E{row}+F{row}+G{row}-C{row})<0.05", {"E": 5, "F": 5, "G": 1, "C": 12}) is False
    # 大小写不敏感的函数名 + 中文列名示例
    assert evaluate_expression("Round(金额-数量*单价, 2) == 0", {"金额": 20, "数量": 2, "单价": 10}) is True
    assert evaluate_expression("MIN(A{row},B{row})>0 and MAX(A,B)<10", {"A": 1, "B": 2}) is True


def test_row_expression_syntax_error_is_structured_invalid_args(tmp_path):
    path = book(tmp_path / "book.xlsx", [["id", "amount"], [1, 10]])
    result = validate_spreadsheet(path.name, [{"kind": "row_expression", "expression": "ABS(amount-{row}"}])
    assert not result.success
    payload = result.value
    assert payload["error_code"] == "INVALID_ARGS"
    assert payload["failure_class"] == "invalid_args"
    message = payload["message"]
    assert "rules[0]" in message and "row_expression" in message and "Data" in message
    assert "语法" in message
    assert "abs(金额-数量*单价)<0.01" in message  # 中文正确示例
    assert "ABS(E{row}+F{row}+G{row}-C{row})<0.05" in message  # Excel 风格正确示例
    assert "invalid syntax" not in message and "SyntaxError" not in message  # 不泄漏 Python 异常
    assert payload["field"] == "expression"

    # 直接调用同样不泄漏 SyntaxError
    with pytest.raises(ExpressionSyntaxError) as excinfo:
        evaluate_expression("金额-", {})
    assert "invalid syntax" not in str(excinfo.value)


def test_row_expression_unknown_column_is_an_args_error_with_headers(tmp_path):
    path = book(tmp_path / "book.xlsx", [["id", "amount"], [1, 10]])
    result = validate_spreadsheet(path.name, [{"kind": "row_expression", "expression": "不存在列>0"}])
    assert not result.success
    message = result.value["message"]
    assert "rules[0]" in message and "不存在列" in message
    assert "amount" in message  # 可用列名


def test_row_expression_a1_reference_is_rejected_with_actionable_contract(tmp_path):
    """回归（真实会话消息 #78）：abs(B7-…) 想作用于单元格时，须说明不支持 A1 引用并给可行替代写法。"""
    path = book(tmp_path / "book.xlsx", [["B", "C"], [1, 2], [3, 4]])
    result = validate_spreadsheet(path.name, [
        {"kind": "row_expression", "sheet": "Data", "expression": "abs(B7-6.707067669172933)<0.0001 and abs(B8-1.433759398496254)<0.0001"},
    ])
    assert not result.success
    payload = result.value
    assert payload["error_code"] == "INVALID_ARGS"  # 错误码语义不变
    assert payload["failure_class"] == "invalid_args"
    assert payload["unknown_names"] == ["B7", "B8"]  # 字段保留
    assert payload["a1_style_references"] == ["B7", "B8"]  # 新增：显式标记 A1 风格引用
    message = payload["message"]
    assert "rules[0]" in message and "row_expression" in message and "Data" in message
    assert "不支持 A1" in message and "B7" in message  # 明确语言不支持 A1 单元格引用
    assert "B 或 B{row}" in message  # 列字母作用于当前行的写法
    assert "observe_spreadsheet" in message  # 单点断言的可执行替代
    remediation = payload["remediation"]
    assert "不支持 A1" in remediation and "observe_spreadsheet" in remediation and "abs(B-" in remediation

    # 去掉行号的当前行写法可用：这是数据结论（第 2 行不满足），不再是参数错误
    per_row = validate_spreadsheet(path.name, [{"kind": "row_expression", "expression": "abs(B-1)<0.001"}])
    assert per_row.success, per_row.model_text
    assert per_row.value["valid"] is False
    assert per_row.value["failure_count"] == 1
    assert per_row.value["failures"][0]["cell"] == "A3"


def test_row_expression_per_row_errors_recorded_not_raised(tmp_path):
    divided_book = book(tmp_path / "div.xlsx", [["名称", "值"], ["a", 3], ["b", 1]])
    divided = validate_spreadsheet(divided_book.name, [{"kind": "row_expression", "expression": "1/(值-3)<2"}])
    assert divided.success, divided.model_text  # ZeroDivisionError 不逃逸
    assert divided.value["valid"] is False
    assert divided.value["rules"][0]["failed"] == 1
    assert divided.value["failures"][0]["cell"] == "A2"

    text_book = book(tmp_path / "text.xlsx", [["名称", "值"], ["a", 3], ["b", 1], ["c", "文本"]])
    typed = validate_spreadsheet(text_book.name, [{"kind": "row_expression", "expression": "值*2>1"}])
    assert typed.success, typed.model_text  # 文本行按 per-row fail 记录
    assert typed.value["valid"] is False
    assert typed.value["rules"][0]["failed"] == 1
    assert typed.value["failures"][0]["cell"] == "A4"


# ── Bug 4：schema 描述与实际能力对齐 ────────────────────────────────────────


def test_input_schema_documents_column_forms_examples_and_kind_requirements():
    tool = next(t for t in get_tools() if t.name == "validate_spreadsheet")
    items = tool.input_schema["properties"]["rules"]["items"]
    props = items["properties"]
    for key in ("column", "columns"):
        description = props[key]["description"]
        assert "列名" in description
        assert "列字母" in description
        assert "1-based" in description
        assert "列名优先" in description
    expression = props["expression"]["description"]
    assert "abs(金额-数量*单价)<0.01" in expression
    assert "ABS(E{row}+F{row}+G{row}-C{row})<0.05" in expression
    assert "{row}" in expression
    kind = props["kind"]["description"]
    assert "total 需要 column + expected" in kind
    assert "row_expression 需要 expression" in kind
    assert "unique/required 需要 columns" in kind
    assert "foreign_key 需要 columns + reference" in kind
    assert "formula_errors 不需要其他字段" in kind
    reference = props["reference"]["description"]
    assert "file_path" in reference and "columns" in reference
    assert "列字母" in reference and "1-based" in reference
    assert items["required"] == ["kind"]
    # cell 规则已进入枚举与字段合同
    assert "cell" in props["kind"]["enum"]
    assert "cell 需要 cell + expected" in kind
    cell_property = props["cell"]["description"]
    assert "A1 单格" in cell_property and "B7" in cell_property
    assert "区域" in cell_property and "INVALID_ARGS" in cell_property
    expected_property = props["expected"]
    for scalar in ("number", "string", "boolean", "null"):
        assert scalar in expected_property["type"]
    assert "type_mismatch" in expected_property["description"]
    assert props["case_sensitive"]["type"] == "boolean"
    assert "case_sensitive" in kind
    assert "case_sensitive 默认 true" in tool.description


# ── Bug 7：cell 规则（断言具体单元格） ──────────────────────────────────────


def regression_book(path):
    """两表工作簿：数据（广告投入/销售额）+ 回归分析（B7=斜率、B8=截距），模拟真实分析产出。"""
    slope = 6.707067669172933
    intercept = 1.433759398496254
    spend = [1.0, 2.0, 3.0, 4.0, 5.0]
    sales = [intercept + slope * x for x in spend]  # 数据由模型生成，OLS 结果与该系数一致
    n = len(spend)
    mean_x = sum(spend) / n
    mean_y = sum(sales) / n
    covariance = sum((x - mean_x) * (y - mean_y) for x, y in zip(spend, sales))
    variance = sum((x - mean_x) ** 2 for x in spend)

    wb = Workbook()
    data = wb.active
    data.title = "数据"
    data.append(["广告投入", "销售额"])
    for x, y in zip(spend, sales):
        data.append([x, y])
    analysis = wb.create_sheet("回归分析")
    analysis.append(["回归指标", "值"])            # 第 1 行：表头
    analysis.append(["样本量", n])                  # 第 2 行
    analysis.append(["R²", 0.999])                  # 第 3 行
    analysis.append(["标准误差", 0.31])             # 第 4 行
    analysis.append(["t 统计量", 21.64])            # 第 5 行
    analysis.append(["系数", "估计值"])             # 第 6 行：系数表头
    analysis.append(["斜率（广告投入）", covariance / variance])  # B7：本次会话想断言的那一格
    analysis.append(["截距", mean_y - covariance / variance * mean_x])  # B8
    wb.save(path)
    wb.close()
    return covariance / variance, mean_y - covariance / variance * mean_x


def formula_cell_book(path, *, cached=None):
    """回归分析!B2 是公式；cached 给定时在 XML 里注入缓存值，模拟已重算过的文件。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "回归分析"
    ws.append(["指标", "值"])
    ws.append(["斜率", "=1.433759398496254*2"])
    wb.save(path)
    wb.close()
    if cached is not None:
        with zipfile.ZipFile(path) as archive:
            items = {name: archive.read(name) for name in archive.namelist()}
        sheet = "xl/worksheets/sheet1.xml"
        xml = items[sheet].decode("utf-8").replace(
            "<f>1.433759398496254*2</f>", f"<f>1.433759398496254*2</f><v>{cached}</v>"
        )
        items[sheet] = xml.encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in items.items():
                archive.writestr(name, data)
    return path


def test_cell_rule_asserts_regression_slope_in_real_two_sheet_workbook(tmp_path):
    """验收回归：两表工作簿上断言 回归分析!B7 == 6.707067669172933（真实会话 #104/#105 没能表达的断言）。"""
    path = tmp_path / "回归.xlsx"
    slope, intercept = regression_book(path)
    result = validate_spreadsheet(path.name, [
        {"kind": "cell", "sheet": "回归分析", "cell": "B7", "expected": 6.707067669172933, "tolerance": 0.0001},
        {"kind": "cell", "sheet": "回归分析", "cell": "B8", "expected": 1.433759398496254, "tolerance": 0.0001},
    ])
    assert result.success, result.model_text
    assert result.value["valid"] is True and result.value["failure_count"] == 0
    assert result.value["validation_status"] == "complete"
    # 同一坐标与工作表内真实计算出的斜率一致（不只是抄常量）：默认 tolerance 也能通过
    computed = validate_spreadsheet(path.name, [
        {"kind": "cell", "sheet": "回归分析", "cell": "B7", "expected": slope},
        {"kind": "cell", "sheet": "回归分析", "cell": "B8", "expected": intercept},
    ])
    assert computed.success, computed.model_text
    assert computed.value["valid"] is True
    first = computed.value["rules"][0]
    assert first["semantics"] == "cell_assertion" and first["cell"] == "B7"
    assert first["checked"] == 1 and first["failed"] == 0 and first["matched"] is True
    assert first["actual"] == pytest.approx(slope, abs=1e-12)


def test_cell_rule_mismatch_returns_machine_readable_failure(tmp_path):
    path = tmp_path / "book.xlsx"
    book(path, [["指标", "值"], ["斜率", 6.707067669172933], ["截距", 1.433759398496254]])
    result = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": 6.7, "tolerance": 0.0001}])
    assert result.success, result.model_text
    assert result.value["valid"] is False and result.value["failure_count"] == 1
    failure = result.value["failures"][0]
    assert set(failure) >= {"rule", "kind", "sheet", "cell", "expected", "actual", "delta", "reason"}
    assert failure["rule"] == 0 and failure["kind"] == "cell" and failure["sheet"] == "Data" and failure["cell"] == "B2"
    assert failure["actual"] == pytest.approx(6.707067669172933)  # 真实数值，不是中文说明串
    assert failure["expected"] == 6.7
    assert failure["delta"] == pytest.approx(0.007067669172933, abs=1e-12)
    assert failure["reason"] == "value_mismatch"
    assert isinstance(failure["actual"], float) and isinstance(failure["expected"], float)
    assert result.value["rules"][0]["matched"] is False
    # 容差放宽到覆盖 delta 即通过
    assert validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": 6.7, "tolerance": 0.01}]).value["valid"] is True


def test_cell_rule_text_boolean_and_empty_assertions(tmp_path):
    path = tmp_path / "book.xlsx"
    book(path, [["名称", "备注", "启用"], ["a", "华东区", True], ["b", "  ", None], ["c", "ABC", False]])
    exact = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": "华东区"}])
    assert exact.success and exact.value["valid"] is True
    case_fail = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B4", "expected": "abc"}])
    assert case_fail.value["valid"] is False
    assert case_fail.value["failures"][0]["actual"] == "ABC"  # 真实文本
    assert case_fail.value["failures"][0]["reason"] == "value_mismatch"
    case_pass = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B4", "expected": "abc", "case_sensitive": False}])
    assert case_pass.value["valid"] is True
    assert case_pass.value["rules"][0]["case_sensitive"] is False
    # 布尔与空值
    assert validate_spreadsheet(path.name, [{"kind": "cell", "cell": "C2", "expected": True}]).value["valid"] is True
    assert validate_spreadsheet(path.name, [{"kind": "cell", "cell": "C2", "expected": False}]).value["valid"] is False
    assert validate_spreadsheet(path.name, [{"kind": "cell", "cell": "C3", "expected": None}]).value["valid"] is True
    assert validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B3", "expected": None}]).value["valid"] is True  # 纯空白视为空
    not_empty = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": None}])
    assert not_empty.value["valid"] is False
    assert not_empty.value["failures"][0]["actual"] == "华东区"
    assert not_empty.value["failures"][0]["reason"] == "value_mismatch"
    # 未使用坐标 = 空单元格（不是参数错误）
    outside = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "Z99", "expected": None}])
    assert outside.success and outside.value["valid"] is True and outside.value["rules"][0]["checked"] == 1


def test_cell_rule_type_mismatch_and_non_numeric_cell_are_machine_readable(tmp_path):
    path = tmp_path / "book.xlsx"
    book(path, [["名称", "值"], ["数值", 6.7071], ["文本", "abc"], ["数字文本", "6.7071"]])
    # 单元格是数值、expected 是文本：不做类型转换，返回真实类型与真实值
    typed = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": "6.7071"}])
    assert typed.success and typed.value["valid"] is False
    failure = typed.value["failures"][0]
    assert failure["reason"] == "type_mismatch"
    assert failure["actual"] == pytest.approx(6.7071) and failure["expected"] == "6.7071"
    assert "delta" not in failure  # 非数值比对不产生 delta
    # 单元格是文本、expected 是数值：报 non_numeric_cell，actual 是真实文本
    text_cell = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B3", "expected": 6.7071}])
    assert text_cell.value["valid"] is False
    assert text_cell.value["failures"][0]["reason"] == "non_numeric_cell"
    assert text_cell.value["failures"][0]["actual"] == "abc"
    assert "delta" not in text_cell.value["failures"][0]
    # 数值型文本（CSV 导入常见）按数值解析，与 total 口径一致
    numeric_text = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B4", "expected": 6.7071, "tolerance": 1e-9}])
    assert numeric_text.success and numeric_text.value["valid"] is True


@pytest.mark.parametrize("coordinate", ["B7:C9", "B", "7", "B0", "A1048577", "XFE1", "ZZZZ1", "", "  ", "回归分析!B7"])
def test_cell_rule_rejects_invalid_or_out_of_range_coordinates(tmp_path, coordinate):
    path = book(tmp_path / "book.xlsx", [["指标", "值"], ["斜率", 1.0]])
    result = validate_spreadsheet(path.name, [{"kind": "cell", "sheet": "Data", "cell": coordinate, "expected": 1.0}])
    assert not result.success, coordinate
    payload = result.value
    assert payload["error_code"] == "INVALID_ARGS"
    assert payload["failure_class"] == "invalid_args"
    assert payload["rule_index"] == 0 and payload["kind"] == "cell" and payload["field"] == "cell"
    assert "rules[0]" in payload["message"] and "Data" in payload["message"]
    assert "A1" in payload["message"]


def test_cell_rule_argument_contract_errors(tmp_path):
    path = book(tmp_path / "book.xlsx", [["指标", "值"], ["斜率", 1.0]])
    missing_expected = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2"}])
    assert not missing_expected.success and missing_expected.value["error_code"] == "INVALID_ARGS"
    assert missing_expected.value["field"] == "expected"
    bad_tolerance = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": 1.0, "tolerance": -1}])
    assert not bad_tolerance.success and bad_tolerance.value["field"] == "tolerance"
    bad_case = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": "x", "case_sensitive": "yes"}])
    assert not bad_case.success and bad_case.value["field"] == "case_sensitive"
    bad_expected = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": [1, 2]}])
    assert not bad_expected.success and bad_expected.value["field"] == "expected"
    unknown_kind = validate_spreadsheet(path.name, [{"kind": "cell_value", "cell": "B2", "expected": 1.0}])
    assert not unknown_kind.success and "cell" in unknown_kind.value["message"]


def test_cell_rule_requires_sheet_in_multi_sheet_workbook(tmp_path):
    path = tmp_path / "book.xlsx"
    regression_book(path)
    result = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B7", "expected": 6.707067669172933}])
    assert not result.success
    payload = result.value
    assert payload["error_code"] == "SHEET_REQUIRED"
    assert payload["rule_index"] == 0 and payload["available_sheets"] == ["数据", "回归分析"]
    assert "需要指定 sheet" in payload["message"]
    # 顶层 sheet 参数同样能满足定位要求
    top_level = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B7", "expected": 6.707067669172933, "tolerance": 0.0001}], sheet="回归分析")
    assert top_level.success and top_level.value["valid"] is True


def test_cell_rule_uncached_formula_keeps_partial_contract(tmp_path):
    path = formula_cell_book(tmp_path / "formula.xlsx")

    partial = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": 2.867518796992508}])
    assert partial.success, partial.model_text
    assert partial.value["status"] == "partial" and partial.value["validation_status"] == "partial"
    assert partial.value["valid"] is None  # 不声称业务校验通过
    assert partial.value["uncalculated_cells"] == ["回归分析!B2"]
    entry = partial.value["rules"][0]
    assert entry["status"] == "uncalculated" and entry["formula_cache"] == "missing" and entry["failed"] == 0
    assert partial.value["failures"] == []

    blocked = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": 2.867518796992508}], allow_uncached=False)
    assert not blocked.success
    assert blocked.error.code == "FORMULA_CACHE_MISSING"
    assert blocked.value["validation_status"] == "blocked" and blocked.value["cells"] == ["回归分析!B2"]


def test_cell_rule_reads_cached_formula_value(tmp_path):
    """已重算的公式单元格：cell 规则按缓存值比对（与 total 同一取值口径），不回退成空值/零。"""
    path = formula_cell_book(tmp_path / "cached.xlsx", cached=2.867518796992508)
    hit = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": 2.867518796992508}])
    assert hit.success, hit.model_text
    assert hit.value["valid"] is True and hit.value["validation_status"] == "complete"
    assert hit.value["rules"][0]["actual"] == pytest.approx(2.867518796992508)
    miss = validate_spreadsheet(path.name, [{"kind": "cell", "cell": "B2", "expected": 3.0}])
    assert miss.value["valid"] is False
    assert miss.value["failures"][0]["actual"] == pytest.approx(2.867518796992508)  # 真实缓存值，不是 None/0
    assert miss.value["failures"][0]["reason"] == "value_mismatch"


def test_cell_rule_does_not_change_other_kinds_or_formula_errors(tmp_path):
    path = book(tmp_path / "book.xlsx", [["id", "amount"], [1, 10], [2, 20]])
    mixed = validate_spreadsheet(path.name, [
        {"kind": "unique", "column": "id"},
        {"kind": "cell", "cell": "B3", "expected": 20},
        {"kind": "total", "column": "B", "expected": 30},
        {"kind": "formula_errors"},
    ])
    assert mixed.success, mixed.model_text
    assert mixed.value["valid"] is True
    assert [entry["kind"] for entry in mixed.value["rules"]] == ["unique", "cell", "total", "formula_errors"]
    assert mixed.value["rules"][2]["semantics"] == "column_total"
    assert "semantics" not in mixed.value["rules"][0]
