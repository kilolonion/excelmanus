"""Real call regressions: copy aliases, formula rules and large-result retrieval."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from excelmanus.engine_core.spill import SpillStore
from excelmanus.tools.context import use_workspace
from excelmanus.tools.registry import ToolRegistry
from excelmanus.workbook_commit import content_version_of_file


@pytest.fixture
def order_book(tmp_path: Path) -> Path:
    path = tmp_path / "orders.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws.append(["订单号", "日期", "产品编号", "数量", "客户", "单价", "金额"])
    for i in range(12):
        ws.append([i + 1, "2026-09-19", f"SKU-{i % 8}", i + 1, "测试", None, None])
    ws["C1"].font = Font(bold=True, color="FFFFFF")
    ws["C1"].fill = PatternFill(patternType="solid", fgColor="336699")
    products = wb.create_sheet("产品目录")
    products.append(["编号", "名称", "单价", "分类"])
    for i in range(8):
        products.append([f"SKU-{i}", f"产品{i}", i + 10, "测试"])
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def registry(tmp_path: Path) -> ToolRegistry:
    registry = ToolRegistry()
    registry.register_builtin_tools(str(tmp_path))
    registry.configure_schema_validation(mode="enforce", canary_percent=100, strict_path=False)
    return registry


def _edit_operations() -> list[dict]:
    operations = [
        {"kind": "insert", "sheet": "订单", "axis": "column", "at": 4, "count": 1},
        {"kind": "copy", "sheet": "订单", "source_sheet": "订单", "source_range": "C1",
         "target_sheet": "订单", "target_start": "D1"},
        {"kind": "write", "sheet": "订单", "start_cell": "D1", "values": [["产品名称"]]},
    ]
    for column, formula in [
        ("D", '=IFERROR(VLOOKUP($C{r},产品目录!$A$2:$D$9,2,0),"未匹配")'),
        ("G", '=IFERROR(VLOOKUP($C{r},产品目录!$A$2:$C$9,3,0),"")'),
        ("H", '=IF($G{r}="","未匹配",$G{r}*$E{r})'),
    ]:
        operations.append({"kind": "write", "sheet": "订单", "start_cell": column + "2",
                           "values": [[formula.format(r=r)] for r in range(2, 14)]})
    return operations


def _format_operations() -> list[dict]:
    return [
        {"kind": "size", "sheet": "订单", "columns": {"D": 14}},
        {"kind": "conditional_format", "sheet": "订单", "range": "A2:H13",
         "rule": {"type": "formula", "formula1": '=$D2="未匹配"',
                  "fill": {"patternType": "solid", "fgColor": "FFC7CE"},
                  "font": {"color": "FF9C0006"}}},
    ]


def test_reported_edit_batch(registry: ToolRegistry, order_book: Path) -> None:
    with use_workspace(order_book.parent):
        result = registry.call_tool("edit_spreadsheet", {
            "file_path": order_book.name, "expected_version": content_version_of_file(order_book),
            "operations": _edit_operations(),
        })
    assert result.success, result.model_text
    wb = load_workbook(order_book)
    try:
        ws = wb["订单"]
        assert ws.max_column == 8
        assert ws["D1"].value == "产品名称"
        assert ws["D1"]._style == ws["C1"]._style
        assert ws["D13"].value == '=IFERROR(VLOOKUP($C13,产品目录!$A$2:$D$9,2,0),"未匹配")'
        assert ws["G2"].value == '=IFERROR(VLOOKUP($C2,产品目录!$A$2:$C$9,3,0),"")'
        assert ws["H13"].value == '=IF($G13="","未匹配",$G13*$E13)'
    finally:
        wb.close()


def test_reported_format_batch(registry: ToolRegistry, order_book: Path) -> None:
    with use_workspace(order_book.parent):
        result = registry.call_tool("format_spreadsheet", {
            "file_path": order_book.name, "expected_version": content_version_of_file(order_book),
            "operations": _format_operations(),
        })
    assert result.success, result.model_text
    wb = load_workbook(order_book)
    try:
        ws = wb["订单"]
        assert ws.column_dimensions["D"].width == 14
        cf = next(iter(ws.conditional_formatting))
        assert str(cf.sqref) == "A2:H13"
        rule = ws.conditional_formatting[cf][0]
        assert rule.formula == ['$D2="未匹配"']
        assert rule.dxf.fill.fgColor.rgb.endswith("FFC7CE")
        assert rule.dxf.font.color.rgb == "FF9C0006"
    finally:
        wb.close()


@pytest.mark.parametrize("tool_name", ["read_text_file", "inspect_spreadsheet"])
@pytest.mark.parametrize("prefix", ["spill", "result_spill", "selection_spill"])
def test_spill_reference_retrieval_across_registry(
    registry: ToolRegistry, tmp_path: Path, tool_name: str, prefix: str,
) -> None:
    original = {"status": "success", "values": [["未匹配" * 3000]]}
    locator = SpillStore(tmp_path).put(json.dumps(original, ensure_ascii=False))
    with use_workspace(tmp_path):
        result = registry.call_tool(tool_name, {"file_path": locator.replace("spill:", prefix + ":")})
    assert result.success, result.model_text
    assert result.value == original
    assert json.loads(result.model_text) == original
    assert result.coverage["spill_retrieve"]


@pytest.mark.asyncio
async def test_reported_calls_through_real_native_dispatcher(registry: ToolRegistry, order_book: Path) -> None:
    from excelmanus.config import ExcelManusConfig
    from excelmanus.engine import AgentEngine

    engine = AgentEngine(ExcelManusConfig(
        api_key="test", base_url="https://test.invalid/v1", model="test",
        workspace_root=str(order_book.parent),
    ), registry)

    async def call(name: str, args: dict):
        tc = SimpleNamespace(id=name, function=SimpleNamespace(name=name, arguments=json.dumps(args)))
        out = await engine._tool_runtime.execute(tc, None, None, 1)
        assert out.success, out.result
        return json.loads(out.result)

    edited = await call("edit_spreadsheet", {
        "file_path": order_book.name, "expected_version": content_version_of_file(order_book),
        "operations": _edit_operations(),
    })
    formatted = await call("format_spreadsheet", {
        "file_path": order_book.name, "expected_version": edited["content_version"],
        "operations": _format_operations(),
    })
    read = await call("inspect_spreadsheet", {
        "file_path": order_book.name, "mode": "range", "sheet": "订单", "range": "A1:H13",
        "include": ["formulas"], "expected_version": formatted["content_version"],
    })
    assert read["result_spill"].startswith("spill:")
    retrieved = await call("read_text_file", {"file_path": read["result_spill"].replace("spill:", "result_spill:")})
    assert retrieved["content_version"] == formatted["content_version"]
    assert retrieved["shape"] == {"rows": 13, "columns": 8}
    assert retrieved["formula_grid"][12][7] == '=IF($G13="","未匹配",$G13*$E13)'


@pytest.mark.parametrize("addresses", [
    {"sheet": "订单"},
    {"sheet_name": "订单"},
    {"source_sheet": "订单"},
    {"target_sheet": "订单"},
    {"sheet": "订单", "source_sheet": "产品目录"},
    {"source_range": "'产品目录'!A1", "target_start": "'订单'!D1"},
])
def test_copy_resolves_sheet_aliases(registry: ToolRegistry, order_book: Path, addresses: dict) -> None:
    op = {"kind": "copy", "source_range": "C1", "target_start": "D1", **addresses}
    with use_workspace(order_book.parent):
        result = registry.call_tool("edit_spreadsheet", {
            "file_path": order_book.name, "expected_version": content_version_of_file(order_book),
            "operations": [op],
        })
    assert result.success, result.model_text
    wb = load_workbook(order_book)
    try:
        expected = "编号" if "!" in op["source_range"] else "单价" if op.get("source_sheet") == "产品目录" else "产品编号"
        assert wb["订单"]["D1"].value == expected
    finally:
        wb.close()


@pytest.mark.parametrize("conflict", [
    {"sheet": "产品目录", "target_sheet": "订单"},
    {"sheet": "订单", "sheet_name": "产品目录"},
    {"source_sheet": "订单", "sourceSheet": "产品目录"},
    {"target_sheet": "订单", "target_start": "产品目录!D1"},
    {"source_sheet": "订单", "source_range": "产品目录!C1"},
])
def test_copy_conflict_aborts_whole_batch(registry: ToolRegistry, order_book: Path, conflict: dict) -> None:
    before = order_book.read_bytes()
    with use_workspace(order_book.parent):
        result = registry.call_tool("edit_spreadsheet", {
            "file_path": order_book.name, "expected_version": content_version_of_file(order_book),
            "operations": [
                {"kind": "write", "sheet": "订单", "start_cell": "A1", "values": [["MUST NOT COMMIT"]]},
                {"kind": "copy", "sheet": "订单", "source_range": "C1", "target_start": "D1", **conflict},
            ],
        })
    assert not result.success
    assert result.value["operation_index"] == 1
    assert result.value["operation_kind"] == "copy"
    assert result.value["committed"] is False
    assert result.value["applied"] == []
    assert "operations[1]" in result.model_text
    assert order_book.read_bytes() == before


@pytest.mark.parametrize("kind, rule_type", [("conditional_format", "formula"), ("data_validation", "custom")])
@pytest.mark.parametrize("formula_fields", [
    {"formula": '=$D2="未匹配"'},
    {"formula1": '=$D2="未匹配"'},
    {"formula": ['$D2="未匹配"']},
    {"formula": ['$D2="未匹配"'], "formula1": '=$D2="未匹配"'},
])
def test_expression_fields_share_semantics(
    registry: ToolRegistry, order_book: Path, kind: str, rule_type: str, formula_fields: dict,
) -> None:
    with use_workspace(order_book.parent):
        result = registry.call_tool("format_spreadsheet", {
            "file_path": order_book.name, "expected_version": content_version_of_file(order_book),
            "operations": [{"kind": kind, "sheet": "订单", "range": "A2:H13",
                            "rule": {"type": rule_type, **formula_fields}}],
        })
    assert result.success, result.model_text
    wb = load_workbook(order_book)
    try:
        ws = wb["订单"]
        if kind == "conditional_format":
            key = next(iter(ws.conditional_formatting))
            assert ws.conditional_formatting[key][0].formula == ['$D2="未匹配"']
        else:
            assert ws.data_validations.dataValidation[0].formula1 == '=$D2="未匹配"'
    finally:
        wb.close()


@pytest.mark.parametrize("kind, rule_type", [("conditional_format", "formula"), ("data_validation", "custom")])
def test_formula_alias_conflict_aborts_format_batch(
    registry: ToolRegistry, order_book: Path, kind: str, rule_type: str,
) -> None:
    before = order_book.read_bytes()
    with use_workspace(order_book.parent):
        result = registry.call_tool("format_spreadsheet", {
            "file_path": order_book.name, "expected_version": content_version_of_file(order_book),
            "operations": [
                {"kind": "size", "sheet": "订单", "columns": {"D": 14}},
                {"kind": kind, "sheet": "订单", "range": "A2:H13",
                 "rule": {"type": rule_type, "formula": "=$D2=1", "formula1": "=$D2=2"}},
            ],
        })
    assert not result.success
    assert "formula1" in result.error.message and "冲突" in result.error.message
    assert result.value["operation_index"] == 1
    assert result.value["committed"] is False
    assert order_book.read_bytes() == before


def test_stale_edit_retry_never_duplicates_insert(registry: ToolRegistry, order_book: Path) -> None:
    version = content_version_of_file(order_book)
    args = {"file_path": order_book.name, "expected_version": version, "operations": _edit_operations()}
    with use_workspace(order_book.parent):
        first = registry.call_tool("edit_spreadsheet", args)
        assert first.success, first.model_text
        after = order_book.read_bytes()
        retry = registry.call_tool("edit_spreadsheet", args)
    assert not retry.success
    assert retry.error.code == "VERSION_CONFLICT"
    assert order_book.read_bytes() == after


def test_spill_error_guidance_and_content_not_reinterpreted(registry: ToolRegistry, tmp_path: Path) -> None:
    from excelmanus.engine_core.spill import extract_spill_locator

    valid_missing = "result_spill:" + "a" * 64
    assert extract_spill_locator({"file_path": "book.xlsx", "query": valid_missing}) is None
    assert extract_spill_locator({"query": valid_missing}) is None
    with use_workspace(tmp_path):
        missing = registry.call_tool("read_text_file", {"file_path": valid_missing})
        invalid = registry.call_tool("read_text_file", {"file_path": "result_spill:../bad"})
    assert missing.error.code == "NOT_FOUND"
    assert missing.coverage["kind"] == "missing"
    assert "只读查询" in missing.value["remediation"]
    assert "不要重放写入" in missing.value["remediation"]
    assert invalid.error.code == "INVALID_ARGS"
    assert invalid.coverage["kind"] == "invalid"


def test_spill_next_call_is_executable(registry: ToolRegistry, tmp_path: Path) -> None:
    from excelmanus.engine_core.spill import expose_spreadsheet_value, project_for_wire
    from excelmanus.engine_core.tool_result import ToolResult

    raw = json.dumps({"status": "success", "values": [["abc" * 10000]]})
    store = SpillStore(tmp_path)
    projected = project_for_wire(raw, store=store).model_text
    native = expose_spreadsheet_value(ToolResult(success=True, value=json.loads(raw), model_text="preview"), store=store).model_text
    for text in [projected, native]:
        next_call = json.loads(text)["next_call"]
        assert next_call["arguments"]["file_path"].startswith("spill:")
        with use_workspace(tmp_path):
            result = registry.call_tool(next_call["tool"], next_call["arguments"])
        assert result.success, result.model_text
        assert json.loads(result.model_text) == json.loads(raw)


@pytest.mark.asyncio
async def test_generated_sdk_executes_same_calls(registry: ToolRegistry, order_book: Path) -> None:
    import sys

    from excelmanus.config import ExcelManusConfig
    from excelmanus.engine import AgentEngine

    engine = AgentEngine(ExcelManusConfig(
        api_key="test", base_url="https://test.invalid/v1", model="test",
        workspace_root=str(order_book.parent),
    ), registry)
    engine._full_access_enabled = True
    locator = SpillStore(order_book.parent).put(json.dumps({"values": [["complete result"]]}))
    code = (
        "import em, json\n"
        "before = em.inspect_spreadsheet(file_path='orders.xlsx', sheet='订单', mode='overview')\n"
        f"edit_ops = json.loads({json.dumps(_edit_operations(), ensure_ascii=False)!r})\n"
        f"format_ops = json.loads({json.dumps(_format_operations(), ensure_ascii=False)!r})\n"
        "edited = em.edit_spreadsheet(path='orders.xlsx', expected_version=before['content_version'], operations=edit_ops)\n"
        "formatted = em.format_spreadsheet(path='orders.xlsx', expected_version=edited['content_version'], operations=format_ops)\n"
        f"stored = em.read_text_file(file_path={locator.replace('spill:', 'result_spill:')!r})\n"
        "assert stored['values'] == [['complete result']]\n"
        "print('SDK calls completed')\n"
    )
    tc = SimpleNamespace(id="sdk", function=SimpleNamespace(name="run_code", arguments=json.dumps({
        "code": code, "python_command": sys.executable,
        "timeout_seconds": 30, "require_excel_deps": False,
    })))
    outcome = await engine._tool_runtime.execute(tc, None, None, 1)
    assert outcome.success, outcome.result
    assert "SDK calls completed" in outcome.result, outcome.result
    wb = load_workbook(order_book)
    try:
        assert wb["订单"]["D1"].value == "产品名称"
        assert wb["订单"].max_column == 8
        assert len(wb["订单"].conditional_formatting) == 1
    finally:
        wb.close()
