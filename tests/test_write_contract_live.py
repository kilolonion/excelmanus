"""真实 dispatcher + 临时工作簿交叉：版本冲突、取消、预算、组合任务。"""

from __future__ import annotations

import asyncio
import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.code_mode import CodeModeSession, build_session_for_run_code
from excelmanus.config import ExcelManusConfig
from excelmanus.engine import AgentEngine
from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.events import EventType
from excelmanus.tools.intent_tools import get_tools as intent_get_tools
from excelmanus.tools.registry import ToolRegistry
from excelmanus.workbook_commit import content_version_of_file


def _make_config(tmp_path: Path, **overrides: object) -> ExcelManusConfig:
    defaults: dict[str, object] = {
        "api_key": "test-key",
        "base_url": "https://test.example.com/v1",
        "model": "test-model",
        "workspace_root": str(tmp_path),
        "backup_enabled": False,
        "max_iterations": 8,
    }
    defaults.update(overrides)
    return ExcelManusConfig(**defaults)


def _live_registry() -> ToolRegistry:
    registry = ToolRegistry()
    registry.register_tools(intent_get_tools())
    return registry


def _inspect_args(path: str) -> dict[str, object]:
    return {"mode": "range", "file_path": path, "sheet_name": "Sheet1", "max_rows": 8}


def _edit_args(path: str, cell: str, value: object) -> dict[str, object]:
    return {
        "file_path": path,
        "operations": [{
            "kind": "write",
            "sheet": "Sheet1",
            "start_cell": cell,
            "values": [[value]],
        }],
    }


def _make_engine(tmp_path: Path) -> AgentEngine:
    from excelmanus.tools import intent_tools
    from excelmanus.workbook import data as data_tools

    intent_tools.init_guard(str(tmp_path))
    data_tools.init_guard(str(tmp_path))
    engine = AgentEngine(_make_config(tmp_path), _live_registry())
    engine._full_access_enabled = True
    return engine


def _seed_book(path: Path, value: str = "seen") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = value
    wb.save(path)
    wb.close()


def _cell(path: Path, ref: str = "A1") -> object:
    wb = load_workbook(path)
    try:
        return wb.active[ref].value
    finally:
        wb.close()


def _error_code(result: ToolResult) -> str | None:
    if result.error is None:
        return None
    return result.error.code


async def _sub(
    engine: AgentEngine,
    tool_name: str,
    arguments: dict[str, object],
    *,
    root_call_id: str = "live",
) -> ToolResult:
    return await engine._tool_dispatcher.execute_subcall(
        tool_name=tool_name,
        arguments=arguments,
        root_call_id=root_call_id,
    )


def _inject_vba(path: Path, marker: bytes) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(buf, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                if "vbaProject.bin" not in text:
                    text = text.replace(
                        "</Types>",
                        '<Override PartName="/xl/vbaProject.bin" '
                        'ContentType="application/vnd.ms-office.vbaProject"/>'
                        "</Types>",
                    )
                data = text.encode("utf-8")
            zout.writestr(info, data)
        zout.writestr("xl/vbaProject.bin", marker)
    path.write_bytes(buf.getvalue())


def _vba_bytes(path: Path) -> bytes | None:
    with zipfile.ZipFile(path, "r") as zf:
        try:
            return zf.read("xl/vbaProject.bin")
        except KeyError:
            return None


@pytest.mark.asyncio
async def test_same_session_read_write_then_reread(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    _seed_book(tmp_path / "book.xlsx", "seen")

    read1 = await _sub(engine, "inspect_spreadsheet", _inspect_args("book.xlsx"))
    assert read1.success
    assert str((read1.value or {}).get("content_version") or "").startswith("sha256:")

    written = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "agent"),
    )
    assert written.success, written.model_text
    assert (written.value or {}).get("file_path") == "book.xlsx"
    assert _cell(tmp_path / "book.xlsx") == "agent"

    read2 = await _sub(engine, "inspect_spreadsheet", _inspect_args("book.xlsx"))
    assert read2.success
    blob = str(read2.value or {})
    assert "agent" in blob


@pytest.mark.asyncio
async def test_human_edit_then_agent_write_conflicts_and_keeps_bytes(
    tmp_path: Path,
) -> None:
    engine = _make_engine(tmp_path)
    book = tmp_path / "book.xlsx"
    _seed_book(book, "seen")

    read1 = await _sub(engine, "inspect_spreadsheet", _inspect_args("book.xlsx"))
    assert read1.success

    outsider = load_workbook(book)
    outsider.active["A1"] = "human"
    outsider.save(book)
    outsider.close()
    before = book.read_bytes()

    conflict = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "agent"),
    )
    assert not conflict.success
    assert _error_code(conflict) == "VERSION_CONFLICT"
    assert book.read_bytes() == before
    assert _cell(book) == "human"

    refreshed = await _sub(engine, "inspect_spreadsheet", _inspect_args("book.xlsx"))
    assert refreshed.success
    ok = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "after-refresh"),
    )
    assert ok.success, ok.model_text
    assert _cell(book) == "after-refresh"


@pytest.mark.asyncio
async def test_two_engines_second_write_conflicts(tmp_path: Path) -> None:
    left = _make_engine(tmp_path)
    right = _make_engine(tmp_path)
    book = tmp_path / "book.xlsx"
    _seed_book(book, "shared")

    assert (await _sub(left, "inspect_spreadsheet", _inspect_args("book.xlsx"))).success
    assert (await _sub(right, "inspect_spreadsheet", _inspect_args("book.xlsx"))).success

    first = await _sub(
        left,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "from-a"),
    )
    assert first.success, first.model_text
    before_b = book.read_bytes()

    second = await _sub(
        right,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "from-b"),
    )
    assert not second.success
    assert _error_code(second) == "VERSION_CONFLICT"
    assert book.read_bytes() == before_b
    assert _cell(book) == "from-a"


@pytest.mark.asyncio
async def test_same_session_rapid_writes_use_new_version(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    book = tmp_path / "book.xlsx"
    _seed_book(book, "v0")

    assert (await _sub(engine, "inspect_spreadsheet", _inspect_args("book.xlsx"))).success
    first = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "v1"),
    )
    assert first.success, first.model_text
    second = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "B1", "v2"),
    )
    assert second.success, second.model_text
    assert _cell(book, "A1") == "v1"
    assert _cell(book, "B1") == "v2"
    assert (first.value or {}).get("content_version") != (
        second.value or {}
    ).get("content_version")


@pytest.mark.asyncio
async def test_same_basename_different_dirs_do_not_share_version(
    tmp_path: Path,
) -> None:
    engine = _make_engine(tmp_path)
    left = tmp_path / "left" / "book.xlsx"
    right = tmp_path / "right" / "book.xlsx"
    _seed_book(left, "L")
    _seed_book(right, "R")

    assert (await _sub(engine, "inspect_spreadsheet", _inspect_args("left/book.xlsx"))).success
    assert (await _sub(engine, "inspect_spreadsheet", _inspect_args("right/book.xlsx"))).success

    written = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("left/book.xlsx", "A1", "L2"),
    )
    assert written.success, written.model_text
    other = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("right/book.xlsx", "A1", "R2"),
    )
    assert other.success, other.model_text
    assert _cell(left) == "L2"
    assert _cell(right) == "R2"


@pytest.mark.asyncio
async def test_write_outside_workspace_is_path_invalid(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    outsider = tmp_path.parent / "excelmanus-live-outside.xlsx"
    _seed_book(outsider, "secret")
    try:
        result = await _sub(
            engine,
            "edit_spreadsheet",
            _edit_args(str(outsider), "A1", "pwn"),
        )
        assert not result.success
        assert _error_code(result) == "PATH_INVALID"
        assert _cell(outsider) == "secret"
    finally:
        outsider.unlink(missing_ok=True)


@pytest.mark.asyncio
async def test_xlsm_write_keeps_vba_bytes(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    book = tmp_path / "macro.xlsm"
    marker = b"EM-VBA-MARKER-9f3c"
    _seed_book(book, "macro")
    _inject_vba(book, marker)
    assert _vba_bytes(book) == marker

    assert (await _sub(engine, "inspect_spreadsheet", _inspect_args("macro.xlsm"))).success
    written = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("macro.xlsm", "A1", "kept"),
    )
    assert written.success, written.model_text
    assert _cell(book) == "kept"
    assert _vba_bytes(book) == marker


@pytest.mark.asyncio
async def test_cancel_rejects_later_subcall_not_just_sleep(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    _seed_book(tmp_path / "book.xlsx", "seen")
    assert (await _sub(engine, "inspect_spreadsheet", _inspect_args("book.xlsx"))).success

    engine._tool_dispatcher.request_cancel()
    blocked = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "nope"),
    )
    assert not blocked.success
    assert _error_code(blocked) == "CANCELLED"
    assert _cell(tmp_path / "book.xlsx") == "seen"

    engine._tool_dispatcher.reset_cancel()
    ok = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "after-reset"),
    )
    assert ok.success, ok.model_text


@pytest.mark.asyncio
async def test_shared_budget_blocks_second_subcall(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    _seed_book(tmp_path / "book.xlsx", "seen")
    dispatcher = engine._tool_dispatcher
    dispatcher.begin_call_budget(1)

    first = await _sub(engine, "inspect_spreadsheet", _inspect_args("book.xlsx"))
    assert first.success
    second = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "over"),
    )
    assert not second.success
    assert _error_code(second) == "BUDGET_EXCEEDED"
    assert _cell(tmp_path / "book.xlsx") == "seen"


@pytest.mark.asyncio
async def test_partial_commit_stays_after_cancel(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    book = tmp_path / "book.xlsx"
    _seed_book(book, "seen")
    assert (await _sub(engine, "inspect_spreadsheet", _inspect_args("book.xlsx"))).success

    first = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "committed"),
    )
    assert first.success, first.model_text
    engine._tool_dispatcher.request_cancel()
    later = await _sub(
        engine,
        "edit_spreadsheet",
        _edit_args("book.xlsx", "B1", "dropped"),
    )
    assert not later.success
    assert _error_code(later) == "CANCELLED"
    assert _cell(book, "A1") == "committed"
    assert _cell(book, "B1") is None


@pytest.mark.asyncio
async def test_code_mode_session_read_write_reread(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    _seed_book(tmp_path / "book.xlsx", "seen")
    events: list[object] = []
    session = build_session_for_run_code(
        engine._tool_dispatcher,
        root_call_id="call_compose",
        on_event=events.append,
    )
    session._loop = asyncio.get_running_loop()

    async def _via_session(tool_name: str, arguments: dict[str, object]) -> ToolResult:
        return await asyncio.to_thread(
            session._call_dispatcher,
            tool_name=tool_name,
            arguments=arguments,
            root_call_id="call_compose",
        )

    read1 = await _via_session("inspect_spreadsheet", _inspect_args("book.xlsx"))
    written = await _via_session(
        "edit_spreadsheet",
        _edit_args("book.xlsx", "A1", "via-sdk"),
    )
    read2 = await _via_session("inspect_spreadsheet", _inspect_args("book.xlsx"))

    assert isinstance(read1, ToolResult) and read1.success
    assert isinstance(written, ToolResult) and written.success, written.model_text
    assert isinstance(read2, ToolResult) and read2.success
    assert _cell(tmp_path / "book.xlsx") == "via-sdk"
    assert "via-sdk" in str(read2.value or {})
    assert session.allocate_subcall_id("inspect_spreadsheet", "call_compose") != (
        session.allocate_subcall_id("edit_spreadsheet", "call_compose")
    )
    start_ids = [
        getattr(event, "tool_call_id", None)
        for event in events
        if getattr(event, "event_type", None) == EventType.TOOL_CALL_START
        and getattr(event, "tool_call_id", None)
    ]
    assert len(set(start_ids)) >= 3


@pytest.mark.asyncio
async def test_code_mode_session_honors_cancel(tmp_path: Path) -> None:
    engine = _make_engine(tmp_path)
    session = CodeModeSession(
        dispatcher=engine._tool_dispatcher,
        root_call_id="call_cancel",
        bridge_dir=tmp_path / ".tmp" / "code_mode" / "cancel",
        call_timeout=8.0,
    )
    session._loop = asyncio.get_running_loop()
    engine._tool_dispatcher.request_cancel()
    blocked = await asyncio.to_thread(
        session._call_dispatcher,
        tool_name="inspect_spreadsheet",
        arguments=_inspect_args("missing.xlsx"),
        root_call_id="call_cancel",
    )
    assert isinstance(blocked, ToolResult)
    assert not blocked.success
    assert _error_code(blocked) == "CANCELLED"
