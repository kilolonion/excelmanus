"""P2 工具结果三分投影契约测试。"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from excelmanus.engine_core.tool_dispatcher import ToolDispatcher
from excelmanus.engine_core.tool_result import ToolResult, ToolUiMeta
from excelmanus.workbook import data as data_tools


def _make_dispatcher(*, vision: bool = True) -> tuple[ToolDispatcher, MagicMock]:
    engine = MagicMock()
    engine.memory = MagicMock()
    engine.is_vision_capable = vision
    engine.state = MagicMock()
    engine.transaction = None
    dispatcher = ToolDispatcher.__new__(ToolDispatcher)
    dispatcher._engine = engine
    dispatcher._deferred_image_injections = []
    dispatcher._injected_image_hashes = set()
    dispatcher._tool_call_store = None
    dispatcher._handlers = []
    return dispatcher, engine


@pytest.fixture(autouse=True)
def _init_guard(tmp_path: Path) -> None:
    data_tools.init_guard(str(tmp_path))


class TestToolResultAdapter:
    def test_from_text_does_not_parse_magic_fields(self) -> None:
        raw = json.dumps({
            "status": "ok",
            "__tool_result_image__": {"base64": "abc", "mime_type": "image/png"},
            "cow_mapping": {"a.xlsx": "outputs/a.xlsx"},
        })
        tr = ToolResult.from_text(raw)
        assert tr.model_text == raw
        assert tr.ui_meta.image is None
        assert tr.ui_meta.cow_mapping is None

    def test_dispatcher_str_result_lifts_image_magic_field(self) -> None:
        dispatcher, engine = _make_dispatcher()
        b64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
        raw = json.dumps({
            "status": "ok",
            "__tool_result_image__": {"base64": b64, "mime_type": "image/png"},
        })
        tr = dispatcher._coerce_tool_result(raw)
        dispatcher._apply_ui_meta_effects(tr)
        assert tr.ui_meta.image is not None
        assert tr.ui_meta.image["base64"] == b64
        assert "__tool_result_image__" not in tr.model_text
        assert len(dispatcher._deferred_image_injections) == 1

    def test_dispatcher_lifts_download_diff_and_cow(self) -> None:
        dispatcher, engine = _make_dispatcher()
        raw = json.dumps({
            "status": "success",
            "file": "a.txt",
            "_file_download": {"file_path": "a.txt", "filename": "a.txt", "description": "d"},
            "_excel_diff": {"file_path": "a.xlsx", "sheet": "S", "changes": [{"cell": "A1", "old": 1, "new": 2}]},
            "_text_diff": {"file_path": "a.txt", "hunks": ["--- a", "+++ b"], "additions": 1, "deletions": 0},
            "cow_mapping": {"src.xlsx": "outputs/src.xlsx"},
        })
        tr = dispatcher._coerce_tool_result(raw)
        assert tr.ui_meta.download["file_path"] == "a.txt"
        assert tr.ui_meta.diff["file_path"] == "a.xlsx"
        assert tr.ui_meta.text_diff["file_path"] == "a.txt"
        assert tr.ui_meta.cow_mapping == {"src.xlsx": "outputs/src.xlsx"}
        for key in ("__tool_result_image__", "_file_download", "_excel_diff", "_text_diff"):
            assert key not in tr.model_text
        assert "cow_mapping" in tr.model_text
        events: list[Any] = []
        engine.emit = lambda _cb, ev: events.append(ev)
        dispatcher._emit_ui_meta_events(
            engine, object(), "call-1", "write_text_file", {}, tr.ui_meta, 0,
        )
        kinds = {getattr(ev, "event_type", None) for ev in events}
        from excelmanus.events import EventType
        assert EventType.FILE_DOWNLOAD in kinds
        assert EventType.EXCEL_DIFF in kinds
        assert EventType.TEXT_DIFF in kinds

    def test_dispatcher_ui_meta_image_injects(self) -> None:
        dispatcher, engine = _make_dispatcher()
        b64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
        tr = ToolResult(
            success=True,
            model_text="图片已加载",
            ui_meta=ToolUiMeta(image={"base64": b64, "mime_type": "image/png", "detail": "auto"}),
        )
        dispatcher._apply_ui_meta_effects(tr)
        assert len(dispatcher._deferred_image_injections) == 1
        dispatcher.flush_deferred_images()
        engine.memory.add_image_message.assert_called_once()


class TestReadExcelToolResult:
    @pytest.fixture()
    def large_excel(self, tmp_path: Path) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.append(["A", "B"])
        for i in range(1, 51):
            ws.append([i, i * 2])
        fp = tmp_path / "large.xlsx"
        wb.save(fp)
        return fp

    def test_returns_tool_result_with_truncation_metadata(self, large_excel: Path) -> None:
        tr = data_tools.read_excel(str(large_excel), max_rows=10)
        assert isinstance(tr, ToolResult)
        assert tr.success
        assert tr.coverage is not None
        assert tr.coverage.get("kind") in {"truncated", "complete", "sampled"}
        assert tr.value is not None
        assert tr.ui_meta.files == ["large.xlsx"]
        assert tr.ui_meta.preview is not None
        # model_text 不应包含完整 50 行 dump
        assert tr.model_text.count('"A"') < 20
        assert "截断" in tr.model_text or tr.coverage.get("kind") != "truncated"
        assert tr.ui_meta.content_version
        assert str(tr.ui_meta.content_version).startswith("sha256:")
        assert tr.value.get("content_version") == tr.ui_meta.content_version


class TestResultConstructors:
    def test_ok_result_lifts_version_and_path(self) -> None:
        from excelmanus.engine_core.tool_result import ok_result

        tr = ok_result(
            {"file_path": "book.xlsx", "content_version": "sha256:abc", "n": 1},
            model_text="ok",
        )
        assert tr.success
        assert tr.value["n"] == 1
        assert tr.ui_meta.files == ["book.xlsx"]
        assert tr.ui_meta.content_version == "sha256:abc"
        assert tr.model_text == "ok"

    def test_error_result_keeps_error_field(self) -> None:
        from excelmanus.engine_core.tool_result import error_result

        tr = error_result("bad page", code="INVALID_ARGS", fields={"error": "bad page"})
        assert tr.success is False
        assert tr.value["error"] == "bad page"
        assert tr.error is not None
        assert tr.error.code == "INVALID_ARGS"


class TestCoerceLegacyError:
    def test_status_error_json_is_failure_not_ok(self) -> None:
        dispatcher, _engine = _make_dispatcher()
        tr = dispatcher._coerce_tool_result(
            json.dumps({"status": "error", "error": "boom", "code": "EXECUTION_FAILED"})
        )
        assert tr.success is False
        assert tr.error is not None
        assert tr.error.code == "EXECUTION_FAILED"
        assert tr.ui_meta.content_version is None

    def test_lifts_content_version_from_legacy_json(self) -> None:
        dispatcher, _engine = _make_dispatcher()
        tr = dispatcher._coerce_tool_result(
            json.dumps({"status": "success", "content_version": "sha256:abc"})
        )
        assert tr.success is True
        assert tr.ui_meta.content_version == "sha256:abc"

    def test_revision_projects_to_sse_ui(self) -> None:
        ui = ToolUiMeta(
            files=["book.xlsx"],
            content_version="sha256:abc",
            revision={"revision_id": "rev_1", "label": "before"},
        )
        payload = ui.to_sse_ui()
        assert payload is not None
        assert payload["revision"]["revision_id"] == "rev_1"
        assert payload["content_version"] == "sha256:abc"


class TestCompareExcelToolResult:
    @pytest.fixture()
    def pair(self, tmp_path: Path) -> tuple[Path, Path]:
        fa = tmp_path / "a.xlsx"
        fb = tmp_path / "b.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "val"])
        ws.append([1, 10])
        ws.append([2, 20])
        wb.save(fa)
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "val"])
        ws.append([1, 10])
        ws.append([2, 99])
        wb.save(fb)
        return fa, fb

    def test_key_compare_ui_meta_and_bounded_model_text(self, pair: tuple[Path, Path]) -> None:
        fa, fb = pair
        tr = data_tools.compare_excel(
            str(fa), str(fb), key_columns=["id"], max_diffs=500,
        )
        assert isinstance(tr, ToolResult)
        assert tr.success
        assert tr.value is not None
        assert tr.value.get("alignment") == "key"
        assert tr.ui_meta.merge is not None
        assert tr.ui_meta.merge.get("key_columns") == ["id"]
        assert tr.ui_meta.diff is not None
        assert len(tr.model_text) < 4000
        assert "sample_diffs" not in tr.model_text


def _assert_bounded_probe(tr: ToolResult) -> None:
    assert isinstance(tr, ToolResult)
    assert tr.success
    assert tr.value is not None
    dumped = json.dumps(tr.value, ensure_ascii=False, default=str)
    assert tr.model_text != dumped
    assert len(tr.model_text) < 4000


class TestSearchExcelValuesToolResult:
    @pytest.fixture()
    def book(self, tmp_path: Path) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "name"])
        for i in range(1, 21):
            ws.append([i, f"user-{i}"])
        fp = tmp_path / "search.xlsx"
        wb.save(fp)
        return fp

    def test_bounded_model_text_and_ui_meta(self, book: Path) -> None:
        tr = data_tools.search_excel_values(str(book), query="user", max_results=5)
        _assert_bounded_probe(tr)
        assert tr.ui_meta.files == ["search.xlsx"]
        assert tr.ui_meta.preview is not None
        assert tr.ui_meta.content_version
        assert str(tr.ui_meta.content_version).startswith("sha256:")
        assert tr.value.get("content_version") == tr.ui_meta.content_version
        assert '"matches"' not in tr.model_text
        assert tr.coverage is not None
        assert tr.coverage.get("kind") in {"truncated", "complete"}

    def test_multi_file_reads_value(self, tmp_path: Path) -> None:
        paths: list[str] = []
        for name, needle in (("a.xlsx", "alpha"), ("b.xlsx", "beta")):
            wb = Workbook()
            ws = wb.active
            ws.append(["name"])
            ws.append([needle])
            fp = tmp_path / name
            wb.save(fp)
            paths.append(str(fp))
        tr = data_tools.search_excel_values(
            query="a", file_paths=paths, match_mode="contains",
        )
        _assert_bounded_probe(tr)
        assert tr.value.get("files_searched") == 2
        assert tr.ui_meta.files


class TestInspectExcelFilesToolResult:
    def test_directory_scan_ui_meta(self, tmp_path: Path) -> None:
        for name in ("a.xlsx", "b.xlsx"):
            wb = Workbook()
            ws = wb.active
            ws.append(["A", "B"])
            ws.append([1, 2])
            wb.save(tmp_path / name)
        tr = data_tools.inspect_excel_files(str(tmp_path), max_files=1)
        _assert_bounded_probe(tr)
        assert tr.value["excel_files_found"] == 1
        assert tr.truncated
        assert tr.coverage is not None
        assert tr.coverage.get("kind") == "truncated"
        assert tr.ui_meta.files
        assert tr.ui_meta.preview is not None
        assert '"preview"' not in tr.model_text
        assert '"file_list"' not in tr.model_text


class TestScanExcelSnapshotToolResult:
    def test_profile_ui_meta_and_bounded_text(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "val"])
        for i in range(1, 12):
            ws.append([i, i * 2])
        fp = tmp_path / "scan.xlsx"
        wb.save(fp)
        tr = data_tools.scan_excel_snapshot(str(fp), max_sample_rows=5)
        _assert_bounded_probe(tr)
        assert tr.ui_meta.files == ["scan.xlsx"]
        assert tr.ui_meta.content_version
        assert str(tr.ui_meta.content_version).startswith("sha256:")
        assert tr.value.get("content_version") == tr.ui_meta.content_version
        assert tr.coverage is not None
        assert tr.coverage.get("kind") in {"truncated", "complete", "sampled"}
        assert '"quality_signals"' not in tr.model_text


class TestFilterDataToolResult:
    def test_filter_preview_and_truncation(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["dept", "amt"])
        for i in range(30):
            ws.append(["sales", i])
        fp = tmp_path / "filter.xlsx"
        wb.save(fp)
        tr = data_tools.filter_data(
            str(fp), column="dept", operator="eq", value="sales", max_rows=5,
        )
        _assert_bounded_probe(tr)
        assert tr.value["filtered_rows"] == 30
        assert tr.value["returned_rows"] == 5
        assert tr.truncated
        assert tr.coverage is not None
        assert tr.coverage.get("kind") == "truncated"
        assert tr.ui_meta.files == ["filter.xlsx"]
        assert tr.ui_meta.preview is not None
        assert tr.ui_meta.content_version
        assert tr.model_text.count("sales") < 10


class TestDiscoverFileRelationshipsToolResult:
    def test_merge_ui_meta_and_bounded_text(self, tmp_path: Path) -> None:
        fa = tmp_path / "left.xlsx"
        fb = tmp_path / "right.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["客户ID", "金额"])
        ws.append(["C001", 10])
        ws.append(["C002", 20])
        ws.append(["C003", 30])
        wb.save(fa)
        wb = Workbook()
        ws = wb.active
        ws.append(["客户ID", "城市"])
        ws.append(["C001", "北京"])
        ws.append(["C002", "上海"])
        ws.append(["C004", "广州"])
        wb.save(fb)
        tr = data_tools.discover_file_relationships(file_paths=[str(fa), str(fb)])
        _assert_bounded_probe(tr)
        assert tr.value.get("file_pairs")
        assert tr.ui_meta.merge is not None
        assert tr.ui_meta.files
        assert '"file_pairs"' not in tr.model_text
        assert '"shared_columns"' not in tr.model_text
