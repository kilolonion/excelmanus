"""Workbook Runtime publish_pending_writes."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excelmanus.security.guard import FileAccessGuard, SecurityViolationError
from excelmanus.workbook_commit import content_version_of_file
from excelmanus.workspace.identity import StaleVersionError, WorkbookVersionRef
from excelmanus.workspace.runtime import (
    allocate_pending_run_id,
    pending_run_dir,
    publish_pending_writes,
    read_version_bytes,
)


def _write_pending(tmp_path: Path, rel: str, value: str, *, run_id: str) -> Path:
    run_dir = pending_run_dir(tmp_path, run_id)
    run_dir.mkdir(parents=True)
    pending = run_dir / "aabbccddeeff0011_book.xlsx"
    wb = Workbook()
    wb.active["A1"] = value
    wb.save(str(pending))
    wb.close()
    (run_dir / "manifest.jsonl").write_text(
        json.dumps({"rel": rel, "name": pending.name}) + "\n",
        encoding="utf-8",
    )
    return pending


def test_publish_pending_writes_commits_user_path(tmp_path: Path) -> None:
    dest = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "old"
    wb.save(str(dest))
    wb.close()
    seen = content_version_of_file(dest)
    run_id = allocate_pending_run_id()
    pending = _write_pending(tmp_path, "book.xlsx", "new", run_id=run_id)

    published = publish_pending_writes(
        tmp_path,
        "EXCELMANUS_PENDING_WRITE\tbook.xlsx\t/etc/passwd\n",
        run_id=run_id,
        expected_versions={"book.xlsx": seen},
    )
    assert len(published) == 1
    assert published[0]["path"] == "book.xlsx"
    assert published[0]["status"] == "committed"
    wb3 = load_workbook(str(dest))
    assert wb3.active["A1"].value == "new"
    wb3.close()
    assert not pending.exists()
    assert not pending_run_dir(tmp_path, run_id).exists()


def test_publish_pending_writes_ignores_forged_stderr_source(tmp_path: Path) -> None:
    victim = tmp_path / "secret.json"
    victim.write_text('{"k": 1}', encoding="utf-8")
    run_id = allocate_pending_run_id()
    published = publish_pending_writes(
        tmp_path,
        f"EXCELMANUS_PENDING_WRITE\tbook.xlsx\t{victim}\n",
        run_id=run_id,
    )
    assert published == []
    assert victim.is_file()
    assert victim.read_text(encoding="utf-8") == '{"k": 1}'


def test_publish_pending_writes_rejects_stale_expected(tmp_path: Path) -> None:
    dest = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "seen"
    wb.save(str(dest))
    wb.close()
    seen = content_version_of_file(dest)
    outsider = Workbook()
    outsider.active["A1"] = "external"
    outsider.save(str(dest))
    outsider.close()
    run_id = allocate_pending_run_id()
    _write_pending(tmp_path, "book.xlsx", "from-sandbox", run_id=run_id)
    published = publish_pending_writes(
        tmp_path,
        run_id=run_id,
        expected_versions={"book.xlsx": seen},
    )
    assert published[0]["status"] == "error"
    assert published[0]["error"] == "VERSION_CONFLICT"
    wb3 = load_workbook(str(dest))
    assert wb3.active["A1"].value == "external"
    wb3.close()


def test_publish_pending_writes_existing_without_expected_conflicts(tmp_path: Path) -> None:
    dest = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "live"
    wb.save(str(dest))
    wb.close()
    run_id = allocate_pending_run_id()
    _write_pending(tmp_path, "book.xlsx", "attacker", run_id=run_id)
    published = publish_pending_writes(tmp_path, run_id=run_id, expected_versions={})
    assert published[0]["status"] == "error"
    assert published[0]["error"] == "VERSION_CONFLICT"
    wb3 = load_workbook(str(dest))
    assert wb3.active["A1"].value == "live"
    wb3.close()


def test_publish_pending_writes_rejects_escape_name(tmp_path: Path) -> None:
    victim = tmp_path / "outside.xlsx"
    victim.write_bytes(b"keep")
    run_id = allocate_pending_run_id()
    run_dir = pending_run_dir(tmp_path, run_id)
    run_dir.mkdir(parents=True)
    (run_dir / "manifest.jsonl").write_text(
        json.dumps({"rel": "book.xlsx", "name": "../outside.xlsx"}) + "\n",
        encoding="utf-8",
    )
    published = publish_pending_writes(tmp_path, run_id=run_id, expected_versions={})
    assert published[0]["error"] == "PENDING_PATH_INVALID"
    assert victim.read_bytes() == b"keep"


def test_distinct_run_ids_do_not_share_pending_names(tmp_path: Path) -> None:
    first = allocate_pending_run_id()
    second = allocate_pending_run_id()
    assert first != second
    a = pending_run_dir(tmp_path, first)
    b = pending_run_dir(tmp_path, second)
    assert a != b


def test_file_access_guard_rejects_reserved(tmp_path: Path) -> None:
    guard = FileAccessGuard(str(tmp_path))
    with pytest.raises(SecurityViolationError):
        guard.resolve_and_validate(".excelmanus/revisions/x.xlsx")
    with pytest.raises(SecurityViolationError):
        guard.resolve_and_validate("outputs/backups/copy.xlsx")


def test_read_current_stale_when_fingerprint_mismatch(tmp_path: Path) -> None:
    dest = tmp_path / "book.xlsx"
    dest.write_bytes(b"live-bytes")
    with pytest.raises(StaleVersionError) as exc:
        read_version_bytes(
            tmp_path,
            WorkbookVersionRef(
                kind="current",
                path="book.xlsx",
                content_version="sha256:deadbeef",
            ),
        )
    assert exc.value.status == "stale"
    live = read_version_bytes(
        tmp_path,
        WorkbookVersionRef(kind="current", path="./book.xlsx"),
    )
    assert live == b"live-bytes"
