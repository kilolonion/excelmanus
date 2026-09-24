"""Offline checks of features supported by the separately bundled interpreter."""
import json
import os
from pathlib import Path
import tempfile
import subprocess
import sys

# A fresh isolated child also starts real spawn workers and checks package data.
subprocess.run([
    sys.executable, "-I", "-B", "-X", "utf8", str(Path(__file__).with_name("check-python-stdlib.py")),
], check=True, timeout=60)

with tempfile.TemporaryDirectory(prefix="excelmanus-runtime-") as directory:
    root = Path(directory)
    os.environ["MPLCONFIGDIR"] = str(root / "matplotlib")
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    import pandas as pd
    import yaml
    import xlrd
    import pyxlsb
    from docx import Document
    from oletools import olevba
    from PIL import Image

    frame = pd.DataFrame({"金额": [10, 20], "数量": [1, 2]})
    for engine in ("openpyxl", "xlsxwriter"):
        path = root / f"中文表格-{engine}.xlsx"
        frame.to_excel(path, index=False, engine=engine)
        pd.testing.assert_frame_equal(pd.read_excel(path), frame)
    document = Document()
    document.add_paragraph("中文文档")
    document.save(root / "文档.docx")
    assert Document(root / "文档.docx").paragraphs[0].text == "中文文档"
    np.testing.assert_allclose(np.array([1., 2.]) * 2, [2., 4.])
    plt.plot([1, 2, 3], [2, 4, 6])
    plt.savefig(root / "chart.png")
    plt.close("all")
    with Image.open(root / "chart.png") as image:
        assert image.width > 0
    assert yaml.safe_load("enabled: true")["enabled"] is True
    assert olevba.VBA_Parser and xlrd.open_workbook and pyxlsb.open_workbook

print("Bundled Python: XLSX read/write, DOCX, PNG, plotting and VBA imports OK")
