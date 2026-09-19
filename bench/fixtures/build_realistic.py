"""生成真实场景评测夹具（确定性随机，附带标准答案）。

与 build_experiential.py 的"脏办公文件"不同，这里模拟普通用户每天真正会丢给
表格助手的东西：几千行的订单导出、三个月的 CSV、两家门店的库存、带公式的模板、
一张 6 万行的流水、一份 Word 简报模板。数据大体干净，只保留真实世界里常见的
少量噪声（重复行、状态列、编码差异、日期写法不一）。

同目录下的 answers.json 记录由同一份随机数据推导出的标准值，suite 断言用
"@answers:key" 引用，避免手抄数字与夹具漂移。
"""
from __future__ import annotations

import csv
import json
import random
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FIXTURE_DIR = Path(__file__).resolve().parent / "realistic"
SEED = 20240914

REGIONS: dict[str, list[tuple[str, str]]] = {
    "华东": [("上海", "上海"), ("江苏", "南京"), ("江苏", "苏州"), ("浙江", "杭州"), ("浙江", "宁波"), ("山东", "青岛")],
    "华南": [("广东", "深圳"), ("广东", "广州"), ("广东", "东莞"), ("福建", "厦门"), ("海南", "海口")],
    "华北": [("北京", "北京"), ("天津", "天津"), ("河北", "石家庄"), ("山西", "太原")],
    "华中": [("湖北", "武汉"), ("湖南", "长沙"), ("河南", "郑州")],
    "西南": [("四川", "成都"), ("重庆", "重庆"), ("云南", "昆明")],
    "东北": [("辽宁", "沈阳"), ("辽宁", "大连"), ("黑龙江", "哈尔滨")],
}
REGION_WEIGHTS = {"华东": 32, "华南": 24, "华北": 18, "华中": 10, "西南": 9, "东北": 7}

# (编码, 名称, 类别, 成本, 标准单价)
PRODUCTS: list[tuple[str, str, str, int, int]] = [
    ("P001", "商务笔记本 14寸", "电脑", 3600, 4999),
    ("P002", "商务笔记本 15.6寸", "电脑", 4100, 5699),
    ("P003", "台式主机 标准版", "电脑", 2800, 3999),
    ("P004", "台式主机 高配版", "电脑", 4500, 6499),
    ("P005", "27寸显示器", "显示设备", 900, 1399),
    ("P006", "24寸显示器", "显示设备", 620, 899),
    ("P007", "32寸曲面显示器", "显示设备", 1500, 2299),
    ("P008", "机械键盘", "外设", 180, 349),
    ("P009", "无线鼠标", "外设", 45, 99),
    ("P010", "无线键鼠套装", "外设", 95, 199),
    ("P011", "USB-C 扩展坞", "外设", 160, 329),
    ("P012", "降噪耳机", "外设", 320, 699),
    ("P013", "激光打印机", "办公设备", 1100, 1699),
    ("P014", "彩色多功能一体机", "办公设备", 2400, 3599),
    ("P015", "高速扫描仪", "办公设备", 1800, 2799),
    ("P016", "会议摄像头", "办公设备", 480, 899),
    ("P017", "A4 复印纸(箱)", "耗材", 95, 139),
    ("P018", "硒鼓 标准", "耗材", 210, 399),
    ("P019", "墨盒 三色", "耗材", 60, 129),
    ("P020", "移动硬盘 2TB", "存储", 380, 599),
]
CATEGORY_OF = {code: cat for code, _, cat, _, _ in PRODUCTS}

SALESPEOPLE = ["王芳", "李强", "张敏", "刘洋", "陈静", "杨帆", "赵磊", "周婷", "吴昊", "郑丽", "孙浩", "马雪"]
SURNAMES = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜"
GIVEN = "伟芳娜敏静丽强磊军洋勇艳杰涛明超秀英华慧建国文平刚桂珍兰云"
COMPANY_PREFIX = [
    "华信", "远洋", "中科", "恒达", "众合", "启明", "博远", "天成", "锦泰", "宏图", "瑞丰", "蓝海", "新宇", "金桥", "永昌",
    "鼎盛", "安泰", "凌云", "汇通", "百川", "长风", "紫光", "前程", "同创", "环宇", "星辰", "卓越", "峻峰", "润泽", "嘉禾",
]
COMPANY_MID = ["科技", "贸易", "电子", "信息", "商贸", "实业", "网络", "智能", "咨询", "物流", "医疗", "建设"]
COMPANY_SUFFIX = ["有限公司", "股份有限公司", "有限责任公司"]
DEPARTMENTS = ["销售部", "市场部", "研发部", "财务部", "人力资源部", "运营部", "客服部", "采购部"]
CHANNELS = ["线上商城", "线下门店", "企业直销", "经销商"]
STORES = ["朝阳店", "海淀店", "浦东店", "南山店", "天河店"]
STORE_GOODS = [("拿铁", 28), ("美式", 22), ("卡布奇诺", 30), ("可颂", 18), ("三明治", 32), ("蛋糕切块", 26), ("矿泉水", 5), ("果汁", 24)]


def _thin() -> Border:
    side = Side(style="thin", color="BFBFBF")
    return Border(top=side, bottom=side, left=side, right=side)


def _style_header(ws, row: int, ncols: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws, widths: dict[str, int] | None = None, default: int = 12) -> None:
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = (widths or {}).get(letter, default)


def _pick_region(rng: random.Random) -> str:
    return rng.choices(list(REGION_WEIGHTS), weights=list(REGION_WEIGHTS.values()))[0]


def _person_name(rng: random.Random) -> str:
    return rng.choice(SURNAMES) + "".join(rng.choice(GIVEN) for _ in range(rng.choice([1, 2])))


def _company_names(rng: random.Random, n: int) -> list[str]:
    """n 个互不重复的公司名（全组合 30×12×3=1080，够用）。"""
    pool = [p + m + s for p in COMPANY_PREFIX for m in COMPANY_MID for s in COMPANY_SUFFIX]
    if n > len(pool):
        raise ValueError(f"公司名池只有 {len(pool)} 个，要 {n} 个")
    return rng.sample(pool, n)


# ── 1. 订单明细_2024.xlsx ────────────────────────────────────


def build_orders(path: Path, rng: random.Random, answers: dict) -> Path:
    """2400 笔订单 + 产品目录 + 区域目标。含 6 行完全重复、状态列（已完成/已取消/退货）。"""
    customers = _company_names(rng, 150)
    rows: list[list] = []
    start = date(2024, 1, 1)
    for i in range(2400):
        region = _pick_region(rng)
        province, city = rng.choice(REGIONS[region])
        code, name, _cat, _cost, std_price = rng.choice(PRODUCTS)
        qty = rng.choice([1, 1, 2, 2, 3, 5, 5, 10, 10, 20, 50])
        price = round(std_price * rng.choice([0.9, 0.95, 1.0, 1.0, 1.0]))
        d = start + timedelta(days=rng.randint(0, 273))  # 2024-01-01 .. 2024-09-30
        status = rng.choices(["已完成", "已取消", "退货"], weights=[92, 5, 3])[0]
        rows.append([
            f"SO2024{i + 1:05d}", datetime(d.year, d.month, d.day), region, province, city,
            rng.choice(SALESPEOPLE), rng.choice(customers), code, name, qty, price, qty * price, status,
        ])
    rows.sort(key=lambda r: r[1])
    # 6 行完全重复（导出时常见）
    dup_indices = rng.sample(range(len(rows)), 6)
    for idx in sorted(dup_indices):
        rows.insert(idx + 1, list(rows[idx]))
    dup_order_ids = sorted({rows[i][0] for i in range(len(rows)) if sum(1 for r in rows if r[0] == rows[i][0]) > 1})

    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    headers = ["订单号", "日期", "区域", "省份", "城市", "销售员", "客户", "产品编码", "产品名称", "数量", "单价", "金额", "状态"]
    # 系统导出原样：无表头样式、无冻结、默认列宽、日期默认格式（R13 要它去整理）
    ws.append(headers)
    for r in rows:
        ws.append(r)

    cat = wb.create_sheet("产品目录")
    cat.append(["产品编码", "产品名称", "类别", "成本", "标准单价"])
    _style_header(cat, 1, 5)
    for p in PRODUCTS:
        cat.append(list(p))
    _autosize(cat, {"B": 22}, 12)

    tgt = wb.create_sheet("区域目标")
    tgt.append(["区域", "月份", "目标金额"])
    _style_header(tgt, 1, 3)
    targets: dict[str, dict[int, int]] = {}
    for region in REGIONS:
        targets[region] = {}
        for month in range(1, 10):
            # 目标与实际同量级（实际月均约 权重×4.5 万），完成率落在 60%–140% 之间才像真的
            base = REGION_WEIGHTS[region] * 45000
            targets[region][month] = int(round(base * rng.uniform(0.85, 1.2), -3))
            tgt.append([region, month, targets[region][month]])
    _autosize(tgt, default=12)
    wb.save(path)

    # 标准答案
    def _sum(pred) -> int:
        return sum(r[11] for r in rows if pred(r))

    completed = lambda r: r[12] == "已完成"  # noqa: E731
    east_aug = lambda r: r[2] == "华东" and r[1].month == 8  # noqa: E731
    region_month: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    region_month_all: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    per_sales: dict[str, int] = defaultdict(int)
    per_cat: dict[str, int] = defaultdict(int)
    for r in rows:
        region_month_all[r[2]][str(r[1].month)] += r[11]
        if completed(r):
            region_month[r[2]][str(r[1].month)] += r[11]
            per_sales[r[5]] += r[11]
            per_cat[CATEGORY_OF[r[7]]] += r[11]
    ranked = sorted(per_sales.items(), key=lambda kv: -kv[1])
    total_completed = _sum(completed)
    per_province: dict[str, int] = defaultdict(int)
    for r in rows:
        if completed(r):
            per_province[r[3]] += r[11]
    ranked_prov = sorted(per_province.items(), key=lambda kv: -kv[1])
    answers.update({
        "orders_rows_with_dups": len(rows),
        "orders_dup_order_ids": dup_order_ids,
        "orders_dup_row_count": len(rows) - len({r[0] for r in rows}),
        "orders_total_all_status": _sum(lambda r: True),
        "orders_total_completed": total_completed,
        "orders_east_aug_all_status": _sum(east_aug),
        "orders_east_aug_completed": _sum(lambda r: east_aug(r) and completed(r)),
        # wave-r5 扩展用例：以下推导只读已生成的 rows，不消耗 rng，
        # 因此不改变既有答案的值（同 seed 下可复现）。
        "orders_east_aug_highprice_completed": _sum(lambda r: east_aug(r) and completed(r) and r[10] >= 1000),
        "orders_province_totals_completed": dict(ranked_prov),
        "orders_top3_provinces": [p for p, _ in ranked_prov[:3]],
        "orders_region_month_completed": {k: dict(v) for k, v in region_month.items()},
        "orders_region_month_all_status": {k: dict(v) for k, v in region_month_all.items()},
        "orders_region_aug_completed": {k: v.get("8", 0) for k, v in region_month.items()},
        "orders_top_salesperson": ranked[0][0],
        "orders_top5_salespeople": [name for name, _ in ranked[:5]],
        "orders_salesperson_totals_completed": dict(ranked),
        "orders_category_totals_completed": dict(per_cat),
        "orders_top_category": max(per_cat.items(), key=lambda kv: kv[1])[0],
        "orders_top_category_share_pct": round(max(per_cat.values()) / total_completed * 100, 1),
        "orders_targets_aug": {k: v[8] for k, v in targets.items()},
        "orders_top3_cost_products": [p[1] for p in sorted(PRODUCTS, key=lambda p: -p[3])[:3]],
    })
    return path


# ── 2. 客户名单_脏.xlsx ──────────────────────────────────────


def _messy_phone(rng: random.Random, digits: str) -> str:
    style = rng.choice(["plain", "plain", "dash", "space", "plus86"])
    if style == "dash":
        return f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"
    if style == "space":
        return f"{digits[:3]} {digits[3:7]} {digits[7:]}"
    if style == "plus86":
        return f"+86 {digits}"
    return digits


def _messy_date(rng: random.Random, d: date):
    style = rng.choice(["iso", "slash", "dot", "cn", "compact", "datetime"])
    if style == "iso":
        return d.isoformat()
    if style == "slash":
        return f"{d.year}/{d.month}/{d.day}"
    if style == "dot":
        return f"{d.year}.{d.month}.{d.day}"
    if style == "cn":
        return f"{d.year}年{d.month}月{d.day}日"
    if style == "compact":
        return f"{d.year}{d.month:02d}{d.day:02d}"
    return datetime(d.year, d.month, d.day)


def build_customers(path: Path, rng: random.Random, answers: dict) -> Path:
    """600 行客户名单：手机写法混杂、日期写法混杂、按手机号有 80 条重复。"""
    base: list[dict] = []
    names = _company_names(rng, 520)
    for idx in range(520):
        digits = "1" + rng.choice("3578") + "".join(rng.choice("0123456789") for _ in range(9))
        base.append({
            "name": names[idx],
            "contact": _person_name(rng),
            "phone": digits,
            "reg": date(2021, 1, 1) + timedelta(days=rng.randint(0, 1300)),
            "city": rng.choice([c for pairs in REGIONS.values() for _, c in pairs]),
        })
    rows: list[list] = []
    for c in base:
        rows.append([c["name"], c["contact"], _messy_phone(rng, c["phone"]), _messy_date(rng, c["reg"]), c["city"], rng.choice(["", "", "", "老客户", "待跟进"])])
    dup_sources = rng.sample(base, 80)
    for c in dup_sources:
        name = c["name"]
        variant = rng.choice(["space", "strip_suffix", "same", "short_suffix"])
        if variant == "space":
            name = name + " "
        elif variant == "strip_suffix":
            for suffix in COMPANY_SUFFIX:
                name = name.removesuffix(suffix)
        elif variant == "short_suffix":
            name = name.replace("股份有限公司", "有限公司").replace("有限责任公司", "有限公司")
        later = c["reg"] + timedelta(days=rng.randint(30, 400))
        rows.append([name, c["contact"], _messy_phone(rng, c["phone"]), _messy_date(rng, later), c["city"], rng.choice(["", "", "老客户", "待跟进"])])
    rng.shuffle(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "客户"
    headers = ["客户名称", "联系人", "手机", "注册日期", "城市", "备注"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for r in rows:
        ws.append(r)
    _autosize(ws, {"A": 28, "C": 18, "D": 14}, 12)
    wb.save(path)
    answers.update({
        "customers_total_rows": len(rows),
        "customers_unique_by_phone": len(base),
        "customers_duplicate_rows": len(rows) - len(base),
    })
    return path


# ── 3. 员工信息.xlsx ─────────────────────────────────────────


def build_employees(path: Path, rng: random.Random, answers: dict) -> Path:
    """320 名员工 + 一张「姓名-工号」挤在一列的导入原始表。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "员工"
    headers = ["工号", "姓名", "部门", "职级", "入职日期", "基本工资", "绩效系数", "手机"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    emps: list[list] = []
    for i in range(320):
        dept = rng.choice(DEPARTMENTS)
        level = rng.choice(["P4", "P5", "P5", "P6", "P6", "P7", "P8"])
        base_salary = {"P4": 8000, "P5": 11000, "P6": 15000, "P7": 21000, "P8": 30000}[level] + rng.randint(0, 20) * 250
        hire = date(2015, 1, 1) + timedelta(days=rng.randint(0, 3500))
        emps.append([
            f"E{1001 + i}", _person_name(rng), dept, level, datetime(hire.year, hire.month, hire.day),
            base_salary, rng.choice([0.8, 0.9, 1.0, 1.0, 1.1, 1.2]),
            "1" + rng.choice("3578") + "".join(rng.choice("0123456789") for _ in range(9)),
        ])
    for e in emps:
        ws.append(e)
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "yyyy-mm-dd"
    _autosize(ws, {"E": 12, "H": 14}, 11)

    raw = wb.create_sheet("导入原始")
    raw.append(["姓名-工号", "部门"])
    _style_header(raw, 1, 2)
    for e in emps[:100]:
        raw.append([f"{e[1]}-{e[0]}", e[2]])
    _autosize(raw, {"A": 18}, 12)
    wb.save(path)

    by_dept: dict[str, list[int]] = defaultdict(list)
    for e in emps:
        by_dept[e[2]].append(e[5])
    cutoff = datetime(2024, 9, 30)
    sales_p6plus = [e[5] for e in emps if e[2] == "销售部" and e[3] in ("P6", "P7", "P8")]
    answers.update({
        "employees_count": len(emps),
        "employees_dept_avg_salary": {k: round(sum(v) / len(v), 2) for k, v in by_dept.items()},
        "employees_sales_dept_avg_salary": round(sum(by_dept["销售部"]) / len(by_dept["销售部"]), 2),
        "employees_sales_p6plus_avg_salary": round(sum(sales_p6plus) / len(sales_p6plus), 2) if sales_p6plus else 0,
        "employees_5yr_plus_count": sum(1 for e in emps if (cutoff - e[4]).days >= 365 * 5),
        "employees_import_rows": 100,
        "employees_import_first": {"name": emps[0][1], "id": emps[0][0]},
        "employees_import_pairs": [
            [e[1], e[0]] for e in (emps[0], emps[50], emps[99])
        ],
    })
    return path


# ── 4. 三个月门店 CSV（编码/列序不一致）────────────────────


def build_store_csvs(dir_path: Path, rng: random.Random, answers: dict) -> list[Path]:
    counts = {1: 380, 2: 410, 3: 395}
    written: list[Path] = []
    total_amount = 0
    per_month_rows: dict[str, int] = {}
    for month, n in counts.items():
        rows = []
        for i in range(n):
            d = date(2024, month, rng.randint(1, 28))
            store = rng.choice(STORES)
            goods, price = rng.choice(STORE_GOODS)
            qty = rng.randint(1, 6)
            rows.append([d.isoformat(), f"T{month:02d}{i + 1:04d}", store, goods, qty, price, qty * price])
            total_amount += qty * price
        per_month_rows[f"{month}月"] = n
        path = dir_path / f"{month}月销售.csv"
        if month == 1:
            with path.open("w", encoding="utf-8-sig", newline="") as fh:
                w = csv.writer(fh)
                w.writerow(["日期", "订单号", "门店", "商品", "数量", "单价", "金额"])
                w.writerows(rows)
        elif month == 2:
            # 财务系统导出：GBK 编码、列序不同
            with path.open("w", encoding="gbk", newline="") as fh:
                w = csv.writer(fh)
                w.writerow(["订单号", "日期", "门店", "商品", "单价", "数量", "金额"])
                for r in rows:
                    w.writerow([r[1], r[0], r[2], r[3], r[5], r[4], r[6]])
        else:
            with path.open("w", encoding="utf-8", newline="") as fh:
                w = csv.writer(fh)
                w.writerow(["日期", "订单号", "门店", "商品", "数量", "单价", "金额", "备注"])
                for r in rows:
                    w.writerow(r + [rng.choice(["", "", "", "会员", "外卖"])])
        written.append(path)
    answers.update({
        "csv_total_rows": sum(counts.values()),
        "csv_rows_per_month": per_month_rows,
        "csv_total_amount": total_amount,
    })
    return written


# ── 5. 两家门店库存 ──────────────────────────────────────────


def build_inventory_pair(dir_path: Path, rng: random.Random, answers: dict) -> list[Path]:
    goods = ["鼠标", "键盘", "耳机", "音箱", "支架", "线材", "底座", "包袋", "贴膜", "充电器"]
    skus = [f"SKU-{i:04d}" for i in range(1, 271)]
    a_skus = skus[:220]
    b_skus = skus[70:270]

    def _write(path: Path, store: str, keep: list[str]) -> dict[str, int]:
        wb = Workbook()
        ws = wb.active
        ws.title = "库存"
        ws.append(["SKU", "品名", "数量", "单价", "门店"])
        _style_header(ws, 1, 5)
        qty_of: dict[str, int] = {}
        for sku in keep:
            idx = int(sku.split("-")[1])
            qty = rng.randint(0, 120)
            qty_of[sku] = qty
            ws.append([sku, f"{goods[idx % len(goods)]}-{idx:03d}", qty, 19 + (idx * 7) % 180, store])
        _autosize(ws, {"B": 14}, 12)
        wb.save(path)
        return qty_of

    qa = _write(dir_path / "库存_门店A.xlsx", "门店A", a_skus)
    qb = _write(dir_path / "库存_门店B.xlsx", "门店B", b_skus)
    answers.update({
        "inventory_union_sku_count": len(set(a_skus) | set(b_skus)),
        "inventory_only_a_count": len(set(a_skus) - set(b_skus)),
        "inventory_only_b_count": len(set(b_skus) - set(a_skus)),
        "inventory_sku_0100_total": qa.get("SKU-0100", 0) + qb.get("SKU-0100", 0),
        "inventory_sku_0100_a": qa.get("SKU-0100", 0),
        "inventory_sku_0100_b": qb.get("SKU-0100", 0),
    })
    return [dir_path / "库存_门店A.xlsx", dir_path / "库存_门店B.xlsx"]


# ── 6. 月报模板.xlsx ─────────────────────────────────────────


def build_monthly_template(path: Path, answers: dict) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "月报"
    ws.merge_cells("A1:F1")
    ws["A1"] = "华夏办公 2024 年 8 月销售月报"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "单位：元"
    headers = ["区域", "目标", "实际", "完成率", "差额", "备注"]
    for col, h in enumerate(headers, start=1):
        ws.cell(3, col, h)
    _style_header(ws, 3, len(headers))
    targets = answers["orders_targets_aug"]
    for i, region in enumerate(REGIONS, start=4):
        ws.cell(i, 1, region)
        ws.cell(i, 2, targets[region])
        ws.cell(i, 4, f'=IF(B{i}=0,"",C{i}/B{i})').number_format = "0.0%"
        ws.cell(i, 5, f"=C{i}-B{i}")
        for col in range(1, 7):
            ws.cell(i, col).border = _thin()
    last = 3 + len(REGIONS)
    ws.cell(last + 1, 1, "合计").font = Font(bold=True)
    ws.cell(last + 1, 2, f"=SUM(B4:B{last})")
    ws.cell(last + 1, 3, f"=SUM(C4:C{last})")
    ws.cell(last + 1, 4, f'=IF(B{last + 1}=0,"",C{last + 1}/B{last + 1})').number_format = "0.0%"
    ws.cell(last + 1, 5, f"=C{last + 1}-B{last + 1}")
    for col in range(1, 7):
        ws.cell(last + 1, col).border = _thin()
        ws.cell(last + 1, col).font = Font(bold=True)
    for r in range(4, last + 2):
        for c in (2, 3, 5):
            ws.cell(r, c).number_format = "#,##0"
    ws.cell(last + 3, 1, "编制：").font = Font(italic=True, color="808080")
    _autosize(ws, {"A": 12, "F": 24}, 14)
    ws.print_area = f"A1:F{last + 3}"
    wb.save(path)
    answers["template_actual_cells"] = {region: f"C{i}" for i, region in enumerate(REGIONS, start=4)}
    answers["template_total_row"] = last + 1
    return path


# ── 7. 大表_交易流水.xlsx（6 万行）─────────────────────────


def build_big_ledger(path: Path, rng: random.Random, answers: dict, n_rows: int = 60_000) -> Path:
    customers = list(zip((f"C{i:04d}" for i in range(1, 301)), _company_names(rng, 300)))
    weights = [rng.uniform(0.3, 3.0) for _ in customers]
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("流水")
    ws.append(["交易ID", "日期", "客户ID", "客户名称", "渠道", "商品类别", "数量", "单价", "金额", "状态"])
    per_customer: dict[str, int] = defaultdict(int)
    name_of = dict(customers)
    total_june = 0
    start = date(2024, 1, 1)
    for i in range(n_rows):
        cid, cname = rng.choices(customers, weights=weights)[0]
        d = start + timedelta(days=rng.randint(0, 273))
        cat = rng.choice(list({p[2] for p in PRODUCTS}))
        qty = rng.randint(1, 30)
        price = rng.choice([99, 129, 199, 349, 599, 899, 1399, 1699, 2299, 3999, 4999])
        amount = qty * price
        status = rng.choices(["成功", "失败"], weights=[97, 3])[0]
        ws.append([f"TX{i + 1:07d}", d.isoformat(), cid, cname, rng.choice(CHANNELS), cat, qty, price, amount, status])
        if status == "成功":
            per_customer[cid] += amount
            if d.month == 6:
                total_june += amount
    wb.save(path)
    ranked = sorted(per_customer.items(), key=lambda kv: -kv[1])
    answers.update({
        "ledger_rows": n_rows,
        "ledger_top20_customers": [[cid, name_of[cid], amt] for cid, amt in ranked[:20]],
        "ledger_top1_customer_name": name_of[ranked[0][0]],
        "ledger_top1_customer_amount": ranked[0][1],
        "ledger_total_2024_06_success": total_june,
    })
    return path


# ── 8. 预算_有公式.xlsx（跨表公式 + 常见错误）──────────────


def build_budget(path: Path, rng: random.Random, answers: dict) -> Path:
    wb = Workbook()
    params = wb.active
    params.title = "参数"
    params["A1"], params["B1"] = "年度增长率", 0.08
    params["A2"], params["B2"] = "人均月成本", 18000
    params["A3"], params["B3"] = "季度调整系数", 1.05
    _autosize(params, {"A": 16}, 12)

    ws = wb.create_sheet("部门预算")
    months = [f"{m}月" for m in range(1, 13)]
    ws.append(["部门", "人数"] + months + ["全年合计"])
    _style_header(ws, 1, 15)
    problems: list[dict] = []
    for i, dept in enumerate(DEPARTMENTS, start=2):
        headcount = rng.randint(6, 40)
        ws.cell(i, 1, dept)
        ws.cell(i, 2, headcount)
        for m in range(12):
            col = 3 + m
            ws.cell(i, col, f"=$B{i}*参数!$B$2*(1+参数!$B$1)")
        ws.cell(i, 15, f"=SUM(C{i}:N{i})")
    # 常见错误：合计范围漏两个月
    ws.cell(4, 15, "=SUM(C4:L4)")
    problems.append({"cell": "部门预算!O4", "issue": "合计范围只到 L 列，漏了 11、12 月"})
    # 常见错误：公式列中间被手填了硬编码数字
    ws.cell(6, 8, 250000)
    problems.append({"cell": "部门预算!H6", "issue": "公式列被硬编码数字覆盖"})
    # 常见错误：引用了不存在的参数行
    ws.cell(7, 5, "=$B7*参数!$B$9*(1+参数!$B$1)")
    problems.append({"cell": "部门预算!E7", "issue": "引用 参数!B9 为空，结果为 0"})
    # 常见错误：除零
    ws.cell(11, 1, "人均预算")
    ws.cell(11, 2, "=O2/B12")
    problems.append({"cell": "部门预算!B11", "issue": "除以空单元格 B12，#DIV/0!"})
    ws.cell(12, 1, "增长后合计")
    ws.cell(12, 3, "=SUM(O2:O9)*参数!B3")
    _autosize(ws, {"A": 14}, 11)
    wb.save(path)
    answers["budget_problem_cells"] = problems
    return path


# ── 9. 报告模板.docx ─────────────────────────────────────────


def build_report_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("{{月份}}月销售简报", level=1)
    doc.add_paragraph("本月公司整体销售额为 {{总销售额}} 元，其中表现最好的区域是 {{最佳区域}}，销售额 {{最佳区域销售额}} 元。")
    doc.add_paragraph("各区域销售情况如下：")
    table = doc.add_table(rows=1 + len(REGIONS), cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "区域"
    table.rows[0].cells[1].text = "销售额（元）"
    for i, region in enumerate(REGIONS, start=1):
        table.rows[i].cells[0].text = region
        table.rows[i].cells[1].text = ""
    doc.add_paragraph("备注：{{备注}}")
    doc.save(path)
    return path


# ── 入口 ─────────────────────────────────────────────────────


def expected_filenames() -> frozenset[str]:
    return frozenset({
        "订单明细_2024.xlsx", "客户名单_脏.xlsx", "员工信息.xlsx",
        "1月销售.csv", "2月销售.csv", "3月销售.csv",
        "库存_门店A.xlsx", "库存_门店B.xlsx", "月报模板.xlsx",
        "大表_交易流水.xlsx", "预算_有公式.xlsx", "报告模板.docx", "answers.json",
    })


def build_all(output_dir: Path | None = None, *, big_rows: int = 60_000, seed: int = SEED) -> list[Path]:
    """生成全部夹具。

    ``seed`` 用于 nightly 隐藏种子轮换：CI 用公开默认 SEED，nightly 传别的
    seed 生成答案不同的同构夹具，防"调参过拟合阈值/背答案"。注意换 seed 后
    必须用新 answers.json 评分（suite 的 answers_file 指向同目录）。
    """
    target = Path(output_dir) if output_dir is not None else FIXTURE_DIR
    target.mkdir(parents=True, exist_ok=True)
    rng = random.Random(seed)
    answers: dict = {"seed": seed, "generated_at": datetime.now().isoformat(timespec="seconds")}
    written = [
        build_orders(target / "订单明细_2024.xlsx", rng, answers),
        build_customers(target / "客户名单_脏.xlsx", rng, answers),
        build_employees(target / "员工信息.xlsx", rng, answers),
        *build_store_csvs(target, rng, answers),
        *build_inventory_pair(target, rng, answers),
        build_monthly_template(target / "月报模板.xlsx", answers),
        build_big_ledger(target / "大表_交易流水.xlsx", rng, answers, n_rows=big_rows),
        build_budget(target / "预算_有公式.xlsx", rng, answers),
        build_report_docx(target / "报告模板.docx"),
    ]
    answers_path = target / "answers.json"
    answers_path.write_text(json.dumps(answers, ensure_ascii=False, indent=2), encoding="utf-8")
    written.append(answers_path)
    missing = expected_filenames() - {p.name for p in written}
    if missing:
        raise RuntimeError(f"夹具漏生成: {sorted(missing)}")
    return written


if __name__ == "__main__":
    import sys

    rows = int(sys.argv[1]) if len(sys.argv) > 1 else 60_000
    seed = int(sys.argv[2]) if len(sys.argv) > 2 else SEED
    for p in build_all(big_rows=rows, seed=seed):
        print(p)
