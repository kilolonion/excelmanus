# -*- coding: utf-8 -*-
"""suite_prompt_contract 的 script 检查。

期望值在检查运行时从夹具 / 产出实时计算，不预置进 answers.json：

- ``s07_moving_avg``：订单明细「已完成」按月汇总金额，最后一个数据月的
  3 个月移动平均，核对回复数字；
- ``s10_spec``：产出工作簿含合并标题、加粗且带边框的表头、状态列下拉
  （list 数据验证）、非默认列宽；
- ``s14_full_sum``：大表「成功」交易金额全量合计，核对回复数字，防止
  把截断/采样读取当成全表事实。

模块经 ``bench_checks._c_script`` 以 ``spec_from_file_location`` 加载，
``check(workdir, result, answers, spec)`` 返回 ``list[dict]``。
"""

from __future__ import annotations

import datetime
from collections import defaultdict
from pathlib import Path
from typing import Any

from excelmanus.bench_checks import extract_numbers, locate_output_workbook

_FIXTURES = Path(__file__).resolve().parent / "fixtures" / "realistic"


def _reply_text(result: dict[str, Any]) -> str:
    return str((result.get("result") or {}).get("reply") or "")


def _near(actual: float, expected: float, rel: float = 0.005) -> bool:
    return abs(actual - expected) <= max(1.0, abs(expected) * rel)


def _sum_orders_by_month(path: Path) -> dict[int, float]:
    from openpyxl import load_workbook

    totals: dict[int, float] = defaultdict(float)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["订单"]
        header: list[str] | None = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                header = [str(c) for c in row]
                continue
            rec = dict(zip(header or (), row))
            if rec.get("状态") == "已完成" and isinstance(rec.get("日期"), datetime.datetime):
                totals[rec["日期"].month] += float(rec.get("金额") or 0)
    finally:
        wb.close()
    return totals


def _s07(workdir: Path, result: dict[str, Any], spec: dict[str, Any]) -> list[dict[str, Any]]:
    totals = _sum_orders_by_month(_FIXTURES / "订单明细_2024.xlsx")
    last = max(totals)
    expected = sum(totals[m] for m in range(last - 2, last + 1)) / 3
    nums = extract_numbers(_reply_text(result))
    ok = any(_near(n, expected, float(spec.get("rel_tolerance", 0.005))) for n in nums)
    return [{
        "name": "s07_moving_avg",
        "passed": ok,
        "message": f"{last}月 MA3≈{expected:,.0f}，回复{'含' if ok else '不含'}该数字",
        "expected": round(expected, 2),
        "actual": sorted(set(nums))[-10:],
    }]


def _s10(workdir: Path, result: dict[str, Any], spec: dict[str, Any]) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    path = locate_output_workbook(workdir, spec.get("file"))
    if path is None:
        return [{"name": "s10_spec", "passed": False, "message": "未找到产出工作簿"}]
    wb = load_workbook(path)
    merged: list[str] = []
    header_styled = False
    list_validation = False
    wide_column = False
    try:
        for ws in wb.worksheets:
            merged.extend(str(r) for r in ws.merged_cells.ranges)
            for row in ws.iter_rows(min_row=1, max_row=6):
                styled = 0
                for cell in row:
                    if cell.value in (None, ""):
                        continue
                    bold = bool(cell.font and cell.font.bold)
                    border = cell.border
                    edged = bool(border and any((
                        border.left.style, border.right.style,
                        border.top.style, border.bottom.style,
                    )))
                    if bold and edged:
                        styled += 1
                if styled >= 4:
                    header_styled = True
            for dv in ws.data_validations.dataValidation:
                if (dv.type or "") == "list" and "未开始" in str(dv.formula1 or ""):
                    list_validation = True
            if ws.sheet_format.defaultColWidth and ws.sheet_format.defaultColWidth >= 12:
                wide_column = True
            for dim in ws.column_dimensions.values():
                if dim.width and dim.width >= 12:
                    wide_column = True
    finally:
        wb.close()
    return [
        {"name": "s10_merged_title", "passed": bool(merged),
         "message": f"合并区域 {merged[:4] or '无'}", "actual": merged[:8]},
        {"name": "s10_header_style", "passed": header_styled,
         "message": "表头加粗+边框" if header_styled else "前 6 行内未见 ≥4 格加粗且带边框的表头"},
        {"name": "s10_status_validation", "passed": list_validation,
         "message": "状态列 list 数据验证" if list_validation else "未见含「未开始」的 list 数据验证"},
        {"name": "s10_column_width", "passed": wide_column,
         "message": "存在 ≥12 的列宽" if wide_column else "未见显式列宽"},
    ]


def _s14(workdir: Path, result: dict[str, Any], spec: dict[str, Any]) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    src = workdir / "uploads" / "大表_交易流水.xlsx"
    if not src.exists():
        src = _FIXTURES / "大表_交易流水.xlsx"
    total = 0.0
    rows_ok = 0
    rows_all = 0
    wb = load_workbook(src, read_only=True, data_only=True)
    try:
        ws = wb["流水"]
        header: list[str] | None = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                header = [str(c) for c in row]
                continue
            rows_all += 1
            rec = dict(zip(header or (), row))
            if rec.get("状态") == "成功":
                total += float(rec.get("金额") or 0)
                rows_ok += 1
    finally:
        wb.close()
    nums = extract_numbers(_reply_text(result))
    ok = any(_near(n, total, float(spec.get("rel_tolerance", 0.005))) for n in nums)
    return [{
        "name": "s14_full_sum",
        "passed": ok,
        "message": (
            f"全表 {rows_all} 行 / 成功 {rows_ok} 行，合计 {total:,.0f}；"
            f"回复{'含' if ok else '不含'}该数字"
        ),
        "expected": round(total, 2),
        "actual": sorted(set(nums))[-10:],
    }]


_SCENARIOS = {
    "s07_moving_avg": _s07,
    "s10_spec": _s10,
    "s14_full_sum": _s14,
}


def check(workdir, result, answers, spec) -> list[dict[str, Any]]:
    scenario = str(spec.get("scenario") or "")
    fn = _SCENARIOS.get(scenario)
    if fn is None:
        return [{"name": "scenario", "passed": False,
                 "message": f"未知 scenario {scenario!r}，可用 {sorted(_SCENARIOS)}"}]
    try:
        return fn(Path(workdir), result, spec)
    except Exception as exc:  # 检查自身故障要可见，不静默判通过
        return [{"name": scenario or "script", "passed": False,
                 "message": f"检查执行异常: {exc!r}"}]
