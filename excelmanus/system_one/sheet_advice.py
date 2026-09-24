"""Bounded column candidates and read suggestions; never read workbook bytes here."""
from __future__ import annotations

import re
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter, range_boundaries

from excelmanus.workspace.identity import IdentityError, resolve_canonical


def cache_stamp(root: Any, path: str) -> str:
    try:
        ident = resolve_canonical(root, path)
        stat = (Path(root) / ident.relative).stat()
        return f"{stat.st_mtime_ns}:{stat.st_size}"
    except (IdentityError, ValueError, TypeError, OSError):
        return ""


def column_candidates(root: Any, targets: list[dict], entries: list[Any], user_text: str) -> tuple[list[dict], bool]:
    """Only use positional metadata from an unchanged cached scan.

    Legacy `headers` remove empty cells and cannot establish column positions.
    The stamp detects normal edits; the eventual tool read owns content_version.
    """
    paths = {row["path"] for row in targets}
    candidates: list[dict] = []
    seen: set[tuple[str, str, int]] = set()
    for entry in entries[:20]:
        try:
            path = resolve_canonical(root, entry.canonical_path).public
        except (IdentityError, ValueError, TypeError, AttributeError, OSError):
            continue
        if path not in paths:
            continue
        stamp = cache_stamp(root, path)
        if not stamp:
            continue
        for sheet in (getattr(entry, "sheet_meta", []) or [])[:10]:
            if not isinstance(sheet, Mapping) or sheet.get("cache_stamp") != stamp:
                continue
            name, row = sheet.get("name"), sheet.get("header_row")
            names = sheet.get("column_names")
            if not isinstance(name, str) or not name or len(name) > 100:
                continue
            if type(row) is not int or not 1 <= row <= 1_048_576 or not isinstance(names, list):
                continue
            related = [t for t in targets if t["path"] == path and (not t.get("sheet") or t["sheet"] == name)]
            if not related:
                continue
            for index, header in enumerate(names[:100], 1):
                if not isinstance(header, str) or not header.strip() or len(header) > 120:
                    continue
                key = path, name, index
                if key in seen:
                    continue
                seen.add(key)
                candidates.append({"path": path, "sheet": name, "header": header,
                                   "column": get_column_letter(index), "header_row": row,
                                   "cache_stamp": stamp, "source": "cached_header"})
    candidates.sort(key=lambda c: (c["header"] not in user_text,
                                   next(i for i, t in enumerate(targets) if t["path"] == c["path"])))
    return candidates[:10], len(candidates) > 10


def bounded_range(value: str) -> str | None:
    # Sheet names are carried separately; reject multi-area, external and entire-column references.
    if not re.fullmatch(r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6})?", value):
        return None
    try:
        left, top, right, bottom = range_boundaries(value.upper())
        if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1_048_576):
            return None
        return f"{get_column_letter(left)}{top}:{get_column_letter(min(right, left + 11))}{min(bottom, top + 29)}"
    except (ValueError, TypeError):
        return None


def column_matches_target(target: Mapping[str, Any] | None, column: Mapping[str, Any] | None) -> bool:
    """Candidate identities must agree before combining their coordinates."""
    return bool(
        target and column and target.get("path") == column.get("path")
        and (not target.get("sheet") or target.get("sheet") == column.get("sheet"))
    )


def read_suggestion(target: Mapping[str, Any] | None, column: Mapping[str, Any] | None,
                    strategy: str) -> dict[str, Any] | None:
    if not target or strategy == "none":
        return None
    if column and not column_matches_target(target, column):
        column, strategy = None, "overview"
    args: dict[str, Any] = {"mode": "overview", "file_path": target["path"],
                            "facets": ["data", "geometry"], "limit": 10}
    if target.get("sheet"):
        args["sheet"] = target["sheet"]
    cell_range = None
    if strategy == "column_sample" and column:
        letter, row = column["column"], column["header_row"]
        cell_range = bounded_range(f"{letter}{row}:{letter}{min(row + 20, 1_048_576)}")
        args["sheet"] = column["sheet"]
    elif strategy in {"selection", "formulas"} and target.get("sheet"):
        cell_range = bounded_range(str(target.get("range") or ""))
    if cell_range:
        args = {"mode": "range", "file_path": target["path"], "sheet": args["sheet"],
                "range": cell_range}
        if strategy == "formulas":
            args["facets"] = ["data", "dependencies"]
    return {"tool": "observe_spreadsheet", "arguments": args,
            "purpose": strategy if cell_range else "overview"}
