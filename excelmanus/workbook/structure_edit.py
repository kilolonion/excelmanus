"""Plan dependent reference edits before changing worksheet coordinates."""
from copy import copy

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange

from excelmanus.workbook.structural_refs import ReferenceTransform, MAX_ROW, MAX_COL


def apply_structure_edit(wb, action, target_sheet, *, at=1, count=1, new_name=None):
    rename = action == "rename_sheet"
    axis = "column" if action.endswith("cols") else "row"
    if not rename and (isinstance(at, bool) or isinstance(count, bool) or not isinstance(at, int) or not isinstance(count, int) or at < 1 or count < 1):
        raise ValueError("at/count 必须为正整数")
    ws = wb[target_sheet]
    transform = ReferenceTransform(target_sheet, axis=axis, at=at, count=count, delete=action.startswith("delete"), new_name=new_name if rename else None)
    if rename and (not new_name or new_name in wb.sheetnames):
        raise ValueError("新表名不能为空或重复")
    if not rename:
        bound = MAX_ROW if axis == "row" else MAX_COL
        extent = ws.max_row if axis == "row" else ws.max_column
        if at > bound or at + count - 1 > bound or (not transform.delete and extent + count > bound):
            raise ValueError("结构修改超出 Excel 行列边界")
    changes = []
    def set_later(obj, attr, value):
        changes.append((obj, attr, value))
    def formula(obj, attr, context):
        value = getattr(obj, attr, None)
        if value:
            set_later(obj, attr, transform.formula(value, context))
    def areas(value):
        result = []
        for part in str(value).split():
            changed = transform.address(part)
            if changed != "#REF!":
                result.append(changed)
        return " ".join(result)
    for sheet in wb.worksheets:
        for cell in list(sheet._cells.values()):
            if cell.data_type == "f":
                formula(cell, "value", sheet.title)
        for name in sheet.defined_names.values():
            formula(name, "attr_text", sheet.title)
        for validation in sheet.data_validations.dataValidation:
            formula(validation, "formula1", sheet.title)
            formula(validation, "formula2", sheet.title)
        for rules in sheet.conditional_formatting._cf_rules.values():
            for rule in rules:
                if rule.formula:
                    set_later(rule, "formula", [transform.formula(f, sheet.title) for f in rule.formula])
        for table in sheet.tables.values():
            for col in table.tableColumns:
                for attr in ("calculatedColumnFormula", "totalsRowFormula"):
                    item = getattr(col, attr, None)
                    if item is not None:
                        formula(item, "attr_text", sheet.title)
        seen = set()
        def chart_refs(obj, depth=0):
            if obj is None or id(obj) in seen or depth > 20:
                return
            seen.add(id(obj))
            if isinstance(getattr(obj, "f", None), str):
                formula(obj, "f", sheet.title)
            for value in getattr(obj, "__dict__", {}).values():
                for child in value if isinstance(value, (list, tuple)) else [value]:
                    if hasattr(child, "__dict__"):
                        chart_refs(child, depth + 1)
        for chart in sheet._charts:
            chart_refs(chart)
        if sheet._pivots:
            for pivot in sheet._pivots:
                source = getattr(getattr(pivot.cache, "cacheSource", None), "worksheetSource", None)
                if sheet is ws or (source and source.sheet == target_sheet):
                    raise ValueError("受影响区域含原生透视表，请重建透视表后再进行结构修改")
    for name in wb.defined_names.values():
        formula(name, "attr_text", None)
    if rename:
        for obj, attr, value in changes:
            setattr(obj, attr, value)
        ws.title = new_name
        return {"action": action, "references_updated": len(changes)}
    from openpyxl.utils.cell import range_boundaries
    for table in ws.tables.values():
        c1, r1, c2, r2 = range_boundaries(table.ref)
        if axis == "column" and c1 < at <= c2:
            raise ValueError("不能隐式修改 Table 列定义；请先使用 table.resize 或移除 Table 后修改列")
        new_ref = transform.address(table.ref)
        if new_ref == "#REF!" or (axis == "row" and transform.delete and at <= r1 < at + count):
            raise ValueError("操作会删除 Table 表头或整个 Table；请先显式删除 Table 对象")
        set_later(table, "ref", new_ref)
        if table.autoFilter:
            set_later(table.autoFilter, "ref", transform.address(table.autoFilter.ref))
    new_merges = [transform.address(str(m)) for m in ws.merged_cells.ranges]
    cf = []
    for key, rules in ws.conditional_formatting._cf_rules.items():
        new_key = copy(key)
        target = areas(key.sqref)
        if target:
            new_key.sqref = target
            cf.append((new_key, rules))
    dv = []
    for item in ws.data_validations.dataValidation:
        target = areas(item.sqref)
        if target:
            set_later(item, "sqref", target)
            dv.append(item)
    if ws.auto_filter.ref:
        value = transform.address(ws.auto_filter.ref)
        set_later(ws.auto_filter, "ref", None if value == "#REF!" else value)
    if ws.freeze_panes:
        value = transform.address(ws.freeze_panes)
        set_later(ws, "freeze_panes", None if value == "#REF!" else value)
    if ws.print_area:
        set_later(ws, "print_area", [transform.address(str(r)) for r in ws._print_area.ranges if transform.address(str(r)) != "#REF!"])
    for attr in ("print_title_rows", "print_title_cols"):
        value = getattr(ws, attr)
        if value:
            value = transform.address(value)
            set_later(ws, attr, None if value == "#REF!" else value)
    for obj in [*ws._charts, *ws._images]:
        anchor = obj.anchor
        if isinstance(anchor, str):
            value = transform.address(anchor)
            set_later(obj, "anchor", value if value != "#REF!" else f"A{at}" if axis == "row" else f"{get_column_letter(at)}1")
        else:
            for marker in (getattr(anchor, "_from", None), getattr(anchor, "to", None)):
                if marker is not None:
                    attr = "row" if axis == "row" else "col"
                    set_later(marker, attr, transform.point(getattr(marker, attr) + 1, anchor=True) - 1)
    dimensions = ws.row_dimensions if axis == "row" else ws.column_dimensions
    moved_dimensions = []
    for key, dim in list(dimensions.items()):
        old = int(key) if axis == "row" else column_index_from_string(key)
        if axis == "column":
            interval = transform.interval(dim.min or old, dim.max or old)
            new = interval[0] if interval else None
            if interval:
                set_later(dim, "min", interval[0]); set_later(dim, "max", interval[1])
        else:
            new = transform.point(old)
        if new is not None:
            moved_dimensions.append((new if axis == "row" else get_column_letter(new), dim))
    # No workbook mutation above this point.
    for obj, attr, value in changes:
        setattr(obj, attr, value)
    for merge in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merge))
    getattr(ws, action)(at, amount=count)
    for merge in new_merges:
        if merge != "#REF!":
            ws.merge_cells(merge)
    dimensions.clear()
    for key, dim in moved_dimensions:
        dim.index = key
        dimensions[key] = dim
    ws.conditional_formatting._cf_rules.clear()
    ws.conditional_formatting._cf_rules.update(cf)
    ws.data_validations.dataValidation = dv
    for cell in ws._cells.values():
        if cell.hyperlink:
            cell.hyperlink.ref = cell.coordinate
    return {"action": action, "references_updated": len(changes)}
