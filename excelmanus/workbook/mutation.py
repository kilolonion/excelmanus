"""Atomic V2 mutation compiler and observations of the published bytes."""

from __future__ import annotations

from io import BytesIO
from typing import Any
from zipfile import ZipFile
from copy import copy

from openpyxl import Workbook, load_workbook

from excelmanus.workbook.contracts import validate_operations
from excelmanus.workbook.geometry import scale_region, region_geometry
from excelmanus.workbook.refs import parse_rect
from excelmanus.workbook.snapshot import SnapshotError
from excelmanus.workspace.file_service import (
    TargetSpec,
    ReadDependency,
    service_for_guard,
)
from excelmanus.workbook_commit import (
    CommitError,
    remember_content_version,
    content_version_of_file,
)
from excelmanus.tools.context import require_guard, operation_id_for
from excelmanus.tools._helpers import prepare_excel_commit_path, MutationAborted


def _apply_object(
    wb: Any, op: dict[str, Any], image_sources: dict | None = None
) -> dict:
    from excelmanus.workbook.objects import KINDS, apply_object_operation
    from excelmanus.workbook.charts import (
        normalize_chart_args,
        add_chart_to_workbook,
        update_chart_in_workbook,
        delete_chart_from_workbook,
    )

    if op["kind"] in KINDS:
        return apply_object_operation(
            wb,
            op,
            guard=require_guard(),
            image_bytes=(image_sources or {}).get(op.get("image_path")),
        )
    if op["kind"] == "update_chart":
        return update_chart_in_workbook(wb, op)
    if op["kind"] == "delete_chart":
        return delete_chart_from_workbook(
            wb,
            sheet_name=op.get("sheet"),
            index=op.get("index"),
            target_cell=op.get("target_cell"),
            title=op.get("title"),
        )
    spec = normalize_chart_args(
        chart_type=op.get("chart_type", ""),
        data_range=op.get("data_range", ""),
        categories_range=op.get("categories_range"),
        sheet_name=op.get("sheet"),
        target_cell=op.get("target_cell", "A1"),
        target_sheet=op.get("target_sheet"),
        title=op.get("title"),
        x_title=op.get("x_title"),
        y_title=op.get("y_title"),
        style=op.get("style"),
        width=op.get("width", 15),
        height=op.get("height", 10),
        from_rows=op.get("from_rows", False),
    )
    from excelmanus.engine_core.tool_result import ToolResult

    if isinstance(spec, ToolResult):
        raise MutationAborted(spec)
    return add_chart_to_workbook(wb, spec)


def execute_operation(
    wb: Any,
    operation: dict,
    *,
    file_path: str = "",
    create: bool = False,
    image_sources: dict | None = None,
) -> dict:
    """The only in-memory executor, including interactive edits and merge review."""
    from excelmanus.workbook import domain

    kind = operation["kind"]
    if kind == "geometry.resize":
        from excelmanus.workbook.geometry import column_native
        from openpyxl.utils import get_column_letter

        column = operation["axis"] == "column"
        sizes = {
            (get_column_letter(int(k)) if column else k): (
                column_native(v) if column else v * 72 / 96
            )
            for k, v in operation["sizes"].items()
        }
        translated = {
            "kind": "size",
            "sheet": operation["sheet"],
            "column_widths" if column else "row_heights": sizes,
        }
        validate_operations([translated])
        domain._apply_format(wb, translated)
        return {
            "kind": kind,
            "requested": operation,
            "native": translated,
            "conversion": "96 DPI, estimated 7px maximum digit width",
        }
    if kind == "cells.patch":
        from excelmanus.workbook.snapshot import require_default_sheet
        from excelmanus.workbook.cells import assign_cell_value
        from openpyxl.styles.cell_style import StyleArray

        ws = wb[require_default_sheet(wb.sheetnames, operation.get("sheet"))]
        for item in operation["cells"]:
            rect = parse_rect(item["cell"])
            if (
                rect.min_row != rect.max_row
                or rect.min_col != rect.max_col
                or rect.sheet not in {None, ws.title}
            ):
                raise ValueError(
                    "cells.patch requires individual cells on the specified sheet"
                )
            if "value" in item:
                assign_cell_value(ws, rect.min_row, rect.min_col, item["value"])
            if "style" in item:
                if item["style"] is None:
                    ws.cell(rect.min_row, rect.min_col)._style = StyleArray()
                else:
                    domain._apply_format(
                        wb,
                        {
                            "kind": "format",
                            "sheet": ws.title,
                            "range": item["cell"],
                            **item["style"],
                        },
                    )
        return {
            "kind": kind,
            "sheet": ws.title,
            "cells_changed": len(operation["cells"]),
        }
    elif kind == "geometry.scale":
        ws = wb[operation["sheet"]]
        rect = _operation_region(ws, operation)
        if (
            rect is None
            or rect.max_row - rect.min_row + rect.max_col - rect.min_col + 2 > 20000
        ):
            raise ValueError("geometry.scale requires at most 20000 affected axes")
        effect = scale_region(
            ws,
            rect,
            operation["x"],
            operation["y"],
            preserve_outside=operation.get("preserve_outside", False),
        )
        return {"kind": kind, "sheet": ws.title, **effect}
    elif kind in domain._FORMAT_KIND_FIELDS:
        label, skipped = domain._apply_format(wb, operation)
        return {"kind": kind, "applied": label, "skipped_merged_non_anchors": skipped}
    elif kind in {
        "chart",
        "update_chart",
        "delete_chart",
        "image",
        "table",
        "defined_name",
        "pivot_table",
        "comment",
        "hyperlink",
    }:
        return {"kind": kind, "object": _apply_object(wb, operation, image_sources)}
    else:
        applied, warnings = domain._apply_edit_operations(
            wb, [operation], file_path=file_path, create_workbook=create
        )
        return {"kind": kind, "applied": applied, "warnings": warnings}


def _preservation_check(
    before: bytes | None, after: bytes, operations: list[dict]
) -> None:
    if before is None:
        return
    with ZipFile(BytesIO(before)) as src, ZipFile(BytesIO(after)) as dst:
        from excelmanus.workbook.adapter import (
            verify_relationships,
            verify_workbook_extensions,
        )

        verify_relationships(src, dst, operations)
        verify_workbook_extensions(src, dst)
        old, new = set(src.namelist()), set(dst.namelist())
        # A removed .rels container is safe only after its individual edges
        # have passed the relationship preservation check above.
        permitted = {"xl/calcChain.xml", "xl/sharedStrings.xml"} | {
            n for n in old if n.endswith(".rels")
        }
        if any(
            op["kind"] == "comment" and op.get("action") == "delete"
            for op in operations
        ):
            permitted.update(
                n
                for n in old
                if n.startswith("xl/comments/")
                or n.startswith("xl/drawings/commentsDrawing")
            )
        for op in operations:
            if op["kind"] == "sheet" and op.get("action") == "delete":
                permitted.update(n for n in old if n.startswith("xl/worksheets/"))
            if op["kind"] in {"delete_chart", "image", "table", "pivot_table"} and (
                op.get("action") == "delete" or op["kind"] == "delete_chart"
            ):
                permitted.update(
                    n
                    for n in old
                    if n.startswith(
                        (
                            "xl/charts/",
                            "xl/drawings/",
                            "xl/media/",
                            "xl/tables/",
                            "xl/pivot",
                        )
                    )
                )
        missing = old - new - permitted
        if missing:
            raise CommitError(
                "UNSUPPORTED_PRESERVATION",
                f"Serialization would discard package parts: {sorted(missing)[:20]}",
            )
        for name in old & new:
            if (
                name.startswith("xl/media/")
                and not any(op["kind"] in {"image", "sheet"} for op in operations)
                and src.read(name) != dst.read(name)
            ):
                raise CommitError(
                    "UNSUPPORTED_PRESERVATION", f"Untouched image bytes changed: {name}"
                )
            if "vbaProject" in name and src.read(name) != dst.read(name):
                raise CommitError(
                    "UNSUPPORTED_PRESERVATION", f"Macro bytes changed: {name}"
                )
            if (
                name.endswith(".xml")
                and b"extLst" in src.read(name)
                and b"extLst" not in dst.read(name)
            ):
                raise CommitError(
                    "UNSUPPORTED_PRESERVATION",
                    f"Serialization would remove extension data: {name}",
                )
        if any(n.startswith("_xmlsignatures/") for n in old):
            raise CommitError(
                "SIGNED_WORKBOOK",
                "Editing would invalidate the workbook package signature",
            )


def _operation_region(ws: Any, op: dict) -> Any | None:
    from openpyxl.utils import get_column_letter, column_index_from_string

    address = op.get("range") or op.get("start_cell")
    if op["kind"] == "geometry.resize":
        indices = [int(k) for k in op["sizes"]]
        address = (
            f"{get_column_letter(min(indices))}1:{get_column_letter(max(indices))}{min(ws.max_row, 30)}"
            if op["axis"] == "column"
            else f"A{min(indices)}:{get_column_letter(min(ws.max_column, 30))}{max(indices)}"
        )
    if not address and op["kind"] == "size":
        cols = [column_index_from_string(k) for k in op.get("column_widths", {})]
        rows = [int(k) for k in op.get("row_heights", {})]
        address = f"{get_column_letter(min(cols, default=1))}{min(rows, default=1)}:{get_column_letter(max(cols, default=min(ws.max_column or 1, 30)))}{max(rows, default=min(ws.max_row or 1, 30))}"
    if not address:
        return None
    rect = parse_rect(address)
    if rect.whole_column or rect.whole_row:
        from excelmanus.workbook.address import resolve_range_to_bounds

        b = resolve_range_to_bounds(
            address, used_max_row=ws.max_row, used_max_col=ws.max_column
        )
        rect = parse_rect(
            f"{get_column_letter(b.min_col)}{b.min_row}:{get_column_letter(b.max_col)}{b.max_row}"
        )
    return rect


def _sample_targets(
    wb: Any, operations: list[dict], limit: int = 64
) -> list[tuple[str, int, int]]:
    targets: dict[tuple[str, int, int], None] = {}
    for op in operations:
        if op.get("selection") is not None:
            from excelmanus.workbook.snapshot import parse_bound_selection

            selection = parse_bound_selection(op["selection"])
            if selection and selection.sheet in wb:
                for row in selection.rows[:limit]:
                    for col in (selection.cols or ())[:limit]:
                        targets[(selection.sheet, row, col)] = None
                        if len(targets) >= limit:
                            return list(targets)[:limit]
        sheet = op.get("sheet") or (
            wb.active.title if len(wb.worksheets) == 1 else None
        )
        if sheet not in wb:
            continue
        addresses = (
            [item["cell"] for item in op.get("cells", [])]
            if op["kind"] == "cells.patch"
            else []
        )
        if op.get("range") and op["kind"] in {
            "format",
            "clear",
            "fill",
            "replace",
            "sort",
        }:
            rect = _operation_region(wb[sheet], op)
            if rect:
                for row in sorted(
                    {rect.min_row, rect.max_row, (rect.min_row + rect.max_row) // 2}
                ):
                    for col in sorted(
                        {rect.min_col, rect.max_col, (rect.min_col + rect.max_col) // 2}
                    ):
                        targets[(sheet, row, col)] = None
        if (
            op["kind"] == "write"
            and op.get("start_cell")
            and isinstance(op.get("values"), list)
        ):
            from excelmanus.workbook.domain import _bind_write_origin

            actual_sheet, first_row, first_col = _bind_write_origin(
                wb, op["start_cell"], sheet
            )
            values = op["values"]
            if values:
                for r in sorted({0, len(values) - 1, len(values) // 2}):
                    if values[r]:
                        for c in sorted({0, len(values[r]) - 1, len(values[r]) // 2}):
                            targets[(actual_sheet, first_row + r, first_col + c)] = None
        for address in addresses[:limit]:
            rect = parse_rect(address)
            targets[(sheet, rect.min_row, rect.min_col)] = None
        if len(targets) >= limit:
            break
    return list(targets)[:limit]


def apply_changes(
    file_path: str,
    operations: list[dict] | None = None,
    *,
    expected_version: str | None = None,
    create: bool = False,
    workbook_spec: dict | None = None,
    workbooks: list[dict] | None = None,
    read_dependencies: list[dict] | None = None,
    event_context: dict | None = None,
    dry_run: bool = False,
) -> dict:
    guard = require_guard()
    from excelmanus.workbook import domain

    batches = (
        workbooks
        if workbooks is not None
        else [
            {
                "file_path": file_path,
                "operations": operations,
                "expected_version": expected_version,
                "create": create,
                "workbook_spec": workbook_spec,
            }
        ]
    )
    if workbooks is not None and (
        file_path or operations or workbook_spec or create or expected_version
    ):
        raise ValueError("workbooks cannot be combined with single-file fields")
    if not batches:
        raise ValueError("At least one workbook is required")
    specs, observations, plans = [], {}, {}
    document_metadata: dict[str, dict[str, Any]] = {}
    dependencies = list(read_dependencies or [])
    for item in batches:
        if "dry_run" in item:
            raise ValueError(
                "dry_run applies to the entire ChangeSet, not an individual workbook"
            )
        if not str(item.get("file_path") or "").strip():
            raise CommitError("PATH_REQUIRED", "file_path is required")
        from excelmanus.security.source_isolation import (
            is_probe_path,
            probe_error_message,
            PROBE_FILE_FORBIDDEN,
        )

        if is_probe_path(item["file_path"]):
            raise CommitError(
                PROBE_FILE_FORBIDDEN, probe_error_message(item["file_path"])
            )
        path, rel = prepare_excel_commit_path(guard, item["file_path"])
        if path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx"}:
            raise ValueError("ChangeSet writes require xlsx/xlsm/xltx")
        dependencies.extend(item.get("read_dependencies") or [])
        ops = item.get("operations")
        document = item.get("workbook_spec")
        creating = bool(item.get("create") or document is not None)
        if document is not None:
            if ops:
                raise ValueError("workbook_spec and operations are mutually exclusive")
            from excelmanus.workbook.spec import (
                validate_workbook_spec,
                document_operations,
            )

            spec = validate_workbook_spec(document)
            document_metadata[rel] = {
                "purpose": spec.purpose,
                "name": spec.name,
                "layout_references": [
                    {
                        "sheet": sheet.name,
                        **sheet.layout_reference.model_dump(),
                        "geometry_basis": "96 DPI; 7px estimated digit width; renderer verification required",
                    }
                    for sheet in spec.sheets
                    if sheet.layout_reference
                ],
                "uncertainties": [
                    item.model_dump(exclude_none=True) for item in spec.uncertainties
                ],
            }
            err = domain._expand_source_csv_sheets(spec, guard, dependencies)
            if err is not None:
                raise MutationAborted(err)
            ops = document_operations(spec)
        ops = validate_operations(ops)
        version = item.get("expected_version")
        if not creating and not version:
            raise CommitError(
                "VERSION_CONFLICT",
                "Existing workbook changes require the observed content_version",
                fields={
                    "path": rel,
                    "reason": "expected_version_missing",
                    "content_version": content_version_of_file(path),
                },
            )
        from excelmanus.workbook.snapshot import (
            validate_selection_target,
            parse_bound_selection,
        )

        for op in ops:
            if "selection" in op:
                validate_selection_target(op["selection"], rel, op.get("sheet"))
                selection = parse_bound_selection(op["selection"])
                if selection is None or selection.snapshot.content_version != version:
                    raise CommitError(
                        "SELECTION_STALE",
                        "Selection and mutation must use the same observed version",
                    )
        image_sources = {}
        from excelmanus.workbook.snapshot import open_snapshot

        for op in ops:
            if op["kind"] == "image" and op.get("image_path"):
                source = open_snapshot(op["image_path"])
                image_sources[op["image_path"]] = source.read_bytes()
                dependencies.append(
                    {"path": source.file.relative, "version": source.content_version}
                )
            join = op.get("join")
            if isinstance(join, dict) and join.get("file_path"):
                source = open_snapshot(
                    join["file_path"], expected_version=join.get("expected_version")
                )
                if source.file.relative != rel:
                    join["expected_version"] = source.content_version
                    dependencies.append(
                        {
                            "path": source.file.relative,
                            "version": source.content_version,
                        }
                    )
        plans[rel] = ops

        def builder(
            before: bytes | None,
            *,
            _ops=ops,
            _path=path,
            _rel=rel,
            _creating=creating,
            _image_sources=image_sources,
        ) -> bytes:
            if before:
                from excelmanus.workbook.adapter import ensure_editable

                ensure_editable(before)
            wb = (
                load_workbook(
                    BytesIO(before),
                    rich_text=True,
                    keep_vba=_path.suffix.lower() == ".xlsm",
                )
                if before
                else Workbook()
            )
            from excelmanus.workbook.formula_values import attach_formula_source

            attach_formula_source(wb, before)
            default_sheet = wb.active.title if before is None else None
            evidence, effects = [], []
            try:
                for index, operation in enumerate(_ops):
                    from excelmanus.tools.spreadsheet_engine_tools import _cancelled

                    _cancelled()
                    try:
                        kind = operation["kind"]
                        sheet = operation.get("sheet")
                        ws = wb[sheet] if sheet and sheet in wb else wb.active
                        if kind in {"size", "geometry.scale", "geometry.resize"}:
                            resolved = domain._format_bound_sheet(
                                operation, operation.get("range", "")
                            )
                            ws = domain._worksheet(wb, resolved)
                        rect = (
                            _operation_region(ws, operation)
                            if kind in {"size", "geometry.scale", "geometry.resize"}
                            else None
                        )
                        if (
                            rect
                            and rect.max_row
                            - rect.min_row
                            + rect.max_col
                            - rect.min_col
                            + 2
                            > 20000
                        ):
                            raise ValueError(
                                "Dimension observation exceeds 20000 axes; split the batch"
                            )
                        before_geometry = region_geometry(ws, rect) if rect else None
                        effects.append(
                            execute_operation(
                                wb,
                                operation,
                                file_path=_rel,
                                create=_creating,
                                image_sources=_image_sources,
                            )
                        )
                        if rect:
                            if kind == "size":
                                affected_axes = []
                                if operation.get("auto_fit"):
                                    affected_axes.append(
                                        "rows"
                                        if operation.get("axis") == "row"
                                        else "columns"
                                    )
                                if operation.get("column_widths"):
                                    affected_axes.append("columns")
                                if operation.get("row_heights"):
                                    affected_axes.append("rows")
                                affected_axes = list(dict.fromkeys(affected_axes)) or [
                                    "columns",
                                    "rows",
                                ]
                            else:
                                affected_axes = ["columns", "rows"]
                            evidence.append(
                                {
                                    "operation_index": index,
                                    "kind": kind,
                                    "sheet": ws.title,
                                    "range": rect.to_a1(include_sheet=False),
                                    "before": before_geometry,
                                    "requested": region_geometry(ws, rect),
                                    "dimensions_requested": copy(operation),
                                    "axes": affected_axes,
                                    "scope": "entire affected rows and columns",
                                }
                            )
                    except MutationAborted as exc:
                        payload = (
                            exc.result.value
                            if isinstance(exc.result.value, dict)
                            else {}
                        )
                        fields = {
                            k: v
                            for k, v in payload.items()
                            if k
                            not in {
                                "status",
                                "message",
                                "error_code",
                                "failure_class",
                                "remediation",
                            }
                        }
                        fields.update(
                            operation_index=index,
                            operation_kind=operation["kind"],
                            committed=False,
                            partial=False,
                            applied=[],
                        )
                        message = (
                            exc.result.error.message
                            if exc.result.error
                            else exc.result.model_text
                        )
                        import re

                        message = re.sub(
                            r"^operations\[\d+\](?: \(kind=[^)]*\))?:\s*", "", message
                        )
                        raise CommitError(
                            exc.result.error.code
                            if exc.result.error
                            else "INVALID_ARGS",
                            f"operations[{index}] (kind={operation['kind']}): {message}",
                            fields=fields,
                        ) from exc
                    except (CommitError, SnapshotError) as exc:
                        # Domain failures (for example an ambiguous chart
                        # source) need the same operation identity and atomic
                        # failure facts as schema/argument failures.
                        raise CommitError(
                            exc.code,
                            f"operations[{index}] (kind={operation['kind']}): {exc}",
                            fields={
                                **exc.fields,
                                "operation_index": index,
                                "operation_kind": operation["kind"],
                                "committed": False,
                                "partial": False,
                                "applied": [],
                            },
                        ) from exc
                    except (ValueError, TypeError, KeyError) as exc:
                        raise CommitError(
                            "INVALID_ARGS",
                            f"operations[{index}]: {exc}",
                            fields={"operation_index": index, "committed": False},
                        ) from exc
                if default_sheet and default_sheet in wb and len(wb.worksheets) > 1:
                    ws = wb[default_sheet]
                    if not any(
                        c.value is not None or c.has_style for c in ws._cells.values()
                    ):
                        del wb[default_sheet]
                samples = []
                for sheet_name, row, col in _sample_targets(wb, _ops):
                    cell = wb[sheet_name].cell(row, col)
                    samples.append(
                        (
                            sheet_name,
                            row,
                            col,
                            cell.value,
                            cell.data_type,
                            {
                                key: copy(getattr(cell, key))
                                for key in (
                                    "font",
                                    "fill",
                                    "border",
                                    "alignment",
                                    "number_format",
                                )
                            },
                        )
                    )
                for entry in evidence:
                    if entry["sheet"] in wb:
                        entry["expected_final"] = region_geometry(
                            wb[entry["sheet"]], parse_rect(entry["range"])
                        )
                    else:
                        entry["status"] = "superseded_by_sheet_operation"
                out = BytesIO()
                wb.save(out)
                from excelmanus.workbook.adapter import (
                    preserve_empty_custom_properties,
                    preserve_workbook_extensions,
                )

                data = preserve_workbook_extensions(before, out.getvalue())
                data = preserve_empty_custom_properties(before, data)
                _preservation_check(before, data, _ops)
                # Verify serialization using a new reader, not the mutable workbook.
                reopened = load_workbook(
                    BytesIO(data), read_only=False, data_only=False, rich_text=True
                )
                try:
                    checks = []
                    from excelmanus.workbook.checks import _values_equal

                    for sheet_name, row, col, expected, data_type, style in samples:
                        cell = reopened[sheet_name].cell(row, col)
                        ok = (
                            _values_equal(expected, cell.value)
                            and data_type == cell.data_type
                            and all(
                                value == copy(getattr(cell, key))
                                for key, value in style.items()
                            )
                        )
                        if not ok:
                            raise CommitError(
                                "SERIALIZATION_MISMATCH",
                                f"Serialized cell differs: {sheet_name}!{cell.coordinate}",
                            )
                        from excelmanus.workbook.observation import _scalar

                        checks.append(
                            {
                                "sheet": sheet_name,
                                "cell": cell.coordinate,
                                "value": _scalar(cell.value),
                                "formula": cell.value
                                if cell.data_type == "f"
                                else None,
                                "verified": True,
                            }
                        )
                    for entry in evidence:
                        if entry.get("status") == "superseded_by_sheet_operation":
                            continue
                        actual = region_geometry(
                            reopened[entry["sheet"]], parse_rect(entry["range"])
                        )
                        entry["persisted"] = actual
                        entry["dimensions_persisted"] = {
                            axis: [
                                {
                                    "index": d["index"],
                                    "native": d["native"],
                                    "unit": d["unit"],
                                }
                                for d in actual[axis]
                            ]
                            for axis in ("rows", "columns")
                        }
                        entry["geometry_delta"] = {
                            "width_px": actual["width_px"]
                            - entry["before"]["width_px"],
                            "height_px": actual["height_px"]
                            - entry["before"]["height_px"],
                            "status": "estimated_pixels",
                        }
                        entry["verified"] = all(
                            [
                                (x["index"], round(x["native"], 8), x["hidden"])
                                for x in actual[axis]
                            ]
                            == [
                                (x["index"], round(x["native"], 8), x["hidden"])
                                for x in entry["expected_final"][axis]
                            ]
                            for axis in entry.get("axes", ("columns", "rows"))
                        )
                        if not entry["verified"]:
                            raise CommitError(
                                "SERIALIZATION_MISMATCH",
                                "Published dimensions would differ from the requested operation",
                            )
                    from excelmanus.workbook.presentation import _collect_print_settings

                    persisted_sheets = [
                        {
                            "name": ws.title,
                            "freeze_panes": str(ws.freeze_panes or ""),
                            "print_settings": _collect_print_settings(ws),
                            "merges": [str(m) for m in ws.merged_cells.ranges],
                        }
                        for ws in reopened
                    ]
                    from excelmanus.workbook.observation import _objects

                    persisted_objects = (
                        {ws.title: _objects(reopened, ws) for ws in reopened}
                        if any(
                            op["kind"]
                            in {
                                "chart",
                                "update_chart",
                                "delete_chart",
                                "table",
                                "defined_name",
                                "image",
                                "comment",
                                "hyperlink",
                                "pivot_table",
                            }
                            for op in _ops
                        )
                        else None
                    )
                finally:
                    reopened.close()
                observations[_rel] = {
                    "geometry_changes": evidence,
                    "visual_observed": False,
                    "operations": effects,
                    "sheets": persisted_sheets,
                    "objects": {"status": "observed", "sheets": persisted_objects}
                    if persisted_objects is not None
                    else {"status": "not_requested"},
                    "cell_checks": checks,
                    "coverage": {
                        "kind": "sampled",
                        "sample_size": len(checks),
                        "max_cells": 64,
                        "scope": "changed values/formulas/styles; geometry separately checked",
                    },
                    "formula_cache": "invalidated_by_serialization; calculate explicitly when needed",
                }
                if _rel in document_metadata:
                    observations[_rel]["document"] = document_metadata[_rel]
                _cancelled()
                return data
            finally:
                wb.close()

        specs.append(
            TargetSpec(
                "create" if creating else "update",
                rel,
                expected_version=version,
                builder=builder,
                event_context=event_context,
                intent={
                    "schema_version": "workbook/2",
                    "operations": ops,
                    "create": creating,
                    "document": document_metadata.get(rel),
                },
            )
        )
    deps = [
        ReadDependency(path=dep["path"], version=dep["version"]) for dep in dependencies
    ]
    service = service_for_guard(guard)
    if dry_run:
        for dep in deps:
            actual = content_version_of_file(guard.resolve_and_validate(dep.path))
            if actual != dep.version:
                raise CommitError(
                    "VERSION_CONFLICT",
                    "Read dependency changed",
                    fields={"path": dep.path, "committed": False},
                )
        for spec in specs:
            target = guard.resolve_and_validate(spec.path)
            if spec.op == "create" and target.exists():
                raise CommitError("FILE_EXISTS", "Dry-run creation target exists")
            before = target.read_bytes() if target.is_file() else None
            from excelmanus.workbook_commit import content_version_of

            if spec.op != "create" and (
                before is None or content_version_of(before) != spec.expected_version
            ):
                raise CommitError(
                    "VERSION_CONFLICT",
                    "Dry-run target changed",
                    fields={"committed": False},
                )
            spec.builder(before)
        return {
            "status": "success",
            "schema_version": "workbook/2",
            "committed": False,
            "dry_run": True,
            "files": [
                {"file_path": p, "observation": o} for p, o in observations.items()
            ],
            "receipt": {},
            "applied": [],
            "planned": plans,
        }
    receipt = service.apply_batch(
        specs, operation_id=operation_id_for("changeset"), read_dependencies=deps
    )
    files = []
    for target in receipt.targets:
        path = target.path
        if target.publish_status != "published":
            files.append(
                {
                    "file_path": path,
                    "content_version": target.before_version,
                    "observation": {
                        "status": "not_published",
                        "visual_observed": False,
                    },
                }
            )
            continue
        remember_content_version(path, target.after_version)
        files.append(
            {
                "file_path": path,
                "content_version": target.after_version,
                "previous_version": target.before_version,
                "observation": {
                    **observations.get(
                        path, {"status": "replayed_receipt", "visual_observed": False}
                    ),
                    "content_version": target.after_version,
                    "snapshot_id": f"{receipt.workspace_key}|{path}|{target.after_version}",
                    "observation_id": f"mutation:{receipt.operation_id}:{path}",
                },
            }
        )
    result = {
        "status": "success" if receipt.state == "committed" else "partial",
        "schema_version": "workbook/2",
        "committed": receipt.state == "committed",
        "operation_id": receipt.operation_id,
        "receipt": receipt.to_dict(),
        "files": files,
        "applied": [op["kind"] for ops in plans.values() for op in ops],
    }
    if receipt.state != "committed":
        result.update(
            status="error",
            error_code=receipt.error_code or "COMMIT_FAILED",
            message=receipt.message,
            recovery_required=receipt.resumeable,
        )
    if len(files) == 1:
        result.update(files[0])
        if isinstance(files[0].get("observation"), dict) and files[0][
            "observation"
        ].get("document"):
            result["document"] = files[0]["observation"]["document"]
    return result
