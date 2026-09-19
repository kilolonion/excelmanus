"""Excel A1 / Sheet!A1 解析：公开函数委托 ``excelmanus.workbook.refs``。

行为变化（相对旧实现）：
- 多区域 ``A1:B2,C3:D4`` 由 ``parse_ref`` 接受，不再视为非法语法；
  ``resolve_range_to_bounds`` 仍只展开**单矩形**（读路径尚未消费并集）。
- ``$`` 绝对引用标记在 ``parse_sheet_address`` / ``parse_ref`` 中保留。
- 命名区域与表引用可解析（工作簿元数据绑定留给上层）。
- 端点包含，行列均为 1-based。

``top_left_cell`` 仍输出 openpyxl 可用坐标（去掉 ``$`` 与表名）。
"""

from __future__ import annotations

from typing import Any, NamedTuple

from openpyxl.utils import column_index_from_string, get_column_letter

from excelmanus.workbook.refs import (
    AreaRef,
    InvalidRefError,
    RectRef,
    parse_rect,
    parse_ref,
)

# 整行/整列按已用范围裁剪后再套此格数上限，避免把 1048576 行读进来。
MAX_WHOLE_RANGE_CELLS = 10_000

_INVALID_COORD_MARKERS = (
    "is not a valid coordinate",
    "is not a valid range",
    "invalid coordinates",
    "无法解析引用",
    "不是合法",
    "空地址",
)


class SheetAddress(NamedTuple):
    sheet: str | None
    address: str


def parse_sheet_address(value: str | None) -> SheetAddress:
    """拆分 ``表!A1:C5`` / ``'表名'!A1`` / ``A1:C5`` / 多区域 / 命名与表引用。

    保留 ``$``；工作表名两侧单引号按 Excel 规则解开（``''`` → ``'``）。
    空字符串仍返回 ``(None, "")``，以便调用方做缺省判断。
    """
    raw = str(value or "").strip()
    if not raw:
        return SheetAddress(None, "")
    return _area_to_sheet_address(parse_ref(raw))


def strip_sheet_qualifier(value: str | None) -> str:
    """只保留 A1 / A1:C5 / 多区域本地部分（含 ``$``）。"""
    return parse_sheet_address(value).address


def combine_sheet_names(explicit: str | None, from_address: str | None) -> str | None:
    """合并独立 sheet 字段与 ``表!A1`` 中的表名；冲突则抛 ValueError。"""
    explicit_name = str(explicit).strip() if explicit not in (None, "") else None
    address_name = str(from_address).strip() if from_address not in (None, "") else None
    if explicit_name and address_name and explicit_name != address_name:
        raise ValueError(
            f"工作表不一致：sheet={explicit_name!r} 与 range 中的 {address_name!r}。"
            "请只保留一处，或两者写成同一个表名。"
        )
    return address_name or explicit_name


def looks_like_coordinate_error(exc: BaseException | str) -> bool:
    if isinstance(exc, InvalidRefError):
        return True
    text = str(exc)
    lowered = text.lower()
    return any(marker in lowered or marker in text for marker in _INVALID_COORD_MARKERS)


def column_map_to_list(raw: dict[object, object]) -> list[float]:
    """``{"A": 9, "C": 14}`` → ``[9, 0, 14]``（缺列填 0，编译时跳过）。"""
    widths: dict[int, float] = {}
    for key, value in raw.items():
        try:
            width = float(value)  # type: ignore[arg-type]
        except (TypeError, ValueError):
            continue
        col = _column_key_to_index(key)
        if col is None:
            continue
        widths[col] = width
    if not widths:
        return []
    max_col = max(widths)
    return [float(widths.get(index, 0.0)) for index in range(1, max_col + 1)]


def _column_key_to_index(key: object) -> int | None:
    if isinstance(key, bool):
        return None
    if isinstance(key, int):
        return key if key >= 1 else key + 1
    text = str(key).strip()
    if not text:
        return None
    if text.isdigit():
        number = int(text)
        return number if number >= 1 else number + 1
    try:
        return int(column_index_from_string(text.upper()))
    except ValueError:
        return None


def top_left_cell(address: str) -> str:
    """``A1:C5`` → ``A1``；单格去掉 ``$`` 后返回。表名忽略。

    语法走 ``parse_ref``。命名/表/并集/整轴没有可写的单一左上角，抛 ``InvalidRefError``。
    """
    from excelmanus.workbook.refs import CellRef, NamedRef, RectRef, TableRef, parse_ref

    text = str(address or "").strip()
    if not text:
        return ""
    area = parse_ref(text)
    if len(area.areas) != 1:
        raise InvalidRefError("并集没有单一左上角。write 请改用单一矩形。")
    part = area.areas[0]
    if isinstance(part, CellRef):
        return f"{get_column_letter(part.col)}{part.row}"
    if isinstance(part, RectRef):
        if part.whole_column or part.whole_row:
            raise InvalidRefError("整轴没有可写的单一左上角。请写有限矩形，例如 A1:A100。")
        return f"{get_column_letter(part.min_col)}{part.min_row}"
    if isinstance(part, (NamedRef, TableRef)):
        raise InvalidRefError("命名区域/表引用必须先绑定工作簿再取左上角。")
    raise InvalidRefError(f"无法取左上角：{text!r}")


class ResolvedRange(NamedTuple):
    min_col: int
    min_row: int
    max_col: int
    max_row: int
    resolved: str
    requested: str
    clipped: bool
    truncated: bool


def worksheet_used_shape(ws: Any) -> tuple[int, int]:
    """工作表已用最大行列（1-based）。空表按 1×1。"""
    max_row = getattr(ws, "max_row", None)
    max_col = getattr(ws, "max_column", None)
    try:
        rows = int(max_row) if max_row not in (None, "") else 1
        cols = int(max_col) if max_col not in (None, "") else 1
    except (TypeError, ValueError):
        return 1, 1
    return max(1, rows), max(1, cols)


def resolve_range_to_bounds(
    address: str,
    *,
    used_max_row: int,
    used_max_col: int,
    max_cells: int | None = None,
) -> ResolvedRange:
    """把 A1 / A:A / 1:1 收成有限矩形。

    整列、整行先裁到工作表已用范围，再套 ``MAX_WHOLE_RANGE_CELLS``。
    多区域请用 ``parse_ref``；本函数只接受单矩形（命名/表引用需上层先解析）。
    有界矩形（``A1:A10``）不按已用范围收缩。端点 1-based 且包含。
    """
    requested_raw = str(address or "").strip()
    if not requested_raw:
        raise InvalidRefError("空地址。正确写法示例：A1、A1:B2、A:A。")
    rect = parse_rect(requested_raw)
    requested = rect.to_a1(include_sheet=False).replace("$", "")
    cap = MAX_WHOLE_RANGE_CELLS if max_cells is None else max(1, int(max_cells))
    used_max_row = max(1, int(used_max_row or 1))
    used_max_col = max(1, int(used_max_col or 1))
    used = RectRef(
        sheet=rect.sheet,
        min_row=1,
        max_row=used_max_row,
        min_col=1,
        max_col=used_max_col,
    )
    whole_col = rect.whole_column
    whole_row = rect.whole_row
    expanded = AreaRef((rect,)).expand(used)
    out = expanded.areas[0]
    assert isinstance(out, RectRef)
    min_col, min_row, max_col, max_row = out.min_col, out.min_row, out.max_col, out.max_row

    clipped = whole_col or whole_row
    truncated = False
    rows = max_row - min_row + 1
    cols = max_col - min_col + 1
    if clipped and rows * cols > cap:
        truncated = True
        if whole_col:
            max_row = min_row + max(1, cap // max(cols, 1)) - 1
            rows = max_row - min_row + 1
        if whole_row or rows * cols > cap:
            max_col = min_col + max(1, cap // max(rows, 1)) - 1

    resolved = (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{max_row}"
    )
    return ResolvedRange(
        min_col=min_col,
        min_row=min_row,
        max_col=max_col,
        max_row=max_row,
        resolved=resolved,
        requested=requested,
        clipped=clipped,
        truncated=truncated,
    )


def _area_to_sheet_address(area: AreaRef) -> SheetAddress:
    sheets = [getattr(part, "sheet", None) for part in area.areas]
    if all(sheet == sheets[0] for sheet in sheets):
        local = ",".join(part.to_a1(include_sheet=False) for part in area.areas)
        return SheetAddress(sheets[0], local)
    return SheetAddress(None, area.to_a1())
