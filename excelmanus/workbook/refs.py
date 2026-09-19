"""单元格 / 区域引用的唯一解析与规范化出口。

本模块只做语法层：把 Excel 惯例的文本变成结构化对象。命名区域与表对象
的工作簿元数据绑定留给上层。端点均为 1-based 且包含。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, replace
from typing import NoReturn

from openpyxl.utils import column_index_from_string, get_column_letter

EXCEL_MAX_ROW = 1_048_576
EXCEL_MAX_COL = 16_384  # XFD

_CELL_TOKEN = re.compile(
    r"^(\$?)([A-Za-z]{1,3})(\$?)([1-9][0-9]{0,6})$"
)
_WHOLE_COL_TOKEN = re.compile(
    r"^(\$?)([A-Za-z]{1,3})\s*:\s*(\$?)([A-Za-z]{1,3})$"
)
_WHOLE_ROW_TOKEN = re.compile(
    r"^(\$?)([1-9][0-9]{0,6})\s*:\s*(\$?)([1-9][0-9]{0,6})$"
)
_R1C1_TOKEN = re.compile(
    r"^[Rr](?:\[-?\d+\]|\d+)?[Cc](?:\[-?\d+\]|\d+)?$"
)
_CELL_LIKE_LEADING_ZERO = re.compile(
    r"^\$?[A-Za-z]{1,3}\$?0[0-9]*$"
)
_NAME_TOKEN = re.compile(
    r"^(?:[\\_]|[^\W\d])[\w.?\\]*$",
    re.UNICODE,
)

_REFERENCE_EXAMPLES: tuple[str, ...] = (
    "A1",
    "$A$1",
    "Sheet1!A1:B2",
    "'My Sheet'!A1",
    "A:A",
    "1:1",
    "A1:B2,C3:D4",
    "Sheet1!A1,Sheet2!B1",
    "MyName",
    "Table1[列]",
)

_SCHEMA_HINT = (
    "整列请写 A:A 而不是 A；整行请写 1:1 而不是 1；"
    "含空格的表名必须用单引号：'My Sheet'!A1；"
    "表名中的单引号写成两个：'O''Brien'!A1；"
    "并集用英文逗号，不要用分号、中文逗号或空格；"
    "不要写 R1C1（请改 A1）；不要前置等号。"
)


class InvalidRefError(ValueError):
    """引用无法解析。消息必须含错误点与正确写法示例。"""

    def __init__(
        self,
        message: str,
        *,
        text: str = "",
        position: int | None = None,
    ) -> None:
        super().__init__(message)
        self.text = text
        self.position = position


def _fail(
    text: str,
    position: int | None,
    reason: str,
    example: str,
) -> NoReturn:
    loc = f"位置 {position}" if position is not None else "整段"
    raise InvalidRefError(
        f"无法解析引用 {text!r}：{reason}（错误点：{loc}）。"
        f"正确写法示例：{example}",
        text=text,
        position=position,
    )


@dataclass(frozen=True)
class CellRef:
    """单格。row/col 为 1-based；abs_* 对应 ``$`` 标记。"""

    sheet: str | None
    row: int
    col: int
    abs_row: bool = False
    abs_col: bool = False

    def __post_init__(self) -> None:
        if not (1 <= self.row <= EXCEL_MAX_ROW and 1 <= self.col <= EXCEL_MAX_COL):
            raise InvalidRefError(
                f"单元格越界 row={self.row} col={self.col}。"
                f"正确写法示例：A1（行 1–{EXCEL_MAX_ROW}，列 A–XFD）。",
            )

    def to_a1(self, *, include_sheet: bool = True) -> str:
        body = _format_cell(self.col, self.row, self.abs_col, self.abs_row)
        if include_sheet and self.sheet:
            return f"{_quote_sheet(self.sheet)}!{body}"
        return body


@dataclass(frozen=True)
class RectRef:
    """单矩形。行列 1-based，端点包含。"""

    sheet: str | None
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    abs_min_col: bool = False
    abs_min_row: bool = False
    abs_max_col: bool = False
    abs_max_row: bool = False
    whole_column: bool = False
    whole_row: bool = False

    def __post_init__(self) -> None:
        if not (
            1 <= self.min_row <= self.max_row <= EXCEL_MAX_ROW
            and 1 <= self.min_col <= self.max_col <= EXCEL_MAX_COL
        ):
            raise InvalidRefError(
                f"矩形越界或端点颠倒 "
                f"rows={self.min_row}:{self.max_row} cols={self.min_col}:{self.max_col}。"
                f"正确写法示例：A1:B2（端点包含，1-based）。",
            )

    def to_a1(self, *, include_sheet: bool = True) -> str:
        if self.whole_column:
            left = _format_col(self.min_col, self.abs_min_col)
            right = _format_col(self.max_col, self.abs_max_col)
            body = f"{left}:{right}"
        elif self.whole_row:
            left = _format_row(self.min_row, self.abs_min_row)
            right = _format_row(self.max_row, self.abs_max_row)
            body = f"{left}:{right}"
        elif self.min_row == self.max_row and self.min_col == self.max_col:
            body = _format_cell(
                self.min_col, self.min_row, self.abs_min_col, self.abs_min_row,
            )
        else:
            start = _format_cell(
                self.min_col, self.min_row, self.abs_min_col, self.abs_min_row,
            )
            end = _format_cell(
                self.max_col, self.max_row, self.abs_max_col, self.abs_max_row,
            )
            body = f"{start}:{end}"
        if include_sheet and self.sheet:
            return f"{_quote_sheet(self.sheet)}!{body}"
        return body

    def to_zero_based(self) -> ZeroBasedRect:
        """0-based **包含**端点，对齐 Excel 快照 ``merged.startRow`` / Univer。"""
        return ZeroBasedRect(
            sheet=self.sheet,
            start_row=self.min_row - 1,
            end_row=self.max_row - 1,
            start_col=self.min_col - 1,
            end_col=self.max_col - 1,
        )

    def top_left_cell(self) -> CellRef:
        return CellRef(
            sheet=self.sheet,
            row=self.min_row,
            col=self.min_col,
            abs_row=self.abs_min_row,
            abs_col=self.abs_min_col,
        )


@dataclass(frozen=True)
class NamedRef:
    """命名区域。工作簿 defined name 解析留给上层。"""

    name: str
    sheet: str | None = None

    def __post_init__(self) -> None:
        if not str(self.name or "").strip():
            raise InvalidRefError("命名区域名为空。正确写法示例：MyName 或 Sheet1!MyName。")

    def to_a1(self, *, include_sheet: bool = True) -> str:
        if include_sheet and self.sheet:
            return f"{_quote_sheet(self.sheet)}!{self.name}"
        return self.name


@dataclass(frozen=True)
class TableRef:
    """表对象引用。表/列元数据解析留给上层。"""

    table: str
    column: str | None = None
    sheet: str | None = None

    def __post_init__(self) -> None:
        if not str(self.table or "").strip():
            raise InvalidRefError(
                "表名为空。正确写法示例：Table1[列] 或 Table1[#All]。",
            )

    def to_a1(self, *, include_sheet: bool = True) -> str:
        body = self.table if self.column is None else f"{self.table}[{self.column}]"
        if include_sheet and self.sheet:
            return f"{_quote_sheet(self.sheet)}!{body}"
        return body


@dataclass(frozen=True)
class ZeroBasedRect:
    """0-based 包含端点。``start_row=0`` 对应 Excel 第 1 行。"""

    sheet: str | None
    start_row: int
    end_row: int
    start_col: int
    end_col: int


@dataclass(frozen=True)
class AreaRef:
    """一个或多个区域（并集）。"""

    areas: tuple[RectRef | NamedRef | TableRef, ...]

    def __post_init__(self) -> None:
        if not self.areas:
            raise InvalidRefError(
                "引用没有区域。正确写法示例：A1 或 A1:B2,C3:D4。",
            )

    def to_a1(self) -> str:
        return ",".join(part.to_a1() for part in self.areas)

    def sheets(self) -> tuple[str, ...]:
        seen: list[str] = []
        for part in self.areas:
            sheet = getattr(part, "sheet", None)
            if sheet and sheet not in seen:
                seen.append(sheet)
        return tuple(seen)

    @property
    def is_whole_column(self) -> bool:
        return bool(self.areas) and all(
            isinstance(part, RectRef) and part.whole_column for part in self.areas
        )

    @property
    def is_whole_row(self) -> bool:
        return bool(self.areas) and all(
            isinstance(part, RectRef) and part.whole_row for part in self.areas
        )

    def to_zero_based(self) -> tuple[ZeroBasedRect, ...]:
        rects: list[ZeroBasedRect] = []
        for part in self.areas:
            if not isinstance(part, RectRef):
                raise InvalidRefError(
                    f"{part.to_a1()!r} 不是矩形，无法换算 0-based。"
                    "命名区域/表引用请先在上层解析成 A1:B2。正确写法示例：A1:B2。",
                )
            rects.append(part.to_zero_based())
        return tuple(rects)

    def expand(self, used_rect: RectRef) -> AreaRef:
        """整轴按已用范围裁剪；有界矩形保持原样。"""
        return AreaRef(tuple(_expand_rect(part, used_rect) for part in self.areas))


def parse_ref(text: str, *, default_sheet: str | None = None) -> AreaRef:
    raw = text if isinstance(text, str) else str(text or "")
    stripped = raw.strip()
    if not stripped:
        _fail(raw, 1 if raw else None, "引用不能为空", "A1、$A$1、Sheet1!A1:B2")
    if stripped.startswith("="):
        _fail(
            raw,
            (raw.find("=") + 1) or 1,
            "不要前置等号",
            "A1 或 Sheet1!A1（不要写成 =A1）",
        )
    offset = raw.find(stripped) + 1
    default = str(default_sheet).strip() if default_sheet not in (None, "") else None
    parts: list[RectRef | NamedRef | TableRef] = []
    for pos, piece in _split_union(stripped, origin=offset):
        parts.append(_parse_one_area(raw, pos, piece, default))
    return AreaRef(tuple(parts))


def parse_rect(text: str, *, default_sheet: str | None = None) -> RectRef:
    area = parse_ref(text, default_sheet=default_sheet)
    if len(area.areas) != 1:
        _fail(
            text,
            None,
            "parse_rect 只接受单个矩形，不接受多区域",
            "A1:B2；多区域请用 parse_ref，例如 A1:B2,C3:D4",
        )
    item = area.areas[0]
    if not isinstance(item, RectRef):
        kind = "命名区域" if isinstance(item, NamedRef) else "表引用"
        _fail(
            text,
            None,
            f"parse_rect 只接受矩形，当前是{kind}",
            "A1:B2、$A$1、Sheet1!A:A",
        )
    return item


def describe_for_schema() -> dict[str, str]:
    """供后续单元生成工具 schema / 错误提示。键与文案保持稳定。"""
    return {
        "syntax": (
            "Excel A1 引用语法（1-based，端点包含）。"
            "单元格：A1、$A$1、$A1、A$1；"
            "矩形：A1:B2；"
            "跨表：Sheet1!A1:B2 或 'My Sheet'!A1；"
            "整列：A:A；整行：1:1；"
            "多区域：A1:B2,C3:D4 或 Sheet1!A1,Sheet2!B1（英文逗号）；"
            "命名区域：MyName（需工作簿定义）；"
            "表引用：Table1[列]、Table1[#All]（需表对象）。"
            "不支持 R1C1、三维引用 Sheet1:Sheet2!A1、外部工作簿 [Book.xlsx]Sheet1!A1。"
        ),
        "examples": "、".join(_REFERENCE_EXAMPLES),
        "common_errors": _SCHEMA_HINT,
        "execution": (
            "语法可解析不等于每个入口都能执行。缺省表：恰好一张表可省略 sheet，多于一张必须显式指定。"
            "range 读：单元格/矩形/整轴/并集/命名/表均可绑定。"
            "write.start_cell：单元格/矩形取左上角；命名/表绑定后取左上角；整轴与并集 REF_UNSUPPORTED。"
            "write.values 必须是矩阵，不能是 records。"
            "format.range：单元格/矩形/同表并集/整列整行（裁到已用范围）；merge/unmerge 只要一个矩形；命名/表 REF_UNSUPPORTED。"
            "Word 源表与 mention：单一矩形；命名绑定后可读；并集/整轴 REF_UNSUPPORTED。"
            "mention 禁止把名称扩成 Name:Name。"
        ),
    }


def reference_examples() -> tuple[str, ...]:
    return _REFERENCE_EXAMPLES


def _split_union(text: str, *, origin: int) -> list[tuple[int, str]]:
    areas: list[tuple[int, str]] = []
    buf: list[str] = []
    start = 0
    depth = 0
    in_quote = False
    i = 0
    while i < len(text):
        ch = text[i]
        if in_quote:
            buf.append(ch)
            if ch == "'":
                if i + 1 < len(text) and text[i + 1] == "'":
                    buf.append("'")
                    i += 2
                    continue
                in_quote = False
            i += 1
            continue
        if ch == "'":
            in_quote = True
            buf.append(ch)
            i += 1
            continue
        if ch == "[":
            depth += 1
            buf.append(ch)
            i += 1
            continue
        if ch == "]":
            depth = max(0, depth - 1)
            buf.append(ch)
            i += 1
            continue
        if depth == 0 and ch == "，":
            _fail(
                text,
                origin + i,
                "并集请用英文逗号 , 不要用中文逗号",
                "A1:B2,C3:D4",
            )
        if depth == 0 and ch == ";":
            _fail(
                text,
                origin + i,
                "并集请用逗号而不是分号",
                "A1:B2,C3:D4",
            )
        if depth == 0 and ch == ",":
            piece = "".join(buf).strip()
            if not piece:
                _fail(text, origin + start, "逗号旁缺少区域", "A1:B2,C3:D4")
            areas.append((origin + start, piece))
            buf = []
            i += 1
            while i < len(text) and text[i] in " \t":
                i += 1
            start = i
            continue
        buf.append(ch)
        i += 1
    if in_quote:
        _fail(text, origin + start, "表名单引号未闭合", "'My Sheet'!A1")
    if depth:
        _fail(text, origin + start, "方括号未闭合", "Table1[列] 或 Table1[#All]")
    piece = "".join(buf).strip()
    if not piece:
        _fail(
            text,
            origin + start if areas else origin,
            "末尾逗号无效" if areas else "引用不能为空",
            "A1:B2,C3:D4" if areas else "A1、$A$1、Sheet1!A1:B2",
        )
    areas.append((origin + start, piece))
    return areas


def _parse_one_area(
    original: str,
    pos: int,
    piece: str,
    default_sheet: str | None,
) -> RectRef | NamedRef | TableRef:
    sheet, rest, rest_pos = _extract_sheet(original, pos, piece)
    if sheet is not None and not sheet:
        _fail(original, pos, "表名为空", "Sheet1!A1 或 'My Sheet'!A1")
    if not rest:
        _fail(original, rest_pos, "表名后缺少引用", "Sheet1!A1 或 'My Sheet'!A1")
    if sheet is not None:
        _reject_bad_sheet(original, pos, sheet)
    if "!" in rest:
        _fail(
            original,
            rest_pos,
            "不支持三维引用或在区域中重复写表名",
            "Sheet1!A1:B2",
        )
    if rest.startswith("["):
        _fail(
            original,
            rest_pos,
            "不支持外部工作簿引用",
            "Sheet1!A1（不要写成 [Book.xlsx]Sheet1!A1）",
        )
    item = _parse_body(original, rest_pos, rest, sheet)
    if default_sheet and getattr(item, "sheet", None) in (None, ""):
        item = replace(item, sheet=default_sheet)
    return item


def _extract_sheet(original: str, pos: int, piece: str) -> tuple[str | None, str, int]:
    if piece.startswith("'"):
        return _extract_quoted_sheet(original, pos, piece)
    bang = piece.find("!")
    if bang < 0:
        return None, piece, pos
    sheet = piece[:bang].strip()
    rest = piece[bang + 1 :].strip()
    if not sheet:
        _fail(original, pos, "表名为空", "Sheet1!A1 或 'My Sheet'!A1")
    rest_offset = bang + 1
    while rest_offset < len(piece) and piece[rest_offset].isspace():
        rest_offset += 1
    return sheet, rest, pos + rest_offset


def _extract_quoted_sheet(
    original: str,
    pos: int,
    piece: str,
) -> tuple[str, str, int]:
    chars: list[str] = []
    i = 1
    while i < len(piece):
        ch = piece[i]
        if ch == "'":
            if i + 1 < len(piece) and piece[i + 1] == "'":
                chars.append("'")
                i += 2
                continue
            i += 1
            while i < len(piece) and piece[i].isspace():
                i += 1
            if i >= len(piece) or piece[i] != "!":
                _fail(
                    original,
                    pos,
                    "引号表名后必须是 !",
                    "'My Sheet'!A1",
                )
            rest = piece[i + 1 :].strip()
            return "".join(chars), rest, pos + i + 1
        chars.append(ch)
        i += 1
    _fail(original, pos, "表名单引号未闭合", "'My Sheet'!A1")


def _reject_bad_sheet(original: str, pos: int, sheet: str) -> None:
    if "[" in sheet or "]" in sheet:
        _fail(
            original,
            pos,
            "不支持外部工作簿引用",
            "Sheet1!A1（不要写成 [Book.xlsx]Sheet1!A1）",
        )
    if ":" in sheet:
        _fail(
            original,
            pos,
            "不支持三维引用 Sheet1:Sheet2!A1",
            "Sheet1!A1:B2",
        )


def _parse_body(
    original: str,
    pos: int,
    body: str,
    sheet: str | None,
) -> RectRef | NamedRef | TableRef:
    if _R1C1_TOKEN.match(body):
        _fail(
            original,
            pos,
            "不支持 R1C1 样式",
            "A1 或 $A$1；跨表用 Sheet1!A1",
        )
    if "[" in body:
        return _parse_table(original, pos, body, sheet)
    if _CELL_LIKE_LEADING_ZERO.match(body):
        _fail(original, pos, "行号不能有前导零", "A1 或 A10")
    rect = _try_parse_rect_body(original, pos, body, sheet)
    if rect is not None:
        return rect
    if _is_only_column_letters(body):
        col = body.upper()
        _fail(
            original,
            pos,
            f"{body!r} 不是完整引用",
            f"{col}1 或 {col}:{col}",
        )
    if body.isdigit():
        _fail(
            original,
            pos,
            f"{body!r} 不是完整引用",
            f"A{body} 或 {body}:{body}",
        )
    if _is_defined_name(body):
        return NamedRef(name=body, sheet=sheet)
    _fail(
        original,
        pos,
        f"{body!r} 不是合法 A1 / 命名区域 / 表引用",
        "A1、$A$1、Sheet1!A1:B2、MyName、Table1[列]",
    )


def _try_parse_rect_body(
    original: str,
    pos: int,
    body: str,
    sheet: str | None,
) -> RectRef | None:
    if ":" not in body:
        cell = _parse_cell_token(body, original=original, pos=pos)
        if cell is None:
            return None
        return _rect_from_cells(sheet, cell, cell, whole_column=False, whole_row=False)

    left, sep, right = body.partition(":")
    if not sep or ":" in right:
        _fail(original, pos, "区域只能有一个冒号", "A1:B2、A:A 或 1:1")
    left, right = left.strip(), right.strip()
    start = _parse_cell_token(left, original=original, pos=pos)
    end = _parse_cell_token(right, original=original, pos=pos)
    if start is not None and end is not None:
        return _rect_from_cells(sheet, start, end, whole_column=False, whole_row=False)

    col_match = _WHOLE_COL_TOKEN.match(body)
    if col_match:
        min_col = _col_index(original, pos, col_match.group(2))
        max_col = _col_index(original, pos, col_match.group(4))
        abs_min = bool(col_match.group(1))
        abs_max = bool(col_match.group(3))
        if min_col > max_col:
            min_col, max_col = max_col, min_col
            abs_min, abs_max = abs_max, abs_min
        return RectRef(
            sheet=sheet,
            min_row=1,
            max_row=EXCEL_MAX_ROW,
            min_col=min_col,
            max_col=max_col,
            abs_min_col=abs_min,
            abs_max_col=abs_max,
            whole_column=True,
        )
    row_match = _WHOLE_ROW_TOKEN.match(body)
    if row_match:
        min_row = _row_index(original, pos, row_match.group(2))
        max_row = _row_index(original, pos, row_match.group(4))
        abs_min = bool(row_match.group(1))
        abs_max = bool(row_match.group(3))
        if min_row > max_row:
            min_row, max_row = max_row, min_row
            abs_min, abs_max = abs_max, abs_min
        return RectRef(
            sheet=sheet,
            min_row=min_row,
            max_row=max_row,
            min_col=1,
            max_col=EXCEL_MAX_COL,
            abs_min_row=abs_min,
            abs_max_row=abs_max,
            whole_row=True,
        )
    if start is not None or end is not None:
        _fail(original, pos, "冒号两侧必须同为单元格、整列或整行", "A1:B2、A:A 或 1:1")
    return None


def _parse_cell_token(
    token: str,
    *,
    original: str | None = None,
    pos: int | None = None,
) -> CellRef | None:
    match = _CELL_TOKEN.match(token.strip())
    if not match:
        return None
    letters = match.group(2)
    row = int(match.group(4))
    try:
        col = int(column_index_from_string(letters.upper()))
    except ValueError:
        col = 0
    if not (1 <= col <= EXCEL_MAX_COL and 1 <= row <= EXCEL_MAX_ROW):
        if original is not None:
            _fail(
                original,
                pos,
                f"单元格 {token.upper()} 超出工作表极限",
                "A1 或 XFD1048576",
            )
        return None
    return CellRef(
        sheet=None,
        row=row,
        col=col,
        abs_col=bool(match.group(1)),
        abs_row=bool(match.group(3)),
    )


def _rect_from_cells(
    sheet: str | None,
    start: CellRef,
    end: CellRef,
    *,
    whole_column: bool,
    whole_row: bool,
) -> RectRef:
    min_col, max_col = start.col, end.col
    min_row, max_row = start.row, end.row
    abs_min_col, abs_max_col = start.abs_col, end.abs_col
    abs_min_row, abs_max_row = start.abs_row, end.abs_row
    if min_col > max_col:
        min_col, max_col = max_col, min_col
        abs_min_col, abs_max_col = abs_max_col, abs_min_col
    if min_row > max_row:
        min_row, max_row = max_row, min_row
        abs_min_row, abs_max_row = abs_max_row, abs_min_row
    return RectRef(
        sheet=sheet,
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        abs_min_col=abs_min_col,
        abs_min_row=abs_min_row,
        abs_max_col=abs_max_col,
        abs_max_row=abs_max_row,
        whole_column=whole_column,
        whole_row=whole_row,
    )


def _parse_table(
    original: str,
    pos: int,
    body: str,
    sheet: str | None,
) -> TableRef:
    open_at = body.find("[")
    if open_at <= 0:
        _fail(original, pos, "表引用方括号不完整", "Table1[列] 或 Table1[#All]")
    depth = 0
    close_at: int | None = None
    for index in range(open_at, len(body)):
        char = body[index]
        if char == "[":
            depth += 1
        elif char == "]":
            depth -= 1
            if depth == 0:
                close_at = index
                break
    if close_at is None:
        _fail(original, pos, "方括号未闭合", "Table1[列] 或 Table1[#All]")
    if close_at != len(body) - 1:
        _fail(original, pos, "表引用后有多余字符", "Table1[列] 或 Table1[#All]")
    table = body[:open_at].strip()
    inner = body[open_at + 1 : close_at].strip()
    if not _is_defined_name(table):
        _fail(original, pos, f"表名 {table!r} 不合法", "Table1[列]")
    if not inner:
        _fail(original, pos, "表引用括号内不能为空", "Table1[列] 或 Table1[#All]")
    return TableRef(table=table, column=inner, sheet=sheet)


def _col_index(original: str, pos: int, letters: str) -> int:
    try:
        col = int(column_index_from_string(letters.upper()))
    except ValueError:
        _fail(original, pos, f"列标 {letters!r} 不合法", "A:A 或 A1（列最大 XFD）")
    if not (1 <= col <= EXCEL_MAX_COL):
        _fail(original, pos, f"列 {letters.upper()} 超过 XFD", "A:A 或 XFD:XFD")
    return col


def _row_index(original: str, pos: int, raw: str) -> int:
    row = int(raw)
    if not (1 <= row <= EXCEL_MAX_ROW):
        _fail(original, pos, f"行号 {row} 越界", f"1:1（行最大 {EXCEL_MAX_ROW}）")
    return row


def _is_only_column_letters(text: str) -> bool:
    if not text.isalpha() or not (1 <= len(text) <= 3):
        return False
    try:
        col = int(column_index_from_string(text.upper()))
    except ValueError:
        return False
    return 1 <= col <= EXCEL_MAX_COL


def _is_defined_name(text: str) -> bool:
    if not text or len(text) > 255:
        return False
    if _CELL_TOKEN.match(text) or _R1C1_TOKEN.match(text):
        return False
    return _NAME_TOKEN.match(text) is not None


def _expand_rect(
    part: RectRef | NamedRef | TableRef,
    used: RectRef,
) -> RectRef:
    if not isinstance(part, RectRef):
        kind = "命名区域" if isinstance(part, NamedRef) else "表引用"
        raise InvalidRefError(
            f"无法展开 {part.to_a1()!r}：{kind}需要工作簿元数据。"
            "请先在上层解析成矩形，或改写为 A1:B2 / A:A。"
            "正确写法示例：A:A 配 used_rect=A1:C5 → A1:A5。",
        )
    if part.sheet and used.sheet and part.sheet != used.sheet:
        raise InvalidRefError(
            f"展开范围工作表不一致：引用 {part.sheet!r} 与已用范围 {used.sheet!r}。"
            "请使用同一工作表的 used_rect。正确写法示例：Sheet1!A:A 配 Sheet1 的已用范围。",
        )
    min_row, max_row = part.min_row, part.max_row
    min_col, max_col = part.min_col, part.max_col
    if part.whole_column:
        min_row = max(min_row, used.min_row)
        max_row = min(max_row, used.max_row)
    if part.whole_row:
        min_col = max(min_col, used.min_col)
        max_col = min(max_col, used.max_col)
    if min_row > max_row or min_col > max_col:
        raise InvalidRefError(
            f"引用 {part.to_a1()!r} 与已用范围 {used.to_a1()!r} 没有交集。"
            "请检查整列/整行是否超出已用范围。正确写法示例：A:A。",
        )
    return RectRef(
        sheet=part.sheet,
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        abs_min_col=part.abs_min_col,
        abs_min_row=part.abs_min_row,
        abs_max_col=part.abs_max_col,
        abs_max_row=part.abs_max_row,
        whole_column=False,
        whole_row=False,
    )


def _format_cell(col: int, row: int, abs_col: bool, abs_row: bool) -> str:
    return f"{_format_col(col, abs_col)}{_format_row(row, abs_row)}"


def _format_col(col: int, absolute: bool) -> str:
    return f"{'$' if absolute else ''}{get_column_letter(col)}"


def _format_row(row: int, absolute: bool) -> str:
    return f"{'$' if absolute else ''}{row}"


def _quote_sheet(name: str) -> str:
    if _sheet_needs_quotes(name):
        return "'" + name.replace("'", "''") + "'"
    return name


def _sheet_needs_quotes(name: str) -> bool:
    if not name or name[0].isdigit():
        return True
    for char in name:
        if char.isalnum() or char == "_" or "\u4e00" <= char <= "\u9fff":
            continue
        return True
    return False
