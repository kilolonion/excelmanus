"""V2 workbook document: validates design input and compiles one ChangeSet."""

from __future__ import annotations


import re


from datetime import date, datetime, time, timezone


from typing import Annotated, Any, Literal


from pydantic import (
    BaseModel,
    BeforeValidator,
    ConfigDict,
    Field,
    ValidationError,
    WithJsonSchema,
    field_validator,
    model_validator,
)


from excelmanus.workbook.layout import PrintLayout, apply_print_layout


_EXCEL_MAX_ROWS = 1_048_576


_EXCEL_MAX_COLS = 16_384


def _strip_a1_field(value: Any) -> Any:
    if isinstance(value, str):
        from excelmanus.workbook.address import strip_sheet_qualifier

        return strip_sheet_qualifier(value)
    return value


def _column_index(key: str) -> int | None:
    """"A"/"AA" → 1/27；无法解析返回 None。"""
    text = key.strip().upper()
    if not text or not text.isalpha():
        return None
    index = 0
    for char in text:
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index


_LEGAL_FILL_TYPES = frozenset(
    {
        "solid",
        "none",
        "darkDown",
        "darkGray",
        "darkGrid",
        "darkHorizontal",
        "darkTrellis",
        "darkUp",
        "darkVertical",
        "gray0625",
        "gray125",
        "lightDown",
        "lightGray",
        "lightGrid",
        "lightHorizontal",
        "lightTrellis",
        "lightUp",
        "lightVertical",
        "mediumGray",
    }
)


class SpecModel(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)


def _normalize_aliases(value: Any, aliases: dict[str, str]) -> Any:
    """把生态内已有的同义键归一到规范键；规范键显式出现时不覆盖。"""
    if not isinstance(value, dict):
        return value
    out = dict(value)
    for alias, canonical in aliases.items():
        if alias in out and canonical not in out:
            out[canonical] = out.pop(alias)
    return out


_FORMULA_CELL_SHAPE_HINT = (
    "公式矩阵单元格只接受字符串（= 开头为公式）、null（跳过该格，不写公式）"
    "或数字/布尔（转成字符串字面量写入该格）；空串与 null 同义（跳过该格）"
)

_VALUE_CELL_SHAPE_HINT = (
    "值矩阵单元格只接受 null（留空，空串同义）、数字、布尔、字符串（= 开头为公式）或日期"
)


def _coerce_formula_cell(value: Any) -> Any:
    """公式矩阵单元格归一化：null 跳过，数字/布尔转字符串字面量，字符串原样保留。"""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        # JSON 的 5.0 与 5 同义：转字面量不得留下 ".0"（5.0 → "5"）
        return str(int(value))
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, str):
        return value
    raise ValueError(_FORMULA_CELL_SHAPE_HINT)


def _coerce_value_cell(value: Any) -> Any:
    """值矩阵单元格只接受标量；容器/对象给出带形状说明的错误。"""
    if value is None or (isinstance(value, str) and value == ""):
        # null/空串都表示留空：不写该格，避免落盘后残留空字符串
        return None
    if isinstance(value, (str, int, float, bool, datetime, date, time)):
        return value
    raise ValueError(_VALUE_CELL_SHAPE_HINT)


# JSON schema 放宽到 string/null/number/boolean，与工具层 schema 校验保持同一合同。
FormulaCellValue = Annotated[
    str | None,
    BeforeValidator(_coerce_formula_cell),
    WithJsonSchema({"type": ["string", "null", "number", "boolean"], "description": _FORMULA_CELL_SHAPE_HINT}),
]

ValueCellValue = Annotated[
    Any,
    BeforeValidator(_coerce_value_cell),
    WithJsonSchema({"type": ["string", "null", "number", "boolean"], "description": _VALUE_CELL_SHAPE_HINT}),
]


class FontSpec(SpecModel):
    model_config = ConfigDict(extra="forbid")
    name: str | None = None
    size: float | None = None
    bold: bool | None = None
    italic: bool | None = None
    color: str | None = None  # 十六进制颜色值，如 "#FF0000"
    underline: str | None = None
    strike: bool | None = None

    @model_validator(mode="before")
    @classmethod
    def _normalize_font_input(cls, value: Any) -> Any:
        # 允许直接传字体名字符串（与提示文案一致）；strikethrough 是 ChangeSet 样式合同的同义键。
        if isinstance(value, str):
            return {"name": value}
        return _normalize_aliases(value, {"strikethrough": "strike"})


_FILL_ALIASES = {
    "fill_type": "type",
    "patternType": "type",
    "pattern": "type",
    "fgColor": "color",
    "fg_color": "color",
    "fgcolor": "color",
    "start_color": "color",
}


class FillSpec(SpecModel):
    model_config = ConfigDict(extra="forbid")
    type: str = "solid"
    color: str | None = None
    end_color: str | None = None

    @model_validator(mode="before")
    @classmethod
    def _normalize_fill_input(cls, value: Any) -> Any:
        # ChangeSet 样式合同（fill_type/patternType/fgColor/...）与 WorkbookSpec 同义；
        # 归一后语义不变，避免同名不同键导致 SPEC_VALIDATION_FAILED。
        return _normalize_aliases(value, _FILL_ALIASES)

    @field_validator("type")
    @classmethod
    def _reject_generic_pattern(cls, value: str) -> str:
        if value == "pattern":
            raise ValueError("fill.type=pattern 不是合法 patternType，请用 solid")
        if value not in _LEGAL_FILL_TYPES:
            raise ValueError(f"fill.type={value!r} 不是合法 patternType")
        return value


class BorderSideSpec(SpecModel):
    model_config = ConfigDict(extra="forbid")
    style: str | None = None
    color: str | None = None


class BorderSpec(SpecModel):
    model_config = ConfigDict(extra="forbid")
    style: str | None = None
    color: str | None = None
    top: BorderSideSpec | None = None
    bottom: BorderSideSpec | None = None
    left: BorderSideSpec | None = None
    right: BorderSideSpec | None = None



class AlignmentSpec(SpecModel):
    model_config = ConfigDict(extra="forbid")
    horizontal: str | None = None
    vertical: str | None = None
    wrap_text: bool | None = None
    text_rotation: int | None = None
    indent: float | None = None
    shrink_to_fit: bool | None = None

    @model_validator(mode="before")
    @classmethod
    def _normalize_alignment_input(cls, value: Any) -> Any:
        # Match ChangeSet style input and openpyxl's camelCase spellings.
        # The wire schema declares these aliases in get_tools(); fold them
        # before ``extra='forbid'`` performs the final shape check.
        return _normalize_aliases(value, {
            "horizontalAlignment": "horizontal",
            "verticalAlignment": "vertical",
            "wrapText": "wrap_text",
            "shrinkToFit": "shrink_to_fit",
            "textRotation": "text_rotation",
        })



class StyleClass(SpecModel):
    model_config = ConfigDict(extra="forbid")
    font: FontSpec | None = None
    fill: FillSpec | None = None
    border: BorderSpec | None = None
    alignment: AlignmentSpec | None = None
    number_format: str | None = None


class CellSpec(SpecModel):
    address: str
    value: Any = None
    value_type: Literal["string", "number", "date", "boolean", "formula", "empty"] = "string"
    display_text: str | None = None
    number_format: str | None = None
    formula_candidate: str | None = None
    style_id: str | None = None
    confidence: float = 1.0



class MergedRange(SpecModel):
    range: str
    confidence: float = 1.0

    @model_validator(mode="before")
    @classmethod
    def _normalize_merge_input(cls, value: Any) -> Any:
        # 提示文案承诺的简写："A1:B1" 等价于 {"range":"A1:B1"}。
        if isinstance(value, str):
            return {"range": value}
        return value



class FormulaPattern(SpecModel):
    column: str
    pattern: str
    confidence: float = 0.5


class SemanticHints(SpecModel):
    header_rows: list[int] = Field(default_factory=list)
    total_rows: list[int] = Field(default_factory=list)
    grouping: Any | None = None
    formula_patterns: list[FormulaPattern] = Field(default_factory=list)


class ConditionalFormatRule(SpecModel):
    """条件格式规则（颜色刻度/数据条/图标集/单元格值）。字段名是 type，不是 kind。"""

    type: Literal["color_scale", "data_bar", "icon_set", "cell_value"] = Field(
        description="WorkbookSpec 条件格式类型；与 apply_spreadsheet_changes 的 rule.type 不同名空间",
    )
    range: str = Field(description="Excel 范围，如 C2:C9")
    min_color: str | None = None
    mid_color: str | None = None
    max_color: str | None = None
    bar_color: str | None = None
    icon_style: str | None = Field(default=None, description="如 3_arrows / 3_traffic_lights")
    operator: str | None = Field(
        default=None,
        description="cell_value：greater_than/less_than/between/equal/not_equal",
    )
    value: Any = None
    value2: Any = None
    font_color: str | None = None
    fill_color: str | None = None
    bold: bool | None = None
    confidence: float = 0.7

    @field_validator("range", mode="before")
    @classmethod
    def _strip_range(cls, value: Any) -> Any:
        return _strip_a1_field(value)


class ObjectsSpec(SpecModel):
    charts: list[dict[str, Any]] = Field(default_factory=list)
    images: list[dict[str, Any]] = Field(default_factory=list)
    shapes: list[dict[str, Any]] = Field(default_factory=list)


class ValueBlock(SpecModel):
    """矩形值块：从 start 锚点展开的二维网格。"""

    start: str = Field(description="左上角 A1，如 A1")
    values: list[list[ValueCellValue]] = Field(
        default_factory=list,
        description="非空矩形二维数组，行等长；单元格为 null/数字/布尔/字符串（= 开头为公式）/日期",
    )

    @field_validator("start", mode="before")
    @classmethod
    def _strip_start(cls, value: Any) -> Any:
        return _strip_a1_field(value)


class FormulaBlock(SpecModel):
    """矩形公式块：从 start 锚点展开的二维公式网格。

    单元格元素语义：
    - "=..." 字符串：公式，写入该格；
    - null（或空串）：跳过该格，不写公式（合并单元格/留空用它表达）；
    - 数字/布尔：转成字符串字面量写入该格（5 → "5"，true → "TRUE"）；
    - 其他字符串：按字符串字面量写入该格。
    """

    start: str = Field(description="左上角 A1")
    formulas: list[list[FormulaCellValue]] = Field(
        default_factory=list,
        description=f"矩形二维数组，行等长；{_FORMULA_CELL_SHAPE_HINT}",
    )

    @field_validator("start", mode="before")
    @classmethod
    def _strip_start(cls, value: Any) -> Any:
        return _strip_a1_field(value)


class StyleRegion(SpecModel):
    """将具名样式应用到矩形区域。"""

    range: str = Field(description="A1 区域，须落在 dimensions 内")
    style_id: str = Field(description="必须是 styles 中已有的键")

    @field_validator("range", mode="before")
    @classmethod
    def _strip_range(cls, value: Any) -> Any:
        return _strip_a1_field(value)


class LayoutReference(SpecModel):
    attachment_id: str = Field(description="已观察图片的 attachment_id；坐标按该附件尺寸")
    table_bbox_px: list[float] = Field(min_length=4, max_length=4, description="表格边界 [x,y,width,height]，附件像素")
    column_edges: list[float] = Field(min_length=2, description="各列分隔位置，表格内部 0..1 相对坐标，从 0 到 1")
    row_edges: list[float] = Field(min_length=2, description="各行分隔位置，表格内部 0..1 相对坐标，从 0 到 1")
    target_width_px: float = Field(gt=0, le=8192, description="工作台目标总宽度，像素估算")

    @model_validator(mode="after")
    def validate_edges(self):
        for edges in (self.column_edges, self.row_edges):
            if edges[0] != 0 or edges[-1] != 1 or any(a >= b for a, b in zip(edges, edges[1:])):
                raise ValueError("Layout edges must increase strictly from 0 to 1")
        x, y, width, height = self.table_bbox_px
        if x < 0 or y < 0 or width <= 0 or height <= 0:
            raise ValueError("Invalid table bounding box")
        return self


class SheetSpec(SpecModel):
    name: str = Field(description="工作表名")
    dimensions: dict[str, int] | None = Field(
        default=None,
        description="规范为 {rows, cols}；带 source_csv 时可省略",
    )
    source_csv: dict[str, Any] | None = Field(
        default=None,
        description="{file_path, encoding?, skip_rows?, start?} 导入工作区 CSV/TSV",
    )
    freeze_panes: str | None = None
    print_layout: PrintLayout | None = None
    layout_reference: LayoutReference | None = None
    cells: list[CellSpec] = Field(default_factory=list, description="单格；公式用 value='=A1' 且 value_type=formula")
    value_blocks: list[ValueBlock] = Field(default_factory=list)
    formula_blocks: list[FormulaBlock] = Field(default_factory=list)
    style_regions: list[StyleRegion] = Field(default_factory=list)
    merged_ranges: list[MergedRange] = Field(default_factory=list)
    styles: dict[str, StyleClass] = Field(
        default_factory=dict,
        description="style_id → 样式；查询 additionalProperties 得 font/fill/border",
    )
    column_widths: list[float] = Field(
        default_factory=list,
        description=(
            "从 A 列起的 Excel 列宽数组；单位为 Excel 字符宽度，不是像素。"
            "图片复刻时按图中各列边界和相对宽度设置；省略时使用默认列宽。"
        ),
    )
    row_heights: dict[str, float] = Field(
        default_factory=dict,
        description=(
            "行号字符串 → Excel 行高；单位为 points，不是像素。"
            "图片复刻时按图中各行边界和相对高度设置；省略时使用默认行高。"
        ),
    )
    conditional_formats: list[ConditionalFormatRule] = Field(default_factory=list)
    objects: ObjectsSpec = Field(default_factory=ObjectsSpec)
    semantic_hints: SemanticHints = Field(default_factory=SemanticHints)

    @model_validator(mode="before")
    @classmethod
    def _normalize_size_input(cls, value: Any) -> Any:
        # 提示文案承诺的简写：column_widths={"A":18,"B":12}、row_heights=[22,15]。
        if not isinstance(value, dict):
            return value
        out = dict(value)
        widths = out.get("column_widths")
        if isinstance(widths, dict):
            ordered: dict[int, float] = {}
            for key, width in widths.items():
                index = _column_index(str(key))
                if index is None:
                    return out  # 无法解析的键交给字段校验报带路径的错
                ordered[index] = width
            out["column_widths"] = [ordered[i] for i in sorted(ordered)]
        heights = out.get("row_heights")
        if isinstance(heights, list):
            out["row_heights"] = {str(i + 1): h for i, h in enumerate(heights)}
        return out

    @field_validator("freeze_panes", mode="before")
    @classmethod
    def _strip_freeze(cls, value: Any) -> Any:
        return _strip_a1_field(value)





class Uncertainty(SpecModel):
    location: str = Field(description="不确定项位置，如 Sheet1!B2")
    reason: str = Field(description="看不清或无法确定的原因")
    candidate_values: list[str] = Field(
        default_factory=list,
        description="候选字符串列表，数字写成 \"4800\"",
    )
    confidence: float = 0.5

    @model_validator(mode="before")
    @classmethod
    def _normalize_uncertainty_input(cls, value: Any) -> Any:
        # 模型常用 field/note、where/message 等同义键描述不确定项；归一到 location/reason。
        return _normalize_aliases(value, {
            "field": "location",
            "where": "location",
            "cell": "location",
            "address": "location",
            "target": "location",
            "position": "location",
            "note": "reason",
            "message": "reason",
            "detail": "reason",
            "description": "reason",
            "comment": "reason",
            "cause": "reason",
            "candidates": "candidate_values",
        })

    @field_validator("candidate_values", mode="before")
    @classmethod
    def _stringify_candidate_values(cls, value: Any) -> Any:
        if not isinstance(value, list):
            return value
        items: list[str] = []
        for item in value:
            if item is None:
                continue
            items.append(str(item))
        return items


class WorkbookSpec(SpecModel):
    """P5 最小工作簿规格：当前模型直接产出，经校验后编译。

    ``uncertainties`` 必填（可为空列表，但不能缺字段）。
    """

    version: Literal["2"] = "2"
    name: str = "replica"
    locale: str | None = None
    default_font: FontSpec | None = Field(
        default=None,
        description="默认字体 {name, size?}",
    )
    theme_hint: str | None = None
    purpose: Literal["data", "visual_replica"] = "data"
    sheets: list[SheetSpec] = Field(description="工作表列表")
    uncertainties: list[Uncertainty] = Field(description="必须出现；没有不确定项时为 []")



class SpecValidationError(ValueError):
    """带字段路径的 WorkbookSpec 校验失败。"""

    def __init__(self, errors: list[dict[str, str]]):
        self.errors = errors
        summary = "; ".join(f"{item['path']}: {item['message']}" for item in errors) or "WorkbookSpec 校验失败"
        super().__init__(summary)

    def to_payload(self) -> dict[str, Any]:
        from excelmanus.engine_core.error_payload import SPEC_VALIDATION_FAILED, make_error_payload

        return make_error_payload(
            "WorkbookSpec 校验失败，请按字段路径修正后重试。",
            error_code=SPEC_VALIDATION_FAILED,
            errors=self.errors,
            shape="WorkbookSpec",
        )


def format_error_path(loc: tuple[Any, ...]) -> str:
    if not loc:
        return "$"
    return ".".join(str(part) for part in loc)


def _friendly_spec_message(path: str, message: str) -> str:
    lower = message.lower()
    if lower.startswith("value error, ") and "只接受" in message:
        # 自定义形状提示去掉 pydantic 的 "Value error, " 前缀，直接呈现合法形状
        return message[len("Value error, "):]
    if "default_font" in path and ("dict" in lower or "font" in lower):
        return 'default_font 必须是 {"name":"微软雅黑","size":11}，或直接传字体名字符串'
    if path.endswith("column_widths") and "list" in lower:
        return 'column_widths 必须是从 A 列起的数字数组 [18,12]，或 {"A":18,"B":12}'
    if path.endswith("row_heights") and "dict" in lower:
        return 'row_heights 必须是 {"1":22} 或从第 1 行起的数字数组 [22,15]'
    if path.endswith("merged_ranges") or ".merged_ranges." in path:
        if "dict" in lower:
            return 'merged_ranges 项必须是 {range:"A1:B1"} 或 "A1:B1"'
    if path == "sheets" and "list" in lower:
        return 'sheets 必须是数组，例如 [{"name":"Sheet1","dimensions":{"rows":10,"cols":5}}]'
    if ".formulas" in path and "list" in lower:
        return f"formulas 必须是行等长的二维数组；{_FORMULA_CELL_SHAPE_HINT}"
    if ".values" in path and "list" in lower:
        return f"values 必须是行等长的二维数组；{_VALUE_CELL_SHAPE_HINT}"
    if ".formulas" in path and ("string" in lower or "valid" in lower):
        return _FORMULA_CELL_SHAPE_HINT
    if "extra inputs" in lower or "not permitted" in lower:
        key = path.rsplit(".", 1)[-1]
        return f"{_extra_field_hint(path)}（不支持的字段 {key}）"
    if "field required" in lower or lower.strip() == "field required":
        if path == "uncertainties":
            return "uncertainties 必填；没有不确定项时传 []"
        if path.startswith("uncertainties."):
            return "uncertainties 项必填 location（位置）与 reason（原因）；field/note 等同义键也可用"
        return f"缺少必填字段：{path.rsplit('.', 1)[-1]}"
    if path.endswith("uncertainties") and ("required" in lower or "field required" in lower):
        return "uncertainties 必填；没有不确定项时传 []"
    return message


_EXTRA_FIELD_LAYER_HINTS = (
    (("fill",), "fill 只接受 {type, color, end_color}；fill_type/patternType/pattern→type，fgColor/fg_color/fgcolor/start_color→color"),
    (("font",), "font 只接受 {name, size, bold, italic, color, underline, strike}；strikethrough→strike"),
    (("border", "top", "bottom", "left", "right"), "border 及其单边只接受 {style, color}"),
    (("alignment",), "alignment 只接受 {horizontal, vertical, wrap_text, text_rotation, indent, shrink_to_fit}"),
    (("uncertainties",), "uncertainties 项只接受 {location, reason, candidate_values, confidence}；field/where/cell→location，note/message→reason"),
)


def _extra_field_hint(path: str) -> str:
    segments = path.split(".")
    for names, hint in _EXTRA_FIELD_LAYER_HINTS:
        if any(segment in names for segment in segments[:-1]):
            return hint
    parent = path.rsplit(".", 1)[0]
    return (
        f"该层只接受 schema 定义的字段；"
        f"可查 introspect_capability(query_type='tool_detail', query='apply_spreadsheet_changes.workbook_spec.{parent}')"
    )


def _pydantic_errors(exc: ValidationError) -> list[dict[str, str]]:
    return [
        {
            "path": format_error_path(tuple(err.get("loc") or ())),
            "message": _friendly_spec_message(
                format_error_path(tuple(err.get("loc") or ())),
                str(err.get("msg") or "校验失败"),
            ),
        }
        for err in exc.errors()
    ]


def _parse_a1(address: str) -> tuple[int, int]:
    """返回 1-indexed (row, col)。"""
    from excelmanus.workbook.refs import CellRef, RectRef, parse_ref

    area = parse_ref(address)
    part = area.areas[0]
    if isinstance(part, CellRef):
        return part.row, part.col
    if isinstance(part, RectRef):
        if part.whole_column or part.whole_row:
            raise ValueError("WorkbookSpec 不支持整轴地址，请写有限单元格如 A1")
        return part.min_row, part.min_col
    raise ValueError(f"WorkbookSpec 地址必须是单元格或矩形：{address!r}")


def _range_bounds(range_str: str) -> tuple[int, int, int, int]:
    """返回 (min_col, min_row, max_col, max_row)，均为 1-indexed。"""
    from excelmanus.workbook.refs import CellRef, RectRef, parse_ref

    area = parse_ref(range_str)
    if len(area.areas) != 1:
        raise ValueError("WorkbookSpec 不支持并集 range")
    part = area.areas[0]
    if isinstance(part, CellRef):
        return part.col, part.row, part.col, part.row
    if isinstance(part, RectRef):
        if part.whole_column or part.whole_row:
            raise ValueError("WorkbookSpec 不支持整轴 range，请写有限矩形")
        return part.min_col, part.min_row, part.max_col, part.max_row
    raise ValueError(f"WorkbookSpec range 必须是单元格或矩形：{range_str!r}")


def _sheet_dimensions(sheet: SheetSpec, sheet_index: int) -> tuple[int, int] | dict[str, str]:
    raw = sheet.dimensions or {}
    try:
        rows = int(raw["rows"])
        cols = int(raw["cols"])
    except (KeyError, TypeError, ValueError):
        return {
            "path": f"sheets.{sheet_index}.dimensions",
            "message": "dimensions 必须包含正整数 rows 与 cols",
        }
    if rows < 1 or cols < 1:
        return {
            "path": f"sheets.{sheet_index}.dimensions",
            "message": f"dimensions 越界: rows={rows}, cols={cols}（须 ≥ 1）",
        }
    if rows > _EXCEL_MAX_ROWS or cols > _EXCEL_MAX_COLS:
        return {
            "path": f"sheets.{sheet_index}.dimensions",
            "message": (
                f"dimensions 超出 Excel 上限: rows={rows}, cols={cols} "
                f"（最大 {_EXCEL_MAX_ROWS}×{_EXCEL_MAX_COLS}）"
            ),
        }
    return rows, cols


def _rect_fits(
    *,
    start_row: int,
    start_col: int,
    height: int,
    width: int,
    rows: int,
    cols: int,
) -> str | None:
    if start_row < 1 or start_col < 1:
        return f"锚点越界: row={start_row}, col={start_col}"
    if height <= 0 or width <= 0:
        return None
    end_row = start_row + height - 1
    end_col = start_col + width - 1
    if end_row > rows or end_col > cols:
        return (
            f"矩形超出 dimensions {rows}×{cols}: "
            f"覆盖 ({start_row},{start_col})–({end_row},{end_col})"
        )
    return None


def collect_layout_errors(spec: WorkbookSpec) -> list[dict[str, str]]:
    """矩形越界、坏样式引用等语义错误（带字段路径）。"""
    errors: list[dict[str, str]] = []
    for i, sheet in enumerate(spec.sheets):
        # source_csv 的 dimensions 由工具层展开后推导，展开前无法做矩形边界检查
        if sheet.source_csv and sheet.dimensions is None:
            continue
        dims = _sheet_dimensions(sheet, i)
        if isinstance(dims, dict):
            errors.append(dims)
            continue
        rows, cols = dims
        estimated_cells = len(sheet.cells) + sum(sum(len(row) for row in block.values) for block in sheet.value_blocks) + sum(sum(len(row) for row in block.formulas) for block in sheet.formula_blocks)
        for region in sheet.style_regions:
            try:
                c0,r0,c1,r1 = _range_bounds(region.range)
                estimated_cells += (c1-c0+1)*(r1-r0+1)
            except ValueError:
                pass
        if estimated_cells > 200000:
            errors.append({"path":f"sheets.{i}","message":"Spec compilation exceeds 200000 cells; use bounded ChangeSets or streaming import"})
            continue
        prefix = f"sheets.{i}"

        for j, block in enumerate(sheet.value_blocks):
            path = f"{prefix}.value_blocks.{j}"
            try:
                start_row, start_col = _parse_a1(block.start)
            except Exception as exc:
                errors.append({"path": f"{path}.start", "message": f"无效锚点 {block.start!r}: {exc}"})
                continue
            height = len(block.values)
            width = max((len(row) for row in block.values), default=0)
            msg = _rect_fits(
                start_row=start_row, start_col=start_col,
                height=height, width=width, rows=rows, cols=cols,
            )
            if msg:
                errors.append({"path": path, "message": msg})

        for j, block in enumerate(sheet.formula_blocks):
            path = f"{prefix}.formula_blocks.{j}"
            try:
                start_row, start_col = _parse_a1(block.start)
            except Exception as exc:
                errors.append({"path": f"{path}.start", "message": f"无效锚点 {block.start!r}: {exc}"})
                continue
            height = len(block.formulas)
            width = max((len(row) for row in block.formulas), default=0)
            msg = _rect_fits(
                start_row=start_row, start_col=start_col,
                height=height, width=width, rows=rows, cols=cols,
            )
            if msg:
                errors.append({"path": path, "message": msg})

        for j, region in enumerate(sheet.style_regions):
            path = f"{prefix}.style_regions.{j}"
            if region.style_id not in sheet.styles:
                errors.append({
                    "path": f"{path}.style_id",
                    "message": f"未知样式引用: {region.style_id!r}",
                })
            try:
                min_col, min_row, max_col, max_row = _range_bounds(region.range)
            except Exception as exc:
                errors.append({"path": f"{path}.range", "message": f"无效区域 {region.range!r}: {exc}"})
                continue
            if min_row < 1 or min_col < 1 or max_row > rows or max_col > cols:
                errors.append({
                    "path": f"{path}.range",
                    "message": (
                        f"样式区域 {region.range} 超出 dimensions {rows}×{cols}"
                    ),
                })

        for j, merged in enumerate(sheet.merged_ranges):
            path = f"{prefix}.merged_ranges.{j}"
            try:
                min_col, min_row, max_col, max_row = _range_bounds(merged.range)
            except Exception as exc:
                errors.append({"path": f"{path}.range", "message": f"无效合并区 {merged.range!r}: {exc}"})
                continue
            if min_row < 1 or min_col < 1 or max_row > rows or max_col > cols:
                errors.append({
                    "path": f"{path}.range",
                    "message": f"合并区 {merged.range} 超出 dimensions {rows}×{cols}",
                })

        for j, cell in enumerate(sheet.cells):
            path = f"{prefix}.cells.{j}"
            try:
                row, col = _parse_a1(cell.address)
            except Exception as exc:
                errors.append({"path": f"{path}.address", "message": f"无效地址 {cell.address!r}: {exc}"})
                continue
            if row > rows or col > cols or row < 1 or col < 1:
                errors.append({
                    "path": f"{path}.address",
                    "message": f"单元格 {cell.address} 超出 dimensions {rows}×{cols}",
                })
            if cell.style_id and cell.style_id not in sheet.styles:
                errors.append({
                    "path": f"{path}.style_id",
                    "message": f"未知样式引用: {cell.style_id!r}",
                })
    return errors


def validate_workbook_spec(data: Any) -> WorkbookSpec:
    """Validate a V2 WorkbookSpec object without coercing legacy wire forms."""
    if isinstance(data, str):
        raise SpecValidationError([{
            "path": "spec",
            "message": "WorkbookSpec 必须是对象；JSON 字符串不是 V2 输入",
        }])
    if not isinstance(data, dict):
        raise SpecValidationError([{
            "path": "spec",
            "message": "WorkbookSpec 必须是 JSON 对象",
        }])
    try:
        spec = WorkbookSpec.model_validate(data)
    except ValidationError as exc:
        raise SpecValidationError(_pydantic_errors(exc)) from exc
    layout_errors = collect_layout_errors(spec)
    if layout_errors:
        raise SpecValidationError(layout_errors)
    return spec


def iter_block_cells(start: str, grid: list[list[Any]]) -> list[tuple[str, Any]]:
    from openpyxl.utils import get_column_letter

    start_row, start_col = _parse_a1(start)
    out: list[tuple[str, Any]] = []
    for r_off, row_vals in enumerate(grid):
        if not isinstance(row_vals, list):
            continue
        for c_off, value in enumerate(row_vals):
            addr = f"{get_column_letter(start_col + c_off)}{start_row + r_off}"
            out.append((addr, value))
    return out


def iter_range_addresses(range_str: str) -> list[str]:
    from openpyxl.utils import get_column_letter

    min_col, min_row, max_col, max_row = _range_bounds(range_str)
    addrs: list[str] = []
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            addrs.append(f"{get_column_letter(col)}{row}")
    return addrs


_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


_ISO_DATETIME_RE = re.compile(
    r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})?$"
)


def _coerce_iso_temporal(value: str) -> datetime | date | None:
    """严格 ISO-8601 字符串 → date/datetime（与 Excel 输入行为一致）。"""
    text = value.strip()
    try:
        if _ISO_DATE_RE.match(text):
            return date.fromisoformat(text)
        if _ISO_DATETIME_RE.match(text):
            return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None
    return None


def materialize_sheet_cells(sheet: SheetSpec) -> list[CellSpec]:
    """把矩形值/公式块与样式区域展开为单元格列表，供 rebuild 编译。"""
    by_addr: dict[str, CellSpec] = {}
    for cell in sheet.cells:
        by_addr[cell.address.upper()] = cell.model_copy(deep=True)

    for block in sheet.value_blocks:
        for addr, value in iter_block_cells(block.start, block.values):
            key = addr.upper()
            existing = by_addr.get(key)
            value_type: Literal["string", "number", "date", "boolean", "formula", "empty"]
            if value is None:
                value_type = "empty"
            elif isinstance(value, bool):
                value_type = "boolean"
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                value_type = "number"
            elif isinstance(value, (datetime, date, time)):
                value_type = "date"
            elif isinstance(value, str) and value.startswith("="):
                value_type = "formula"
            elif isinstance(value, str) and (_temporal := _coerce_iso_temporal(value)) is not None:
                value = _temporal
                value_type = "date"
            else:
                value_type = "string"
            if existing is None:
                by_addr[key] = CellSpec(address=addr, value=value, value_type=value_type)
            else:
                existing.value = value
                existing.value_type = value_type

    for block in sheet.formula_blocks:
        for addr, formula in iter_block_cells(block.start, list(block.formulas)):
            if formula is None or formula == "":
                # null/空串：跳过该格，不写公式也不覆盖已有单元格
                continue
            key = addr.upper()
            text = str(formula)
            value_type: Literal["formula", "string"] = "formula" if text.startswith("=") else "string"
            existing = by_addr.get(key)
            if existing is None:
                by_addr[key] = CellSpec(address=addr, value=text, value_type=value_type)
            else:
                existing.value = text
                existing.value_type = value_type

    for region in sheet.style_regions:
        try:
            addrs = iter_range_addresses(region.range)
        except Exception:
            continue
        for addr in addrs:
            key = addr.upper()
            existing = by_addr.get(key)
            if existing is None:
                by_addr[key] = CellSpec(
                    address=addr, value=None, value_type="empty", style_id=region.style_id,
                )
            else:
                existing.style_id = region.style_id

    return list(by_addr.values())


def workbook_spec_json_schema() -> dict[str, Any]:
    """从 Pydantic 模型生成供工具 schema / 字段查询使用的结构。"""
    schema = WorkbookSpec.model_json_schema(mode="validation")
    from excelmanus.tools.workbook_examples import workbook_creation_example

    schema["examples"] = [workbook_creation_example()]
    schema["description"] = (
        "创建用 WorkbookSpec，与 operations 互斥。必填 sheets 与 uncertainties。"
        "V2 只接受规范对象；尺寸和样式与统一 ChangeSet 共用执行语义。"
        "图片/版式还原类创建请显式传 purpose=visual_replica（要求完整行高列宽或 layout_reference，"
        "并启用视觉核验义务）；默认 purpose=data 只表示数据抽取。"
    )
    return schema


def document_operations(spec: WorkbookSpec) -> list[dict[str, Any]]:
    """Compile a document into the same operations used for existing files."""
    from openpyxl.utils import get_column_letter
    # Declare every sheet before populating data, then bind drawings. A chart
    # on an earlier sheet may target a later sheet without implicitly creating
    # it twice; whole-column references also see the final data dimensions.
    operations: list[dict[str, Any]] = [
        {"kind": "sheet", "action": "create", "new_name": sheet.name} for sheet in spec.sheets
    ]
    object_operations: list[dict[str, Any]] = []
    for sheet in spec.sheets:
        resolve_document_geometry(sheet, spec.purpose)
        materialized = materialize_sheet_cells(sheet)
        for cell in materialized:
            if cell.value_type == "string" and isinstance(cell.value, str) and cell.value.startswith("="):
                raise ValueError("Literal strings beginning with = require an explicit text cell operation")
            if cell.value_type == "date" and isinstance(cell.value, str):
                converted = _coerce_iso_temporal(cell.value)
                if converted is None:
                    raise ValueError(f"{sheet.name}!{cell.address}: date must be ISO-8601")
                cell.value = converted
            if cell.value_type == "formula" and not str(cell.value or "").startswith("="):
                raise ValueError(f"{sheet.name}!{cell.address}: formula must start with =")
            if cell.value is not None:
                operations.append({"kind": "write", "sheet": sheet.name, "start_cell": cell.address,
                                   "values": [[cell.value]]})
            style = sheet.styles.get(cell.style_id) if cell.style_id else None
            fields = style.model_dump(exclude_none=True) if style else {}
            if spec.default_font:
                fields["font"] = {**spec.default_font.model_dump(exclude_none=True), **fields.get("font", {})}
            if cell.number_format:
                fields["number_format"] = cell.number_format
            if fields:
                operations.append({"kind": "format", "sheet": sheet.name, "range": cell.address, **fields})
        for merged in sheet.merged_ranges:
            operations.append({"kind": "merge", "sheet": sheet.name, "range": merged.range})
        if sheet.column_widths or sheet.row_heights:
            operations.append({"kind": "size", "sheet": sheet.name,
                "column_widths": {get_column_letter(i + 1): width for i, width in enumerate(sheet.column_widths)},
                "row_heights": sheet.row_heights})
        if sheet.freeze_panes:
            operations.append({"kind": "freeze", "sheet": sheet.name, "freeze_panes": sheet.freeze_panes})
        if sheet.print_layout:
            operations.append({"kind": "print_layout", "sheet": sheet.name,
                               "print_layout": sheet.print_layout.model_dump(exclude_none=True)})
        for rule in sheet.conditional_formats:
            cfg = rule.model_dump(exclude_none=True)
            cfg.pop("confidence", None)
            address = cfg.pop("range")
            if cfg.get("fill_color"):
                cfg["fill"] = {"color": cfg.pop("fill_color")}
            font = {k: cfg.pop(key) for key, k in (("font_color", "color"), ("bold", "bold")) if key in cfg}
            if font:
                cfg["font"] = font
            operations.append({"kind": "conditional_format", "sheet": sheet.name, "range": address, "rule": cfg})
        if sheet.objects.shapes:
            raise ValueError("Drawing shapes are unsupported; no file was created")
        for kind, items in (("chart", sheet.objects.charts), ("image", sheet.objects.images)):
            for item in items:
                if item.get("kind", kind) != kind or item.get("sheet", sheet.name) != sheet.name:
                    raise ValueError("Document object kind/sheet conflicts with its container")
                object_operations.append({**item, "kind":kind, "sheet":sheet.name})
    return operations + object_operations


def resolve_document_geometry(sheet: SheetSpec, purpose: str) -> None:
    from excelmanus.workbook.geometry import column_native
    ref = sheet.layout_reference
    dims = sheet.dimensions or {}
    rows, cols = dims.get("rows", 0), dims.get("cols", 0)
    if ref:
        from excelmanus.attachments.store import get_attachment_store
        from excelmanus.tools.context import require_call
        if ref.attachment_id not in require_call().durable_attachment_ids:
            raise ValueError("Layout reference must identify an image already observed in this session")
        image = get_attachment_store().get_ref(ref.attachment_id)
        if image is None:
            raise ValueError("Layout reference image is unavailable")
        x, y, width, height = ref.table_bbox_px
        if x + width > image.width or y + height > image.height:
            raise ValueError("Table bounds exceed the reference attachment")
        if len(ref.column_edges) != cols + 1 or len(ref.row_edges) != rows + 1:
            raise ValueError("Layout edges must cover every declared row and column")
        if sheet.column_widths or sheet.row_heights:
            raise ValueError("Specify layout_reference or native sizes, not both")
        sheet.column_widths = [column_native((b - a) * ref.target_width_px) for a, b in zip(ref.column_edges, ref.column_edges[1:])]
        target_height = ref.target_width_px * height / width
        sheet.row_heights = {str(i + 1): (b - a) * target_height * 72 / 96 for i, (a, b) in enumerate(zip(ref.row_edges, ref.row_edges[1:]))}
    if purpose == "visual_replica" and (len(sheet.column_widths) != cols or any(str(i) not in sheet.row_heights for i in range(1, rows + 1))):
        raise ValueError("Visual replica requires complete row/column geometry or a layout_reference; data extraction may use purpose=data")
