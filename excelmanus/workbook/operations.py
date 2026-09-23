"""Range operations shared by native edit and Code Mode; never evaluate Python."""
from copy import copy
from datetime import date, datetime, timedelta
from numbers import Number
import re

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange


KINDS = {
    "sort": "range by ascending header visible_only",
    "fill": "range source_range mode value start step visible_only",
    "replace": "range find replacement regex match_case include_formulas",
    "clear": "range mode",
    "paste_special": "source_sheet source_range target_start mode transpose",
    "append": "values columns table start_cell",
}


def schema_fields():
    return {
        "range": {"type": "string"}, "by": {"type": ["array", "string"], "items": {}},
        "ascending": {"type": ["boolean", "array"], "items": {"type": "boolean"}},
        "header": {"type": "boolean"}, "visible_only": {"type": "boolean"},
        "mode": {"type": "string", "description": "clear: values/formats/all; paste_special: values/formulas/formats/all; fill: copy/series/value"},
        "find": {}, "replacement": {}, "regex": {"type": "boolean"},
        "match_case": {"type": "boolean"}, "include_formulas": {"type": "boolean"},
        "transpose": {"type": "boolean"}, "table": {"type": "string"}, "step": {"type": "number"},
    }


def _bounds(ws, raw):
    from excelmanus.workbook.address import resolve_range_to_bounds
    b = resolve_range_to_bounds(str(raw), used_max_row=ws.max_row, used_max_col=ws.max_column)
    target = CellRange(min_col=b.min_col, min_row=b.min_row, max_col=b.max_col, max_row=b.max_row)
    if target.size["rows"] * target.size["columns"] > 2_000_000:
        raise ValueError("一次区域操作最多 200 万单元格，请拆分范围")
    if any(not target.isdisjoint(m) for m in ws.merged_cells.ranges):
        raise ValueError("区域含合并单元格，请先取消合并或选择独立数据区域")
    return target


def _paste(source, target, mode="all", *, transpose=False):
    if mode in {"all", "formulas"}:
        value = source.value
        if source.data_type == "f":
            value = Translator(value, origin=source.coordinate).translate_formula(target.coordinate)
        target.value = value
        if source.data_type != "f":
            target.data_type = source.data_type
    if mode in {"all", "formats"}:
        target._style = copy(source._style)
    if mode == "all":
        target.comment = copy(source.comment)
        target.hyperlink = copy(source.hyperlink)


def apply_range_operation(wb, op):
    name = op.get("sheet") or op.get("sheet_name")
    if not name:
        if len(wb.worksheets) != 1:
            raise ValueError("需要明确指定 sheet")
        name = wb.active.title
    ws = wb[name]
    kind = op["kind"]
    changed = 0
    if kind == "append":
        rows = op.get("values")
        if not isinstance(rows, list) or not rows or not all(isinstance(r, list) for r in rows) or len({len(r) for r in rows}) != 1:
            raise ValueError("append.values 需要非空矩形二维数组")
        table = ws.tables.get(op.get("table")) if op.get("table") else None
        if op.get("table") and table is None:
            raise ValueError("Table 不存在")
        if table:
            region = CellRange(table.ref)
            if table.totalsRowCount:
                raise ValueError("带汇总行的 Table 请先移除汇总行再追加")
            start_row, start_col = region.max_row + 1, region.min_col
            if len(rows[0]) != region.max_col - region.min_col + 1:
                raise ValueError("追加列数与 Table 不符")
        elif op.get("start_cell"):
            anchor = ws[op["start_cell"]]
            start_row, start_col = anchor.row, anchor.column
        else:
            start_row, start_col = ws.max_row + (1 if any(c.value is not None for c in ws._cells.values()) else 0), 1
        if op.get("columns"):
            headers = [ws.cell(1 if table is None else region.min_row, start_col + i).value for i in range(len(rows[0]))]
            if len(set(headers)) != len(headers) or set(op["columns"]) != set(headers):
                raise ValueError("columns 必须与目标表头一一对应")
            rows = [[row[op["columns"].index(h)] for h in headers] for row in rows]
        if start_row + len(rows) - 1 > 1048576 or start_col + len(rows[0]) - 1 > 16384:
            raise ValueError("追加超出 Excel 边界")
        if any(ws.cell(start_row+r, start_col+c).value is not None for r in range(len(rows)) for c in range(len(rows[0]))):
            raise ValueError("追加目标已有内容")
        for r, row in enumerate(rows, start_row):
            for c, value in enumerate(row, start_col):
                ws.cell(r,c).value = value; changed += 1
        if table:
            table.ref = f"{get_column_letter(region.min_col)}{region.min_row}:{get_column_letter(region.max_col)}{start_row+len(rows)-1}"
            if table.autoFilter:
                table.autoFilter.ref = table.ref
        return f"append:{ws.title}:{changed}"
    if kind == "paste_special":
        source_ws = wb[op.get("source_sheet") or name]
        source = _bounds(source_ws, op["source_range"])
        seed = [[copy(c) for c in row] for row in source_ws.iter_rows(min_row=source.min_row,max_row=source.max_row,min_col=source.min_col,max_col=source.max_col)]
        anchor = ws[op.get("target_start", "A1")]
        height, width = (source.size["columns"], source.size["rows"]) if op.get("transpose") else (source.size["rows"], source.size["columns"])
        _bounds(ws, f"{anchor.coordinate}:{get_column_letter(anchor.column+width-1)}{anchor.row+height-1}")
        mode = op.get("mode", "all")
        if mode not in {"all", "values", "formulas", "formats"}:
            raise ValueError("paste_special.mode 无效")
        cache = None
        if mode == "values":
            from excelmanus.workbook.formula_values import analysis_rows
            cache = analysis_rows(source_ws)
        for r, row in enumerate(seed):
            for c, cell in enumerate(row):
                target = ws.cell(anchor.row + (c if op.get("transpose") else r), anchor.column + (r if op.get("transpose") else c))
                if mode == "values":
                    target.value = cache[cell.row-1][cell.column-1]
                    if isinstance(target.value, str): target.data_type = "s"
                else:
                    _paste(cell, target, mode)
                changed += 1
        return f"paste_special:{ws.title}:{changed}"
    region = _bounds(ws, op["range"])
    cells = [[ws.cell(r,c) for c in range(region.min_col,region.max_col+1)] for r in range(region.min_row,region.max_row+1)]
    if kind == "sort":
        has_header = op.get("header", True)
        header = [c.value for c in cells[0]]
        by = op.get("by")
        by = [by] if isinstance(by, (str,int)) else by
        if not by: raise ValueError("sort.by 不能为空")
        columns = [header.index(v) if isinstance(v,str) and v in header else int(v)-1 for v in by]
        if any(c < 0 or c >= len(header) for c in columns): raise ValueError("排序键超出区域")
        positions = [r for r in range(region.min_row+int(has_header),region.max_row+1) if not op.get("visible_only") or not ws.row_dimensions[r].hidden]
        source = {r:[copy(ws.cell(r,c)) for c in range(region.min_col,region.max_col+1)] for r in positions}
        from excelmanus.workbook.formula_values import analysis_rows
        values = analysis_rows(ws)
        def key(r,c):
            v = values[r-1][region.min_col+c-1]
            if isinstance(v,Number): return (0,float(v))
            if isinstance(v,(datetime,date)): return (1,v.isoformat())
            return (2,str(v).casefold())
        order = positions[:]
        directions = op.get("ascending", True)
        directions = directions if isinstance(directions,list) else [directions]*len(columns)
        if len(directions)!=len(columns): raise ValueError("ascending 数量须与 by 相同")
        for col, asc in reversed(list(zip(columns,directions))):
            nonempty = [r for r in order if values[r-1][region.min_col+col-1] is not None]
            empty = [r for r in order if values[r-1][region.min_col+col-1] is None]
            order = sorted(nonempty,key=lambda r:key(r,col),reverse=not asc)+empty
        for target_row, source_row in zip(positions,order):
            for c,cell in enumerate(source[source_row],region.min_col):
                _paste(cell,ws.cell(target_row,c)); changed += 1
    elif kind == "fill":
        mode = op.get("mode", "copy" if op.get("source_range") else "value")
        seed = None
        if mode == "copy":
            source = _bounds(ws,op["source_range"])
            seed = [[copy(c) for c in row] for row in ws.iter_rows(min_row=source.min_row,max_row=source.max_row,min_col=source.min_col,max_col=source.max_col)]
        elif mode not in {"series","value"}: raise ValueError("fill.mode 无效")
        for r,row in enumerate(cells):
            if op.get("visible_only") and ws.row_dimensions[row[0].row].hidden: continue
            for c,cell in enumerate(row):
                if seed:
                    _paste(seed[r%len(seed)][c%len(seed[0])],cell)
                elif mode=="series":
                    cell.value = op.get("start",0)+changed*op.get("step",1)
                else: cell.value=op.get("value")
                changed += 1
    elif kind == "clear":
        mode = op.get("mode","values")
        if mode not in {"values","formats","all"}: raise ValueError("clear.mode 无效")
        for row in cells:
            for cell in row:
                if mode in {"values","all"}: cell.value=None
                if mode in {"formats","all"}: cell._style=None
                if mode=="all": cell.comment=None; cell.hyperlink=None
                changed += 1
    elif kind == "replace":
        find=op.get("find"); replacement=op.get("replacement")
        if find is None: raise ValueError("replace.find 不能为空")
        pattern=re.compile(str(find) if op.get("regex") else re.escape(str(find)), 0 if op.get("match_case",True) else re.I)
        for row in cells:
            for cell in row:
                if cell.data_type=="f" and not op.get("include_formulas"): continue
                old=cell.value
                if isinstance(old,str):
                    new=pattern.sub(str(replacement or ""),old) if op.get("regex") else pattern.sub(lambda m: str(replacement if replacement is not None else ""),old)
                else: new=replacement if old==find else old
                if new!=old: cell.value=new; changed+=1
    return f"{kind}:{ws.title}:{changed}"
