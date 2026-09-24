"""Workbook presentation and object metadata shared by observation and analysis."""
from __future__ import annotations
from typing import Any

def _color_to_hex_short(color: Any) -> str | None:
    """将 openpyxl Color 对象转为 6 位十六进制字符串，无效或默认色返回 None。"""
    if color is None:
        return None
    from excelmanus.tools._style_extract import resolve_color

    color_type = getattr(color, "type", None)
    if color_type == "theme":
        resolved = resolve_color(color)
        return resolved.lstrip("#") if resolved else None
    if color_type == "indexed":
        resolved = resolve_color(color)
        return resolved.lstrip("#") if resolved else None
    if color_type == "rgb":
        resolved = resolve_color(color)
        if resolved is None:
            return None
        hex_val = resolved.lstrip("#")
        if hex_val in {"000000", "FFFFFF"} and str(getattr(color, "rgb", "")).upper() in {
            "00000000",
            "FFFFFFFF",
        }:
            return None
        return hex_val
    resolved = resolve_color(color)
    return resolved.lstrip("#") if resolved else None


def _extract_style_tuple(cell: Any) -> tuple | None:
    """从单元格提取样式关键属性元组，全默认样式返回 None。"""
    parts: list[Any] = []
    has_custom = False

    # 字体
    f = cell.font
    if f:
        font_info: dict[str, Any] = {}
        if f.name and f.name != "Calibri":
            font_info["name"] = f.name
        if f.size and f.size != 11:
            font_info["size"] = f.size
        if f.bold:
            font_info["bold"] = True
        if f.italic:
            font_info["italic"] = True
        if f.underline and f.underline != "none":
            font_info["underline"] = f.underline
        if f.strike:
            font_info["strike"] = True
        c = _color_to_hex_short(f.color)
        if c and c != "000000":
            font_info["color"] = c
        if font_info:
            has_custom = True
        parts.append(tuple(sorted(font_info.items())) if font_info else ())
    else:
        parts.append(())

    # 填充
    fl = cell.fill
    if fl:
        fill_type = fl.fill_type or fl.patternType
        if fill_type and fill_type != "none":
            fg = _color_to_hex_short(fl.fgColor)
            parts.append(("fill", fill_type, fg))
            has_custom = True
        else:
            parts.append(())
    else:
        parts.append(())

    # 边框
    b = cell.border
    if b:
        border_parts: list[tuple[str, str]] = []
        for side_name in ("left", "right", "top", "bottom"):
            side = getattr(b, side_name, None)
            if side and side.style and side.style != "none":
                border_parts.append((side_name, side.style))
        if border_parts:
            has_custom = True
        parts.append(tuple(border_parts))
    else:
        parts.append(())

    # 对齐
    a = cell.alignment
    if a:
        align_info: dict[str, Any] = {}
        if a.horizontal and a.horizontal != "general":
            align_info["horizontal"] = a.horizontal
        if a.vertical and a.vertical != "bottom":
            align_info["vertical"] = a.vertical
        if a.wrap_text:
            align_info["wrap_text"] = True
        if align_info:
            has_custom = True
        parts.append(tuple(sorted(align_info.items())) if align_info else ())
    else:
        parts.append(())

    # 数字格式
    nf = cell.number_format
    if nf and nf != "General":
        parts.append(nf)
        has_custom = True
    else:
        parts.append("")

    if not has_custom:
        return None
    return tuple(parts)


def _style_tuple_to_dict(st: tuple) -> dict[str, Any]:
    """将样式元组还原为可读字典。"""
    result: dict[str, Any] = {}
    font_parts, fill_parts, border_parts, align_parts, num_fmt = st

    if font_parts:
        result["font"] = dict(font_parts)
    if fill_parts:
        _, fill_type, fg = fill_parts
        info: dict[str, Any] = {"type": fill_type}
        if fg:
            info["color"] = fg
        result["fill"] = info
    if border_parts:
        result["border"] = {side: style for side, style in border_parts}
    if align_parts:
        result["alignment"] = dict(align_parts)
    if num_fmt:
        result["number_format"] = num_fmt
    return result


def _collect_styles_compressed(
    ws: Any,
    max_rows: int = 200,
) -> dict[str, Any]:
    """扫描工作表，以 Style Classes 压缩方式返回样式信息。

    算法：
    1. 逐单元格提取样式元组
    2. 为唯一组合分配 sN ID
    3. 按列扫描合并连续相同样式的单元格为范围

    Returns:
        包含 style_classes, cell_style_map, merged_ranges 的字典。
    """
    from openpyxl.utils import get_column_letter

    scan_rows = min(ws.max_row or 0, max_rows)
    scan_cols = ws.max_column or 0
    if scan_rows == 0 or scan_cols == 0:
        return {
            "style_classes": {},
            "cell_style_map": {},
            "merged_ranges": [str(mr) for mr in ws.merged_cells.ranges],
            "rows_scanned": scan_rows,
            "truncated": (ws.max_row or 0) > max_rows,
        }

    # 第一遍：收集所有样式元组，分配 ID
    style_to_id: dict[tuple, str] = {}
    # cell_map[col_idx][row_idx] = style_id
    cell_map: dict[int, dict[int, str]] = {}
    id_counter = 0

    for row in ws.iter_rows(min_row=1, max_row=scan_rows, min_col=1, max_col=scan_cols):
        for cell in row:
            st = _extract_style_tuple(cell)
            if st is None:
                continue
            if st not in style_to_id:
                style_to_id[st] = f"s{id_counter}"
                id_counter += 1
            sid = style_to_id[st]
            col_idx = cell.column
            row_idx = cell.row
            if col_idx not in cell_map:
                cell_map[col_idx] = {}
            cell_map[col_idx][row_idx] = sid

    # 构建 style_classes 字典
    style_classes = {sid: _style_tuple_to_dict(st) for st, sid in style_to_id.items()}

    # 第二遍：按列合并连续相同 style_id 为范围
    range_map: dict[str, str] = {}  # "A1:A10" -> "s0"

    for col_idx in sorted(cell_map.keys()):
        col_letter = get_column_letter(col_idx)
        rows_dict = cell_map[col_idx]
        sorted_rows = sorted(rows_dict.keys())
        if not sorted_rows:
            continue

        # 合并连续行
        start_row = sorted_rows[0]
        current_sid = rows_dict[start_row]
        prev_row = start_row

        for r in sorted_rows[1:]:
            sid = rows_dict[r]
            if sid == current_sid and r == prev_row + 1:
                # 连续且相同
                prev_row = r
            else:
                # 输出前一段
                if start_row == prev_row:
                    range_map[f"{col_letter}{start_row}"] = current_sid
                else:
                    range_map[f"{col_letter}{start_row}:{col_letter}{prev_row}"] = current_sid
                start_row = r
                current_sid = sid
                prev_row = r

        # 输出最后一段
        if start_row == prev_row:
            range_map[f"{col_letter}{start_row}"] = current_sid
        else:
            range_map[f"{col_letter}{start_row}:{col_letter}{prev_row}"] = current_sid

    merged_ranges = [str(mr) for mr in ws.merged_cells.ranges]

    return {
        "style_classes": style_classes,
        "cell_style_map": range_map,
        "merged_ranges": merged_ranges,
        "rows_scanned": scan_rows,
        "truncated": (ws.max_row or 0) > max_rows,
    }


def _collect_charts(ws: Any) -> list[dict[str, Any]]:
    """检测工作表中嵌入的图表，返回元信息列表。"""
    charts_info: list[dict[str, Any]] = []
    chart_list = getattr(ws, "_charts", [])
    for chart in chart_list:
        info: dict[str, Any] = {}
        # 图表类型
        type_name = type(chart).__name__.replace("Chart", "").lower()
        info["type"] = type_name
        if hasattr(chart, "title") and chart.title:
            title = chart.title
            if isinstance(title, str):
                info["title"] = title
            else:
                # openpyxl Title/Text object: drill into rich text paragraphs
                text_obj = title.text if hasattr(title, "text") else title
                rich = getattr(text_obj, "rich", None)
                if rich is not None:
                    parts: list[str] = []
                    for p in getattr(rich, "p", []):
                        for r in (getattr(p, "r", None) or []):
                            if getattr(r, "t", None):
                                parts.append(r.t)
                    if parts:
                        info["title"] = "".join(parts)
        info["series_count"] = len(chart.series) if hasattr(chart, "series") else 0
        # 锚点位置
        if hasattr(chart, "anchor") and chart.anchor:
            anchor = chart.anchor
            if hasattr(anchor, "_from") and anchor._from:
                f = anchor._from
                from openpyxl.utils import get_column_letter as gcl
                info["anchor_cell"] = f"{gcl(f.col + 1)}{f.row + 1}"
        charts_info.append(info)
    return charts_info


def _collect_images(ws: Any) -> list[dict[str, Any]]:
    """检测工作表中嵌入的图片，返回元信息列表。"""
    images_info: list[dict[str, Any]] = []
    image_list = getattr(ws, "_images", [])
    for img in image_list:
        info: dict[str, Any] = {}
        if hasattr(img, "width") and img.width:
            info["width_px"] = img.width
        if hasattr(img, "height") and img.height:
            info["height_px"] = img.height
        # 图片格式
        if hasattr(img, "format"):
            info["format"] = img.format
        elif hasattr(img, "path") and img.path:
            ext = str(img.path).rsplit(".", 1)[-1] if "." in str(img.path) else "unknown"
            info["format"] = ext
        # 锚点位置
        if hasattr(img, "anchor") and img.anchor:
            anchor = img.anchor
            if isinstance(anchor, str):
                info["anchor_cell"] = anchor
            elif hasattr(anchor, "_from") and anchor._from:
                f = anchor._from
                from openpyxl.utils import get_column_letter as gcl
                info["anchor_cell"] = f"{gcl(f.col + 1)}{f.row + 1}"
        images_info.append(info)
    return images_info


def _collect_freeze_panes(ws: Any) -> str | None:
    """返回冻结窗格位置（如 'A4'），未冻结返回 None。"""
    fp = ws.freeze_panes
    return str(fp) if fp else None


def _collect_conditional_formatting(ws: Any) -> list[dict[str, Any]]:
    """收集条件格式规则列表。"""
    rules_info: list[dict[str, Any]] = []
    for cf in ws.conditional_formatting:
        ranges_str = str(cf)
        for rule in cf.rules:
            from openpyxl.xml.functions import tostring
            info: dict[str, Any] = {"range": str(cf.sqref), "raw_xml": tostring(rule.to_tree()).decode(), "evaluated": False}
            if rule.dxf is not None:
                info["differential_style_xml"] = tostring(rule.dxf.to_tree()).decode()
            if hasattr(rule, "type") and rule.type:
                info["type"] = rule.type
            if hasattr(rule, "priority") and rule.priority is not None:
                info["priority"] = rule.priority
            if hasattr(rule, "formula") and rule.formula:
                info["formula"] = list(rule.formula) if not isinstance(rule.formula, str) else [rule.formula]
            if hasattr(rule, "operator") and rule.operator:
                info["operator"] = rule.operator
            rules_info.append(info)
    return rules_info


def _collect_data_validation(ws: Any) -> list[dict[str, Any]]:
    """收集数据验证规则列表。"""
    validations: list[dict[str, Any]] = []
    dv_list = getattr(ws, "data_validations", None)
    if dv_list is None:
        return validations
    dv_items = getattr(dv_list, "dataValidation", [])
    for dv in dv_items:
        info: dict[str, Any] = {}
        if hasattr(dv, "sqref") and dv.sqref:
            info["range"] = str(dv.sqref)
        if hasattr(dv, "type") and dv.type:
            info["type"] = dv.type
        if hasattr(dv, "formula1") and dv.formula1:
            info["formula1"] = str(dv.formula1)
        if hasattr(dv, "formula2") and dv.formula2:
            info["formula2"] = str(dv.formula2)
        if hasattr(dv, "allow_blank") and dv.allow_blank is not None:
            info["allow_blank"] = bool(dv.allow_blank)
        if hasattr(dv, "showDropDown") and dv.showDropDown is not None:
            info["show_dropdown"] = not bool(dv.showDropDown)
        validations.append(info)
    return validations


def _collect_print_settings(ws: Any) -> dict[str, Any]:
    """收集打印设置信息。"""
    info: dict[str, Any] = {}
    setup = ws.sheet_properties.pageSetUpPr
    fit_to_page = bool(setup is not None and setup.fitToPage)
    info["fit_to_page"] = fit_to_page
    info["scaling_mode"] = "fit_to_page" if fit_to_page else "scale"
    if ws.print_area:
        info["print_area"] = ws.print_area
    ps = ws.page_setup
    if ps:
        if ps.orientation:
            info["orientation"] = ps.orientation
        if ps.paperSize is not None:
            info["paper_size"] = ps.paperSize
        if ps.fitToWidth is not None:
            info["fit_to_width"] = ps.fitToWidth
        if ps.fitToHeight is not None:
            info["fit_to_height"] = ps.fitToHeight
        if ps.scale is not None:
            info["scale"] = ps.scale
    if ws.print_title_rows:
        info["repeat_rows"] = ws.print_title_rows
    if ws.print_title_cols:
        info["repeat_columns"] = ws.print_title_cols
    # Preserve legacy OOXML fields above; this named subset can be passed
    # directly to apply_spreadsheet_changes(kind=print_layout). Never collapse
    # multiple print areas into a single range or activate inactive fit sizes.
    layout: dict[str, Any] = {"fit_to_page": fit_to_page}
    from openpyxl.worksheet.print_settings import PrintArea

    ranges = sorted(str(area) for area in PrintArea.from_string(str(ws.print_area)).ranges) if ws.print_area else []
    unsupported: list[str] = []
    if len(ranges) <= 1:
        layout["print_area"] = ranges[0] if ranges else ""
    else:
        unsupported.append("multiple_print_areas")
    if ps.orientation in {"portrait", "landscape"}:
        layout["orientation"] = ps.orientation
    paper = {"1": "Letter", "5": "Legal", "8": "A3", "9": "A4", "11": "A5"}.get(str(ps.paperSize))
    if paper:
        layout["paper_size"] = paper
    elif ps.paperSize is not None:
        unsupported.append("paper_size")
    if fit_to_page:
        for key, value in (("fit_to_width", ps.fitToWidth), ("fit_to_height", ps.fitToHeight)):
            if value is not None and 0 <= value <= 32767:
                layout[key] = value
    elif ps.scale is not None:
        if 10 <= ps.scale <= 400:
            layout["scale"] = ps.scale
        else:
            unsupported.append("scale")
    info["print_layout"] = layout
    if unsupported:
        info["print_layout_omitted"] = unsupported
    return info


def _collect_column_widths(ws: Any) -> dict[str, float]:
    """收集非默认列宽映射。"""
    widths: dict[str, float] = {}
    from openpyxl.utils.cell import column_index_from_string, get_column_letter

    for col_letter, dim in ws.column_dimensions.items():
        if dim.width is not None:
            first = dim.min or column_index_from_string(col_letter)
            last = dim.max or first
            for col in range(first, min(last, 16384) + 1):
                widths[get_column_letter(col)] = round(dim.width, 2)
    return widths


def _collect_row_heights(ws: Any) -> dict[str, float]:
    return {
        str(row): round(dim.height, 2)
        for row, dim in ws.row_dimensions.items() if dim.height is not None
    }


def _collect_formulas(ws: Any, max_rows: int = 200) -> dict[str, Any]:
    """收集含公式的单元格位置和公式内容。"""
    from openpyxl.utils import get_column_letter

    formulas: list[dict[str, str]] = []
    total_rows = ws.max_row or 0
    scan_rows = min(total_rows, max_rows)
    scan_cols = ws.max_column or 0
    if scan_rows == 0 or scan_cols == 0:
        return {
            "items": formulas,
            "rows_scanned": scan_rows,
            "truncated": False,
        }

    for row in ws.iter_rows(min_row=1, max_row=scan_rows, min_col=1, max_col=scan_cols):
        for cell in row:
            val = cell.value
            if isinstance(val, str) and val.startswith("="):
                coord = f"{get_column_letter(cell.column)}{cell.row}"
                formulas.append({"cell": coord, "formula": val})
    return {
        "items": formulas,
        "rows_scanned": scan_rows,
        "truncated": total_rows > max_rows,
    }


