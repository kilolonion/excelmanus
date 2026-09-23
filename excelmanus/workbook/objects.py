"""Native workbook objects, including real OOXML PivotTable/cache parts."""
from copy import copy
from datetime import date, datetime
import math
import re

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

KINDS = {"table", "defined_name", "image", "hyperlink", "comment", "pivot_table"}
FIELDS = "kind action sheet sheet_name name ref range cell scope refers_to hidden comment columns totals_row style show_first_column show_last_column show_row_stripes show_column_stripes image_path target_cell width height target location display tooltip text author index source_sheet source_range target_sheet rows values refresh_on_load"


def object_schema_fields():
    fields = {k: {"type":"string"} for k in FIELDS.split()}
    for k in ("hidden","totals_row","show_first_column","show_last_column","show_row_stripes","show_column_stripes","refresh_on_load"):
        fields[k]={"type":"boolean"}
    for k in ("columns","rows","values"):
        fields[k]={"type":"array","items":{}}
    for k in ("width","height"): fields[k]={"type":"number","exclusiveMinimum":0}
    fields["index"]={"type":"integer","minimum":0}
    fields["style"]={"type":["string","integer"]}
    return fields


def list_workbook_objects(wb, sheet_name=None):
    result=[]
    for ws in wb.worksheets:
        if sheet_name and ws.title!=sheet_name: continue
        for table in ws.tables.values():
            result.append({"kind":"table","sheet":ws.title,"name":table.name,"ref":table.ref,"columns":table.column_names})
        for name in ws.defined_names.values():
            result.append({"kind":"defined_name","name":name.name,"scope":ws.title,"refers_to":name.attr_text})
        for i,pivot in enumerate(ws._pivots):
            source=getattr(pivot.cache.cacheSource,"worksheetSource",None)
            result.append({"kind":"pivot_table","sheet":ws.title,"index":i,"name":pivot.name,"ref":pivot.location.ref,"source_sheet":getattr(source,"sheet",None),"source_range":getattr(source,"ref",None)})
        for kind,items in (("chart",ws._charts),("image",ws._images)):
            for i,item in enumerate(items):
                anchor=item.anchor
                point=getattr(anchor,"_from",None)
                result.append({"kind":kind,"sheet":ws.title,"index":i,"target_cell":anchor if isinstance(anchor,str) else f"{get_column_letter(point.col+1)}{point.row+1}" if point else None})
        for cell in ws._cells.values():
            if cell.comment: result.append({"kind":"comment","sheet":ws.title,"cell":cell.coordinate,"text":cell.comment.text,"author":cell.comment.author})
            if cell.hyperlink: result.append({"kind":"hyperlink","sheet":ws.title,"cell":cell.coordinate,"target":cell.hyperlink.target,"location":cell.hyperlink.location})
    for name in wb.defined_names.values():
        result.append({"kind":"defined_name","name":name.name,"scope":"workbook","refers_to":name.attr_text})
    return result


def _pivot(wb, ws, op, action):
    name=op.get("name")
    current=next((p for p in ws._pivots if p.name==name),None)
    if action=="delete":
        if current is None: raise ValueError("透视表不存在")
        ws._pivots.remove(current)
        return {"kind":"pivot_table","name":name,"deleted":True}
    if not name or not re.fullmatch(r"[A-Za-z_\u4e00-\u9fff][\w.]*",name): raise ValueError("需要合法 pivot name")
    if action=="create" and any(p.name==name for sheet in wb.worksheets for p in sheet._pivots): raise ValueError("透视表名称重复")
    if action in {"update","refresh"} and current is None: raise ValueError("透视表不存在")
    source_sheet=op.get("source_sheet")
    source_range=op.get("source_range")
    row_fields=op.get("rows") or []
    col_fields=op.get("columns") or []
    measures=op.get("values") or []
    if current and action=="refresh":
        src=current.cache.cacheSource.worksheetSource
        source_sheet,source_range=src.sheet,src.ref
        names=[f.name for f in current.cache.cacheFields]
        row_fields=[names[f.x] for f in current.rowFields]
        col_fields=[names[f.x] for f in current.colFields if f.x>=0]
        measures=[{"field":names[f.fld],"aggregation":f.subtotal,"name":f.name} for f in current.dataFields]
    if not source_sheet or not source_range or not row_fields or len(measures)!=1 or len(col_fields)>1:
        raise ValueError("原生透视表需要 source_sheet/source_range、rows，当前支持一个值字段、至多一个列字段")
    src=wb[source_sheet]; region=CellRange(source_range)
    from excelmanus.workbook.formula_values import analysis_rows
    matrix=analysis_rows(src)
    source_rows=[list(row[region.min_col-1:region.max_col]) for row in matrix[region.min_row-1:region.max_row]]
    headers=[str(v) if v is not None else "" for v in source_rows[0]]
    if not all(headers) or len(set(headers))!=len(headers): raise ValueError("透视源表头必须非空且唯一")
    measure=measures[0] if isinstance(measures[0],dict) else {"field":measures[0]}
    field=measure["field"]; agg=measure.get("aggregation","sum")
    functions={"sum":"sum","count":"count","average":"mean","min":"min","max":"max","product":"prod"}
    if agg not in functions: raise ValueError("透视聚合支持 sum/count/average/min/max/product")
    for label in [*row_fields,*col_fields,field]:
        if label not in headers: raise ValueError(f"透视字段不存在: {label}")
    if field in row_fields or field in col_fields or set(row_fields)&set(col_fields): raise ValueError("透视维度与值字段不能重叠")
    import pandas as pd
    frame=pd.DataFrame(source_rows[1:],columns=headers)
    from excelmanus.workbook.numeric import parse_number
    if agg!="count":
        converted=frame[field].map(parse_number)
        if (frame[field].notna() & converted.isna()).any(): raise ValueError("值字段存在非数值，请先规范化")
        frame[field]=converted
    result=pd.pivot_table(frame,index=row_fields,columns=col_fields or None,values=field,aggfunc=functions[agg],fill_value=0,dropna=False)
    table=result.reset_index()
    output=[list(map(str,table.columns))]+[[None if pd.isna(v) else v.item() if hasattr(v,"item") else v for v in row] for row in table.itertuples(index=False,name=None)]
    target=op.get("target_cell") or (current.location.ref.split(":")[0] if current else "A1")
    anchor=ws[target]
    footprint=CellRange(min_col=anchor.column,min_row=anchor.row,max_col=anchor.column+len(output[0])-1,max_row=anchor.row+len(output)-1)
    old_region=CellRange(current.location.ref) if current else None
    if src is ws and not region.isdisjoint(footprint): raise ValueError("透视输出不能覆盖源数据")
    for row in ws.iter_rows(min_row=footprint.min_row,max_row=footprint.max_row,min_col=footprint.min_col,max_col=footprint.max_col):
        for c in row:
            if c.value is not None and (old_region is None or c.coordinate not in old_region): raise ValueError("透视输出区域已有内容")
    from openpyxl.pivot.cache import CacheDefinition,CacheSource,WorksheetSource,CacheField,SharedItems
    from openpyxl.pivot.fields import Text,Number,Missing,Boolean,DateTimeField,Index
    from openpyxl.pivot.record import Record,RecordList
    from openpyxl.pivot.table import TableDefinition,Location,PivotField,RowColField,DataField,FieldItem
    shared=[]; mappings=[]
    def key(v): return (type(v).__name__,str(v))
    def item(v):
        if v is None: return Missing()
        if isinstance(v,bool): return Boolean(v=v)
        if isinstance(v,(int,float)) and math.isfinite(v): return Number(v=v)
        if isinstance(v,datetime): return DateTimeField(v=v)
        return Text(v=str(v))
    for col in range(len(headers)):
        unique=[]; mapping={}
        for row in source_rows[1:]:
            v=row[col]
            if key(v) not in mapping: mapping[key(v)]=len(unique); unique.append(v)
        shared.append(unique); mappings.append(mapping)
    cache=CacheDefinition(cacheSource=CacheSource(type="worksheet",worksheetSource=WorksheetSource(ref=region.coord,sheet=src.title)),cacheFields=[CacheField(name=h,sharedItems=SharedItems(_fields=[item(v) for v in shared[i]])) for i,h in enumerate(headers)],recordCount=len(source_rows)-1,saveData=True,enableRefresh=True,refreshOnLoad=op.get("refresh_on_load",True),createdVersion=6,refreshedVersion=6,minRefreshableVersion=3)
    cache.records=RecordList(r=[Record(_fields=[Index(v=mappings[i][key(v)]) for i,v in enumerate(row)]) for row in source_rows[1:]])
    pivot=TableDefinition(name=name,cacheId=1,dataCaption="Values",location=Location(ref=footprint.coord,firstHeaderRow=1,firstDataRow=1,firstDataCol=len(row_fields)),rowGrandTotals=False,colGrandTotals=False,compact=False,compactData=False,gridDropZones=True,pivotFields=[PivotField(axis="axisRow" if h in row_fields else "axisCol" if h in col_fields else None,dataField=h==field,defaultSubtotal=False,items=[FieldItem(x=j) for j in range(len(shared[i]))] if h in row_fields+col_fields else []) for i,h in enumerate(headers)],rowFields=[RowColField(x=headers.index(h)) for h in row_fields],colFields=[RowColField(x=headers.index(h)) for h in col_fields],dataFields=[DataField(name=measure.get("name",f"{agg}_{field}"),fld=headers.index(field),subtotal=agg)])
    pivot.cache=cache
    if current:
        ws._pivots.remove(current)
        for row in ws.iter_rows(min_row=old_region.min_row,max_row=old_region.max_row,min_col=old_region.min_col,max_col=old_region.max_col):
            for c in row: c.value=None
    for r,row in enumerate(output,anchor.row):
        for c,v in enumerate(row,anchor.column):
            ws.cell(r,c).value=v
            if isinstance(v,str): ws.cell(r,c).data_type="s"
    ws.add_pivot(pivot)
    return {"kind":"pivot_table","name":name,"sheet":ws.title,"ref":footprint.coord,"cache_records":len(source_rows)-1,"native":True}


def apply_object_operation(wb,op,*,guard=None):
    kind=op["kind"]; action=op.get("action","create")
    if action not in {"create","update","resize","delete","refresh"}: raise ValueError("对象 action 无效")
    sheet=op.get("target_sheet") if kind=="pivot_table" else None
    sheet=sheet or op.get("sheet") or op.get("sheet_name")
    if not sheet and len(wb.worksheets)!=1: raise ValueError("对象操作需要 sheet")
    ws=wb[sheet] if sheet else wb.active
    if kind=="pivot_table": return _pivot(wb,ws,op,action)
    name=op.get("name")
    if kind=="table":
        from openpyxl.worksheet.table import Table,TableStyleInfo,TableColumn,TableFormula
        from openpyxl.worksheet.filters import AutoFilter
        existing=ws.tables.get(name)
        if action=="delete":
            if existing is None: raise ValueError("Table 不存在")
            del ws.tables[name]
        else:
            if action=="create" and any(name in s.tables for s in wb.worksheets): raise ValueError("Table 名称已存在")
            if action!="create" and existing is None: raise ValueError("Table 不存在")
            if not name or not re.fullmatch(r"[A-Za-z_\u4e00-\u9fff][\w.]*",name): raise ValueError("Table name 无效")
            region=CellRange(op.get("ref") or op.get("range") or existing.ref)
            if any(not region.isdisjoint(CellRange(t.ref)) for t in ws.tables.values() if t is not existing): raise ValueError("Table 区域重叠")
            if any(not region.isdisjoint(m) for m in ws.merged_cells.ranges): raise ValueError("Table 区域不能含合并格")
            configs=op.get("columns")
            width=region.max_col-region.min_col+1
            headers=[str(ws.cell(region.min_row,c).value or "") for c in range(region.min_col,region.max_col+1)]
            if configs:
                if len(configs)!=width: raise ValueError("Table columns 数量不符")
                headers=[c if isinstance(c,str) else c["name"] for c in configs]
            if not all(headers) or len(set(headers))!=width: raise ValueError("Table 表头必须非空且唯一")
            table=existing or Table(displayName=name,ref=region.coord)
            table.ref=region.coord
            old={c.name:c for c in table.tableColumns}
            table.tableColumns=[copy(old[h]) if h in old else TableColumn(id=i+1,name=h) for i,h in enumerate(headers)]
            for i,h in enumerate(headers):
                ws.cell(region.min_row,region.min_col+i).value=h
                cfg=configs[i] if configs and isinstance(configs[i],dict) else {}
                col=table.tableColumns[i]; col.id=i+1
                if "formula" in cfg: col.calculatedColumnFormula=TableFormula(attr_text=cfg["formula"].lstrip("="))
                if "totals_function" in cfg: col.totalsRowFunction=cfg["totals_function"]
                if "totals_label" in cfg: col.totalsRowLabel=cfg["totals_label"]
                if "formula" in cfg:
                    for r in range(region.min_row+1,region.max_row+1-int(op.get("totals_row",bool(table.totalsRowCount)))):
                        ws.cell(r,region.min_col+i).value="="+cfg["formula"].lstrip("=")
            if "totals_row" in op: table.totalsRowCount=int(op["totals_row"]); table.totalsRowShown=op["totals_row"]
            table.autoFilter=AutoFilter(ref=f"{get_column_letter(region.min_col)}{region.min_row}:{get_column_letter(region.max_col)}{region.max_row-int(bool(table.totalsRowCount))}")
            if "style" in op or table.tableStyleInfo is None: table.tableStyleInfo=TableStyleInfo(name=op.get("style","TableStyleMedium2"),showRowStripes=True)
            for field,attr in (("show_first_column","showFirstColumn"),("show_last_column","showLastColumn"),("show_row_stripes","showRowStripes"),("show_column_stripes","showColumnStripes")):
                if field in op: setattr(table.tableStyleInfo,attr,op[field])
            if existing is None: ws.add_table(table)
    elif kind=="defined_name":
        from openpyxl.workbook.defined_name import DefinedName
        scope=op.get("scope","workbook")
        names=wb.defined_names if scope=="workbook" else wb[scope].defined_names
        if not name or not re.fullmatch(r"[A-Za-z_\u4e00-\u9fff][\w.]*",name): raise ValueError("名称无效")
        if action=="delete":
            if name not in names: raise ValueError("名称不存在")
            del names[name]
        else:
            if action=="create" and name in names: raise ValueError("名称已存在")
            if action!="create" and name not in names: raise ValueError("名称不存在")
            value=op.get("refers_to") or op.get("ref")
            if not value: raise ValueError("名称需要 refers_to")
            names.add(DefinedName(name,attr_text=value.lstrip("="),hidden=op.get("hidden",False),comment=op.get("comment")))
    elif kind in {"comment","hyperlink"}:
        from openpyxl.comments import Comment
        from openpyxl.worksheet.hyperlink import Hyperlink
        cell=ws[op["cell"]]
        if action=="delete": setattr(cell,kind,None)
        elif kind=="comment":
            cell.comment=Comment(op["text"],op.get("author","ExcelManus"))
            if "width" in op: cell.comment.width=op["width"]
            if "height" in op: cell.comment.height=op["height"]
        else:
            if not op.get("target") and not op.get("location"): raise ValueError("hyperlink 需要 target 或 location")
            cell.hyperlink=Hyperlink(ref=cell.coordinate,target=op.get("target"),location=op.get("location"),display=op.get("display"),tooltip=op.get("tooltip"))
            if "display" in op: cell.value=op["display"]
    elif kind=="image":
        from openpyxl.drawing.image import Image
        index=op.get("index")
        if action in {"update","delete"}:
            if not isinstance(index,int) or index<0 or index>=len(ws._images): raise ValueError("image.index 不存在")
            image=ws._images[index]
            if action=="delete": ws._images.remove(image); return {"kind":kind,"index":index,"deleted":True}
        else: image=None
        if op.get("image_path"):
            if guard is None: raise ValueError("图片操作需要工作区 guard")
            new=Image(str(guard.resolve_and_validate(op["image_path"])))
            if image: new.anchor=image.anchor; ws._images[index]=new
            else: ws.add_image(new,op.get("target_cell","A1"))
            image=new
        if image is None: raise ValueError("image.create 需要 image_path")
        if "target_cell" in op: image.anchor=op["target_cell"]
        for key in ("width","height"):
            if key in op: setattr(image,key,op[key])
    else: raise ValueError(f"不支持的对象类型 {kind}")
    return {"kind":kind,"action":action,"sheet":ws.title,"name":name,"cell":op.get("cell")}
