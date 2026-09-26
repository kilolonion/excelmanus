"""Excel 原生图表实现：供 `apply_spreadsheet_changes` 内部调用。

模型面不注册本模块；PNG / matplotlib 出图走 `run_code`。
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from excelmanus.engine_core.tool_result import ToolResult, error_result
from excelmanus.logger import get_logger
from excelmanus.tools.context import bind_workspace
from excelmanus.tools._helpers import (
    MutationAborted,
    get_worksheet,
    resolve_sheet_name,
)
from excelmanus.workbook.refs import InvalidRefError

logger = get_logger("tools.chart")

EXCEL_CHART_TYPES = ("bar", "line", "pie", "scatter", "area")
_CHART_TYPE_ALIASES = {
    "column": "bar",
    "col": "bar",
    "barchart": "bar",
    "columnclustered": "bar",
    "clusteredcolumn": "bar",
    "linechart": "line",
    "piechart": "pie",
    "areachart": "area",
    "scatterchart": "scatter",
    "xy": "scatter",
}


def init_guard(workspace_root: str) -> None:
    bind_workspace(workspace_root)


@dataclass(frozen=True)
class ChartSpec:
    chart_type: str
    data_range: str
    categories_range: str | None
    sheet_name: str | None
    target_cell: str
    target_sheet: str | None
    title: str | None
    x_title: str | None
    y_title: str | None
    style: int | None
    width: float
    height: float
    from_rows: bool


def normalize_chart_args(
    *,
    chart_type: str,
    data_range: str,
    categories_range: str | None = None,
    sheet_name: str | None = None,
    target_cell: str = "A1",
    target_sheet: str | None = None,
    title: str | None = None,
    x_title: str | None = None,
    y_title: str | None = None,
    style: int | None = None,
    width: float = 15.0,
    height: float = 10.0,
    from_rows: bool = False,
) -> ChartSpec | ToolResult:
    """校验图表参数；失败返回 ToolResult，成功返回 ChartSpec。"""
    from excelmanus.workbook.address import (
        combine_sheet_names,
        parse_sheet_address,
        top_left_cell,
    )

    normalized_type = _CHART_TYPE_ALIASES.get(
        (chart_type or "").strip().lower(),
        (chart_type or "").strip().lower(),
    )
    import math

    if any(isinstance(n, bool) or not isinstance(n, (int, float)) or not math.isfinite(n) or n <= 0 for n in (width, height)):
        return error_result("图表 width/height 必须为正数，单位厘米", code="INVALID_ARGS")
    if normalized_type not in EXCEL_CHART_TYPES:
        return error_result(
            f"不支持的图表类型 '{chart_type}'，支持: {list(EXCEL_CHART_TYPES)}；column 视为 bar",
            code="INVALID_ARGS",
        )
    try:
        parsed_data = parse_sheet_address(data_range)
        sheet_name = combine_sheet_names(sheet_name, parsed_data.sheet)
        data_range = parsed_data.address
        if categories_range:
            parsed_cats = parse_sheet_address(categories_range)
            categories_range = parsed_cats.address
            sheet_name = combine_sheet_names(sheet_name, parsed_cats.sheet)
        parsed_target = parse_sheet_address(target_cell)
        if parsed_target.sheet:
            target_sheet = combine_sheet_names(target_sheet, parsed_target.sheet)
        target_cell = parsed_target.address or "A1"
        from openpyxl.utils.cell import range_boundaries as _rb

        bounds = _rb(data_range)
        if all(value is None for value in bounds):
            return error_result(f"data_range 不是合法区域: {data_range!r}", code="RANGE_INVALID")
        if categories_range:
            cat_bounds = _rb(categories_range)
            if all(value is None for value in cat_bounds):
                return error_result(f"categories_range 不是合法区域: {categories_range!r}", code="RANGE_INVALID")
        target_cell = top_left_cell(target_cell) or "A1"
    except InvalidRefError as exc:
        return error_result(str(exc), code="RANGE_INVALID")
    except ValueError as exc:
        return error_result(
            str(exc) + " 图表的 sheet 是数据源表；跨表放置用 target_sheet。",
            code="INVALID_ARGS",
        )
    except Exception as exc:
        return error_result(
            f"data_range/categories_range 不是合法坐标：{exc}。"
            "请用 A1:B12，表名放在 sheet，或写成 数据!A1:B12。",
            code="RANGE_INVALID",
        )
    return ChartSpec(
        chart_type=normalized_type,
        data_range=data_range,
        categories_range=categories_range,
        sheet_name=sheet_name,
        target_cell=target_cell,
        target_sheet=target_sheet,
        title=title,
        x_title=x_title,
        y_title=y_title,
        style=style,
        width=width,
        height=height,
        from_rows=from_rows,
    )


def add_chart_to_workbook(wb: Any, spec: ChartSpec) -> dict[str, Any]:
    """在已打开的工作簿上插入一张图，不提交。"""
    from openpyxl.chart import (
        AreaChart,
        BarChart,
        LineChart,
        PieChart,
        Reference,
        ScatterChart,
    )
    from openpyxl.utils.cell import range_boundaries

    def _bounds(address: str) -> tuple[int, int, int, int]:
        raw = range_boundaries(address)
        min_col, min_row, max_col, max_row = raw
        # openpyxl leaves one axis open for A:A / 1:1.  Clip it to the used
        # region so charts remain finite and portable across Excel engines.
        min_col = int(min_col or 1)
        min_row = int(min_row or 1)
        max_col = int(max_col or ws.max_column or min_col)
        max_row = int(max_row or ws.max_row or min_row)
        if ":" not in address:
            max_col = max_col or min_col
            max_row = max_row or min_row
        return min_col, min_row, max_col, max_row

    try:
        from excelmanus.workbook.snapshot import require_default_sheet
        ws = wb[require_default_sheet(wb.sheetnames, spec.sheet_name)]
    except ValueError as exc:
        raise MutationAborted(error_result(str(exc), code="INVALID_ARGS")) from exc

    chart_class_map = {
        "bar": BarChart,
        "line": LineChart,
        "pie": PieChart,
        "scatter": ScatterChart,
        "area": AreaChart,
    }
    chart = chart_class_map[spec.chart_type]()
    if spec.title:
        chart.title = spec.title
    if spec.style is not None:
        chart.style = spec.style
    chart.width = spec.width
    chart.height = spec.height
    if spec.chart_type not in ("pie",):
        if spec.x_title:
            chart.x_axis.title = spec.x_title
        if spec.y_title:
            chart.y_axis.title = spec.y_title

    min_col, min_row, max_col, max_row = _bounds(spec.data_range)
    cats_ref = None
    if spec.categories_range:
        c_min_col, c_min_row, c_max_col, c_max_row = _bounds(spec.categories_range)
        cats_ref = Reference(
            ws, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row
        )
        # A table-shaped range may include the explicitly designated category
        # axis. Exclude that edge vector instead of plotting labels as a series.
        # 散点图的行方向在下方单独处理，不在这里改写 min_row/max_row。
        if not spec.from_rows and c_min_col == c_max_col and (c_min_row, c_max_row) == (min_row + 1, max_row):
            if c_min_col == min_col:
                min_col += 1
            elif c_min_col == max_col:
                max_col -= 1
        elif spec.from_rows and spec.chart_type != "scatter" and c_min_row == c_max_row and (c_min_col, c_max_col) == (min_col + 1, max_col):
            if c_min_row == min_row:
                min_row += 1
            elif c_min_row == max_row:
                max_row -= 1
    if min_col > max_col or min_row > max_row:
        raise ValueError("data_range 除类别范围外必须至少包含一个数据系列")
    data_ref = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)

    if spec.chart_type == "scatter":
        from openpyxl.chart import Series as ChartSeries
        from openpyxl.chart.data_source import StrRef
        from openpyxl.chart.marker import Marker
        from openpyxl.chart.series import SeriesLabel
        from openpyxl.chart.shapes import GraphicalProperties
        from openpyxl.drawing.line import LineProperties
        from openpyxl.utils import get_column_letter, quote_sheetname

        def _point_style(series: Any) -> None:
            """散点图显式默认：只有数据点、不连线。

            Excel 与 LibreOffice 对“未声明 marker/line”的散点序列各有默认解释，
            同一文件会出现“点”或“点+线”两种观感；这里把观察到的正确形态写进
            对象，避免交付物依赖渲染器默认值。
            """
            series.marker = Marker(symbol="circle", size=7)
            series.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))

        def _series_title(row: int, col: int) -> SeriesLabel:
            """系列名取表头单元格的动态引用，而不是让渲染器回退成 Column B/Column C。"""
            ref = f"{quote_sheetname(ws.title)}!${get_column_letter(col)}${row}"
            return SeriesLabel(strRef=StrRef(f=ref))

        if spec.from_rows:
            # 行方向：每一行是一个序列，行首标签列是系列名，其余单元格是数据点；
            # 类别轴取 categories_range，缺省用数据区首行（表头行）。
            label_col = min_col
            data_min_col = min_col + 1
            if data_min_col > max_col:
                raise ValueError("行方向散点图除系列名外必须至少包含一列数值")
            if cats_ref is not None:
                c_min_col, c_min_row, c_max_col, c_max_row = _bounds(spec.categories_range)
                if c_min_row != c_max_row:
                    raise ValueError(
                        "行方向散点图的 categories_range 必须是单行（类别是每个数据点的 x 值）；"
                        "行列方向相反时请改用 from_rows=false 或转置数据。"
                    )
                if (c_max_col - c_min_col + 1) != (max_col - data_min_col + 1):
                    raise ValueError(
                        f"categories_range 有 {c_max_col - c_min_col + 1} 个类别，"
                        f"而每行有 {max_col - data_min_col + 1} 个数据点，数量必须一致。"
                    )
                x_values = cats_ref
            else:
                x_values = Reference(ws, min_col=data_min_col, max_col=max_col, min_row=min_row, max_row=min_row)
            for row_idx in range(min_row + 1, max_row + 1):
                y_values = Reference(ws, min_col=data_min_col, max_col=max_col, min_row=row_idx, max_row=row_idx)
                series = ChartSeries(y_values, xvalues=x_values, title_from_data=False)
                series.tx = _series_title(row_idx, label_col)
                _point_style(series)
                chart.series.append(series)
        else:
            if cats_ref is not None:
                c_min_col, c_min_row, c_max_col, c_max_row = _bounds(spec.categories_range)
                if c_min_col != c_max_col:
                    raise ValueError(
                        "列方向散点图的 categories_range 必须是单列（类别是每个数据点的 x 值）；"
                        "行列方向相反时请改用 from_rows=true 或转置数据。"
                    )
                if (c_max_row - c_min_row + 1) != (max_row - min_row):
                    raise ValueError(
                        f"categories_range 有 {c_max_row - c_min_row + 1} 个类别，"
                        f"而每个序列有 {max_row - min_row} 个数据点（不含表头），数量必须一致。"
                    )
                x_values = cats_ref
            else:
                x_values = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
            for col_idx in range(min_col if cats_ref else min_col + 1, max_col + 1):
                y_values = Reference(ws, min_col=col_idx, min_row=min_row + 1, max_row=max_row)
                series = ChartSeries(y_values, xvalues=x_values, title_from_data=False)
                series.tx = _series_title(min_row, col_idx)
                _point_style(series)
                chart.series.append(series)
    else:
        chart.add_data(data_ref, titles_from_data=True, from_rows=spec.from_rows)
        if cats_ref is not None:
            chart.set_categories(cats_ref)

    target_ws = ws
    if spec.target_sheet:
        resolved_target = resolve_sheet_name(spec.target_sheet, wb.sheetnames)
        if resolved_target:
            target_ws = wb[resolved_target]
        else:
            target_ws = wb.create_sheet(title=spec.target_sheet)
    target_ws.add_chart(chart, spec.target_cell)

    from excelmanus.workbook.presentation import _collect_charts

    chart_info = _collect_charts(target_ws)
    return {
        "chart_type": spec.chart_type,
        "data_range": spec.data_range,
        "target_sheet": target_ws.title,
        "target_cell": spec.target_cell,
        "chart_info": chart_info[-1] if chart_info else {},
        "total_charts": len(chart_info),
    }


def _select_chart(ws: Any, *, index: int | None = None, target_cell: str | None = None, title: str | None = None) -> tuple[int, Any] | None:
    charts = list(getattr(ws, "_charts", None) or [])
    if index is not None:
        if 0 <= int(index) < len(charts):
            return int(index), charts[int(index)]
        return None
    for pos, chart in enumerate(charts):
        if title and _chart_title(chart) == title:
            return pos, chart
        if target_cell and isinstance(chart.anchor, str) and chart.anchor.upper() == target_cell.upper():
            return pos, chart
        anchor = getattr(getattr(chart, "anchor", None), "_from", None)
        if target_cell and anchor is not None:
            from openpyxl.utils import get_column_letter
            cell = f"{get_column_letter(anchor.col + 1)}{anchor.row + 1}"
            if cell.upper() == str(target_cell).upper():
                return pos, chart
    return None


def _chart_title(chart: Any) -> str:
    tx = getattr(getattr(chart, "title", None), "tx", None)
    rich = getattr(tx, "rich", None)
    return "".join(str(run.t or "") for paragraph in getattr(rich, "p", ()) for run in getattr(paragraph, "r", ()))


def update_chart_in_workbook(wb: Any, op: dict[str, Any]) -> dict[str, Any]:
    """Patch properties/series on the existing chart, preserving other settings."""
    ws = get_worksheet(wb, op.get("sheet") or op.get("sheet_name"))
    found = _select_chart(ws, index=op.get("index", op.get("chart_index")), target_cell=op.get("target_cell"), title=op.get("old_title"))
    if found is None:
        raise ValueError("未找到要更新的图表；请提供 index、old_title 或 target_cell")
    index, chart = found
    from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, AreaChart
    classes = {"bar": BarChart, "line": LineChart, "pie": PieChart, "scatter": ScatterChart, "area": AreaChart}
    old_type = next((name for name, cls in classes.items() if isinstance(chart, cls)), None)
    new_type = _CHART_TYPE_ALIASES.get(str(op.get("chart_type", "")).lower(), op.get("chart_type")) or old_type
    if new_type is None:
        raise ValueError("该图表类型暂不支持原生修改")
    if new_type != old_type:
        raise ValueError("修改图表类型请显式删除并创建；update_chart 保留原对象的其余属性")
    if "data_range" in op or "categories_range" in op:
        if not op.get("data_range"):
            raise ValueError("修改数据系列需要 data_range")
        spec = normalize_chart_args(chart_type=new_type, data_range=op["data_range"], categories_range=op.get("categories_range"), sheet_name=ws.title, from_rows=bool(op.get("from_rows")))
        if isinstance(spec, ToolResult):
            raise ValueError(spec.model_text)
        add_chart_to_workbook(wb, spec)
        new_chart = ws._charts.pop()
        old_series = chart.series
        for pos, series in enumerate(new_chart.series):
            if pos < len(old_series):
                from copy import deepcopy
                # 旧版本散点序列没有显式 marker（symbol=none）也从未声明“点+线”意图，
                # 那是实现默认值而不是用户选择；把它复制回来会把刚修好的默认样式
                # 重新污染成“无点带线”。只有显式点样式才视为用户意图并保留。
                legacy_scatter_style = (
                    new_type == "scatter"
                    and str(getattr(getattr(old_series[pos], "marker", None), "symbol", "") or "") in ("", "none")
                )
                attrs = ("dLbls", "trendline", "errBars") if legacy_scatter_style else (
                    "graphicalProperties", "marker", "dLbls", "trendline", "errBars")
                for attr in attrs:
                    if hasattr(series, attr) and hasattr(old_series[pos], attr):
                        setattr(series, attr, deepcopy(getattr(old_series[pos], attr)))
        chart.series = new_chart.series
    for field in ("title", "style"):
        if field in op:
            setattr(chart, field, op[field])
    for field, axis in (("x_title", "x_axis"), ("y_title", "y_axis")):
        if field in op:
            if not hasattr(chart, axis):
                raise ValueError("该图表没有坐标轴")
            getattr(chart, axis).title = op[field]
    if "target_cell" in op:
        chart.anchor = op["target_cell"]
    from excelmanus.workbook.geometry import resize_drawing
    resize_drawing(chart, width=op.get("width"), height=op.get("height"), unit="cm")
    return {"target_sheet": ws.title, "index": index, "chart_type": new_type, "total_charts": len(ws._charts)}


def delete_chart_from_workbook(
    wb: Any, *, sheet_name: str | None = None, index: int | None = None,
    target_cell: str | None = None, title: str | None = None,
) -> dict[str, Any]:
    ws = get_worksheet(wb, sheet_name)
    selected = _select_chart(ws, index=index, target_cell=target_cell, title=title)
    if selected is None:
        raise MutationAborted(error_result("未找到要删除的图表", code="NOT_FOUND"))
    pos, chart = selected
    ws._charts.remove(chart)
    return {"deleted_index": pos, "target_sheet": ws.title, "remaining_charts": len(ws._charts)}
