"""Version-bound workbook observations for tools, SDK and the workbench."""

from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from typing import Any
from datetime import date, datetime, time

from openpyxl.utils import get_column_letter

from excelmanus.workbook.geometry import (
    region_geometry,
    defaults,
    axis_offset,
    column_dimension,
)
from excelmanus.workbook.refs import parse_rect
from excelmanus.workbook.snapshot import (
    WorkbookSnapshot,
    SnapshotError,
    _cached_workbook_pair,
    cell_fact_from_openpyxl,
    require_default_sheet,
    selection_from_rows,
)
from excelmanus.tools._style_extract import extract_cell_style

from excelmanus.workbook.protocol import SCHEMA_VERSION, FACETS, ObservationRequest

MAX_CELLS = 20_000


def observation_schema() -> dict[str, Any]:
    """The discoverable output contract, including geometry units and coverage."""
    axis = {
        "type": "object",
        "properties": {
            "index": {"type": "integer", "description": "Excel 1-based"},
            "native": {"type": "number"},
            "unit": {"type": "string", "enum": ["excel_char", "pt"]},
            "pixels": {"type": "number"},
            "hidden": {"type": "boolean"},
            "source": {"type": "string"},
            "pixel_status": {"type": "string"},
        },
    }
    geometry = {
        "type": "object",
        "properties": {
            "columns": {"type": "array", "items": axis},
            "rows": {"type": "array", "items": axis},
            "defaults": {"type": "object"},
            "width_px": {"type": "number"},
            "height_px": {"type": "number"},
            "aspect_ratio": {"type": ["number", "null"]},
            "status": {"type": "string"},
            "dpi": {"type": "integer"},
        },
    }
    cell = {
        "type": "object",
        "properties": {
            "v": {},
            "raw_value": {},
            "f": {"type": "string"},
            "t": {"type": "string"},
            "cached": {"type": "string", "enum": ["yes", "no", "unknown"]},
            "value_source": {"type": "string"},
            "s": {"type": "object", "description": "展开的工作台显示样式"},
            "display": {"type": "object"},
            "cached_value": {},
            "serial_value": {"type": "number"},
            "calculation_provenance": {"type": "object"},
            "formula_metadata": {"type":"object"},
        },
    }
    region = {
        "type": "object",
        "required": ["sheet", "range", "rect", "cells", "coverage"],
        "properties": {
            "sheet": {"type": "string"},
            "range": {"type": "string"},
            "rect": {"type": "object"},
            "cells": {
                "type": "object",
                "additionalProperties": cell,
                "description": "键为 1-based row,column",
            },
            "geometry": geometry,
            "coverage": {
                "type": "object",
                "description": "按 facet 报告 complete/partial/not_requested/unsupported",
            },
            "objects": {"type": "array", "items": {"type": "object"}},
            "selection": {"type": "object"},
            "merges": {"type": "array", "items": {"type": "object"}},
            "merge_anchors": {"type": "object"},
            "conditional_formatting": {"type": "object"},
            "print_settings": {"type": "object"},
        },
    }
    return {
        "type": "object",
        "required": [
            "status",
            "schema_version",
            "file_path",
            "content_version",
            "snapshot_id",
            "observation_id",
        ],
        "properties": {
            "status": {"type": "string"},
            "schema_version": {"type": "string", "enum": [SCHEMA_VERSION]},
            "file_path": {"type": "string"},
            "content_version": {"type": "string"},
            "snapshot_id": {"type": "string"},
            "observation_id": {"type": "string"},
            "regions": {"type": "array", "items": region},
            "sheets": {"type": "array", "items": {"type": "object"}},
            "coverage": {"type": "object"},
            "capabilities": {"type": "object"},
            "provenance": {"type": "object"},
            "file": {"type": "object"},
            "request": {"type": "object"},
            "matches": {"type": "array"},
            "active_sheet": {"type": "string"},
        },
    }


def _scalar(value: Any) -> Any:
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if value is not None and not isinstance(value, (str, int, float, bool)):
        return str(value)
    return value


def _bounds(cells: list) -> dict | None:
    if not cells:
        return None
    return {
        "r0": min(c.row for c in cells),
        "r1": max(c.row for c in cells),
        "c0": min(c.column for c in cells),
        "c1": max(c.column for c in cells),
    }


def _manifest(wb: Any) -> list[dict]:
    cached = getattr(wb, "_observation_manifest", None)
    if cached is not None:
        return deepcopy(cached)
    manifest = []
    for ws in wb:
        occupied, formatted = [], []
        for cell in ws._cells.values():
            if cell.value is not None:
                occupied.append(cell)
            if cell.has_style:
                formatted.append(cell)
        manifest.append(
            {
                "name": ws.title,
                "sheet_id": ws.title,
                "state": ws.sheet_state,
                "used": {"rows": ws.max_row or 1, "cols": ws.max_column or 1},
                "defaults": defaults(ws),
                "bounds": {
                    "content": _bounds(occupied),
                    "formatted": _bounds(formatted),
                    "drawing": {"status": "query_objects"},
                    "visible": {
                        "status": "query_geometry",
                        "reason": "hidden axes resolved per region",
                    },
                    "print": str(ws.print_area) or None,
                },
                "freeze_panes": str(ws.freeze_panes or ""),
                "available": {
                    "merges": len(ws.merged_cells.ranges),
                    "charts": len(ws._charts),
                    "images": len(ws._images),
                    "tables": len(ws.tables),
                    "conditional_rules": len(ws.conditional_formatting),
                },
            }
        )
    wb._observation_manifest = manifest
    return deepcopy(manifest)


def _rect_dict(rect: Any) -> dict[str, int]:
    return {
        "r0": rect.min_row,
        "c0": rect.min_col,
        "r1": rect.max_row,
        "c1": rect.max_col,
    }


def _objects(wb: Any, ws: Any) -> list[dict[str, Any]]:
    from excelmanus.workbook.objects import list_workbook_objects

    items = list_workbook_objects(wb, ws.title)
    for obj in items:
        obj["id"] = (
            f"{ws.title}:{obj['kind']}:{obj.get('name', obj.get('cell', obj.get('index', 0)))}"
        )
        obj["identity_scope"] = "snapshot"
        if obj["kind"] not in {"image", "chart"}:
            continue
        item = (ws._images if obj["kind"] == "image" else ws._charts)[obj["index"]]
        anchor = item.anchor
        start, end = getattr(anchor, "_from", None), getattr(anchor, "to", None)
        extent = getattr(anchor, "ext", None)
        obj["anchor"] = {
            "from": {
                "row": start.row + 1,
                "column": start.col + 1,
                "row_offset_emu": start.rowOff,
                "column_offset_emu": start.colOff,
            }
            if start
            else None,
            "to": {
                "row": end.row + 1,
                "column": end.col + 1,
                "row_offset_emu": end.rowOff,
                "column_offset_emu": end.colOff,
            }
            if end
            else None,
            "extent_emu": {"width": extent.cx, "height": extent.cy} if extent else None,
        }
        if start:
            x = axis_offset(ws, start.col + 1, axis="column") + start.colOff / 9525
            y = axis_offset(ws, start.row + 1, axis="row") + start.rowOff / 9525
            if end:
                width = (
                    axis_offset(ws, end.col + 1, axis="column") + end.colOff / 9525 - x
                )
                height = (
                    axis_offset(ws, end.row + 1, axis="row") + end.rowOff / 9525 - y
                )
            elif extent:
                width, height = extent.cx / 9525, extent.cy / 9525
            else:
                width = height = None
            obj["bounds"] = {
                "x": x,
                "y": y,
                "width": width,
                "height": height,
                "unit": "css_px",
                "status": "estimated",
                "coordinate_space": "sheet",
            }
        if not start:
            position = getattr(anchor, "pos", None)
            if position is not None and extent is not None:
                obj["bounds"] = {
                    "x": position.x / 9525,
                    "y": position.y / 9525,
                    "width": extent.cx / 9525,
                    "height": extent.cy / 9525,
                    "unit": "css_px",
                    "status": "unit_conversion",
                    "coordinate_space": "sheet",
                }
            else:
                obj["bounds_status"] = "unsupported_anchor"
        if obj["kind"] == "image":
            obj["source_size_px"] = {"width": item.width, "height": item.height}
            obj["format"] = item.format
        else:
            obj["chart_type"] = type(item).__name__
            obj["series"] = [
                {
                    "index": i,
                    "references": [
                        getattr(
                            getattr(getattr(series, key, None), "numRef", None),
                            "f",
                            None,
                        )
                        for key in ("val", "cat", "xVal", "yVal")
                    ],
                }
                for i, series in enumerate(item.series)
            ]
    return items


def observe_snapshot(
    snapshot: WorkbookSnapshot,
    *,
    sheet: str | None = None,
    range: str | None = None,
    facets: list[str] | None = None,
    mode: str = "overview",
    query: str = "",
    offset: int = 0,
    limit: int = 50,
    locale: str = "zh-CN",
) -> dict[str, Any]:
    chosen = list(dict.fromkeys(facets if facets is not None else ["data", "geometry"]))
    if set(chosen) - set(FACETS):
        raise ValueError(
            f"Unknown observation facets: {sorted(set(chosen) - set(FACETS))}"
        )
    if mode not in {"overview", "range", "search", "objects", "dependencies"}:
        raise ValueError("mode must be overview/range/search/objects/dependencies")
    if offset < 0 or not 1 <= limit <= 500:
        raise ValueError("offset must be nonnegative; limit must be 1..500")
    if mode in {"objects", "dependencies"} and mode not in chosen:
        chosen.append(mode)
    if mode == "range" and not range:
        from excelmanus.workbook_commit import CommitError

        raise CommitError(
            "INVALID_ARGS",
            "range mode requires a finite A1 range",
            fields={"missing_fields": ["range"]},
        )
    request = ObservationRequest(
        sheet=sheet,
        range=range,
        facets=chosen,
        mode=mode,
        query=query,
        offset=offset,
        limit=limit,
        locale=locale,
    ).canonical()
    identity = hashlib.sha256(
        (
            snapshot.id.key() + json.dumps(request, sort_keys=True, ensure_ascii=False)
        ).encode()
    ).hexdigest()
    result: dict[str, Any] = {
        "status": "success",
        "schema_version": SCHEMA_VERSION,
        "observation_id": f"obs:{identity}",
        "file_path": snapshot.file.relative,
        "content_version": snapshot.content_version,
        "snapshot_id": snapshot.id.key(),
        "file": {
            "workspaceKey": snapshot.file.workspace.identity_key(),
            "relative": snapshot.file.relative,
            "observedVersion": snapshot.content_version,
        },
        "request": request,
        "sheets": [],
        "regions": [],
        "coverage": {"facets": {}},
    }
    from importlib.metadata import version as package_version

    result["provenance"] = {
        "adapter_version": package_version("openpyxl")
        if not snapshot.is_csv()
        else "csv/1",
        "projection_version": "workbook-observation/2.1",
        "adapter": "openpyxl" if not snapshot.is_csv() else "csv",
        "calculation": "not_recalculated",
        "locale": locale,
    }
    result["file"]["lineage_id"] = snapshot.lineage_id
    if mode == "search":
        return _search_snapshot(snapshot, result, request)
    if snapshot.is_csv():
        return _observe_csv(snapshot, result, request)
    from excelmanus.workbook.adapter import package_inventory

    result["capabilities"] = package_inventory(snapshot.read_bytes())
    if range:
        from excelmanus.workbook.refs import parse_ref

        parts = parse_ref(range).areas
        if len(parts) > 1:
            if len(parts) > 32:
                raise ValueError("At most 32 observation areas are allowed")
            regions, count = [], 0
            for part in parts:
                child = observe_snapshot(
                    snapshot,
                    sheet=sheet,
                    range=part.to_a1(),
                    facets=chosen,
                    mode="range",
                    offset=offset,
                    limit=limit,
                    locale=locale,
                )
                regions.extend(child["regions"])
                count += sum(
                    (r["rect"]["r1"] - r["rect"]["r0"] + 1)
                    * (r["rect"]["c1"] - r["rect"]["c0"] + 1)
                    for r in child["regions"]
                )
                if count > MAX_CELLS:
                    raise ValueError(f"Combined observation exceeds {MAX_CELLS} cells")
            result.update(
                sheets=child["sheets"],
                active_sheet=child["active_sheet"],
                regions=regions,
            )
            return _finish_coverage(result, chosen)
    wb, cached, lock = _cached_workbook_pair(snapshot, True)
    with lock:
        names = wb.sheetnames
        if range:
            rect = parse_rect(range)
            if sheet and rect.sheet and sheet != rect.sheet:
                raise ValueError("sheet conflicts with the range qualifier")
            sheet = sheet or rect.sheet
        if sheet or mode != "overview":
            selected = [require_default_sheet(names, sheet)]
        else:
            selected = names[offset : offset + limit]
        result["active_sheet"] = (wb.active or wb.worksheets[0]).title
        # The manifest is lightweight; data/style regions are independently bounded.
        result["sheets"] = _manifest(wb)
        for name in selected:
            ws = wb[name]
            address = (
                range
                or f"A1:{get_column_letter(min(ws.max_column or 1, 12))}{min(ws.max_row or 1, 20)}"
            )
            rect = parse_rect(address)
            requested_address = address
            bounds = None
            if rect.whole_column or rect.whole_row:
                from excelmanus.workbook.address import resolve_range_to_bounds

                bounds = resolve_range_to_bounds(
                    address, used_max_row=ws.max_row, used_max_col=ws.max_column
                )
                address = bounds.resolved
                rect = parse_rect(address)
            if (rect.max_row - rect.min_row + 1) * (
                rect.max_col - rect.min_col + 1
            ) > MAX_CELLS:
                raise ValueError(
                    f"Observation region exceeds {MAX_CELLS} cells; request smaller windows"
                )
            region: dict[str, Any] = {
                "sheet": name,
                "range": rect.to_a1(include_sheet=False),
                "rect": _rect_dict(rect),
                "cells": {},
                "coverage": {f: {"status": "not_requested"} for f in FACETS},
            }
            if bounds:
                region["requested_range"] = requested_address
                region["truncated"] = bounds.truncated
                if bounds.truncated:
                    region["next_range"] = (
                        f"{get_column_letter(rect.min_col)}{rect.max_row + 1}:{get_column_letter(rect.max_col)}{min(ws.max_row, rect.max_row + rect.max_row - rect.min_row + 1)}"
                    )
            if "geometry" in chosen:
                region["geometry"] = region_geometry(ws, rect)
                region["coverage"]["geometry"] = {
                    "status": "complete",
                    "pixel_status": "estimated",
                }
            if "data" in chosen or "presentation" in chosen:
                max_row, max_col = (
                    min(rect.max_row, ws.max_row),
                    min(rect.max_col, ws.max_column),
                )
                cached_values = {
                    (c.row, c.column): c.value
                    for cells in cached[name].iter_rows(
                        min_row=rect.min_row,
                        max_row=max_row,
                        min_col=rect.min_col,
                        max_col=max_col,
                    )
                    for c in cells
                    if hasattr(c, "row")
                }
                from openpyxl.cell.cell import Cell

                for row_index in __builtins_range(rect.min_row, rect.max_row + 1):
                    for col_index in __builtins_range(rect.min_col, rect.max_col + 1):
                        cell = ws._cells.get((row_index, col_index))
                        if cell is None:
                            cell = Cell(ws, row=row_index, column=col_index)
                        inherited = [
                            ws.row_dimensions.get(cell.row),
                            column_dimension(ws, cell.column),
                        ]
                        if (
                            cell.value is None
                            and not cell.has_style
                            and not (
                                "presentation" in chosen
                                and any(
                                    d is not None and d.has_style for d in inherited
                                )
                            )
                        ):
                            continue
                        cv = (
                            cached_values.get((cell.row, cell.column))
                            if cell.data_type == "f"
                            else cell.value
                        )
                        fact = cell_fact_from_openpyxl(
                            name, cell.row, cell.column, cell, cv
                        )
                        payload = {
                            "t": fact.t,
                            "v": fact.v,
                            "cached": fact.cached,
                            "raw_value": _scalar(cell.value),
                            "value_source": (
                                "saved_cache"
                                if fact.cached == "yes"
                                else "formula_text"
                            )
                            if cell.data_type == "f"
                            else "literal",
                        }
                        if isinstance(cv, (date, datetime, time)):
                            from openpyxl.utils.datetime import to_excel

                            payload["serial_value"] = to_excel(cv, wb.epoch)
                        if cell.data_type == "f":
                            if fact.f is not None:
                                payload["f"] = fact.f
                            payload["raw_value"] = fact.f
                            if not isinstance(cell.value, str):
                                payload["formula_metadata"] = dict(cell.value)
                            payload["cached_value"] = _scalar(cv)
                            payload["calculation_provenance"] = {
                                "source": "saved_cache"
                                if cv is not None
                                else "missing_cache",
                                "engine": "unknown",
                                "input_version": snapshot.content_version,
                                "freshness": "unknown"
                                if cv is not None
                                else "unavailable",
                            }
                        elif isinstance(cell.value, str) and cell.value.startswith("="):
                            payload.update(
                                t="s",
                                v=cell.value,
                                cached="yes",
                                value_source="literal",
                            )
                        if cell.data_type == "f" and cv in (
                            "#NULL!",
                            "#DIV/0!",
                            "#VALUE!",
                            "#REF!",
                            "#NAME?",
                            "#NUM!",
                            "#N/A",
                        ):
                            payload.update(t="e", e=cv)
                        if cell.data_type == "e":
                            payload.update(t="e", e=cell.value)
                        if "presentation" in chosen:
                            from copy import copy

                            effective = copy(cell)
                            row_dim, col_dim = (
                                ws.row_dimensions.get(cell.row),
                                column_dimension(ws, cell.column),
                            )
                            if not cell.has_style:
                                if row_dim is not None and row_dim.has_style:
                                    effective._style = row_dim._style
                                elif col_dim is not None and col_dim.has_style:
                                    effective._style = col_dim._style
                            payload["s"] = extract_cell_style(effective) or {}
                            unsupported_style = (
                                getattr(effective.fill, "patternType", None)
                                not in {None, "none", "solid"}
                                or type(effective.fill).__name__ == "GradientFill"
                                or effective.alignment.horizontal
                                in {"distributed", "centerContinuous", "fill"}
                                or effective.alignment.textRotation == 255
                                or effective.font.underline
                                in {"double", "doubleAccounting"}
                                or effective.font.vertAlign
                                in {"superscript", "subscript"}
                                or bool(
                                    getattr(effective.border.diagonal, "style", None)
                                )
                            )
                            if unsupported_style:
                                from openpyxl.xml.functions import tostring

                                payload["raw_style"] = {
                                    key: tostring(
                                        getattr(effective, key).to_tree()
                                    ).decode()
                                    for key in ("font", "fill", "border", "alignment")
                                }
                                payload["style_projection"] = (
                                    "partial; raw_style preserves original properties"
                                )
                            payload["display"] = {
                                "status": "renderer_required",
                                "number_format": effective.number_format,
                                "locale": locale,
                                "displayed_text": None,
                                "renderer": None,
                            }
                        if "data" not in chosen:
                            payload = {
                                key: value
                                for key, value in payload.items()
                                if key
                                in {"s", "display", "raw_style", "style_projection"}
                            }
                        region["cells"][f"{cell.row},{cell.column}"] = payload
                for facet in ("data", "presentation"):
                    if facet in chosen:
                        region["coverage"][facet] = {
                            "status": "partial"
                            if region.get("truncated")
                            else "complete",
                            "range": region["range"],
                        }
                if "data" in chosen:
                    region["selection"] = selection_from_rows(
                        snapshot,
                        sheet=name,
                        rows=list(__builtins_range(rect.min_row, rect.max_row + 1)),
                        cols=list(__builtins_range(rect.min_col, rect.max_col + 1)),
                        origin="observe",
                    ).to_json()
            if "presentation" in chosen or "geometry" in chosen or "data" in chosen:
                merges = [
                    m
                    for m in ws.merged_cells.ranges
                    if not (
                        m.max_col < rect.min_col
                        or m.min_col > rect.max_col
                        or m.max_row < rect.min_row
                        or m.min_row > rect.max_row
                    )
                ]
                region["merges"] = [
                    {
                        "min_row": m.min_row,
                        "max_row": m.max_row,
                        "min_col": m.min_col,
                        "max_col": m.max_col,
                        "anchor": f"{get_column_letter(m.min_col)}{m.min_row}",
                    }
                    for m in merges
                ]
                region["merge_anchors"] = {
                    m.start_cell.coordinate: {
                        "value": _scalar(m.start_cell.value),
                        "formula": m.start_cell.value
                        if m.start_cell.data_type == "f"
                        else None,
                        "s": extract_cell_style(m.start_cell),
                    }
                    for m in merges
                }
            if "presentation" in chosen:
                from excelmanus.workbook.presentation import (
                    _collect_print_settings,
                    _collect_conditional_formatting,
                    _collect_data_validation,
                )

                region["print_settings"] = _collect_print_settings(ws)
                region["conditional_formatting"] = {
                    "status": "rules_only",
                    "rules": _collect_conditional_formatting(ws),
                }
                region["data_validation"] = _collect_data_validation(ws)
                region["coverage"]["presentation"]["displayed_text"] = (
                    "renderer_required"
                )
                region["coverage"]["presentation"]["limitations"] = [
                    "Conditional rules are not evaluated; displayed text requires a renderer"
                ]
                if len(ws.conditional_formatting) or any(
                    c.get("raw_style") for c in region["cells"].values()
                ):
                    region["coverage"]["presentation"]["status"] = "partial"
            if "objects" in chosen:
                objects = _objects(wb, ws)
                region["objects"] = objects[offset : offset + limit]
                if "geometry" in region:
                    boxes = [
                        obj["bounds"]
                        for obj in objects
                        if obj.get("bounds")
                        and obj["bounds"].get("width") is not None
                        and obj["bounds"].get("height") is not None
                    ]
                    region["geometry"]["drawing_bounds"] = (
                        {
                            "x": min(b["x"] for b in boxes),
                            "y": min(b["y"] for b in boxes),
                            "right": max(b["x"] + b["width"] for b in boxes),
                            "bottom": max(b["y"] + b["height"] for b in boxes),
                            "coordinate_space": "sheet_css_px",
                            "status": "estimated",
                        }
                        if boxes
                        else None
                    )
                region["coverage"]["objects"] = {
                    "status": "unsupported"
                    if result["capabilities"]["unsupported_features"]
                    else "truncated"
                    if offset + limit < len(objects)
                    else "complete",
                    "scope": "sheet",
                    "total": len(objects),
                    "next_offset": offset + limit
                    if offset + limit < len(objects)
                    else None,
                }
            if "dependencies" in chosen:
                region["dependencies"] = [
                    {
                        "cell": c.coordinate,
                        "formula": c.value
                        if isinstance(c.value, str)
                        else getattr(c.value, "text", None),
                    }
                    for c in ws._cells.values()
                    if c.data_type == "f"
                    and rect.min_row <= c.row <= rect.max_row
                    and rect.min_col <= c.column <= rect.max_col
                ]
                region["coverage"]["dependencies"] = {
                    "status": "partial",
                    "scope": "formula_text",
                    "reason": "Use formula trace for transitive dependencies",
                }
            result["regions"].append(region)
        result["coverage"]["next_offset"] = (
            offset + limit
            if mode == "overview" and not sheet and offset + limit < len(names)
            else None
        )
    return _finish_coverage(result, chosen)


from builtins import range as __builtins_range


def _observe_csv(snapshot: WorkbookSnapshot, result: dict, request: dict) -> dict:
    from excelmanus.workbook.csv_index import csv_index

    index = csv_index(snapshot)
    if request["sheet"] not in (None, "Sheet1"):
        raise SnapshotError("CSV has one logical sheet: Sheet1", code="SHEET_NOT_FOUND")
    address = request["range"] or "A1:L20"
    rect = parse_rect(address)
    if (
        rect.whole_column
        or rect.whole_row
        or (rect.max_row - rect.min_row + 1) * (rect.max_col - rect.min_col + 1)
        > MAX_CELLS
    ):
        raise ValueError(
            "CSV observation requires a bounded region of at most 20000 cells"
        )
    cells = {}
    for row, values in enumerate(
        index.window(rect.min_row, rect.max_row) if "data" in request["facets"] else [],
        rect.min_row,
    ):
        for col in __builtins_range(rect.min_col, min(rect.max_col, len(values)) + 1):
            value = values[col - 1]
            if value != "":
                # CSV has no stored types; preserve identifiers and leading zeros.
                cells[f"{row},{col}"] = {
                    "t": "s",
                    "v": value,
                    "raw_value": value,
                    "cached": "yes",
                    "value_source": "literal",
                }
    coverage = {
        f: {
            "status": "not_requested"
            if f not in request["facets"]
            else "complete"
            if f == "data"
            else "unsupported"
        }
        for f in FACETS
    }
    result["active_sheet"] = "Sheet1"
    result["sheets"] = [
        {
            "name": "Sheet1",
            "sheet_id": "Sheet1",
            "state": "visible",
            "used": {"rows": index.rows, "cols": index.columns},
        }
    ]
    result["regions"] = [
        {
            "sheet": "Sheet1",
            "range": address,
            "rect": _rect_dict(rect),
            "cells": cells,
            "coverage": coverage,
        }
    ]
    if "data" in request["facets"]:
        result["regions"][0]["selection"] = selection_from_rows(
            snapshot,
            sheet="Sheet1",
            rows=list(__builtins_range(rect.min_row, rect.max_row + 1)),
            cols=list(__builtins_range(rect.min_col, rect.max_col + 1)),
            origin="observe",
        ).to_json()
    return _finish_coverage(result, request["facets"])


def _finish_coverage(result: dict, chosen: list[str]) -> dict:
    """Exact rectangle subtraction, including holes between disjoint windows."""
    unloaded = []
    for meta in result["sheets"]:
        pending = [
            {"r0": 1, "c0": 1, "r1": meta["used"]["rows"], "c1": meta["used"]["cols"]}
        ]
        for region in result["regions"]:
            if region["sheet"] != meta["name"]:
                continue
            cut = region["rect"]
            remaining = []
            for box in pending:
                r0, r1 = max(box["r0"], cut["r0"]), min(box["r1"], cut["r1"])
                c0, c1 = max(box["c0"], cut["c0"]), min(box["c1"], cut["c1"])
                if r0 > r1 or c0 > c1:
                    remaining.append(box)
                    continue
                for a, b, c, d in [
                    (box["r0"], r0 - 1, box["c0"], box["c1"]),
                    (r1 + 1, box["r1"], box["c0"], box["c1"]),
                    (r0, r1, box["c0"], c0 - 1),
                    (r0, r1, c1 + 1, box["c1"]),
                ]:
                    if a <= b and c <= d:
                        remaining.append({"r0": a, "r1": b, "c0": c, "c1": d})
            pending = remaining
        unloaded.extend({"sheet": meta["name"], **box} for box in pending)
    result["coverage"].update(
        loaded=[{"sheet": r["sheet"], **r["rect"]} for r in result["regions"]],
        unloaded=unloaded,
        scope="requested regions; manifest covers workbook",
        has_more=bool(unloaded) or bool(result["coverage"].get("next_offset")),
        facets={
            f: {"status": "region_scoped" if f in chosen else "not_requested"}
            for f in FACETS
        },
    )
    return result


def _search_snapshot(snapshot, result, request):
    """Search the pinned bytes; a cursor never reopens the live file."""
    query, offset, limit = (
        request["query"].casefold(),
        request["offset"],
        request["limit"],
    )
    matches, total = [], 0

    def scan(name, rows):
        nonlocal total
        for row_number, row in enumerate(rows, 1):
            if row_number % 512 == 0:
                from excelmanus.tools.spreadsheet_engine_tools import _cancelled

                _cancelled()
            for col, value in enumerate(row, 1):
                if value is not None and query and query in str(value).casefold():
                    if offset <= total < offset + limit:
                        matches.append(
                            {
                                "sheet": name,
                                "row": row_number,
                                "column_index": col,
                                "cell_ref": f"{get_column_letter(col)}{row_number}",
                                "value": _scalar(value),
                            }
                        )
                    total += 1

    if snapshot.is_csv():
        from excelmanus.workbook.csv_index import csv_index

        index = csv_index(snapshot)
        if request["sheet"] not in (None, "Sheet1"):
            raise SnapshotError(
                "CSV has one logical sheet: Sheet1", code="SHEET_NOT_FOUND"
            )
        scan("Sheet1", index.window(1, index.rows))
    else:
        book = snapshot.open_workbook(data_only=False, read_only=True)
        try:
            names = (
                [require_default_sheet(book.sheetnames, request["sheet"])]
                if request["sheet"]
                else book.sheetnames
            )
            for name in names:
                scan(name, book[name].iter_rows(values_only=True))
        finally:
            book.close()
    result.update(
        matches=matches,
        total_matches=total,
        returned=len(matches),
        truncated=offset + limit < total,
    )
    result["coverage"] = {
        "scope": "literal values and formula text",
        "facets": {
            f: {
                "status": "complete"
                if f == "data"
                else "unsupported"
                if f in request["facets"]
                else "not_requested"
            }
            for f in FACETS
        },
        "next_offset": offset + limit if offset + limit < total else None,
        "has_more": offset + limit < total,
        "offset": offset,
        "loaded": [],
        "unloaded": [],
    }
    return result
