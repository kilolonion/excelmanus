"""Workbook algorithms used by the V2 domain service; no tool registration."""

from __future__ import annotations


import json


import re


from typing import Any, NoReturn


from openpyxl.utils import column_index_from_string, get_column_letter


from openpyxl.utils.cell import coordinate_to_tuple


from excelmanus.engine_core.tool_result import ToolResult, ToolUiMeta, error_result, from_payload


from excelmanus.logger import get_logger


from excelmanus.prompt.canonical import TOOL_DESCRIPTIONS


from excelmanus.security import FileAccessGuard, SecurityViolationError


from excelmanus.tools.context import bind_workspace, require_guard, current_call


from excelmanus.tools._helpers import (
    MutationAborted,
    OUTSIDE_WORKSPACE_MESSAGE,
    commit_error_result,
    commit_workbook_tool,
    get_worksheet,
    prepare_excel_commit_path,
    resolve_sheet_name,
    unwrap_mutation_abort,
    workspace_relpath,
)


from excelmanus.workbook.address import (
    combine_sheet_names,
    looks_like_coordinate_error,
    parse_sheet_address,
    resolve_range_to_bounds,
    top_left_cell,
    worksheet_used_shape,
)


from excelmanus.workbook.refs import (
    CellRef,
    InvalidRefError,
    NamedRef,
    RectRef,
    TableRef,
    describe_for_schema,
    parse_ref,
)


from excelmanus.workbook.snapshot import parse_bound_selection


from excelmanus.workbook.cells import assign_cell_value, _resolve_merged_cell


from excelmanus.workbook.styles import (
    _build_border,
    _build_fill,
    _patch_alignment,
    _patch_font,
    apply_column_sizes,
    apply_freeze_panes,
    apply_row_sizes,
)


from excelmanus.tools.registry import ToolDef


from excelmanus.workbook_commit import (
    CommitError,
    commit_bytes,
    content_version_of_file,
    peek_seen_content_version,
)






def _get_guard() -> FileAccessGuard:
    return require_guard()


def _op_get(op: dict[str, Any], *keys: str, default: Any = None) -> Any:
    for key in keys:
        if key in op and op[key] is not None:
            return op[key]
    return default


def _ui_from_value(payload: dict[str, Any]) -> ToolUiMeta:
    ui = ToolUiMeta()
    version = payload.get("content_version")
    if isinstance(version, str) and version:
        ui.content_version = version
    file_path = payload.get("file_path") or payload.get("path")
    if isinstance(file_path, str) and file_path.strip():
        ui.files = [file_path]
    revision = payload.get("revision")
    if isinstance(revision, dict):
        ui.revision = dict(revision)
    restored = payload.get("restored_revision")
    if isinstance(restored, str) and restored:
        extra = dict(ui.revision or {})
        extra.setdefault("revision_id", restored)
        ui.revision = extra
    return ui


def _summarize_payload(payload: dict[str, Any]) -> str:
    """给模型的有界摘要：能决策即可，禁止整包 JSON。"""
    if payload.get("mode") == "capabilities":
        tools = payload.get("model_facing") or []
        return "表格意图: " + ", ".join(str(t) for t in tools)
    parts: list[str] = []
    path = payload.get("file_path") or payload.get("path") or ""
    version = payload.get("content_version") or ""
    if path:
        parts.append(str(path))
    if version:
        parts.append(str(version))
    applied = payload.get("applied")
    if isinstance(applied, list) and applied:
        shown = ", ".join(str(item) for item in applied[:12])
        extra = f" (+{len(applied) - 12})" if len(applied) > 12 else ""
        parts.append(f"applied: {shown}{extra}")
    revision = payload.get("revision")
    if isinstance(revision, dict) and revision.get("revision_id"):
        parts.append(f"revision={revision['revision_id']}")
    revisions = payload.get("revisions")
    if isinstance(revisions, list):
        parts.append(f"revisions={len(revisions)}（本次返回）")
        parts.extend(
            f"revision={item.get('revision_id')} {item.get('reason')} {item.get('label', '')} {item.get('content_version')}"
            for item in revisions[-5:] if isinstance(item, dict)
        )
    if payload.get("restored_revision"):
        parts.append(f"restored_revision={payload['restored_revision']}")
    uncertainties = payload.get("uncertainties")
    if isinstance(uncertainties, list):
        parts.append(f"uncertainties={len(uncertainties)}")
    for key in ("files", "warnings", "verification", "appearance", "skipped_merged_non_anchors"):
        if payload.get(key):
            parts.append(f"{key}: {json.dumps(payload[key], ensure_ascii=False, default=str)}")
    if parts:
        return " ".join(parts)
    keys = [str(key) for key in payload if key != "status"]
    return "ok " + " ".join(keys[:8]) if keys else "ok"


def _sheet_missing_message(name: str, available: list[str]) -> str:
    """工作表不存在时的统一错误文案：附最接近候选，便于一轮纠正。"""
    from difflib import SequenceMatcher

    best, ratio = None, 0.0
    target = name.lower()
    for candidate in available:
        r = SequenceMatcher(None, target, candidate.lower()).ratio()
        if r > ratio:
            best, ratio = candidate, r
    base = f"工作表 '{name}' 不存在。该文件包含: {list(available)}"
    if best is not None and ratio >= 0.5:
        return f"{base}。最接近的是 '{best}'（相似度 {ratio:.0%}），请确认名称后重试。"
    return f"{base}。"


def _sheet_not_found(name: str, available: list[str]) -> ToolResult:
    return _invalid(
        _sheet_missing_message(name, available),
        code="SHEET_NOT_FOUND",
        available_sheets=list(available),
    )


def _invalid(message: str, *, code: str = "INVALID_ARGS", **extra: Any) -> ToolResult:
    return error_result(message, code=code, fields=extra or None)


_JSON_TRAILING_JUNK_RE = re.compile(r"^[\s}\],]*$")


_JSON_RECOVER_MISSING = object()




_ANALYZE_MODE_FIELDS: dict[str, frozenset[str]] = {
    "profile": frozenset({
        "request", "mode", "file_path", "path", "sheet", "sheet_name", "max_rows",
        "sample_rows", "limit",
        "expected_version", "content_version", "header_row",
    }),
    "quality": frozenset({
        "request", "mode", "file_path", "path", "sheet", "sheet_name", "max_rows",
        "sample_rows", "limit",
        "expected_version", "content_version", "header_row",
    }),
    "filter": frozenset({
        "request", "mode", "file_path", "path", "sheet", "sheet_name", "header_row",
        "column", "operator", "value", "conditions", "logic", "columns",
        "max_rows", "sort_by", "ascending", "limit",
        "expected_version", "content_version",
    }),
    "aggregate": frozenset({
        "request", "mode", "file_path", "path", "sheet", "sheet_name", "header_row",
        "group_by", "aggregations", "conditions", "logic", "join",
        "column", "operator", "value", "sort_by", "ascending", "limit", "max_rows",
        "expected_version", "content_version",
    }),
    "distinct": frozenset({
        "request", "mode", "file_path", "path", "sheet", "sheet_name", "header_row",
        "column", "conditions", "logic", "limit", "max_rows", "dup_only",
        "expected_version", "content_version",
    }),
    "relationships": frozenset({
        "request", "mode", "file_path", "path", "file_paths", "paths",
        "directory", "max_files", "sample_rows",
    }),
    "pivot": frozenset({
        "request", "mode", "file_path", "path", "sheet", "sheet_name", "header_row",
        "index", "columns", "values", "aggfunc", "group_by",
        "conditions", "logic", "join", "column", "operator", "value", "limit", "max_rows",
        "margins", "margins_name", "totals", "totals_name", "grand_total",
        "expected_version", "content_version",
    }),
    "files": frozenset({
        "request", "mode", "directory", "include", "query", "max_files", "file_path", "path",
    }),
}


_ANALYZE_DEFAULTS = {
    "logic": "and",
    "ascending": True,
    "directory": ".",
}


_TRACE_MODE_FIELDS: dict[str, frozenset[str]] = {
    "map": frozenset({"request", "mode", "file_path", "path", "detail"}),
    "trace": frozenset({
        "request", "mode", "file_path", "path", "target", "direction", "depth",
    }),
    "impact": frozenset({
        "request", "mode", "file_path", "path", "target", "scope",
    }),
}


_TRACE_DEFAULTS = {
    "direction": "both",
    "depth": 2,
    "detail": "summary",
    "scope": "all",
}


_EDIT_KIND_FIELDS = {
    "write": "sheet sheet_name start_cell startCell start cell values selection source_rows content_version",
    "insert": "sheet sheet_name axis at row column count",
    "delete_rows": "sheet sheet_name at row column count selection source_rows content_version",
    "delete_columns": "sheet sheet_name at row column count",
    "sheet": "sheet sheet_name action new_name newName",
    "copy": "sheet sheet_name source_sheet sourceSheet source_range sourceRange target_sheet targetSheet target_start targetStart",
    "pivot": "sheet sheet_name target_sheet new_name index columns pivot_columns values pivot_values group_by aggfunc header_row join column operator value conditions logic margins margins_name totals totals_name grand_total overwrite refresh",
    "pivot_refresh": "target_sheet new_name overwrite",
    "transform": "sheet sheet_name header_row action transform key_columns key_normalizers keep order_by column delimiter sep new_columns into",
}


from excelmanus.workbook.operations import KINDS as _RANGE_KINDS


_FORMAT_KIND_FIELDS = {
    "format": "range font fill border alignment number_format",
    "size": "range column_widths row_heights auto_fit axis",
    "freeze": "range freeze_panes rows cols",
    "print_layout": "print_layout",
    "merge": "range allow_data_loss",
    "unmerge": "range",
    "conditional_format": "range rule remove",
    "data_validation": "range rule remove",
}


def _reject_operation_fields(op: dict[str, Any], label: str, fields: str) -> None:
    allowed = set(fields.split()) | {"kind", "sheet", "sheet_name"}
    extra = sorted(set(op) - allowed)
    if extra:
        raise MutationAborted(_invalid(
            f"{label} 不接受字段 {extra}；请使用该操作的字段，其他操作分成独立项。",
            invalid_fields=extra, accepted_fields=sorted(allowed),
        ))


def _reject_edit_fields(op: dict[str, Any], kind: str) -> None:
    fields = _EDIT_KIND_FIELDS.get(kind)
    if fields is None:
        return  # the dispatch branch returns the unsupported-kind error
    allowed = set(fields.split()) | {"kind"}
    extra = sorted(set(op) - allowed)
    if extra:
        raise MutationAborted(_invalid(
            f"edit.kind={kind} 不接受字段 {extra}；请用该操作的字段或对应意图工具。",
            invalid_fields=extra, accepted_fields=sorted(allowed),
        ))


def _abort_operation(exc: MutationAborted, index: int, kind: str) -> NoReturn:
    """A failed atomic batch must identify the bad operation, not imply a partial write."""
    payload = dict(exc.result.value or {})
    message = payload.get("message") or exc.result.model_text
    payload.update(
        message=f"operations[{index}] (kind={kind}): {message}",
        operation_index=index,
        operation_kind=kind,
        committed=False,
        partial=False,
        applied=[],
    )
    raise MutationAborted(from_payload(payload)) from exc


def _reject_mode_fields(
    args: dict[str, Any],
    mode: str,
    allowed: frozenset[str],
    defaults: dict[str, Any],
    *,
    required_for_mode: list[str] | None = None,
    missing_required: list[str] | None = None,
) -> ToolResult | None:
    extras: list[str] = []
    for key, value in args.items():
        if key in allowed:
            continue
        if key in defaults and value == defaults[key]:
            continue
        if value in (None, "", [], {}):
            continue
        extras.append(str(key))
    missing = [str(item) for item in (missing_required or []) if str(item)]
    if not extras and not missing:
        return None
    bits: list[str] = []
    if extras:
        bits.append(
            f"mode={mode} 不接受字段：{', '.join(extras)}；"
            f"可用字段: {', '.join(sorted(allowed))}"
        )
    if missing:
        bits.append(f"mode={mode} 需要 {', '.join(missing)}")
    payload: dict[str, Any] = {
        "accepted_fields": sorted(allowed),
        "required_for_mode": list(required_for_mode or []),
    }
    if extras:
        payload["ignored_fields"] = extras
    if missing:
        payload["missing_fields"] = missing
    return _invalid("；".join(bits), **payload)


def _arg_provided(args: dict[str, Any], key: str) -> bool:
    if key not in args:
        return False
    value = args[key]
    if value is None or value == "":
        return False
    if isinstance(value, dict) and not value:
        return False
    if isinstance(value, list) and not value and key != "conditions":
        return False
    return True


_ANALYZE_REQUIRED_FOR_MODE: dict[str, list[str]] = {
    "pivot": ["index", "columns", "values"],
    "filter": ["column+operator+value 或 conditions"],
}


_INSPECT_ONLY_ANALYZE_MODES = frozenset({"overview", "range", "search", "capabilities"})


def _analyze_missing_required(mode: str, args: dict[str, Any]) -> list[str]:
    if mode == "pivot":
        missing: list[str] = []
        if not (_arg_provided(args, "index") or _arg_provided(args, "group_by")):
            missing.append("index")
        if not _arg_provided(args, "columns"):
            missing.append("columns")
        if not _arg_provided(args, "values"):
            missing.append("values")
        return missing
    if mode == "filter":
        has_conditions = _arg_provided(args, "conditions")
        has_triple = all(_arg_provided(args, key) for key in ("column", "operator", "value"))
        if _arg_provided(args, "column") and args.get("operator") in {"isnull", "notnull", "is_null", "not_null", "is_not_null", "empty", "not_empty", "null"}:
            has_triple = True
        if has_conditions or has_triple:
            return []
        return ["column+operator+value 或 conditions"]
    return []


def _address_has_sheet(raw: Any) -> bool:
    if raw in (None, ""):
        return False
    try:
        area = parse_ref(str(raw))
    except InvalidRefError:
        return False
    return any(bool(getattr(part, "sheet", None)) for part in area.areas)


def _workbook_sheetnames(wb: Any) -> list[str]:
    return [str(n) for n in list(getattr(wb, "sheetnames", []) or []) if str(n)]


def _op_has_sheet(
    op: dict[str, Any],
    *address_keys: str,
    sheet_keys: tuple[str, ...] = ("sheet", "sheet_name"),
) -> bool:
    if _op_get(op, *sheet_keys):
        return True
    return any(_address_has_sheet(_op_get(op, key)) for key in address_keys)


def _maybe_bind_unique_sheet(
    wb: Any,
    op: dict[str, Any],
    *address_keys: str,
    sheet_keys: tuple[str, ...] = ("sheet", "sheet_name"),
) -> None:
    """I8：恰好一张表且未指定 sheet 时自动绑定。多表绝不挑选。"""
    if wb is None or _op_has_sheet(op, *address_keys, sheet_keys=sheet_keys):
        return
    names = _workbook_sheetnames(wb)
    if len(names) == 1:
        op.setdefault("sheet", names[0])


def _require_explicit_sheet(
    op: dict[str, Any],
    action: str,
    *address_keys: str,
    sheet_keys: tuple[str, ...] = ("sheet", "sheet_name"),
    wb: Any = None,
) -> None:
    _maybe_bind_unique_sheet(wb, op, *address_keys, sheet_keys=sheet_keys)
    if _op_has_sheet(op, *address_keys, sheet_keys=sheet_keys):
        return
    names = _workbook_sheetnames(wb) if wb is not None else []
    extra: dict[str, Any] = {}
    if names:
        extra["available_sheets"] = names
    code = "SHEET_REQUIRED" if len(names) > 1 else "INVALID_ARGS"
    raise MutationAborted(
        _invalid(f"{action} 必须提供 sheet，或在地址里写 表!A1", code=code, **extra)
    )


def _with_resolved_sheet(
    result: ToolResult,
    sheet: str | None = None,
    sheets: list[str] | None = None,
) -> ToolResult:
    if not result.success or not isinstance(result.value, dict):
        return result
    if sheet:
        result.value.setdefault("resolved_sheet", sheet)
    if sheets:
        result.value.setdefault("resolved_sheets", sheets)
    return result


def _worksheet(wb: Any, sheet: str | None) -> Any:
    try:
        return get_worksheet(wb, sheet)
    except ValueError as exc:
        names = list(getattr(wb, "sheetnames", []) or [])
        raise MutationAborted(_sheet_not_found(str(sheet or ""), names)) from exc


def _split_op_address(op: dict[str, Any], raw: str) -> tuple[str | None, str]:
    try:
        parsed = parse_sheet_address(raw)
    except InvalidRefError as exc:
        raise MutationAborted(_invalid(str(exc), code="RANGE_INVALID")) from exc
    explicit = _op_get(op, "sheet")
    try:
        sheet = combine_sheet_names(
            str(explicit) if explicit not in (None, "") else None,
            parsed.sheet,
        )
    except ValueError as exc:
        raise MutationAborted(_invalid(str(exc))) from exc
    return sheet, parsed.address


def _abort_bad_address(field: str, raw: str, exc: BaseException | None = None) -> NoReturn:
    if isinstance(exc, InvalidRefError):
        raise MutationAborted(_invalid(str(exc), code="RANGE_INVALID"))
    extra = f"：{exc}" if exc is not None else ""
    raise MutationAborted(
        _invalid(
            f"{field}={raw!r} 不是合法坐标{extra}。"
            "请写 A1、A1:C5，或同表并集 A4,A11。",
            code="RANGE_INVALID",
        )
    )


def _format_rects(
    op: dict[str, Any],
    raw_range: str,
    *,
    allow_union: bool,
) -> tuple[str | None, list[str]]:
    try:
        area = parse_ref(raw_range)
    except InvalidRefError as exc:
        raise MutationAborted(_invalid(str(exc), code="RANGE_INVALID")) from exc
    if not allow_union and len(area.areas) != 1:
        raise MutationAborted(
            _invalid(
                f"merge/unmerge 只要一个矩形，不要写并集。当前 range={raw_range!r}。",
                code="REF_UNSUPPORTED",
            )
        )
    explicit = _op_get(op, "sheet")
    explicit_text = str(explicit) if explicit not in (None, "") else None
    sheets: list[str | None] = []
    locals_: list[str] = []
    for part in area.areas:
        if isinstance(part, (NamedRef, TableRef)):
            raise MutationAborted(
                _invalid(
                    "format 不接受命名区域或表引用。请写成 A1:C5 或同表并集 A4,A11。",
                    code="REF_UNSUPPORTED",
                )
            )
        if not isinstance(part, (CellRef, RectRef)):
            raise MutationAborted(
                _invalid(f"format 不支持该引用：{raw_range!r}。", code="REF_UNSUPPORTED")
            )
        try:
            sheet = combine_sheet_names(explicit_text, part.sheet)
        except ValueError as exc:
            raise MutationAborted(_invalid(str(exc))) from exc
        sheets.append(sheet)
        locals_.append(part.to_a1(include_sheet=False).replace("$", ""))
    unique_sheets = {name for name in sheets}
    if len(unique_sheets) > 1:
        raise MutationAborted(
            _invalid(
                "format 并集必须在同一张表。跨表请拆成多次 format。",
                code="REF_UNSUPPORTED",
            )
        )
    return sheets[0], locals_


def _clip_format_local_range(ws: Any, local_a1: str) -> str:
    """整列/整行裁到已用范围；有界矩形原样返回。"""
    used_rows, used_cols = worksheet_used_shape(ws)
    bounds = resolve_range_to_bounds(
        local_a1,
        used_max_row=used_rows,
        used_max_col=used_cols,
        max_cells=max(used_rows * used_cols, 1),
    )
    return bounds.resolved


def _freeze_cell_from_op(op: dict[str, Any]) -> str:
    """kind=freeze 的目标格。A2=冻结首行；空串取消。"""
    raw = _op_get(op, "freeze_panes")
    rows_raw = _op_get(op, "rows")
    cols_raw = _op_get(op, "cols")
    counts_provided = rows_raw is not None or cols_raw is not None
    if raw is not None:
        if counts_provided:
            raise ValueError("freeze_panes 与 rows/cols 不能同时指定")
        if not isinstance(raw, str):
            raise ValueError("freeze_panes 需要单格地址字符串；空字符串取消冻结")
        text = raw.strip()
        aliases = {"first_row": "A2", "首行": "A2", "row1": "A2",
                   "first_col": "B1", "first_column": "B1", "首列": "B1", "col1": "B1"}
        return aliases.get(text.lower(), text)
    if not counts_provided:
        raise ValueError('freeze 需要 freeze_panes="A2" 等单格地址；取消冻结用 freeze_panes=""')
    freeze_rows = 0
    freeze_cols = 0
    for name, raw_value in (("rows", rows_raw), ("cols", cols_raw)):
        if raw_value is None:
            continue
        limit = 1048575 if name == "rows" else 16383
        if isinstance(raw_value, bool) or not isinstance(raw_value, int) or not 0 <= raw_value <= limit:
            raise ValueError(f"freeze.{name} 必须是 0 到 {limit} 的整数")
    freeze_rows = int(rows_raw or 0)
    freeze_cols = int(cols_raw or 0)
    if freeze_rows <= 0 and freeze_cols <= 0:
        return ""
    col_letter = get_column_letter(freeze_cols + 1)
    return f"{col_letter}{freeze_rows + 1}"


def _paint_format_cells(
    ws: Any,
    cell_range: str,
    raw_range: str,
    *,
    font_cfg: Any,
    fill: Any,
    border: Any,
    align_cfg: Any,
    number_format: Any,
) -> list[str]:
    try:
        data = ws[cell_range]
    except Exception as exc:
        _abort_bad_address("range", raw_range or cell_range, exc)
    if not isinstance(data, tuple):
        rows_data = ((data,),)
    elif data and not isinstance(data[0], tuple):
        rows_data = (data,)
    else:
        rows_data = data
    written: set[str] = set()
    skipped: list[str] = []
    from openpyxl.cell.cell import MergedCell

    for row in rows_data:
        for cell in row if isinstance(row, tuple) else (row,):
            coord = getattr(cell, "coordinate", "")
            if isinstance(cell, MergedCell):
                skipped.append(coord)
                cell = _merged_anchor_cell(ws, cell)
                if cell is None:
                    continue
                coord = getattr(cell, "coordinate", "")
            if coord in written:
                continue
            if font_cfg:
                cell.font = _patch_font(cell.font, font_cfg)
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border
            if align_cfg:
                cell.alignment = _patch_alignment(cell.alignment, align_cfg)
            if number_format:
                cell.number_format = str(number_format)
            if coord:
                written.add(coord)
    return skipped


def _success(payload: dict[str, Any]) -> ToolResult:
    payload.setdefault("status", "success")
    return ToolResult(
        success=True,
        model_text=_summarize_payload(payload),
        value=payload,
        ui_meta=_ui_from_value(payload),
    )


def _merge_request(request: dict[str, Any] | None, **kwargs: Any) -> dict[str, Any]:
    merged: dict[str, Any] = dict(request or {})
    for key, value in kwargs.items():
        if value is None:
            continue
        if value == "" and key not in merged:
            continue
        if key not in merged or merged[key] in (None, ""):
            merged[key] = value
    if "file_path" not in merged or not merged.get("file_path"):
        merged["file_path"] = merged.get("path") or merged.get("file_a") or ""
    if "sheet_name" not in merged or not merged.get("sheet_name"):
        merged["sheet_name"] = merged.get("sheet")
    return merged


def _bind_write_origin(
    wb: Any,
    raw_start: str,
    explicit_sheet: Any,
    *,
    create_sheet_if_missing: bool = False,
) -> tuple[str, int, int]:
    from excelmanus.workbook.data import WorkbookRefBindError, _bind_area_in_workbook
    from excelmanus.workbook.refs import parse_ref
    from excelmanus.workbook.snapshot import SnapshotError

    default = str(explicit_sheet).strip() if explicit_sheet not in (None, "") else None
    try:
        area = parse_ref(raw_start, default_sheet=default)
    except InvalidRefError as exc:
        _abort_bad_address("start_cell", raw_start, exc)
    if len(area.areas) != 1:
        raise MutationAborted(
            _invalid(
                "write.start_cell 不支持并集。一次 write 只写一个矩形。",
                code="REF_UNSUPPORTED",
            )
        )
    try:
        rects = _bind_area_in_workbook(wb, area, default_sheet=default)
    except WorkbookRefBindError as exc:
        # create_workbook 场景下目标表尚不存在是正常用法：按 write 指定的表名建表后重绑一次
        target_sheets = area.sheets()
        if (
            create_sheet_if_missing
            and getattr(exc, "code", None) == "SHEET_NOT_FOUND"
            and target_sheets
        ):
            wb.create_sheet(title=target_sheets[0])
            try:
                rects = _bind_area_in_workbook(wb, area, default_sheet=default)
            except (WorkbookRefBindError, SnapshotError) as exc2:
                raise MutationAborted(
                    _invalid(str(exc2), code=getattr(exc2, "code", None))
                ) from exc2
        else:
            raise MutationAborted(_invalid(str(exc), code=exc.code)) from exc
    except SnapshotError as exc:
        raise MutationAborted(_invalid(str(exc), code=exc.code)) from exc
    if len(rects) != 1:
        raise MutationAborted(
            _invalid(
                "write.start_cell 命名/表引用解析为多区域，不能一次写入。",
                code="REF_UNSUPPORTED",
            )
        )
    rect = rects[0]
    if rect.whole_column or rect.whole_row:
        raise MutationAborted(
            _invalid(
                "write 不支持整轴引用。请写有限矩形，例如 A1:A100。",
                code="REF_UNSUPPORTED",
            )
        )
    return str(rect.sheet), int(rect.min_row), int(rect.min_col)


def _write_matrix(
    wb: Any,
    *,
    sheet: str,
    row0: int,
    col0: int,
    values: list[Any],
) -> str:
    ws = _worksheet(wb, sheet)
    widths = [len(row) if isinstance(row, list) else 1 for row in values]
    if not widths or not widths[0] or len(set(widths)) != 1:
        raise MutationAborted(_invalid(
            "values 必须是非空矩形；不同宽度请拆成多次 write。null 表示清空该格。"
        ))
    width = widths[0]
    _write_selection_cells(ws, rows=list(range(row0, row0 + len(values))),
                           cols=list(range(col0, col0 + width)), values=values)
    return f"{get_column_letter(col0)}{row0}:{get_column_letter(col0 + width - 1)}{row0 + len(values) - 1}"


def _write_selection_cells(
    ws: Any,
    *,
    rows: list[int],
    cols: list[int],
    values: list[Any],
) -> None:
    """Apply a selection write through the same merged-cell contract as write."""
    placements: list[tuple[int, int, int, int, Any, bool]] = []
    for r_idx, excel_row in enumerate(rows):
        cells = values[r_idx] if isinstance(values[r_idx], list) else [values[r_idx]]
        for c_idx, raw in enumerate(cells):
            actual_row, actual_col, redirected = _resolve_merged_cell(
                ws, excel_row, cols[c_idx]
            )
            placements.append(
                (excel_row, cols[c_idx], actual_row, actual_col, raw, redirected)
            )

    collisions: dict[tuple[int, int], list[str]] = {}
    for src_row, src_col, actual_row, actual_col, raw, _redirected in placements:
        if raw is None:
            continue
        collisions.setdefault((actual_row, actual_col), []).append(
            f"{get_column_letter(src_col)}{src_row}"
        )
    conflicted = {
        anchor: sources for anchor, sources in collisions.items() if len(sources) > 1
    }
    if conflicted:
        details = []
        for (ar, ac), sources in conflicted.items():
            details.append(f"{get_column_letter(ac)}{ar} <- {', '.join(sources)}")
        raise MutationAborted(
            _invalid(
                "合并区多个非空值指向同一锚点："
                + "; ".join(details)
                + "。请只写锚点格，或先 unmerge。"
            )
        )

    for _src_row, _src_col, actual_row, actual_col, raw, redirected in placements:
        if redirected and raw is None:
            continue
        assign_cell_value(ws, actual_row, actual_col, raw)


def _write_from_selection(wb: Any, op: dict[str, Any], values: list[Any]) -> str:
    sel = parse_bound_selection(_op_get(op, "selection"))
    if sel is None:
        raise MutationAborted(_invalid("selection 无法解析", code="SELECTION_STALE"))
    rows = [int(r) for r in sel.rows]
    if len(values) != len(rows):
        raise MutationAborted(
            _invalid(
                f"selection 有 {len(rows)} 行，values 有 {len(values)} 行，必须对齐。",
                code="COORD_CONTRACT",
            )
        )
    if sel.cols:
        cols = [int(c) for c in sel.cols]
        width = None
        for row in values:
            cells = row if isinstance(row, list) else [row]
            if width is None:
                width = len(cells)
            elif len(cells) != width:
                raise MutationAborted(_invalid("values 必须是矩形", code="COORD_CONTRACT"))
        if width != len(cols):
            raise MutationAborted(
                _invalid(
                    f"selection 有 {len(cols)} 列，values 宽度 {width}，必须对齐。",
                    code="COORD_CONTRACT",
                )
            )
        ws = _worksheet(wb, sel.sheet)
        _write_selection_cells(ws, rows=rows, cols=cols, values=values)
        return f"selection:{sel.sheet}:{len(rows)}x{len(cols)}"
    start_col = 1
    raw_start = str(_op_get(op, "start_cell") or "")
    if raw_start:
        _sheet, _row0, start_col = _bind_write_origin(wb, raw_start, sel.sheet)
    ws = _worksheet(wb, sel.sheet)
    width = None
    for r_idx, excel_row in enumerate(rows):
        cells = values[r_idx] if isinstance(values[r_idx], list) else [values[r_idx]]
        if width is None:
            width = len(cells)
        elif len(cells) != width:
            raise MutationAborted(_invalid("values 必须是矩形", code="COORD_CONTRACT"))
    _write_selection_cells(
        ws,
        rows=rows,
        cols=list(range(start_col, start_col + int(width or 0))),
        values=values,
    )
    return f"selection:{sel.sheet}:{len(rows)}x{width}"


def _refuse_unmaintained_structure(
    ws: Any,
    action: str,
    target_sheet: str | None = None,
    *,
    axis: str | None = None,
    at: int | None = None,
    count: int = 1,
) -> None:
    from excelmanus.workbook.structure import assert_structure_supported

    try:
        assert_structure_supported(
            ws.parent, action, target_sheet=target_sheet,
            axis=axis, at=at, count=count,
        )
    except ValueError as exc:
        raise MutationAborted(_invalid(str(exc))) from exc


def _resolve_write_values(raw: Any) -> Any:
    """write.values：矩形数组或矩阵 spill 句柄。

    spill 是读侧大结果的外置句柄；写侧直接消费可免去模型把大矩阵
    逐字重序列化进参数——这是 analyze/spill → 落盘的免重排通道。
    """
    if not isinstance(raw, str):
        return raw
    text = raw.strip()
    from excelmanus.engine_core.spill import SpillNotFound, SpillStore, is_spill_reference

    if is_spill_reference(text):
        try:
            resolved = SpillStore(_get_guard().workspace_root).get(text)
        except (SpillNotFound, ValueError) as exc:
            raise MutationAborted(_invalid(f"values 的 spill 句柄无法取回：{exc}"))
        try:
            payload = json.loads(resolved)
        except json.JSONDecodeError:
            raise MutationAborted(
                _invalid("spill 句柄内容不是 JSON；write.values 只接受二维矩阵（或其 JSON / spill 句柄）。")
            )
        if isinstance(payload, dict):
            for key in ("values", "matrix"):
                if isinstance(payload.get(key), list):
                    return payload[key]
            raise MutationAborted(
                _invalid(
                    "spill 内容是对象载荷而非二维矩阵；write.values 需要二维数组。"
                    "records 类结果请消费其 selection，或用 analyze/pivot 产出矩阵。"
                )
            )
        return payload
    raise MutationAborted(_invalid("write.values accepts a matrix or a spill: handle, not encoded JSON"))


def _apply_write(
    wb: Any, op: dict[str, Any], *, create_sheet_if_missing: bool = False
) -> str:
    values = _resolve_write_values(_op_get(op, "values"))
    if not isinstance(values, list) or not values:
        raise MutationAborted(
            _invalid("write 需要非空矩形 values（二维数组或 spill 句柄）")
        )
    if all(isinstance(row, dict) for row in values):
        raise MutationAborted(
            _invalid(
                "write.values 必须是二维矩阵，不能把 filter/search 的 records 对象数组当写入值。"
                "请使用返回的 selection，并提供与选择同行同列的矩阵。",
                code="COORD_CONTRACT",
            )
        )
    if parse_bound_selection(_op_get(op, "selection")) is not None:
        return _write_from_selection(wb, op, values)
    raw_start = str(_op_get(op, "start_cell") or "")
    if not raw_start:
        raise MutationAborted(_invalid("write 需要 start_cell 与非空矩形 values"))
    explicit_sheet = _op_get(op, "sheet")
    sheet, row0, col0 = _bind_write_origin(
        wb, raw_start, explicit_sheet,
        create_sheet_if_missing=create_sheet_if_missing,
    )
    return _write_matrix(wb, sheet=sheet, row0=row0, col0=col0, values=values)


def _parse_at_count(op: dict[str, Any], action: str) -> tuple[int, int]:
    at_raw = _op_get(op, "at")
    if isinstance(at_raw, bool) or (isinstance(at_raw, float) and not at_raw.is_integer()):
        raise MutationAborted(_invalid(f"{action}.at 必须是整数行列号，或列字母"))
    if isinstance(at_raw, str) and at_raw.strip().isalpha():
        try:
            at = column_index_from_string(at_raw.strip().upper())
        except ValueError as exc:
            raise MutationAborted(_invalid(f"{action} 列字母无效：{at_raw}")) from exc
    else:
        try:
            at = int(at_raw or 0)
        except (TypeError, ValueError) as exc:
            raise MutationAborted(_invalid(f"{action} 的 at 必须是列字母或正整数")) from exc
    count_raw = _op_get(op, "count")
    if isinstance(count_raw, bool) or (isinstance(count_raw, float) and not count_raw.is_integer()):
        raise MutationAborted(_invalid(f"{action}.count 必须是正整数"))
    try:
        count = 1 if count_raw is None else int(count_raw)
    except (TypeError, ValueError) as exc:
        raise MutationAborted(_invalid(f"{action} 的 count 必须是正整数")) from exc
    if at < 1 or count < 1:
        raise MutationAborted(_invalid(f"{action} 的 at/count 必须 >= 1"))
    column_axis = action == "delete_columns" or (action == "insert" and str(op.get("axis") or "").lower() in {"column", "columns", "col", "cols"})
    if not column_axis and isinstance(at_raw, str) and at_raw.strip().isalpha():
        raise MutationAborted(_invalid("行操作的 at 需要 Excel 行号，不能用列字母"))
    limit = 16384 if column_axis else 1048576
    if at + count - 1 > limit:
        raise MutationAborted(_invalid("at/count 超出 Excel 行列上限"))
    return at, count


def _structure(wb, action, sheet, **kwargs):
    from excelmanus.workbook.structure_edit import apply_structure_edit
    try:
        return apply_structure_edit(wb, action, sheet, **kwargs)
    except (ValueError, TypeError) as exc:
        raise MutationAborted(_invalid(str(exc), code="STRUCTURE_UNSUPPORTED")) from exc


def _apply_insert(wb: Any, op: dict[str, Any]) -> str:
    _require_explicit_sheet(op, "insert", wb=wb)
    ws = _worksheet(wb, _op_get(op, "sheet"))
    axis = str(_op_get(op, "axis") or "").lower()
    at, count = _parse_at_count(op, "insert")
    if axis not in {"row", "rows", "column", "columns", "col", "cols"}:
        raise MutationAborted(_invalid("insert.axis 必须是 row 或 column"))
    action = "insert_rows" if axis in {"row", "rows"} else "insert_cols"
    _structure(wb, action, ws.title, at=at, count=count)
    return f"{axis}s@{at}+{count}"


def _apply_delete_rows(wb: Any, op: dict[str, Any]) -> str:
    sel = parse_bound_selection(_op_get(op, "selection"))
    if sel is not None:
        ws = _worksheet(wb, sel.sheet)
        rows = sorted({int(r) for r in sel.rows}, reverse=True)
        for row in rows:
            _structure(wb, "delete_rows", ws.title, at=row, count=1)
        return f"delete_rows:selection:{len(rows)}"
    _require_explicit_sheet(op, "delete_rows", wb=wb)
    ws = _worksheet(wb, _op_get(op, "sheet"))
    at, count = _parse_at_count(op, "delete_rows")
    _structure(wb, "delete_rows", ws.title, at=at, count=count)
    return f"delete_rows@{at}x{count}"


def _apply_delete_columns(wb: Any, op: dict[str, Any]) -> str:
    _require_explicit_sheet(op, "delete_columns", wb=wb)
    ws = _worksheet(wb, _op_get(op, "sheet"))
    at, count = _parse_at_count(op, "delete_columns")
    _structure(wb, "delete_cols", ws.title, at=at, count=count)
    return f"delete_columns@{at}x{count}"


def _apply_pivot(
    wb: Any,
    op: dict[str, Any],
    file_path: str,
    warnings: list[str] | None = None,
) -> str:
    from excelmanus.workbook.data import (
        _apply_join,
        _build_condition_mask,
        _materialize_derived_keys,
        _normalize_conditions,
        _normalize_group_keys,
        _normalize_join,
        _pivot_frame,
        dataframe_from_worksheet,
        write_dataframe_to_worksheet,
    )

    # ``pivot_refresh`` and ``pivot(refresh=true)`` replay the recorded source
    # contract instead of treating the current matrix as a new source.
    if str(_op_get(op, "kind") or "") == "pivot_refresh" or bool(_op_get(op, "refresh")):
        target_name = str(_op_get(op, "target_sheet") or "")
        meta_ws = wb["__excelmanus_pivot_meta"] if "__excelmanus_pivot_meta" in wb.sheetnames else None
        stored = None
        if meta_ws is not None:
            for row in meta_ws.iter_rows(min_row=2, values_only=True):
                if str(row[0] or "") == target_name:
                    try:
                        stored = json.loads(str(row[1] or ""))
                    except (TypeError, json.JSONDecodeError):
                        stored = None
                    break
        if not isinstance(stored, dict):
            raise MutationAborted(_invalid(f"找不到目标表 {target_name!r} 的可刷新透视定义", code="NOT_FOUND"))
        replay = dict(stored)
        replay["target_sheet"] = target_name
        replay["overwrite"] = True
        replay.pop("refresh", None)
        op = replay

    _require_explicit_sheet(op, "pivot", wb=wb)
    source = str(_op_get(op, "sheet") or "")
    target = str(_op_get(op, "target_sheet") or source)
    overwrite = bool(_op_get(op, "overwrite"))
    if target == source and not overwrite:
        raise MutationAborted(
            _invalid(
                "pivot 默认不会覆盖源工作表；请提供不同的 target_sheet，"
                "或显式传 overwrite=true 确认替换源表内容。"
            )
        )
    header_row = _op_get(op, "header_row") or 1
    src = _worksheet(wb, source)
    if not isinstance(header_row, int) or isinstance(header_row, bool) or header_row < 1:
        raise MutationAborted(_invalid("pivot.header_row 必须是 Excel 正整数行号；表单数据请先转换为带表头的数据表"))
    df = dataframe_from_worksheet(src, header_row=int(header_row), cached_formulas=True)
    if df.empty and source == target:
        raise MutationAborted(
            _invalid(
                f"pivot 的 sheet '{source}' 是空表；sheet 要指向含数据的源表，"
                "写入目标表由 target_sheet 指定（不存在会自动创建）"
            )
        )
    join_spec, join_err = _normalize_join(_op_get(op, "join"), df)
    if join_err is not None:
        raise MutationAborted(join_err)
    if join_spec is not None:
        right_frame = None
        right_file = join_spec["right_file"] or file_path
        if _get_guard().resolve_and_validate(right_file) == _get_guard().resolve_and_validate(file_path):
            if not join_spec["right_sheet"]:
                raise MutationAborted(_invalid("同簿 pivot.join 必须提供右侧 sheet"))
            right_frame = dataframe_from_worksheet(_worksheet(wb, join_spec["right_sheet"]), header_row=join_spec["right_header"] or 1, cached_formulas=True)
        merged, _unmatched, merge_err = _apply_join(df, join_spec, file_path, right_frame=right_frame)
        if merge_err is not None:
            raise MutationAborted(merge_err)
        assert merged is not None
        df = merged
    keys, derived_specs, keys_err = _normalize_group_keys(_op_get(op, "group_by"), df.columns)
    if keys_err is not None:
        raise MutationAborted(keys_err)
    idx_keys, idx_derived, idx_err = _normalize_group_keys(_op_get(op, "index"), df.columns)
    if idx_err is not None:
        raise MutationAborted(idx_err)
    col_keys, col_derived, col_err = _normalize_group_keys(
        _op_get(op, "columns"), df.columns
    )
    if col_err is not None:
        raise MutationAborted(col_err)
    assert derived_specs is not None and idx_derived is not None and col_derived is not None
    work, mat_err = _materialize_derived_keys(df, [*derived_specs, *idx_derived, *col_derived])
    if mat_err is not None:
        raise MutationAborted(mat_err)
    assert work is not None
    cond_list, cond_err = _normalize_conditions(
        work,
        column=_op_get(op, "column"),
        operator=_op_get(op, "operator"),
        value=_op_get(op, "value"),
        conditions=_op_get(op, "conditions"),
        logic=str(_op_get(op, "logic") or "and"),
        require=False,
    )
    if cond_err is not None:
        raise MutationAborted(cond_err)
    assert cond_list is not None
    mask, mask_err = _build_condition_mask(work, cond_list, str(_op_get(op, "logic") or "and"))
    if mask_err is not None:
        raise MutationAborted(mask_err)
    filtered = work[mask] if mask is not None else work
    index = idx_keys or keys
    table, pivot_err = _pivot_frame(
        filtered,
        index=index,
        columns=col_keys,
        values=_op_get(op, "values"),
        aggfunc=str(_op_get(op, "aggfunc") or "sum"),
        margins=_op_get(op, "margins"),
        margins_name=_op_get(op, "margins_name") or "合计",
    )
    if pivot_err is not None:
        raise MutationAborted(_invalid(pivot_err))
    assert table is not None
    existing_target = resolve_sheet_name(target, wb.sheetnames)
    if existing_target is not None:
        target = existing_target
        dest = _worksheet(wb, target)
        if any(cell.value is not None for row in dest.iter_rows() for cell in row) and not overwrite:
            raise MutationAborted(_invalid("target_sheet 已有内容；选择新的工作表，或传 overwrite=true 替换整张目标表。"))
        if dest.merged_cells.ranges or dest.tables or dest._charts or dest._images or dest.conditional_formatting or dest.data_validations.dataValidation:
            raise MutationAborted(_invalid("pivot 覆盖不能维护目标表的合并、图表、表对象或规则，请写到新表。"))
    if existing_target is None:
        wb.create_sheet(title=target)
    elif warnings is not None:
        warnings.append(f"已按 overwrite=true 替换整张目标表 {target} 的值。")
    dest = _worksheet(wb, target)
    write_dataframe_to_worksheet(dest, table)
    # Persist a compact, refreshable source definition in a hidden sheet.  The
    # output matrix remains ordinary cells for compatibility, while refresh is
    # deterministic and uses the same CAS transaction as the original edit.
    meta_name = "__excelmanus_pivot_meta"
    if meta_name not in wb.sheetnames:
        meta_ws = wb.create_sheet(meta_name)
        meta_ws.sheet_state = "hidden"
        meta_ws.append(["target_sheet", "definition"])
    else:
        meta_ws = wb[meta_name]
    definition = dict(op)
    definition.pop("kind", None)
    definition["file_path"] = file_path
    definition["target_sheet"] = target
    definition["overwrite"] = True
    found_row = None
    for row in range(2, meta_ws.max_row + 1):
        if str(meta_ws.cell(row=row, column=1).value or "") == target:
            found_row = row
            break
    if found_row is None:
        found_row = meta_ws.max_row + 1
    meta_ws.cell(row=found_row, column=1, value=target)
    meta_ws.cell(row=found_row, column=2, value=json.dumps(definition, ensure_ascii=False, default=str, sort_keys=True))
    return f"pivot:{target}:{len(table)}:refreshable"


def _apply_transform(
    wb: Any,
    op: dict[str, Any],
    warnings: list[str] | None = None,
) -> str:
    from excelmanus.workbook.data import apply_transform_frame, dataframe_from_worksheet, write_dataframe_to_worksheet

    _require_explicit_sheet(op, "transform", wb=wb)
    sheet = str(_op_get(op, "sheet") or "")
    header_row = _op_get(op, "header_row") or 1
    action = str(_op_get(op, "action") or "").strip().lower()
    if not isinstance(header_row, int) or isinstance(header_row, bool) or header_row < 1:
        raise MutationAborted(_invalid("transform.header_row 必须是 Excel 正整数行号"))
    if not action:
        # 模型常省略 action 只给特征字段——按字段推断，避免无意义重试。
        if _op_get(op, "key_columns") is not None:
            action = "dedupe"
        elif (
            _op_get(op, "delimiter") is not None
            or _op_get(op, "new_columns") is not None
            or _op_get(op, "into") is not None
        ):
            action = "split"
    ws = _worksheet(wb, sheet)
    value_column = next((i + 1 for i, cell in enumerate(ws[header_row]) if cell.value == _op_get(op, "column")), 0)
    if action in {"dedupe", "split"}:
        _refuse_unmaintained_structure(ws, f"transform.{action}")
        if ws.conditional_formatting or ws.data_validations.dataValidation or any(m.max_row >= int(header_row) for m in ws.merged_cells.ranges):
            raise MutationAborted(_invalid("结构清洗不能自动维护数据区合并、条件格式或验证范围，请先移除规则或输出到新表。"))
    if any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in ws.iter_rows()
        for cell in row
        if action in {"dedupe", "split"} or cell.column == value_column
    ):
        raise MutationAborted(
            _invalid(
                "transform 遇到公式单元格时拒绝整表重写，以免删除/重排后公式引用失真。"
                "请用显式 write/copy 操作处理公式区域，或先将公式冻结为字面值。"
            )
        )
    df = dataframe_from_worksheet(ws, header_row=int(header_row))
    # into 在别处语义是目标锚点（{"start_cell": "A1"}），只有列表/字符串才视作新列名。
    _into = _op_get(op, "into")
    _names = _op_get(op, "new_columns")
    if _names is None and isinstance(_into, (list, str)):
        _names = _into
    out, err = apply_transform_frame(
        df,
        action,
        key_columns=_op_get(op, "key_columns"),
        key_normalizers=_op_get(op, "key_normalizers"),
        keep=str(_op_get(op, "keep") or "first"),
        order_by=_op_get(op, "order_by"),
        column=_op_get(op, "column"),
        delimiter=str(_op_get(op, "delimiter") or ","),
        into=_names,
    )
    if err is not None:
        raise MutationAborted(_invalid(err))
    assert out is not None
    if action in {"normalize_date", "normalize_phone"}:
        target_col = list(df.columns).index(_op_get(op, "column")) + 1
        for index, value in enumerate(out[_op_get(op, "column")], int(header_row) + 1):
            from excelmanus.workbook.data import _py_scalar
            assign_cell_value(ws, index, target_col, _py_scalar(value))
        return f"transform:{action}:{sheet}!{get_column_letter(target_col)}{int(header_row) + 1}:{get_column_letter(target_col)}{int(header_row) + len(out)}"
    if int(header_row) != 1 and warnings is not None:
        warnings.append(
            "transform_preserved_prefix: rows above header_row were preserved; formulas and objects are not reference-rewritten"
        )
    elif warnings is not None:
        warnings.append(
            "transform_rewrites_data_block: styles may be retained for existing cells, but formula/object references are not recalculated"
        )
    source_cols = [list(df.columns).index(c) + 1 if c in df.columns else list(df.columns).index(_op_get(op, "column")) + 1 for c in out.columns]
    source_rows = [int(header_row) + 1 + i for i in out.attrs.get("source_positions", list(range(len(out))))]
    write_dataframe_to_worksheet(ws, out, start_row=int(header_row), source_rows=source_rows, source_columns=source_cols)
    if action == "dedupe":
        keys = _op_get(op, "key_columns") or []
        normalizers = _op_get(op, "key_normalizers") or {}
        keep = str(_op_get(op, "keep") or "first")
        order_by = _op_get(op, "order_by")
        rule = f"keep={keep}" + (f" order_by={order_by}" if order_by else "")
        return (
            f"transform:dedupe rows={len(df)}→{len(out)} removed={len(df) - len(out)} "
            f"keys={keys} key_normalizers={normalizers} {rule} "
            f"output_range={get_column_letter(1)}{int(header_row)}:"
            f"{get_column_letter(max(1, len(out.columns)))}{int(header_row) + len(out)}"
        )
    return f"transform:{action}:{len(out)}"


def _apply_sheet(wb: Any, op: dict[str, Any]) -> str:
    action = str(_op_get(op, "action") or "")
    name = _op_get(op, "sheet")
    new_name = _op_get(op, "new_name")
    if action == "create":
        title = new_name or name
        if not title:
            raise MutationAborted(
                _invalid("sheet.create 需要新表名：new_name，或把新表名写在 sheet / sheet_name")
            )
        if resolve_sheet_name(title, wb.sheetnames) is not None:
            raise MutationAborted(_invalid(f"工作表已存在：{title}"))
        wb.create_sheet(title=str(title))
        return f"create:{title}"
    if action == "rename":
        if not name or not new_name:
            raise MutationAborted(_invalid("sheet.rename 需要 sheet 与 new_name"))
        resolved = resolve_sheet_name(name, wb.sheetnames)
        if resolved is None:
            raise MutationAborted(_sheet_not_found(name, wb.sheetnames))
        if resolve_sheet_name(new_name, wb.sheetnames) is not None:
            raise MutationAborted(_invalid(f"工作表已存在：{new_name}"))
        _structure(wb, "rename_sheet", resolved, new_name=str(new_name))
        return f"rename:{resolved}->{new_name}"
    if action == "delete":
        if not name:
            raise MutationAborted(_invalid("sheet.delete 需要 sheet"))
        resolved = resolve_sheet_name(name, wb.sheetnames)
        if resolved is None:
            raise MutationAborted(_sheet_not_found(name, wb.sheetnames))
        if len(wb.sheetnames) <= 1:
            raise MutationAborted(_invalid("不能删除唯一的工作表"))
        _refuse_unmaintained_structure(wb[resolved], "sheet.delete", target_sheet=resolved)
        del wb[resolved]
        return f"delete:{resolved}"
    if action == "copy":
        if not name or not new_name:
            raise MutationAborted(_invalid("sheet.copy 需要 sheet 与 new_name"))
        resolved = resolve_sheet_name(name, wb.sheetnames)
        if resolved is None:
            raise MutationAborted(_sheet_not_found(name, wb.sheetnames))
        if resolve_sheet_name(new_name, wb.sheetnames) is not None:
            raise MutationAborted(_invalid(f"工作表已存在：{new_name}"))
        source_ws = wb[resolved]
        unsupported = {
            "charts": len(source_ws._charts), "images": len(source_ws._images),
            "tables": len(source_ws.tables), "local_names": len(source_ws.defined_names),
        }
        if any(unsupported.values()):
            raise MutationAborted(_invalid("sheet.copy 不能完整复制这些对象；请保留原表或只复制所需区域。", unsupported_objects=unsupported))
        from copy import deepcopy
        from excelmanus.workbook.structure import _formula_mentions_sheet
        for row in source_ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and _formula_mentions_sheet(str(cell.value), resolved):
                    raise MutationAborted(_invalid("sheet.copy 无法自动重绑显式引用源表名的公式；请使用区域复制并显式指定公式。"))
        copied = wb.copy_worksheet(source_ws)
        copied.freeze_panes = source_ws.freeze_panes
        copied.views = deepcopy(source_ws.views)
        copied.conditional_formatting = deepcopy(source_ws.conditional_formatting)
        copied.data_validations = deepcopy(source_ws.data_validations)
        copied.title = str(new_name)
        return f"copy:{name}->{new_name}"
    raise MutationAborted(_invalid("sheet.action 必须是 create/copy/rename/delete"))


def _apply_copy(wb: Any, op: dict[str, Any]) -> str:
    for names in (
        ("sheet", "sheet_name"), ("source_sheet", "sourceSheet"),
        ("target_sheet", "targetSheet"), ("source_range", "sourceRange"),
        ("target_start", "targetStart"),
    ):
        values = [op[name] for name in names if op.get(name) not in (None, "")]
        if len(values) > 1 and values[0] != values[1]:
            raise MutationAborted(_invalid(f"copy 别名冲突：{' 与 '.join(names)} 值不同", invalid_fields=list(names)))
    src_sheet = _op_get(op, "source_sheet")
    raw_src_range = str(_op_get(op, "source_range") or "")
    dst_sheet = _op_get(op, "target_sheet")
    raw_dst_start = str(_op_get(op, "target_start") or "A1")
    try:
        src_from_range, src_range = parse_sheet_address(raw_src_range)
        dst_from_start, dst_start = parse_sheet_address(raw_dst_start)
        src_sheet = combine_sheet_names(
            str(src_sheet) if src_sheet not in (None, "") else None,
            src_from_range,
        )
        dst_sheet = combine_sheet_names(
            str(dst_sheet) if dst_sheet not in (None, "") else None,
            dst_from_start,
        )
        # sheet remains the operation's destination sheet, as for write/format.
        # A missing side uses the explicitly resolved other side (same-sheet copy).
        dst_sheet = combine_sheet_names(_op_get(op, "sheet"), dst_sheet)
    except ValueError as exc:
        raise MutationAborted(_invalid(str(exc))) from exc
    dst_sheet = dst_sheet or src_sheet
    src_sheet = src_sheet or dst_sheet
    if not src_sheet:
        if len(wb.sheetnames) != 1:
            raise MutationAborted(_invalid(
                "copy 需要 sheet（同表复制），或 source_sheet/target_sheet（跨表复制）；也可在地址里写 表!A1。",
                code="SHEET_REQUIRED", available_sheets=list(wb.sheetnames),
            ))
        src_sheet = dst_sheet = wb.sheetnames[0]
    if not src_range:
        raise MutationAborted(_invalid("copy 需要 source_range，如 C1 或 A1:B10"))
    try:
        dst_start = top_left_cell(dst_start) or "A1"
        start_row, start_col = coordinate_to_tuple(dst_start.upper())
    except Exception as exc:
        _abort_bad_address("target_start", raw_dst_start, exc)
    src = _worksheet(wb, src_sheet)
    dst = _worksheet(wb, dst_sheet)
    used_max_row, used_max_col = worksheet_used_shape(src)
    try:
        bounds = resolve_range_to_bounds(
            src_range,
            used_max_row=used_max_row,
            used_max_col=used_max_col,
        )
    except ValueError as exc:
        raise MutationAborted(_invalid(str(exc), code="RANGE_INVALID")) from exc
    min_col, min_row, max_col, max_row = (
        bounds.min_col, bounds.min_row, bounds.max_col, bounds.max_row,
    )
    import copy as _copy

    snapshot = [
        (row, col, _copy.copy(src.cell(row=row, column=col)))
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    ]
    for row, col, source_cell in snapshot:
        value, style = source_cell.value, source_cell._style
        target_row = start_row + row - min_row
        target_col = start_col + col - min_col
        target = dst.cell(row=target_row, column=target_col)
        if source_cell.data_type == "f":
            try:
                from openpyxl.formula.translate import Translator

                value = Translator(value, origin=src.cell(row=row, column=col).coordinate).translate_formula(
                    target.coordinate
                )
            except Exception as exc:
                raise MutationAborted(_invalid(f"无法平移公式 {source_cell.coordinate} 到 {target.coordinate}: {exc}")) from exc
        target.value = value
        if source_cell.data_type != "f" and isinstance(value, str):
            target.data_type = source_cell.data_type
        target._style = _copy.copy(style)
        target.comment = _copy.copy(source_cell.comment)
        target._hyperlink = None
        if source_cell.hyperlink is not None:
            target.hyperlink = _copy.copy(source_cell.hyperlink)
    applied = f"{src_sheet}!{bounds.resolved}->{dst_sheet}!{dst_start}"
    if bounds.requested != bounds.resolved:
        applied = f"{applied} (from {bounds.requested})"
    return applied


def _merged_anchor_cell(ws: Any, cell: Any) -> Any | None:
    from openpyxl.cell.cell import MergedCell

    if not isinstance(cell, MergedCell):
        return cell
    coord = getattr(cell, "coordinate", "")
    for merged in ws.merged_cells.ranges:
        if coord in merged:
            return ws.cell(row=merged.min_row, column=merged.min_col)
    return None


def _canonical_format_kind(kind: str) -> str:
    aliases = {
        "conditionalformat": "conditional_format",
        "cf": "conditional_format",
        "datavalidation": "data_validation",
        "dv": "data_validation",
    }
    return aliases.get(kind, kind)


def _truthy_op_field(op: dict[str, Any], *keys: str) -> bool:
    value = _op_get(op, *keys)
    if value in (None, "", False):
        return False
    if value in ([], {}):
        return False
    return True


def _format_range_optional(kind: str, op: dict[str, Any]) -> bool:
    if kind == "size":
        return (
            _truthy_op_field(op, "auto_fit")
            or _truthy_op_field(op, "column_widths")
            or _truthy_op_field(op, "row_heights")
        )
    if kind in {"freeze", "print_layout"}:
        return True
    return False


def _format_example_op(kind: str, sheet: str) -> dict[str, Any]:
    canon = _canonical_format_kind(kind)
    if canon == "print_layout":
        return {"kind": "print_layout", "sheet": sheet, "print_layout": {"fit_to_width": 1, "fit_to_height": 1}}
    if canon == "size":
        return {"kind": "size", "sheet": sheet, "column_widths": {"A": 18}}
    if canon == "freeze":
        return {"kind": "freeze", "sheet": sheet, "freeze_panes": "A2"}
    if canon in {"merge", "unmerge"}:
        return {"kind": canon, "sheet": sheet, "range": "A1:B1"}
    if canon == "conditional_format":
        return {
            "kind": "conditional_format",
            "sheet": sheet,
            "range": "A1:K1",
            "rule": {"type": "cell_value", "operator": "greaterThan", "value": 0},
        }
    if canon == "data_validation":
        return {
            "kind": "data_validation",
            "sheet": sheet,
            "range": "A2:A10",
            "rule": {"type": "list", "values": ["是", "否"]},
        }
    return {
        "kind": "format",
        "sheet": sheet,
        "range": "A1:K1",
        "font": {"bold": True},
    }


def _format_contract_abort(
    *,
    kind: str,
    kind_omitted: bool,
    missing: list[str],
    wb: Any,
) -> None:
    names = _workbook_sheetnames(wb)
    sheet_for_example = names[0] if names else "Sheet1"
    example = _format_example_op(kind, sheet_for_example)
    labels = {
        "sheet": "sheet（或在地址里写 表!A1）",
        "range": "range",
        "column_widths/row_heights 或 auto_fit=true": "column_widths/row_heights 或 auto_fit=true（同时提供 axis 和 range）",
    }
    parts: list[str] = []
    if kind_omitted:
        parts.append("kind 未提供（缺省为 format）")
    if missing:
        listed = "、".join(labels.get(item, item) for item in missing)
        parts.append(f"{kind} 必须提供 {listed}")
    parts.append(f"最小合法示例：{json.dumps(example, ensure_ascii=False)}")
    only_sheet = missing == ["sheet"]
    code = "SHEET_REQUIRED" if only_sheet and len(names) > 1 else "INVALID_ARGS"
    extra: dict[str, Any] = {
        "example": example,
        "missing_fields": list(missing),
    }
    if names:
        extra["available_sheets"] = names
    raise MutationAborted(_invalid("。".join(parts), code=code, **extra))


def _format_bound_sheet(op: dict[str, Any], raw_range: str) -> str | None:
    """size/freeze 在已绑定 sheet 时不要把表名当 range 再解析。"""
    explicit = _op_get(op, "sheet")
    if explicit not in (None, "") and (
        not str(raw_range or "").strip() or not _address_has_sheet(raw_range)
    ):
        return str(explicit)
    sheet, _local = _split_op_address(op, raw_range)
    return sheet


def _apply_format(wb: Any, op: dict[str, Any]) -> tuple[str, list[str]]:
    kind_raw = _op_get(op, "kind")
    kind_omitted = kind_raw in (None, "")
    kind = _canonical_format_kind(str(kind_raw or "format"))
    fields = _FORMAT_KIND_FIELDS.get(kind)
    if fields is not None:
        _reject_operation_fields(op, f"format.kind={kind}", fields)
    _maybe_bind_unique_sheet(wb, op, "range", "cell_range")
    raw_range = str(_op_get(op, "range") or "")
    missing: list[str] = []
    if not _op_has_sheet(op, "range", "cell_range"):
        missing.append("sheet")
    if kind == "size":
        if not _format_range_optional(kind, op) and not raw_range.strip():
            missing.append("column_widths/row_heights 或 auto_fit=true")
    elif not _format_range_optional(kind, op) and not raw_range.strip():
        missing.append("range")
    if missing:
        _format_contract_abort(
            kind=kind, kind_omitted=kind_omitted, missing=missing, wb=wb,
        )
    if kind not in {"format", "merge", "unmerge", "size", "freeze", "print_layout", "conditional_format", "conditionalformat", "cf", "data_validation", "datavalidation", "dv"}:
        raise MutationAborted(
            _invalid(
                f"不支持的 format.kind={kind}。可用：format / merge / unmerge / size / freeze / print_layout / conditional_format / data_validation"
            )
        )
    if kind == "print_layout":
        from excelmanus.workbook.layout import PrintLayout, apply_print_layout

        layout = PrintLayout.model_validate(_op_get(op, "print_layout"))
        if not layout.model_dump(exclude_none=True):
            raise ValueError("print_layout 至少提供一个打印设置")
        ws = _worksheet(wb, _format_bound_sheet(op, raw_range))
        apply_print_layout(ws, layout)
        return "print_layout", []
    if kind == "freeze":
        sheet = _format_bound_sheet(op, raw_range)
        ws = _worksheet(wb, sheet)
        try:
            cell = _freeze_cell_from_op(op)
            applied = apply_freeze_panes(ws, cell or None)
        except ValueError as exc:
            raise MutationAborted(_invalid(str(exc))) from exc
        return f"freeze:{applied or 'off'}", []
    if kind == "size":
        from excelmanus.workbook.data import _maybe_json

        sheet = _format_bound_sheet(op, raw_range)
        ws = _worksheet(wb, sheet)
        from excelmanus.workbook.styles import _size_entries

        columns = _maybe_json(op.get("column_widths"))
        rows = _maybe_json(op.get("row_heights"))
        columns = {} if columns is None else columns
        rows = {} if rows is None else rows
        auto_fit = _op_get(op, "auto_fit", default=False)
        if not isinstance(auto_fit, bool):
            raise ValueError("size.auto_fit 必须是布尔值")
        axis = str(_op_get(op, "axis") or "").lower()
        axis = {"columns": "column", "col": "column", "cols": "column", "rows": "row"}.get(axis, axis)
        if axis not in {"", "column", "row", "both"}:
            raise ValueError("size.axis 必须是 row 或 column")
        letters = None
        if isinstance(columns, list):
            if columns and all(isinstance(c, str) and c.strip().isascii() and c.strip().isalpha() for c in columns):
                letters = set(_size_entries({c.strip(): 1 for c in columns}, axis="column"))
                columns = {}
            else:
                columns = {str(i + 1): v for i, v in enumerate(columns)}
        if isinstance(rows, list):
            rows = {str(i + 1): v for i, v in enumerate(rows)}
        if not isinstance(columns, dict) or not isinstance(rows, dict):
            raise ValueError("size.column_widths/row_heights 需要尺寸字典；冻结用 kind=freeze 与 freeze_panes")
        columns = _size_entries(columns, axis="column")
        rows = _size_entries(rows, axis="row")
        if axis == "row" and (columns or letters) or axis == "column" and rows:
            raise ValueError("size.axis 与提供的行列尺寸冲突；同时改行列时省略 axis")
        range_cols = range_rows = None
        if raw_range.strip() and raw_range.strip() != ws.title:
            _, rects = _format_rects(op, raw_range, allow_union=True)
            range_cols, range_rows = set(), set()
            for rect in rects:
                bounds = resolve_range_to_bounds(rect, used_max_row=ws.max_row, used_max_col=ws.max_column)
                range_cols.update(get_column_letter(c) for c in range(bounds.min_col, bounds.max_col + 1))
                range_rows.update(range(bounds.min_row, bounds.max_row + 1))
            if set(columns) - range_cols or {int(r) for r in rows} - range_rows or (letters is not None and letters - range_cols):
                raise ValueError("尺寸目标超出 range；请统一范围与 column_widths/row_heights")
        if not auto_fit and not columns and not rows and letters is None:
            raise ValueError('kind=size 需要 column_widths={"A":18} 或 row_heights={"1":24}；auto_fit=true 需同时提供 axis 和 range')
        if auto_fit or letters is not None:
            if axis != "row":
                apply_column_sizes(ws, auto_fit=True, letters=letters if letters is not None else range_cols)
            if auto_fit and axis != "column" and letters is None:
                apply_row_sizes(ws, auto_fit=True, row_numbers=range_rows)
        if columns:
            apply_column_sizes(ws, columns)
        if rows:
            apply_row_sizes(ws, rows)
        return "size:auto_fit" if auto_fit or letters is not None else "size", []
    if not raw_range.strip():
        raise MutationAborted(_invalid(f"{kind} 需要 range"))
    sheet, rects = _format_rects(
        op, raw_range, allow_union=(kind in {"format", "conditional_format", "conditionalformat", "cf", "data_validation", "datavalidation", "dv"})
    )
    ws = _worksheet(wb, sheet)
    rects = [_clip_format_local_range(ws, rect) for rect in rects]
    cell_range = rects[0]
    if kind in {"data_validation", "datavalidation", "dv"}:
        from openpyxl.worksheet.cell_range import CellRange

        from excelmanus.workbook.data import _maybe_json
        from excelmanus.workbook.styles import build_data_validation

        target_ranges = [CellRange(rect) for rect in rects]

        def _intersects(a: CellRange, b: CellRange) -> bool:
            return not (
                a.max_col < b.min_col or a.min_col > b.max_col
                or a.max_row < b.min_row or a.min_row > b.max_row
            )

        if _op_get(op, "remove"):
            dv_list = getattr(ws.data_validations, "dataValidation", []) or []
            kept: list[Any] = []
            removed = 0
            for existing in dv_list:
                sqref = getattr(existing, "sqref", None)
                hit = False
                if sqref is not None:
                    for rng in getattr(sqref, "ranges", []):
                        if any(_intersects(rng, target) for target in target_ranges):
                            hit = True
                            break
                if hit:
                    removed += 1
                else:
                    kept.append(existing)
            ws.data_validations.dataValidation = kept
            return f"data_validation:removed:{removed}", []

        rule_spec = _maybe_json(
            _op_get(op, "rule")
        )
        if not isinstance(rule_spec, dict):
            rule_spec = {k: v for k, v in op.items() if k not in {"kind", "op", "range", "cell_range", "sheet"}}
        try:
            dv = build_data_validation(rule_spec)
        except (ValueError, TypeError) as exc:
            raise MutationAborted(_invalid(str(exc))) from exc
        sqref = " ".join(rects)
        try:
            ws.add_data_validation(dv)
            dv.add(sqref)
        except Exception as exc:
            _abort_bad_address("range", raw_range, exc)
        return f"data_validation:{sqref}", []
    if kind in {"conditional_format", "conditionalformat", "cf"}:
        from openpyxl.worksheet.cell_range import CellRange

        from excelmanus.workbook.data import _maybe_json
        from excelmanus.workbook.styles import build_conditional_format_rule

        if _op_get(op, "remove"):
            targets = [CellRange(rect) for rect in rects]
            cf_list = ws.conditional_formatting
            removed = 0
            for cf in list(cf_list):
                hit = any(
                    not (
                        r.max_col < t.min_col or r.min_col > t.max_col
                        or r.max_row < t.min_row or r.min_row > t.max_row
                    )
                    for r in cf.sqref.ranges
                    for t in targets
                )
                if hit:
                    removed += len(cf.rules)
                    del cf_list[cf.sqref]
            return f"conditional_format:removed:{removed}", []

        rule_spec = _maybe_json(_op_get(op, "rule"))
        if not isinstance(rule_spec, dict):
            rule_spec = op
        try:
            anchor = cell_range.split(":")[0]
            rule = build_conditional_format_rule(rule_spec, anchor=anchor)
        except (ValueError, TypeError) as exc:
            raise MutationAborted(_invalid(str(exc))) from exc
        sqref = " ".join(rects)
        try:
            ws.conditional_formatting.add(sqref, rule)
        except Exception as exc:
            _abort_bad_address("range", raw_range, exc)
        return f"conditional_format:{sqref}", []
    if kind == "merge":
        occupied = [cell.coordinate for row in ws[cell_range] for cell in (row if isinstance(row, tuple) else (row,))
                    if cell.coordinate != cell_range.split(":")[0] and cell.value is not None]
        if occupied and not _op_get(op, "allow_data_loss"):
            raise MutationAborted(_invalid(
                "合并会删除非锚点单元格的值。先合并内容，或用 allow_data_loss=true 明确舍弃这些值。",
                affected_cells=occupied[:20], affected_count=len(occupied),
            ))
        try:
            ws.merge_cells(cell_range)
        except Exception as exc:
            _abort_bad_address("range", raw_range, exc)
        return f"merge:{cell_range}", []
    if kind == "unmerge":
        try:
            ws.unmerge_cells(cell_range)
        except Exception as exc:
            _abort_bad_address("range", raw_range, exc)
        return f"unmerge:{cell_range}", []
    font_cfg = _op_get(op, "font")
    fill_cfg = _op_get(op, "fill")
    border_cfg = _op_get(op, "border")
    align_cfg = _op_get(op, "alignment")
    try:
        fill = _build_fill(fill_cfg) if fill_cfg else None
        border = _build_border(border_cfg) if border_cfg else None
    except ValueError as exc:
        raise MutationAborted(_invalid(str(exc))) from exc
    number_format = _op_get(op, "number_format")
    skipped: list[str] = []
    for rect in rects:
        try:
            skipped.extend(
                _paint_format_cells(
                    ws,
                    rect,
                    raw_range,
                    font_cfg=font_cfg,
                    fill=fill,
                    border=border,
                    align_cfg=align_cfg,
                    number_format=number_format,
                )
            )
        except (ValueError, TypeError) as exc:
            raise MutationAborted(_invalid(str(exc))) from exc
    return f"format:{','.join(rects)}", skipped


def _csv_cell_value(value: Any) -> Any:
    """CSV 单元格按需转数值；"007" 类前导零编号保持字符串。"""
    if not isinstance(value, str):
        return value
    s = value.strip()
    if not s:
        return value
    if s[0] == "0" and len(s) > 1 and s.isdigit():
        return value
    try:
        f = float(s)
    except ValueError:
        return value
    if f.is_integer() and "." not in s and "e" not in s.lower():
        return int(f)
    return f


def _expand_source_csv_sheets(spec: Any, guard: Any, dependencies: list | None = None) -> ToolResult | None:
    """把 sheets[].source_csv 展开为 A1 锚点 value_block 与 dimensions（就地修改）。"""
    from pathlib import Path

    import pandas as pd

    from excelmanus.workbook.spec import SpecValidationError, ValueBlock, collect_layout_errors, _parse_a1
    from excelmanus.workbook.data import _detect_csv_encoding
    from excelmanus.workbook.snapshot import csv_separator_for

    for i, sheet in enumerate(spec.sheets):
        src = sheet.source_csv
        if src is None:
            continue
        if not isinstance(src, dict):
            return _invalid(f"sheets.{i}.source_csv 必须是对象", code="SPEC_INVALID")
        csv_path = src.get("file_path")
        if not isinstance(csv_path, str) or not csv_path:
            return _invalid(f"sheets.{i}.source_csv 需要 file_path", code="SPEC_INVALID")
        try:
            live = guard.resolve_and_validate(csv_path)
        except SecurityViolationError:
            return _invalid(OUTSIDE_WORKSPACE_MESSAGE, code="PATH_INVALID")
        live = Path(live)
        if not live.exists() or not live.is_file():
            return _invalid(f"CSV 不存在: {csv_path}", code="NOT_FOUND")
        if live.suffix.lower() not in {".csv", ".tsv", ".txt"}:
            return _invalid(
                f"source_csv 只接受 CSV/TSV/TXT 文本文件，{csv_path} 是 {live.suffix or '无扩展名'}。"
                "xlsx 数据源请在 spec 里用 value_blocks 写数据，或用 analyze/inspect 先读。",
                code="INVALID_ARGS",
            )
        from excelmanus.workbook.snapshot import open_snapshot
        from io import BytesIO
        source_snapshot = open_snapshot(csv_path)
        if dependencies is not None:
            dependencies.append({"path": source_snapshot.file.relative, "version": source_snapshot.content_version})
        encoding = src.get("encoding") or _detect_csv_encoding(source_snapshot.backing_path)
        start = src.get("start") or "A1"
        skip_rows = int(src.get("skip_rows") or 0)
        try:
            df = pd.read_csv(
                BytesIO(source_snapshot.read_bytes()), sep=csv_separator_for(live),
                encoding=encoding, header=None, skiprows=skip_rows, dtype=object,
            )
        except Exception as exc:
            return _invalid(f"CSV 读取失败: {csv_path}: {exc}", code="CSV_READ_FAILED")
        df = df.where(pd.notna(df), None)
        values = [[_csv_cell_value(v) for v in row] for row in df.values.tolist()]
        if not values:
            return _invalid(f"CSV 为空: {csv_path}", code="CSV_EMPTY")
        try:
            s_row, s_col = _parse_a1(start)
        except Exception:
            return _invalid(f"source_csv.start 锚点无效: {start!r}", code="SPEC_INVALID")
        sheet.value_blocks.append(ValueBlock(start=start, values=values))
        if not sheet.dimensions:
            sheet.dimensions = {
                "rows": s_row - 1 + len(values),
                "cols": s_col - 1 + max(len(r) for r in values),
            }

    layout_errors = collect_layout_errors(spec)
    if layout_errors:
        exc = SpecValidationError(layout_errors)
        return from_payload(exc.to_payload())
    return None


def _apply_edit_operations(
    wb: Any,
    operations: list[Any],
    *,
    file_path: str,
    create_workbook: bool,
) -> tuple[list[str], list[str]]:
    """Apply one edit operation list to an in-memory workbook.

    Keeping this mutation loop independent from the commit wrapper lets the
    single-file and cross-file entry points share exactly the same operation
    semantics while the latter publishes all workbooks in one transaction.
    """
    applied: list[str] = []
    mutation_warnings: list[str] = []
    from excelmanus.workbook.formula_values import FormulaValueError
    for index, raw in enumerate(operations):
        if not isinstance(raw, dict):
            raise MutationAborted(_invalid(f"operations[{index}] 必须是对象"))
        kind = str(_op_get(raw, "kind") or "")
        try:
            _reject_edit_fields(raw, kind)
            if kind == "write":
                applied.append(_apply_write(wb, raw, create_sheet_if_missing=create_workbook))
            elif kind == "insert":
                applied.append(_apply_insert(wb, raw))
            elif kind == "sheet":
                applied.append(_apply_sheet(wb, raw))
            elif kind == "copy":
                applied.append(_apply_copy(wb, raw))
            elif kind == "delete_rows":
                applied.append(_apply_delete_rows(wb, raw))
            elif kind == "delete_columns":
                applied.append(_apply_delete_columns(wb, raw))
            elif kind in {"pivot", "pivot_refresh"}:
                applied.append(_apply_pivot(wb, raw, file_path, mutation_warnings))
            elif kind == "transform":
                applied.append(_apply_transform(wb, raw, mutation_warnings))
            elif kind in _RANGE_KINDS:
                from excelmanus.workbook.operations import apply_range_operation
                applied.append(apply_range_operation(wb, raw))
            else:
                raise MutationAborted(
                    _invalid(
                        f"不支持的 edit.kind={kind}。"
                        "值/表结构用 write|insert|sheet|copy|delete_rows|delete_columns|pivot|transform；"
                        "外观用 apply_spreadsheet_changes"
                    )
                )
        except (ValueError, TypeError, KeyError) as exc:
            if isinstance(exc, FormulaValueError):
                _abort_operation(MutationAborted(_invalid(str(exc), code=exc.code, cells=exc.cells)), index, kind)
            _abort_operation(MutationAborted(_invalid(str(exc))), index, kind)
        except MutationAborted as exc:
            _abort_operation(exc, index, kind)
    return applied, mutation_warnings


def manage_spreadsheet_versions(
    file_path: str,
    action: str,
    revision_id: str | None = None,
    expected_version: str | None = None,
    label: str | None = None,
    limit: int = 50,
) -> ToolResult:
    """列出当前版本与检查点，创建检查点，或按检查点恢复。

    后台只打 RevisionStore（``.excelmanus/revisions/``）。不再读取
    ``intent_revisions.json`` / ``outputs/.versions/rev_*``。
    """
    from excelmanus.workspace.revisions import RevisionIntegrityError

    guard = _get_guard()
    try:
        dest = guard.resolve_and_validate(file_path)
    except SecurityViolationError:
        return commit_error_result(CommitError("PATH_INVALID", OUTSIDE_WORKSPACE_MESSAGE))
    rel = workspace_relpath(guard, dest).replace("\\", "/")
    current = content_version_of_file(dest) if dest.is_file() else None
    seen = peek_seen_content_version(rel)
    from excelmanus.workspace.file_service import WorkspaceFileService

    svc = WorkspaceFileService(guard.workspace_root)

    action = str(action or "").strip().lower()
    if action not in {"list", "checkpoint", "restore", "delete", "delete_checkpoint"}:
        return _invalid("action 必须是 list / checkpoint / restore / delete")

    if action == "list":
        records = [rec.to_public_dict() for rec in svc.list_history(rel)]
        cap = max(1, int(limit))
        return _success(
            {
                "file_path": rel,
                "content_version": current,
                "seen_version": seen,
                "revisions": records[-cap:],
            }
        )

    if action == "checkpoint":
        if not dest.is_file() or not current:
            return _invalid("文件不存在，无法建立检查点", code="PATH_INVALID")
        checkpoint_version = current
        try:
            rec = svc.checkpoint(
                rel,
                expected_version=expected_version or checkpoint_version or "",
                label=label,
            )
        except CommitError as exc:
            return commit_error_result(exc)
        entry = rec.to_public_dict()
        return _success({"file_path": rel, "content_version": checkpoint_version, "revision": entry})

    if action in {"delete", "delete_checkpoint"}:
        if not revision_id:
            return _invalid("delete 需要 revision_id")
        try:
            svc.delete_checkpoint(rel, revision_id)
        except CommitError as exc:
            return commit_error_result(exc)
        return _success({"file_path": rel, "deleted_revision": revision_id, "summary": f"已删除检查点 {revision_id}"})

    if action == "restore":
        if not revision_id:
            return _invalid("restore 需要 revision_id")
        live = dest.is_file()
        if live and not (expected_version or "").strip():
            return _invalid("restore 必须提供 expected_version", code="VERSION_CONFLICT")
        try:
            receipt = svc.restore(
                rel,
                revision_id,
                expected_version=expected_version,
                restore_missing=not live,
            )
            svc.raise_if_failed(receipt)
        except RevisionIntegrityError:
            return _invalid("检查点快照损坏", code="NOT_FOUND")
        except CommitError as exc:
            return commit_error_result(exc)
        return _success(
            {
                "file_path": receipt.primary_path() or rel,
                "content_version": receipt.primary_version(),
                "restored_revision": revision_id,
                "lineage_id": receipt.targets[-1].lineage_id if receipt.targets else None,
                "exists_after": receipt.targets[-1].exists_after if receipt.targets else True,
                "summary": (
                    f"已恢复 {receipt.primary_path() or rel} 到修订 {revision_id}"
                ),
            }
        )

    return _invalid("action 必须是 list / checkpoint / restore / delete")


def analyze_spreadsheet(
    request: dict[str, Any] | None = None,
    mode: str | None = None,
    file_path: str = "",
    path: str = "",
    sheet_name: str | None = None,
    sheet: str | None = None,
    header_row: int | None = None,
    column: str | None = None,
    operator: str | None = None,
    value: Any = None,
    conditions: list[dict[str, Any]] | None = None,
    logic: str = "and",
    columns: list[str] | None = None,
    max_rows: int | None = None,
    sort_by: str | None = None,
    ascending: bool = True,
    limit: int | None = None,
    directory: str = ".",
    file_paths: list[str] | None = None,
    paths: list[str] | None = None,
    max_files: int | None = None,
    query: str | None = None,
    include: list[str] | None = None,
    sample_rows: int | None = None,
    group_by: Any = None,
    aggregations: Any = None,
    dup_only: bool | None = None,
    join: Any = None,
    index: Any = None,
    values: Any = None,
    aggfunc: str | None = None,
    margins: Any = None,
    margins_name: str | None = None,
    totals: Any = None,
    totals_name: str | None = None,
    grand_total: Any = None,
    expected_version: str | None = None,
    content_version: str | None = None,
) -> ToolResult:
    """只读分析：profile / quality / filter / aggregate / distinct / pivot / relationships / files。"""
    args = _merge_request(
        request,
        mode=mode,
        file_path=file_path or path,
        sheet_name=sheet_name or sheet,
        header_row=header_row,
        column=column,
        operator=operator,
        value=value,
        conditions=conditions,
        logic=logic,
        columns=columns,
        max_rows=max_rows,
        sort_by=sort_by,
        ascending=ascending,
        limit=limit,
        directory=directory,
        file_paths=file_paths or paths,
        max_files=max_files,
        query=query,
        include=include,
        sample_rows=sample_rows,
        group_by=group_by,
        aggregations=aggregations,
        dup_only=dup_only,
        join=join,
        index=index,
        values=values,
        aggfunc=aggfunc,
        margins=margins,
        margins_name=margins_name,
        totals=totals,
        totals_name=totals_name,
        grand_total=grand_total,
        expected_version=expected_version or content_version,
    )
    chosen = str(args.get("mode") or "profile")
    target = str(args.get("file_path") or "")
    if chosen in _INSPECT_ONLY_ANALYZE_MODES:
        return _invalid(
            f"mode={chosen} 属于 observe_spreadsheet，请改用 observe_spreadsheet(mode=\"{chosen}\")。"
            "analyze_spreadsheet 可用：profile / quality / filter / aggregate / distinct / pivot / relationships / files"
        )
    mode_err = _reject_mode_fields(
        args, chosen, _ANALYZE_MODE_FIELDS.get(chosen, frozenset()), _ANALYZE_DEFAULTS,
        required_for_mode=_ANALYZE_REQUIRED_FOR_MODE.get(chosen) or [],
        missing_required=_analyze_missing_required(chosen, args),
    )
    if mode_err is not None:
        return mode_err

    if chosen in {"profile", "quality"}:
        from excelmanus.workbook.data import scan_excel_snapshot

        if not target:
            return _invalid(f"{chosen} 需要 file_path / path")
        return scan_excel_snapshot(
            file_path=target,
            max_sample_rows=int(
                args.get("max_rows") or args.get("maxRows")
                or args.get("sample_rows") or args.get("sampleRows")
                or args.get("limit") or 500
            ),
            include_relationships=True,
            sheet_name=args.get("sheet_name"),
            expected_version=args.get("expected_version"),
            header_row=args.get("header_row"),
        )

    if chosen == "filter":
        from excelmanus.workbook.data import filter_data

        if not target:
            return _invalid("filter 需要 file_path / path")
        return filter_data(
            file_path=target,
            column=args.get("column"),
            operator=args.get("operator"),
            value=args.get("value"),
            sheet_name=args.get("sheet_name"),
            header_row=args.get("header_row") if args.get("header_row") is not None else args.get("headerRow"),
            columns=args.get("columns"),
            conditions=args.get("conditions"),
            logic=str(args.get("logic") or "and"),
            max_rows=args.get("max_rows") or args.get("maxRows"),
            sort_by=args.get("sort_by") or args.get("sortBy"),
            ascending=bool(args.get("ascending", True)),
            limit=args.get("max_rows") or args.get("maxRows") or args.get("limit"),
            expected_version=args.get("expected_version"),
        )

    if chosen == "aggregate":
        from excelmanus.workbook.data import aggregate_data

        if not target:
            return _invalid("aggregate 需要 file_path / path")
        return aggregate_data(
            file_path=target,
            group_by=args.get("group_by") or args.get("groupBy"),
            aggregations=args.get("aggregations") or args.get("aggs"),
            sheet_name=args.get("sheet_name"),
            header_row=args.get("header_row") if args.get("header_row") is not None else args.get("headerRow"),
            column=args.get("column"),
            operator=args.get("operator"),
            value=args.get("value"),
            conditions=args.get("conditions"),
            logic=str(args.get("logic") or "and"),
            sort_by=args.get("sort_by") or args.get("sortBy"),
            ascending=bool(args.get("ascending", False)),
            limit=args.get("max_rows") or args.get("maxRows") or args.get("limit"),
            join=args.get("join") or args.get("lookup"),
            expected_version=args.get("expected_version"),
        )

    if chosen == "pivot":
        from excelmanus.workbook.data import pivot_data

        if not target:
            return _invalid("pivot 需要 file_path / path")
        return pivot_data(
            file_path=target,
            index=args.get("index"),
            columns=args.get("columns"),
            values=args.get("values"),
            aggfunc=str(args.get("aggfunc") or "sum"),
            sheet_name=args.get("sheet_name"),
            header_row=args.get("header_row") if args.get("header_row") is not None else args.get("headerRow"),
            column=args.get("column"),
            operator=args.get("operator"),
            value=args.get("value"),
            conditions=args.get("conditions"),
            logic=str(args.get("logic") or "and"),
            join=args.get("join") or args.get("lookup"),
            group_by=args.get("group_by") or args.get("groupBy"),
            limit=args.get("max_rows") or args.get("maxRows") or args.get("limit"),
            margins=args.get("margins") or args.get("totals") or args.get("grand_total"),
            margins_name=args.get("margins_name") or args.get("totals_name") or "合计",
            expected_version=args.get("expected_version"),
        )

    if chosen == "distinct":
        from excelmanus.workbook.data import distinct_data

        if not target:
            return _invalid("distinct 需要 file_path / path")
        return distinct_data(
            file_path=target,
            column=args.get("column"),
            sheet_name=args.get("sheet_name"),
            header_row=args.get("header_row") if args.get("header_row") is not None else args.get("headerRow"),
            conditions=args.get("conditions"),
            logic=str(args.get("logic") or "and"),
            limit=args.get("limit") or args.get("max_rows") or args.get("maxRows"),
            dup_only=bool(args.get("dup_only") or args.get("dupOnly") or False),
            expected_version=args.get("expected_version"),
        )

    if chosen == "relationships":
        from excelmanus.workbook.data import _maybe_json, discover_file_relationships

        relationship_paths = _maybe_json(args.get("file_paths") or args.get("paths"))
        if not relationship_paths and args.get("file_path"):
            relationship_paths = [args.get("file_path")]
        return discover_file_relationships(
            file_paths=relationship_paths,
            directory=str(args.get("directory") or "."),
            max_files=int(args.get("max_files") or args.get("maxFiles") or 5),
            sample_rows=int(args.get("sample_rows") or args.get("sampleRows") or 200),
        )

    if chosen == "files":
        from excelmanus.workbook.data import _maybe_json, inspect_excel_files

        files_directory = str(args.get("directory") or ".")
        files_target = str(args.get("file_path") or args.get("path") or "").strip()
        files_query = args.get("query") or args.get("search")
        if files_target and files_directory == ".":
            from pathlib import Path

            target_path = Path(files_target)
            files_directory = str(target_path.parent if target_path.suffix else target_path)
            if target_path.suffix and not files_query:
                files_query = target_path.name
        return inspect_excel_files(
            directory=files_directory,
            max_files=int(args.get("max_files") or args.get("maxFiles") or 20),
            include=_maybe_json(args.get("include")),
            search=files_query,
            sheet_name=args.get("sheet_name"),
        )

    return _invalid(
        f"不支持的 analyze.mode={chosen}。"
        "可用：profile / quality / filter / aggregate / distinct / pivot / relationships / files"
    )


def compare_spreadsheets(
    request: dict[str, Any] | None = None,
    file_a: str = "",
    file_b: str = "",
    path: str = "",
    other_path: str = "",
    file_path: str = "",
    sheet_a: str = "",
    sheet_b: str = "",
    sheet: str = "",
    other_sheet: str = "",
    alignment: str = "position",
    key_columns: list[str] | None = None,
    max_diffs: int = 500,
    ignore_style: bool = True,
) -> ToolResult:
    """只读对比：alignment=position 按坐标，alignment=key 按关键列。"""
    args = _merge_request(
        request,
        file_a=file_a or path or file_path,
        file_b=file_b or other_path,
        path=path,
        other_path=other_path,
        sheet_a=sheet_a or sheet,
        sheet_b=sheet_b or other_sheet,
        alignment=alignment,
        key_columns=key_columns,
        max_diffs=max_diffs,
        ignore_style=ignore_style,
    )
    left = str(args.get("file_a") or args.get("path") or args.get("file_path") or "")
    right_explicit = str(args.get("file_b") or args.get("other_path") or args.get("otherPath") or "")
    right = right_explicit or left
    if not left:
        return _invalid("compare 需要 file_a / path")
    if not right_explicit:
        left_sheet = str(args.get("sheet_a") or args.get("sheet") or "")
        right_sheet = str(args.get("sheet_b") or args.get("other_sheet") or args.get("otherSheet") or "")
        if not left_sheet or not right_sheet or left_sheet == right_sheet:
            return _invalid("compare 需要 file_b，或在同一文件内提供不同的 sheet_a 与 sheet_b")
    align = str(args.get("alignment") or "position").strip().lower()
    keys = args.get("key_columns") or args.get("keyColumns")
    if not bool(args.get("ignore_style", True)) and current_call() is None:
        return _invalid("样式和对象对比需要绑定工作区上下文；在直接脚本中使用 run_code 或工具调用入口", next_step="run_code")
    if align == "key" and not keys:
        return _invalid("alignment=key 需要 key_columns")
    if align == "position" and keys:
        return _invalid("alignment=position 不能同时提供 key_columns；按键对齐请用 alignment=key")
    from excelmanus.workbook.data import compare_excel

    result = compare_excel(
            file_a=left,
            file_b=right,
            sheet_a=str(args.get("sheet_a") or args.get("sheet") or ""),
            sheet_b=str(args.get("sheet_b") or args.get("other_sheet") or args.get("otherSheet") or ""),
            ignore_style=True,
            alignment=align,
            key_columns=list(keys) if keys else None,
            max_diffs=int(args.get("max_diffs") or args.get("maxDifferences") or 500),
        )
    if result.success and not bool(args.get("ignore_style", True)):
        from excelmanus.workbook.snapshot import open_snapshot
        from excelmanus.workbook.appearance import compare_appearance
        sa = open_snapshot(left, expected_version=result.value.get("content_version_a"))
        sb = open_snapshot(right, expected_version=result.value.get("content_version_b"))
        wa = sa.open_workbook(data_only=False, read_only=False)
        try:
            wb = sb.open_workbook(data_only=False, read_only=False)
            try:
                appearance = compare_appearance(wa, wb, args.get("sheet_a"), args.get("sheet_b"), int(args.get("max_diffs") or 500))
            finally:
                wb.close()
        finally:
            wa.close()
        return _success({**result.value, "appearance":appearance, "appearance_versions":{"left":sa.content_version,"right":sb.content_version}})
    return result


def trace_spreadsheet_formulas(
    request: dict[str, Any] | None = None,
    mode: str | None = None,
    file_path: str = "",
    path: str = "",
    target: str = "",
    direction: str = "both",
    depth: int = 2,
    detail: str = "summary",
    scope: str = "all",
) -> ToolResult:
    """公式引用：map 全景、trace 单元格、impact 影响面。"""
    args = _merge_request(
        request,
        mode=mode,
        file_path=file_path or path,
        target=target,
        direction=direction,
        depth=depth,
        detail=detail,
        scope=scope,
    )
    chosen = str(args.get("mode") or "map")
    target_path = str(args.get("file_path") or "")
    if not target_path:
        return _invalid("需要 file_path / path")
    mode_err = _reject_mode_fields(
        args, chosen, _TRACE_MODE_FIELDS.get(chosen, frozenset()), _TRACE_DEFAULTS,
        required_for_mode=["target"] if chosen in {"trace", "impact"} else [],
    )
    if mode_err is not None:
        return mode_err
    if chosen == "impact" and str(args.get("scope") or "all").strip().lower() not in {"all", "sheet"}:
        return _invalid("impact.scope 仅支持 all 或 sheet")
    from excelmanus.tools.reference_tools import (
        get_impact_analysis,
        get_reference_map,
        trace_references,
    )

    if chosen == "map":
        return get_reference_map(
            file_path=target_path,
            detail=str(args.get("detail") or "summary"),
        )
    cell = str(args.get("target") or "")
    try:
        parsed = parse_sheet_address(cell)
    except InvalidRefError as exc:
        return _invalid(str(exc), code="RANGE_INVALID")
    if chosen == "trace":
        if not cell:
            return _invalid("trace 需要 target，如 Sheet1!B2")
        requested_depth = args.get("depth", 2)
        if isinstance(requested_depth, bool) or not isinstance(requested_depth, int) or not 1 <= requested_depth <= 5:
            return _invalid("trace.depth 必须为 1 到 5 的整数")
        return _with_resolved_sheet(
            trace_references(
                file_path=target_path,
                target=cell,
                direction=str(args.get("direction") or "both"),
                depth=int(args.get("depth") or 2),
            ),
            sheet=parsed.sheet,
        )
    if chosen == "impact":
        if not cell:
            return _invalid("impact 需要 target")
        return _with_resolved_sheet(
            get_impact_analysis(
                file_path=target_path,
                target=cell,
                scope=str(args.get("scope") or "all"),
            ),
            sheet=parsed.sheet,
        )
    return _invalid(f"不支持的 trace.mode={chosen}。可用：map / trace / impact")


_SPLIT_MAX_FILES_DEFAULT = 50


def _split_safe_filename(key: Any, *, max_len: int = 60) -> str:
    """把拆分键值转成安全文件名片段；空/NaN → '空白'。"""
    import re

    import pandas as pd

    if key is None or (not isinstance(key, str) and pd.isna(key)):
        text = ""
    else:
        text = str(key).strip()
    if not text:
        text = "空白"
    text = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", text).strip(" ._") or "空白"
    return text[:max_len]


def split_spreadsheet(
    file_path: str = "",
    path: str = "",
    by_column: str = "",
    column: str = "",
    sheet_name: str | None = None,
    sheet: str | None = None,
    output_dir: str = "outputs",
    filename_template: str = "{key}",
    header_row: int | None = None,
    max_files: int = 0,
    expected_version: str | None = None,
    content_version: str | None = None,
) -> ToolResult:
    """按某列拆分为新 xlsx，保留原始值/公式/单元格样式的可复制部分。

    分组身份与安全文件名分离；所有目标字节先构造完，再批量提交。
    不能复制的宏、合并、图表或表对象会在 warnings 中明确列出。
    """
    file_path = str(file_path or path or "").strip()
    sheet_name = sheet_name or sheet
    by_column = str(by_column or column or "").strip()
    expected_version = expected_version or content_version
    if not file_path:
        return error_result("缺少 file_path（源文件路径）", code="INVALID_ARGS")
    if not by_column:
        return error_result(
            "缺少 by_column（按哪列拆分，如 '省份'）",
            code="INVALID_ARGS",
            fields={
                "accepted_fields": [
                    "file_path", "by_column", "sheet_name", "output_dir",
                    "filename_template", "header_row", "max_files", "expected_version",
                ]
            },
        )

    from excelmanus.workbook.data import _detect_csv_encoding, _load_df_for_tool

    ctx, err = _load_df_for_tool(
        file_path, sheet_name, header_row, column_hints=[by_column],
        expected_version=expected_version,
    )
    if err is not None:
        return err
    assert ctx is not None
    if expected_version and str(ctx["bound_version"]) != str(expected_version):
        return error_result(
            f"{file_path} 版本已变化：期望 {expected_version}，实际 {ctx['bound_version']}",
            code="STALE_SNAPSHOT",
            fields={
                "expected_version": expected_version,
                "content_version": ctx["bound_version"],
                "file_path": ctx["rel_path"],
            },
        )

    df = ctx["df"]
    sheet_name = ctx["sheet_name"]
    if by_column not in df.columns:
        return error_result(
            f"拆分列 '{by_column}' 不存在。可用列：{list(df.columns)}",
            code="INVALID_ARGS",
            fields={"by_column": by_column, "columns": [str(c) for c in df.columns]},
        )
    key_idx = list(df.columns).index(by_column)
    width = len(df.columns)
    snap = ctx["snap"]
    # _read_df returns the internal zero-based header index.
    effective_header = int(ctx["effective_header"]) + 1
    warnings: list[str] = []

    # Each row carries its original Excel coordinate when the source is xlsx.
    # This lets us copy formula/style metadata without using cached values.
    raw_header: list[Any]
    body_rows: list[tuple[int | None, list[Any], list[Any] | None]] = []
    if snap.is_csv():
        import csv
        from excelmanus.workbook.snapshot import csv_separator_for

        encoding = _detect_csv_encoding(snap.backing_path)
        with open(snap.backing_path, newline="", encoding=encoding) as fh:
            all_rows = list(csv.reader(fh, delimiter=csv_separator_for(snap.backing_path)))
        hdr = effective_header - 1
        raw_header = (all_rows[hdr] if hdr >= 0 and hdr < len(all_rows) else list(df.columns))[:width]
        raw_header = raw_header + [None] * max(0, width - len(raw_header))
        for row in all_rows[hdr + 1 :]:
            cells = (list(row[:width]) + [None] * max(0, width - len(row)))[:width]
            body_rows.append((None, cells, None))
    else:
        wb_src = snap.open_workbook(data_only=False, read_only=False)
        try:
            ws_src = wb_src[sheet_name]
            raw_header = [
                ws_src.cell(row=effective_header, column=col)
                for col in range(1, width + 1)
            ] if effective_header else list(df.columns)
            for row_index in range(effective_header + 1, (ws_src.max_row or 0) + 1):
                cells = [
                    ws_src.cell(row=row_index, column=col)
                    for col in range(1, width + 1)
                ]
                values = [cell.value for cell in cells]
                body_rows.append((row_index, values, cells))
            if ws_src.merged_cells.ranges:
                warnings.append("merged_ranges_not_copied")
            if getattr(ws_src, "_charts", None):
                warnings.append("charts_not_copied")
            if getattr(ws_src, "_images", None):
                warnings.append("images_not_copied")
            if getattr(ws_src, "tables", None):
                warnings.append("tables_not_copied")
            if getattr(ws_src, "data_validations", None) and getattr(
                ws_src.data_validations, "dataValidation", None
            ):
                warnings.append("data_validations_not_copied")
            if ws_src.conditional_formatting:
                warnings.append("conditional_formats_not_copied")
            if effective_header > 1:
                warnings.append("prefix_rows_not_copied: 输出从表头开始，不含源表表头前的标题区")
            if str(getattr(snap, "suffix", "")).lower() in {".xlsm", ".xltm"}:
                warnings.append("vba_not_copied_to_xlsx")
            source_ws = ws_src
        finally:
            wb_src.close()

    # Drop physically empty rows, matching the previous split contract.
    body_rows = [
        item for item in body_rows
        if any(value is not None and str(value) != "" for value in item[1])
    ]

    def _group_identity(value: Any) -> tuple[str, str]:
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return ("blank", "")
        text = str(value)
        return ("blank", "") if not text.strip() else (type(value).__name__, text)

    import pandas as pd

    groups: list[dict[str, Any]] = []
    by_identity: dict[tuple[str, str], dict[str, Any]] = {}
    cached_keys: dict[int, Any] = {}
    if not snap.is_csv() and any(cells and cells[key_idx].data_type == "f" for _, _, cells in body_rows):
        cached_wb = snap.open_workbook(data_only=True, read_only=True)
        try:
            for row_index, row in enumerate(cached_wb[sheet_name].iter_rows(min_row=effective_header + 1, min_col=key_idx + 1, max_col=key_idx + 1, values_only=True), effective_header + 1):
                cached_keys[row_index] = row[0]
        finally:
            cached_wb.close()
    blank_rows = 0
    for row_no, values, cells in body_rows:
        key_value = values[key_idx] if key_idx < len(values) else None
        if cells is not None and cells[key_idx].data_type == "f":
            key_value = cached_keys.get(row_no)
            if key_value is None:
                return _invalid("拆分键含未计算的公式；请先在 Excel 重算保存，不能按公式文本分组。")
        identity = _group_identity(key_value)
        if identity[0] == "blank":
            blank_rows += 1
        group = by_identity.get(identity)
        if group is None:
            display = "空白" if identity[0] == "blank" else identity[1]
            group = {"identity": identity, "display": display, "key": None if identity[0] == "blank" else key_value, "rows": []}
            by_identity[identity] = group
            groups.append(group)
        group["rows"].append((row_no, values, cells))

    limit = int(max_files) if max_files else _SPLIT_MAX_FILES_DEFAULT
    if len(groups) > limit:
        return error_result(
            f"按 '{by_column}' 拆分将产生 {len(groups)} 个文件，超过上限 {limit}。请提高 max_files 或改用更粗粒度列。",
            code="INVALID_ARGS",
            fields={"groups": len(groups), "max_files": limit},
        )

    guard = _get_guard()
    rel_source = ctx["rel_path"]
    source_stem = rel_source.replace("\\", "/").rsplit("/", 1)[-1].rsplit(".", 1)[0]
    template = str(filename_template or "{key}")
    if "{key}" not in template and "{stem}" not in template:
        return error_result(
            "filename_template 必须包含 {key} 或 {stem} 占位符",
            code="INVALID_ARGS",
            fields={"filename_template": template},
        )

    used_names: set[str] = set()
    import unicodedata
    resolved: list[tuple[dict[str, Any], str]] = []
    conflicts: list[str] = []
    for group in groups:
        try:
            name = template.format(key=group["display"], stem=source_stem)
        except (KeyError, IndexError, ValueError) as exc:
            return error_result(
                f"filename_template 占位符非法：{exc}。仅支持 {{key}} 与 {{stem}}",
                code="INVALID_ARGS",
                fields={"filename_template": template},
            )
        name = _split_safe_filename(name)
        if not name.lower().endswith(".xlsx"):
            name = f"{name}.xlsx"
        base = name[:-5] if name.lower().endswith(".xlsx") else name
        candidate = name
        suffix = 2
        while unicodedata.normalize("NFC", candidate).casefold() in used_names:
            candidate = f"{base}_{suffix}.xlsx"
            suffix += 1
        used_names.add(unicodedata.normalize("NFC", candidate).casefold())
        out_rel = f"{str(output_dir or 'outputs').strip().rstrip('/')}/{candidate}"
        try:
            dest = guard.resolve_and_validate(out_rel)
        except SecurityViolationError as exc:
            return error_result(f"输出路径越界：{exc}", code="PATH_INVALID", fields={"path": out_rel})
        rel_posix = str(dest.relative_to(guard.workspace_root)).replace("\\", "/")
        if rel_posix == "uploads" or rel_posix.startswith("uploads/"):
            return error_result(
                "output_dir 不能指向 uploads/（只读附件目录）",
                code="PATH_INVALID",
                fields={"path": rel_posix},
            )
        if dest.exists():
            conflicts.append(rel_posix)
        resolved.append((group, rel_posix))
    if conflicts:
        return error_result(
            f"{len(conflicts)} 个目标文件已存在，拆分已取消（不会覆盖）：{conflicts[:5]}",
            code="VERSION_CONFLICT",
            fields={"conflicts": conflicts},
        )

    import copy
    import io
    import re
    from openpyxl import Workbook
    from openpyxl.formula.translate import Translator

    headers = [str(c.value) if hasattr(c, "value") else str(c) for c in raw_header]
    safe_sheet = re.sub(r"[\[\]\\/*?:]+", "_", str(sheet_name or "Sheet1")).strip()[:31] or "Sheet1"
    rendered: list[tuple[dict[str, Any], str, bytes]] = []

    def _set_cell(dst: Any, src: Any, value: Any, *, origin: str | None = None) -> None:
        if src is not None and src.data_type == "f" and origin:
            try:
                value = Translator(value, origin=origin).translate_formula(dst.coordinate)
            except Exception as exc:
                raise ValueError(f"无法平移公式 {origin} 到 {dst.coordinate}: {exc}") from exc
        dst.value = value
        if isinstance(value, str) and (src is None or src.data_type != "f"):
            dst.data_type = "s"
        if src is not None and hasattr(src, "_style"):
            # Do not transplant the source workbook's indexed style array;
            # openpyxl style IDs belong to their workbook.  Copy components so
            # the new workbook builds a valid independent style table.
            dst.font = copy.copy(src.font)
            dst.fill = copy.copy(src.fill)
            dst.border = copy.copy(src.border)
            dst.alignment = copy.copy(src.alignment)
            dst.protection = copy.copy(src.protection)
            dst.number_format = src.number_format
            if src.comment is not None:
                dst.comment = copy.copy(src.comment)
            if src.hyperlink is not None:
                dst._hyperlink = copy.copy(src.hyperlink)

    for group, rel_posix in resolved:
        out_wb = Workbook()
        out_ws = out_wb.active
        out_ws.title = safe_sheet
        if not snap.is_csv():
            out_wb.loaded_theme = source_ws.parent.loaded_theme
            out_wb._colors = copy.copy(source_ws.parent._colors)
            for letter, dim in source_ws.column_dimensions.items():
                out_ws.column_dimensions[letter].width = dim.width
            if source_ws.freeze_panes:
                freeze_row, freeze_col = coordinate_to_tuple(str(source_ws.freeze_panes))
                out_ws.freeze_panes = f"{get_column_letter(freeze_col)}{max(1, freeze_row - max(0, effective_header - 1))}"
        for col, header in enumerate(raw_header, start=1):
            dst = out_ws.cell(row=1, column=col)
            if hasattr(header, "value"):
                _set_cell(dst, header, header.value, origin=header.coordinate)
            else:
                dst.value = header
        for dest_row, (src_row, values, cells) in enumerate(group["rows"], start=2):
            for col, value in enumerate(values, start=1):
                dst = out_ws.cell(row=dest_row, column=col)
                src = cells[col - 1] if cells is not None else None
                _set_cell(
                    dst,
                    src,
                    value,
                    origin=src.coordinate if src is not None else None,
                )
        buf = io.BytesIO()
        out_wb.save(buf)
        out_wb.close()
        rendered.append((group, rel_posix, buf.getvalue()))

    # Use the existing workspace transaction: check all destinations under
    # lock, preserve history, and expose recoverable partial publication.
    from excelmanus.workspace.file_service import ReadDependency, TargetSpec, service_for_guard

    svc = service_for_guard(guard)
    try:
        receipt = (
            svc.apply_batch(
                [TargetSpec(op="create", path=rel, data=data) for _, rel, data in rendered],
                # The source snapshot is a read dependency of every output.
                # If a user edits the source while rendering is in progress,
                # abort the complete fan-out instead of publishing stale files.
                read_dependencies=[ReadDependency(rel_source, snap.content_version)],
                actor="split_spreadsheet",
            )
            if rendered
            else None
        )
    except CommitError as exc:
        return commit_error_result(exc)
    if receipt is not None and receipt.state != "committed":
        return error_result(
            "拆分产物未全部发布；已提交文件见 committed_files，事务由工作区恢复流程接管，不要重放整批。",
            code="PARTIAL_COMMIT",
            remediation="保留 operation_id/tx_id；工作区恢复后核对产物和版本，不要重新拆分到同一路径。",
            fields={
                "partial": True,
                "committed_files": [t.path for t in receipt.targets if t.publish_status == "published"],
                "operation_id": receipt.operation_id,
                "tx_id": receipt.tx_id,
                "recovery_required": receipt.resumeable,
                "receipt": receipt.to_dict(),
            },
        )
    versions = {t.path: t.after_version for t in receipt.targets} if receipt else {}
    files_out = [{"file_path": rel, "key": group["key"], "rows": len(group["rows"]), "content_version": versions[rel]} for group, rel, _ in rendered]
    files_out.sort(key=lambda f: f["file_path"])
    return _success(
        {
            "status": "ok",
            "file_path": rel_source,
            "sheet_name": sheet_name,
            "by_column": by_column,
            "groups": len(groups),
            "total_rows": len(body_rows),
            "blank_key_rows": blank_rows,
            "output_dir": str(output_dir or "outputs").strip().rstrip("/"),
            "files": files_out,
            "source_content_version": snap.content_version,
            "copy_semantics": "行复制：保留公式并相对平移；跨表依赖和原生对象不随单表产物迁移",
            "warnings": sorted(set(warnings)),
        }
    )
