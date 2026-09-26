"""Workbook geometry in native units and CSS pixels, shared by every observer.

Pixel widths are estimates until a renderer measures the specified font. Native
sizes and their origin always survive projection; hidden axes occupy no pixels.
"""

from __future__ import annotations

import math
from typing import Any

from openpyxl.utils import column_index_from_string, get_column_letter


def column_pixels(width: float, digit_width: float = 7.0) -> float:
    return float(
        math.floor(width * (digit_width + 5) if width < 1 else width * digit_width + 5)
    )


def column_native(pixels: float, digit_width: float = 7.0) -> float:
    return (
        pixels / (digit_width + 5)
        if pixels < digit_width + 5
        else (pixels - 5) / digit_width
    )


def defaults(ws: Any) -> dict[str, Any]:
    fmt = ws.sheet_format
    font = ws.parent._fonts[0]
    width = fmt.defaultColWidth
    height = fmt.defaultRowHeight
    return {
        "column_width": {
            "native": width,
            "effective_native": float(width if width is not None else 8.43),
            "pixels": column_pixels(float(width if width is not None else 8.43)),
            "unit": "excel_char",
            "source": "declared" if width is not None else "estimated",
        },
        "row_height": {
            "native": height,
            "effective_native": float(height if height is not None else 15),
            "pixels": float(height if height is not None else 15) * 96 / 72,
            "unit": "pt",
            "source": "declared" if height is not None else "estimated",
        },
        "font": {"name": font.name, "size_pt": font.sz},
        "pixel_basis": {
            "dpi": 96,
            "maximum_digit_width": 7.0,
            "status": "estimated",
            "reason": "font metrics must be measured by the renderer",
        },
    }


def column_intervals(ws: Any) -> list[tuple[int, int, Any]]:
    return [
        (
            dim.min or column_index_from_string(letter),
            dim.max or dim.min or column_index_from_string(letter),
            dim,
        )
        for letter, dim in ws.column_dimensions.items()
    ]


def column_dimension(ws: Any, index: int, intervals: list | None = None) -> Any:
    matches = [
        (end - start, order, dim)
        for order, (start, end, dim) in enumerate(
            intervals if intervals is not None else column_intervals(ws)
        )
        if start <= index <= end
    ]
    return min(matches, key=lambda item: (item[0], -item[1]))[2] if matches else None


def axis_sizes(ws: Any, start: int, end: int, *, axis: str) -> list[dict[str, Any]]:
    base = defaults(ws)
    columns = axis == "column"
    default = base["column_width" if columns else "row_height"]
    intervals = column_intervals(ws) if columns else []
    sizes = []
    for index in range(start, end + 1):
        dim = (
            column_dimension(ws, index, intervals)
            if columns
            else ws.row_dimensions.get(index)
        )
        explicit = getattr(dim, "width" if columns else "height", None)
        native = float(
            explicit if explicit is not None else default["effective_native"]
        )
        hidden = bool(getattr(dim, "hidden", False)) or (
            not columns and bool(ws.sheet_format.zeroHeight) and dim is None
        )
        pixels = column_pixels(native) if columns else native * 96 / 72
        sizes.append(
            {
                "index": index,
                "native": native,
                "unit": default["unit"],
                "pixels": 0.0 if hidden else pixels,
                "hidden": hidden,
                "outline_level": int(getattr(dim, "outlineLevel", 0) or 0),
                "source": "explicit" if explicit is not None else default["source"],
                "pixel_status": "estimated" if columns else "unit_conversion",
            }
        )
    return sizes


def region_geometry(ws: Any, rect: Any) -> dict[str, Any]:
    columns = axis_sizes(ws, rect.min_col, rect.max_col, axis="column")
    rows = axis_sizes(ws, rect.min_row, rect.max_row, axis="row")
    width = sum(item["pixels"] for item in columns)
    height = sum(item["pixels"] for item in rows)
    return {
        "defaults": defaults(ws),
        "columns": columns,
        "rows": rows,
        "width_px": width,
        "height_px": height,
        "aspect_ratio": width / height if height else None,
        "status": "estimated",
        "coordinate_space": "region_relative_css_px",
        "dpi": 96,
    }


def axis_offset(ws: Any, before: int, *, axis: str) -> float:
    """Prefix size without expanding a million default rows."""
    if axis == "column":
        return sum(item["pixels"] for item in axis_sizes(ws, 1, before - 1, axis=axis))
    base = defaults(ws)["row_height"]["effective_native"] * 96 / 72
    default = 0 if ws.sheet_format.zeroHeight else base
    total = max(0, before - 1) * default
    for index, dim in ws.row_dimensions.items():
        if index >= before:
            continue
        size = (
            0
            if dim.hidden
            else (dim.height * 96 / 72 if dim.height is not None else base)
        )
        total += size - default
    return total


def drawing_preview_range(ws: Any, obj: dict) -> str | None:
    """Smallest cell rectangle from a drawing anchor through its pixel extent."""
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    box = obj.get("bounds") or {}
    if not obj.get("target_cell") or any(box.get(k) is None for k in ("x", "y", "width", "height")):
        return None
    row, column = coordinate_to_tuple(obj["target_cell"])

    def end(start, edge, axis, maximum):
        hi = start
        while axis_offset(ws, hi + 1, axis=axis) < edge and hi < maximum:
            hi = min(maximum, hi + max(1, hi - start + 1))
        lo = start
        while lo < hi:
            mid = (lo + hi) // 2
            if axis_offset(ws, mid + 1, axis=axis) >= edge:
                hi = mid
            else:
                lo = mid + 1
        return lo

    last_col = end(column, box["x"] + box["width"], "column", 16384)
    last_row = end(row, box["y"] + box["height"], "row", 1048576)
    return f"{get_column_letter(column)}{row}:{get_column_letter(last_col)}{last_row}"


def scale_region(
    ws: Any, rect: Any, x: float, y: float, *, preserve_outside: bool = False
) -> dict[str, Any]:
    for value in (x, y):
        if (
            isinstance(value, bool)
            or not isinstance(value, (int, float))
            or not math.isfinite(value)
            or value <= 0
        ):
            raise ValueError("geometry.scale x/y must be positive finite numbers")
    impact = axis_impact(ws, rect, columns=x != 1, rows=y != 1)
    if preserve_outside and impact["outside_count"]:
        from excelmanus.workbook_commit import CommitError

        raise CommitError(
            "LAYOUT_SCOPE_CONFLICT",
            "Row/column resizing would affect content outside the selected rectangle; move the table to independent axes or allow the reported scope",
            fields={"impact": impact, "committed": False},
        )
    before = region_geometry(ws, rect)

    def scaled_width(c):
        target = column_native(c["pixels"] * x) if c["pixels"] else c["native"] * x
        # Pixel rounding must never reverse an explicit native axis direction.
        return max(c["native"] * x, target) if x > 1 else min(c["native"] * x, target)

    widths = {
        get_column_letter(c["index"]): scaled_width(c)
        for c in before["columns"]
        if not c["hidden"] and x != 1
    }
    heights = {
        r["index"]: r["native"] * y
        for r in before["rows"]
        if not r["hidden"] and y != 1
    }
    if any(not 0 < value <= 255 for value in widths.values()) or any(
        not 0 < value <= 409.5 for value in heights.values()
    ):
        raise ValueError(
            "Scaled dimensions exceed Excel limits (column 255 characters, row 409.5 pt)"
        )
    # Materialize every member of intersecting grouped column spans before
    # changing an individual member; overlapping <col> records are ambiguous.
    materialize_column_spans(
        ws, [c["index"] for c in before["columns"]] if x != 1 else []
    )
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for row, height in heights.items():
        ws.row_dimensions[row].height = height
    return {
        "before": before,
        "after": region_geometry(ws, rect),
        "affected_axes": {"columns": list(widths), "rows": list(heights)},
        "impact": impact,
        "clipping_risk": y < 1,
        "requested_scale": {"x": x, "y": y},
        "scope_note": "Column widths and row heights affect the entire axes, including other regions.",
        "text_fit": "not_measured; preview to inspect wrapping and clipping",
    }


def axis_impact(ws: Any, rect: Any, *, columns=True, rows=True) -> dict:
    outside, count = [], 0
    for cell in ws._cells.values():
        if cell.value is None and not cell.has_style:
            continue
        inside = (
            rect.min_row <= cell.row <= rect.max_row
            and rect.min_col <= cell.column <= rect.max_col
        )
        affected = (columns and rect.min_col <= cell.column <= rect.max_col) or (
            rows and rect.min_row <= cell.row <= rect.max_row
        )
        if affected and not inside:
            count += 1
            if len(outside) < 50:
                outside.append(cell.coordinate)
    objects = []
    for kind, drawings in (("chart", ws._charts), ("image", ws._images)):
        for index, drawing in enumerate(drawings):
            start = getattr(drawing.anchor, "_from", None)
            if start is not None:
                row, col = start.row + 1, start.col + 1
                affected = (columns and rect.min_col <= col <= rect.max_col) or (
                    rows and rect.min_row <= row <= rect.max_row
                )
                inside = (
                    rect.min_row <= row <= rect.max_row
                    and rect.min_col <= col <= rect.max_col
                )
                if affected and not inside:
                    count += 1
                    if len(objects) < 50:
                        objects.append(
                            {
                                "kind": kind,
                                "index": index,
                                "anchor": f"{get_column_letter(col)}{row}",
                            }
                        )
    return {
        "outside_objects_sample": objects,
        "scope": "entire_axes",
        "columns": [rect.min_col, rect.max_col] if columns else None,
        "rows": [rect.min_row, rect.max_row] if rows else None,
        "outside_count": count,
        "outside_cells_sample": outside,
        "clipping": "renderer_required",
        "classification": "occupied/styled cells; no inferred table boundaries",
    }


def materialize_column_spans(ws, indices):
    if not indices:
        return
    for start, end, dim in list(column_intervals(ws)):
        if end <= start or end < min(indices) or start > max(indices):
            continue
        from copy import copy

        del ws.column_dimensions[get_column_letter(start)]
        for index in range(start, end + 1):
            letter = get_column_letter(index)
            clone = copy(dim)
            clone.index, clone.min, clone.max = letter, index, index
            ws.column_dimensions[letter] = clone


def resize_drawing(drawing, *, width=None, height=None, unit="px"):
    """Persist size changes in existing one-cell/absolute anchors, not just Python attrs."""
    if width is None and height is None:
        return
    extent = getattr(drawing.anchor, "ext", None)
    if not isinstance(drawing.anchor, str) and extent is None:
        from excelmanus.workbook_commit import CommitError

        raise CommitError(
            "UNSUPPORTED_PRESERVATION",
            "Resizing a two-cell drawing requires an explicit target_cell to use a fixed-size anchor",
            fields={"committed": False},
        )
    factor = 360000 if unit == "cm" else 9525
    for name, value, axis in (("width", width, "cx"), ("height", height, "cy")):
        if value is not None:
            if (
                not isinstance(value, (float, int))
                or isinstance(value, bool)
                or not math.isfinite(value)
                or value <= 0
            ):
                raise ValueError("Drawing dimensions must be positive finite numbers")
            setattr(drawing, name, value)
            if extent is not None:
                setattr(extent, axis, round(value * factor))
