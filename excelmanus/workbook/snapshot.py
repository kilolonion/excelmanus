"""版本绑定工作簿快照：文件身份、内容版本与不可变 backing 的唯一拥有者。

所有正式读（工具、Word 源表、图谱、UI、写后核验）必须经 ``open_snapshot`` /
``open_snapshot_at``，之后只读 backing，禁止再把活路径当权威。
"""

from __future__ import annotations

import io
import json
import os
import tempfile
import threading
import time
import hashlib
from weakref import WeakValueDictionary
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

from excelmanus.workbook.read_cache import ReadCache
from excelmanus.workbook.refs import RectRef
from excelmanus.workbook_commit import content_version_of, remember_content_version
from excelmanus.workspace.refs import FileRef, WorkspaceRef

SNAPSHOT_BLOB_THRESHOLD = 8 * 1024 * 1024
_SNAPSHOT_DIRNAME = ".excelmanus"
_SNAPSHOTS = "snapshots"
_live_backings: WeakValueDictionary[int, "SnapshotBacking"] = WeakValueDictionary()
_workbook_views: ReadCache[tuple[Any, Any, threading.Lock]] = ReadCache(
    max_bytes=256 * 1024 * 1024, max_entries=4
)
_materialize_locks: dict[str, threading.Lock] = {}
_materialize_locks_mu = threading.Lock()

CellType = Literal["n", "s", "b", "d", "e", "z"]
CachedState = Literal["yes", "no", "unknown"]
FactSource = Literal["cached", "formula_text", "unknown_cache", "derived"]
CoverageKind = Literal["complete", "window", "sampled", "truncated"]
ResultKind = Literal["matrix", "records", "areas", "selection"]


class SnapshotError(Exception):
    """开簿或绑定失败。"""

    def __init__(self, message: str, *, code: str, fields: dict[str, Any] | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.fields = fields or {}


class SnapshotStale(SnapshotError):
    def __init__(self, message: str, *, fields: dict[str, Any] | None = None) -> None:
        super().__init__(message, code="STALE_SNAPSHOT", fields=fields)


class SheetRequired(SnapshotError):
    def __init__(self, names: list[str]) -> None:
        super().__init__(
            f"工作簿有 {len(names)} 张表，必须显式指定 sheet。可用：{names}",
            code="SHEET_REQUIRED",
            fields={"available_sheets": names},
        )


class RefUnsupported(SnapshotError):
    def __init__(self, message: str, *, fields: dict[str, Any] | None = None) -> None:
        super().__init__(message, code="REF_UNSUPPORTED", fields=fields)


@dataclass(frozen=True)
class SnapshotId:
    workspace_key: str
    relative: str
    content_version: str

    def key(self) -> str:
        return f"{self.workspace_key}|{self.relative}|{self.content_version}"

    def file_prefix(self) -> str:
        return f"{self.workspace_key}|{self.relative}|"


@dataclass(frozen=True)
class SnapshotBacking:
    kind: Literal["bytes", "blob"]
    sha256: str
    suffix: str
    payload: bytes | None
    blob_path: Path | None

    def __post_init__(self) -> None:
        _live_backings[id(self)] = self

    def read_bytes(self) -> bytes:
        if self.kind == "bytes" and self.payload is not None:
            return self.payload
        if self.blob_path is not None:
            return self.blob_path.read_bytes()
        raise SnapshotError("快照 backing 为空", code="TOOL_ERROR")

    def path(self) -> Path:
        if self.blob_path is not None:
            return self.blob_path
        raise SnapshotError("内存 backing 没有路径；先 materialize", code="TOOL_ERROR")


@dataclass(frozen=True)
class WorkbookSnapshot:
    """ReadSnapshot 的实现。"""

    id: SnapshotId
    file: FileRef
    content_version: str
    backing: SnapshotBacking
    suffix: str

    @property
    def backing_path(self) -> Path:
        return self.backing.path()

    def read_bytes(self) -> bytes:
        return self.backing.read_bytes()

    def open_bytes_io(self) -> io.BytesIO:
        return io.BytesIO(self.read_bytes())

    def open_workbook(self, *, data_only: bool, read_only: bool = True) -> Any:
        from openpyxl import load_workbook

        if self.suffix in {".csv", ".tsv", ".txt"}:
            raise SnapshotError("CSV/TSV 不能当工作簿打开", code="INVALID_ARGS")
        return load_workbook(self.backing_path, data_only=data_only, read_only=read_only)

    def is_csv(self) -> bool:
        return self.suffix in {".csv", ".tsv", ".txt"}

    def csv_separator(self) -> str:
        return "\t" if self.suffix == ".tsv" else ","


@dataclass(frozen=True)
class CellFact:
    sheet: str
    row: int
    col: int
    t: CellType
    v: Any
    text: str | None = None
    f: str | None = None
    cached: CachedState = "unknown"
    source: FactSource = "cached"
    e: str | None = None

    def to_json(self) -> dict[str, Any]:
        payload: dict[str, Any] = {
            "sheet": self.sheet,
            "row": self.row,
            "col": self.col,
            "t": self.t,
            "v": self.v,
            "cached": self.cached,
            "source": self.source,
        }
        if self.text is not None:
            payload["text"] = self.text
        if self.f:
            payload["f"] = self.f
        if self.e:
            payload["e"] = self.e
        return payload


@dataclass(frozen=True)
class Coverage:
    kind: CoverageKind
    returned_rows: int
    total_rows: int | None
    offset: int = 0
    truncated_reason: str | None = None

    def to_json(self) -> dict[str, Any]:
        return {
            "kind": self.kind,
            "returned_rows": self.returned_rows,
            "total_rows": self.total_rows,
            "offset": self.offset,
            "truncated_reason": self.truncated_reason,
        }


@dataclass(frozen=True)
class BoundSelection:
    snapshot: SnapshotId
    file: FileRef
    sheet: str
    rows: tuple[int, ...]
    cols: tuple[int, ...] | None
    origin: str
    header_row: int | str | None = None

    def to_json(self) -> dict[str, Any]:
        return {
            "result_kind": "selection",
            "snapshot_id": self.snapshot.key(),
            "content_version": self.snapshot.content_version,
            "file": self.file.relative,
            "sheet": self.sheet,
            "rows": list(self.rows),
            "cols": list(self.cols) if self.cols is not None else None,
            "origin": self.origin,
            "header_row": self.header_row,
        }


def snapshot_cache_key(snapshot: WorkbookSnapshot) -> str:
    return snapshot.id.key()


def csv_separator_for(path: Path | str) -> str:
    return "\t" if Path(path).suffix.lower() == ".tsv" else ","


def require_default_sheet(sheetnames: list[str], requested: str | None) -> str:
    """I8：恰好一张表可省略；多于一张必须显式 sheet。"""
    from excelmanus.tools._helpers import resolve_sheet_name

    names = [str(n) for n in sheetnames if str(n)]
    if requested not in (None, ""):
        resolved = resolve_sheet_name(str(requested), names)
        if resolved is None:
            raise SnapshotError(
                f"工作表 '{requested}' 不存在。该文件包含: {names}",
                code="SHEET_NOT_FOUND",
                fields={"available_sheets": names, "requested_sheet": requested},
            )
        return resolved
    if len(names) == 1:
        return names[0]
    if not names:
        raise SnapshotError("工作簿没有工作表", code="SHEET_NOT_FOUND")
    raise SheetRequired(names)


# ── 列感知消歧（P0-B）─────────────────────────────────────
#
# 安全红线：静默选错表给出“看起来合理但全错的数字”比报错更危险。
# 因此只允许“唯一精确命中才默认”，其余一律硬失败 SheetRequired。
# 隐藏表参与碰撞检测、永不自动选中；表单/空表/读错的表排除出候选。

_DISAMBIG_MAX_SHEETS = 10


def _normalize_column_name(text: str) -> str:
    return str(text).strip().casefold()


def resolve_sheet_by_visibility(
    sheetnames: list[str],
    hidden: dict[str, bool],
    *,
    max_sheets: int = _DISAMBIG_MAX_SHEETS,
) -> tuple[str, dict[str, Any]]:
    """无列证据时，仅自动绑定唯一可见工作表。"""
    names = [str(n) for n in sheetnames if str(n)]
    if len(names) > max_sheets:
        err = SheetRequired(names)
        err.fields["disambiguation_skipped_reason"] = (
            f"候选表 {len(names)} 超上限 {max_sheets}，直接硬失败"
        )
        raise err
    visible = [name for name in names if not hidden.get(name, False)]
    if len(visible) != 1:
        err = SheetRequired(names)
        err.fields["visible_sheets"] = visible
        err.fields["disambiguation_skipped_reason"] = (
            "没有唯一可见工作表，需显式指定 sheet"
        )
        raise err
    matched = visible[0]
    return matched, {
        "strategy": "sole_visible_sheet",
        "hints": [],
        "matched_sheet": matched,
        "visible_sheets": visible,
        "hidden_sheets": [name for name in names if hidden.get(name, False)],
        "normalized_match": False,
    }


def resolve_sheet_by_columns(
    sheetnames: list[str],
    hints: list[str],
    probe: Any,
    *,
    max_sheets: int = _DISAMBIG_MAX_SHEETS,
) -> tuple[str, dict[str, Any]]:
    """多表省略 sheet 时的列感知消歧。

    Args:
        sheetnames: 工作簿全部表名。
        hints: 目标列名（已清洗：非空、非 "*"、非 Unnamed 前缀）。
        probe: ``probe(sheet) -> (columns | None, excluded | None, hidden)``。
            columns 为该表列名列表；excluded 为 "form"/"empty"/"error" 之一；
            hidden 为该表是否隐藏。探针必须对每张表传显式表名，
            禁止传 None 触发 wb.active 回退。

    Returns:
        (绑定表名, 消歧信息 dict)。信息含 strategy/hints/matched_sheet/
        candidates/excluded/normalized_match，供结果回填与评测断言。

    Raises:
        SheetRequired: 零命中、多命中、仅隐藏命中、超限、探针失败时。
            fields 含 available_sheets/hint_match_matrix/missing_per_sheet。
    """
    names = [str(n) for n in sheetnames if str(n)]
    clean_hints = [str(h).strip() for h in (hints or []) if str(h).strip()]
    if len(names) > max_sheets:
        err = SheetRequired(names)
        err.fields["disambiguation_skipped_reason"] = f"候选表 {len(names)} 超上限 {max_sheets}，直接硬失败"
        raise err
    if not clean_hints:
        raise SheetRequired(names)

    candidates: dict[str, list[str]] = {}
    excluded: dict[str, list[str]] = {"form": [], "empty": [], "error": []}
    hidden: dict[str, bool] = {}
    hint_matrix: dict[str, list[str]] = {}
    missing: dict[str, list[str]] = {}
    for name in names:
        try:
            columns, excluded_reason, is_hidden = probe(name)
        except Exception:
            columns, excluded_reason, is_hidden = None, "error", False
        hidden[name] = bool(is_hidden)
        if excluded_reason or not columns:
            excluded[excluded_reason or "empty"].append(name)
            continue
        cols = [str(c) for c in columns]
        candidates[name] = cols
        exact = set(cols)
        normed = {_normalize_column_name(c) for c in cols}
        hit_exact = [h for h in clean_hints if h in exact]
        hit_norm = [h for h in clean_hints if _normalize_column_name(h) in normed]
        hint_matrix[name] = hit_exact
        missing[name] = [h for h in clean_hints if h not in hit_exact]

    full_exact = [n for n, cols in candidates.items() if all(h in set(cols) for h in clean_hints)]
    full_norm = [
        n for n, cols in candidates.items()
        if all(_normalize_column_name(h) in {_normalize_column_name(c) for c in cols} for h in clean_hints)
    ]
    # 精确匹配蕴含归一匹配（full_exact ⊆ full_norm 恒成立）。
    # 归一仅用于容忍大小写/首尾空格：若归一制造了新的歧义（ winners 变化），判失败。
    # 隐藏表参与碰撞检测：任一隐藏表命中即视为歧义（永不自动选中隐藏表）。
    hidden_hit = [n for n in (set(full_exact) | set(full_norm)) if hidden.get(n)]

    def _fail(reason: str) -> SheetRequired:
        err = SheetRequired(names)
        err.fields["hint_match_matrix"] = hint_matrix
        err.fields["missing_per_sheet"] = missing
        err.fields["disambiguation_skipped_reason"] = reason
        return err

    if hidden_hit:
        raise _fail("命中涉及隐藏表，需显式指定 sheet")
    if len(full_norm) == 1:
        matched = full_norm[0]
        if full_exact and full_exact != [matched]:
            raise _fail("归一化改变了命中结果，存在歧义")
        normalized = full_exact != [matched]
        return matched, {
            "strategy": "column_unique_match",
            "hints": clean_hints,
            "matched_sheet": matched,
            "candidates": candidates,
            "excluded": {k: v for k, v in excluded.items() if v},
            "hidden_hit": hidden_hit,
            "normalized_match": normalized,
        }
    if not full_norm:
        raise _fail("目标列无唯一命中表")
    raise _fail("目标列在多张表同时命中")


def _hex_of_version(version: str) -> str:
    text = str(version)
    if text.startswith("sha256:"):
        return text[7:]
    return text


def _lock_for_backing(dest: Path) -> threading.Lock:
    key = str(dest)
    with _materialize_locks_mu:
        lock = _materialize_locks.get(key)
        if lock is None:
            lock = threading.Lock()
            _materialize_locks[key] = lock
        return lock


def _backing_ready(dest: Path, size: int) -> bool:
    try:
        return dest.is_file() and dest.stat().st_size == size
    except OSError:
        return False


def _touch_quiet(dest: Path) -> None:
    try:
        dest.touch()
    except OSError:
        pass


def _materialize_backing(
    workspace: WorkspaceRef,
    data: bytes,
    suffix: str,
    version: str,
) -> SnapshotBacking:
    digest = _hex_of_version(version)
    root = Path(workspace.root) / _SNAPSHOT_DIRNAME / _SNAPSHOTS
    root.mkdir(parents=True, exist_ok=True)
    dest = root / f"{digest}{suffix or '.bin'}"
    expected = len(data)
    with _lock_for_backing(dest):
        if not _backing_ready(dest, expected):
            from excelmanus.workspace.txlog import replace_with_retry

            fd, name = tempfile.mkstemp(prefix=".snapshot-", dir=str(root))
            try:
                with os.fdopen(fd, "wb") as handle:
                    handle.write(data)
                    handle.flush()
                    os.fsync(handle.fileno())
                try:
                    replace_with_retry(name, str(dest), retries=8, delay=0.05)
                    name = ""
                except PermissionError:
                    if not _backing_ready(dest, expected):
                        raise
            finally:
                if name:
                    Path(name).unlink(missing_ok=True)
        _touch_quiet(dest)
    if expected < SNAPSHOT_BLOB_THRESHOLD:
        return SnapshotBacking(
            kind="bytes",
            sha256=digest,
            suffix=suffix,
            payload=data,
            blob_path=dest,
        )
    return SnapshotBacking(
        kind="blob",
        sha256=digest,
        suffix=suffix,
        payload=None,
        blob_path=dest,
    )


def prune_snapshot_cache(workspace_root: str | Path, *, keep: int = 40, max_age_seconds: float = 86400) -> int:
    """Only discard old derived backing; live Python snapshots retain a lease."""
    workspace = Path(workspace_root).resolve()
    directory = (workspace / _SNAPSHOT_DIRNAME / _SNAPSHOTS).resolve()
    directory.relative_to(workspace)
    if not directory.is_dir():
        return 0
    active = {b.blob_path for b in list(_live_backings.values()) if b.blob_path}
    files = sorted((p for p in directory.iterdir() if p.is_file() and not p.name.startswith(".")), key=lambda p: p.stat().st_mtime, reverse=True)
    removed = 0
    for path in files[max(0, keep):]:
        if path in active or time.time() - path.stat().st_mtime < max_age_seconds:
            continue
        path.unlink(missing_ok=True)
        removed += 1
    return removed


def _read_and_convert(abs_path: Path, workspace: WorkspaceRef) -> tuple[bytes, str]:
    from excelmanus.xls_converter import ensure_xlsx, needs_conversion

    if needs_conversion(abs_path):
        converted, _ = ensure_xlsx(abs_path, workspace_root=workspace.root)
        return Path(converted).read_bytes(), ".xlsx"
    suffix = abs_path.suffix.lower() or ".bin"
    return abs_path.read_bytes(), suffix


def open_snapshot_at(
    abs_path: Path | str,
    *,
    relative: str,
    workspace: WorkspaceRef,
    expected_version: str | None = None,
) -> WorkbookSnapshot:
    """HTTP / 核验等无 ToolCallContext 的开簿。仍禁止之后再读活路径。"""
    path = Path(abs_path)
    if not path.is_file():
        raise SnapshotError(f"文件不存在: {relative or path}", code="PATH_INVALID")
    data, suffix = _read_and_convert(path, workspace)
    version = content_version_of(data)
    if expected_version and expected_version != version:
        raise SnapshotStale(
            f"{relative} 版本已变化：期望 {expected_version}，实际 {version}",
            fields={"expected_version": expected_version, "content_version": version, "path": relative},
        )
    file_ref = FileRef(workspace=workspace, relative=relative.replace("\\", "/"), observed_version=version)
    snap_id = SnapshotId(
        workspace_key=workspace.identity_key(),
        relative=file_ref.relative,
        content_version=version,
    )
    backing = _materialize_backing(workspace, data, suffix, version)
    remember_content_version(file_ref.relative, version)
    return WorkbookSnapshot(
        id=snap_id,
        file=file_ref,
        content_version=version,
        backing=backing,
        suffix=suffix,
    )


def open_snapshot_bytes(
    data: bytes,
    *,
    relative: str,
    workspace: WorkspaceRef,
    suffix: str | None = None,
    expected_version: str | None = None,
) -> WorkbookSnapshot:
    """Open immutable bytes that do not have a live user path (for history preview)."""
    file_suffix = suffix or Path(relative).suffix.lower() or ".bin"
    version = content_version_of(data)
    if expected_version and expected_version != version:
        raise SnapshotStale(
            f"{relative} 修订版本不匹配：期望 {expected_version}，实际 {version}",
            fields={"expected_version": expected_version, "content_version": version, "path": relative},
        )
    file_ref = FileRef(workspace=workspace, relative=relative.replace("\\", "/"), observed_version=version)
    backing = _materialize_backing(workspace, data, file_suffix, version)
    return WorkbookSnapshot(
        id=SnapshotId(workspace_key=workspace.identity_key(), relative=file_ref.relative, content_version=version),
        file=file_ref,
        content_version=version,
        backing=backing,
        suffix=file_suffix,
    )


def open_snapshot(raw: str, *, expected_version: str | None = None) -> WorkbookSnapshot:
    """正式工具开簿：解析 FileRef，读字节一次，钉 backing。"""
    from excelmanus.tools.context import require_guard, resolve_file_ref

    ref = resolve_file_ref(raw, observed_version=expected_version)
    guard = require_guard()
    abs_path = Path(guard.resolve_and_validate(raw))
    expected = expected_version or ref.observed_version
    return open_snapshot_at(
        abs_path,
        relative=ref.relative,
        workspace=ref.workspace,
        expected_version=expected,
    )


def bind_rect_sheet(wb: Any, rect: RectRef, default_sheet: str | None) -> RectRef:
    """把矩形绑到唯一缺省表规则（I8）。"""
    from excelmanus.tools._helpers import resolve_sheet_name

    names = list(wb.sheetnames)
    title = rect.sheet or default_sheet
    if title:
        resolved = resolve_sheet_name(title, names)
        if resolved is None:
            raise SnapshotError(
                f"工作表 '{title}' 不存在。该文件包含: {names}",
                code="SHEET_NOT_FOUND",
                fields={"available_sheets": names},
            )
        return rect if rect.sheet == resolved else replace_rect(rect, resolved)
    return replace_rect(rect, require_default_sheet(names, None))


def replace_rect(rect: RectRef, sheet: str) -> RectRef:
    from dataclasses import replace

    return replace(rect, sheet=sheet)


def selection_from_rows(
    snapshot: WorkbookSnapshot,
    *,
    sheet: str,
    rows: list[int],
    cols: list[int] | None = None,
    origin: str,
    header_row: int | str | None = None,
) -> BoundSelection:
    return BoundSelection(
        snapshot=snapshot.id,
        file=snapshot.file,
        sheet=sheet,
        rows=tuple(int(r) for r in rows),
        cols=tuple(int(c) for c in cols) if cols else None,
        origin=origin,
        header_row=header_row,
    )


def parse_bound_selection(raw: Any) -> BoundSelection | None:
    if isinstance(raw, str):
        text = raw.strip()
        # Large filter/range results expose a spill locator instead of putting
        # thousands of row numbers in the model-facing message.  Accept the
        # locator anywhere a normal selection object is accepted so the native
        # read -> write path remains composable.
        try:
            from excelmanus.engine_core.spill import (
                SpillNotFound,
                SpillStore,
                is_spill_reference,
            )

            if is_spill_reference(text):
                from excelmanus.tools.context import current_call

                call = current_call()
                if call is None:
                    return None
                payload = json.loads(SpillStore(call.binding.workspace.root).get(text))
                if isinstance(payload, dict) and isinstance(payload.get("selection"), dict):
                    raw = payload["selection"]
                else:
                    raw = payload
        except (SpillNotFound, ValueError, TypeError, json.JSONDecodeError):
            return None
        if not isinstance(raw, str):
            text = ""
        if text[:1] == "{":
            try:
                raw = json.loads(text)
            except (ValueError, TypeError):
                return None
    if not isinstance(raw, dict):
        return None
    if raw.get("result_kind") not in (None, "selection") and "rows" not in raw:
        return None
    snapshot_id = str(raw.get("snapshot_id") or "")
    version = str(raw.get("content_version") or "")
    relative = str(raw.get("file") or raw.get("file_path") or "")
    sheet = str(raw.get("sheet") or "")
    rows = raw.get("rows") or raw.get("source_rows")
    if not version or not relative or not sheet or not isinstance(rows, list) or not rows:
        return None
    parts = snapshot_id.split("|") if snapshot_id else []
    workspace_key = parts[0] if len(parts) >= 3 else ""
    from excelmanus.tools.context import current_call

    ctx = current_call()
    workspace = ctx.binding.workspace if ctx is not None else WorkspaceRef.from_root(".")
    if not workspace_key:
        workspace_key = workspace.identity_key()
    file_ref = FileRef(workspace=workspace, relative=relative, observed_version=version)
    cols = raw.get("cols") or raw.get("source_cols")
    col_tuple: tuple[int, ...] | None
    if isinstance(cols, dict):
        col_tuple = tuple(int(v) for v in cols.values() if str(v).isdigit() or isinstance(v, int))
    elif isinstance(cols, list):
        col_tuple = tuple(int(c) for c in cols)
    else:
        col_tuple = None
    return BoundSelection(
        snapshot=SnapshotId(workspace_key=workspace_key, relative=relative, content_version=version),
        file=file_ref,
        sheet=sheet,
        rows=tuple(int(r) for r in rows),
        cols=col_tuple,
        origin=str(raw.get("origin") or "filter"),
        header_row=raw.get("header_row"),
    )


def validate_selection_target(raw: Any, target_path: str, explicit_sheet: str | None = None) -> BoundSelection:
    """A matching content hash does not authorize reusing another file's rows."""
    from excelmanus.tools.context import resolve_file_ref

    try:
        selection = parse_bound_selection(raw)
        if selection is None:
            raise ValueError("selection 缺少文件、表、版本或行号")
        target = resolve_file_ref(target_path)
        source = resolve_file_ref(selection.file.relative)
        if selection.snapshot.workspace_key != target.workspace.identity_key() or source.relative != target.relative:
            raise ValueError("selection 所属工作区/文件与写入目标不一致")
        if explicit_sheet and explicit_sheet.casefold() != selection.sheet.casefold():
            raise ValueError("selection 所属工作表与写入目标不一致")
        if any(r < 1 or r > 1_048_576 for r in selection.rows) or len(set(selection.rows)) != len(selection.rows):
            raise ValueError("selection 行号越界或重复")
        if selection.cols is not None and (not selection.cols or any(c < 1 or c > 16_384 for c in selection.cols) or len(set(selection.cols)) != len(selection.cols)):
            raise ValueError("selection 列号越界或重复")
        return selection
    except (ValueError, TypeError) as exc:
        raise SnapshotError(str(exc), code="SELECTION_STALE") from exc


def apply_read_contract(
    payload: dict[str, Any],
    *,
    snapshot: WorkbookSnapshot,
    result_kind: ResultKind,
    sheet: str | None,
    header_row: Any = None,
    source_rows: list[int] | None = None,
    source_cols: Any = None,
    coverage: Coverage | None = None,
    formulas_uncached: Any = "unknown",
    selection: BoundSelection | None = None,
    meta_kind: str | None = None,
) -> dict[str, Any]:
    """类型化读合同。不猜测 overview 必为 sampled。"""
    payload["status"] = "success"
    payload["result_kind"] = result_kind
    payload["content_version"] = snapshot.content_version
    payload["snapshot_id"] = snapshot.id.key()
    payload["file_path"] = snapshot.file.relative
    if sheet:
        payload["resolved_sheet"] = sheet
    if header_row is not None:
        payload["header_row"] = header_row
        payload.setdefault("detected_header_row", header_row)
    if source_rows is not None:
        payload["source_rows"] = source_rows
    if source_cols is not None:
        payload["source_cols"] = source_cols
    cov = coverage or Coverage(
        kind="complete",
        returned_rows=len(source_rows or payload.get("values") or payload.get("data") or []),
        total_rows=None,
    )
    payload["coverage"] = cov.to_json()
    if selection is not None:
        payload["selection"] = selection.to_json()
    sampled = cov.kind == "sampled"
    truncated = cov.kind == "truncated" or bool(
        payload.get("is_truncated") or payload.get("truncated") or payload.get("has_more")
    )
    payload["meta"] = {
        "sheet": sheet or payload.get("resolved_sheet"),
        "header_row": header_row if header_row is not None else payload.get("header_row"),
        "truncated": truncated,
        "sampled": sampled,
        "formulas_uncached": formulas_uncached,
        "kind": meta_kind or result_kind,
    }
    if "range" not in payload:
        payload["range"] = payload.get("resolved_range")
    if payload.get("formulas") is None and payload.get("formula_grid") is not None:
        payload["formulas"] = payload["formula_grid"]
    if payload.get("values") is None:
        data = payload.get("data")
        preview = payload.get("preview")
        if isinstance(data, list):
            payload["values"] = data
        elif isinstance(preview, list):
            payload["values"] = preview
        else:
            payload["values"] = []
    return payload


def cell_fact_from_openpyxl(sheet: str, row: int, col: int, cell: Any, cached_value: Any) -> CellFact:
    raw = getattr(cell, "value", None)
    formula = raw if isinstance(raw, str) and raw.startswith("=") else None
    err = None
    if cached_value is not None and type(cached_value).__name__ == "Error":
        err = str(cached_value)
        cached_value = None
    if formula and cached_value is None:
        cached: CachedState = "no"
        source: FactSource = "formula_text"
        t: CellType = "z"
        text = None
        v = None
    elif formula:
        cached = "yes"
        source = "cached"
        t, v, text = _classify_value(cached_value)
    elif raw is None and cached_value is None:
        cached = "yes"
        source = "cached"
        t, v, text = "z", None, None
    else:
        cached = "yes"
        source = "cached"
        t, v, text = _classify_value(cached_value if cached_value is not None else raw)
    if err:
        t = "e"
    return CellFact(
        sheet=sheet,
        row=row,
        col=col,
        t=t,
        v=v,
        text=text,
        f=formula,
        cached=cached,
        source=source,
        e=err,
    )


def _classify_value(value: Any) -> tuple[CellType, Any, str | None]:
    if value is None:
        return "z", None, None
    if isinstance(value, bool):
        return "b", value, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "n", value, None
    from datetime import date, datetime

    if isinstance(value, (date, datetime)):
        return "d", value.isoformat(), str(value)
    text = str(value)
    return "s", text, text


def _unloaded_rects(
    used_rows: int,
    used_cols: int,
    loaded: list[dict[str, int]],
) -> list[dict[str, int]]:
    if used_rows <= 0 or used_cols <= 0:
        return []
    if not loaded:
        return [{"r0": 1, "c0": 1, "r1": used_rows, "c1": used_cols}]
    r0 = min(int(w["r0"]) for w in loaded)
    c0 = min(int(w["c0"]) for w in loaded)
    r1 = max(int(w["r1"]) for w in loaded)
    c1 = max(int(w["c1"]) for w in loaded)
    rects: list[dict[str, int]] = []
    if r0 > 1:
        rects.append({"r0": 1, "c0": 1, "r1": r0 - 1, "c1": used_cols})
    if c0 > 1:
        rects.append({
            "r0": max(r0, 1),
            "c0": 1,
            "r1": min(r1, used_rows),
            "c1": c0 - 1,
        })
    if c1 < used_cols:
        rects.append({
            "r0": max(r0, 1),
            "c0": c1 + 1,
            "r1": min(r1, used_rows),
            "c1": used_cols,
        })
    if r1 < used_rows:
        rects.append({"r0": r1 + 1, "c0": 1, "r1": used_rows, "c1": used_cols})
    return rects


def _decode_text_bytes(data: bytes) -> str:
    if not data:
        return ""
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("latin-1")


def _window_dims(ws: Any, rect: RectRef) -> tuple[dict[str, float], dict[str, float]]:
    from openpyxl.utils import get_column_letter

    col_widths: dict[str, float] = {}
    for col in range(rect.min_col, rect.max_col + 1):
        letter = get_column_letter(col)
        dim = ws.column_dimensions.get(letter)
        if dim is not None and dim.width:
            col_widths[letter] = float(dim.width)
    row_heights: dict[str, float] = {}
    for row in range(rect.min_row, rect.max_row + 1):
        dim = ws.row_dimensions.get(row)
        if dim is not None and dim.height:
            row_heights[str(row)] = float(dim.height)
    return col_widths, row_heights


def _window_merges(ws: Any, rect: RectRef) -> list[dict[str, int]]:
    from excelmanus.tools._style_extract import extract_merge_ranges

    merges: list[dict[str, int]] = []
    for item in extract_merge_ranges(ws):
        if (
            item["max_row"] < rect.min_row
            or item["min_row"] > rect.max_row
            or item["max_col"] < rect.min_col
            or item["min_col"] > rect.max_col
        ):
            continue
        merges.append(item)
    return merges


def project_csv_view(
    snapshot: WorkbookSnapshot,
    windows: list[RectRef],
) -> dict[str, Any]:
    """CSV/TSV 网格：第 1 行就是 R1，total 是全文行数。"""
    from excelmanus.workbook.csv_index import csv_index

    index = csv_index(snapshot)
    used_rows = index.rows
    used_cols = index.columns
    sheet_name = "Sheet1"
    sheets = [{"name": sheet_name, "sheet_id": sheet_name, "used": {"rows": used_rows, "cols": used_cols}}]
    projected: list[dict[str, Any]] = []
    loaded: list[dict[str, int]] = []
    for rect in windows:
        if rect.sheet and rect.sheet != sheet_name:
            raise SnapshotError(f"工作表 '{rect.sheet}' 不存在", code="SHEET_NOT_FOUND")
        r0 = max(1, rect.min_row)
        c0 = max(1, rect.min_col)
        r1 = min(rect.max_row, max(used_rows, 1))
        c1 = min(rect.max_col, max(used_cols, 1))
        cells: dict[str, Any] = {}
        for row, src in enumerate(index.window(r0, r1), r0):
            for col in range(c0, c1 + 1):
                raw = src[col - 1] if col - 1 < len(src) else ""
                if raw == "":
                    continue
                try:
                    value: Any = int(raw)
                    kind = "n"
                except ValueError:
                    try:
                        value = float(raw)
                        kind = "n"
                    except ValueError:
                        value = raw
                        kind = "s"
                cells[f"{row},{col}"] = {"t": kind, "v": value, "cached": "yes"}
        window_rect = {"r0": rect.min_row, "c0": rect.min_col, "r1": rect.max_row, "c1": rect.max_col}
        loaded.append({**window_rect, "sheet": sheet_name})
        projected.append({
            "sheet": sheet_name,
            "rect": window_rect,
            "cells": cells,
            "merges": [],
            "col_widths": {},
            "row_heights": {},
        })
    return {
        "file": {
            "workspaceKey": snapshot.file.workspace.identity_key(),
            "relative": snapshot.file.relative,
            "observedVersion": snapshot.content_version,
        },
        "content_version": snapshot.content_version,
        "snapshot_id": snapshot.id.key(),
        "active_sheet": sheet_name,
        "with_styles": True,
        "sheets": sheets,
        "windows": projected,
        "coverage": {
            "loaded": loaded,
            "unloaded": _unloaded_rects(used_rows, used_cols, loaded),
            "truncated_reason": "window" if used_rows > (loaded[0]["r1"] if loaded else 0) else None,
        },
    }


def _cached_workbook_pair(
    snapshot: WorkbookSnapshot, with_styles: bool
) -> tuple[Any, Any, threading.Lock]:
    """按不可变快照复用已解析的工作簿对；read_only 共享同一 ZipFile，须持锁遍历。"""
    from openpyxl import load_workbook

    def build() -> tuple[tuple[Any, Any, threading.Lock], int]:
        raw = snapshot.read_bytes()
        # BytesIO 而非文件路径：缓存期间不占用 OS 句柄，Windows 上不会锁死源文件。
        wb_f = load_workbook(io.BytesIO(raw), data_only=False, read_only=not with_styles)
        wb_v = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        cost = max(len(raw), 1) * (8 if with_styles else 3)
        return (wb_f, wb_v, threading.Lock()), cost

    return _workbook_views.get_or_create((snapshot.id.key(), with_styles), build)


def project_view(
    snapshot: WorkbookSnapshot,
    windows: list[RectRef],
    *,
    data_only: bool = False,
    with_styles: bool = True,
    active_sheet_default: bool = False,
) -> dict[str, Any]:
    """第 5 批 WorkbookViewSnapshot 的服务端投影。"""
    if snapshot.is_csv():
        return project_csv_view(snapshot, windows)
    from excelmanus.tools._style_extract import extract_cell_style

    wb_f, wb_v, wb_lock = _cached_workbook_pair(snapshot, with_styles)
    wb_lock.acquire()
    try:
        if active_sheet_default:
            active = wb_f.active or wb_f.worksheets[0]
            windows = [replace_rect(rect, rect.sheet or active.title) for rect in windows]
        sheets = []
        used_by_sheet: dict[str, tuple[int, int]] = {}
        for name in wb_f.sheetnames:
            ws = wb_f[name]
            used = (int(ws.max_row or 0), int(ws.max_column or 0))
            used_by_sheet[name] = used
            sheets.append({
                "name": name,
                "sheet_id": name,
                "used": {"rows": used[0], "cols": used[1]},
            })
        projected = []
        loaded_by_sheet: dict[str, list[dict[str, int]]] = {name: [] for name in wb_f.sheetnames}
        for rect in windows:
            title = require_default_sheet(list(wb_f.sheetnames), rect.sheet)
            ws_f = wb_f[title]
            ws_v = wb_v[title]
            cells: dict[str, Any] = {}
            # ReadOnlyWorksheet.cell() starts an XML scan on every call. Walk
            # both snapshots once, and omit default blanks from the wire data.
            from itertools import zip_longest

            used_rows, used_cols = used_by_sheet[title]
            bounds = dict(min_row=rect.min_row, max_row=min(rect.max_row, used_rows),
                          min_col=rect.min_col, max_col=min(rect.max_col, used_cols))
            rows_f = ws_f.iter_rows(**bounds) if bounds["max_row"] >= rect.min_row and bounds["max_col"] >= rect.min_col else ()
            rows_v = ws_v.iter_rows(**bounds) if bounds["max_row"] >= rect.min_row and bounds["max_col"] >= rect.min_col else ()
            for row, pair in enumerate(zip_longest(rows_f, rows_v, fillvalue=()), rect.min_row):
                for col, (cell, cached) in enumerate(zip_longest(*pair), rect.min_col):
                    cached_value = getattr(cached, "value", None)
                    has_style = with_styles and getattr(cell, "has_style", False)
                    if getattr(cell, "value", None) is None and cached_value is None and not has_style:
                        continue
                    fact = cell_fact_from_openpyxl(
                        title, row, col, cell, cached_value,
                    )
                    payload: dict[str, Any] = {
                        "t": fact.t,
                        "v": fact.v,
                        "cached": fact.cached,
                    }
                    if fact.f:
                        payload["f"] = fact.f
                    if fact.e:
                        payload["e"] = fact.e
                    if has_style:
                        style = extract_cell_style(cell)
                        if style:
                            payload["s"] = style
                    cells[f"{row},{col}"] = payload
            window_rect = {
                "r0": rect.min_row,
                "c0": rect.min_col,
                "r1": rect.max_row,
                "c1": rect.max_col,
            }
            loaded_by_sheet[title].append(window_rect)
            col_widths, row_heights = _window_dims(ws_f, rect) if with_styles else ({}, {})
            projected.append({
                "sheet": title,
                "rect": window_rect,
                "cells": cells,
                "merges": _window_merges(ws_f, rect),
                "col_widths": col_widths,
                "row_heights": row_heights,
            })
        unloaded: list[dict[str, Any]] = []
        for name, used in used_by_sheet.items():
            for item in _unloaded_rects(used[0], used[1], loaded_by_sheet.get(name) or []):
                unloaded.append({"sheet": name, **item})
        used_rows = max((u[0] for u in used_by_sheet.values()), default=0)
        loaded_max = max((w.max_row for w in windows), default=0)
        return {
            "file": {
                "workspaceKey": snapshot.file.workspace.identity_key(),
                "relative": snapshot.file.relative,
                "observedVersion": snapshot.content_version,
            },
            "content_version": snapshot.content_version,
            "snapshot_id": snapshot.id.key(),
            "active_sheet": (wb_f.active or wb_f.worksheets[0]).title,
            "with_styles": with_styles,
            "sheets": sheets,
            "windows": projected,
            "coverage": {
                "loaded": [
                    {
                        "sheet": require_default_sheet(list(wb_f.sheetnames), w.sheet),
                        "r0": w.min_row,
                        "c0": w.min_col,
                        "r1": w.max_row,
                        "c1": w.max_col,
                    }
                    for w in windows
                ],
                "unloaded": unloaded,
                "truncated_reason": "window" if used_rows > loaded_max else None,
            },
        }
    finally:
        wb_lock.release()
