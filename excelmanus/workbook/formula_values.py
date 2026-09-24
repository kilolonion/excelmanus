"""Use cached formula values only for an unchanged calculation input snapshot."""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json

from openpyxl import load_workbook


class FormulaValueError(ValueError):
    def __init__(self, message, *, code="FORMULA_CACHE_MISSING", cells=None):
        super().__init__(message)
        self.code = code
        self.cells = cells or []


def calculation_fingerprint(wb):
    facts = []
    for ws in wb.worksheets:
        facts.append((ws.title, [(pos, c.data_type, c.value) for pos, c in sorted(ws._cells.items()) if c.value is not None],
                      [(t.name, t.ref) for t in ws.tables.values()],
                      [(n.name, n.attr_text) for n in ws.defined_names.values()]))
    facts.append([(n.name, n.attr_text) for n in wb.defined_names.values()])
    return sha256(json.dumps(facts, default=str, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def attach_formula_source(wb, source_bytes):
    wb._em_formula_source = source_bytes
    wb._em_calculation_fingerprint = calculation_fingerprint(wb)


def analysis_rows(ws):
    """Return values, refusing missing/error/stale caches instead of coercing to 0."""
    formulas = [c for c in ws._cells.values() if c.data_type == "f"]
    if not formulas:
        return list(ws.iter_rows(values_only=True))
    wb = ws.parent
    source = getattr(wb, "_em_formula_source", None)
    if source is None:
        raise FormulaValueError("公式分析需要已计算的源文件缓存；请先调用 calculate_spreadsheet。", cells=[f"{ws.title}!{c.coordinate}" for c in formulas[:100]])
    if calculation_fingerprint(wb) != getattr(wb, "_em_calculation_fingerprint", None):
        raise FormulaValueError("当前批次修改了计算输入，旧公式缓存已失效；请提交并重算后再分析。", code="FORMULA_CACHE_STALE")
    cached = load_workbook(BytesIO(source), data_only=True, read_only=True)
    try:
        rows = list(ws.iter_rows(values_only=True))
        rows = [list(row) for row in rows]
        missing = []
        cached_cells = {(c.row, c.column): c for row in cached[ws.title].iter_rows() for c in row if c.value is not None}
        for cell in formulas:
            value_cell = cached_cells.get((cell.row, cell.column))
            if value_cell is None or value_cell.data_type == "e":
                missing.append(f"{ws.title}!{cell.coordinate}")
            else:
                rows[cell.row - 1][cell.column - 1] = value_cell.value
        if missing:
            raise FormulaValueError("公式缓存缺失或包含错误；请先重算，不能将其视为零。", cells=missing[:100])
        return rows
    finally:
        cached.close()


def snapshot_cache_status(snapshot, sheet, *, max_row=None):
    """Inspect cached values in a streaming pair; never infer missing formulas as zero."""
    from excelmanus.workbook.snapshot import _cached_workbook_pair
    formula, values, lock = _cached_workbook_pair(snapshot, False)
    missing, columns, count = [], set(), 0
    with lock:
        for index, (fr, vr) in enumerate(zip(formula[sheet].iter_rows(max_row=max_row), values[sheet].iter_rows(max_row=max_row))):
            if index % 512 == 0:
                from excelmanus.tools.spreadsheet_engine_tools import _cancelled
                _cancelled()
            for cell, value in zip(fr, vr):
                if cell.data_type == 'f' and (value.value is None or value.data_type == 'e'):
                    count += 1
                    columns.add(cell.column)
                    if len(missing) < 100:
                        missing.append(cell.coordinate)
    return {"status":"partial" if count else "complete", "missing_or_error_count":count,
            "cells":missing,"columns":sorted(columns),"scope":"saved_cache", "max_row":max_row}
