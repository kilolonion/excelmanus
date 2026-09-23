"""Apply UI workbook operations inside one openpyxl workbook (one CAS)."""

from __future__ import annotations

from typing import Any

from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from excelmanus.tools._style_extract import _BORDER_STYLE_MAP


_H_ALIGN = {1: "left", 2: "center", 3: "right", 4: "justify"}
_V_ALIGN = {1: "top", 2: "center", 3: "bottom"}
_BORDER_STYLE = {value: name for name, value in _BORDER_STYLE_MAP.items()}


def _rgb(value: Any) -> str | None:
    text = str(value or "").strip().lstrip("#")
    if len(text) == 8:
        text = text[2:]
    if len(text) != 6:
        return None
    return text.upper()


def apply_univer_style(cell: Any, style: dict[str, Any] | None) -> None:
    if style is None:
        from openpyxl.styles.cell_style import StyleArray
        cell._style = StyleArray()
        return
    if not isinstance(style, dict):
        raise ValueError("style 必须是展开的样式对象")
    if not style:
        return
    font_kw: dict[str, Any] = {}
    if "bl" in style:
        font_kw["bold"] = bool(style["bl"])
    if "it" in style:
        font_kw["italic"] = bool(style["it"])
    if "ul" in style:
        value = style["ul"]
        font_kw["underline"] = "single" if (value.get("s") if isinstance(value, dict) else value) else None
    if "st" in style:
        value = style["st"]
        font_kw["strike"] = bool(value.get("s") if isinstance(value, dict) else value)
    if style.get("fs"):
        font_kw["size"] = style["fs"]
    if style.get("ff"):
        font_kw["name"] = style["ff"]
    color = _rgb((style.get("cl") or {}).get("rgb"))
    if color:
        font_kw["color"] = Color(rgb=f"FF{color}")
    if font_kw:
        font = copy(cell.font) if cell.font else Font()
        for name, value in font_kw.items():
            setattr(font, name, value)
        cell.font = font
    bg = _rgb((style.get("bg") or {}).get("rgb"))
    if bg:
        cell.fill = PatternFill(fill_type="solid", fgColor=bg)
    align_kw: dict[str, Any] = {}
    if "ht" in style:
        align_kw["horizontal"] = _H_ALIGN.get(int(style["ht"]))
    if "vt" in style:
        align_kw["vertical"] = _V_ALIGN.get(int(style["vt"]))
    if "tb" in style:
        align_kw["wrap_text"] = style["tb"] == 3
    if (style.get("tr") or {}).get("a") is not None:
        align_kw["text_rotation"] = style["tr"]["a"]
    if (style.get("pd") or {}).get("l"):
        align_kw["indent"] = style["pd"]["l"]
    if "sk" in style:
        align_kw["shrink_to_fit"] = bool(style["sk"])
    if align_kw:
        alignment = copy(cell.alignment)
        for name, value in align_kw.items():
            setattr(alignment, name, value)
        cell.alignment = alignment
    if style.get("n") and style["n"].get("pattern"):
        cell.number_format = str(style["n"]["pattern"])
    bd = style.get("bd") or {}
    if bd:
        sides: dict[str, Any] = {}
        for key, attr in (("l", "left"), ("r", "right"), ("t", "top"), ("b", "bottom")):
            item = bd.get(key)
            if not item:
                continue
            color = _rgb((item.get("cl") or {}).get("rgb"))
            sides[attr] = Side(
                style=_BORDER_STYLE.get(int(item.get("s") or 1), "thin"),
                color=color,
            )
        if sides:
            border = copy(cell.border)
            for name, value in sides.items():
                setattr(border, name, value)
            cell.border = border


def _sheet(wb: Any, name: str | None) -> Any:
    from excelmanus.workbook.snapshot import require_default_sheet
    if name:
        return wb[require_default_sheet(list(wb.sheetnames), name)]
    if name:
        raise KeyError(name)
    ws = wb[require_default_sheet(list(wb.sheetnames), None)]
    if ws is None:
        raise KeyError("active")
    return ws


def apply_cell_value(ws: Any, cell_ref: str, value: Any) -> None:
    from excelmanus.workbook.refs import parse_rect
    from excelmanus.workbook.cells import assign_cell_value

    ref = parse_rect(str(cell_ref), default_sheet=ws.title)
    if ref.sheet != ws.title or ref.min_row != ref.max_row or ref.min_col != ref.max_col:
        raise ValueError("cell 必须是当前表的单元格地址")
    assign_cell_value(ws, ref.min_row, ref.min_col, value)


def apply_workbook_operations(wb: Any, operations: list[dict[str, Any]]) -> int:
    """Apply UI ops. Returns the number of cell-level writes (values or styles)."""
    written = 0
    from excelmanus.workbook.structure import assert_structure_supported
    allowed = {
        "set_values": {"cells"}, "set_styles": {"cells"}, "set_dims": {"rows", "columns"},
        "merge": {"range"}, "unmerge": {"range"},
        "insert_axis": {"axis", "index", "count"}, "delete_axis": {"axis", "index", "count"},
        "sheet_add": {"name"}, "sheet_rename": {"from", "to", "name"}, "sheet_delete": {"name"},
    }
    for raw in operations:
        if not isinstance(raw, dict):
            raise ValueError("operation 必须是对象")
        kind = str(raw.get("op") or raw.get("kind") or "").strip()
        if kind not in allowed or set(raw) - allowed[kind] - {"op", "kind", "sheet"}:
            raise ValueError(f"操作名称或字段无效：{kind}")
        if kind == "set_values":
            ws = _sheet(wb, raw.get("sheet"))
            for item in raw.get("cells") or []:
                if not isinstance(item, dict) or not item.get("cell"):
                    raise ValueError("cells 项必须提供 cell")
                if set(item) - {"cell", "value", "style"}:
                    raise ValueError("cells 包含未知字段")
                apply_cell_value(ws, str(item["cell"]), item.get("value"))
                if "style" in item:
                    apply_univer_style(ws[str(item["cell"]).upper()], item["style"])
                written += 1
        elif kind == "set_styles":
            ws = _sheet(wb, raw.get("sheet"))
            for item in raw.get("cells") or []:
                if not isinstance(item, dict) or not item.get("cell"):
                    raise ValueError("cells 项必须提供 cell")
                if "style" in item:
                    apply_univer_style(ws[str(item["cell"]).upper()], item["style"])
                written += 1
        elif kind == "set_dims":
            ws = _sheet(wb, raw.get("sheet"))
            for letter, width in (raw.get("columns") or {}).items():
                ws.column_dimensions[str(letter).upper()].width = float(width)
            for row_key, height in (raw.get("rows") or {}).items():
                ws.row_dimensions[int(row_key)].height = float(height)
        elif kind == "merge":
            ws = _sheet(wb, raw.get("sheet"))
            ws.merge_cells(str(raw.get("range") or ""))
        elif kind == "unmerge":
            ws = _sheet(wb, raw.get("sheet"))
            ws.unmerge_cells(str(raw.get("range") or ""))
        elif kind in {"insert_axis", "delete_axis"}:
            from excelmanus.workbook.structure_edit import apply_structure_edit
            ws = _sheet(wb, raw.get("sheet"))
            axis = str(raw.get("axis") or "row")
            index = int(raw.get("index") or 1)
            count = max(1, int(raw.get("count") or 1))
            if axis not in {"row", "col"} or index < 1 or int(raw.get("count", 1)) < 1:
                raise ValueError("axis/index/count 无效")
            action = ("insert_" if kind == "insert_axis" else "delete_") + ("cols" if axis == "col" else "rows")
            apply_structure_edit(wb, action, ws.title, at=index, count=count)
        elif kind == "sheet_add":
            title = str(raw.get("name") or "Sheet")
            wb.create_sheet(title=title)
        elif kind == "sheet_rename":
            from excelmanus.workbook.structure_edit import apply_structure_edit
            src = str(raw.get("from") or raw.get("sheet") or "")
            dest = str(raw.get("to") or raw.get("name") or "")
            if src in wb.sheetnames and dest:
                apply_structure_edit(wb, "rename_sheet", src, new_name=dest)
            else:
                raise ValueError("重命名的源表或目标名无效")
        elif kind == "sheet_delete":
            name = str(raw.get("name") or raw.get("sheet") or "")
            if name not in wb.sheetnames or len(wb.sheetnames) <= 1:
                raise ValueError("不能删除该工作表")
            assert_structure_supported(wb, kind, target_sheet=name)
            del wb[name]
        else:
            raise ValueError(f"不支持的工作簿操作：{kind}")
    return written


def changes_to_operations(
    changes: list[dict[str, Any]],
    default_sheet: str | None,
) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for change in changes:
        if not isinstance(change, dict) or not change.get("cell"):
            continue
        sheet = str(change.get("sheet") or default_sheet or "")
        grouped.setdefault(sheet, []).append(
            {"cell": change["cell"], "value": change.get("value"), **({"style": change["style"]} if "style" in change else {})}
        )
    return [
        {"op": "set_values", "sheet": sheet or None, "cells": cells}
        for sheet, cells in grouped.items()
    ]


def a1_range(min_row: int, min_col: int, max_row: int, max_col: int) -> str:
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
