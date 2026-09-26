"""格式化内部辅助：样式构建、样式读取、行列尺寸估算。

写入走 `apply_spreadsheet_changes`；本模块不再提交工作簿。
"""

from __future__ import annotations

import math
import re
import unicodedata
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from excelmanus.engine_core.tool_result import ToolResult, from_payload
from excelmanus.logger import get_logger
from excelmanus.security import FileAccessGuard
from excelmanus.tools.context import bind_workspace, require_guard
from excelmanus.tools._helpers import get_worksheet

logger = get_logger("tools.format")

# ── 中文颜色名 → 十六进制映射 ────────────────────────────

COLOR_NAME_MAP: dict[str, str] = {
    # 基础色
    "红": "FF0000", "红色": "FF0000", "red": "FF0000",
    "绿": "00B050", "绿色": "00B050", "green": "00B050",
    "蓝": "0000FF", "蓝色": "0000FF", "blue": "0000FF",
    "黄": "FFFF00", "黄色": "FFFF00", "yellow": "FFFF00",
    "白": "FFFFFF", "白色": "FFFFFF", "white": "FFFFFF",
    "黑": "000000", "黑色": "000000", "black": "000000",
    # 常用色
    "橙": "FFC000", "橙色": "FFC000", "orange": "FFC000",
    "紫": "7030A0", "紫色": "7030A0", "purple": "7030A0",
    "粉": "FF69B4", "粉色": "FF69B4", "pink": "FF69B4",
    "棕": "8B4513", "棕色": "8B4513", "brown": "8B4513",
    "灰": "808080", "灰色": "808080", "gray": "808080", "grey": "808080",
    "青": "00CED1", "青色": "00CED1", "cyan": "00FFFF",
    # 浅色系
    "浅蓝": "5B9BD5", "浅蓝色": "5B9BD5", "lightblue": "ADD8E6",
    "浅绿": "92D050", "浅绿色": "92D050", "lightgreen": "90EE90",
    "浅黄": "FFF2CC", "浅黄色": "FFF2CC", "lightyellow": "FFFFE0",
    "浅灰": "D9D9D9", "浅灰色": "D9D9D9", "lightgray": "D3D3D3",
    "浅红": "FF7F7F", "浅红色": "FF7F7F",
    "浅紫": "B4A7D6", "浅紫色": "B4A7D6",
    # 深色系
    "深蓝": "002060", "深蓝色": "002060", "darkblue": "00008B",
    "深绿": "006100", "深绿色": "006100", "darkgreen": "006400",
    "深红": "C00000", "深红色": "C00000", "darkred": "8B0000",
    "深灰": "404040", "深灰色": "404040", "darkgray": "A9A9A9",
    # Excel 主题常用色
    "金": "FFD700", "金色": "FFD700", "gold": "FFD700",
    "银": "C0C0C0", "银色": "C0C0C0", "silver": "C0C0C0",
    "天蓝": "4472C4", "天蓝色": "4472C4",
    "草绿": "70AD47", "草绿色": "70AD47",
    "珊瑚": "FF7F50", "珊瑚色": "FF7F50", "coral": "FF7F50",
}

def _get_guard() -> FileAccessGuard:
    return require_guard()


def init_guard(workspace_root: str) -> None:
    bind_workspace(workspace_root)


# ── 只读样式探查 ──────────────────────────────────────────

# ── 列宽 / 行高智能估算 ─────────────────────────────────

_MIN_COL_WIDTH = 8.0
_MAX_COL_WIDTH = 60.0
_COL_PADDING = 2.5
_MIN_ROW_HEIGHT = 15.0


def _is_wide_char(ch: str) -> bool:
    """判断字符是否为东亚宽字符（CJK / 全角）。"""
    ea = unicodedata.east_asian_width(ch)
    return ea in ("W", "F")


def _display_char_width(text: str) -> float:
    """计算文本的显示字符宽度（CJK 字符按 2.0 计）。"""
    width = 0.0
    for ch in text:
        width += 2.0 if _is_wide_char(ch) else 1.0
    return width


def _format_number_display(value: Any, number_format: str | None) -> str | None:
    """尝试模拟 number_format 后的显示文本长度。

    不追求 100% 精确——仅用于列宽估算。
    """
    if value is None or number_format is None or number_format == "General":
        return None
    try:
        num = float(value)
    except (ValueError, TypeError):
        return None

    fmt = number_format
    result_len_hint: str | None = None

    # 百分比：0.85 → "85%" or "85.0%"
    if "%" in fmt:
        pct = num * 100
        dec_match = re.search(r"0\.(0+)%", fmt)
        decimals = len(dec_match.group(1)) if dec_match else 0
        result_len_hint = f"{pct:,.{decimals}f}%" if "#,##" in fmt else f"{pct:.{decimals}f}%"
        return result_len_hint

    # 千分位 + 小数
    dec_match = re.search(r"0\.(0+)", fmt)
    decimals = len(dec_match.group(1)) if dec_match else 0
    has_comma = "#,##" in fmt or "," in fmt

    if has_comma:
        result_len_hint = f"{num:,.{decimals}f}"
    elif decimals > 0:
        result_len_hint = f"{num:.{decimals}f}"

    # 货币前缀/后缀
    for sym in ("$", "¥", "€", "£", "₩"):
        if sym in fmt and result_len_hint:
            result_len_hint = sym + result_len_hint
            break

    return result_len_hint


def _estimate_display_width(
    value: Any,
    font_size: float = 11.0,
    is_bold: bool = False,
    number_format: str | None = None,
) -> float:
    """估算单元格内容的显示字符宽度。

    返回以 Excel 列宽单位（≈字符数）计的宽度。
    """
    if value is None:
        return 0.0

    # 优先使用 number_format 模拟的显示文本
    display = _format_number_display(value, number_format)
    text = display if display is not None else str(value)

    # 多行取最长行
    lines = text.split("\n")
    char_width = max(_display_char_width(line) for line in lines) if lines else 0.0

    # 字体缩放
    scale = font_size / 11.0
    if is_bold:
        scale *= 1.07

    return char_width * scale


def _estimate_row_height(
    row_cells: tuple | list,
    col_widths: dict[str, float] | None = None,
) -> float:
    """估算一行的合适行高（pt）。

    考虑字体大小和 wrap_text 多行。
    """
    max_height = _MIN_ROW_HEIGHT

    for cell in row_cells:
        font_size = 11.0
        if cell.font and cell.font.size:
            font_size = float(cell.font.size)

        # 基础单行行高
        has_cjk = False
        val_str = str(cell.value) if cell.value is not None else ""
        if any(_is_wide_char(ch) for ch in val_str):
            has_cjk = True
        line_height = font_size * (1.45 if has_cjk else 1.35)

        # wrap_text 多行估算
        wrap = cell.alignment and cell.alignment.wrap_text
        if wrap and cell.value is not None and col_widths:
            col_letter = get_column_letter(cell.column)
            col_w = col_widths.get(col_letter, 8.0)
            # 可用字符宽度 ≈ 列宽 - padding
            usable = max(col_w - 1.0, 4.0)
            display_w = _estimate_display_width(
                cell.value,
                font_size=font_size,
                is_bold=bool(cell.font and cell.font.bold),
                number_format=cell.number_format if cell.number_format != "General" else None,
            )
            num_lines = max(1, math.ceil(display_w / usable))
            # 也考虑显式换行符
            explicit_lines = val_str.count("\n") + 1
            num_lines = max(num_lines, explicit_lines)
            cell_height = line_height * num_lines
        else:
            # 非 wrap 也考虑显式换行符
            explicit_lines = val_str.count("\n") + 1
            cell_height = line_height * explicit_lines

        if cell_height > max_height:
            max_height = cell_height

    return round(max_height, 1)


def apply_freeze_panes(ws: Any, freeze_panes: str | None) -> str:
    """在已打开的工作表上设置或取消冻结窗格，不提交文件。

    空字符串 / None 取消冻结。返回实际冻结单元格（如 ``A2``），取消时为空串。
    """
    text = "" if freeze_panes is None else str(freeze_panes).strip().replace("$", "")
    if not text or text.lower() in {"none", "null", "false", "0"}:
        ws.freeze_panes = None
        return ""
    from openpyxl.utils.cell import coordinate_to_tuple

    try:
        row, col = coordinate_to_tuple(text.upper())
        if not (1 <= row <= 1048576 and 1 <= col <= 16384):
            raise ValueError("超出 Excel 行列上限")
    except (ValueError, KeyError) as exc:
        raise ValueError("freeze_panes 需要单个 A1 单元格，或空字符串取消冻结") from exc
    ws.freeze_panes = text.upper()
    applied = getattr(ws, "freeze_panes", None)
    return str(applied) if applied else ""


def _size_entries(sizes: dict[Any, Any], *, axis: str) -> dict[str, float]:
    """Validate every entry before applying any dimensions; never drop invalid entries."""
    import math
    from openpyxl.utils import column_index_from_string

    out: dict[str, float] = {}
    for key, value in sizes.items():
        text = str(key).strip()
        try:
            index = column_index_from_string(text.upper()) if axis == "column" and text.isalpha() else int(text)
        except (ValueError, TypeError) as exc:
            raise ValueError(f"{axis} 尺寸键 {key!r} 必须是 Excel 1-based 行列号或列字母") from exc
        if isinstance(key, bool) or index < 1 or index > (16384 if axis == "column" else 1048576):
            raise ValueError(f"{axis} 尺寸键 {key!r} 超出 Excel 行列上限")
        max_size = 255 if axis == "column" else 409.5
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 < value <= max_size:
            raise ValueError(f"{axis}[{key!r}] 必须为 0 < 尺寸 <= {max_size} 的有限数值；列宽用字符宽度，行高用 points")
        canonical = get_column_letter(index) if axis == "column" else str(index)
        if canonical in out and out[canonical] != value:
            raise ValueError(f"{axis} 尺寸键 {key!r} 与 {canonical} 冲突")
        out[canonical] = float(value)
    return out


def apply_column_sizes(
    ws: Any,
    columns: dict[str, float] | None = None,
    auto_fit: bool = False,
    letters: set[str] | None = None,
) -> dict[str, float]:
    """在已打开的工作表上调整列宽，不提交文件。"""
    adjusted: dict[str, float] = {}
    if columns:
        from excelmanus.workbook.geometry import materialize_column_spans
        from openpyxl.utils import column_index_from_string
        normalized = _size_entries(columns, axis="column")
        materialize_column_spans(ws, [column_index_from_string(k) for k in normalized])
        for col_letter, width in normalized.items():
            ws.column_dimensions[str(col_letter).upper()].width = float(width)
            adjusted[str(col_letter).upper()] = float(width)
        return adjusted
    if not auto_fit:
        return adjusted

    merged_non_anchor: set[str] = set()
    for mr in ws.merged_cells.ranges:
        anchor = f"{get_column_letter(mr.min_col)}{mr.min_row}"
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                addr = f"{get_column_letter(c)}{r}"
                if addr != anchor:
                    merged_non_anchor.add(addr)

    for col_cells in ws.iter_cols(min_row=1, max_row=ws.max_row):
        max_w = 0.0
        col_letter = get_column_letter(col_cells[0].column)
        if letters is not None and col_letter not in letters:
            continue
        for cell in col_cells:
            coord = f"{col_letter}{cell.row}"
            if coord in merged_non_anchor or cell.value is None:
                continue
            font_size = 11.0
            is_bold = False
            if cell.font:
                if cell.font.size:
                    font_size = float(cell.font.size)
                if cell.font.bold:
                    is_bold = True
            nf = cell.number_format if cell.number_format != "General" else None
            w = _estimate_display_width(cell.value, font_size, is_bold, nf)
            if w > max_w:
                max_w = w
        width = max(_MIN_COL_WIDTH, min(max_w + _COL_PADDING, _MAX_COL_WIDTH))
        ws.column_dimensions[col_letter].width = width
        adjusted[col_letter] = round(width, 2)
    return adjusted


def apply_row_sizes(
    ws: Any,
    rows: dict[str, float] | None = None,
    auto_fit: bool = False,
    row_numbers: set[int] | None = None,
) -> dict[str, float]:
    """在已打开的工作表上调整行高，不提交文件。"""
    adjusted: dict[str, float] = {}
    if rows:
        for row_num_str, height in _size_entries(rows, axis="row").items():
            row_num = int(row_num_str)
            ws.row_dimensions[row_num].height = float(height)
            adjusted[str(row_num)] = float(height)
        return adjusted
    if not auto_fit:
        return adjusted

    current_col_widths: dict[str, float] = {}
    for col_letter, dim in ws.column_dimensions.items():
        if dim.width is not None:
            current_col_widths[col_letter] = dim.width
    for row_idx in range(1, (ws.max_row or 0) + 1):
        if row_numbers is not None and row_idx not in row_numbers:
            continue
        row_cells = [cell for cell in ws[row_idx] if not isinstance(cell, MergedCell)]
        if not row_cells:
            continue
        height = _estimate_row_height(row_cells, current_col_widths)
        ws.row_dimensions[row_idx].height = height
        adjusted[str(row_idx)] = height
    return adjusted


# ── 内部辅助函数 ──────────────────────────────────────────


def _resolve_color(value: str | None) -> str | None:
    """将颜色名称或十六进制码统一解析为十六进制码。"""
    if value is None:
        return None
    normalized = value.strip().lower()
    if normalized in COLOR_NAME_MAP:
        return COLOR_NAME_MAP[normalized]
    # 去除可能的 # 前缀
    hex_value = value.strip().lstrip("#")
    if len(hex_value) in (6, 8) and all(c in "0123456789abcdefABCDEF" for c in hex_value):
        return hex_value.upper()
    return value


LEGAL_FILL_PATTERN_TYPES: frozenset[str] = frozenset(
    {
        "solid",
        "none",
        "darkDown",
        "darkGray",
        "darkGrid",
        "darkHorizontal",
        "darkTrellis",
        "darkUp",
        "darkVertical",
        "gray0625",
        "gray125",
        "lightDown",
        "lightGray",
        "lightGrid",
        "lightHorizontal",
        "lightTrellis",
        "lightUp",
        "lightVertical",
        "mediumGray",
    }
)


def _build_font(config: dict[str, Any]) -> Font:
    """从配置字典构建 openpyxl Font 对象。"""
    if isinstance(config, str):
        return Font(name=config)
    if "strike" in config and "strikethrough" in config and config["strike"] != config["strikethrough"]:
        raise ValueError("font 别名冲突: strike 与 strikethrough 值不同")
    return Font(
        name=config.get("name"),
        size=config.get("size"),
        bold=config.get("bold"),
        italic=config.get("italic"),
        color=_resolve_color(config.get("color")),
        underline=config.get("underline"),
        strike=config.get("strikethrough", config.get("strike")),
    )


def _patch_font(existing: Any, config: dict[str, Any] | str) -> Font:
    """Patch only declared fields, preserving all other font attributes."""
    from copy import copy

    if isinstance(config, str):
        config = {"name": config}
    allowed = {"name", "size", "bold", "italic", "color", "underline", "strike", "strikethrough", "vertAlign"}
    if not isinstance(config, dict) or set(config) - allowed:
        raise ValueError(f"font 支持的字段: {sorted(allowed)}")
    if "strike" in config and "strikethrough" in config and config["strike"] != config["strikethrough"]:
        raise ValueError("font 别名冲突: strike 与 strikethrough 值不同")
    font = copy(existing) if existing is not None else Font()
    for key, value in config.items():
        key = "strike" if key == "strikethrough" else key
        setattr(font, key, _resolve_color(value) if key == "color" else value)
    return font


def _build_fill(config: dict[str, Any]) -> PatternFill:
    """从配置字典构建 openpyxl PatternFill 对象。不接受笼统的 type=pattern。"""
    color_keys = ("color", "fgColor", "start_color", "fg_color")
    color_values = [config[key] for key in color_keys if config.get(key) not in (None, "")]
    if len({str(_resolve_color(value)) for value in color_values}) > 1:
        raise ValueError("fill 别名冲突：color/fgColor/start_color/fg_color 值不同")
    color = _resolve_color(color_values[0] if color_values else None) or "FFFFFF"
    type_keys = ("type", "fill_type", "patternType", "pattern")
    type_values = [config[key] for key in type_keys if config.get(key) not in (None, "")]
    if len({str(value) for value in type_values}) > 1:
        raise ValueError("fill 别名冲突：type/fill_type/patternType/pattern 值不同")
    fill_type = (type_values[0] if type_values else "solid")
    fill_type = str(fill_type)
    if fill_type == "pattern":
        raise ValueError(
            "fill.type=pattern 不是合法 patternType。请用 solid 或 openpyxl 的 patternType 名称。"
        )
    if fill_type not in LEGAL_FILL_PATTERN_TYPES:
        raise ValueError(
            f"fill.type={fill_type!r} 不是合法 patternType。"
            f"可用：{sorted(LEGAL_FILL_PATTERN_TYPES - {'none'})}"
        )
    return PatternFill(
        start_color=color,
        end_color=_resolve_color(config.get("end_color")) or color,
        fill_type=None if fill_type == "none" else fill_type,
    )


def _build_side(side_config: dict[str, Any] | str) -> Side:
    """从配置构建单个 Side 对象。"""
    if isinstance(side_config, str):
        return Side(style=None if side_config == "none" else side_config, color="000000")
    return Side(
        style=None if side_config.get("style") == "none" else side_config.get("style", "thin"),
        color=_resolve_color(side_config.get("color")) or "000000",
    )


def _build_border(config: dict[str, Any]) -> Border:
    """从配置字典构建 openpyxl Border 对象。支持统一设置或单边差异化。"""
    # 如果指定了 left/right/top/bottom 中任意一个，使用单边模式
    has_sides = any(k in config for k in ("left", "right", "top", "bottom"))
    if has_sides:
        return Border(
            left=_build_side(config["left"]) if "left" in config else Side(),
            right=_build_side(config["right"]) if "right" in config else Side(),
            top=_build_side(config["top"]) if "top" in config else Side(),
            bottom=_build_side(config["bottom"]) if "bottom" in config else Side(),
        )
    # 统一模式：四边相同
    style = config.get("style", "thin")
    color = _resolve_color(config.get("color")) or "000000"
    side = Side(style=None if style == "none" else style, color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def _build_alignment(config: dict[str, Any]) -> Alignment:
    """从配置字典构建 openpyxl Alignment 对象。"""
    aliases = {
        "horizontal": ("horizontal", "horizontalAlignment"),
        "vertical": ("vertical", "verticalAlignment"),
        "wrap_text": ("wrap_text", "wrapText"),
    }
    normalized: dict[str, Any] = {}
    for canonical, keys in aliases.items():
        values = [config[key] for key in keys if key in config and config[key] is not None]
        if len(values) > 1 and values[0] != values[1]:
            raise ValueError(f"alignment 别名冲突: {keys[0]} 与 {keys[1]} 值不同")
        if values:
            normalized[canonical] = values[0]
    return Alignment(
        horizontal=normalized.get("horizontal"),
        vertical=normalized.get("vertical"),
        wrap_text=normalized.get("wrap_text"),
    )


def _patch_alignment(existing: Any, config: dict[str, Any]) -> Alignment:
    """Patch alignment using Excel/openpyxl field names and public aliases."""
    from copy import copy

    aliases = {"horizontalAlignment": "horizontal", "verticalAlignment": "vertical",
               "wrap_text": "wrapText", "shrink_to_fit": "shrinkToFit", "text_rotation": "textRotation"}
    allowed = {"horizontal", "vertical", "wrapText", "shrinkToFit", "textRotation", "indent", "readingOrder"}
    if not isinstance(config, dict):
        raise ValueError("alignment 必须是对象")
    alignment = copy(existing) if existing is not None else Alignment()
    normalized = {}
    for key, value in config.items():
        name = aliases.get(key, key)
        if name not in allowed:
            raise ValueError(f"alignment 不支持字段 {key}；可用 {sorted(allowed | set(aliases))}")
        if name in normalized and normalized[name] != value:
            raise ValueError(f"alignment 别名冲突: {name}")
        normalized[name] = value
    for name, value in normalized.items():
        setattr(alignment, name, value)
    return alignment


# ── 条件格式规则构建 ─────────────────────────────────────
#
# apply_spreadsheet_changes kind=conditional_format 与 workbook_spec
# conditional_formats 共用的规则构建。输入是模型面 dict（允许
# 常用别名/大小写混写），输出 openpyxl 规则对象；非法输入抛 ValueError。

_CF_OPERATOR_MAP: dict[str, str] = {
    "gt": "greaterThan", "greaterthan": "greaterThan",
    "greater_than": "greaterThan", ">": "greaterThan", "more_than": "greaterThan",
    "ge": "greaterThanOrEqual", "gte": "greaterThanOrEqual",
    "greaterthanorequal": "greaterThanOrEqual", "greaterthanorequals": "greaterThanOrEqual",
    "greater_than_or_equal": "greaterThanOrEqual", "greater_than_or_equals": "greaterThanOrEqual",
    "greater_or_equal": "greaterThanOrEqual", "greater_or_equal_to": "greaterThanOrEqual",
    "greater_than_or_equal_to": "greaterThanOrEqual",
    "greater_than_equal": "greaterThanOrEqual", ">=": "greaterThanOrEqual",
    "more_than_or_equal": "greaterThanOrEqual",
    "lt": "lessThan", "lessthan": "lessThan", "less_than": "lessThan",
    "<": "lessThan",
    "le": "lessThanOrEqual", "lte": "lessThanOrEqual",
    "lessthanorequal": "lessThanOrEqual", "lessthanorequals": "lessThanOrEqual",
    "less_than_or_equal": "lessThanOrEqual", "less_than_or_equals": "lessThanOrEqual",
    "less_or_equal": "lessThanOrEqual", "less_or_equal_to": "lessThanOrEqual",
    "less_than_or_equal_to": "lessThanOrEqual",
    "less_than_equal": "lessThanOrEqual", "<=": "lessThanOrEqual",
    "eq": "equal", "equal": "equal", "equals": "equal", "equal_to": "equal",
    "=": "equal", "==": "equal",
    "ne": "notEqual", "notequal": "notEqual", "not_equal": "notEqual",
    "not_equals": "notEqual", "not_equal_to": "notEqual", "!=": "notEqual", "<>": "notEqual",
    "between": "between", "not_between": "notBetween", "notbetween": "notBetween",
}


def _normalize_operator_name(raw_op: Any) -> str | None:
    """宽松归一化比较操作符（容忍 snake_case、空格、下划线、大小写与常见别名）。"""
    if not raw_op:
        return None
    s = str(raw_op).strip()
    if s in _CF_OPERATOR_MAP:
        return _CF_OPERATOR_MAP[s]
    lower = s.lower()
    if lower in _CF_OPERATOR_MAP:
        return _CF_OPERATOR_MAP[lower]
    clean = lower.replace("_", "").replace(" ", "").replace("-", "")
    return _CF_OPERATOR_MAP.get(clean)


_CF_ICON_STYLE_MAP: dict[str, str] = {
    "3_arrows": "3Arrows", "3arrows": "3Arrows",
    "3_arrows_gray": "3ArrowsGray", "3arrowsgray": "3ArrowsGray",
    "3_flags": "3Flags", "3flags": "3Flags",
    "3_traffic_lights": "3TrafficLights1",
    "3_traffic_lights_1": "3TrafficLights1", "3trafficlights1": "3TrafficLights1",
    "3_traffic_lights_2": "3TrafficLights2", "3trafficlights2": "3TrafficLights2",
    "3_signs": "3Signs", "3signs": "3Signs",
    "3_symbols": "3Symbols", "3symbols": "3Symbols",
    "3_symbols_2": "3Symbols2", "3symbols2": "3Symbols2",
    "4_arrows": "4Arrows", "4arrows": "4Arrows",
    "4_arrows_gray": "4ArrowsGray", "4arrowsgray": "4ArrowsGray",
    "4_traffic_lights": "4TrafficLights", "4trafficlights": "4TrafficLights",
    "4_ratings": "4Ratings", "4ratings": "4Ratings",
    "5_arrows": "5Arrows", "5arrows": "5Arrows",
    "5_arrows_gray": "5ArrowsGray", "5arrowsgray": "5ArrowsGray",
    "5_ratings": "5Ratings", "5ratings": "5Ratings",
    "5_quarters": "5Quarters", "5quarters": "5Quarters",
}

_CF_RULE_TYPES = (
    "cell_value / text / formula / duplicate / unique / top_n / bottom_n / "
    "color_scale / data_bar / icon_set"
)


def _cf_differential_kwargs(rule_spec: dict[str, Any]) -> dict[str, Any]:
    """条件格式命中的字体/填充/边框（dxf 样式），复用 format 构建器。支持顶层扁平样式属性容错。"""
    out: dict[str, Any] = {}
    font_cfg = rule_spec.get("font")
    if font_cfg is None:
        fc = rule_spec.get("font_color") or rule_spec.get("text_color")
        if fc:
            font_cfg = {"color": fc}
    fill_cfg = rule_spec.get("fill")
    if fill_cfg is None:
        bg = rule_spec.get("bg_color") or rule_spec.get("background") or rule_spec.get("fill_color")
        if bg:
            fill_cfg = {"color": bg, "fill_type": "solid"}
    border_cfg = rule_spec.get("border")
    if font_cfg:
        out["font"] = _patch_font(None, font_cfg)
    if fill_cfg:
        out["fill"] = _build_fill(fill_cfg)
    if border_cfg:
        out["border"] = _build_border(border_cfg)
    return out


def _cf_dxf(rule_spec: dict[str, Any]) -> Any:
    from openpyxl.styles.differential import DifferentialStyle

    kwargs = _cf_differential_kwargs(rule_spec)
    if not kwargs:
        return None
    return DifferentialStyle(**kwargs)


def _rule_formula(rule_spec: dict[str, Any]) -> str:
    """Shared expression input for conditional formatting and custom validation."""
    formulas: list[str] = []
    for key in ("formula", "formula1"):
        raw = rule_spec.get(key)
        if raw is None:
            continue
        if isinstance(raw, list) and len(raw) == 1:
            raw = raw[0]
        if not isinstance(raw, str) or not raw.strip():
            raise ValueError(f"rule.{key} 需要非空公式字符串（或单元素字符串数组）")
        formula = raw.strip().removeprefix("=").strip()
        if not formula:
            raise ValueError(f"rule.{key} 不能为空公式")
        formulas.append(formula)
    if not formulas:
        raise ValueError("需要 rule.formula（也接受同义字段 formula1），例如 =$D2=\"未匹配\"")
    if len(set(formulas)) != 1:
        raise ValueError("rule.formula 与 rule.formula1 冲突；只传一个，或使用相同公式")
    return formulas[0]


def build_conditional_format_rule(
    rule_spec: dict[str, Any], *, anchor: str = "A1"
) -> Any:
    """按模型面 dict 构建 openpyxl 条件格式规则。

    ``anchor`` 为规则范围的左上角单元格，仅 text 类型拼 SEARCH 公式用。
    非法输入抛 ValueError。
    """
    from openpyxl.formatting.rule import (
        CellIsRule,
        ColorScaleRule,
        DataBarRule,
        FormulaRule,
        IconSetRule,
        Rule,
    )

    if not isinstance(rule_spec, dict):
        raise ValueError("conditional_format 需要 rule 对象")
    rtype = str(rule_spec.get("type") or "cell_value").strip().lower()
    style_kwargs = _cf_differential_kwargs(rule_spec)

    if rtype in {"cell_value", "cellvalue", "cell_is", "value"}:
        raw_op = str(rule_spec.get("operator") or "ge").strip()
        operator = _normalize_operator_name(raw_op)
        if not operator:
            raise ValueError(
                f"conditional_format.rule.operator={raw_op!r} 不支持；"
                f"可用 {sorted(set(_CF_OPERATOR_MAP.values()))}"
            )
        if rule_spec.get("value") is not None and rule_spec.get("formula1") is not None and str(rule_spec["value"]) != str(rule_spec["formula1"]):
            raise ValueError("rule.value 与 rule.formula1 冲突；只传一个，或使用相同值")
        if rule_spec.get("value2") is not None and rule_spec.get("formula2") is not None and str(rule_spec["value2"]) != str(rule_spec["formula2"]):
            raise ValueError("rule.value2 与 rule.formula2 冲突；只传一个，或使用相同值")
        v1 = (
            rule_spec.get("value")
            if rule_spec.get("value") is not None
            else rule_spec.get("formula1")
        )
        if v1 is None:
            v1 = rule_spec.get("min")
        if v1 is None:
            v1 = rule_spec.get("val")
        if v1 is None:
            v1 = rule_spec.get("threshold")
        if v1 is None:
            v1 = rule_spec.get("target")
        if v1 is None:
            v1 = rule_spec.get("cutoff")
        if v1 is None:
            v1 = rule_spec.get("limit")

        v2 = (
            rule_spec.get("value2")
            if rule_spec.get("value2") is not None
            else rule_spec.get("formula2")
        )
        if v2 is None:
            v2 = rule_spec.get("max")
        formulas = [
            str(v)
            for v in (v1, v2)
            if v is not None and str(v) != ""
        ]
        if operator in {"between", "notBetween"}:
            if len(formulas) < 2:
                raise ValueError(
                    f"operator={operator} 需要两个边界值：rule.value 与 rule.value2（或 min 与 max）"
                )
        elif not formulas:
            raise ValueError(
                f"operator={operator} 需要 rule.value 阈值（between/notBetween 用 value+value2）"
            )
        return CellIsRule(operator=operator, formula=formulas[:2], **style_kwargs)

    if rtype in {"formula", "expression"}:
        formula = _rule_formula(rule_spec)
        return FormulaRule(formula=[formula], **style_kwargs)

    if rtype in {"text", "contains_text", "containstext", "text_contains"}:
        text = str(rule_spec.get("text") or rule_spec.get("value") or "")
        if not text:
            raise ValueError("type=text 需要 rule.text")
        rule = Rule(
            type="containsText",
            operator="containsText",
            text=text,
            dxf=_cf_dxf(rule_spec),
        )
        escaped = text.replace('"', '""')
        rule.formula = [f'NOT(ISERROR(SEARCH("{escaped}",{anchor})))']
        return rule

    if rtype in {"duplicate", "duplicates", "duplicate_values", "duplicatevalues"}:
        return Rule(type="duplicateValues", dxf=_cf_dxf(rule_spec))
    if rtype in {"unique", "unique_values", "uniquevalues"}:
        return Rule(type="uniqueValues", dxf=_cf_dxf(rule_spec))

    if rtype in {"top", "top_n", "topn", "top10", "bottom", "bottom_n", "bottomn"}:
        n = rule_spec.get("n")
        if n is None:
            n = rule_spec.get("value")
        try:
            rank = int(n)
        except (TypeError, ValueError):
            rank = 10
        bottom = rtype in {"bottom", "bottom_n", "bottomn"}
        return Rule(
            type="top10",
            rank=max(1, rank),
            bottom=bottom,
            percent=bool(rule_spec.get("percent")),
            dxf=_cf_dxf(rule_spec),
        )

    if rtype in {"color_scale", "colorscale", "scale", "3_color_scale", "2_color_scale"}:
        min_c = _resolve_color(rule_spec.get("min_color") or rule_spec.get("start_color") or "F8696B")
        max_c = _resolve_color(rule_spec.get("max_color") or rule_spec.get("end_color") or "63BE7B")
        kwargs: dict[str, Any] = {
            "start_type": "min",
            "start_color": min_c,
            "end_type": "max",
            "end_color": max_c,
        }
        mid_c = _resolve_color(rule_spec.get("mid_color"))
        if mid_c:
            kwargs["mid_type"] = "percentile"
            kwargs["mid_value"] = 50
            kwargs["mid_color"] = mid_c
        return ColorScaleRule(**kwargs)

    if rtype in {"data_bar", "databar", "bar"}:
        return DataBarRule(
            start_type="min",
            end_type="max",
            color=_resolve_color(rule_spec.get("bar_color") or rule_spec.get("color") or "638EC6"),
            showValue=True,
        )

    if rtype in {"icon_set", "iconset", "icon"}:
        raw_style = str(rule_spec.get("icon_style") or rule_spec.get("style") or "3_arrows").strip()
        icon_style = _CF_ICON_STYLE_MAP.get(raw_style.lower()) or _CF_ICON_STYLE_MAP.get(raw_style) or raw_style
        values = rule_spec.get("values")
        if not (isinstance(values, list) and values):
            values = [0, 33, 67] if icon_style.startswith("3") else [0, 25, 50, 75] if icon_style.startswith("4") else [0, 20, 40, 60, 80]
        return IconSetRule(
            icon_style=icon_style,
            type=str(rule_spec.get("value_type") or "percent"),
            values=values,
            showValue=True,
            percent=bool(rule_spec.get("percent", True)),
        )

    raise ValueError(
        f"conditional_format.rule.type={rtype!r} 不支持；可用 {_CF_RULE_TYPES}"
    )


# ── 数据验证（data_validation）──────────────────────────────

_DV_TYPE_MAP: dict[str, str] = {
    "list": "list",
    "dropdown": "list",
    "whole": "whole",
    "int": "whole",
    "integer": "whole",
    "decimal": "decimal",
    "float": "decimal",
    "number": "decimal",
    "date": "date",
    "time": "time",
    "textlength": "textLength",
    "text_length": "textLength",
    "length": "textLength",
    "custom": "custom",
    "formula": "custom",
}
_DV_TYPES = sorted(set(_DV_TYPE_MAP.values()))

_DV_OPERATOR_MAP: dict[str, str] = {
    "between": "between",
    "notbetween": "notBetween",
    "not_between": "notBetween",
    "equal": "equal",
    "eq": "equal",
    "=": "equal",
    "==": "equal",
    "notequal": "notEqual",
    "not_equal": "notEqual",
    "ne": "notEqual",
    "!=": "notEqual",
    "<>": "notEqual",
    "greaterthan": "greaterThan",
    "gt": "greaterThan",
    ">": "greaterThan",
    "lessthan": "lessThan",
    "lt": "lessThan",
    "<": "lessThan",
    "greaterthanorequal": "greaterThanOrEqual",
    "ge": "greaterThanOrEqual",
    "gte": "greaterThanOrEqual",
    ">=": "greaterThanOrEqual",
    "lessthanorequal": "lessThanOrEqual",
    "le": "lessThanOrEqual",
    "lte": "lessThanOrEqual",
    "<=": "lessThanOrEqual",
}

_DV_ERROR_STYLE_MAP: dict[str, str] = {
    "stop": "stop",
    "warning": "warning",
    "information": "information",
    "info": "information",
}

_ISO_DATE_RE = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})$")
_ISO_TIME_RE = re.compile(r"^(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?$")


def _dv_formula_value(raw: Any, *, dv_type: str) -> str:
    """数值/日期/时间边界值 → Excel 公式串；ISO 日期时间转 DATE/TIME 公式。"""
    text = str(raw).strip()
    if dv_type == "date":
        m = _ISO_DATE_RE.match(text)
        if m:
            return f"DATE({m.group(1)},{int(m.group(2))},{int(m.group(3))})"
    if dv_type == "time":
        m = _ISO_TIME_RE.match(text)
        if m:
            return f"TIME({int(m.group(1))},{int(m.group(2))},{int(m.group(3) or 0)})"
    return text


def build_data_validation(rule_spec: dict[str, Any]) -> Any:
    """按模型面 dict 构建 openpyxl DataValidation。非法输入抛 ValueError。

    模型面字段：type / operator / values|formula1|formula2（min|max、value|value2）
    / allow_blank / show_dropdown / error_style / prompt_title|prompt
    / error_title|error / show_input_message / show_error_message。
    """
    from openpyxl.worksheet.datavalidation import DataValidation

    if not isinstance(rule_spec, dict):
        raise ValueError("data_validation 需要 rule 对象")
    rtype = _DV_TYPE_MAP.get(str(rule_spec.get("type") or "list").strip().lower())
    if rtype is None:
        raise ValueError(
            f"data_validation.rule.type={rule_spec.get('type')!r} 不支持；可用 {_DV_TYPES}"
        )

    kwargs: dict[str, Any] = {
        "type": rtype,
        "allow_blank": bool(rule_spec.get("allow_blank", rule_spec.get("allowBlank", False))),
    }

    if rtype == "list":
        values = rule_spec.get("values")
        formula1 = rule_spec.get("formula1", rule_spec.get("source"))
        if isinstance(values, list) and values and formula1 not in (None, ""):
            raise ValueError("data_validation type=list 的 values 与 formula1 不能同时指定")
        if isinstance(values, list) and values:
            formula1 = '"' + ",".join(str(v) for v in values) + '"'
        elif isinstance(values, str) and values.strip() and not formula1:
            formula1 = values
        if formula1 is None or str(formula1).strip() == "":
            raise ValueError(
                "type=list 需要 values 数组或 formula1"
                '（内联形如 "a,b,c"，区域引用形如 Sheet!A1:A5）'
            )
        f1 = str(formula1).strip()
        if not f1.startswith('"') and not f1.startswith("="):
            f1 = "=" + f1
        kwargs["formula1"] = f1
    elif rtype == "custom":
        kwargs["formula1"] = "=" + _rule_formula(rule_spec)
    else:
        raw_op = str(rule_spec.get("operator") or "between").strip()
        operator = _normalize_operator_name(raw_op) or _DV_OPERATOR_MAP.get(raw_op.lower())
        if not operator:
            raise ValueError(
                f"data_validation.rule.operator={raw_op!r} 不支持；"
                f"可用 {sorted(set(_DV_OPERATOR_MAP.values()))}"
            )
        formula1 = rule_spec.get("formula1", rule_spec.get("value", rule_spec.get("min")))
        formula2 = rule_spec.get("formula2", rule_spec.get("value2", rule_spec.get("max")))
        if operator in {"between", "notBetween"}:
            if formula1 is None or formula2 is None:
                raise ValueError(
                    f"operator={operator} 需要上下界：value/min 与 value2/max（或 formula1/formula2）"
                )
            kwargs["formula1"] = _dv_formula_value(formula1, dv_type=rtype)
            kwargs["formula2"] = _dv_formula_value(formula2, dv_type=rtype)
        else:
            if formula1 is None:
                raise ValueError(f"operator={operator} 需要 value/formula1")
            kwargs["formula1"] = _dv_formula_value(formula1, dv_type=rtype)
        kwargs["operator"] = operator

    show_dropdown = rule_spec.get("show_dropdown", rule_spec.get("showDropDown"))
    if show_dropdown is not None:
        # OOXML showDropDown=1 是"隐藏下拉箭头"的遗留语义；模型面 show_dropdown
        # 保持直觉语义（True=显示下拉），这里取反。
        kwargs["showDropDown"] = not bool(show_dropdown)

    error_style = rule_spec.get("error_style", rule_spec.get("errorStyle"))
    if error_style is not None:
        kwargs["errorStyle"] = _DV_ERROR_STYLE_MAP.get(str(error_style).strip().lower(), "stop")
    for spec_key, attr in (
        ("prompt_title", "promptTitle"),
        ("prompt", "prompt"),
        ("error_title", "errorTitle"),
        ("error", "error"),
    ):
        if rule_spec.get(spec_key) is not None:
            kwargs[attr] = str(rule_spec[spec_key])
    if rule_spec.get("show_input_message") is not None:
        kwargs["showInputMessage"] = bool(rule_spec["show_input_message"])
    elif kwargs.get("prompt") or kwargs.get("promptTitle"):
        kwargs["showInputMessage"] = True
    if rule_spec.get("show_error_message") is not None:
        kwargs["showErrorMessage"] = bool(rule_spec["show_error_message"])

    return DataValidation(**kwargs)


# ── 样式提取辅助函数──────────────


def _color_to_hex(color: Any) -> str | None:
    """将 openpyxl Color 对象转换为十六进制字符串。"""
    if color is None:
        return None
    if hasattr(color, "rgb") and color.rgb and color.rgb != "00000000":
        rgb = str(color.rgb)
        # openpyxl 的 rgb 可能是 AARRGGBB 格式
        if len(rgb) == 8:
            return rgb[2:]  # 去掉 alpha 通道
        return rgb
    if hasattr(color, "theme") and color.theme is not None:
        return f"theme:{color.theme}"
    if hasattr(color, "indexed") and color.indexed is not None:
        return f"indexed:{color.indexed}"
    return None


def _extract_font(font: Font | None) -> dict[str, Any] | None:
    """从 openpyxl Font 提取非默认属性字典。"""
    if font is None:
        return None
    info: dict[str, Any] = {}
    if font.name and font.name != "Calibri":
        info["name"] = font.name
    if font.size and font.size != 11:
        info["size"] = font.size
    if font.bold:
        info["bold"] = True
    if font.italic:
        info["italic"] = True
    if font.underline and font.underline != "none":
        info["underline"] = font.underline
    if font.strike:
        info["strikethrough"] = True
    color_hex = _color_to_hex(font.color)
    if color_hex and color_hex != "000000":
        info["color"] = color_hex
    return info or None


def _extract_fill(fill: PatternFill | None) -> dict[str, Any] | None:
    """从 openpyxl PatternFill 提取非默认属性字典。"""
    if fill is None:
        return None
    fill_type = fill.fill_type or fill.patternType
    if not fill_type or fill_type == "none":
        return None
    info: dict[str, Any] = {"type": fill_type}
    color_hex = _color_to_hex(fill.fgColor)
    if color_hex:
        info["color"] = color_hex
    return info


def _extract_border(border: Border | None) -> dict[str, Any] | None:
    """从 openpyxl Border 提取非默认属性字典。"""
    if border is None:
        return None
    info: dict[str, Any] = {}
    for side_name in ("left", "right", "top", "bottom"):
        side: Side = getattr(border, side_name, None)
        if side and side.style and side.style != "none":
            info[side_name] = side.style
    return info or None


def _extract_alignment(alignment: Alignment | None) -> dict[str, Any] | None:
    """从 openpyxl Alignment 提取非默认属性字典。"""
    if alignment is None:
        return None
    info: dict[str, Any] = {}
    if alignment.horizontal and alignment.horizontal != "general":
        info["horizontal"] = alignment.horizontal
    if alignment.vertical and alignment.vertical != "bottom":
        info["vertical"] = alignment.vertical
    if alignment.wrap_text:
        info["wrap_text"] = True
    return info or None
