"""Bench 结果断言：检查 agent 交付物（工作簿 / docx / 回复）是否符合任务要求。

由 ``bench_validator.validate_case`` 调用。断言写在 suite/case 的 ``assertions`` 里：

    "answers_file": "bench/fixtures/realistic/answers.json",   # 可选，标准答案
    "uploads_unchanged": true,                                  # uploads/ 原件不得被改
    "output_checks": [ {"type": "...", ...}, ... ]

任何检查里的值都可以写成 ``"@answers:key.sub"`` 从 answers_file 取。

output_checks 类型：
- file_exists      {glob, min=1, max?, min_rows?}   min_rows: 有 ≥min_rows 个数据行的文件才计入 min
- sheet_exists     {file?, name_regex}
- cell_value       {file?, sheet?, cell, equals | approx, tolerance=0.5, rel_tolerance=0}
- formula          {file?, sheet?, range, regex, min=1}        range 可为 N2 / N2:N10 / N:N
- row_count        {file?, sheet?, header_rows=1, equals | min | max}
- matrix_contains  {file?, sheet?, row_label, value, tolerance, rel_tolerance, label_mode=exact|contains}
- cell_exists      {file?, sheet?, value}                        某值单独占一格
- column_values    {file?, sheet?, column, unique_count | all_match_regex | no_blank, header_rows=1}
                   column 可为序号 / 表头名 / "~正则" / 列字母
- sheet_props      {file?, sheet?, freeze_panes?, header_bold?, number_format?{column, contains},
                    conditional_formatting_min?, column_width_min?{col: n}, print_area?, print_ready?}
- charts           {file?, min=1, types?[line|bar|pie|area|scatter], min_series?}  min_series: 每张图至少绑几个数据系列
- docx             {file?, no_placeholder?, contains?[], contains_any?[], contains_number?}
- reply_number     {value | any_of, tolerance=0.5, rel_tolerance=0.005}
- reply_regex      {pattern}
- reply_contains_any {values[]}
- script           {path}   模块需定义 check(workdir, result, answers, spec) -> list[dict(name, passed, message)]

``file`` 为相对工作区的 glob；省略时在 outputs/ 下找最近修改的工作簿，再退到工作区根。
uploads/、.excelmanus/、.tmp/ 永远不当作产出。
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

logger = logging.getLogger(__name__)

_EXCLUDED_DIRS = ("uploads", ".excelmanus", ".tmp", "__pycache__")
_WORKBOOK_SUFFIXES = (".xlsx", ".xlsm")
_NUMBER_RE = re.compile(r"(-?\d[\d,]*(?:\.\d+)?)\s*(万亿|亿|万|w|W|k|K)?")
_ANSWER_REF = "@answers:"


@dataclass
class CheckResult:
    name: str
    passed: bool
    message: str = ""
    expected: Any = None
    actual: Any = None
    severity: str = "error"  # "error" | "warning"


# ── 标准答案引用 ─────────────────────────────────────────────


def load_answers(path: str | Path | None) -> dict[str, Any]:
    if not path:
        return {}
    p = Path(path)
    if not p.is_absolute():
        for candidate in (p, Path.cwd() / p):
            if candidate.exists():
                p = candidate
                break
    if not p.exists():
        logger.warning("answers_file 不存在: %s", path)
        return {}
    return json.loads(p.read_text(encoding="utf-8"))


def resolve_value(value: Any, answers: dict[str, Any]) -> Any:
    """把 "@answers:a.b.c" 递归替换成 answers 中的值；其它值原样返回。"""
    if isinstance(value, str) and value.startswith(_ANSWER_REF):
        node: Any = answers
        for part in value[len(_ANSWER_REF):].split("."):
            if isinstance(node, dict):
                node = node.get(part)
            elif isinstance(node, list) and part.isdigit():
                node = node[int(part)]
            else:
                node = None
            if node is None:
                raise KeyError(f"answers 里没有 {value!r}")
        return node
    if isinstance(value, dict):
        return {k: resolve_value(v, answers) for k, v in value.items()}
    if isinstance(value, list):
        return [resolve_value(v, answers) for v in value]
    return value


# ── 文件定位 ─────────────────────────────────────────────────


def _is_excluded(path: Path, root: Path) -> bool:
    try:
        parts = path.relative_to(root).parts
    except ValueError:
        return True
    return any(part in _EXCLUDED_DIRS for part in parts[:-1]) or path.name.startswith("~$")


def find_outputs(workdir: Path, pattern: str | None, *, suffixes: tuple[str, ...] = _WORKBOOK_SUFFIXES) -> list[Path]:
    """按 glob 找产出文件；无 glob 时先 outputs/ 再根目录，按修改时间新→旧。"""
    if pattern:
        found = [p for p in workdir.glob(pattern) if p.is_file() and not _is_excluded(p, workdir)]
    else:
        found = [p for p in workdir.glob("outputs/**/*") if p.is_file() and p.suffix.lower() in suffixes]
        found = [p for p in found if not _is_excluded(p, workdir)]
        if not found:
            found = [
                p for p in workdir.rglob("*")
                if p.is_file() and p.suffix.lower() in suffixes and not _is_excluded(p, workdir)
            ]
    return sorted(found, key=lambda p: p.stat().st_mtime, reverse=True)


def locate_output_workbook(workdir: Path, pattern: str | None = None) -> Path | None:
    found = find_outputs(workdir, pattern)
    return found[0] if found else None


class _WorkbookCache:
    """同一 case 内多条检查共用已加载的工作簿。"""

    def __init__(self) -> None:
        self._cache: dict[tuple[str, bool], Any] = {}

    def load(self, path: Path, *, data_only: bool) -> Any:
        from openpyxl import load_workbook

        key = (str(path), data_only)
        if key not in self._cache:
            self._cache[key] = load_workbook(path, data_only=data_only)
        return self._cache[key]

    def close(self) -> None:
        for wb in self._cache.values():
            try:
                wb.close()
            except Exception:
                pass
        self._cache.clear()


def _match_sheets(wb: Any, spec: Any) -> list[Any]:
    """sheet 省略 → 全部；"~regex" → 正则；否则精确名。"""
    if spec in (None, "", "*"):
        return list(wb.worksheets)
    if isinstance(spec, str) and spec.startswith("~"):
        rx = re.compile(spec[1:])
        return [ws for ws in wb.worksheets if rx.search(ws.title)]
    return [ws for ws in wb.worksheets if ws.title == spec]


def _approx(actual: Any, expected: Any, tolerance: float, rel_tolerance: float) -> bool:
    try:
        a = float(actual)
        e = float(expected)
    except (TypeError, ValueError):
        return str(actual).strip() == str(expected).strip()
    return abs(a - e) <= max(tolerance, abs(e) * rel_tolerance)


def _row_is_empty(row: tuple) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def _iter_range_cells(ws: Any, rng: str):
    from openpyxl.utils import column_index_from_string, range_boundaries

    rng = rng.strip()
    if re.fullmatch(r"[A-Za-z]{1,3}(:[A-Za-z]{1,3})?", rng):  # 整列 N 或 N:N
        first, _, last = rng.partition(":")
        c1 = column_index_from_string(first.upper())
        c2 = column_index_from_string((last or first).upper())
        for row in ws.iter_rows(min_col=c1, max_col=c2, min_row=1, max_row=ws.max_row):
            yield from row
        return
    min_col, min_row, max_col, max_row = range_boundaries(rng)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        yield from row


# ── 回复里的数字 ─────────────────────────────────────────────


def extract_numbers(text: str) -> list[float]:
    """抽出回复里的所有数字，中文单位（万/亿）展开；同时保留未展开值。"""
    out: list[float] = []
    for raw, unit in _NUMBER_RE.findall(text or ""):
        try:
            base = float(raw.replace(",", ""))
        except ValueError:
            continue
        out.append(base)
        mult = {"万": 1e4, "w": 1e4, "W": 1e4, "亿": 1e8, "万亿": 1e12, "k": 1e3, "K": 1e3}.get(unit)
        if mult:
            out.append(base * mult)
    return out


# ── 各类检查 ─────────────────────────────────────────────────


def _c_file_exists(ctx: _Ctx, spec: dict) -> CheckResult:
    pattern = spec.get("glob") or "outputs/**/*"
    files = [p for p in ctx.workdir.glob(pattern) if p.is_file() and not _is_excluded(p, ctx.workdir)]
    lo, hi = int(spec.get("min", 1)), spec.get("max")
    # min_rows：防"建空文件凑数"（R24 曾只计数不验内容）。达标文件数才计入 min。
    min_rows = spec.get("min_rows")
    qualified = len(files)
    detail = ""
    if min_rows is not None:
        need = int(min_rows)
        qualified = sum(1 for p in files if _file_data_rows(p) >= need)
        detail = f"，其中 ≥{need} 数据行的 {qualified} 个"
    ok = qualified >= lo and (hi is None or len(files) <= int(hi))
    return CheckResult("file_exists", ok, f"匹配 {pattern!r} 的文件 {len(files)} 个{detail}",
                       expected=f">={lo}" + (f", <={hi}" if hi is not None else "") + (f", min_rows={min_rows}" if min_rows is not None else ""),
                       actual=[str(p.relative_to(ctx.workdir)) for p in files[:10]])


def _file_data_rows(path: Path, header_rows: int = 1) -> int:
    """单个工作簿的数据行数（取各表最大值）；打不开/非工作簿记 0，不抛异常。"""
    if path.suffix.lower() not in _WORKBOOK_SUFFIXES:
        return 0
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return 0
    try:
        best = 0
        for ws in wb.worksheets:
            best = max(best, _sheet_row_count(ws, header_rows))
        return best
    except Exception:
        return 0
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _c_sheet_exists(ctx: _Ctx, spec: dict) -> CheckResult:
    wb, path = ctx.workbook(spec.get("file"), data_only=True)
    if wb is None:
        return CheckResult("sheet_exists", False, "未找到产出工作簿")
    rx = re.compile(spec["name_regex"])
    hit = [ws.title for ws in wb.worksheets if rx.search(ws.title)]
    return CheckResult("sheet_exists", bool(hit), f"{path.name}: 匹配 {spec['name_regex']!r} 的工作表 {hit or '无'}",
                       expected=spec["name_regex"], actual=wb.sheetnames)


def _c_cell_value(ctx: _Ctx, spec: dict) -> CheckResult:
    wb, path = ctx.workbook(spec.get("file"), data_only=True)
    if wb is None:
        return CheckResult("cell_value", False, "未找到产出工作簿")
    expected = spec.get("equals", spec.get("approx"))
    tol = float(spec.get("tolerance", 0.5 if "approx" in spec else 0))
    rel = float(spec.get("rel_tolerance", 0))
    or_formula = bool(spec.get("or_formula"))
    wbf = None
    seen: list[Any] = []
    for ws in _match_sheets(wb, spec.get("sheet")):
        actual = ws[spec["cell"]].value
        seen.append(f"{ws.title}!{spec['cell']}={actual!r}")
        if _approx(actual, expected, tol, rel):
            return CheckResult("cell_value", True, f"{path.name} {ws.title}!{spec['cell']} = {actual!r}", expected=expected, actual=actual)
        if or_formula and actual is None:
            if wbf is None:
                wbf, _ = ctx.workbook(spec.get("file"), data_only=False)
            raw = wbf[ws.title][spec["cell"]].value if wbf and ws.title in wbf.sheetnames else None
            if isinstance(raw, str) and raw.startswith("="):
                return CheckResult(
                    "cell_value", True,
                    f"{path.name} {ws.title}!{spec['cell']} 为未缓存公式 {raw[:60]!r}（值未验证）",
                    expected=expected, actual=raw,
                )
    return CheckResult("cell_value", False, f"{path.name}: {'; '.join(seen) or '无匹配工作表'} ≠ {expected!r}", expected=expected, actual=seen)


def _c_formula(ctx: _Ctx, spec: dict) -> CheckResult:
    wb, path = ctx.workbook(spec.get("file"), data_only=False)
    if wb is None:
        return CheckResult("formula", False, "未找到产出工作簿")
    rx = re.compile(spec["regex"], re.IGNORECASE)
    need = int(spec.get("min", 1))
    hits: list[str] = []
    for ws in _match_sheets(wb, spec.get("sheet")):
        for cell in _iter_range_cells(ws, spec["range"]):
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and rx.search(v):
                hits.append(f"{ws.title}!{cell.coordinate}: {v[:80]}")
                if len(hits) >= max(need, 5):
                    break
    ok = len(hits) >= need
    return CheckResult("formula", ok, f"{path.name}: 范围 {spec['range']} 内匹配 /{spec['regex']}/ 的公式 {len(hits)} 个（需 ≥{need}）",
                       expected=spec["regex"], actual=hits[:5])


def _sheet_row_count(ws: Any, header_rows: int) -> int:
    n = 0
    for row in ws.iter_rows(values_only=True):
        if not _row_is_empty(row):
            n += 1
    return max(0, n - header_rows)


def _c_row_count(ctx: _Ctx, spec: dict) -> CheckResult:
    wb, path = ctx.workbook(spec.get("file"), data_only=True)
    if wb is None:
        return CheckResult("row_count", False, "未找到产出工作簿")
    header = int(spec.get("header_rows", 1))
    sheets = _match_sheets(wb, spec.get("sheet"))
    if not sheets:
        return CheckResult("row_count", False, f"{path.name}: 没有匹配 {spec.get('sheet')!r} 的工作表", actual=wb.sheetnames)
    counts = {ws.title: _sheet_row_count(ws, header) for ws in sheets}
    # 未指定 sheet 时取数据最多的一张
    title, n = max(counts.items(), key=lambda kv: kv[1])
    ok = True
    if "equals" in spec:
        ok = n == int(spec["equals"])
    if "min" in spec:
        ok = ok and n >= int(spec["min"])
    if "max" in spec:
        ok = ok and n <= int(spec["max"])
    expect = {k: spec[k] for k in ("equals", "min", "max") if k in spec}
    return CheckResult("row_count", ok, f"{path.name} {title}: 数据行 {n}（期望 {expect}）", expected=expect, actual=counts)


def _c_matrix_contains(ctx: _Ctx, spec: dict) -> CheckResult:
    wb, path = ctx.workbook(spec.get("file"), data_only=True)
    if wb is None:
        return CheckResult("matrix_contains", False, "未找到产出工作簿")
    label = str(spec["row_label"]).strip()
    mode = spec.get("label_mode", "exact")
    value = spec["value"]
    tol = float(spec.get("tolerance", 0.5))
    rel = float(spec.get("rel_tolerance", 0))
    or_formula = bool(spec.get("or_formula"))
    wbf = None
    near: list[str] = []
    for ws in _match_sheets(wb, spec.get("sheet")):
        for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            labels = [str(v).strip() for v in row if v is not None]
            hit = label in labels if mode == "exact" else any(label in s for s in labels)
            if not hit:
                continue
            nums = [v for v in row if isinstance(v, (int, float)) and not isinstance(v, bool)]
            if any(_approx(v, value, tol, rel) for v in nums):
                return CheckResult("matrix_contains", True, f"{path.name} {ws.title}: 行 {label!r} 含 ≈{value}", expected=value, actual=nums)
            near.append(f"{ws.title}:{nums[:6]}")
            if or_formula:
                if wbf is None:
                    wbf, _ = ctx.workbook(spec.get("file"), data_only=False)
                if wbf is not None and ws.title in wbf.sheetnames:
                    wsf = wbf[ws.title]
                    formulas = [
                        c.value for c in wsf[ridx]
                        if isinstance(c.value, str) and c.value.startswith("=")
                    ]
                    if formulas:
                        return CheckResult(
                            "matrix_contains", True,
                            f"{path.name} {ws.title}: 行 {label!r} 第 {ridx} 行含未缓存公式"
                            f" {formulas[0][:60]!r}（值未验证）",
                            expected=value, actual=formulas[:3],
                        )
    return CheckResult("matrix_contains", False, f"{path.name}: 没有哪一行同时含 {label!r} 与 ≈{value}；候选行 {near[:4] or '无'}",
                       expected=value, actual=near[:4])


def _find_column(ws: Any, col: Any, header_row: int) -> int | None:
    """列定位：数字序号 / 表头名 / "~正则" / 列字母。"""
    from openpyxl.utils import column_index_from_string

    if isinstance(col, int) or (isinstance(col, str) and col.isdigit()):
        return int(col)
    text = str(col)
    rx = re.compile(text[1:]) if text.startswith("~") else None
    for c in range(1, ws.max_column + 1):
        head = str(ws.cell(header_row, c).value or "").strip()
        if (rx and rx.search(head)) or (not rx and head == text):
            return c
    if not rx and re.fullmatch(r"[A-Za-z]{1,3}", text):
        return column_index_from_string(text.upper())
    return None


def _c_row_contains(ctx: _Ctx, spec: dict) -> CheckResult | list[CheckResult]:
    """同一行内必须同时出现所有给定值（校验配对关系未被拆分打乱）。

    spec.values: 单个值数组，或值数组的数组（逐对校验，每对独立一条结果）。
    """
    wb, path = ctx.workbook(spec.get("file"), data_only=True)
    if wb is None:
        return CheckResult("row_contains", False, "未找到产出工作簿")
    values = spec["values"]
    groups = [values] if values and not isinstance(values[0], list) else values
    out: list[CheckResult] = []
    sheets = list(_match_sheets(wb, spec.get("sheet")))
    for group in groups:
        wants = [str(v).strip() for v in group]
        found_at = None
        for ws in sheets:
            for row in ws.iter_rows(values_only=True):
                cells = {str(v).strip() for v in row if v is not None}
                if all(w in cells for w in wants):
                    found_at = f"{ws.title}"
                    break
            if found_at:
                break
        if found_at:
            out.append(CheckResult("row_contains", True, f"{path.name} {found_at}: 同行含 {wants}", expected=wants))
        else:
            out.append(CheckResult("row_contains", False, f"{path.name}: 没有哪一行同时含 {wants}", expected=wants))
    return out


def _c_cell_exists(ctx: _Ctx, spec: dict) -> CheckResult:
    """某个值是否单独占一个单元格（例如拆列后工号应独立成格）。"""
    wb, path = ctx.workbook(spec.get("file"), data_only=True)
    if wb is None:
        return CheckResult("cell_exists", False, "未找到产出工作簿")
    want = str(spec["value"]).strip()
    for ws in _match_sheets(wb, spec.get("sheet")):
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and str(cell.value).strip() == want:
                    return CheckResult("cell_exists", True, f"{path.name} {ws.title}!{cell.coordinate} = {want!r}", expected=want)
    return CheckResult("cell_exists", False, f"{path.name}: 没有单元格恰好等于 {want!r}", expected=want)


def _c_column_values(ctx: _Ctx, spec: dict) -> CheckResult:
    wb, path = ctx.workbook(spec.get("file"), data_only=True)
    if wb is None:
        return CheckResult("column_values", False, "未找到产出工作簿")
    header = int(spec.get("header_rows", 1))
    col = spec["column"]
    problems: list[str] = []
    for ws in _match_sheets(wb, spec.get("sheet")):
        idx = _find_column(ws, col, header)
        if idx is None:
            problems.append(f"{ws.title}: 找不到列 {col!r}")
            continue
        values = [ws.cell(r, idx).value for r in range(header + 1, ws.max_row + 1)]
        values = [v for v in values if v is not None and str(v).strip() != ""]
        if "unique_count" in spec:
            n = len({str(v).strip() for v in values})
            if n == int(spec["unique_count"]):
                return CheckResult("column_values", True, f"{path.name} {ws.title} 列 {col}: 唯一值 {n}", expected=spec["unique_count"], actual=n)
            problems.append(f"{ws.title}: 唯一值 {n} ≠ {spec['unique_count']}")
        if "all_match_regex" in spec:
            rx = re.compile(spec["all_match_regex"])
            bad = [v for v in values if not rx.fullmatch(str(v).strip())]
            if not bad and values:
                return CheckResult("column_values", True, f"{path.name} {ws.title} 列 {col}: {len(values)} 个值全部匹配", expected=spec["all_match_regex"], actual=len(values))
            problems.append(f"{ws.title}: {len(bad)}/{len(values)} 不匹配，例如 {bad[:3]}")
        if spec.get("no_blank"):
            total = ws.max_row - header
            if len(values) == total and total > 0:
                return CheckResult("column_values", True, f"{path.name} {ws.title} 列 {col}: 无空值", actual=total)
            problems.append(f"{ws.title}: {total - len(values)} 个空值")
    return CheckResult("column_values", False, f"{path.name}: {'; '.join(problems) or '无匹配工作表'}", actual=problems)


def _c_sheet_props(ctx: _Ctx, spec: dict) -> CheckResult:
    from openpyxl.utils import column_index_from_string

    wb, path = ctx.workbook(spec.get("file"), data_only=False)
    if wb is None:
        return CheckResult("sheet_props", False, "未找到产出工作簿")
    sheets = _match_sheets(wb, spec.get("sheet"))
    if not sheets:
        return CheckResult("sheet_props", False, f"{path.name}: 没有匹配 {spec.get('sheet')!r} 的工作表", actual=wb.sheetnames)
    best_fail: list[str] = []
    for ws in sheets:
        fails: list[str] = []
        if "freeze_panes" in spec:
            want = spec["freeze_panes"]
            got = ws.freeze_panes
            if (want is True and not got) or (isinstance(want, str) and got != want):
                fails.append(f"freeze_panes={got!r}≠{want!r}")
        if spec.get("header_bold"):
            row = int(spec.get("header_row", 1))
            cells = [ws.cell(row, c) for c in range(1, min(ws.max_column, 30) + 1) if ws.cell(row, c).value is not None]
            if not cells or not all(bool(c.font and c.font.bold) for c in cells):
                fails.append("表头未全部加粗")
        nf = spec.get("number_format")
        if nf:
            col = nf["column"]
            idx = _find_column(ws, col, int(spec.get("header_row", 1)))
            if idx is None:
                fails.append(f"找不到列 {col!r}")
            else:
                sample = [ws.cell(r, idx).number_format for r in range(2, min(ws.max_row, 40) + 1) if ws.cell(r, idx).value is not None]
                if not sample or not all(nf["contains"] in fmt for fmt in sample):
                    fails.append(f"列 {col} number_format 样本 {sorted(set(sample))[:3]} 不含 {nf['contains']!r}")
        if "conditional_formatting_min" in spec:
            n = sum(len(rules) for rules in ws.conditional_formatting._cf_rules.values()) if hasattr(ws.conditional_formatting, "_cf_rules") else len(list(ws.conditional_formatting))
            if n < int(spec["conditional_formatting_min"]):
                fails.append(f"条件格式 {n} 条 < {spec['conditional_formatting_min']}")
        for col, minw in (spec.get("column_width_min") or {}).items():
            w = ws.column_dimensions[col].width
            if not w or w < float(minw):
                fails.append(f"列 {col} 宽 {w} < {minw}")
        if spec.get("print_area") and not ws.print_area:
            fails.append("未设置打印区域")
        if spec.get("print_ready"):
            fit = getattr(getattr(ws.sheet_properties, "pageSetUpPr", None), "fitToPage", None)
            setup = ws.page_setup
            ready = bool(ws.print_area) or bool(fit) or bool(setup.fitToWidth or setup.fitToHeight) or bool(setup.paperSize)
            if not ready:
                fails.append("没有打印区域 / 缩放到一页 / 纸张设置")
        if not fails:
            return CheckResult("sheet_props", True, f"{path.name} {ws.title}: 全部属性满足")
        best_fail = fails if not best_fail or len(fails) < len(best_fail) else best_fail
    return CheckResult("sheet_props", False, f"{path.name}: {'; '.join(best_fail)}", actual=best_fail)


_CHART_TAGS = {
    "line": ("lineChart", "line3DChart"),
    "bar": ("barChart", "bar3DChart"),
    "pie": ("pieChart", "pie3DChart", "doughnutChart"),
    "area": ("areaChart", "area3DChart"),
    "scatter": ("scatterChart",),
}


def _chart_kind(xml: str) -> str:
    """chart xml 的类型标签可能带命名空间前缀（<c:lineChart> 或裸 <lineChart>）。"""
    for kind, tags in _CHART_TAGS.items():
        for tag in tags:
            if re.search(rf"<(?:\w+:)?{tag}\b", xml):
                return kind
    return "other"


def inspect_charts(path: Path) -> list[str]:
    """读 xlsx 包里的 chart xml，返回每张图的类型（openpyxl 打开后会丢图表，故直接读 zip）。"""
    kinds: list[str] = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if not re.fullmatch(r"xl/charts/chart\d+\.xml", name):
                continue
            xml = zf.read(name).decode("utf-8", errors="ignore")
            kinds.append(_chart_kind(xml))
    return kinds


def inspect_chart_series(path: Path) -> list[int]:
    """每张图的数据系列数（<c:ser> 个数）。空图/未绑定数据的图为 0。"""
    counts: list[int] = []
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if not re.fullmatch(r"xl/charts/chart\d+\.xml", name):
                    continue
                xml = zf.read(name).decode("utf-8", errors="ignore")
                counts.append(len(re.findall(r"<(?:\w+:)?ser\b", xml)))
    except Exception:
        pass
    return counts


def _c_charts(ctx: _Ctx, spec: dict) -> CheckResult:
    path = locate_output_workbook(ctx.workdir, spec.get("file"))
    if path is None:
        return CheckResult("charts", False, "未找到产出工作簿")
    # 多个产出时取图最多的那个
    candidates = find_outputs(ctx.workdir, spec.get("file"))
    best_path, best = path, inspect_charts(path)
    for cand in candidates[1:6]:
        kinds = inspect_charts(cand)
        if len(kinds) > len(best):
            best_path, best = cand, kinds
    need = int(spec.get("min", 1))
    want_types = set(spec.get("types") or [])
    ok = len(best) >= need and want_types <= set(best)
    msg = f"{best_path.name}: 图表 {best}（需 ≥{need}，类型 {sorted(want_types) or '任意'}）"
    # min_series：空图（只插图框、未绑定数据区）也应判 fail
    need_series = int(spec.get("min_series", 0))
    series: list[int] = []
    if need_series > 0:
        series = inspect_chart_series(best_path)
        thin = [i + 1 for i, n in enumerate(series) if n < need_series]
        if thin:
            ok = False
            msg += f"；第 {thin} 张图数据系列 {series} < min_series={need_series}"
        else:
            msg += f"；数据系列 {series} 满足 min_series={need_series}"
    return CheckResult("charts", ok, msg, expected=spec, actual={"kinds": best, "series": series})


def _c_docx(ctx: _Ctx, spec: dict) -> CheckResult:
    from docx import Document

    files = find_outputs(ctx.workdir, spec.get("file"), suffixes=(".docx",))
    if not files:
        return CheckResult("docx", False, "未找到产出 docx")
    path = files[0]
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    full = "\n".join(texts)
    fails: list[str] = []
    if spec.get("no_placeholder") and re.search(r"\{\{.*?\}\}", full):
        placeholders = re.findall(r"\{\{.*?\}\}", full)[:5]
        fails.append(f"仍有占位符: {placeholders}")
    for kw in spec.get("contains") or []:
        if str(kw) not in full:
            fails.append(f"缺少 {kw!r}")
    any_kw = spec.get("contains_any") or []
    if any_kw and not any(str(kw) in full for kw in any_kw):
        fails.append(f"不含任一 {any_kw}")
    num = spec.get("contains_number")
    if num is not None:
        tol = float(spec.get("tolerance", 0.5))
        rel = float(spec.get("rel_tolerance", 0.005))
        if not any(_approx(n, num, tol, rel) for n in extract_numbers(full)):
            fails.append(f"正文里没有 ≈{num} 的数字")
    return CheckResult("docx", not fails, f"{path.name}: {'; '.join(fails) or 'OK'} （{len(full)} 字）", actual=fails)


def _c_reply_number(ctx: _Ctx, spec: dict) -> CheckResult:
    reply = ctx.reply
    candidates = spec.get("any_of") or [spec["value"]]
    tol = float(spec.get("tolerance", 0.5))
    rel = float(spec.get("rel_tolerance", 0.005))
    nums = extract_numbers(reply)
    for want in candidates:
        if any(_approx(n, want, tol, rel) for n in nums):
            return CheckResult("reply_number", True, f"回复含 ≈{want}", expected=candidates, actual=want)
    return CheckResult("reply_number", False, f"回复里没有 ≈{candidates} 的数字；抽到 {sorted(set(nums))[-8:]}", expected=candidates, actual=sorted(set(nums))[-12:])


def _c_reply_regex(ctx: _Ctx, spec: dict) -> CheckResult:
    ok = re.search(spec["pattern"], ctx.reply or "", re.IGNORECASE | re.DOTALL) is not None
    return CheckResult("reply_regex", ok, f"回复{'匹配' if ok else '不匹配'} /{spec['pattern']}/", expected=spec["pattern"], actual=(ctx.reply or "")[:200])


def _c_reply_contains_any(ctx: _Ctx, spec: dict) -> CheckResult:
    values = [str(v) for v in spec.get("values") or []]
    hit = [v for v in values if v in (ctx.reply or "")]
    need = int(spec.get("min", 1))
    return CheckResult("reply_contains_any", len(hit) >= need, f"回复命中 {hit}（需 ≥{need}）", expected=values, actual=hit)


def _c_script(ctx: _Ctx, spec: dict) -> CheckResult | list[CheckResult]:
    path = Path(spec["path"])
    if not path.is_absolute():
        path = Path.cwd() / path
    module_spec = importlib.util.spec_from_file_location(f"bench_check_{path.stem}", path)
    if module_spec is None or module_spec.loader is None:
        return CheckResult("script", False, f"无法加载 {path}")
    module = importlib.util.module_from_spec(module_spec)
    module_spec.loader.exec_module(module)
    raw = module.check(ctx.workdir, ctx.result, ctx.answers, spec)
    out: list[CheckResult] = []
    for item in raw or []:
        if isinstance(item, dict):
            out.append(CheckResult(str(item.get("name", path.stem)), bool(item.get("passed")), str(item.get("message", "")),
                                   expected=item.get("expected"), actual=item.get("actual")))
        else:
            name, passed, message = item
            out.append(CheckResult(str(name), bool(passed), str(message)))
    return out


_CHECKS: dict[str, Callable[["_Ctx", dict], CheckResult | list[CheckResult]]] = {
    "file_exists": _c_file_exists,
    "sheet_exists": _c_sheet_exists,
    "cell_value": _c_cell_value,
    "formula": _c_formula,
    "row_count": _c_row_count,
    "matrix_contains": _c_matrix_contains,
    "row_contains": _c_row_contains,
    "cell_exists": _c_cell_exists,
    "column_values": _c_column_values,
    "sheet_props": _c_sheet_props,
    "charts": _c_charts,
    "docx": _c_docx,
    "reply_number": _c_reply_number,
    "reply_regex": _c_reply_regex,
    "reply_contains_any": _c_reply_contains_any,
    "script": _c_script,
}


class _Ctx:
    def __init__(self, workdir: Path, result: dict[str, Any], answers: dict[str, Any]) -> None:
        self.workdir = workdir
        self.result = result
        self.answers = answers
        self.reply = str((result.get("result") or {}).get("reply") or "")
        self._books = _WorkbookCache()

    def workbook(self, pattern: str | None, *, data_only: bool) -> tuple[Any | None, Path | None]:
        path = locate_output_workbook(self.workdir, pattern)
        if path is None:
            return None, None
        return self._books.load(path, data_only=data_only), path

    def close(self) -> None:
        self._books.close()


def run_output_checks(
    workdir: Path | None,
    result: dict[str, Any],
    checks: list[dict[str, Any]],
    answers: dict[str, Any],
) -> list[CheckResult]:
    if workdir is None:
        return [CheckResult("output_checks", False, "workfile_dir 未提供，无法检查产出")]
    ctx = _Ctx(workdir, result, answers)
    out: list[CheckResult] = []
    try:
        for i, raw in enumerate(checks):
            try:
                spec = resolve_value(dict(raw), answers)
                kind = str(spec.get("type") or "")
                fn = _CHECKS.get(kind)
                if fn is None:
                    out.append(CheckResult(f"output_checks[{i}]", False, f"未知检查类型 {kind!r}"))
                    continue
                res = fn(ctx, spec)
                items = res if isinstance(res, list) else [res]
                label = spec.get("label")
                spec_severity = spec.get("severity")
                for item in items:
                    if label:
                        item.name = f"{item.name}:{label}"
                    if spec_severity:
                        item.severity = str(spec_severity)
                    out.append(item)
            except Exception as exc:  # 单条检查出错不拖垮其余
                out.append(CheckResult(f"output_checks[{i}]:{raw.get('type')}", False, f"检查执行异常: {type(exc).__name__}: {exc}"))
    finally:
        ctx.close()
    return out


# ── uploads 原件未被改动 ─────────────────────────────────────


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def check_uploads_unchanged(workdir: Path | None, result: dict[str, Any]) -> CheckResult:
    """对照 fake_frontend 记录的附件来源，确认 uploads/ 里的副本字节未变。"""
    if workdir is None:
        return CheckResult("uploads_unchanged", False, "workfile_dir 未提供")
    attachments = ((result.get("artifacts") or {}).get("conversation_export") or {}).get("attachments") or []
    if not attachments:
        return CheckResult("uploads_unchanged", True, "无附件，跳过")
    changed: list[str] = []
    missing: list[str] = []
    for att in attachments:
        src = att.get("source")
        rel = str(att.get("path") or "").lstrip("./")
        if not src or not rel:
            continue
        dst = workdir / rel
        src_path = Path(src)
        if not dst.exists():
            missing.append(rel)
            continue
        if not src_path.exists():
            continue
        # .xls 上传后转成 .xlsx，字节必然不同，跳过
        if src_path.suffix.lower() != dst.suffix.lower():
            continue
        if _sha256(src_path) != _sha256(dst):
            changed.append(rel)
    ok = not changed and not missing
    msg = []
    if changed:
        msg.append(f"被改动: {changed}")
    if missing:
        msg.append(f"被删除/移动: {missing}")
    return CheckResult("uploads_unchanged", ok, "; ".join(msg) or f"{len(attachments)} 个附件原件未变", actual={"changed": changed, "missing": missing})
