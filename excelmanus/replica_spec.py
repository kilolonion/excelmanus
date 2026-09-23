"""ReplicaSpec / WorkbookSpec 数据协议：图片→Excel 复刻的结构化中间格式。"""

from __future__ import annotations

import json
import re
from datetime import date, datetime, time, timezone
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator, model_validator

from excelmanus.workbook.layout import PrintLayout, apply_print_layout

_EXCEL_MAX_ROWS = 1_048_576
_EXCEL_MAX_COLS = 16_384


def _strip_a1_field(value: Any) -> Any:
    if isinstance(value, str):
        from excelmanus.workbook.address import strip_sheet_qualifier

        return strip_sheet_qualifier(value)
    return value


class Provenance(BaseModel):
    source_image_hash: str
    model: str
    timestamp: str
    extraction_params: dict[str, Any] = Field(default_factory=dict)


_LEGAL_FILL_TYPES = frozenset(
    {
        "solid",
        "none",
        "darkDown",
        "darkGray",
        "darkGrid",
        "darkHorizontal",
        "darkTrellis",
        "darkUp",
        "darkVertical",
        "gray0625",
        "gray125",
        "lightDown",
        "lightGray",
        "lightGrid",
        "lightHorizontal",
        "lightTrellis",
        "lightUp",
        "lightVertical",
        "mediumGray",
    }
)


class FontSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str | None = None
    size: float | None = None
    bold: bool | None = None
    italic: bool | None = None
    color: str | None = None  # 十六进制颜色值，如 "#FF0000"
    underline: str | None = None
    strike: bool | None = None

    @model_validator(mode="before")
    @classmethod
    def _coerce_font_name(cls, data: Any) -> Any:
        if isinstance(data, str) and data.strip():
            return {"name": data.strip()}
        return data


class FillSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")
    type: str = "solid"
    color: str | None = None

    @model_validator(mode="before")
    @classmethod
    def _coerce_fill_aliases(cls, data: Any) -> Any:
        if not isinstance(data, dict):
            return data
        payload = dict(data)
        alias = payload.pop("pattern", None) or payload.pop("fill_type", None) or payload.pop("patternType", None)
        if payload.get("type") in (None, ""):
            payload["type"] = alias or "solid"
        if payload.get("color") in (None, ""):
            for key in ("fgColor", "start_color", "fg_color"):
                if payload.get(key) not in (None, ""):
                    payload["color"] = payload[key]
                    break
        for key in ("fgColor", "start_color", "fg_color"):
            payload.pop(key, None)
        return payload

    @field_validator("type")
    @classmethod
    def _reject_generic_pattern(cls, value: str) -> str:
        if value == "pattern":
            raise ValueError("fill.type=pattern 不是合法 patternType，请用 solid")
        if value not in _LEGAL_FILL_TYPES:
            raise ValueError(f"fill.type={value!r} 不是合法 patternType")
        return value


class BorderSideSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")
    style: str | None = None
    color: str | None = None


class BorderSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")
    style: str | None = None
    color: str | None = None
    top: BorderSideSpec | None = None
    bottom: BorderSideSpec | None = None
    left: BorderSideSpec | None = None
    right: BorderSideSpec | None = None

    @field_validator("top", "bottom", "left", "right", mode="before")
    @classmethod
    def _coerce_side(cls, value: Any) -> Any:
        if isinstance(value, str) and value.strip():
            return {"style": value.strip()}
        return value


class AlignmentSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")
    horizontal: str | None = None
    vertical: str | None = None
    wrap_text: bool | None = None

    @model_validator(mode="before")
    @classmethod
    def _coerce_wrap_text(cls, data: Any) -> Any:
        if not isinstance(data, dict):
            return data
        payload = dict(data)
        if "wrapText" in payload and payload.get("wrap_text") is None:
            payload["wrap_text"] = payload.pop("wrapText")
        else:
            payload.pop("wrapText", None)
        return payload


class StyleClass(BaseModel):
    model_config = ConfigDict(extra="forbid")
    font: FontSpec | None = None
    fill: FillSpec | None = None
    border: BorderSpec | None = None
    alignment: AlignmentSpec | None = None
    number_format: str | None = None


class CellSpec(BaseModel):
    address: str
    value: Any = None
    value_type: Literal["string", "number", "date", "boolean", "formula", "empty"] = "string"
    display_text: str | None = None
    number_format: str | None = None
    formula_candidate: str | None = None
    style_id: str | None = None
    confidence: float = 1.0

    @model_validator(mode="before")
    @classmethod
    def _coerce_formula_and_address(cls, data: Any) -> Any:
        if not isinstance(data, dict):
            return data
        payload = dict(data)
        from excelmanus.workbook.address import strip_sheet_qualifier

        if isinstance(payload.get("address"), str):
            payload["address"] = strip_sheet_qualifier(payload["address"])
        formula = payload.pop("formula", None)
        if formula not in (None, "") and payload.get("value") in (None, ""):
            payload["value"] = formula
            payload["value_type"] = "formula"
        if payload.get("formula_candidate") not in (None, "") and payload.get("value") in (None, ""):
            payload["value"] = payload["formula_candidate"]
            payload["value_type"] = "formula"
        return payload


class MergedRange(BaseModel):
    range: str
    confidence: float = 1.0

    @model_validator(mode="before")
    @classmethod
    def _coerce_range_string(cls, data: Any) -> Any:
        if isinstance(data, str):
            data = {"range": data}
        if isinstance(data, dict) and isinstance(data.get("range"), str):
            from excelmanus.workbook.address import strip_sheet_qualifier

            payload = dict(data)
            payload["range"] = strip_sheet_qualifier(payload["range"])
            return payload
        return data


class FormulaPattern(BaseModel):
    column: str
    pattern: str
    confidence: float = 0.5


class SemanticHints(BaseModel):
    header_rows: list[int] = Field(default_factory=list)
    total_rows: list[int] = Field(default_factory=list)
    grouping: Any | None = None
    formula_patterns: list[FormulaPattern] = Field(default_factory=list)


class ConditionalFormatRule(BaseModel):
    """条件格式规则（颜色刻度/数据条/图标集/单元格值）。字段名是 type，不是 kind。"""

    type: Literal["color_scale", "data_bar", "icon_set", "cell_value"] = Field(
        description="WorkbookSpec 条件格式类型；与 format_spreadsheet 的 rule.type 不同名空间",
    )
    range: str = Field(description="Excel 范围，如 C2:C9")
    min_color: str | None = None
    mid_color: str | None = None
    max_color: str | None = None
    bar_color: str | None = None
    icon_style: str | None = Field(default=None, description="如 3_arrows / 3_traffic_lights")
    operator: str | None = Field(
        default=None,
        description="cell_value：greater_than/less_than/between/equal/not_equal",
    )
    value: Any = None
    value2: Any = None
    font_color: str | None = None
    fill_color: str | None = None
    bold: bool | None = None
    confidence: float = 0.7

    @field_validator("range", mode="before")
    @classmethod
    def _strip_range(cls, value: Any) -> Any:
        return _strip_a1_field(value)


class ObjectsSpec(BaseModel):
    charts: list[Any] = Field(default_factory=list)
    images: list[Any] = Field(default_factory=list)
    shapes: list[Any] = Field(default_factory=list)


class ValueBlock(BaseModel):
    """矩形值块：从 start 锚点展开的二维网格。"""

    start: str = Field(description="左上角 A1，如 A1")
    values: list[list[Any]] = Field(default_factory=list, description="非空矩形二维数组，行等长")

    @field_validator("start", mode="before")
    @classmethod
    def _strip_start(cls, value: Any) -> Any:
        return _strip_a1_field(value)


class FormulaBlock(BaseModel):
    """矩形公式块：从 start 锚点展开的二维公式网格。"""

    start: str = Field(description="左上角 A1")
    formulas: list[list[str]] = Field(default_factory=list, description="非空矩形二维公式数组，单元格以 = 开头")

    @field_validator("start", mode="before")
    @classmethod
    def _strip_start(cls, value: Any) -> Any:
        return _strip_a1_field(value)


class StyleRegion(BaseModel):
    """将具名样式应用到矩形区域。"""

    range: str = Field(description="A1 区域，须落在 dimensions 内")
    style_id: str = Field(description="必须是 styles 中已有的键")

    @field_validator("range", mode="before")
    @classmethod
    def _strip_range(cls, value: Any) -> Any:
        return _strip_a1_field(value)


class SheetSpec(BaseModel):
    name: str = Field(description="工作表名")
    dimensions: dict[str, int] | None = Field(
        default=None,
        description="规范为 {rows, cols}；带 source_csv 时可省略",
    )
    source_csv: dict[str, Any] | None = Field(
        default=None,
        description="{file_path, encoding?, skip_rows?, start?} 导入工作区 CSV/TSV",
    )
    freeze_panes: str | None = None
    print_layout: PrintLayout | None = None
    cells: list[CellSpec] = Field(default_factory=list, description="单格；公式用 value='=A1' 且 value_type=formula")
    value_blocks: list[ValueBlock] = Field(default_factory=list)
    formula_blocks: list[FormulaBlock] = Field(default_factory=list)
    style_regions: list[StyleRegion] = Field(default_factory=list)
    merged_ranges: list[MergedRange] = Field(default_factory=list)
    styles: dict[str, StyleClass] = Field(
        default_factory=dict,
        description="style_id → 样式；查询 additionalProperties 得 font/fill/border",
    )
    column_widths: list[float] = Field(
        default_factory=list,
        description="规范为从 A 起的数字数组；也接受 {\"A\":18}",
    )
    row_heights: dict[str, float] = Field(
        default_factory=dict,
        description="行号字符串 → 行高；也接受数组",
    )
    conditional_formats: list[ConditionalFormatRule] = Field(default_factory=list)
    objects: ObjectsSpec = Field(default_factory=ObjectsSpec)
    semantic_hints: SemanticHints = Field(default_factory=SemanticHints)

    @field_validator("freeze_panes", mode="before")
    @classmethod
    def _strip_freeze(cls, value: Any) -> Any:
        return _strip_a1_field(value)

    @field_validator("column_widths", mode="before")
    @classmethod
    def _coerce_column_widths(cls, value: Any) -> Any:
        if value is None:
            return []
        if isinstance(value, dict):
            from excelmanus.workbook.address import column_map_to_list

            return column_map_to_list(value)
        return value

    @field_validator("row_heights", mode="before")
    @classmethod
    def _coerce_row_heights(cls, value: Any) -> Any:
        if value is None:
            return {}
        if isinstance(value, list):
            heights: dict[str, float] = {}
            for index, height in enumerate(value):
                try:
                    heights[str(index + 1)] = float(height)
                except (TypeError, ValueError):
                    continue
            return heights
        if isinstance(value, dict):
            return {str(key): item for key, item in value.items()}
        return value

    @field_validator("dimensions", mode="before")
    @classmethod
    def _coerce_dimensions(cls, value: Any) -> Any:
        if isinstance(value, (list, tuple)) and len(value) >= 2:
            return {"rows": value[0], "cols": value[1]}
        if not isinstance(value, dict):
            return value
        payload = dict(value)
        if "rows" not in payload and "row" in payload:
            payload["rows"] = payload["row"]
        if "cols" not in payload:
            if "col" in payload:
                payload["cols"] = payload["col"]
            elif "columns" in payload:
                payload["cols"] = payload["columns"]
        return payload


class WorkbookMeta(BaseModel):
    """ReplicaSpec 内嵌的工作簿元数据（名称/字体）。"""

    name: str = "replica"
    locale: str | None = None
    default_font: FontSpec | None = None
    theme_hint: str | None = None


class Uncertainty(BaseModel):
    location: str = Field(description="不确定项位置，如 Sheet1!B2")
    reason: str = Field(description="看不清或无法确定的原因")
    candidate_values: list[str] = Field(
        default_factory=list,
        description="候选字符串列表，数字写成 \"4800\"",
    )
    confidence: float = 0.5

    @field_validator("candidate_values", mode="before")
    @classmethod
    def _stringify_candidate_values(cls, value: Any) -> Any:
        if not isinstance(value, list):
            return value
        items: list[str] = []
        for item in value:
            if item is None:
                continue
            items.append(str(item))
        return items


class WorkbookSpec(BaseModel):
    """P5 最小工作簿规格：当前模型直接产出，经校验后编译。

    ``uncertainties`` 必填（可为空列表，但不能缺字段）。
    """

    version: str = "1.0"
    name: str = "replica"
    locale: str | None = None
    default_font: FontSpec | None = Field(
        default=None,
        description="规范为 {name, size?}；也接受字符串字体名",
    )
    theme_hint: str | None = None
    sheets: list[SheetSpec] = Field(description="工作表列表")
    uncertainties: list[Uncertainty] = Field(description="必须出现；没有不确定项时为 []")

    @model_validator(mode="before")
    @classmethod
    def _lift_nested_workbook(cls, data: Any) -> Any:
        if not isinstance(data, dict):
            return data
        payload = dict(data)
        nested = payload.get("workbook")
        if isinstance(nested, dict):
            for key in ("name", "locale", "default_font", "theme_hint"):
                if payload.get(key) in (None, "") and nested.get(key) not in (None, ""):
                    payload[key] = nested[key]
        return payload


class ReplicaSpec(BaseModel):
    version: str = "1.0"
    provenance: Provenance
    workbook: WorkbookMeta = Field(default_factory=WorkbookMeta)
    sheets: list[SheetSpec]
    uncertainties: list[Uncertainty] = Field(default_factory=list)


class SpecValidationError(ValueError):
    """带字段路径的 WorkbookSpec 校验失败。"""

    def __init__(self, errors: list[dict[str, str]]):
        self.errors = errors
        summary = "; ".join(f"{item['path']}: {item['message']}" for item in errors) or "WorkbookSpec 校验失败"
        super().__init__(summary)

    def to_payload(self) -> dict[str, Any]:
        from excelmanus.engine_core.error_payload import SPEC_VALIDATION_FAILED, make_error_payload

        return make_error_payload(
            "WorkbookSpec 校验失败，请按字段路径修正后重试。",
            error_code=SPEC_VALIDATION_FAILED,
            errors=self.errors,
            shape="WorkbookSpec",
        )


def format_error_path(loc: tuple[Any, ...]) -> str:
    if not loc:
        return "$"
    return ".".join(str(part) for part in loc)


def _friendly_spec_message(path: str, message: str) -> str:
    lower = message.lower()
    if "default_font" in path and ("dict" in lower or "font" in lower):
        return 'default_font 必须是 {"name":"微软雅黑","size":11}，或直接传字体名字符串'
    if path.endswith("column_widths") and "list" in lower:
        return 'column_widths 必须是从 A 列起的数字数组 [18,12]，或 {"A":18,"B":12}'
    if path.endswith("row_heights") and "dict" in lower:
        return 'row_heights 必须是 {"1":22} 或从第 1 行起的数字数组 [22,15]'
    if path.endswith("merged_ranges") or ".merged_ranges." in path:
        if "dict" in lower:
            return 'merged_ranges 项必须是 {range:"A1:B1"} 或 "A1:B1"'
    if path == "sheets" and "list" in lower:
        return 'sheets 必须是数组，例如 [{"name":"Sheet1","dimensions":{"rows":10,"cols":5}}]'
    if path.endswith("uncertainties") and ("required" in lower or "field required" in lower):
        return "uncertainties 必填；没有不确定项时传 []"
    return message


def _pydantic_errors(exc: ValidationError) -> list[dict[str, str]]:
    return [
        {
            "path": format_error_path(tuple(err.get("loc") or ())),
            "message": _friendly_spec_message(
                format_error_path(tuple(err.get("loc") or ())),
                str(err.get("msg") or "校验失败"),
            ),
        }
        for err in exc.errors()
    ]


def _parse_a1(address: str) -> tuple[int, int]:
    """返回 1-indexed (row, col)。"""
    from excelmanus.workbook.refs import CellRef, RectRef, parse_ref

    area = parse_ref(address)
    part = area.areas[0]
    if isinstance(part, CellRef):
        return part.row, part.col
    if isinstance(part, RectRef):
        if part.whole_column or part.whole_row:
            raise ValueError("WorkbookSpec 不支持整轴地址，请写有限单元格如 A1")
        return part.min_row, part.min_col
    raise ValueError(f"WorkbookSpec 地址必须是单元格或矩形：{address!r}")


def _range_bounds(range_str: str) -> tuple[int, int, int, int]:
    """返回 (min_col, min_row, max_col, max_row)，均为 1-indexed。"""
    from excelmanus.workbook.refs import CellRef, RectRef, parse_ref

    area = parse_ref(range_str)
    if len(area.areas) != 1:
        raise ValueError("WorkbookSpec 不支持并集 range")
    part = area.areas[0]
    if isinstance(part, CellRef):
        return part.col, part.row, part.col, part.row
    if isinstance(part, RectRef):
        if part.whole_column or part.whole_row:
            raise ValueError("WorkbookSpec 不支持整轴 range，请写有限矩形")
        return part.min_col, part.min_row, part.max_col, part.max_row
    raise ValueError(f"WorkbookSpec range 必须是单元格或矩形：{range_str!r}")


def _sheet_dimensions(sheet: SheetSpec, sheet_index: int) -> tuple[int, int] | dict[str, str]:
    raw = sheet.dimensions or {}
    try:
        rows = int(raw["rows"])
        cols = int(raw["cols"])
    except (KeyError, TypeError, ValueError):
        return {
            "path": f"sheets.{sheet_index}.dimensions",
            "message": "dimensions 必须包含正整数 rows 与 cols",
        }
    if rows < 1 or cols < 1:
        return {
            "path": f"sheets.{sheet_index}.dimensions",
            "message": f"dimensions 越界: rows={rows}, cols={cols}（须 ≥ 1）",
        }
    if rows > _EXCEL_MAX_ROWS or cols > _EXCEL_MAX_COLS:
        return {
            "path": f"sheets.{sheet_index}.dimensions",
            "message": (
                f"dimensions 超出 Excel 上限: rows={rows}, cols={cols} "
                f"（最大 {_EXCEL_MAX_ROWS}×{_EXCEL_MAX_COLS}）"
            ),
        }
    return rows, cols


def _rect_fits(
    *,
    start_row: int,
    start_col: int,
    height: int,
    width: int,
    rows: int,
    cols: int,
) -> str | None:
    if start_row < 1 or start_col < 1:
        return f"锚点越界: row={start_row}, col={start_col}"
    if height <= 0 or width <= 0:
        return None
    end_row = start_row + height - 1
    end_col = start_col + width - 1
    if end_row > rows or end_col > cols:
        return (
            f"矩形超出 dimensions {rows}×{cols}: "
            f"覆盖 ({start_row},{start_col})–({end_row},{end_col})"
        )
    return None


def collect_layout_errors(spec: WorkbookSpec) -> list[dict[str, str]]:
    """矩形越界、坏样式引用等语义错误（带字段路径）。"""
    errors: list[dict[str, str]] = []
    for i, sheet in enumerate(spec.sheets):
        # source_csv 的 dimensions 由工具层展开后推导，展开前无法做矩形边界检查
        if sheet.source_csv and sheet.dimensions is None:
            continue
        dims = _sheet_dimensions(sheet, i)
        if isinstance(dims, dict):
            errors.append(dims)
            continue
        rows, cols = dims
        prefix = f"sheets.{i}"

        for j, block in enumerate(sheet.value_blocks):
            path = f"{prefix}.value_blocks.{j}"
            try:
                start_row, start_col = _parse_a1(block.start)
            except Exception as exc:
                errors.append({"path": f"{path}.start", "message": f"无效锚点 {block.start!r}: {exc}"})
                continue
            height = len(block.values)
            width = max((len(row) for row in block.values), default=0)
            msg = _rect_fits(
                start_row=start_row, start_col=start_col,
                height=height, width=width, rows=rows, cols=cols,
            )
            if msg:
                errors.append({"path": path, "message": msg})

        for j, block in enumerate(sheet.formula_blocks):
            path = f"{prefix}.formula_blocks.{j}"
            try:
                start_row, start_col = _parse_a1(block.start)
            except Exception as exc:
                errors.append({"path": f"{path}.start", "message": f"无效锚点 {block.start!r}: {exc}"})
                continue
            height = len(block.formulas)
            width = max((len(row) for row in block.formulas), default=0)
            msg = _rect_fits(
                start_row=start_row, start_col=start_col,
                height=height, width=width, rows=rows, cols=cols,
            )
            if msg:
                errors.append({"path": path, "message": msg})

        for j, region in enumerate(sheet.style_regions):
            path = f"{prefix}.style_regions.{j}"
            if region.style_id not in sheet.styles:
                errors.append({
                    "path": f"{path}.style_id",
                    "message": f"未知样式引用: {region.style_id!r}",
                })
            try:
                min_col, min_row, max_col, max_row = _range_bounds(region.range)
            except Exception as exc:
                errors.append({"path": f"{path}.range", "message": f"无效区域 {region.range!r}: {exc}"})
                continue
            if min_row < 1 or min_col < 1 or max_row > rows or max_col > cols:
                errors.append({
                    "path": f"{path}.range",
                    "message": (
                        f"样式区域 {region.range} 超出 dimensions {rows}×{cols}"
                    ),
                })

        for j, merged in enumerate(sheet.merged_ranges):
            path = f"{prefix}.merged_ranges.{j}"
            try:
                min_col, min_row, max_col, max_row = _range_bounds(merged.range)
            except Exception as exc:
                errors.append({"path": f"{path}.range", "message": f"无效合并区 {merged.range!r}: {exc}"})
                continue
            if min_row < 1 or min_col < 1 or max_row > rows or max_col > cols:
                errors.append({
                    "path": f"{path}.range",
                    "message": f"合并区 {merged.range} 超出 dimensions {rows}×{cols}",
                })

        for j, cell in enumerate(sheet.cells):
            path = f"{prefix}.cells.{j}"
            try:
                row, col = _parse_a1(cell.address)
            except Exception as exc:
                errors.append({"path": f"{path}.address", "message": f"无效地址 {cell.address!r}: {exc}"})
                continue
            if row > rows or col > cols or row < 1 or col < 1:
                errors.append({
                    "path": f"{path}.address",
                    "message": f"单元格 {cell.address} 超出 dimensions {rows}×{cols}",
                })
            if cell.style_id and cell.style_id not in sheet.styles:
                errors.append({
                    "path": f"{path}.style_id",
                    "message": f"未知样式引用: {cell.style_id!r}",
                })
    return errors


def validate_workbook_spec(data: Any) -> WorkbookSpec:
    """校验最小 WorkbookSpec；失败时抛出带字段路径的 SpecValidationError。

    接受与 inspect / Excel 常见写法等价的形状（字体名字符串、列字母列宽映射等），
    但不补缺失的必填字段。
    """
    if isinstance(data, str):
        try:
            data = json.loads(data)
        except json.JSONDecodeError as exc:
            if exc.pos >= len(data.rstrip()) - 2:
                message = (
                    f"JSON 在末尾被截断（{exc}）：workbook_spec 数据量大时不要整表一次传入，"
                    "改用 sheets[].source_csv 导入数据文件，或先建骨架再用 operations 分块写入。"
                )
            else:
                message = f"不是合法 JSON: {exc}"
            raise SpecValidationError([{
                "path": "spec",
                "message": message,
            }]) from exc
    if not isinstance(data, dict):
        raise SpecValidationError([{
            "path": "spec",
            "message": "WorkbookSpec 必须是 JSON 对象",
        }])
    try:
        spec = WorkbookSpec.model_validate(data)
    except ValidationError as exc:
        raise SpecValidationError(_pydantic_errors(exc)) from exc
    layout_errors = collect_layout_errors(spec)
    if layout_errors:
        raise SpecValidationError(layout_errors)
    return spec


def iter_block_cells(start: str, grid: list[list[Any]]) -> list[tuple[str, Any]]:
    from openpyxl.utils import get_column_letter

    start_row, start_col = _parse_a1(start)
    out: list[tuple[str, Any]] = []
    for r_off, row_vals in enumerate(grid):
        if not isinstance(row_vals, list):
            continue
        for c_off, value in enumerate(row_vals):
            addr = f"{get_column_letter(start_col + c_off)}{start_row + r_off}"
            out.append((addr, value))
    return out


def iter_range_addresses(range_str: str) -> list[str]:
    from openpyxl.utils import get_column_letter

    min_col, min_row, max_col, max_row = _range_bounds(range_str)
    addrs: list[str] = []
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            addrs.append(f"{get_column_letter(col)}{row}")
    return addrs


_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_ISO_DATETIME_RE = re.compile(
    r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})?$"
)


def _coerce_iso_temporal(value: str) -> datetime | date | None:
    """严格 ISO-8601 字符串 → date/datetime（与 Excel 输入行为一致）。"""
    text = value.strip()
    try:
        if _ISO_DATE_RE.match(text):
            return date.fromisoformat(text)
        if _ISO_DATETIME_RE.match(text):
            return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None
    return None


def materialize_sheet_cells(sheet: SheetSpec) -> list[CellSpec]:
    """把矩形值/公式块与样式区域展开为单元格列表，供 rebuild 编译。"""
    by_addr: dict[str, CellSpec] = {}
    for cell in sheet.cells:
        by_addr[cell.address.upper()] = cell.model_copy(deep=True)

    for block in sheet.value_blocks:
        for addr, value in iter_block_cells(block.start, block.values):
            key = addr.upper()
            existing = by_addr.get(key)
            value_type: Literal["string", "number", "date", "boolean", "formula", "empty"]
            if value is None:
                value_type = "empty"
            elif isinstance(value, bool):
                value_type = "boolean"
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                value_type = "number"
            elif isinstance(value, (datetime, date, time)):
                value_type = "date"
            elif isinstance(value, str) and value.startswith("="):
                value_type = "formula"
            elif isinstance(value, str) and (_temporal := _coerce_iso_temporal(value)) is not None:
                value = _temporal
                value_type = "date"
            else:
                value_type = "string"
            if existing is None:
                by_addr[key] = CellSpec(address=addr, value=value, value_type=value_type)
            else:
                existing.value = value
                existing.value_type = value_type

    for block in sheet.formula_blocks:
        for addr, formula in iter_block_cells(block.start, list(block.formulas)):
            key = addr.upper()
            text = str(formula) if formula is not None else ""
            existing = by_addr.get(key)
            if existing is None:
                by_addr[key] = CellSpec(address=addr, value=text, value_type="formula")
            else:
                existing.value = text
                existing.value_type = "formula"

    for region in sheet.style_regions:
        try:
            addrs = iter_range_addresses(region.range)
        except Exception:
            continue
        for addr in addrs:
            key = addr.upper()
            existing = by_addr.get(key)
            if existing is None:
                by_addr[key] = CellSpec(
                    address=addr, value=None, value_type="empty", style_id=region.style_id,
                )
            else:
                existing.style_id = region.style_id

    return list(by_addr.values())


def workbook_spec_to_replica(spec: WorkbookSpec, provenance: Provenance | None = None) -> ReplicaSpec:
    sheets: list[SheetSpec] = []
    for sheet in spec.sheets:
        dumped = sheet.model_copy(deep=True)
        dumped.cells = materialize_sheet_cells(sheet)
        sheets.append(dumped)
    if provenance is None:
        provenance = Provenance(
            source_image_hash="",
            model="workbook-spec",
            timestamp=datetime.now(timezone.utc).isoformat(),
        )
    return ReplicaSpec(
        version=spec.version,
        provenance=provenance,
        workbook=WorkbookMeta(
            name=spec.name,
            locale=spec.locale,
            default_font=spec.default_font,
            theme_hint=spec.theme_hint,
        ),
        sheets=sheets,
        uncertainties=list(spec.uncertainties),
    )


def looks_like_replica_spec(data: dict[str, Any]) -> bool:
    return "provenance" in data or (
        isinstance(data.get("workbook"), dict) and "sheets" in data
    )


def load_spec_document(text: str) -> ReplicaSpec:
    """加载 ReplicaSpec 或最小 WorkbookSpec，统一为 ReplicaSpec 供编译。"""
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Spec 不是合法 JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("Spec 必须是 JSON 对象")
    if looks_like_replica_spec(data):
        return ReplicaSpec.model_validate(data)
    wb = validate_workbook_spec(data)
    return workbook_spec_to_replica(wb)


def compile_replica_to_bytes(replica: ReplicaSpec) -> tuple[bytes, dict[str, Any]]:
    """把 ReplicaSpec 编译为 xlsx 字节。WorkbookSpec 先转 replica 再走这里。"""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def _side(part: BorderSideSpec | None, fallback_style: str | None, fallback_color: str | None) -> Side:
        style = ((part.style if part else None) or fallback_style or "thin")
        color = ((part.color if part else None) or fallback_color)
        kwargs: dict[str, Any] = {"style": style}
        if color:
            kwargs["color"] = str(color).lstrip("#")
        return Side(**kwargs)

    def _border(spec: BorderSpec) -> Border:
        return Border(
            left=_side(spec.left, spec.style, spec.color),
            right=_side(spec.right, spec.style, spec.color),
            top=_side(spec.top, spec.style, spec.color),
            bottom=_side(spec.bottom, spec.style, spec.color),
        )

    wb = Workbook()
    if wb.sheetnames:
        del wb[wb.sheetnames[0]]
    cells_written = 0
    formulas_written = 0
    merges_applied = 0
    default_font_obj = None
    if replica.workbook.default_font is not None:
        spec_font = replica.workbook.default_font
        default_font_obj = Font(
            name=spec_font.name,
            size=spec_font.size,
            bold=spec_font.bold,
            italic=spec_font.italic,
            color=spec_font.color.lstrip("#") if spec_font.color else None,
            underline=spec_font.underline,
            strike=spec_font.strike,
        )
        try:
            if spec_font.name:
                wb._fonts[0].name = spec_font.name
            if spec_font.size:
                wb._fonts[0].sz = spec_font.size
        except Exception:
            pass
    for sheet in replica.sheets:
        ws = wb.create_sheet(title=sheet.name)
        for cell in materialize_sheet_cells(sheet):
            target = ws[cell.address]
            if cell.value_type == "formula" and cell.value:
                target.value = str(cell.value)
                formulas_written += 1
            elif cell.value_type == "number" and cell.value is not None:
                try:
                    target.value = float(cell.value) if "." in str(cell.value) else int(cell.value)
                except (TypeError, ValueError):
                    target.value = cell.value
            elif cell.value_type == "empty":
                target.value = None
            else:
                target.value = cell.value
            if cell.number_format:
                target.number_format = cell.number_format
            style = sheet.styles.get(cell.style_id) if cell.style_id else None
            if style is not None and style.font is not None:
                fallback = default_font_obj
                target.font = Font(
                    name=style.font.name or (fallback.name if fallback else None),
                    size=style.font.size if style.font.size is not None else (
                        fallback.size if fallback else None
                    ),
                    bold=style.font.bold,
                    italic=style.font.italic,
                    color=style.font.color.lstrip("#") if style.font.color else None,
                    underline=style.font.underline,
                    strike=style.font.strike,
                )
            elif default_font_obj is not None:
                target.font = default_font_obj
            if (
                style is not None
                and style.fill is not None
                and style.fill.color
                and style.fill.type != "none"
            ):
                target.fill = PatternFill(
                    patternType="solid",
                    fgColor=style.fill.color.lstrip("#"),
                )
            if style is not None and style.alignment is not None:
                target.alignment = Alignment(
                    horizontal=style.alignment.horizontal,
                    vertical=style.alignment.vertical,
                    wrap_text=style.alignment.wrap_text,
                )
            if style is not None and style.border is not None:
                target.border = _border(style.border)
            if style is not None and style.number_format:
                target.number_format = style.number_format
            cells_written += 1
        merge_errors: list[dict[str, str]] = []
        for merged in sheet.merged_ranges:
            try:
                ws.merge_cells(merged.range)
                merges_applied += 1
            except Exception as exc:
                merge_errors.append({"range": merged.range, "error": str(exc)})
        if merge_errors:
            raise ValueError(f"合并失败: {merge_errors}")
        for index, width in enumerate(sheet.column_widths or []):
            if isinstance(width, (int, float)) and width > 0:
                ws.column_dimensions[get_column_letter(index + 1)].width = float(width)
        for row_key, height in (sheet.row_heights or {}).items():
            try:
                ws.row_dimensions[int(row_key)].height = float(height)
            except (TypeError, ValueError):
                continue
        if sheet.freeze_panes:
            from excelmanus.workbook.styles import apply_freeze_panes

            apply_freeze_panes(ws, sheet.freeze_panes)
        if sheet.print_layout is not None:
            apply_print_layout(ws, sheet.print_layout)
        for cf_index, cf_rule in enumerate(sheet.conditional_formats or []):
            try:
                from excelmanus.workbook.styles import build_conditional_format_rule

                rule = build_conditional_format_rule(
                    {
                        "type": cf_rule.type,
                        "operator": cf_rule.operator,
                        "value": cf_rule.value,
                        "value2": cf_rule.value2,
                        "min_color": cf_rule.min_color,
                        "mid_color": cf_rule.mid_color,
                        "max_color": cf_rule.max_color,
                        "bar_color": cf_rule.bar_color,
                        "icon_style": cf_rule.icon_style,
                        "fill": {"color": cf_rule.fill_color} if cf_rule.fill_color else None,
                        "font": (
                            {"color": cf_rule.font_color, "bold": cf_rule.bold}
                            if (cf_rule.font_color or cf_rule.bold)
                            else None
                        ),
                    },
                    anchor=str(cf_rule.range).split()[0].split(":")[0],
                )
                ws.conditional_formatting.add(cf_rule.range, rule)
            except Exception as exc:
                raise ValueError(
                    f"sheet {sheet.name!r} conditional_formats[{cf_index}] "
                    f"应用失败 range={cf_rule.range!r}: {exc}"
                ) from exc
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), {
        "cells_written": cells_written,
        "formulas_written": formulas_written,
        "merges_applied": merges_applied,
    }


def compile_workbook_spec_to_bytes(spec: WorkbookSpec) -> tuple[bytes, dict[str, Any]]:
    """把已校验的 WorkbookSpec 编译为 xlsx 字节，供同一道 commit_* 提交。"""
    return compile_replica_to_bytes(workbook_spec_to_replica(spec))


def compile_spec_text_to_bytes(text: str) -> tuple[bytes, dict[str, Any]]:
    """从 ReplicaSpec / WorkbookSpec JSON 文本编译为 xlsx 字节。"""
    return compile_replica_to_bytes(load_spec_document(text))


def workbook_spec_json_schema() -> dict[str, Any]:
    """从 Pydantic 模型生成供工具 schema / 字段查询使用的结构。"""
    schema = WorkbookSpec.model_json_schema(mode="validation")
    from excelmanus.tools.workbook_examples import workbook_creation_example

    schema["examples"] = [workbook_creation_example()]
    schema["description"] = (
        "创建用 WorkbookSpec，与 operations 互斥。必填 sheets 与 uncertainties。"
        "规范输入见 properties；字符串字体、列宽字典、CSV 省略 dimensions 等兼容形由校验器接受。"
    )
    return schema
