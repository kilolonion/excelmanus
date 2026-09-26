"""Deterministic checks and disk-backed SQL for spreadsheet datasets."""
from __future__ import annotations
import ast
import csv
from io import BytesIO
import json
import operator
from pathlib import Path
import re
import sqlite3
import tempfile
import time

from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from excelmanus.prompt.canonical import TOOL_DESCRIPTIONS
from excelmanus.tools.registry import ToolDef
from excelmanus.tools.context import bind_workspace, require_guard
from excelmanus.engine_core.tool_result import from_payload, error_result
from excelmanus.tools.spreadsheet_engine_tools import _source, _publish, _cancelled, _failure
from excelmanus.workbook.formula_values import FormulaValueError


def init_guard(workspace_root): bind_workspace(workspace_root)


_EXPRESSION_HELP = (
    "表达式语法：可引用当前行的列名或列字母（如 金额、E；{row} 是可选占位符，E{row} 等价于 E），"
    "支持 + - * / % 算术、== != < <= > >= 比较、and/or/not 逻辑与括号，"
    "以及大小写不敏感的 abs/round/min/max 函数。"
    "本语言不支持 A1 单元格引用：abs(B7-1)<0.001 里的 B7 是列名而不是第 7 行的 B 列；"
    "要断言某个具体单元格的值，请用 cell 规则（如 {\"kind\":\"cell\",\"sheet\":\"...\",\"cell\":\"B7\",\"expected\":6.707}），"
    "或用 observe_spreadsheet 回读该坐标，或把期望常量写进同一行的数据列再用 row_expression 断言该列。"
    "正确示例：abs(金额-数量*单价)<0.01；ABS(E{row}+F{row}+G{row}-C{row})<0.05"
)

# 形如 B7 / AA12 的名字：在 Python 表达式里会被解析成列名，但调用方几乎总是想写 A1 单元格引用。
_A1_STYLE_REFERENCE = re.compile(r"[A-Za-z]{1,3}\d+")

_A1_EXPRESSION_REMEDIATION = (
    "row_expression 的变量只能是当前行的列（列名或列字母），不支持 A1 单元格引用：B7 会被当成列名，而不是第 7 行的 B 列。"
    "首选：若要断言具体单元格，保留坐标并改用 cell + expected + tolerance，"
    "例如 {\"kind\":\"cell\",\"sheet\":\"回归分析\",\"cell\":\"B7\",\"expected\":6.707067669172933,\"tolerance\":0.0001}；"
    "B8 应另建一条 cell 规则，使用它自己的 expected；不确定坐标或值时先 observe_spreadsheet 回读。"
    "仅在确实要逐行校验时，才写 B 或 B{row}（表示当前行的 B 列），如 abs(B-6.707067669172933)<0.0001，"
    "或把期望常量写进同行数据列再比较；不要为了消除报错而去掉单点坐标的行号。"
    "不要用 total 替代单点断言：total 校验整列合计。不要把同一组参数原样重试。"
)

_EXPRESSION_FUNCTIONS = {"abs": abs, "round": round, "min": min, "max": max}


class ExpressionSyntaxError(ValueError):
    """表达式语法/结构不合法：参数类错误（INVALID_ARGS），不是数据问题。"""
    code = "INVALID_ARGS"


def normalize_expression(expression):
    """支持 {row} 占位符：E{row} → E，即引用当前行的 E 列。"""
    return re.sub(r"\{\s*row\s*\}", "", str(expression), flags=re.IGNORECASE)


def parse_expression(expression):
    """解析表达式为 AST；语法错误抛结构化 ExpressionSyntaxError，绝不泄漏 SyntaxError。"""
    if not isinstance(expression, str) or not expression.strip():
        raise ExpressionSyntaxError(f"expression 不能为空。{_EXPRESSION_HELP}")
    try:
        tree = ast.parse(normalize_expression(expression), mode="eval")
    except (SyntaxError, ValueError) as exc:
        raise ExpressionSyntaxError(f"expression 语法无效：{_EXPRESSION_HELP}") from exc
    if len(list(ast.walk(tree))) > 100:
        raise ExpressionSyntaxError(f"expression 过长（最多 100 个语法节点）：{_EXPRESSION_HELP}")
    return tree


def expression_names(tree):
    """表达式里作为变量引用的列名/列字母（不含函数名）。"""
    called = {node.func.id for node in ast.walk(tree) if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)}
    return [node.id for node in ast.walk(tree) if isinstance(node, ast.Name) and node.id not in called]


def evaluate_expression(expression, values):
    """Small arithmetic/predicate language; no attributes, imports or arbitrary calls."""
    return evaluate_parsed_expression(parse_expression(expression), values)


def evaluate_parsed_expression(tree, values):
    binary={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv,ast.Mod:operator.mod}
    compare={ast.Eq:operator.eq,ast.NotEq:operator.ne,ast.Lt:operator.lt,ast.LtE:operator.le,ast.Gt:operator.gt,ast.GtE:operator.ge}
    def visit(node):
        if isinstance(node,ast.Expression): return visit(node.body)
        if isinstance(node,ast.Constant) and isinstance(node.value,(str,int,float,bool,type(None))): return node.value
        if isinstance(node,ast.Name) and node.id in values: return values[node.id]
        if isinstance(node,ast.Name): raise ValueError(f"表达式引用了不存在的列 {node.id!r}；{_EXPRESSION_HELP}")
        if isinstance(node,ast.BinOp) and type(node.op) in binary:
            a,b=visit(node.left),visit(node.right)
            if not isinstance(a,(int,float)) or not isinstance(b,(int,float)): raise ValueError("算术表达式仅接受数值")
            return binary[type(node.op)](a,b)
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.USub,ast.UAdd,ast.Not)):
            v=visit(node.operand)
            return not v if isinstance(node.op,ast.Not) else -v if isinstance(node.op,ast.USub) else +v
        if isinstance(node,ast.BoolOp):
            return all(visit(v) for v in node.values) if isinstance(node.op,ast.And) else any(visit(v) for v in node.values)
        if isinstance(node,ast.Compare):
            left=visit(node.left)
            for op,right in zip(node.ops,node.comparators):
                value=visit(right)
                if type(op) not in compare or not compare[type(op)](left,value): return False
                left=value
            return True
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id.lower() in _EXPRESSION_FUNCTIONS and not node.keywords:
            args=[visit(a) for a in node.args]
            if not all(isinstance(a,(int,float)) for a in args): raise ValueError("函数参数必须是数值")
            return _EXPRESSION_FUNCTIONS[node.func.id.lower()](*args)
        raise ValueError(f"不支持的表达式；{_EXPRESSION_HELP}")
    return visit(tree)


def _limited(names, cap=100):
    return list(names)[:cap]


def _rule_arg_error(number, kind, sheet, message, code="INVALID_ARGS", **fields):
    """参数类错误：message 固定带 rules[i] 下标、sheet、涉及列；failure_class=invalid_args。"""
    error = ValueError(f"rules[{number}] ({kind}) sheet={sheet!r}：{message}")
    error.code = code
    error.fields = {"rule_index": number, "kind": kind, "sheet": sheet, **fields}
    return error


def _finite_rule_number(value, *, number, kind, sheet, field, nonnegative=False):
    """只接受有限 JSON 数值；布尔、NaN/Infinity 和浮点溢出均为参数错误。"""
    from math import isfinite

    parsed = None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = float(value)
        except (ValueError, OverflowError):
            pass
    if parsed is None or not isfinite(parsed) or (nonnegative and parsed < 0):
        requirement = "有限非负数值（默认 1e-9）" if nonnegative else "有限数值"
        raise _rule_arg_error(
            number, kind, sheet,
            f"{field} 无效：必须是{requirement}，不接受布尔、NaN、Infinity 或超出浮点范围的数值",
            field=field,
        )
    return parsed


def _resolve_column(value, *, names, number, kind, sheet, field):
    """列引用解析：表头列名 > 列字母(A/AA) > 1-based 列号；失败时报可用表头名而非 Python 异常。"""
    index = None
    if isinstance(value, str):
        text = value.strip()
        if value in names or text in names:
            return names.index(value if value in names else text)
        letter = text.replace("$", "")
        if re.fullmatch(r"[A-Za-z]{1,3}", letter):
            try:
                index = column_index_from_string(letter) - 1
            except ValueError:
                index = None
        elif re.fullmatch(r"\d+", text):
            index = int(text) - 1
    elif isinstance(value, int) and not isinstance(value, bool):
        index = value - 1
    if index is not None and 0 <= index < len(names):
        return index
    raise _rule_arg_error(
        number, kind, sheet,
        f"{field} 无法解析列引用 {value!r}：支持表头列名（与列字母同名时列名优先）、列字母（A/B/AA）或 1-based 列号；可用表头名：{_limited(names)}",
        field=field, column_ref=value, available_columns=_limited(names),
    )


# cell 规则：A1 单格坐标（允许 $B$7），越界/区域/带工作表前缀一律报 INVALID_ARGS。
_A1_CELL = re.compile(r"\$?([A-Za-z]{1,3})\$?(\d{1,7})")
_MAX_ROW, _MAX_COLUMN = 1048576, 16384


def _parse_cell_reference(value, *, number, kind, sheet):
    """解析 cell 规则的坐标，返回 (row, column, "B7")；只接受 A1 单格。"""
    if not isinstance(value, str) or not value.strip():
        raise _rule_arg_error(number, kind, sheet, 'cell 需要 A1 单格坐标（如 "B7"），不接受空值', field="cell", cell_ref=value)
    text = value.strip()
    if "!" in text:
        raise _rule_arg_error(
            number, kind, sheet,
            f"cell={value!r} 不能带工作表前缀：cell 只写 A1 单格坐标，工作表请写在规则的 sheet 字段里，"
            f'例如 {{"kind":"cell","sheet":"回归分析","cell":"B7","expected":6.707}}',
            field="cell", cell_ref=value,
        )
    match = _A1_CELL.fullmatch(text)
    if not match:
        raise _rule_arg_error(
            number, kind, sheet,
            f"cell={value!r} 不是合法的 A1 单格坐标：写法是列字母+行号（B7、AA12，可带 $）；"
            "不接受区域（B7:C9）、纯列名（B）或纯行号（7）",
            field="cell", cell_ref=value,
        )
    column = column_index_from_string(match.group(1).upper())
    row = int(match.group(2))
    if not (1 <= row <= _MAX_ROW and 1 <= column <= _MAX_COLUMN):
        raise _rule_arg_error(
            number, kind, sheet,
            f"cell={value!r} 不是有效的 A1 单格坐标：超出工作表范围（合法行 1..{_MAX_ROW}，合法列 A..{get_column_letter(_MAX_COLUMN)}）",
            field="cell", cell_ref=value,
        )
    return row, column, f"{get_column_letter(column)}{row}"


def validate_spreadsheet(
    file_path,
    rules,
    sheet=None,
    expected_version=None,
    max_failures=200,
    allow_uncached=True,
):
    book=None
    try:
        if not isinstance(rules,list) or not rules or not 1<=max_failures<=2000: raise ValueError("rules/max_failures 无效")
        source,data,version=_source(file_path,expected_version)
        book=load_workbook(BytesIO(data),data_only=False)
        from excelmanus.workbook.formula_values import attach_formula_source,analysis_rows
        attach_formula_source(book,data)
        results=[]; failures=[]; failure_count=0; versions={file_path:version}
        uncalculated_cells=[]
        validation_status="complete"

        def note_uncalculated(cells):
            nonlocal validation_status
            validation_status = "partial"
            for cell in cells or []:
                cell = str(cell)
                if cell and cell not in uncalculated_cells:
                    uncalculated_cells.append(cell)

        def rows_or_uncached(number, rule_kind, ws):
            """analysis_rows 的统一入口：公式缓存缺失时沿用既有口径（allow_uncached=false 阻断，否则 partial）。"""
            try:
                return analysis_rows(ws), None
            except FormulaValueError as exc:
                if not allow_uncached:
                    raise
                note_uncalculated(exc.cells)
                return None, {
                    "rule":number,
                    "sheet":ws.title,
                    "kind":rule_kind,
                    "checked":0,
                    "failed":0,
                    "status":"uncalculated",
                    "formula_cache":"missing",
                    "cells":list(exc.cells or [])[:max_failures],
                }

        checks = []
        for number, rule in enumerate(rules):
            if not isinstance(rule, dict) or not rule.get("kind"):
                raise _rule_arg_error(number, rule.get("kind") if isinstance(rule, dict) else None, sheet, "每条规则必须是带 kind 字段的对象", field="kind")
            selected = rule.get("sheet") or sheet
            if selected:
                if selected not in book.sheetnames:
                    raise _rule_arg_error(number, rule.get("kind"), selected, f"sheet={selected!r} 不存在，可用工作表：{book.sheetnames}", code="SHEET_NOT_FOUND", available_sheets=list(book.sheetnames))
                targets = [book[selected]]
            elif rule.get("kind") == "formula_errors":
                targets = book.worksheets
            elif len(book.worksheets) == 1:
                targets = [book.active]
            else:
                return error_result(
                    f"rules[{number}] ({rule.get('kind')}) 需要指定 sheet",
                    code="SHEET_REQUIRED",
                    fields={"rule_index": number, "available_sheets": book.sheetnames},
                )
            checks.extend((number, rule, ws) for ws in targets)
        for number, rule, ws in checks:
            kind=rule.get("kind"); failed=0; checked=0; skipped=0; skipped_cells=[]; extras={}
            def fail(address,actual,expected=None,**fields):
                nonlocal failed,failure_count
                failed+=1; failure_count+=1
                if len(failures)<max_failures:
                    item={"rule":number,"kind":kind,"sheet":ws.title,"cell":address,"actual":actual,"expected":expected}
                    item.update({key:value for key,value in fields.items() if value is not None})
                    failures.append(item)
            if kind not in {"formula_errors","unique","required","foreign_key","row_expression","total","cell"}:
                raise _rule_arg_error(number,kind,ws.title,f"不支持的校验规则 kind={kind!r}，可用：unique/required/foreign_key/row_expression/total/cell/formula_errors",field="kind")
            if kind=="formula_errors":
                cached=load_workbook(BytesIO(data),data_only=True,read_only=True)
                try:
                    matrix=list(cached[ws.title].iter_rows())
                    for row in matrix:
                        for cell in row:
                            checked+=1
                            if cell.data_type=="e": fail(cell.coordinate,cell.value,"no formula error")
                    missing=[]
                    for cell in ws._cells.values():
                        if cell.data_type=="f" and (cell.row>len(matrix) or cell.column>len(matrix[cell.row-1]) or matrix[cell.row-1][cell.column-1].value is None):
                            missing.append(f"{ws.title}!{cell.coordinate}")
                    if missing:
                        if allow_uncached:
                            note_uncalculated(missing)
                        else:
                            for cell in missing:
                                fail(cell.split("!", 1)[-1],"missing_cache","calculated value")
                finally: cached.close()
            elif kind=="cell":
                from excelmanus.workbook.numeric import parse_number
                from math import isfinite
                cell_row, cell_column, address = _parse_cell_reference(rule.get("cell"), number=number, kind=kind, sheet=ws.title)
                if "expected" not in rule:
                    raise _rule_arg_error(number,kind,ws.title,'cell 需要 expected（期望值：数值/文本/布尔/null；断言空单元格用 null）',field="expected")
                expected=rule.get("expected")
                is_number=isinstance(expected,(int,float)) and not isinstance(expected,bool)
                if not (is_number or expected is None or isinstance(expected,(str,bool))):
                    raise _rule_arg_error(number,kind,ws.title,f"cell 的 expected 类型不支持：{type(expected).__name__}；只支持数值/文本/布尔/null",field="expected")
                if is_number:
                    _finite_rule_number(expected, number=number, kind=kind, sheet=ws.title, field="expected")
                tolerance=_finite_rule_number(
                    rule.get("tolerance",1e-9), number=number, kind=kind, sheet=ws.title,
                    field="tolerance", nonnegative=True,
                )
                case_sensitive=rule.get("case_sensitive",True)
                if not isinstance(case_sensitive,bool):
                    raise _rule_arg_error(number,kind,ws.title,f"case_sensitive={case_sensitive!r} 无效：必须是布尔值",field="case_sensitive")
                rows, uncached = rows_or_uncached(number, kind, ws)
                if uncached is not None:
                    results.append(uncached)
                    continue
                row_values=rows[cell_row-1] if 1<=cell_row<=len(rows) else ()
                actual=row_values[cell_column-1] if 0<=cell_column-1<len(row_values) else None
                if hasattr(actual,"isoformat"): actual=actual.isoformat()  # 日期/时间按 ISO 文本比对，保持载荷 JSON 安全
                checked+=1
                blank=actual is None or (isinstance(actual,str) and not actual.strip())
                matched=False; reason=None; delta=None; message=None
                if expected is None:
                    matched=blank
                    if not matched:
                        reason="value_mismatch"; message=f"{address} 不是空单元格（actual={actual!r}）；expected=null 表示断言该格为空。"
                elif is_number:
                    value=parse_number(actual)
                    if not isfinite(value):
                        if blank:
                            reason="empty_cell"; message=f"{address} 是空单元格，无法与数值 expected={expected} 比较。"
                        else:
                            reason="non_numeric_cell"; message=f"{address} 的实际值 {actual!r} 无法解析为数值；若该格确实是文本，请把 expected 写成字符串。"
                    else:
                        delta=value-float(expected)
                        matched=abs(delta)<=tolerance
                        if not matched:
                            reason="value_mismatch"; message=f"{address}={value}，与 expected={expected} 相差 {delta}（tolerance={tolerance}）。"
                elif isinstance(expected,bool):
                    if blank:
                        reason="empty_cell"; message=f"{address} 是空单元格，无法与布尔 expected={expected} 比较。"
                    elif not isinstance(actual,bool):
                        reason="type_mismatch"; message=f"{address} 的实际值类型是 {type(actual).__name__}（actual={actual!r}），expected 是布尔值。"
                    elif actual is not expected:
                        reason="value_mismatch"; message=f"{address}={actual}，与 expected={expected} 不符。"
                    else: matched=True
                else:
                    if blank:
                        reason="empty_cell"; message=f"{address} 是空单元格，无法与文本 expected={expected!r} 精确比对。"
                    elif not isinstance(actual,str):
                        reason="type_mismatch"; message=f"{address} 的实际值类型是 {type(actual).__name__}（actual={actual!r}），而 expected 是文本；数值断言请把 expected 写成 JSON number。"
                    else:
                        matched=actual==expected if case_sensitive else actual.casefold()==expected.casefold()
                        if not matched:
                            reason="value_mismatch"; message=f"{address}={actual!r}，与 expected={expected!r} 不相等（case_sensitive={case_sensitive}）。"
                if not matched:
                    fail(address,actual,expected,reason=reason,delta=delta,message=message)
                extras.update(semantics="cell_assertion",cell=address,expected=expected,actual=actual,matched=matched)
                if is_number: extras["tolerance"]=tolerance
                if isinstance(expected,str) and not case_sensitive: extras["case_sensitive"]=False
            else:
                rows, uncached = rows_or_uncached(number, kind, ws)
                if uncached is not None:
                    results.append(uncached)
                    continue
                raw_header=rule.get("header_row",1)
                try:
                    header_row=int(raw_header)
                except (TypeError,ValueError):
                    raise _rule_arg_error(number,kind,ws.title,f"header_row={raw_header!r} 无效，必须是 ≥1 的整数",field="header_row") from None
                if not 1<=header_row<=len(rows):
                    raise _rule_arg_error(number,kind,ws.title,f"header_row={header_row} 超出实际行数 {len(rows)}",field="header_row")
                header=rows[header_row-1]
                names=[str(h) if h is not None else get_column_letter(i+1) for i,h in enumerate(header)]
                raw_columns=rule.get("columns") or ([rule["column"]] if "column" in rule else [])
                if not isinstance(raw_columns,list):
                    raise _rule_arg_error(number,kind,ws.title,"columns 必须是列表（单列可用 column）",field="columns")
                column_field="columns" if rule.get("columns") else "column"
                indices=[_resolve_column(v,names=names,number=number,kind=kind,sheet=ws.title,field=column_field) for v in raw_columns]
                body=rows[header_row:]
                def cell(row,c): return row[c] if 0<=c<len(row) else None
                if kind in {"unique","required","foreign_key"}:
                    if not indices: raise _rule_arg_error(number,kind,ws.title,f"{kind} 需要 columns（或 column）指定列：列名、列字母或 1-based 列号",field="columns")
                    seen=set(); reference=set()
                    if kind=="foreign_key":
                        ref=rule.get("reference")
                        if not isinstance(ref,dict) or not ref.get("file_path") or not isinstance(ref.get("columns"),list) or not ref["columns"]:
                            raise _rule_arg_error(number,kind,ws.title,"foreign_key 需要 reference={file_path, columns, sheet?, header_row?, expected_version?}",field="reference")
                        path,raw,ver=_source(ref["file_path"],ref.get("expected_version")); versions[ref["file_path"]]=ver
                        other=load_workbook(BytesIO(raw),data_only=False)
                        try:
                            attach_formula_source(other,raw)
                            if ref.get("sheet"):
                                if ref["sheet"] not in other.sheetnames:
                                    raise _rule_arg_error(number,kind,ws.title,f"reference.sheet={ref['sheet']!r} 不存在，可用工作表：{other.sheetnames}",code="SHEET_NOT_FOUND",field="reference.sheet",available_sheets=list(other.sheetnames))
                                other_ws=other[ref["sheet"]]
                            else:
                                other_ws=other.active
                            other_rows=analysis_rows(other_ws); h=int(ref.get("header_row",1))
                            ref_header=other_rows[h-1] if 0<h<=len(other_rows) else ()
                            headers=[str(x) if x is not None else get_column_letter(i+1) for i,x in enumerate(ref_header)]
                            rc=[_resolve_column(c,names=headers,number=number,kind=kind,sheet=ws.title,field="reference.columns") for c in ref["columns"]]
                            reference={tuple(cell(row,c) for c in rc) for row in other_rows[h:]}
                        finally: other.close()
                    for row_number,row in enumerate(body,header_row+1):
                        key=tuple(cell(row,c) for c in indices); checked+=1
                        bad=any(v is None or v=="" for v in key) if kind=="required" else key in seen if kind=="unique" else key not in reference
                        if bad: fail(f"{get_column_letter(indices[0]+1)}{row_number}",list(key),kind)
                        seen.add(key)
                elif kind=="row_expression":
                    expression=rule.get("expression")
                    try:
                        tree=parse_expression(expression)
                    except ExpressionSyntaxError as exc:
                        raise _rule_arg_error(number,kind,ws.title,f"expression 语法无效：{exc}",field="expression",expression=str(expression))
                    identifier_names=[n for n in names if isinstance(n,str) and n.isidentifier()]
                    usable={get_column_letter(i+1) for i in range(len(names))}|set(identifier_names)
                    unknown=[n for n in expression_names(tree) if n not in usable]
                    if unknown:
                        a1_style=[n for n in unknown if _A1_STYLE_REFERENCE.fullmatch(n)]
                        detail=[
                            f"expression 引用了不存在的列 {unknown}；可用列名：{_limited(identifier_names)}（也可用列字母 A、B、C…）",
                        ]
                        if a1_style:
                            detail.append(f"检测到 A1 风格引用 {a1_style}。")
                        detail.append(_A1_EXPRESSION_REMEDIATION)
                        raise _rule_arg_error(
                            number,kind,ws.title,"".join(detail),
                            field="expression",expression=str(expression),unknown_names=unknown,
                            a1_style_references=a1_style,remediation=_A1_EXPRESSION_REMEDIATION,
                        )
                    for row_number,row in enumerate(body,header_row+1):
                        checked+=1
                        env={get_column_letter(i+1):v for i,v in enumerate(row)}; env.update({k:v for k,v in zip(names,row) if isinstance(k,str) and k.isidentifier()})
                        try: ok=bool(evaluate_parsed_expression(tree,env))
                        except Exception as exc:
                            fail(f"A{row_number}",None,str(expression),reason="expression_error",expression=str(expression),message=f"该行无法求值（值类型或运算不匹配）：{exc}")
                            continue
                        if not ok: fail(f"A{row_number}",list(row),str(expression),reason="expression_false")
                elif kind=="total":
                    from excelmanus.workbook.numeric import parse_number
                    from math import fsum,isfinite
                    if len(indices)!=1: raise _rule_arg_error(number,kind,ws.title,"total 需要且仅需要一个 column：列名、列字母或 1-based 列号",field="column")
                    expected=rule.get("expected")
                    _finite_rule_number(expected, number=number, kind=kind, sheet=ws.title, field="expected")
                    tolerance=_finite_rule_number(
                        rule.get("tolerance",1e-9), number=number, kind=kind, sheet=ws.title,
                        field="tolerance", nonnegative=True,
                    )
                    column=indices[0]; letter=get_column_letter(column+1)
                    values=[]; excluded_cells=[]; excluded_count=0; expected_matches=[]; last_row=header_row+1
                    for row_number,row in enumerate(body,header_row+1):
                        last_row=row_number
                        raw=cell(row,column)
                        if raw is None or (isinstance(raw,str) and not raw.strip()):
                            skipped+=1
                            if len(skipped_cells)<max_failures: skipped_cells.append(f"{letter}{row_number}")
                            continue
                        checked+=1
                        address=f"{letter}{row_number}"
                        value=parse_number(raw)
                        if not isfinite(value):
                            excluded_count+=1
                            if len(excluded_cells)<max_failures: excluded_cells.append(address)
                            fail(
                                address,None,None,
                                reason="non_numeric_text",
                                raw_value=raw if isinstance(raw,str) else str(raw),
                                message=f"{address} 是非数值文本，未计入 {letter} 列合计（原因码 non_numeric_text）；要计入合计请把该格改成数值。",
                            )
                            continue
                        values.append(value)
                        # expected 恰好等于列内某个单元格的值：典型的"把单点断言写成整列求和"信号（交付门按事实采信）。
                        if abs(value-expected)<=tolerance and len(expected_matches)<max_failures:
                            expected_matches.append(address)
                    total_value=fsum(values)
                    extras.update(
                        semantics="column_total", column=letter, expected=expected, tolerance=tolerance,
                        total=total_value, excluded_non_numeric=excluded_count, excluded_non_numeric_cells=excluded_cells,
                    )
                    if abs(total_value-expected)>tolerance:
                        if expected_matches: extras["expected_matches_cell"]=expected_matches
                        hint=(
                            f"注意：expected={expected} 与列内单元格 {'、'.join(expected_matches[:5])} 的值相等，"
                            f"说明你想断言的可能是单元格的值而不是整列合计；请改用 cell 规则"
                            f'（{{"kind":"cell","sheet":"{ws.title}","cell":"'
                            f'{expected_matches[0]}","expected":{expected},"tolerance":{tolerance}}}）直接断言该单元格，'
                            "或先用 observe_spreadsheet 回读坐标；仅在确实要逐行校验时才用 row_expression。"
                        ) if expected_matches else (
                            "total 校验的是整列合计，失败项的 cell 字段只是定位用的首个数据行；"
                            f'若要断言单个单元格，请先用 observe_spreadsheet 确认坐标，再用 cell + expected + tolerance'
                            f'（例如 {{"kind":"cell","sheet":"{ws.title}","cell":"B7","expected":{expected},"tolerance":{tolerance}}}，坐标按实际调整）；'
                            "仅在确实要逐行校验时才用 row_expression。不要用 total 替代单点断言。"
                        )
                        fail(
                            f"{letter}{header_row+1}",total_value,expected,
                            reason="total_mismatch",data_range=f"{letter}{header_row+1}:{letter}{last_row}",
                            expected_matches_cell=expected_matches or None,
                            message=(
                                f"{letter} 列求和={total_value}，与 expected={expected} 不符（tolerance={tolerance}；"
                                f"区间 {letter}{header_row+1}:{letter}{last_row}，已排除空值 {skipped} 个与非数值文本 {excluded_count} 个）。"
                                f"{hint}"
                            ),
                        )
                else: raise _rule_arg_error(number,kind,ws.title,f"不支持的校验规则 {kind}",field="kind")
            entry={"rule":number,"kind":kind,"sheet":ws.title,"checked":checked,"failed":failed}
            if skipped or skipped_cells: entry.update(skipped=skipped,skipped_cells=skipped_cells)
            entry.update(extras)
            results.append(entry)
        # openpyxl returns ``date``/``datetime`` objects for date-formatted
        # cells.  Keep the tool value JSON-safe as well as the model text;
        # callers frequently persist ``result.value`` directly and should not
        # need a second serializer just because a failed rule included a date.
        def _json_safe(value):
            if hasattr(value, "isoformat") and not isinstance(value, (str, bytes)):
                try:
                    return value.isoformat()
                except (TypeError, ValueError):
                    pass
            if isinstance(value, dict):
                return {str(key): _json_safe(item) for key, item in value.items()}
            if isinstance(value, (list, tuple)):
                return [_json_safe(item) for item in value]
            return value

        results = _json_safe(results)
        failures = _json_safe(failures)
        uncalculated_cells = _json_safe(uncalculated_cells)
        return from_payload({
            "status":"partial" if validation_status != "complete" else "success",
            "file_path":file_path,
            "content_version":version,
            "valid":None if validation_status != "complete" else failure_count==0,
            "validation_status":validation_status,
            "uncalculated_cells":uncalculated_cells,
            "rules":results,
            "failure_count":failure_count,
            "failures":failures,
            "truncated":failure_count>len(failures),
            "source_versions":versions,
        })
    except FormulaValueError as exc:
        return error_result(
            str(exc),
            code=exc.code,
            fields={
                "cells":list(exc.cells or []),
                "validation_status":"blocked",
                "formula_cache":"missing",
                "executed":True,
                "committed":False,
            },
        )
    except (ValueError,KeyError,TypeError,OSError) as exc: return _failure(exc)
    except Exception as exc: return error_result(str(exc) or type(exc).__name__, code=getattr(exc,"code",None) or "TOOL_ERROR", fields=getattr(exc,"fields",None))
    finally:
        if book is not None: book.close()


def query_spreadsheet(sources, sql, output_path=None, max_rows=200, timeout_seconds=60, expected_output_version=None):
    """Stream source rows into a temporary SQLite database and query/export them."""
    try:
        if not isinstance(sources,list) or not 1<=len(sources)<=20 or not 1<=max_rows<=2000 or not 1<=timeout_seconds<=300: raise ValueError("sources/max_rows/timeout_seconds 无效")
        from excelmanus.workbook_commit import content_version_of_file
        from excelmanus.workspace.file_service import TargetSpec,ReadDependency,service_for_guard
        import re
        if not re.match(r"^\s*(SELECT|WITH)\b",sql,re.I): raise ValueError("仅支持 SELECT/WITH 查询")
        if output_path and Path(output_path).suffix.lower() not in {".xlsx",".csv"}: raise ValueError("输出只支持 xlsx 或 csv")
        started=time.monotonic(); dependencies=[]; summary=[]
        def check():
            _cancelled()
            if time.monotonic()-started>timeout_seconds: raise ValueError("数据查询超时")
        with tempfile.TemporaryDirectory(prefix="excelmanus-query-") as raw:
            root=Path(raw); connection=sqlite3.connect(root/"data.sqlite")
            try:
                for index,source in enumerate(sources):
                    alias=source.get("name",f"data{index+1}")
                    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*",alias): raise ValueError("源 name 仅支持字母数字下划线")
                    path=require_guard().resolve_and_validate(source["file_path"])
                    version=content_version_of_file(path)
                    if source.get("expected_version") and source["expected_version"]!=version: raise ValueError("源文件版本冲突")
                    # Snapshot on disk avoids holding the workbook ZIP and rows in RAM.
                    snapshot=root/(alias+path.suffix); shutil_copy(path,snapshot)
                    if content_version_of_file(snapshot)!=version: raise ValueError("读取期间源文件改变")
                    dependencies.append(ReadDependency(path.relative_to(require_guard().workspace_root).as_posix(),version))
                    book=None; stream=None; formulas=None; formula_iter=None
                    try:
                        if path.suffix.lower() in {".csv",".tsv"}:
                            stream=snapshot.open(encoding=source.get("encoding","utf-8-sig"),newline="")
                            iterator=csv.reader(stream,delimiter="\t" if path.suffix.lower()==".tsv" else ",")
                        else:
                            book=load_workbook(snapshot,read_only=True,data_only=True)
                            if not source.get("sheet") and len(book.worksheets)>1: raise ValueError("数据源需指定 sheet")
                            ws=book[source["sheet"]] if source.get("sheet") else book.active
                            iterator=ws.iter_rows(values_only=True)
                            formulas=load_workbook(snapshot,read_only=True,data_only=False)
                            formula_iter=formulas[ws.title].iter_rows()
                        header_row=int(source.get("header_row",1))
                        if header_row<1: raise ValueError("header_row 必须为正数")
                        for _ in range(header_row):
                            header=next(iterator)
                            if formulas: next(formula_iter)
                        headers=[str(v) if v is not None else "" for v in header]
                        if not all(headers) or len(set(headers))!=len(headers): raise ValueError("数据源表头必须非空且唯一")
                        selected=source.get("columns") or headers
                        positions=[headers.index(c) for c in selected]
                        quote=lambda v:'"'+v.replace('"','""')+'"'
                        connection.execute(f'CREATE TABLE "{alias}" ({",".join(quote(c) for c in selected)})')
                        command=f'INSERT INTO "{alias}" VALUES ({",".join("?" for _ in selected)})'
                        batch=[]; count=0
                        types=source.get("types") or {}
                        for row in iterator:
                            if formulas:
                                formula_row=next(formula_iter)
                                if any(formula_row[i].data_type=="f" and row[i] is None for i in positions): raise ValueError("源数据存在缺失公式缓存，请先重算")
                            values=[]
                            for c,i in zip(selected,positions):
                                v=row[i] if i<len(row) else None
                                if v is not None and types.get(c)=="number":
                                    from excelmanus.workbook.numeric import parse_number
                                    import math
                                    v=parse_number(v)
                                    if not math.isfinite(v): raise ValueError(f"{alias}.{c} 包含无法转换的数值")
                                elif v is not None and types.get(c)=="text": v=str(v)
                                elif hasattr(v,"isoformat"): v=v.isoformat()
                                values.append(v)
                            batch.append(values); count+=1
                            if len(batch)>=1000: check(); connection.executemany(command,batch); batch.clear()
                        if batch: connection.executemany(command,batch)
                        connection.commit(); summary.append({"name":alias,"file_path":source["file_path"],"content_version":version,"rows":count,"columns":selected})
                    finally:
                        if formula_iter is not None:
                            formula_iter = None
                        iterator = None
                        if book: book.close()
                        if formulas: formulas.close()
                        if stream: stream.close()
                allowed={sqlite3.SQLITE_SELECT,sqlite3.SQLITE_READ,sqlite3.SQLITE_FUNCTION,sqlite3.SQLITE_RECURSIVE}
                def authorizer(action,a,b,db,trigger):
                    if action==sqlite3.SQLITE_FUNCTION and str(b or a).lower() in {"load_extension","readfile","writefile"}: return sqlite3.SQLITE_DENY
                    return sqlite3.SQLITE_OK if action in allowed else sqlite3.SQLITE_DENY
                connection.set_authorizer(authorizer)
                def progress():
                    try: check(); return 0
                    except Exception: return 1
                connection.set_progress_handler(progress,10000)
                cursor=connection.execute(sql)
                columns=[d[0] for d in cursor.description]; preview=[]; total=0
                workbook=None; csv_stream=None; writer=None
                artifact=root/"result.xlsx" if output_path and output_path.endswith(".xlsx") else root/"result.csv"
                if output_path:
                    if artifact.suffix==".xlsx": workbook=Workbook(write_only=True); ws=workbook.create_sheet("Result"); ws.append(columns)
                    else: csv_stream=artifact.open("w",encoding="utf-8-sig",newline=""); writer=csv.writer(csv_stream); writer.writerow(columns)
                try:
                    for row in cursor:
                        total+=1
                        if len(preview)<max_rows: preview.append(list(row))
                        if workbook:
                            from openpyxl.cell import WriteOnlyCell
                            items=[]
                            for value in row:
                                c=WriteOnlyCell(ws,value=value)
                                if isinstance(value,str): c.data_type="s"
                                items.append(c)
                            if total>=1048576: raise ValueError("结果超出 Excel 最大行数，请输出 CSV")
                            ws.append(items)
                        if writer: writer.writerow(row)
                        if total%1000==0: check()
                    if workbook: workbook.save(artifact)
                finally:
                    if csv_stream: csv_stream.close()
                    if workbook: workbook.close()
                payload={"status":"success","columns":columns,"values":preview,"total_rows":total,"truncated":total>len(preview),"sources":summary,"engine":"sqlite","storage":"disk","elapsed_seconds":round(time.monotonic()-started,3)}
                if output_path:
                    dest=require_guard().resolve_and_validate(output_path)
                    if any(dest==require_guard().workspace_root/dep.path for dep in dependencies): raise ValueError("query 输出不能覆盖数据源")
                    if dest.exists() and expected_output_version is None: raise ValueError("输出已存在，需要 expected_output_version")
                    check()
                    receipt=service_for_guard(require_guard()).apply_batch([TargetSpec("update" if dest.exists() else "create",str(dest.relative_to(require_guard().workspace_root)),data=artifact.read_bytes(),expected_version=expected_output_version)],read_dependencies=dependencies)
                    payload.update(file_path=receipt.primary_path(),content_version=receipt.primary_version(),receipt=receipt.to_dict())
                return from_payload(payload)
            finally: connection.close()
    except (ValueError,KeyError,TypeError,OSError,sqlite3.Error,StopIteration) as exc: return _failure(exc)


def shutil_copy(src,dest):
    import shutil
    shutil.copyfile(src,dest)


def get_tools():
    return [ToolDef(name="validate_spreadsheet",description=TOOL_DESCRIPTIONS["validate_spreadsheet"],func=validate_spreadsheet,write_effect="none",
        input_schema={
            "type":"object",
            "additionalProperties":False,
            "properties":{
                "file_path":{"type":"string"},
                "sheet":{"type":"string","description":"默认工作表；多工作表工作簿中除 formula_errors 外每条规则都需要（规则内 sheet 可覆盖）。"},
                "expected_version":{"type":"string"},
                "max_failures":{"type":"integer","minimum":1,"maximum":2000},
                "allow_uncached":{"type":"boolean","description":"公式缓存缺失时是否允许 partial/valid=null；默认 true，不声称业务校验已通过。"},
                "rules":{
                    "type":"array",
                    "minItems":1,
                    "items":{
                        "type":"object",
                        "properties":{
                            "kind":{
                                "type":"string",
                                "enum":["unique","required","foreign_key","row_expression","total","cell","formula_errors"],
                                "description":"规则类型与所需字段：unique/required 需要 columns（或 column）；foreign_key 需要 columns + reference；row_expression 需要 expression（当前行列级断言，不支持 A1 单元格引用）；total 需要 column + expected（可选 tolerance；校验整列求和，空值跳过、非数值文本记为 excluded_non_numeric，失败项 actual/expected 只放数值或 null）；cell 需要 cell + expected（单点断言：定位到 A1 单格，数值按 tolerance、文本精确比对、expected=null 断言空值，可选 tolerance/case_sensitive；多表工作簿必须给 sheet）；formula_errors 不需要其他字段（检查公式错误值与缺失缓存）。",
                            },
                            "sheet":{"type":"string","description":"该规则作用的工作表，覆盖顶层 sheet。"},
                            "header_row":{"type":"integer","minimum":1,"description":"表头行号（1-based），默认 1。"},
                            "columns":{"type":"array","items":{"type":["string","integer"]},"description":"多列引用。每个元素是列名、列字母（如 \"A\"、\"AA\"）或 1-based 列号；与表头列名同名时列名优先。unique/required/foreign_key 用。"},
                            "column":{"type":["string","integer"],"description":"单列引用：列名、列字母（如 \"C\"）或 1-based 列号；与表头列名同名时列名优先。total 的合计列，或 unique/required 的单列写法。"},
                            "cell":{"type":"string","description":"cell 规则的 A1 单格坐标（如 \"B7\"、\"$B$7\"；列字母+行号）：只接受单格，不接受区域（B7:C9）、纯列（B）或纯行（7），也不接受 \"表名!B7\"（工作表写在 sheet 字段）。越界（B0、XFE1、A1048577）或非法坐标返回 INVALID_ARGS。"},
                            "case_sensitive":{"type":"boolean","description":"cell 规则文本比对是否区分大小写，默认 true（数值/布尔/null 比对不受影响）。"},
                            "expression":{"type":"string","description":"row_expression 的布尔表达式：引用当前行的列名或列字母（如 金额、E；{row} 是可选占位符，E{row} 等价于 E），支持 + - * / %、== != < <= > >=、and/or/not、括号和大小写不敏感的 abs/round/min/max。不支持 A1 单元格引用：B7 会被当成列名而不是单元格，要断言单个单元格的值请用 cell 规则（cell + expected）或 observe_spreadsheet 回读，或把期望常量写进同行数据列再断言该列。示例：abs(金额-数量*单价)<0.01；ABS(E{row}+F{row}+G{row}-C{row})<0.05"},
                            "expected":{"type":["number","string","boolean","null"],"description":"期望值。数值 expected 必须有限（不接受 NaN/Infinity 或超出浮点范围的数值）。total：数值合计（整列求和，不是某个单元格的值）。cell：数值（按 tolerance 容差比对，默认 1e-9；千分位/万/亿/百分号等按与 total 相同的数值解析，数值型文本如 CSV 导入的 \"6.7\" 也能按数值比较）/文本（精确比对，见 case_sensitive；不做类型转换：单元格是数值而 expected 是文本会返回 reason=type_mismatch）/布尔/null（null 表示断言该格为空）。"},
                            "tolerance":{"type":"number","minimum":0,"description":"total/cell 允许的绝对误差，默认 1e-9；必须是有限数（不接受 NaN/Infinity），且非负；cell 只作用于数值 expected。"},
                            "reference":{"type":"object","description":"foreign_key 的对照表：{file_path, columns（列名/列字母/1-based 列号）, sheet?, header_row?, expected_version?}。"},
                        },
                        "required":["kind"],
                        "additionalProperties":False,
                    },
                },
            },
            "required":["file_path","rules"],
        },max_result_chars=8000),
            ToolDef(name="query_spreadsheet",description="将 Excel/CSV 按行流入磁盘 SQLite，执行只读 SQL；支持多表 join/union/计算列/窗口统计。返回有限预览，可将全部结果原子导出 xlsx/csv，源公式必须有缓存。",func=query_spreadsheet,write_effect="workspace_write",input_schema={"type":"object","additionalProperties":False,"properties":{"sources":{"type":"array","minItems":1,"maxItems":20,"items":{"type":"object","properties":{"file_path":{"type":"string"},"name":{"type":"string"},"sheet":{"type":"string"},"header_row":{"type":"integer","minimum":1},"columns":{"type":"array","items":{"type":"string"}},"types":{"type":"object","additionalProperties":{"enum":["number","text"]}},"encoding":{"type":"string"},"expected_version":{"type":"string"}},"required":["file_path"],"additionalProperties":False}},"sql":{"type":"string","description":"SELECT/WITH；源表名默认 data1/data2，可用 sources.name 指定。列名建议用双引号。"},"output_path":{"type":"string"},"expected_output_version":{"type":"string"},"max_rows":{"type":"integer","minimum":1,"maximum":2000},"timeout_seconds":{"type":"integer","minimum":1,"maximum":300}},"required":["sources","sql"]},max_result_chars=8000)]
