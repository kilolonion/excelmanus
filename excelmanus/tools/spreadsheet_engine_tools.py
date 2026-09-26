"""Explicit spreadsheet calculation, rendering and conversion with CAS publication."""
from __future__ import annotations
from io import BytesIO
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from zipfile import ZipFile
import csv
import html

from openpyxl import load_workbook, Workbook

from excelmanus.engine_core.tool_result import from_payload, error_result
from excelmanus.tools.context import require_guard, bind_workspace, operation_id_for
from excelmanus.tools.registry import ToolDef
from excelmanus.workbook_commit import content_version_of, CommitError


def init_guard(workspace_root):
    bind_workspace(workspace_root)


def _source(path, expected_version=None):
    from excelmanus.workbook.snapshot import open_snapshot
    source = require_guard().resolve_and_validate(path)
    snapshot = open_snapshot(path, expected_version=expected_version)
    return source, snapshot.read_bytes(), snapshot.content_version


def _cancelled():
    from excelmanus.tools.runtime import current_execution
    execution=current_execution()
    if execution and execution.cancel_event.is_set():
        raise CommitError("CANCELLED","操作已取消，未发布结果")


def _run(args, timeout=90):
    from excelmanus.tools.runtime import current_execution, register_killable_process, unregister_killable_process, _terminate_process_tree
    _cancelled()
    execution=current_execution(); execution_id=getattr(execution,"execution_id",None)
    flags=(subprocess.CREATE_NO_WINDOW|subprocess.CREATE_NEW_PROCESS_GROUP) if os.name=="nt" else 0
    process=subprocess.Popen(args,stdin=subprocess.DEVNULL,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,encoding="utf-8",errors="replace",creationflags=flags,start_new_session=os.name!="nt")
    register_killable_process(execution_id,process)
    try:
        try: out,err=process.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            _terminate_process_tree(process)
            try: process.communicate(timeout=3)
            except (OSError,subprocess.SubprocessError): pass
            raise CommitError("ENGINE_TIMEOUT",f"转换超过 {timeout} 秒")
        _cancelled()
        if process.returncode: raise CommitError("ENGINE_FAILED",(err or out or "转换失败")[:1000])
        return out
    finally:
        unregister_killable_process(execution_id,process)


def _office_convert(source, destination, extension, *, timeout=90):
    from excelmanus.runtime_capabilities import office_executable
    engine=office_executable()
    if not engine: raise CommitError("ENGINE_UNAVAILABLE","未找到 LibreOffice/soffice；未生成或修改文件")
    profile=destination.parent/"profile"; profile.mkdir(exist_ok=True)
    _run([engine,"--headless","--nolockcheck","--nodefault","--nologo","--nofirststartwizard",f"-env:UserInstallation={profile.as_uri()}","--convert-to",extension,"--outdir",str(destination),str(source)],timeout)
    target=destination/(source.stem+"."+extension.split(":")[0])
    if not target.is_file() or not target.stat().st_size: raise CommitError("ENGINE_FAILED","引擎未生成有效产物")
    return target,engine


def _publish(source, source_version, outputs, *, expected_output_version=None):
    from excelmanus.workspace.file_service import TargetSpec, ReadDependency, service_for_guard
    guard=require_guard(); specs=[]
    for path,data in outputs:
        dest=guard.resolve_and_validate(path)
        rel=dest.relative_to(guard.workspace_root).as_posix()
        if dest==source and len(outputs)>1: raise ValueError("多产物输出不能覆盖源文件")
        version=source_version if dest==source else expected_output_version if len(outputs)==1 else None
        if dest.exists() and version is None: raise CommitError("VERSION_CONFLICT",f"输出已存在，请提供 expected_output_version 或使用新路径：{rel}")
        specs.append(TargetSpec("update" if dest.exists() else "create",rel,data=data,expected_version=version))
    _cancelled()
    receipt=service_for_guard(guard).apply_batch(specs,operation_id=operation_id_for("spreadsheet_engine"),read_dependencies=[ReadDependency(source.relative_to(guard.workspace_root).as_posix(),source_version)])
    return receipt.to_dict()


def _failure(exc):
    return error_result(str(exc),code=getattr(exc,"code","INVALID_ARGS"),fields=getattr(exc,"fields",None))


def calculate_spreadsheet(file_path, expected_version=None, output_path=None, expected_output_version=None, allow_formula_errors=False):
    try:
        source,data,version=_source(file_path,expected_version)
        from excelmanus.workbook_commit import recalculate_workbook_bytes
        output,info=recalculate_workbook_bytes(data,suffix=source.suffix)
        if info["status"]!="recalculated":
            return error_result("公式重算未完成",code="CALCULATION_UNAVAILABLE" if info["status"] in {"unavailable","unsupported_format","disabled"} else "CALCULATION_FAILED",fields={"formula_recalculation":info,"source_version":version,"committed":False})
        if info.get("errors") and not allow_formula_errors:
            return error_result(
                "重算发现单元格公式错误，验收未通过，未发布结果",
                code="FORMULA_ERRORS",
                fields={
                    "file_path": file_path,
                    "formula_recalculation": info,
                    "source_version": version,
                    "validation_status": "failed",
                    "committed": False,
                },
            )
        receipt=_publish(source,version,[(output_path or file_path,output)],expected_output_version=expected_output_version)
        return from_payload({"status":"success","file_path":receipt["published_paths"][0],"content_version":receipt["targets"][0]["after_version"],"source_version":version,"formula_recalculation":info,"receipt":receipt})
    except (ValueError,OSError,CommitError) as exc: return _failure(exc)


def render_spreadsheet(file_path, output_path, sheet=None, range=None, format="pdf", expected_version=None, expected_output_version=None, dpi=120, max_pages=50):
    try:
        if format not in {"pdf","png"}: raise ValueError("format 必须是 pdf 或 png")
        if not isinstance(dpi,int) or not 72<=dpi<=300 or not isinstance(max_pages,int) or not 1<=max_pages<=200: raise ValueError("dpi/max_pages 超出范围")
        source,data,version=_source(file_path,expected_version)
        if source.suffix.lower() not in {".xlsx",".xlsm",".xltx"}: raise ValueError("请先转换为 xlsx 再渲染")
        if Path(output_path).suffix.lower()!="."+format: raise ValueError("output_path 扩展名与 format 不符")
        with tempfile.TemporaryDirectory(prefix="excelmanus-render-") as raw:
            root=Path(raw); src=root/"input.xlsx"; out=root/"converted"; out.mkdir()
            if sheet or range:
                from excelmanus.workbook.ooxml import render_selection
                src.write_bytes(render_selection(data, sheet, range))
            else: src.write_bytes(data)
            pdf,engine=_office_convert(src,out,"pdf:calc_pdf_Export")
            page_count=None
            info_bin=shutil.which("pdfinfo")
            if info_bin:
                match=re.search(r"Pages:\s*(\d+)",_run([info_bin,str(pdf)],30))
                if match: page_count=int(match[1])
            if page_count is not None and page_count > max_pages:
                raise CommitError(
                    "LIMIT_EXCEEDED",
                    f"共 {page_count} 页，超过 max_pages={max_pages}；请将 max_pages 调整为至少 {page_count}，或缩小 range。",
                    fields={
                        "page_count": page_count,
                        "max_pages": max_pages,
                        "suggested_max_pages": page_count,
                        "format": format,
                        "committed": False,
                        "remediation": "缩小 range，或在 1–200 范围内将 max_pages 提高到实际页数；只检查单页可用 preview_spreadsheet(surface='print', page=...)。",
                    },
                )
            if format=="pdf": outputs=[(output_path,pdf.read_bytes())]
            else:
                executable=shutil.which("pdftoppm")
                if not executable: raise CommitError("ENGINE_UNAVAILABLE","PNG 导出需要 pdftoppm")
                if page_count is None: raise CommitError("ENGINE_UNAVAILABLE","PNG 分页导出需要 pdfinfo 获取实际页数")
                prefix=root/"page"
                _run([executable,"-png","-r",str(dpi),str(pdf),str(prefix)],90)
                pages=sorted(root.glob("page-*.png"),key=lambda p:int(p.stem.split("-")[-1]))
                if len(pages)!=page_count: raise CommitError("ENGINE_FAILED","生成页数与 PDF 不符")
                dest=Path(output_path)
                outputs=[(output_path if len(pages)==1 else str(dest.with_name(f"{dest.stem}-{i+1}{dest.suffix}")),p.read_bytes()) for i,p in enumerate(pages)]
            receipt=_publish(source,version,outputs,expected_output_version=expected_output_version)
            return from_payload({"status":"success","files":receipt["published_paths"],"source_version":version,"engine":engine,"format":format,"page_count":page_count,"sheet":sheet,"range":range,"fidelity":"LibreOffice rendering; font substitution not measured","receipt":receipt})
    except (ValueError,KeyError,OSError,CommitError) as exc: return _failure(exc)


def _inventory(data):
    try:
        with ZipFile(BytesIO(data)) as z:
            patterns={"charts":"xl/charts/chart", "images":"xl/media/", "pivot_tables":"xl/pivotTables/pivotTable", "vba":"xl/vbaProject.bin", "tables":"xl/tables/table", "connections":"xl/connections.xml", "external_links":"xl/externalLinks/externalLink"}
            return {key:sum(n.startswith(prefix) and not n.endswith(".rels") for n in z.namelist()) for key,prefix in patterns.items()}
    except Exception: return None


def _export_values(data: bytes, output_path: str, *, sheet: str | None, source: Path, version: str, expected_output_version: str | None, data_only: bool = False):
    """Export a single worksheet as CSV or self-contained HTML."""
    book = load_workbook(BytesIO(data), data_only=data_only, read_only=False)
    try:
        selected = sheet or book.active.title
        if selected not in book.sheetnames:
            raise ValueError(f"工作表不存在: {selected}")
        ws = book[selected]
        rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row or 1, min_col=1, max_col=ws.max_column or 1)]
        ext = Path(output_path).suffix.lower()
        if ext == ".csv":
            stream = BytesIO()
            import io
            text = io.StringIO(newline="")
            writer = csv.writer(text)
            writer.writerows(rows)
            payload = text.getvalue().encode("utf-8-sig")
        else:
            cells = []
            for row in rows:
                cells.append("<tr>" + "".join(f"<td>{html.escape('' if value is None else str(value))}</td>" for value in row) + "</tr>")
            drawings = []
            import base64
            import mimetypes
            for image in getattr(ws, "_images", []) or []:
                raw = image._data() if callable(getattr(image, "_data", None)) else b""
                if raw:
                    fmt = str(getattr(image, "format", "png") or "png").lower()
                    media = {"jpg": "image/jpeg", "jpeg": "image/jpeg"}.get(fmt, f"image/{fmt}")
                    anchor = getattr(image, "anchor", None)
                    target = getattr(getattr(anchor, "_from", None), "row", 0) + 1
                    col = getattr(getattr(anchor, "_from", None), "col", 0) + 1
                    drawings.append(f"<figure data-cell='{target},{col}'><img alt='embedded image' src='data:{media};base64,{base64.b64encode(raw).decode()}' style='max-width:{int(getattr(image, 'width', 240) or 240)}px;max-height:{int(getattr(image, 'height', 160) or 160)}px'></figure>")
            for chart in getattr(ws, "_charts", []) or []:
                anchor = getattr(chart, "anchor", None)
                target = getattr(getattr(anchor, "_from", None), "row", 0) + 1
                col = getattr(getattr(anchor, "_from", None), "col", 0) + 1
                drawings.append(f"<figure data-chart-cell='{target},{col}'><figcaption>{html.escape(type(chart).__name__)}</figcaption><div class='chart-placeholder'>Embedded Excel chart ({len(getattr(chart, 'series', []) or [])} series)</div></figure>")
            payload = ("<!doctype html><meta charset='utf-8'><title>" + html.escape(selected) + "</title>"
                       + "<style>table{border-collapse:collapse;font:13px sans-serif}td{border:1px solid #d1d5db;padding:4px 8px;white-space:nowrap}</style>"
                       + "<style>figure{display:inline-block;vertical-align:top;margin:12px}.chart-placeholder{border:1px solid #94a3b8;padding:28px;color:#475569}</style>"
                       + "<table data-sheet='" + html.escape(selected, quote=True) + "'>" + "".join(cells) + "</table>" + "".join(drawings)).encode("utf-8")
    finally:
        book.close()
    receipt = _publish(source, version, [(output_path, payload)], expected_output_version=expected_output_version)
    return from_payload({"status": "success", "file_path": receipt["published_paths"][0], "source_version": version, "format": Path(output_path).suffix.lower().lstrip("."), "sheet": selected, "receipt": receipt})


def convert_spreadsheet(file_path, output_path, mode="preserve", expected_version=None, expected_output_version=None, sheet=None):
    try:
        if mode not in {"preserve","data_only"}: raise ValueError("mode 必须为 preserve 或 data_only")
        source,data,version=_source(file_path,expected_version)
        target_ext = Path(output_path).suffix.lower()
        if target_ext in {".pdf", ".png"}:
            return render_spreadsheet(file_path, output_path, sheet=sheet, format=target_ext.lstrip("."), expected_version=expected_version, expected_output_version=expected_output_version)
        if target_ext in {".csv", ".html", ".htm"}:
            if target_ext == ".htm":
                output_path = str(Path(output_path).with_suffix(".html"))
            return _export_values(data, output_path, sheet=sheet, source=source, version=version, expected_output_version=expected_output_version, data_only=mode == "data_only")
        if target_ext == ".ods":
            with tempfile.TemporaryDirectory(prefix="excelmanus-ods-") as raw:
                root = Path(raw); src = root / ("input" + source.suffix); src.write_bytes(data); out = root / "out"; out.mkdir()
                converted, engine = _office_convert(src, out, "ods")
                receipt = _publish(source, version, [(output_path, converted.read_bytes())], expected_output_version=expected_output_version)
            return from_payload({"status": "success", "file_path": receipt["published_paths"][0], "source_version": version, "format": "ods", "engine": engine, "receipt": receipt})
        if target_ext != ".xlsx": raise ValueError("转换输出支持 .xlsx、.csv、.html、.ods、.pdf 或 .png")
        if require_guard().resolve_and_validate(output_path)==source: raise ValueError("转换须输出到新文件以保留源文件")
        warnings=[]; before=_inventory(data)
        with tempfile.TemporaryDirectory(prefix="excelmanus-convert-") as raw:
            root=Path(raw); src=root/("input"+source.suffix); src.write_bytes(data); out=root/"out"; out.mkdir()
            if mode=="preserve":
                converted,engine=_office_convert(src,out,"xlsx")
                output=converted.read_bytes()
                warnings.append("引擎转换不保证所有 Excel 扩展对象或像素布局相同；请查看对象数量差异")
                if source.suffix.lower() in {".xlsm",".xlsb",".xls"}: warnings.append("xlsx 不能保留 VBA；源格式中的宏需保留在原件")
            else:
                wb=Workbook(); wb.remove(wb.active)
                suffix=source.suffix.lower()
                if suffix==".xls":
                    import xlrd
                    book=xlrd.open_workbook(file_contents=data)
                    try:
                        for sheet in book.sheets():
                            ws=wb.create_sheet(sheet.name)
                            for row in sheet.get_rows(): ws.append([c.value for c in row])
                    finally: book.release_resources()
                    engine="xlrd"
                elif suffix==".xlsb":
                    from pyxlsb import open_workbook
                    with open_workbook(src) as book:
                        for name in book.sheets:
                            ws=wb.create_sheet(name)
                            with book.get_sheet(name) as sheet:
                                for row in sheet.rows():
                                    for cell in row: ws.cell(cell.r+1,cell.c+1).value=cell.v
                    engine="pyxlsb"
                elif suffix in {".xlsx",".xlsm",".xltx"}:
                    book=load_workbook(BytesIO(data),data_only=False)
                    from excelmanus.workbook.formula_values import attach_formula_source,analysis_rows
                    attach_formula_source(book,data)
                    try:
                        for sheet in book:
                            ws=wb.create_sheet(sheet.title)
                            for row in analysis_rows(sheet):
                                ws.append(row)
                                for cell in ws[ws.max_row]:
                                    if isinstance(cell.value,str): cell.data_type="s"
                    finally: book.close()
                    engine="openpyxl"
                else: raise ValueError("data_only 支持 xls/xlsb/xlsx/xlsm/xltx")
                stream=BytesIO(); wb.save(stream); wb.close(); output=stream.getvalue()
                warnings.append("仅保留单元格值和工作表；公式、样式、宏、图表、图片、验证、条件格式和连接不迁移")
            check=load_workbook(BytesIO(output),read_only=True); check.close()
            after=_inventory(output)
            receipt=_publish(source,version,[(output_path,output)],expected_output_version=expected_output_version)
            return from_payload({"status":"success","file_path":receipt["published_paths"][0],"content_version":receipt["targets"][0]["after_version"],"source_version":version,"engine":engine,"mode":mode,"loss_report":{"before":before,"after":after,"decreased":{k:before[k]-after.get(k,0) for k in before if before[k]>after.get(k,0)} if before and after else None,"unknown_source_objects":before is None,"warnings":warnings},"receipt":receipt})
    except (ValueError,KeyError,OSError,CommitError) as exc: return _failure(exc)


def get_tools():
    common={"file_path":{"type":"string"},"output_path":{"type":"string"},"expected_version":{"type":"string"},"expected_output_version":{"type":"string"}}
    def make(name,description,func,extra,required):
        return ToolDef(name=name,description=description,func=func,write_effect="workspace_write",input_schema={"type":"object","additionalProperties":False,"properties":{**common,**extra},"required":required},max_result_chars=6000)
    return [make("calculate_spreadsheet","显式重算并检查公式错误；需要可用 LibreOffice。失败不发布，返回计算引擎、缓存结果与源版本。",calculate_spreadsheet,{"allow_formula_errors":{"type":"boolean"}},["file_path"]),
            make("render_spreadsheet","渲染指定工作表/区域为 PDF 或分页 PNG，返回实际页数、引擎和源版本；需要 soffice，PNG 另需 Poppler。",render_spreadsheet,{"sheet":{"type":"string"},"range":{"type":"string"},"format":{"type":"string","enum":["pdf","png"]},"dpi":{"type":"integer","minimum":72,"maximum":300},"max_pages":{"type":"integer","minimum":1,"maximum":200}},["file_path","output_path"]),
            make("convert_spreadsheet","将 Excel 转为 xlsx、CSV、HTML、ODS、PDF 或 PNG；xlsx preserve 使用 LibreOffice，data_only 仅迁移值；返回图片/图表对象损失报告。",convert_spreadsheet,{"mode":{"type":"string","enum":["preserve","data_only"]},"sheet":{"type":"string"}},["file_path","output_path"])]
