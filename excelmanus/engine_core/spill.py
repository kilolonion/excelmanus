"""大对象 spill + 写后语义校验。

Wire / durable 不再内联超大工具结果：正文进 workspace 内 content-addressed
仓库，模型只看到摘要 + opaque locator。取回是内部函数；dispatcher 把既有
读工具的 ``file_path`` 识别为句柄后走同一通路，不新增工具名。

写后校验产出有界的 ``meta.write_verification``（值/公式分列），替代
max_row 维度假象。``write_operations_log`` 仍是会话级索引，不平行再造一份
单元格日志。
"""

from __future__ import annotations

import datetime
import hashlib
import json
import os
import re
import tempfile
from collections.abc import Mapping
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Literal

from excelmanus.engine_core.error_payload import (
    FAILURE_INTERNAL,
    FAILURE_NOT_FOUND,
    NOT_FOUND,
    RESULT_UNCERTAIN,
    SHEET_NOT_FOUND,
    TOOL_ERROR,
    make_error_payload,
)
from excelmanus.engine_core.tool_result import ToolResult, error_result
from excelmanus.logger import get_logger

logger = get_logger("spill")

# ── 阈值 ──────────────────────────────────────────────────
# 默认硬截断是 12000 字符（config.tool_result_hard_cap_chars）。spill 必须
# 独立、且默认低于该值，这样超大结果先外置再截断，句柄不会被砍掉。
# 8000 字符 ≈ 一次 inspect 预览的数倍、远小于整表 dump；2000 token 按
# utf-8 字节/4 估算，与 ASCII 8000 字符同量级，CJK 会更早外置（更安全）。
DEFAULT_SPILL_CHAR_THRESHOLD = 8000
DEFAULT_SPILL_BYTE_THRESHOLD = 24_000
DEFAULT_SPILL_TOKEN_THRESHOLD = 2000
DEFAULT_PREVIEW_CHARS = 400
DEFAULT_PREVIEW_LINES = 12

# 写后校验：最多展示 20 条单元格，采样时最多走 64 个意图坐标。
# 不扫整表 used range；只读 operations 声明过的格子。
MAX_WRITE_VERIFY_ENTRIES = 20
MAX_WRITE_VERIFY_SAMPLE = 64

SPILL_DIR_REL = ".excelmanus/spill"
LOCATOR_PREFIX = "spill:"
_LOCATOR_RE = re.compile(r"^(?:spill|result_spill|selection_spill):(?P<digest>[0-9a-f]{64})$")
_REFERENCE_PREFIXES = ("spill:", "result_spill:", "selection_spill:")
_HOST_PATH_RE = re.compile(r"(?:^[A-Za-z]:\\)|(?:/Users/)|(?:/home/)")
_CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]{0,6})$")

ChangeKind = Literal["value", "formula"]


class SpillNotFound(KeyError):
    """句柄不存在或已损坏。"""


class SpillLocator(str):
    """Opaque content-addressed 句柄，形如 ``spill:<sha256>``。

    故意不携带任何文件系统路径。创建后不可再拼接目录。
    """

    __slots__ = ()

    def hex_digest(self) -> str:
        return str(self)[len(LOCATOR_PREFIX) :]


def is_spill_locator(text: str | None) -> bool:
    raw = str(text or "").strip()
    if not _LOCATOR_RE.fullmatch(raw):
        return False
    return _HOST_PATH_RE.search(raw) is None


def is_spill_reference(text: str | None) -> bool:
    """Recognize reserved handles, including malformed ones, before path resolution."""
    return str(text or "").strip().startswith(_REFERENCE_PREFIXES)


def parse_locator(text: str) -> SpillLocator:
    raw = str(text or "").strip()
    match = _LOCATOR_RE.fullmatch(raw)
    if match is None:
        raise ValueError(f"非法 spill 句柄: {text!r}")
    # Output keys are often copied as prefixes. Accept those exact aliases but
    # keep one storage/return format; never interpret the digest as a file path.
    return SpillLocator(LOCATOR_PREFIX + match.group("digest"))


def estimate_tokens(text: str) -> int:
    """无 tokenizer 时的保守估算：utf-8 字节 / 4（CJK 偏高 → 更早 spill）。"""
    raw = text if isinstance(text, str) else str(text or "")
    return max(1, (len(raw.encode("utf-8")) + 3) // 4) if raw else 0


def should_spill(
    text: str,
    *,
    char_threshold: int = DEFAULT_SPILL_CHAR_THRESHOLD,
    byte_threshold: int = DEFAULT_SPILL_BYTE_THRESHOLD,
    token_threshold: int = DEFAULT_SPILL_TOKEN_THRESHOLD,
) -> bool:
    raw = text if isinstance(text, str) else str(text or "")
    if not raw:
        return False
    if char_threshold > 0 and len(raw) > char_threshold:
        return True
    data = raw.encode("utf-8")
    if byte_threshold > 0 and len(data) > byte_threshold:
        return True
    if token_threshold > 0 and estimate_tokens(raw) > token_threshold:
        return True
    return False


def extract_spill_locator(arguments: dict[str, Any] | None) -> str | None:
    """从工具参数里取出 spill 句柄（优先 file_path / path）。"""
    if not isinstance(arguments, dict):
        return None
    for key in ("file_path", "path", "locator", "spill", "result_spill", "selection_spill"):
        value = arguments.get(key)
        if isinstance(value, str) and is_spill_reference(value):
            return value.strip()
        if key in {"file_path", "path"} and value:
            return None
    # A search query/cell value that looks like a handle is still user data.
    return None


@dataclass(frozen=True)
class SpillProjection:
    model_text: str
    locator: SpillLocator | None
    spilled: bool
    chars: int
    bytes: int
    estimated_tokens: int
    kind: str = "text"
    preview: str = ""

    def summary_payload(self) -> dict[str, Any]:
        payload: dict[str, Any] = {
            "spilled": self.spilled,
            "spill": self.locator,
            "kind": self.kind,
            "chars": self.chars,
            "bytes": self.bytes,
            "tokens": self.estimated_tokens,
            "preview": self.preview,
            "retrieve": (
                "按 next_call 取回原文；将 spill 字段的完整值原样传入 file_path，"
                "不要把字段名拼到哈希前。句柄不是磁盘路径。"
            ),
            "next_call": {"tool": "read_text_file", "arguments": {"file_path": str(self.locator)}},
        }
        return payload


class SpillStore:
    """workspace 内 content-addressed 外置仓库。

    目录：``{workspace}/.excelmanus/spill/{sha256}``。
    句柄只有 ``spill:{sha256}``，文件名也只有 hex，从不写入绝对路径。
    """

    def __init__(self, workspace_root: str | Path) -> None:
        root = Path(workspace_root).expanduser()
        self._workspace_root = root
        self._dir = root / ".excelmanus" / "spill"

    @property
    def directory(self) -> Path:
        return self._dir

    def put(self, content: str) -> SpillLocator:
        data = (content if isinstance(content, str) else str(content or "")).encode("utf-8")
        digest = hashlib.sha256(data).hexdigest()
        locator = SpillLocator(f"{LOCATOR_PREFIX}{digest}")
        self._dir.mkdir(parents=True, exist_ok=True)
        dest = self._dir / digest
        if dest.is_file() and dest.stat().st_size == len(data):
            return locator
        fd, tmp_name = tempfile.mkstemp(prefix=".spill_", dir=str(self._dir))
        try:
            with os.fdopen(fd, "wb") as handle:
                handle.write(data)
            os.replace(tmp_name, str(dest))
        except BaseException:
            try:
                os.unlink(tmp_name)
            except OSError:
                pass
            raise
        return locator

    def get(self, locator: str) -> str:
        parsed = parse_locator(locator)
        path = self._dir / parsed.hex_digest()
        try:
            data = path.read_bytes()
        except FileNotFoundError as exc:
            raise SpillNotFound(str(parsed)) from exc
        except OSError as exc:
            raise SpillNotFound(str(parsed)) from exc
        actual = hashlib.sha256(data).hexdigest()
        if actual != parsed.hex_digest():
            raise SpillNotFound(str(parsed))
        return data.decode("utf-8")

    def exists(self, locator: str) -> bool:
        try:
            parsed = parse_locator(locator)
        except ValueError:
            return False
        return (self._dir / parsed.hex_digest()).is_file()


def retrieve_spill(locator: str, *, workspace_root: str | Path) -> str:
    """按句柄取回原文。内部通路；不要新增面向用户的工具名。"""
    return SpillStore(workspace_root).get(locator)


def retrieve_spill_result(locator: str, *, workspace_root: str | Path) -> ToolResult:
    """给 dispatcher 复用读工具时的 ToolResult 投影。"""
    try:
        text = retrieve_spill(locator, workspace_root=workspace_root)
    except ValueError:
        return replace(error_result(
            f"结果句柄格式无效：{locator}", code="INVALID_ARGS",
            fields={
                "locator": str(locator),
                "remediation": "原样复制返回的 spill/result_spill/selection_spill 字段值（spill: 加 64 位小写十六进制摘要），用 read_text_file(file_path=该值)；不要拼接字段名或目录。",
            },
        ), coverage={"spill_retrieve": True, "kind": "invalid"})
    except SpillNotFound:
        return replace(error_result(
            f"spill 句柄不存在或已失效：{locator}",
            code=NOT_FOUND,
            fields={
                "locator": str(parse_locator(locator)),
                "remediation": "该结果句柄在当前工作区不存在或内容已损坏。重新执行产生它的只读查询以取得新句柄；不要列目录寻找句柄，也不要重放写入操作。",
            },
        ), coverage={"spill_retrieve": True, "kind": "missing"})
    parsed: Any = text
    stripped = text.strip()
    if (stripped.startswith("{") and stripped.endswith("}")) or (stripped.startswith("[") and stripped.endswith("]")):
        try:
            loaded = json.loads(stripped)
        except (json.JSONDecodeError, TypeError, ValueError):
            loaded = None
        if isinstance(loaded, (dict, list)):
            parsed = loaded
    return ToolResult(
        success=True,
        model_text=text,
        value=parsed,
        coverage={"spill_retrieve": True, "kind": "complete"},
    )


def _preview_text(text: str, *, max_chars: int, max_lines: int) -> tuple[str, int, str]:
    lines = text.splitlines()
    kind = "text"
    stripped = text.strip()
    if stripped.startswith("{") or stripped.startswith("["):
        try:
            json.loads(stripped)
            kind = "json"
        except (json.JSONDecodeError, TypeError, ValueError):
            kind = "text"
    if len(lines) > max_lines:
        chunk = "\n".join(lines[:max_lines])
        if len(chunk) > max_chars:
            chunk = chunk[:max_chars]
        preview = f"{chunk}\n… ({len(lines) - max_lines} more lines)"
    elif len(text) > max_chars:
        preview = text[:max_chars] + "…"
    else:
        preview = text
    return preview, len(lines), kind


def _copy_small_fields(original: str, envelope: dict[str, Any]) -> None:
    stripped = original.strip()
    if not (stripped.startswith("{") and stripped.endswith("}")):
        return
    try:
        parsed = json.loads(stripped)
    except (json.JSONDecodeError, TypeError, ValueError):
        return
    if not isinstance(parsed, dict):
        return
    for key in ("status", "file_path", "path", "content_version", "sheet", "range"):
        if key in parsed and parsed[key] is not None:
            envelope[key] = parsed[key]
    meta = parsed.get("meta")
    if isinstance(meta, dict):
        slim = {
            key: value
            for key, value in meta.items()
            if key != "spill" and not isinstance(value, (list, dict))
        }
        if slim:
            envelope.setdefault("meta", {}).update(slim)


def project_for_wire(
    text: str,
    *,
    store: SpillStore,
    char_threshold: int = DEFAULT_SPILL_CHAR_THRESHOLD,
    byte_threshold: int = DEFAULT_SPILL_BYTE_THRESHOLD,
    token_threshold: int = DEFAULT_SPILL_TOKEN_THRESHOLD,
    preview_chars: int = DEFAULT_PREVIEW_CHARS,
    preview_lines: int = DEFAULT_PREVIEW_LINES,
) -> SpillProjection:
    raw = text if isinstance(text, str) else str(text or "")
    chars = len(raw)
    data_len = len(raw.encode("utf-8"))
    tokens = estimate_tokens(raw)
    if not should_spill(
        raw,
        char_threshold=char_threshold,
        byte_threshold=byte_threshold,
        token_threshold=token_threshold,
    ):
        return SpillProjection(
            model_text=raw,
            locator=None,
            spilled=False,
            chars=chars,
            bytes=data_len,
            estimated_tokens=tokens,
        )
    locator = store.put(raw)
    preview, _line_count, kind = _preview_text(
        raw, max_chars=preview_chars, max_lines=preview_lines,
    )
    projection = SpillProjection(
        model_text="",
        locator=locator,
        spilled=True,
        chars=chars,
        bytes=data_len,
        estimated_tokens=tokens,
        kind=kind,
        preview=preview,
    )
    envelope = projection.summary_payload()
    _copy_small_fields(raw, envelope)
    # 句柄放在最前，避免后续硬截断切掉 locator。
    ordered = {"spill": locator, **{k: v for k, v in envelope.items() if k != "spill"}}
    model_text = json.dumps(ordered, ensure_ascii=False, default=str)
    return replace(projection, model_text=model_text)


def apply_spill(
    result: ToolResult,
    *,
    store: SpillStore,
    char_threshold: int = DEFAULT_SPILL_CHAR_THRESHOLD,
    byte_threshold: int = DEFAULT_SPILL_BYTE_THRESHOLD,
    token_threshold: int = DEFAULT_SPILL_TOKEN_THRESHOLD,
) -> ToolResult:
    """只改 model_text（呈现）；value 原字段保留，外置信息进 meta.spill。"""
    if result.coverage and result.coverage.get("spill_retrieve"):
        return result
    projection = project_for_wire(
        result.model_text,
        store=store,
        char_threshold=char_threshold,
        byte_threshold=byte_threshold,
        token_threshold=token_threshold,
    )
    if not projection.spilled:
        return result
    value = result.value
    model_text = projection.model_text
    if isinstance(value, dict):
        value = dict(value)
        meta = dict(value.get("meta") or {})
        meta["spill"] = {
            "locator": projection.locator,
            "kind": projection.kind,
            "chars": projection.chars,
            "bytes": projection.bytes,
        }
        value["meta"] = meta
        write_ver = meta.get("write_verification")
        if write_ver:
            try:
                envelope = json.loads(model_text)
            except (json.JSONDecodeError, TypeError, ValueError):
                envelope = None
            if isinstance(envelope, dict):
                envelope["write_verification"] = write_ver
                model_text = json.dumps(envelope, ensure_ascii=False, default=str)
    coverage = dict(result.coverage or {})
    coverage["spilled"] = True
    coverage["declared"] = True
    return replace(
        result,
        model_text=model_text,
        value=value,
        truncated=True,
        coverage=coverage,
    )


# ── 写后语义校验 ─────────────────────────────────────────


def _cell_addr(row: int, col: int) -> str:
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(col)}{row}"


def _parse_cell(token: str) -> tuple[int, int] | None:
    from openpyxl.utils import column_index_from_string

    text = str(token or "").strip().split("!")[-1].replace("$", "")
    match = _CELL_RE.match(text)
    if not match:
        return None
    try:
        return int(match.group(2)), column_index_from_string(match.group(1))
    except (ValueError, TypeError):
        return None


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _date_semantic_equal(a: Any, b: Any) -> bool:
    """date/datetime 语义等价：混比时要求 datetime 侧为午夜，避免吞掉时间差异。"""
    if isinstance(a, datetime.datetime) and isinstance(b, datetime.datetime):
        return a == b
    a_date = a.date() if isinstance(a, datetime.datetime) else a
    b_date = b.date() if isinstance(b, datetime.datetime) else b
    if a_date != b_date:
        return False
    for value in (a, b):
        if isinstance(value, datetime.datetime) and value.time() != datetime.time(0, 0):
            return False
    return True


def _str_date_equal(text: str, dt_value: Any) -> bool:
    """ISO 日期/时间字符串 vs 读回的日期值（序列化边界会把 date 降级为字符串）。"""
    s = text.strip()
    try:
        parsed: Any = (
            datetime.datetime.fromisoformat(s)
            if "T" in s or " " in s
            else datetime.date.fromisoformat(s)
        )
    except ValueError:
        return False
    return _date_semantic_equal(parsed, dt_value)


def _values_equal(expected: Any, actual: Any) -> bool:
    if expected is actual:
        return True
    if expected is None or actual is None:
        return expected is None and actual is None
    if _is_formula(expected) or _is_formula(actual):
        return str(expected).strip() == str(actual).strip()
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return float(expected) == float(actual)
    date_types = (datetime.date, datetime.datetime)
    if isinstance(expected, date_types) and isinstance(actual, date_types):
        return _date_semantic_equal(expected, actual)
    if isinstance(expected, str) and isinstance(actual, date_types):
        return _str_date_equal(expected, actual)
    if isinstance(actual, str) and isinstance(expected, date_types):
        return _str_date_equal(actual, expected)
    return expected == actual


def _sample_grid_coords(n_rows: int, n_cols: int, limit: int) -> list[tuple[int, int]]:
    if n_rows <= 0 or n_cols <= 0:
        return []
    total = n_rows * n_cols
    if total <= limit:
        return [(r, c) for r in range(n_rows) for c in range(n_cols)]
    picked: list[tuple[int, int]] = []
    seen: set[tuple[int, int]] = set()

    def add(r: int, c: int) -> None:
        key = (max(0, min(n_rows - 1, r)), max(0, min(n_cols - 1, c)))
        if key not in seen:
            seen.add(key)
            picked.append(key)

    add(0, 0)
    add(0, n_cols - 1)
    add(n_rows - 1, 0)
    add(n_rows - 1, n_cols - 1)
    remaining = limit - len(picked)
    if remaining > 0:
        stride = max(1, total // (remaining + 1))
        for index in range(stride, total, stride):
            add(index // n_cols, index % n_cols)
            if len(picked) >= limit:
                break
    return picked[:limit]


def _op_get(op: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in op and op[key] is not None:
            return op[key]
    return None


def _intended_from_write_op(op: dict[str, Any], default_sheet: str) -> tuple[list[dict[str, Any]], int]:
    from excelmanus.workbook.snapshot import parse_bound_selection

    selection = parse_bound_selection(op.get("selection"))
    start_raw = str(_op_get(op, "start_cell", "startCell", "cell", "start") or "")
    sheet = str(_op_get(op, "sheet", "sheet_name") or default_sheet or "")
    if "!" in start_raw:
        maybe_sheet, start_raw = start_raw.split("!", 1)
        sheet = sheet or maybe_sheet.strip("'")
    parsed = _parse_cell(start_raw)
    values = _op_get(op, "values")
    if not isinstance(values, list) or not values or (parsed is None and selection is None):
        return [], 0
    row0, col0 = parsed or (1, 1)
    if selection is not None:
        sheet = selection.sheet
    n_rows = len(values)
    first = values[0] if values else []
    n_cols = len(first) if isinstance(first, list) else 1
    total = 0
    for row in values:
        cells = row if isinstance(row, list) else [row]
        total += len(cells)
    coords = _sample_grid_coords(n_rows, n_cols, MAX_WRITE_VERIFY_SAMPLE)
    intended: list[dict[str, Any]] = []
    for r_off, c_off in coords:
        row = values[r_off]
        cells = row if isinstance(row, list) else [row]
        if c_off >= len(cells):
            continue
        raw = cells[c_off]
        intended.append({
            "sheet": sheet,
            "row": selection.rows[r_off] if selection is not None else row0 + r_off,
            "col": selection.cols[c_off] if selection is not None and selection.cols else col0 + c_off,
            "expected": raw,
            "kind": "formula" if _is_formula(raw) else "value",
        })
    return intended, total


def collect_intended_cells(
    arguments: dict[str, Any],
    *,
    limit: int = MAX_WRITE_VERIFY_SAMPLE,
) -> tuple[list[dict[str, Any]], int, list[dict[str, Any]]]:
    """从写入参数抽出有界意图格子 + 结构改动。不打开工作簿。"""
    default_sheet = str(arguments.get("sheet") or arguments.get("sheet_name") or "")
    ops = arguments.get("operations")
    intended: list[dict[str, Any]] = []
    total = 0
    structure: list[dict[str, Any]] = []
    if not isinstance(ops, list):
        return intended, total, structure
    for raw in ops:
        if not isinstance(raw, dict):
            continue
        kind = str(_op_get(raw, "kind") or "")
        if kind == "write" or (kind == "" and _op_get(raw, "values") is not None):
            cells, count = _intended_from_write_op(raw, default_sheet)
            intended.extend(cells)
            total += count
        elif kind in {"insert", "delete_rows", "delete_columns"}:
            structure.append({
                "kind": kind,
                "sheet": str(_op_get(raw, "sheet", "sheet_name") or default_sheet),
                "at": _op_get(raw, "at", "row", "column"),
                "count": _op_get(raw, "count") or 1,
            })
        elif kind == "copy":
            start = str(_op_get(raw, "target_start", "targetStart") or "A1")
            parsed = _parse_cell(start)
            if parsed:
                intended.append({
                    "sheet": str(
                        _op_get(raw, "target_sheet", "targetSheet") or default_sheet
                    ),
                    "row": parsed[0],
                    "col": parsed[1],
                    "expected": None,
                    "kind": "value",
                    "note": "copy-dest",
                })
                total += 1
        if len(intended) >= limit:
            intended = intended[:limit]
            break
    return intended[:limit], total, structure[:MAX_WRITE_VERIFY_ENTRIES]


def _resolve_workbook_path(file_path: str, workspace_root: str) -> Path | None:
    if not file_path:
        return None
    candidate = Path(file_path)
    if not candidate.is_absolute():
        candidate = Path(workspace_root) / file_path
    try:
        return candidate.resolve()
    except OSError:
        return None


def _error_verification(
    message: str,
    *,
    error_code: str,
    failure_class: str | None = None,
    **fields: Any,
) -> dict[str, Any]:
    payload = make_error_payload(
        message,
        error_code=error_code,
        failure_class=failure_class,
        **fields,
    )
    payload.setdefault("value_changes", [])
    payload.setdefault("formula_changes", [])
    payload.setdefault("total_changes", 0)
    payload.setdefault("shown", 0)
    payload.setdefault("truncated", False)
    return payload


def _open_post_write_workbook(
    abs_path: Path,
    file_path: str,
    workspace_root: str,
    arguments: dict[str, Any],
) -> tuple[Any, Any] | dict[str, Any]:
    """按写后版本打开工作簿快照；失败时返回 error payload 而不是抛异常。"""
    try:
        from excelmanus.workbook.snapshot import SnapshotError, open_snapshot_at
        from excelmanus.workspace.refs import WorkspaceRef

        rel = str(file_path).replace("\\", "/")
        try:
            rel = str(abs_path.resolve().relative_to(Path(workspace_root).resolve())).replace("\\", "/")
        except ValueError:
            pass
        # content_version 是写前乐观锁版本，写后必然过期——只认提交管线回传的 after_version
        expected = str(arguments.get("after_version") or "").strip() or None
        snap = open_snapshot_at(
            abs_path,
            relative=rel,
            workspace=WorkspaceRef.from_root(workspace_root),
            expected_version=expected,
        )
        wb = snap.open_workbook(data_only=False, read_only=False)
        return snap, wb
    except SnapshotError as exc:
        return _error_verification(
            f"写后回读打开失败：{exc}",
            error_code=getattr(exc, "code", TOOL_ERROR),
            failure_class=FAILURE_INTERNAL,
            file_path=file_path,
        )
    except Exception as exc:
        return _error_verification(
            f"写后回读打开失败：{exc}",
            error_code=TOOL_ERROR,
            failure_class=FAILURE_INTERNAL,
            file_path=file_path,
        )


def _short_scalar(value: Any, limit: int = 80) -> Any:
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    text = str(value)
    return text[:limit]


def _style_color_tail(value: Any) -> str | None:
    """归一颜色到末 6 位 hex（大写）；非 hex（theme:/indexed:）返回 None。"""
    if value in (None, ""):
        return None
    from excelmanus.workbook.styles import _resolve_color

    resolved = _resolve_color(str(value).strip())
    if not resolved:
        return None
    text = str(resolved).strip().lstrip("#")
    if len(text) >= 6 and all(c in "0123456789abcdefABCDEF" for c in text):
        return text[-6:].upper()
    return None


def _check_style_cell(cell: Any, op: dict[str, Any]) -> tuple[list[str], list[dict[str, Any]]]:
    """对单格核对 op 声明的样式属性；返回 (ok_props, mismatches)。"""
    from excelmanus.workbook.styles import _extract_alignment, _extract_border, _extract_fill, _extract_font

    ok: list[str] = []
    bad: list[dict[str, Any]] = []

    def record(prop: str, matched: bool, expected: Any, actual: Any) -> None:
        if matched:
            ok.append(prop)
        else:
            bad.append({
                "prop": prop,
                "expected": _short_scalar(expected),
                "actual": _short_scalar(actual),
            })

    fill_cfg = _op_get(op, "fill")
    if fill_cfg:
        requested = (
            _op_get(fill_cfg, "color", "fgColor", "fg_color", "start_color")
            if isinstance(fill_cfg, Mapping)
            else fill_cfg if isinstance(fill_cfg, str) else None
        )
        info = _extract_fill(cell.fill)
        if requested is not None:
            exp_tail = _style_color_tail(requested)
            act_tail = (
                _style_color_tail(info.get("color"))
                if isinstance(info, Mapping)
                else None
            )
            record(
                "fill",
                exp_tail is not None and act_tail is not None and exp_tail == act_tail,
                requested,
                (info.get("color") if isinstance(info, Mapping) and info.get("color") else "none"),
            )
        else:
            record("fill", bool(info), "non-empty fill", "none" if not info else "present")

    font_cfg = _op_get(op, "font")
    if isinstance(font_cfg, Mapping):
        info = _extract_font(cell.font) or {}
        for prop in ("bold", "italic", "underline", "size", "color", "name"):
            if prop not in font_cfg:
                continue
            expected = font_cfg[prop]
            actual = info.get(prop)
            if prop in {"bold", "italic"}:
                matched = bool(expected) == bool(actual)
            elif prop == "underline":
                if isinstance(expected, bool):
                    matched = bool(actual) == expected
                else:
                    matched = str(actual or "") == str(expected)
            elif prop == "size":
                try:
                    matched = float(expected) == float(actual)
                except (TypeError, ValueError):
                    matched = False
            elif prop == "color":
                matched = (
                    _style_color_tail(expected) is not None
                    and _style_color_tail(expected) == _style_color_tail(actual)
                )
            else:
                matched = str(actual or "") == str(expected)
            record(f"font.{prop}", matched, expected, actual if actual is not None else "default")

    number_format = _op_get(op, "number_format", "numberFormat", "numFmt")
    if number_format is not None:
        record(
            "number_format",
            str(cell.number_format) == str(number_format),
            number_format,
            cell.number_format,
        )

    if _op_get(op, "border"):
        info = _extract_border(cell.border)
        record("border", bool(info), "non-empty border", "none" if not info else "present")

    align_cfg = _op_get(op, "alignment")
    if isinstance(align_cfg, Mapping):
        info = _extract_alignment(cell.alignment) or {}
        aliases = {"wrapText": "wrap_text", "horizontalAlignment": "horizontal", "verticalAlignment": "vertical"}
        seen: set[str] = set()
        for raw_prop in ("horizontal", "vertical", "wrap_text", "wrapText"):
            prop = aliases.get(raw_prop, raw_prop)
            if prop in seen or raw_prop not in align_cfg:
                continue
            seen.add(prop)
            expected = align_cfg[raw_prop]
            actual = info.get(prop)
            default = {"horizontal": "general", "vertical": "bottom", "wrap_text": False}[prop]
            if actual is None:
                actual = default
            if prop == "wrap_text":
                matched = bool(expected) == bool(actual)
            else:
                matched = str(actual) == str(expected)
            record(f"alignment.{prop}", matched, expected, actual)

    return ok, bad


def _rect_bounds(rect: Any, ws: Any) -> tuple[int, int, int, int]:
    """RectRef → 有界 min_row/min_col/max_row/max_col；整行整列裁到已用范围。"""
    min_row, max_row = int(rect.min_row), int(rect.max_row)
    min_col, max_col = int(rect.min_col), int(rect.max_col)
    if getattr(rect, "whole_column", False):
        max_row = min(max_row, max(1, int(ws.max_row or 1)))
    if getattr(rect, "whole_row", False):
        max_col = min(max_col, max(1, int(ws.max_column or 1)))
    return min_row, min_col, max_row, max_col


def _ranges_intersect(a_min_col: int, a_min_row: int, a_max_col: int, a_max_row: int, other: Any) -> bool:
    return not (
        other.max_col < a_min_col or other.min_col > a_max_col
        or other.max_row < a_min_row or other.min_row > a_max_row
    )


def _verify_write_style(
    arguments: dict[str, Any],
    file_path: str,
    abs_path: Path,
    workspace_root: str,
    entry_limit: int,
) -> dict[str, Any]:
    """样式/结构类写入的回读：抽样核对声明的样式属性与结构改动。"""
    opened = _open_post_write_workbook(abs_path, file_path, workspace_root, arguments)
    if isinstance(opened, dict):
        return opened
    snap, wb = opened

    from excelmanus.workbook.refs import parse_rect
    from excelmanus.workbook.snapshot import require_default_sheet

    try:
        from excelmanus.tools.intent_tools import (
            _canonical_format_kind,
            _format_rects,
            _freeze_cell_from_op,
        )
    except Exception:  # pragma: no cover - 防御性回退
        _aliases = {
            "conditionalformat": "conditional_format",
            "cf": "conditional_format",
            "datavalidation": "data_validation",
            "dv": "data_validation",
        }

        def _canonical_format_kind(kind: str) -> str:  # type: ignore[no-redef]
            return _aliases.get(kind, kind)

        _format_rects = None  # type: ignore[assignment]
        _freeze_cell_from_op = None  # type: ignore[assignment]

    default_sheet = str(arguments.get("sheet") or arguments.get("sheet_name") or "")
    style_changes: list[dict[str, Any]] = []
    style_mismatches: list[dict[str, Any]] = []
    structure_changes: list[dict[str, Any]] = []
    sampled = 0
    sample_capped = False
    sheets_available = list(wb.sheetnames)
    ops = arguments.get("operations")
    ops = ops if isinstance(ops, list) else []

    def _resolve_sheet(op: dict[str, Any], raw_range: str) -> tuple[str | None, list[str]]:
        """(sheet, local rects)。失败返回 (sheet or None, [])。"""
        explicit = _op_get(op, "sheet", "sheet_name")
        if _format_rects is not None and str(raw_range or "").strip():
            try:
                sheet, locals_ = _format_rects(op, raw_range, allow_union=True)
                return sheet or (str(explicit) if explicit else None), locals_
            except Exception:
                return (str(explicit) if explicit else None), []
        locals_: list[str] = []
        sheet = str(explicit) if explicit else None
        for part in re.split(r"[\s,]+", str(raw_range or "").strip()):
            if not part:
                continue
            if "!" in part:
                prefix, part = part.split("!", 1)
                sheet = sheet or prefix.strip("'")
            locals_.append(part.replace("$", ""))
        return sheet, locals_

    try:
        for raw in ops:
            if not isinstance(raw, dict):
                continue
            kind_raw = str(_op_get(raw, "kind") or "format")
            kind = _canonical_format_kind(kind_raw)
            raw_range = str(_op_get(raw, "range", "cell_range") or "")
            sheet_name, locals_ = _resolve_sheet(raw, raw_range)
            try:
                title = require_default_sheet(sheets_available, sheet_name or default_sheet)
            except Exception:
                structure_changes.append({
                    "kind": kind,
                    "sheet": str(sheet_name or default_sheet or ""),
                    "range": raw_range[:80],
                    "verified": False,
                })
                continue
            if title not in wb.sheetnames:
                structure_changes.append({
                    "kind": kind,
                    "sheet": str(title),
                    "range": raw_range[:80],
                    "verified": False,
                })
                continue
            ws = wb[title]

            if kind == "format":
                if not locals_:
                    structure_changes.append({
                        "kind": "format",
                        "sheet": ws.title,
                        "range": raw_range[:80],
                        "verified": False,
                    })
                    continue
                for local in locals_:
                    try:
                        rect = parse_rect(local)
                    except Exception:
                        structure_changes.append({
                            "kind": "format",
                            "sheet": ws.title,
                            "range": local[:80],
                            "verified": False,
                        })
                        continue
                    min_row, min_col, max_row, max_col = _rect_bounds(rect, ws)
                    remaining = MAX_WRITE_VERIFY_SAMPLE - sampled
                    if remaining <= 0:
                        sample_capped = True
                        break
                    coords = _sample_grid_coords(max_row - min_row + 1, max_col - min_col + 1, remaining)
                    for r_off, c_off in coords:
                        row = min_row + r_off
                        col = min_col + c_off
                        cell = ws.cell(row=row, column=col)
                        ok, bad = _check_style_cell(cell, raw)
                        addr = _cell_addr(row, col)
                        sampled += 1
                        if ok:
                            style_changes.append({"sheet": ws.title, "cell": addr, "ok": ok})
                        for entry in bad:
                            entry["sheet"] = ws.title
                            entry["cell"] = addr
                            style_mismatches.append(entry)
            elif kind in {"merge", "unmerge"}:
                local = locals_[0] if locals_ else raw_range.replace("$", "")
                try:
                    rect = parse_rect(local)
                    target = rect.to_a1(include_sheet=False).replace("$", "")
                except Exception:
                    target = local.replace("$", "")
                merged = {str(r) for r in ws.merged_cells.ranges}
                present = target in merged
                structure_changes.append({
                    "kind": kind,
                    "sheet": ws.title,
                    "range": target[:80],
                    "verified": present if kind == "merge" else not present,
                })
            elif kind == "freeze":
                requested: str | None = None
                loose = False
                if _freeze_cell_from_op is not None:
                    try:
                        requested = _freeze_cell_from_op(raw)
                    except Exception:
                        requested = None
                if requested is None:
                    loose = True
                panes = str(ws.freeze_panes or "")
                if loose:
                    verified = bool(panes)
                elif requested == "":
                    verified = not panes or panes == "A1"
                else:
                    verified = panes == requested.replace("$", "").upper()
                entry = {
                    "kind": "freeze",
                    "sheet": ws.title,
                    "range": str(requested or "")[:80],
                    "verified": verified,
                }
                if loose:
                    entry["loose"] = True
                structure_changes.append(entry)
            elif kind in {"conditional_format", "data_validation"}:
                remove = bool(_op_get(raw, "remove", "delete", "clear"))
                rule_count = 0
                if kind == "conditional_format":
                    configured = [
                        rng
                        for cf in ws.conditional_formatting
                        for rng in getattr(cf.sqref, "ranges", [])
                    ]
                else:
                    configured = [
                        rng
                        for dv in getattr(ws.data_validations, "dataValidation", [])
                        for rng in getattr(dv.sqref, "ranges", [])
                    ]
                for local in locals_ or [raw_range.replace("$", "")]:
                    try:
                        rect = parse_rect(local)
                        min_row, min_col, max_row, max_col = _rect_bounds(rect, ws)
                    except Exception:
                        continue
                    rule_count += sum(
                        1
                        for rng in configured
                        if _ranges_intersect(min_col, min_row, max_col, max_row, rng)
                    )
                structure_changes.append({
                    "kind": kind,
                    "sheet": ws.title,
                    "range": raw_range[:80],
                    "verified": (rule_count == 0) if remove else (rule_count > 0),
                    "rule_count": rule_count,
                })
            elif kind == "size":
                structure_changes.append({"kind": "size", "sheet": ws.title, "verified": None})
            if sample_capped:
                break
    finally:
        try:
            wb.close()
        except Exception:
            pass

    mismatch_count = len(style_mismatches)
    shown_changes = style_changes[:entry_limit]
    shown_mismatches = style_mismatches[:entry_limit]
    shown_structure = structure_changes[:entry_limit]
    total_changes = sampled + len(structure_changes)
    truncated = (
        sample_capped
        or len(style_changes) > len(shown_changes)
        or len(style_mismatches) > len(shown_mismatches)
        or len(structure_changes) > len(shown_structure)
    )
    sheet_hint = str(
        arguments.get("sheet")
        or arguments.get("sheet_name")
        or (style_changes[0].get("sheet") if style_changes else "")
        or (structure_changes[0].get("sheet") if structure_changes else "")
        or ""
    )
    base: dict[str, Any] = {
        "sheet": sheet_hint,
        "verification_kind": "style",
        "sampled": True,
        "style_changes": shown_changes,
        "style_mismatches": shown_mismatches,
        "structure_changes": shown_structure,
        "value_changes": [],
        "formula_changes": [],
        "total_changes": total_changes,
        "shown": len(shown_changes) + len(shown_mismatches) + len(shown_structure),
        "mismatch_count": mismatch_count,
        "truncated": truncated,
        "coverage": {
            "kind": "sampled",
            "sample_size": sampled,
            "snapshot_id": snap.id.key(),
        },
        "snapshot_id": snap.id.key(),
        "content_version": snap.content_version,
    }
    unverified_structure = sum(
        1 for entry in structure_changes if entry.get("verified") is False
    )
    if mismatch_count or unverified_structure:
        return _error_verification(
            f"写后样式抽样发现 {mismatch_count} 项属性不符、"
            f"{unverified_structure} 项结构未确认（抽样 {sampled} 格 / 结构 {len(structure_changes)} 项）",
            error_code=RESULT_UNCERTAIN,
            **base,
        )
    base["status"] = "success"
    return base


def verify_write(
    tool_name: str,
    arguments: dict[str, Any] | None,
    *,
    workspace_root: str,
    entry_limit: int = MAX_WRITE_VERIFY_ENTRIES,
) -> dict[str, Any]:
    """抽样回读：只读 operations 声明的格子，不上整表扫描。

    抽样：矩形 write 取四角 + 均匀步进，上限 ``MAX_WRITE_VERIFY_SAMPLE``；
    展示再截到 ``entry_limit``（默认 20）并给出总数。
    """
    arguments = arguments or {}
    excel_tools = {
        "edit_spreadsheet",
        "format_spreadsheet",
        "manage_spreadsheet_objects",
        "manage_spreadsheet_versions",
    }
    if tool_name == "write_word":
        return _verify_write_word(arguments, workspace_root)
    if tool_name not in excel_tools:
        return {"skipped": True, "status": "success"}

    file_path = str(arguments.get("file_path") or arguments.get("path") or "").strip()
    if not file_path:
        return {"skipped": True, "status": "success"}

    abs_path = _resolve_workbook_path(file_path, workspace_root)
    if abs_path is None or not abs_path.is_file():
        return _error_verification(
            f"写后回读找不到文件：{file_path}",
            error_code=NOT_FOUND,
            failure_class=FAILURE_NOT_FOUND,
            file_path=file_path,
        )

    if tool_name == "format_spreadsheet":
        return _verify_write_style(arguments, file_path, abs_path, workspace_root, entry_limit)

    intended, total_intended, structure = collect_intended_cells(arguments)
    opened = _open_post_write_workbook(abs_path, file_path, workspace_root, arguments)
    if isinstance(opened, dict):
        return opened
    snap, wb = opened
    from excelmanus.workbook.snapshot import SnapshotError, require_default_sheet

    value_changes: list[dict[str, Any]] = []
    formula_changes: list[dict[str, Any]] = []
    mismatches: list[dict[str, Any]] = []
    try:
        sheets_available = list(wb.sheetnames)
        for item in intended:
            sheet_name = item.get("sheet") or arguments.get("sheet") or arguments.get("sheet_name")
            try:
                title = require_default_sheet(sheets_available, sheet_name)
            except SnapshotError as exc:
                return _error_verification(
                    str(exc),
                    error_code=getattr(exc, "code", SHEET_NOT_FOUND),
                    failure_class=FAILURE_NOT_FOUND,
                    sheet=sheet_name,
                    available_sheets=sheets_available,
                )
            if title not in wb.sheetnames:
                return _error_verification(
                    f"写后回读工作表不存在：{title}",
                    error_code=SHEET_NOT_FOUND,
                    failure_class=FAILURE_NOT_FOUND,
                    sheet=title,
                    available_sheets=sheets_available,
                )
            ws = wb[title]
            row = int(item["row"])
            col = int(item["col"])
            cell = ws.cell(row=row, column=col)
            actual = cell.value
            addr = _cell_addr(row, col)
            kind: ChangeKind = "formula" if _is_formula(item.get("expected")) or _is_formula(actual) else "value"
            if item.get("note") == "copy-dest":
                entry = {
                    "sheet": ws.title,
                    "cell": addr,
                    "value": actual if kind == "value" else None,
                    "formula": actual if kind == "formula" else None,
                }
                (formula_changes if kind == "formula" else value_changes).append(entry)
                continue
            expected = item.get("expected")
            if not _values_equal(expected, actual):
                mismatches.append({
                    "sheet": ws.title,
                    "cell": addr,
                    "kind": kind,
                    "expected": expected,
                    "actual": actual,
                })
                continue
            if kind == "formula":
                formula_changes.append({
                    "sheet": ws.title,
                    "cell": addr,
                    "formula": actual,
                })
            else:
                value_changes.append({
                    "sheet": ws.title,
                    "cell": addr,
                    "value": actual,
                })
    finally:
        try:
            wb.close()
        except Exception:
            pass

    shown_values = value_changes[:entry_limit]
    remaining = max(0, entry_limit - len(shown_values))
    shown_formulas = formula_changes[:remaining]
    shown_mismatches = mismatches[: max(0, entry_limit - len(shown_values) - len(shown_formulas))]
    listed = len(shown_values) + len(shown_formulas) + len(shown_mismatches)
    total_changes = total_intended or (len(value_changes) + len(formula_changes) + len(mismatches))
    truncated = total_changes > listed or len(value_changes) > len(shown_values) or len(formula_changes) > len(shown_formulas)

    sheet_hint = str(
        arguments.get("sheet")
        or arguments.get("sheet_name")
        or ""
    )
    if not sheet_hint:
        ops = arguments.get("operations") or []
        if isinstance(ops, list):
            for raw in ops:
                if isinstance(raw, dict):
                    sheet_hint = str(_op_get(raw, "sheet", "sheet_name") or "")
                    if sheet_hint:
                        break
    if not sheet_hint:
        for bucket in (value_changes, formula_changes, mismatches):
            if bucket:
                sheet_hint = str(bucket[0].get("sheet") or "")
                break
    formula_intended = sum(1 for item in intended if item.get("kind") == "formula")
    formula_verified = len(formula_changes)
    formula_overwritten = sum(
        1
        for item in mismatches
        if item.get("kind") == "formula" and not _is_formula(item.get("actual"))
    )
    base: dict[str, Any] = {
        "sheet": sheet_hint,
        "value_changes": shown_values,
        "formula_changes": shown_formulas,
        "formula_intended_count": formula_intended,
        "formula_verified_count": formula_verified,
        "formula_overwritten_count": formula_overwritten,
        "structure_changes": structure[:entry_limit],
        "total_changes": total_changes,
        "shown": listed,
        "truncated": truncated,
        "sampled": True,
        "coverage": {
            "kind": "sampled",
            "sample_size": len(intended),
            "snapshot_id": snap.id.key(),
        },
        "snapshot_id": snap.id.key(),
        "content_version": snap.content_version,
    }
    if mismatches:
        extra = dict(base)
        extra["mismatches"] = shown_mismatches
        extra["mismatch_count"] = len(mismatches)
        return _error_verification(
            f"写后抽样发现 {len(mismatches)} 格与意图不符（抽样 {len(intended)} / 意图 {total_changes}）",
            error_code=RESULT_UNCERTAIN,
            **extra,
        )
    if not intended:
        base["verification_kind"] = "readability"
        base["reason"] = "未提取到可逐格核验的值意图；仅确认提交版本可打开"
    base["status"] = "success"
    return base


def _verify_write_word(arguments: dict[str, Any], workspace_root: str) -> dict[str, Any]:
    file_path = str(arguments.get("file_path") or "").strip()
    if not file_path:
        return {"skipped": True, "status": "success"}
    abs_path = _resolve_workbook_path(file_path, workspace_root)
    if abs_path is None or not abs_path.is_file():
        return _error_verification(
            f"写后回读找不到 Word 文件：{file_path}",
            error_code=NOT_FOUND,
            failure_class=FAILURE_NOT_FOUND,
            file_path=file_path,
        )
    try:
        from docx import Document as _Document

        doc = _Document(str(abs_path))
        return {
            "status": "success",
            "kind": "word",
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "value_changes": [],
            "formula_changes": [],
            "total_changes": 0,
            "shown": 0,
            "truncated": False,
        }
    except Exception as exc:
        return _error_verification(
            f"写后回读 Word 失败：{exc}",
            error_code=TOOL_ERROR,
            failure_class=FAILURE_INTERNAL,
            file_path=file_path,
        )


def format_write_verification_line(payload: dict[str, Any]) -> str:
    if not payload:
        return ""
    if payload.get("skipped"):
        if payload.get("verification_kind") == "style":
            sheet = str(payload.get("sheet") or "")
            prefix = f"{sheet} " if sheet else ""
            return f"\n未核验: {prefix}样式属性未做写后回读，不能据此声称样式已验证。"
        return ""
    if payload.get("status") == "error" or payload.get("error_code"):
        code = payload.get("error_code") or "TOOL_ERROR"
        message = payload.get("message") or "写后校验失败"
        n_mis = payload.get("mismatch_count") or len(payload.get("mismatches") or [])
        return f"\n回读确认: {code} {message}" + (f" mismatches={n_mis}" if n_mis else "")
    if payload.get("kind") == "word":
        return (
            f"\n回读确认: Word 文档可读，"
            f"当前段落数={payload.get('paragraphs', 0)}，"
            f"表格数={payload.get('tables', 0)}"
        )
    if payload.get("verification_kind") == "style":
        changes = payload.get("style_changes") or []
        mismatches = payload.get("style_mismatches") or []
        structures = payload.get("structure_changes") or []
        ok_count = sum(len(item.get("ok") or []) for item in changes)
        n_mis = payload.get("mismatch_count") or len(mismatches)
        n_structure_ok = sum(1 for item in structures if item.get("verified") is True)
        sample_size = (payload.get("coverage") or {}).get("sample_size") or len(changes) + len(mismatches)
        sheet = str(payload.get("sheet") or "")
        prefix = f"{sheet} " if sheet else ""
        line = (
            f"\n样式回读: {prefix}抽样 {sample_size} 格，属性 {ok_count} 项通过，"
            f"{n_mis} 项不符；结构 {n_structure_ok}/{len(structures)} 项已确认"
        )
        if mismatches:
            head = "；".join(
                f"{item.get('cell','')}.{item.get('prop','')} {item.get('expected','')}→{item.get('actual','')}"
                for item in mismatches[:3]
            )
            line += f"（不符: {head}）"
        return line
    if payload.get("verification_kind") == "readability":
        return f"\n回读确认: {payload.get('sheet', '')} 文件可打开；未核验本次对象、结构或样式效果。"
    n_val = len(payload.get("value_changes") or [])
    n_fml = len(payload.get("formula_changes") or [])
    total = payload.get("total_changes") or (n_val + n_fml)
    sheet = str(payload.get("sheet") or "")
    if not sheet:
        for bucket in (payload.get("value_changes"), payload.get("formula_changes")):
            if bucket:
                sheet = str(bucket[0].get("sheet") or "")
                break
    sheet_bit = f"{sheet} " if sheet else ""
    trunc = " 已截断" if payload.get("truncated") else ""
    line = (
        f"\n回读确认: {sheet_bit}值核验 {n_val} 公式核验 {n_fml}（仅写后值，无编辑前值）"
        f"（共 {total}，显示 {payload.get('shown', n_val + n_fml)}）{trunc}"
    )
    formula_intended = int(payload.get("formula_intended_count") or 0)
    if formula_intended:
        line += f"；公式 保留 {int(payload.get('formula_verified_count') or 0)}/{formula_intended}"
    return line.replace("  ", " ")


def compact_write_verification(payload: dict[str, Any] | None) -> str:
    if not payload or payload.get("skipped"):
        return ""
    if payload.get("status") == "error" or payload.get("error_code"):
        return f"verify {payload.get('error_code')}"
    if payload.get("verification_kind") == "style":
        ok_count = sum(len(item.get("ok") or []) for item in payload.get("style_changes") or [])
        n_mis = payload.get("mismatch_count") or len(payload.get("style_mismatches") or [])
        n_structure = len(payload.get("structure_changes") or [])
        return f"verify style_ok={ok_count} style_mismatch={n_mis} structure={n_structure} total={payload.get('total_changes', 0)}"
    n_val = len(payload.get("value_changes") or [])
    n_fml = len(payload.get("formula_changes") or [])
    return f"verify value={n_val} formula={n_fml} total={payload.get('total_changes', n_val + n_fml)}"


def attach_write_verification(
    result: ToolResult | None,
    verification: dict[str, Any],
    result_str: str,
) -> tuple[ToolResult | None, str]:
    """把校验写入 value.meta.write_verification，并追加一行给模型。"""
    if not verification:
        return result, result_str
    line = format_write_verification_line(verification)
    entries = [*(verification.get("value_changes") or []), *(verification.get("formula_changes") or [])]
    if entries:
        line += "\n回读样本: " + json.dumps(entries[:8], ensure_ascii=False, default=str)
    if line and line not in result_str:
        result_str = result_str + line
    if result is None:
        return result, result_str
    value = result.value
    if isinstance(value, dict):
        value = dict(value)
        meta = dict(value.get("meta") or {})
        meta["write_verification"] = verification
        value["meta"] = meta
        result = replace(result, value=value, model_text=result_str)
    else:
        result = result.with_model_text(result_str)
    return result, result_str


def spill_result_text(
    result_str: str,
    structured: ToolResult | None,
    *,
    store: SpillStore | None,
    char_threshold: int = DEFAULT_SPILL_CHAR_THRESHOLD,
    byte_threshold: int = DEFAULT_SPILL_BYTE_THRESHOLD,
    token_threshold: int = DEFAULT_SPILL_TOKEN_THRESHOLD,
) -> tuple[str, ToolResult | None]:
    if store is None:
        return result_str, structured
    if structured is not None and structured.coverage and structured.coverage.get("spill_retrieve"):
        return result_str, structured
    target = structured.with_model_text(result_str) if structured is not None else ToolResult.from_text(result_str)
    try:
        spilled = apply_spill(
            target,
            store=store,
            char_threshold=char_threshold,
            byte_threshold=byte_threshold,
            token_threshold=token_threshold,
        )
    except OSError:
        logger.debug("spill 写入失败，回退原文", exc_info=True)
        return result_str, structured
    if structured is None:
        return spilled.model_text, None
    return spilled.model_text, spilled


def expose_spreadsheet_value(result: ToolResult, *, store: SpillStore) -> ToolResult:
    """Keep native tool results as usable as their SDK value, with bounded text.

    Small results carry the actual payload; large results carry an opaque
    handle to that same payload. Never spill only a preview while calling it
    the complete result. The existing read_text_file path retrieves handles.
    """
    if not isinstance(result.value, dict) or (result.coverage or {}).get("spill_retrieve"):
        return result
    payload = result.value
    raw = json.dumps(payload, ensure_ascii=False, default=str)
    if not should_spill(raw):
        return result.with_model_text(raw)
    locator = store.put(raw)
    envelope = {
        "result_spill": str(locator),
        "read_result": "按 next_call 取回完整 JSON；file_path 原样使用 result_spill 字段的值（spill:…），不要把 result_spill 当作前缀或目录。",
        "next_call": {"tool": "read_text_file", "arguments": {"file_path": str(locator)}},
    }
    for key in (
        "status", "file_path", "file_a", "file_b", "content_version",
        "content_version_a", "content_version_b", "resolved_sheet", "scope",
        "coverage", "selection_spill", "warnings", "error_code", "message", "remediation",
        "committed", "partial", "tx_id", "operation_id", "recovery_required",
        "operation_index", "operation_kind",
    ):
        if key in payload:
            envelope[key] = payload[key]
    # A short summary is navigation only; the complete structured payload is
    # always recoverable, including artifact lists, selections and warnings.
    envelope["preview"] = result.model_text[:DEFAULT_PREVIEW_CHARS]
    envelope["result_projection"] = "partial; full payload in result_spill"
    value = {**payload, "result_spill": str(locator)}
    return replace(result, value=value, model_text=json.dumps(envelope, ensure_ascii=False, default=str))
