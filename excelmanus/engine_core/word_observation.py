"""Post-write Word observations; workbook evidence belongs to WorkbookService."""
from __future__ import annotations
import json
from dataclasses import replace
from pathlib import Path
from typing import Any
from excelmanus.engine_core.tool_result import ToolResult
from excelmanus.engine_core.error_payload import NOT_FOUND, TOOL_ERROR, FAILURE_INTERNAL, FAILURE_NOT_FOUND, make_error_payload
def _resolve_document_path(file_path: str, workspace_root: str) -> Path | None:
    if not file_path:
        return None
    candidate = Path(file_path)
    if not candidate.is_absolute():
        candidate = Path(workspace_root) / file_path
    try:
        return candidate.resolve()
    except OSError:
        return None


def _error_verification(
    message: str,
    *,
    error_code: str,
    failure_class: str | None = None,
    **fields: Any,
) -> dict[str, Any]:
    payload = make_error_payload(
        message,
        error_code=error_code,
        failure_class=failure_class,
        **fields,
    )
    payload.setdefault("value_changes", [])
    payload.setdefault("formula_changes", [])
    payload.setdefault("total_changes", 0)
    payload.setdefault("shown", 0)
    payload.setdefault("truncated", False)
    return payload


def _verify_write_word(arguments: dict[str, Any], workspace_root: str) -> dict[str, Any]:
    file_path = str(arguments.get("file_path") or "").strip()
    if not file_path:
        return {"skipped": True, "status": "success"}
    abs_path = _resolve_document_path(file_path, workspace_root)
    if abs_path is None or not abs_path.is_file():
        return _error_verification(
            f"写后回读找不到 Word 文件：{file_path}",
            error_code=NOT_FOUND,
            failure_class=FAILURE_NOT_FOUND,
            file_path=file_path,
        )
    try:
        from docx import Document as _Document

        doc = _Document(str(abs_path))
        return {
            "status": "success",
            "kind": "word",
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "value_changes": [],
            "formula_changes": [],
            "total_changes": 0,
            "shown": 0,
            "truncated": False,
        }
    except Exception as exc:
        return _error_verification(
            f"写后回读 Word 失败：{exc}",
            error_code=TOOL_ERROR,
            failure_class=FAILURE_INTERNAL,
            file_path=file_path,
        )


def format_write_verification_line(payload: dict[str, Any]) -> str:
    if not payload or payload.get("skipped"):
        return ""
    if payload.get("error_code"):
        return f"\nWord 回读失败: {payload.get('message', '')}"
    return f"\nWord 回读: {payload.get('paragraphs', 0)} 段落，{payload.get('tables', 0)} 表格"


def compact_write_verification(payload: dict[str, Any] | None) -> str:
    return format_write_verification_line(payload or {}).strip()


def attach_write_verification(
    result: ToolResult | None,
    verification: dict[str, Any],
    result_str: str,
) -> tuple[ToolResult | None, str]:
    """把校验写入 value.meta.write_verification，并追加一行给模型。"""
    if not verification:
        return result, result_str
    line = format_write_verification_line(verification)
    entries = [*(verification.get("value_changes") or []), *(verification.get("formula_changes") or [])]
    if entries:
        line += "\n回读样本: " + json.dumps(entries[:8], ensure_ascii=False, default=str)
    if line and line not in result_str:
        result_str = result_str + line
    if result is None:
        return result, result_str
    value = result.value
    if isinstance(value, dict):
        value = dict(value)
        meta = dict(value.get("meta") or {})
        meta["write_verification"] = verification
        value["meta"] = meta
        result = replace(result, value=value, model_text=result_str)
    else:
        result = result.with_model_text(result_str)
    return result, result_str


def observe_word_write(tool_name: str, arguments: dict, *, workspace_root: str) -> dict:
    return _verify_write_word(arguments, workspace_root) if tool_name == "write_word" else {"skipped": True, "status": "success"}
