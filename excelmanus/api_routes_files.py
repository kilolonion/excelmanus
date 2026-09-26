"""文件相关 API：Excel / Word / 工作区 / 下载 / 文件组。

从 api.py 抽出的独立路由模块。运行时状态只从 api_app_state 读取，
禁止 ``from excelmanus.api import _config`` 反向导入。
由 api.py 在 create_app 中 include_router 注册。
"""

from __future__ import annotations

import os
import uuid
from pathlib import Path
from typing import Any, Literal

from fastapi import APIRouter, Request
from fastapi.responses import JSONResponse, StreamingResponse
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel, ConfigDict, Field

from excelmanus.api_app_state import (
    UnicodeJSONResponse,
    error_json_response as _error_json_response,
    get_config,
    get_file_registry as _get_file_registry,
    get_session_manager,
    make_content_disposition as _make_content_disposition,
    resolve_excel_path as _resolve_excel_path,
    resolve_workspace as _resolve_workspace_impl,
    resolve_workspace_root as _resolve_workspace_root_impl,
)
from excelmanus.logger import get_logger
from excelmanus.workspace.identity import IdentityError, resolve_canonical

logger = get_logger("api.files")

router = APIRouter()


@router.get("/api/v1/workbooks/observe")
async def get_workbook_observation(request: Request) -> JSONResponse:
    """The V2 observation endpoint; workspace authorization precedes projection."""
    from excelmanus.workbook.service import WorkbookService
    from excelmanus.workbook.snapshot import SnapshotError
    from excelmanus.workbook_commit import CommitError
    from zipfile import BadZipFile
    from xml.etree.ElementTree import ParseError
    try:
        from lxml.etree import XMLSyntaxError
    except ImportError:
        XMLSyntaxError = ParseError
    params = request.query_params
    protocol = request.headers.get("x-workbook-protocol", "workbook/2")
    if protocol != "workbook/2":
        return _error_json_response(409, "工作簿协议版本不匹配，请刷新客户端", code="WORKBOOK_PROTOCOL_MISMATCH")
    path, sheet = params.get("path", ""), params.get("sheet")
    root, scope_error = _file_workspace_root(request, params.get("session_id"), params.get("workspace_id"))
    if scope_error is not None:
        return scope_error
    resolved = _resolve_excel_path(path, params.get("session_id"), workspace_root=root)
    if resolved is None:
        return _error_json_response(404, "文件不存在或路径非法")
    def read():
        snap = _open_route_snapshot(resolved, path, root, params.get("workspace_id"))
        expected = params.get("expected_version")
        if expected and expected != snap.content_version:
            return UnicodeJSONResponse(status_code=409, content={"error": "观察版本已变化", "code": "STALE_VIEW", "content_version": snap.content_version})
        selected = sheet
        mode = params.get("mode", "range")
        if not selected and mode == "range":
            if snap.is_csv():
                selected = "Sheet1"
            else:
                wb = snap.open_workbook(data_only=False, read_only=True)
                try:
                    selected = (wb.active or wb.worksheets[0]).title
                finally:
                    wb.close()
        facets = params.get("facets", "data,presentation,geometry,objects").split(",")
        payload = WorkbookService().observe_snapshot(snap, sheet=selected, range=params.get("range", "A1:Z80" if mode == "range" else None), mode=mode, facets=facets, query=params.get("query", ""), offset=int(params.get("offset", 0)), limit=int(params.get("limit", 50)))
        return JSONResponse(content=payload)
    try:
        return await run_in_threadpool(read)
    except (SnapshotError, CommitError, BadZipFile, XMLSyntaxError, ParseError, ValueError, KeyError) as exc:
        return _error_json_response(400, str(exc), code=getattr(exc, "code", "INVALID_ARGS"))


@router.get("/api/v1/workbooks/object-image")
@router.get("/api/v1/workbooks/image")
async def get_workbook_object_image(request: Request) -> StreamingResponse:
    """Serve one embedded workbook image from an immutable snapshot.

    Workbook observations intentionally carry metadata rather than base64.  The
    browser requests the binary only for objects in the current viewport, while
    the version token prevents an old image being shown after a workbook edit.
    """
    from excelmanus.workbook.snapshot import SnapshotError

    params = request.query_params
    path, sheet = params.get("path", ""), params.get("sheet", "")
    kind = params.get("kind", "image")
    try:
        index = int(params.get("index", "0"))
    except ValueError:
        return _error_json_response(400, "index 必须是非负整数", code="INVALID_ARGS")  # type: ignore[return-value]
    if not path or not sheet or kind != "image" or index < 0:
        return _error_json_response(400, "需要 path、sheet、kind=image、非负 index", code="INVALID_ARGS")  # type: ignore[return-value]
    root, scope_error = _file_workspace_root(request, params.get("session_id"), params.get("workspace_id"))
    if scope_error is not None:
        return scope_error  # type: ignore[return-value]
    resolved = _resolve_excel_path(path, params.get("session_id"), workspace_root=root)
    if resolved is None:
        return _error_json_response(404, "工作簿不存在")  # type: ignore[return-value]

    def read_image():
        snap = _open_route_snapshot(resolved, path, root, params.get("workspace_id"))
        expected = params.get("expected_version")
        if expected and expected != snap.content_version:
            return _error_json_response(409, "对象图片版本已变化", code="STALE_VIEW")
        wb = snap.open_workbook(data_only=False, read_only=False)
        try:
            if sheet not in wb.sheetnames:
                return _error_json_response(404, "工作表不存在")
            images = list(getattr(wb[sheet], "_images", []) or [])
            if index >= len(images):
                return _error_json_response(404, "图片对象不存在")
            image = images[index]
            payload = image._data() if callable(getattr(image, "_data", None)) else None
            if not payload:
                return _error_json_response(404, "图片数据不可用")
            fmt = str(getattr(image, "format", "") or "").lower().lstrip(".")
            media_type = {
                "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
                "gif": "image/gif", "bmp": "image/bmp", "webp": "image/webp",
                "tif": "image/tiff", "tiff": "image/tiff",
            }.get(fmt, "application/octet-stream")
            headers = {
                "Cache-Control": "private, max-age=300",
                "ETag": f'"{snap.content_version}:{sheet}:{index}"',
                "X-Workbook-Version": snap.content_version,
            }
            return StreamingResponse(iter((payload,)), media_type=media_type, headers=headers)
        finally:
            wb.close()

    try:
        return await run_in_threadpool(read_image)
    except (SnapshotError, ValueError, KeyError, OSError) as exc:
        return _error_json_response(400, str(exc), code=getattr(exc, "code", "INVALID_ARGS"))  # type: ignore[return-value]


class WorkbookChangesRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    session_id: str | None = None
    workspace_id: str | None = None
    path: str
    operations: list[dict[str, Any]] = Field(min_length=1, max_length=20000)
    expected_version: str | None = None
    operation_id: str | None = None


@router.post("/api/v1/workbooks/changes")
async def apply_workbook_changes(request: WorkbookChangesRequest, raw_request: Request) -> JSONResponse:
    from excelmanus.tools.context import use_workspace, ToolCallContext, bind_call, reset_call
    from excelmanus.workbook.service import WorkbookService
    from excelmanus.workbook_commit import CommitError
    from excelmanus.tools._helpers import MutationAborted
    from excelmanus.workbook.user_edits import summarize_operations
    root, scope_error = _file_workspace_root(raw_request, request.session_id, request.workspace_id)
    if scope_error is not None:
        return scope_error
    def apply():
        with use_workspace(root, workspace_id=request.workspace_id) as ctx:
            token = bind_call(ToolCallContext(binding=ctx.binding, call_id=request.operation_id or uuid.uuid4().hex,
                                             tool_name="apply_spreadsheet_changes"))
            try:
                return WorkbookService().apply(request.path, operations=request.operations, expected_version=request.expected_version,
                    event_context={"source": "user", "session_id": request.session_id, "summary": summarize_operations(request.operations)})
            finally:
                reset_call(token)
    try:
        payload = await run_in_threadpool(apply)
        payload["cells_written"] = sum(len(op.get("cells", [])) for op in request.operations if op["kind"] == "cells.patch")
        manager = get_session_manager()
        if manager is not None:
            try:
                manager.drain_workspace_events()
            except Exception:
                logger.warning("Workbook committed; event delivery pending", exc_info=True)
        return JSONResponse(content=payload, status_code=200 if payload.get("committed") else 409)
    except CommitError as exc:
        return UnicodeJSONResponse(status_code=409 if exc.code in {"VERSION_CONFLICT", "VERSION_REQUIRED"} else 400,
            content={"error": str(exc), "code": exc.code, **exc.fields})
    except MutationAborted as exc:
        return JSONResponse(status_code=400, content=exc.result.value)
    except (ValueError, KeyError) as exc:
        return _error_json_response(400, str(exc), code="INVALID_ARGS")


@router.get("/api/v1/workbooks/compare")
async def get_workbook_comparison(request: Request) -> JSONResponse:
    """Return V2 observations plus the canonical comparison/relationship facts."""
    from excelmanus.workbook.service import WorkbookService
    from excelmanus.workbook.service import WorkbookService
    from excelmanus.tools.context import use_workspace
    params = request.query_params
    root, scope_error = _file_workspace_root(request, params.get("session_id"), params.get("workspace_id"))
    if scope_error is not None:
        return scope_error
    file_a, file_b = params.get("file_a", ""), params.get("file_b", "")
    if not file_a or not file_b:
        return _error_json_response(400, "需要 file_a 和 file_b", code="INVALID_ARGS")
    snapshots = {}
    for key, path in (("file_a", file_a), ("file_b", file_b)):
        resolved = _resolve_excel_path(path, params.get("session_id"), workspace_root=root)
        if resolved is None:
            return _error_json_response(404, f"文件不存在或路径非法: {path}")
        snapshots[key] = await run_in_threadpool(_open_route_snapshot, resolved, path, root, params.get("workspace_id"))

    def compare():
        with use_workspace(root, workspace_id=params.get("workspace_id")):
            compare_result = WorkbookService().query("compare_spreadsheets", {
                "file_a": file_a,
                "file_b": file_b,
                "alignment": "position",
                "ignore_style": True,
                "max_diffs": int(params.get("max_diffs", "500")),
            })
            relationship_result = WorkbookService().query("analyze_spreadsheet", {
                "mode": "relationships", "file_paths": [file_a, file_b], "max_files": 2,
                "sample_rows": int(params.get("max_rows", "200")),
            })
        return compare_result, relationship_result

    try:
        compare_result, relationship_result = await run_in_threadpool(compare)
        if compare_result.success:
            for side in ("a","b"):
                if compare_result.value.get(f"content_version_{side}") != snapshots[f"file_{side}"].content_version:
                    return _error_json_response(409, "文件在比较期间发生变化，请重试", code="STALE_VIEW")
        if relationship_result.success:
            observed = relationship_result.value.get("source_versions") or {}
            if any(observed.get(snap.file.relative, snap.content_version) != snap.content_version for snap in snapshots.values()):
                return _error_json_response(409, "文件在关联分析期间发生变化，请重试", code="STALE_VIEW")
    except (ValueError, TypeError) as exc:
        return _error_json_response(400, str(exc), code="INVALID_ARGS")
    body = {
        key: await run_in_threadpool(WorkbookService().observe_snapshot, snapshots[key], facets=[], mode="overview")
        for key in ("file_a", "file_b")
    }
    relationships: dict[str, Any] = {"shared_columns": []}
    if getattr(relationship_result, "success", False):
        value = relationship_result.value if isinstance(relationship_result.value, dict) else {}
        pairs = value.get("file_pairs") or []
        if pairs:
            relationships["shared_columns"] = pairs[0].get("shared_columns", [])
        hints = value.get("merge_hints") or []
        if hints:
            relationships["merge_hint"] = hints[0]
    body["relationships"] = relationships
    body["comparison"] = compare_result.value if getattr(compare_result, "success", False) else {
        "status": "error", "message": getattr(compare_result, "model_text", "comparison failed"),
    }
    return JSONResponse(content=body)

_TEXT_PREVIEW_SUFFIXES = frozenset({
    ".txt", ".md", ".markdown", ".json", ".js", ".jsx", ".ts", ".tsx",
    ".py", ".rb", ".go", ".rs", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".php", ".swift", ".kt", ".scala", ".sh", ".bash", ".zsh",
    ".sql", ".html", ".css", ".scss", ".less", ".xml", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".log", ".env", ".gitignore",
    ".dockerignore", ".graphql", ".gql", ".vue", ".svelte", ".ex",
    ".exs", ".erl", ".hs", ".ml", ".fs", ".clj", ".lua", ".r", ".dart",
    ".groovy", ".tsv",
})
_TEXT_PREVIEW_NAMES = frozenset({".gitignore", ".dockerignore", ".env"})
_NOT_TEXT_SUFFIXES = frozenset({
    ".xlsx", ".xls", ".xlsm", ".xlsb", ".csv",
    ".docx", ".doc", ".pptx", ".ppt",
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg",
    ".pdf", ".zip", ".gz", ".tar", ".tgz", ".7z", ".rar", ".bin", ".exe",
})
_NO_STORE_HEADERS = {"Cache-Control": "private, no-store"}


def _display_filename(served_path: str | Path, ws_root: str | Path | None) -> str:
    """Download label: strip ``uploads/{8hex}_`` prefixes and backup suffixes."""
    from excelmanus.workspace.identity import display_name_for

    p = Path(served_path)
    rel = p.name
    if ws_root:
        try:
            rel = p.resolve().relative_to(Path(ws_root).resolve()).as_posix()
        except (ValueError, OSError):
            pass
    return display_name_for(rel) or p.name


def is_text_preview_file(file_path: Path) -> bool:
    """Whether /files/read may return this path as UTF-8 text.

    Known text suffixes and special names pass. Spreadsheet/word/image/archive
    suffixes fail. Unknown suffixes are sniffed (no NUL, UTF-8) so Makefile
    and .env.local 不会因为 suffix 漂移 404.
    """
    name = file_path.name.lower()
    suffix = file_path.suffix.lower()
    if name in _TEXT_PREVIEW_NAMES or name.startswith(".env."):
        return True
    if suffix in _TEXT_PREVIEW_SUFFIXES:
        return True
    if suffix in _NOT_TEXT_SUFFIXES:
        return False
    try:
        sample = file_path.read_bytes()[:8192]
    except OSError:
        return False
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return True



def _resolve_workspace_root(
    request: Request,
    session_id: str | None = None,
    workspace_id: str | None = None,
    *,
    require_scope: bool = True,
) -> str:
    """File-module workspace root: fail-closed. Tests may monkeypatch this name."""
    from fastapi import HTTPException

    try:
        return _resolve_workspace_root_impl(
            request,
            session_id=session_id,
            workspace_id=workspace_id,
            require_scope=require_scope,
        )
    except TypeError:
        return _resolve_workspace_root_impl(request, session_id=session_id)
    except HTTPException as exc:
        detail = exc.detail
        message = (
            str(detail.get("error") or "无法确定工作区，请从会话重新打开文件")
            if isinstance(detail, dict)
            else str(detail)
        )
        code = (
            str(detail.get("code") or "FILE_SCOPE_REQUIRED")
            if isinstance(detail, dict)
            else "FILE_SCOPE_REQUIRED"
        )
        raise HTTPException(
            status_code=exc.status_code,
            detail={"error": message, "code": code},
        ) from exc


def _resolve_workspace(
    request: Request,
    session_id: str | None = None,
    workspace_id: str | None = None,
    *,
    require_scope: bool = True,
) -> Any:
    from fastapi import HTTPException

    try:
        return _resolve_workspace_impl(
            request,
            session_id=session_id,
            workspace_id=workspace_id,
            require_scope=require_scope,
        )
    except TypeError:
        return _resolve_workspace_impl(request, session_id=session_id)
    except HTTPException:
        raise


def _file_workspace_root(
    request: Request,
    session_id: str | None = None,
    workspace_id: str | None = None,
) -> tuple[str | None, JSONResponse | None]:
    from fastapi import HTTPException

    try:
        return _resolve_workspace_root(request, session_id=session_id, workspace_id=workspace_id), None
    except TypeError:
        return _resolve_workspace_root(request, session_id=session_id), None
    except HTTPException as exc:
        detail = exc.detail if isinstance(exc.detail, dict) else {}
        return None, _error_json_response(
            exc.status_code,
            str(detail.get("error") or exc.detail),
            code=str(detail.get("code") or "FILE_SCOPE_REQUIRED"),
        )


def _open_route_snapshot(
    resolved: str,
    display_path: str,
    ws_root: str,
    workspace_id: str | None = None,
):
    """HTTP 开簿：吸收原 _bind_file_bytes，与工具层共用 open_snapshot_at。"""
    from excelmanus.workbook.snapshot import open_snapshot_at
    from excelmanus.workspace.refs import WorkspaceRef

    workspace = WorkspaceRef.from_root(ws_root, workspace_id=workspace_id)
    try:
        rel = str(Path(resolved).resolve().relative_to(Path(ws_root).resolve())).replace("\\", "/")
    except ValueError:
        rel = str(display_path or Path(resolved).name).replace("\\", "/")
    return open_snapshot_at(resolved, relative=rel, workspace=workspace)


def _with_bound_version(payload: dict[str, Any], version: str) -> dict[str, Any]:
    payload["content_version"] = version
    return payload


def _apply_cell_write(ws: Any, cell_ref: str, value: Any) -> None:
    """写入或清空单元格。``ws.cell(..., value=None)`` 不会改已有值。"""
    from openpyxl.utils.cell import coordinate_to_tuple

    row, col = coordinate_to_tuple(str(cell_ref).upper())
    cell = ws.cell(row=row, column=col)
    if value is None or value == "":
        cell.value = None
    else:
        cell.value = value


@router.get("/api/v1/files/excel")
async def get_excel_file(request: Request) -> StreamingResponse:
    """返回 xlsx 文件二进制流供前端 Univer 加载。"""
    assert get_config() is not None, "服务未初始化"

    path = request.query_params.get("path", "")
    session_id = request.query_params.get("session_id")
    workspace_id = request.query_params.get("workspace_id")
    if not path:
        return _error_json_response(400, "缺少 path 参数")  # type: ignore[return-value]

    ws_root, scope_error = _file_workspace_root(request, session_id, workspace_id)
    if scope_error is not None:
        return scope_error  # type: ignore[return-value]
    resolved = _resolve_excel_path(path, session_id, workspace_root=ws_root)
    if resolved is None:
        return _error_json_response(404, f"文件不存在或路径非法: {path}")  # type: ignore[return-value]

    from pathlib import Path as _Path

    file_path = _Path(resolved)
    suffix = file_path.suffix.lower()
    if suffix not in {".xlsx", ".xls", ".xlsb", ".xlsm", ".csv"}:
        return _error_json_response(400, f"不支持的文件格式: {suffix}")  # type: ignore[return-value]

    # .xls/.xlsb → 透明转换为 xlsx 供前端 Univer 加载
    actual_file = resolved
    converted_bytes: bytes | None = None
    from excelmanus.xls_converter import needs_conversion as _nc2
    if _nc2(resolved):
        try:
            snap = await run_in_threadpool(_open_route_snapshot, resolved, path, ws_root, workspace_id)
            converted_bytes = snap.read_bytes()
            suffix = ".xlsx"
        except Exception as exc:
            logger.warning("excel 流转换失败: %s", resolved, exc_info=True)
            return _error_json_response(422, f"旧版工作簿无法转换为 .xlsx: {exc}", code="CONVERSION_FAILED")  # type: ignore[return-value]

    content_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if suffix in (".xlsx", ".xlsm")
        else "text/csv"
    )

    def _iter_file():
        if converted_bytes is not None:
            yield converted_bytes
            return
        with open(actual_file, "rb") as f:  # type: ignore[arg-type]
            while chunk := f.read(65536):
                yield chunk

    return StreamingResponse(
        _iter_file(),
        media_type=content_type,
        headers={
            "Content-Disposition": _make_content_disposition(_display_filename(file_path, ws_root)),
            **_NO_STORE_HEADERS,
        },
    )









class WorkbookDraftBatch(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expected_version: str = Field(pattern=r"^sha256:[0-9a-f]{64}$")
    operations: list[dict[str, Any]] = Field(max_length=5000)


class WorkbookMergeRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    session_id: str | None = None
    workspace_id: str | None = None
    path: str = Field(max_length=300)
    batches: list[WorkbookDraftBatch] = Field(min_length=1, max_length=100)
    apply: bool = False
    expected_version: str | None = Field(default=None, pattern=r"^sha256:[0-9a-f]{64}$")
    choices: dict[str, Literal["local", "remote"]] = Field(default_factory=dict)
    operation_id: str | None = None


@router.post("/api/v1/workbooks/merge-review")
async def merge_workbook_draft(request: WorkbookMergeRequest, raw_request: Request) -> JSONResponse:
    """Review a draft, then publish the chosen V2 ChangeSet through WorkbookService."""
    import json
    from excelmanus.workbook.merge import review_merge
    from excelmanus.workbook_commit import CommitError
    from excelmanus.workspace.file_service import WorkspaceFileService

    ws_root, scope_error = _file_workspace_root(raw_request, request.session_id, request.workspace_id)
    if scope_error is not None:
        return scope_error
    resolved = _resolve_excel_path(request.path, request.session_id, workspace_root=ws_root)
    if resolved is None:
        return _error_json_response(404, "文件不存在或路径非法")
    operations = [op for batch in request.batches for op in batch.operations]
    if len(json.dumps(request.model_dump(), ensure_ascii=False)) > 1_000_000:
        return _error_json_response(400, "草稿过大，请分批处理")

    def run_merge():
        dest = Path(resolved).resolve()
        rel = dest.relative_to(Path(ws_root).resolve()).as_posix()
        history = WorkspaceFileService(ws_root)
        if dest.suffix.lower() not in {".xlsx", ".xlsm"}:
            return {"status": "replan", "reason": "此文件格式需要由 Agent 核对并重新制定方案。"}, 200
        versions = {batch.expected_version for batch in request.batches}
        if len(versions) != 1:
            return {"status": "replan", "reason": "草稿跨越多个版本，需要由 Agent 分批核对修改。"}, 200
        version = next(iter(versions))
        snap = _open_route_snapshot(resolved, request.path, ws_root, request.workspace_id)
        current = snap.read_bytes()
        current_version = snap.content_version
        base = current if version == current_version else history.store.read_blob(rel, version.removeprefix("sha256:"))
        if base is None:
            return {"status": "replan", "reason": "原始版本已不在历史中，无法可靠地自动合并。草稿已保留，可交给 Agent 核对。", "content_version": current_version}, 200
        if request.apply and request.expected_version != current_version:
            return {"error": "文件再次变化，请重新核对冲突", "code": "VERSION_CONFLICT", "content_version": current_version}, 409
        options = {"keep_vba": dest.suffix.lower() == ".xlsm"}
        preview, _ = review_merge(base, current, operations, **options)
        preview["content_version"] = current_version
        if not request.apply or preview["status"] != "review":
            return preview, 200

        # The reviewer returns a filtered canonical ChangeSet.  The service
        # rechecks the same current version and performs the actual write.
        review, _ = review_merge(base, current, operations, choices=request.choices, apply=True, **options)
        merged_operations = review.get("operations") or []
        if not merged_operations:
            return {"status": "merged", "content_version": current_version,
                    "operation_id": request.operation_id, "observation": {"status": "no_op"}}, 200
        from excelmanus.tools.context import ToolCallContext, bind_call, reset_call, use_workspace
        from excelmanus.workbook.user_edits import summarize_operations
        from excelmanus.workbook.service import WorkbookService
        with use_workspace(ws_root, workspace_id=request.workspace_id) as ctx:
            token = bind_call(ToolCallContext(binding=ctx.binding, call_id=request.operation_id or uuid.uuid4().hex,
                                              tool_name="apply_spreadsheet_changes"))
            try:
                applied = WorkbookService().apply(request.path, operations=merged_operations,
                    expected_version=current_version,
                    event_context={"source": "user", "session_id": request.session_id,
                                   "summary": "用户已核对并合并并发草稿；" + summarize_operations(merged_operations),
                                   "merge_base": version, "choices": request.choices})
            finally:
                reset_call(token)
        if not applied.get("committed"):
            return {"error": "合并未提交", "code": applied.get("code", "COMMIT_FAILED")}, 409
        return {"status": "merged", "content_version": applied.get("content_version"),
                "operation_id": applied.get("receipt", {}).get("operation_id", request.operation_id),
                "observation": applied.get("observation", {})}, 200

    try:
        body, status = await run_in_threadpool(run_merge)
        if body.get("status") == "merged":
            manager = get_session_manager()
            if manager is not None:
                try:
                    manager.drain_workspace_events()
                except Exception:
                    logger.warning("合并已保存，改动通知等待重试", exc_info=True)
        return JSONResponse(content=body, status_code=status)
    except (ValueError, KeyError, CommitError) as exc:
        return _error_json_response(409 if getattr(exc, "code", "") == "VERSION_CONFLICT" else 400, str(exc))


@router.get("/api/v1/files/excel/list")
async def list_excel_files(request: Request) -> JSONResponse:
    return await run_in_threadpool(_list_excel_files, request)


def _list_excel_files(request: Request) -> JSONResponse:
    """List workbook identities using the same fast scanner as the file tree."""
    assert get_config() is not None, "服务未初始化"
    from excelmanus.workspace.listing import scan_workspace

    files, _ = scan_workspace(_resolve_workspace_root(request), limit=None, excel_only=True)
    results = [{"path": f"./{f['path']}", "filename": f["filename"],
                "modified_at": f["modified_at"]} for f in files]
    results.sort(key=lambda item: item["modified_at"], reverse=True)
    return JSONResponse(content={"files": results})


_MAX_WORKSPACE_FILES = 2000


@router.get("/api/v1/files/workspace/list")
async def list_workspace_files(request: Request) -> JSONResponse:
    return await run_in_threadpool(_list_workspace_files, request)


def _list_workspace_files(request: Request) -> JSONResponse:
    assert get_config() is not None, "服务未初始化"
    from excelmanus.workspace.listing import scan_workspace

    ws_root = str(_resolve_workspace(request).root_dir)
    files, truncated = scan_workspace(ws_root, limit=_MAX_WORKSPACE_FILES)
    files.sort(key=lambda item: (not item["is_dir"], item["path"].lower()))
    return JSONResponse(content={
        "files": files, "truncated": truncated, "workspace_path": ws_root,
    })


@router.get("/api/v1/files/registry")
async def get_file_registry(request: Request) -> JSONResponse:
    """返回 FileRegistry 全量文件列表 + 可选事件历史。

    Query params:
      - include_deleted: bool (default false) — 是否包含已软删除的文件
      - include_events: bool (default false) — 是否附带每个文件的事件历史
      - file_id: str (optional) — 仅返回指定文件及其事件/谱系
    """
    assert get_config() is not None, "服务未初始化"

    ws_root = _resolve_workspace_root(request)
    registry = _get_file_registry(ws_root)
    if registry is None:
        return JSONResponse(content={"files": [], "total": 0})

    include_deleted = request.query_params.get("include_deleted", "").lower() in ("1", "true")
    include_events = request.query_params.get("include_events", "").lower() in ("1", "true")
    file_id = request.query_params.get("file_id", "").strip()

    def _event_to_dict(evt: Any) -> dict[str, Any]:
        return {
            "id": evt.id,
            "file_id": evt.file_id,
            "event_type": evt.event_type,
            "session_id": evt.session_id,
            "turn": evt.turn,
            "tool_name": evt.tool_name,
            "details": evt.details,
            "created_at": evt.created_at,
        }

    # 单文件查询模式
    if file_id:
        entry = registry.get_by_id(file_id)
        if entry is None:
            # 退而求其次：按路径查
            entry = registry.get_by_path(file_id)
        if entry is None:
            return _error_json_response(404, f"文件未找到: {file_id}")
        file_dict = entry.to_dict()
        if include_events:
            file_dict["events"] = [_event_to_dict(e) for e in registry.get_events(entry.id)]
        children = registry.get_children(entry.id)
        file_dict["children"] = [c.to_dict() for c in children]
        lineage = registry.get_lineage(entry.id)
        file_dict["lineage"] = [a.to_dict() for a in lineage]
        return JSONResponse(content={"file": file_dict})

    # 全量列表模式
    entries = registry.list_all(include_deleted=include_deleted)
    files: list[dict[str, Any]] = []
    for entry in entries:
        d = entry.to_dict()
        if include_events:
            d["events"] = [_event_to_dict(e) for e in registry.get_events(entry.id)]
        files.append(d)

    return JSONResponse(content={"files": files, "total": len(files)})

# ── File Groups API ──────────────────────────────────────────

@router.get("/api/v1/files/groups")
async def list_file_groups(request: Request) -> JSONResponse:
    """列出当前工作区的所有文件组（附成员摘要）。"""
    assert get_config() is not None, "服务未初始化"
    ws_root = _resolve_workspace_root(request)
    registry = _get_file_registry(ws_root)
    if registry is None:
        return JSONResponse(content={"groups": []})

    try:
        groups = registry.list_groups()
        result = []
        for g in groups:
            members = registry.get_group_files(g.id)
            result.append({
                **g.to_dict(),
                "members": members,
            })
        return JSONResponse(content={"groups": result})
    except Exception:
        return JSONResponse(content={"groups": []})

@router.put("/api/v1/files/groups/{group_id}")
async def update_file_group(group_id: str, request: Request) -> JSONResponse:
    """更新文件组名称/描述。

    Body: {name?: str, description?: str}
    """
    assert get_config() is not None, "服务未初始化"
    ws_root = _resolve_workspace_root(request)
    registry = _get_file_registry(ws_root)
    if registry is None:
        return _error_json_response(500, "FileRegistry 不可用")

    body = await request.json()
    name = body.get("name")
    description = body.get("description")

    group = registry.update_group(group_id, name=name, description=description)
    if group is None:
        return _error_json_response(404, f"文件组未找到: {group_id}")

    members = registry.get_group_files(group.id)
    return JSONResponse(content={**group.to_dict(), "members": members})

@router.delete("/api/v1/files/groups/{group_id}")
async def delete_file_group(group_id: str, request: Request) -> JSONResponse:
    """删除文件组。"""
    assert get_config() is not None, "服务未初始化"
    ws_root = _resolve_workspace_root(request)
    registry = _get_file_registry(ws_root)
    if registry is None:
        return _error_json_response(500, "FileRegistry 不可用")

    ok = registry.delete_group(group_id)
    if not ok:
        return _error_json_response(404, f"文件组未找到: {group_id}")
    return JSONResponse(content={"status": "deleted", "group_id": group_id})

@router.put("/api/v1/files/groups/{group_id}/members")
async def update_file_group_members(group_id: str, request: Request) -> JSONResponse:
    """管理文件组成员。

    Body: {add?: [{file_id: str, role?: str}], remove?: [str]}
    """
    assert get_config() is not None, "服务未初始化"
    ws_root = _resolve_workspace_root(request)
    registry = _get_file_registry(ws_root)
    if registry is None:
        return _error_json_response(500, "FileRegistry 不可用")

    group = registry.get_group(group_id)
    if group is None:
        return _error_json_response(404, f"文件组未找到: {group_id}")

    body = await request.json()
    to_add = body.get("add", [])
    to_remove = body.get("remove", [])

    for item in to_add:
        if isinstance(item, dict):
            fid = item.get("file_id", "")
            role = item.get("role", "member")
        else:
            fid = str(item)
            role = "member"
        if fid:
            registry.add_to_group(group_id, fid, role)

    for fid in to_remove:
        if fid:
            registry.remove_from_group(group_id, str(fid))

    members = registry.get_group_files(group_id)
    return JSONResponse(content={**group.to_dict(), "members": members})

@router.get("/api/v1/files/spec")
async def get_spec_file(request: Request) -> JSONResponse:
    """返回 workspace 内 JSON spec 文件内容（WorkbookSpec 等）。"""
    assert get_config() is not None, "服务未初始化"

    path = request.query_params.get("path", "")
    if not path:
        return _error_json_response(400, "缺少 path 参数")  # type: ignore[return-value]

    from excelmanus.security.guard import FileAccessGuard, SecurityViolationError

    ws_root = _resolve_workspace_root(request)
    try:
        file_path = FileAccessGuard(ws_root).resolve_and_validate(path)
    except SecurityViolationError:
        return _error_json_response(404, f"Spec 文件不存在: {path}")  # type: ignore[return-value]

    if not file_path.is_file() or file_path.suffix.lower() != ".json":
        return _error_json_response(404, f"Spec 文件不存在: {path}")  # type: ignore[return-value]

    if file_path.stat().st_size > 2 * 1024 * 1024:
        return _error_json_response(400, "Spec 文件过大")  # type: ignore[return-value]

    import json as _json

    try:
        content = file_path.read_text(encoding="utf-8")
        data = _json.loads(content)
    except Exception as exc:
        return _error_json_response(500, f"读取 spec 失败: {exc}")  # type: ignore[return-value]
    if not isinstance(data, (dict, list)):
        return _error_json_response(400, "Spec 必须是 JSON 对象或数组")  # type: ignore[return-value]

    return JSONResponse(content=data)

class AdmitAttachmentRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    data: str
    media_type: str = "image/png"
    name: str | None = None


@router.post("/api/v1/attachments")
async def admit_attachment(body: AdmitAttachmentRequest) -> JSONResponse:
    """Admit one image into the content-addressed store. Chat then sends attachment_id."""
    from excelmanus.attachments.admit import admit_image_bytes, decode_image_payload
    from excelmanus.attachments.types import AttachmentError

    if not (body.data or "").strip():
        return _error_json_response(400, "缺少图片数据")
    try:
        raw = decode_image_payload(body.data)
        ref = admit_image_bytes(raw, media_type=body.media_type, name=body.name)
    except AttachmentError as exc:
        return _error_json_response(400, str(exc))
    except Exception as exc:
        logger.warning("附件准入失败", exc_info=True)
        return _error_json_response(400, f"图片无法读取: {exc}")
    return JSONResponse(content={"attachment": ref.to_dict()})


@router.get("/api/v1/files/image")
async def get_image_file(request: Request) -> StreamingResponse:
    """返回 workspace 内图片文件的二进制流。"""
    assert get_config() is not None, "服务未初始化"

    path = request.query_params.get("path", "")
    session_id = request.query_params.get("session_id")
    workspace_id = request.query_params.get("workspace_id")
    if not path:
        return _error_json_response(400, "缺少 path 参数")  # type: ignore[return-value]

    from pathlib import Path as _Path
    import mimetypes

    ws_root, scope_error = _file_workspace_root(request, session_id, workspace_id)
    if scope_error is not None:
        return scope_error  # type: ignore[return-value]

    # 使用与下载相同的路径解析逻辑
    resolved = _resolve_excel_path(path, session_id, workspace_root=ws_root)
    if resolved is None:
        return _error_json_response(404, f"图片文件不存在: {path}")  # type: ignore[return-value]

    file_path = _Path(resolved)

    _IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg"}
    if not file_path.is_file() or file_path.suffix.lower() not in _IMAGE_EXTENSIONS:
        return _error_json_response(404, f"图片文件不存在: {path}")  # type: ignore[return-value]

    content_type = mimetypes.guess_type(file_path.name)[0] or "image/png"

    def _iter_image():
        with open(file_path, "rb") as f:
            while chunk := f.read(65536):
                yield chunk

    return StreamingResponse(
        _iter_image(),
        media_type=content_type,
        headers={
            "Content-Disposition": _make_content_disposition(_display_filename(file_path, ws_root)),
            **_NO_STORE_HEADERS,
        },
    )

@router.get("/api/v1/files/read")
async def read_text_file(request: Request) -> JSONResponse:
    """返回 workspace 内文本文件的内容（供代码/MD 预览）。"""
    assert get_config() is not None, "服务未初始化"

    path = request.query_params.get("path", "")
    session_id = request.query_params.get("session_id")
    workspace_id = request.query_params.get("workspace_id")
    if not path:
        return _error_json_response(400, "缺少 path 参数")  # type: ignore[return-value]

    from pathlib import Path as _Path

    ws_root, scope_error = _file_workspace_root(request, session_id, workspace_id)
    if scope_error is not None:
        return scope_error  # type: ignore[return-value]

    # 使用与下载相同的路径解析逻辑
    resolved = _resolve_excel_path(path, session_id, workspace_root=ws_root)
    if resolved is None:
        return _error_json_response(404, f"文本文件不存在: {path}")  # type: ignore[return-value]

    file_path = _Path(resolved)

    if not file_path.is_file() or not is_text_preview_file(file_path):
        return _error_json_response(404, f"文本文件不存在: {path}")  # type: ignore[return-value]

    # 限制文件大小（最大 1MB）
    if file_path.stat().st_size > 1024 * 1024:
        return _error_json_response(400, "文件过大，无法预览")  # type: ignore[return-value]

    try:
        content = file_path.read_text(encoding="utf-8")
        return JSONResponse(content={"content": content}, headers=dict(_NO_STORE_HEADERS))
    except UnicodeDecodeError:
        return _error_json_response(400, "文件编码不支持，请使用 UTF-8 编码")  # type: ignore[return-value]
    except Exception as exc:
        return _error_json_response(500, f"读取文件失败: {exc}")  # type: ignore[return-value]

@router.get("/api/v1/files/download")
async def download_file(request: Request) -> StreamingResponse:
    """通用文件下载：返回 workspace 内任意文件的二进制流。"""
    assert get_config() is not None, "服务未初始化"

    path = request.query_params.get("path", "")
    session_id = request.query_params.get("session_id")
    workspace_id = request.query_params.get("workspace_id")
    if not path:
        return _error_json_response(400, "缺少 path 参数")  # type: ignore[return-value]

    ws_root, scope_error = _file_workspace_root(request, session_id, workspace_id)
    if scope_error is not None:
        return scope_error  # type: ignore[return-value]
    resolved = _resolve_excel_path(path, session_id, workspace_root=ws_root)
    if resolved is None:
        return _error_json_response(404, f"文件不存在或路径非法: {path}")  # type: ignore[return-value]

    from pathlib import Path as _Path
    import mimetypes

    file_path = _Path(resolved)
    content_type = mimetypes.guess_type(file_path.name)[0] or "application/octet-stream"

    def _iter_file():
        with open(resolved, "rb") as f:  # type: ignore[arg-type]
            while chunk := f.read(65536):
                yield chunk

    return StreamingResponse(
        _iter_file(),
        media_type=content_type,
        headers={
            "Content-Disposition": _make_content_disposition(_display_filename(file_path, ws_root)),
            **_NO_STORE_HEADERS,
        },
    )


@router.get("/api/v1/files/relationships")
async def get_file_relationships(request: Request) -> JSONResponse:
    """发现工作区内 Excel 文件之间的列关联关系。

    参数:
      - directory: 扫描目录（可选，默认 "."）
    """
    assert get_config() is not None, "服务未初始化"

    directory = request.query_params.get("directory", ".")
    ws_root = _resolve_workspace_root(request)
    if not directory or directory == ".":
        directory = ws_root

    import asyncio

    try:
        from excelmanus.tools.context import use_workspace
        from excelmanus.workbook.data import discover_file_relationships as _dfr

        with use_workspace(ws_root):
            rel_result = await asyncio.to_thread(_dfr, directory=directory, max_files=5)
        rel_data = rel_result.value if isinstance(getattr(rel_result, "value", None), dict) else {}
        return JSONResponse(content=rel_data)
    except Exception as exc:
        logger.error("文件关系发现失败: %s", exc, exc_info=True)
        return _error_json_response(500, f"分析失败: {exc}")

# ── Word 文件 API ───────────────────────────────────────────────

def _resolve_supported_word_file(
    requested_path: str,
    resolved_path: str | None,
) -> tuple["Path | None", JSONResponse | None]:
    from pathlib import Path as _Path

    if resolved_path is None:
        return None, _error_json_response(404, f"Word 文件不存在: {requested_path}")

    file_path = _Path(resolved_path)
    if not file_path.is_file():
        return None, _error_json_response(404, f"Word 文件不存在: {requested_path}")

    if file_path.suffix.lower() != ".docx":
        return None, _error_json_response(404, f"Word 文件格式不支持（仅支持 .docx）: {requested_path}")

    return file_path, None

@router.get("/api/v1/files/word")
async def get_word_file(request: Request) -> StreamingResponse:
    """返回 .docx 文件二进制流（下载用）。"""
    assert get_config() is not None, "服务未初始化"

    path = request.query_params.get("path", "")
    session_id = request.query_params.get("session_id")
    workspace_id = request.query_params.get("workspace_id")
    if not path:
        return _error_json_response(400, "缺少 path 参数")  # type: ignore[return-value]

    ws_root, scope_error = _file_workspace_root(request, session_id, workspace_id)
    if scope_error is not None:
        return scope_error  # type: ignore[return-value]
    resolved = _resolve_excel_path(path, session_id, workspace_root=ws_root)
    file_path, error_response = _resolve_supported_word_file(path, resolved)
    if error_response is not None:
        return error_response  # type: ignore[return-value]

    def _iter_file():
        with open(resolved, "rb") as f:  # type: ignore[arg-type]
            while chunk := f.read(65536):
                yield chunk

    return StreamingResponse(
        _iter_file(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": _make_content_disposition(_display_filename(file_path, ws_root)),
            **_NO_STORE_HEADERS,
        },
    )

@router.get("/api/v1/files/word/snapshot")
async def get_word_snapshot(request: Request) -> JSONResponse:
    """返回 Word 文档的 JSON 快照（段落+样式+表格），供前端只读快照预览。"""
    assert get_config() is not None, "服务未初始化"

    path = request.query_params.get("path", "")
    session_id = request.query_params.get("session_id")
    workspace_id = request.query_params.get("workspace_id")
    max_paragraphs = int(request.query_params.get("max_paragraphs", "500"))
    if not path:
        return _error_json_response(400, "缺少 path 参数")

    ws_root, scope_error = _file_workspace_root(request, session_id, workspace_id)
    if scope_error is not None:
        return scope_error
    resolved = _resolve_excel_path(path, session_id, workspace_root=ws_root)
    file_path, error_response = _resolve_supported_word_file(path, resolved)
    if error_response is not None:
        return error_response

    try:
        import asyncio
        snap = _open_route_snapshot(resolved, path, ws_root, workspace_id)
        snapshot = await asyncio.to_thread(_build_word_snapshot, str(file_path), max_paragraphs)
        from excelmanus.workspace.identity import display_name_for
        snapshot["file"] = display_name_for(snap.file.relative)
        snapshot["content_version"] = snap.content_version
        return JSONResponse(content=snapshot)
    except Exception as exc:
        logger.error("Word snapshot 生成失败: %s", exc, exc_info=True)
        return _error_json_response(500, f"读取 Word 文件失败: {exc}")

def _build_word_snapshot(file_path: str, max_paragraphs: int = 500) -> dict[str, Any]:
    """构建 Word 文档 JSON 快照。

    返回结构化数据，包含段落列表（含文本、样式、行内格式）和表格，
    供前端只读快照渲染（正文 + 表格）。
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(file_path)

    alignment_map = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }

    paragraphs: list[dict[str, Any]] = []
    for i, para in enumerate(doc.paragraphs):
        if i >= max_paragraphs:
            break
        style_name = para.style.name if para.style else "Normal"

        entry: dict[str, Any] = {
            "text": para.text,
            "style": style_name,
        }

        if style_name.startswith("Heading"):
            try:
                entry["heading_level"] = int(style_name.split()[-1])
            except (ValueError, IndexError):
                pass
        elif style_name == "Title":
            entry["heading_level"] = 0

        if para.alignment is not None:
            entry["alignment"] = alignment_map.get(para.alignment, "left")

        runs: list[dict[str, Any]] = []
        for run in para.runs:
            rd: dict[str, Any] = {"text": run.text}
            if run.bold:
                rd["bold"] = True
            if run.italic:
                rd["italic"] = True
            if run.underline:
                rd["underline"] = True
            if run.font.size:
                rd["size_pt"] = round(run.font.size.pt, 1)
            if run.font.name:
                rd["font_name"] = run.font.name
            if run.font.color and run.font.color.rgb:
                rd["color"] = str(run.font.color.rgb)
            runs.append(rd)

        if runs:
            entry["runs"] = runs

        paragraphs.append(entry)

    tables: list[dict[str, Any]] = []
    for idx, tbl in enumerate(doc.tables):
        rows_data: list[list[str]] = []
        for row in tbl.rows:
            rows_data.append([cell.text.strip() for cell in row.cells])
        tables.append({
            "index": idx,
            "rows": len(tbl.rows),
            "columns": len(tbl.columns),
            "data": rows_data,
        })

    properties: dict[str, Any] = {}
    core = doc.core_properties
    if core.title:
        properties["title"] = core.title
    if core.author:
        properties["author"] = core.author

    return {
        "total_paragraphs": len(doc.paragraphs),
        "returned_paragraphs": len(paragraphs),
        "truncated": len(doc.paragraphs) > max_paragraphs,
        "paragraphs": paragraphs,
        "tables": tables,
        "total_tables": len(doc.tables),
        "sections": len(doc.sections),
        "properties": properties,
    }


class WordWriteRequest(BaseModel):
    """Word 文档写入请求。"""

    model_config = ConfigDict(extra="forbid")
    session_id: str | None = None
    workspace_id: str | None = None
    path: str
    operations: list[dict[str, Any]]
    expected_version: str


@router.post("/api/v1/files/word/write")
async def write_word_content(request: WordWriteRequest, raw_request: Request) -> JSONResponse:
    """前端编辑回写：将内容变更写入 Word 文件。"""
    assert get_config() is not None, "服务未初始化"

    if not request.expected_version.strip():
        return _error_json_response(400, "write 必须提供 expected_version")

    ws_root = _resolve_workspace_root(
        raw_request,
        session_id=request.session_id,
        workspace_id=request.workspace_id,
    )
    resolved = _resolve_excel_path(request.path, request.session_id, workspace_root=ws_root)
    file_path, error_response = _resolve_supported_word_file(request.path, resolved)
    if error_response is not None:
        return error_response

    try:
        import asyncio
        result = await asyncio.to_thread(
            _apply_word_write,
            str(file_path),
            request.operations,
            request.expected_version,
            str(ws_root),
        )
        status = 200
        if isinstance(result, dict) and result.get("code") == "VERSION_CONFLICT":
            status = 409
        elif isinstance(result, dict) and result.get("status") == "error" and result.get("code"):
            status = 400
        return JSONResponse(status_code=status, content=result)
    except Exception as exc:
        logger.error("Word write 失败: %s", exc, exc_info=True)
        return _error_json_response(500, f"写入失败: {exc}")


def _apply_word_write(
    file_path: str,
    operations: list[dict[str, Any]],
    expected_version: str,
    workspace_root: str,
) -> dict[str, Any]:
    """执行 Word 文档写入：任一操作失败则不落盘。"""
    from docx import Document

    from excelmanus.security.guard import FileAccessGuard
    from excelmanus.tools.word_tools import apply_word_operations, _docx_bytes
    from excelmanus.workbook_commit import CommitError, commit_bytes

    doc = Document(file_path)
    applied, errors = apply_word_operations(doc, operations)
    if errors:
        return {
            "status": "error",
            "applied": applied,
            "applied_count": len(applied),
            "errors": errors,
        }

    guard = FileAccessGuard(workspace_root)
    from pathlib import Path as _Path
    dest = _Path(file_path).resolve()
    rel = str(dest.relative_to(guard.workspace_root)).replace("\\", "/")
    try:
        cr = commit_bytes(
            guard=guard,
            file_path=rel,
            data=_docx_bytes(doc),
            expected_version=expected_version,
        )
    except CommitError as exc:
        return {
            "status": "error",
            "code": exc.code,
            "error": exc.message,
            "message": exc.message,
            "fields": exc.fields,
        }
    return {
        "status": "success",
        "applied": applied,
        "applied_count": len(applied),
        "content_version": cr.content_version,
        "path": cr.path,
    }

# ── 文件管理 API ─────────────────────────────────────

def _workspace_file_service(
    request: Request,
    session_id: str | None,
    workspace_id: str | None = None,
):
    from excelmanus.workspace.file_service import WorkspaceFileService

    ws = _resolve_workspace(request, session_id=session_id, workspace_id=workspace_id)
    return ws, WorkspaceFileService(ws.root_dir)


def _commit_http_error(exc: Exception) -> JSONResponse:
    from excelmanus.workbook_commit import CommitError

    if not isinstance(exc, CommitError):
        return _error_json_response(400, str(exc))
    if exc.code in {"NOT_FOUND"}:
        return _error_json_response(404, exc.message)
    if exc.code in {"VERSION_CONFLICT", "FILE_EXISTS", "PATH_OCCUPIED"}:
        return _error_json_response(409, exc.message)
    return _error_json_response(400, exc.message)


@router.post("/api/v1/files/workspace/mkdir")
async def workspace_mkdir(request: Request) -> JSONResponse:
    """在工作区相对路径创建目录。"""
    assert get_config() is not None, "服务未初始化"
    body = await request.json()
    path = body.get("path", "").strip()
    if not path:
        return _error_json_response(400, "缺少 path 参数")

    _ws, svc = _workspace_file_service(request, body.get("session_id") or None, body.get("workspace_id") or None)
    try:
        rel = resolve_canonical(svc.root, path).relative
    except IdentityError as exc:
        return _error_json_response(400, str(exc))
    dest = svc.root / rel
    if dest.exists():
        return _error_json_response(409, "目录已存在" if dest.is_dir() else "目标已存在")
    try:
        svc.raise_if_failed(svc.mkdir(path))
    except Exception as exc:
        return _commit_http_error(exc)
    return JSONResponse(content={"status": "created", "path": f"./{rel}"})

@router.post("/api/v1/files/workspace/create")
async def workspace_create_file(request: Request) -> JSONResponse:
    """在工作区相对路径创建空文件。"""
    assert get_config() is not None, "服务未初始化"
    body = await request.json()
    path = body.get("path", "").strip()
    if not path:
        return _error_json_response(400, "缺少 path 参数")

    _ws, svc = _workspace_file_service(request, body.get("session_id") or None, body.get("workspace_id") or None)
    try:
        rel = resolve_canonical(svc.root, path).relative
    except IdentityError as exc:
        return _error_json_response(400, str(exc))
    try:
        svc.raise_if_failed(svc.create(path, b""))
    except Exception as exc:
        return _commit_http_error(exc)
    return JSONResponse(content={"status": "created", "path": f"./{rel}"})

@router.delete("/api/v1/files/workspace/item")
async def workspace_delete_item(request: Request) -> JSONResponse:
    """删除工作区相对路径上的文件或文件夹。"""
    assert get_config() is not None, "服务未初始化"
    body = await request.json()
    path = body.get("path", "").strip()
    if not path:
        return _error_json_response(400, "缺少 path 参数")

    _ws, svc = _workspace_file_service(request, body.get("session_id") or None, body.get("workspace_id") or None)
    try:
        rel = resolve_canonical(svc.root, path).relative
    except IdentityError as exc:
        return _error_json_response(400, str(exc))
    dest = svc.root / rel
    try:
        if dest.is_dir():
            receipt = svc.delete_tree(path, observe_live=True)
        else:
            receipt = svc.delete(path, expected_version=None, observe_live=True)
        svc.raise_if_failed(receipt)
    except Exception as exc:
        return _commit_http_error(exc)
    manager = get_session_manager()
    if manager is not None:
        manager.notify_mutation(receipt.to_dict())
    return JSONResponse(content={"status": "deleted", "path": f"./{rel}"})

@router.post("/api/v1/files/workspace/rename")
async def workspace_rename_item(request: Request) -> JSONResponse:
    """重命名工作区相对路径上的文件或文件夹。"""
    assert get_config() is not None, "服务未初始化"
    body = await request.json()
    old_path = body.get("old_path", "").strip()
    new_path = body.get("new_path", "").strip()
    if not old_path or not new_path:
        return _error_json_response(400, "缺少路径参数")
    _ws, svc = _workspace_file_service(request, body.get("session_id") or None, body.get("workspace_id") or None)
    try:
        old_rel = resolve_canonical(svc.root, old_path).relative
        new_rel = resolve_canonical(svc.root, new_path).relative
    except IdentityError as exc:
        return _error_json_response(400, str(exc))
    src = svc.root / old_rel
    try:
        if src.is_dir():
            receipt = svc.move_tree(old_path, new_path, observe_live=True)
        else:
            receipt = svc.move(old_path, new_path, expected_version=None, observe_live=True)
        svc.raise_if_failed(receipt)
    except Exception as exc:
        return _commit_http_error(exc)
    manager = get_session_manager()
    if manager is not None:
        manager.notify_mutation(receipt.to_dict())
    return JSONResponse(content={"status": "renamed", "old_path": f"./{old_rel}", "new_path": f"./{new_rel}"})

@router.post("/api/v1/files/reveal")
async def reveal_file(request: Request) -> JSONResponse:
    """在本地文件管理器中打开文件所在目录。"""
    # 服务器/Docker 模式下此功能无意义（打开的是服务器端文件管理器，用户不可见）
    if get_config() is not None and get_config().is_server:
        return _error_json_response(
            400,
            "此功能仅在本地部署模式下可用。服务器模式下无法打开本地文件管理器。",
        )

    import platform
    import subprocess

    body = await request.json()
    file_path = body.get("path", "").strip()
    session_id = body.get("session_id") or None
    workspace_id = body.get("workspace_id") or None
    if not file_path:
        return _error_json_response(400, "缺少 path 参数")

    ws_root, scope_error = _file_workspace_root(request, session_id, workspace_id)
    if scope_error is not None:
        return scope_error
    from excelmanus.security.guard import FileAccessGuard, SecurityViolationError

    try:
        target_path = FileAccessGuard(ws_root).resolve_and_validate(file_path)
    except SecurityViolationError:
        return _error_json_response(403, "路径不在工作区范围内")
    target = str(target_path)
    if not os.path.exists(target):
        return _error_json_response(404, f"路径不存在: {target}")

    try:
        system = platform.system()
        if system == "Darwin":
            subprocess.Popen(["open", "-R", target])
        elif system == "Windows":
            subprocess.Popen(["explorer", "/select,", target])
        else:
            parent = os.path.dirname(target) if os.path.isfile(target) else target
            subprocess.Popen(["xdg-open", parent])
    except Exception as exc:
        logger.warning("打开文件管理器失败: %s", exc)
        return _error_json_response(500, f"打开失败: {exc}")

    return JSONResponse(content={"status": "ok", "path": target})
